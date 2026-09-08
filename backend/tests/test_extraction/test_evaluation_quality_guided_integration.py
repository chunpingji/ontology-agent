"""Quality runner integration through production extraction and compact citations."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.quality_guided_variant import build_quality_guided_variant
from app.schemas.evidence import TaskBudget
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:Report": {
        "iri": "urn:Report",
        "label": "Report",
        "parents": [],
        "properties": [],
        "relationships": [
            {
                "iri": "urn:describes",
                "label": "describes",
                "description": "The product described by this report",
                "range": ["urn:Product"],
            }
        ],
    },
    "urn:Product": {
        "iri": "urn:Product",
        "label": "Product",
        "parents": [],
        "properties": [
            {"iri": "urn:dose", "label": "Dose", "datatype": "decimal", "canonical_unit": "mg"}
        ],
        "relationships": [],
    },
    "urn:Unreachable": {
        "iri": "urn:Unreachable",
        "label": "Unreachable",
        "parents": [],
        "properties": [],
        "relationships": [],
    },
}


@pytest.fixture
def source(tmp_path):
    document = Document()
    document.add_heading("Product report", 1)
    document.add_paragraph("This report describes Product A.")
    document.add_heading("Planned doses", 1)
    document.add_paragraph(
        "The product column identifies the products whose doses are planned here."
    )
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [["Product", "Dose (mg)"], ["Product A", "5 mg"], ["Product B", "9 mg"]],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    path = tmp_path / "quality.docx"
    document.save(path)
    return analyze_word_core(path).ir


def fixture_model(calls, *, reject_root=False, invalid_citation=False, id_only_assertions=False):
    def model(system, user, schema, budget):
        request = json.loads(user)
        if "local_menu" in request:
            calls.append(("route_records", request))
            routes = []
            for record_id, record in request["records"].items():
                text = " ".join(
                    record.get("source", [])
                    + [text for cell in record.get("cells", []) for text in cell["fragments"]]
                )
                selected = [
                    key
                    for key, predicate in request["local_menu"].items()
                    if predicate["iri"] == "urn:describes"
                    and "This report describes" in text
                    or predicate["iri"] == "urn:dose"
                    and record["kind"] == "table_row"
                ]
                routes.append({"record_id": record_id, "predicate_ids": selected})
            return {"records": routes}
        context = json.loads(request["context"])
        task, stage = context["task"], request["stage"]
        calls.append((stage, request))
        targets = [fragment for fragment in context["fragments"] if fragment["fact_eligible"]]

        def fragment(text, *, exact=True):
            return next(
                f
                for f in context["fragments"]
                if (f["text"] == text if exact else text in f["text"])
            )

        def quote(item, text=None):
            if id_only_assertions and text is None:
                return {"evidence_id": item["anchor"]["evidence_id"]}
            return {
                "evidence_id": item["anchor"]["evidence_id"],
                "text": item["text"] if text is None else text,
            }

        if stage == "verify_entity_types":
            assert "discovery_relation" not in task["predicate_definition"]
            return {
                "decisions": [
                    {
                        "candidate_id": c["candidate_id"],
                        "supported": True,
                        "identity_supported": False,
                        "reason": "Explicit named product",
                    }
                    for c in request["candidate"]["proposed_entities"]
                ]
            }
        if stage == "recall" and task["task_kind"] == "entity":
            classes = task["predicate_definition"]["classes"]
            assert all(item["name"] != "urn:Unreachable" for item in classes.values())
            key = next(key for key, value in classes.items() if value["name"] == "urn:Product")
            return {
                "entities": [
                    {
                        "class_iri": key,
                        "mention": quote(fragment("This report describes Product A."), "Product A"),
                        "supported": True,
                        "reason": "Explicit product",
                    }
                ]
            }
        if stage == "recall" and task["task_kind"] == "relationship":
            return {
                "assertions": [
                    {
                        "object_candidate_id": ref["candidate_id"],
                        "assertion_status": "affirmed",
                        "assertion_spans": [quote(targets[0])],
                    }
                    for ref in task["object_candidates"]
                ]
            }
        if stage == "recall":
            dose = next(f for f in targets if f["text"] in {"5 mg", "9 mg"})
            value = quote(dose, dose["text"][0])
            if invalid_citation:
                value["context"] = "Product A | " + dose["text"]
            return {"assertions": [{"value": value, "assertion_status": "affirmed"}]}
        if stage == "verify_binding":
            proposed = request["candidate"]
            decision = {
                "supported": not (reject_root and task["task_kind"] == "relationship"),
                "subject_candidate_id": proposed["subject"]["candidate_id"],
                "object_candidate_id": (proposed.get("object") or {}).get("candidate_id"),
                "assertion_status": proposed["assertion_status"],
                "record_mapping": {},
            }
            if task["task_kind"] == "relationship":
                return {
                    **decision,
                    "method": "explicit_assertion",
                    "assertion_spans": [quote(targets[0])],
                }
            dose = next(f for f in targets if f["text"] in {"5 mg", "9 mg"})
            row_owner = "Product A" if dose["text"] == "5 mg" else "Product B"
            # Deliberately accept the B row here: independent reference verification
            # must still reject assigning it to the current Product A subject.
            return {
                **decision,
                "method": "table_record",
                "assertion_spans": [
                    quote(fragment(row_owner)),
                    quote(fragment("Dose (mg)")),
                    quote(dose),
                ],
                "source_unit": quote(dose, "mg"),
            }
        assert stage == "verify_reference", stage
        dose = next(f for f in targets if f["text"] in {"5 mg", "9 mg"})
        return {
            "supported": dose["text"] == "5 mg",
            "subject_candidate_id": task["subject"]["candidate_id"],
            "assertion_spans": [quote(fragment("This report describes Product A.")), quote(dose)],
            "reason": "A row supports A; B row belongs to a different product",
        }

    return model


def runner(source, calls, **options):
    base = GenericExtractionRunner(
        deepcopy(SCHEMA),
        TestTokenizer(),
        fixture_model(calls, **options),
        model_identity="quality-integration-fixture",
        compact_identifiers=True,
        budget=TaskBudget(
            max_input_tokens=100000,
            max_output_tokens=4096,
            max_regions_per_task=32,
            max_tasks=100,
            max_hops=3,
        ),
    )
    return build_quality_guided_variant(base, source)


@pytest.mark.parametrize("id_only_assertions", [False, True])
def test_root_edge_then_child_property_pass_with_independent_wrong_row_rejection(
    source, id_only_assertions
):
    calls = []
    original = source.model_dump(mode="json")
    experiment = runner(source, calls, id_only_assertions=id_only_assertions)
    result = experiment.run(source, effective_class="urn:Report")
    positive = [c for c in result.candidates if c.positive_eligible]
    edges = [c for c in positive if c.kind == "relationship"]
    properties = [c for c in positive if c.kind == "property"]
    assert len(edges) == 1, result.diagnostics
    assert edges[0].predicate_iri == "urn:describes"
    assert len(properties) == 1, result.diagnostics
    assert properties[0].literal.normalized_value == "5"
    assert properties[0].literal.canonical_unit == "mg"
    assert properties[0].subject == edges[0].object
    assert edges[0].candidate_id in {ref.candidate_id for ref in properties[0].dependency_refs}
    assert "reference_not_supported" in result.diagnostics
    assert any(stage == "verify_reference" for stage, _ in calls)
    assert result.completion == "incomplete"
    assert source.model_dump(mode="json") == original
    assert all("output_contract" in request for stage, request in calls if stage != "route_records")
    assert any(item.get("reference_verification_required") for item in experiment.coverage)


def test_rejected_root_relation_does_not_expand_discovered_product(source):
    calls = []
    result = runner(source, calls, reject_root=True).run(source, effective_class="urn:Report")
    assert any(c.kind == "entity" and c.text == "Product A" for c in result.candidates)
    assert not any(c.kind == "relationship" and c.positive_eligible for c in result.candidates)
    assert not any(c.kind == "property" for c in result.candidates)
    assert not any(
        stage == "route_records" and request["subject"]["class_iri"] == "urn:Product"
        for stage, request in calls
    )


def test_cross_cell_context_is_rejected_before_property_binding(source):
    calls = []
    result = runner(source, calls, invalid_citation=True).run(source, effective_class="urn:Report")
    assert any(c.kind == "relationship" and c.positive_eligible for c in result.candidates)
    assert not any(c.kind == "property" and c.positive_eligible for c in result.candidates)
    assert not any(
        stage == "verify_binding"
        and json.loads(request["context"])["task"]["task_kind"] == "property"
        for stage, request in calls
    )
    assert result.completion == "incomplete"
