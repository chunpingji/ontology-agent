"""017 US1 — shared Word IR heading, title, table, and provenance contracts."""

from __future__ import annotations

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.extraction.docx_structure import (
    infer_heading_level,
    parse_docx_structure,
)


def _add_outline_level(paragraph, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), str(level - 1))
    ppr.append(node)


def test_toc_outline_and_bold_numbered_paragraphs_share_heading_inference(tmp_path):
    doc = Document()
    toc = doc.styles.add_style("toc 2", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("产品的基本性质", style=toc)
    doc.add_paragraph("制剂剂型：口服片剂")

    numbered = doc.add_paragraph()
    run = numbered.add_run("4.2.3产品毒性信息")
    run.bold = True
    doc.add_paragraph("是否是高致敏药物：否")

    outline = doc.add_paragraph("设备需求")
    _add_outline_level(outline, 3)
    doc.add_paragraph("设备名称：反应釜")

    emphasis = doc.add_paragraph()
    run = emphasis.add_run("本段为较长的正文强调句，并非一个章节标题，因此不应切断当前章节。")
    run.bold = True

    path = tmp_path / "stored-uuid.docx"
    doc.save(path)

    structure = parse_docx_structure(path, source_filename="HRS-1597 CMC报告.docx")

    headings = {s.heading: s.level for s in structure.sections if s.heading}
    assert headings["产品的基本性质"] == 2
    assert headings["4.2.3产品毒性信息"] == 3
    assert headings["设备需求"] == 3
    assert "本段为较长的正文强调句，并非一个章节标题，因此不应切断当前章节。" not in headings

    toxicity = structure.find_section("产品毒性信息")
    assert toxicity is not None
    assert toxicity.paras == ["是否是高致敏药物：否"]
    assert toxicity.heading_index is not None
    assert len(toxicity.para_indices) == 1


def test_tables_keep_nearest_heading_and_word_section_path(tmp_path):
    doc = Document()
    doc.add_heading("工艺描述", level=2)
    doc.add_heading("3.1.2 工艺描述", level=3)
    doc.add_heading("HRS-1597结晶纯化", level=4)
    detail = doc.add_table(rows=2, cols=2)
    detail.cell(0, 0).text = "投料"
    detail.cell(0, 1).text = "加入原料和溶剂，升温搅拌至完全溶清。"
    detail.cell(1, 0).text = "干燥"
    detail.cell(1, 1).text = "滤饼转入真空干燥箱并按规定温度干燥。"
    doc.add_heading("得量收率范围", level=3)
    summary = doc.add_table(rows=2, cols=3)
    for index, value in enumerate(("名称", "参考得量范围", "参考收率范围")):
        summary.cell(0, index).text = value
    for index, value in enumerate(("HRS-1597", "3.0~6.0kg", "50%~85%")):
        summary.cell(1, index).text = value
    path = tmp_path / "process-tables.docx"
    doc.save(path)

    structure = parse_docx_structure(path)

    assert structure.tables[0].section_heading == "HRS-1597结晶纯化"
    assert structure.tables[0].heading_index == 2
    assert structure.tables[0].section_path == [
        "工艺描述", "3.1.2 工艺描述", "HRS-1597结晶纯化",
    ]
    assert structure.tables[1].section_heading == "得量收率范围"


def test_title_falls_back_to_visual_heading_then_original_filename(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("原料药 HRS-1597 临床备样生产信息")
    run.bold = True
    run.font.size = None
    doc.add_paragraph("正文")
    path = tmp_path / "8c4f2ff8-d88a-4f84-8e79-b6472174cc27.docx"
    doc.save(path)

    structure = parse_docx_structure(path, source_filename="HRS-1597报告.docx")
    assert structure.title == "原料药 HRS-1597 临床备样生产信息"
    assert "8c4f2ff8" not in structure.title

    blank = Document()
    blank.add_paragraph("普通正文。")
    blank_path = tmp_path / "uuid-only.docx"
    blank.save(blank_path)
    fallback = parse_docx_structure(blank_path, source_filename="HRS-1597报告.docx")
    assert fallback.title == "HRS-1597报告"


def test_outline_hierarchy_and_bold_key_values_preserve_product_section(tmp_path):
    doc = Document()
    toc = doc.styles.add_style("toc 2", WD_STYLE_TYPE.PARAGRAPH)

    parent = doc.add_paragraph("产品的基本性质", style=toc)
    _add_outline_level(parent, 1)
    child = doc.add_paragraph("产品的结构", style=toc)
    _add_outline_level(child, 2)

    appearance = doc.add_paragraph()
    run = appearance.add_run("性状：类白色到白色粉末")
    run.bold = True
    sensitizing = doc.add_paragraph()
    run = sensitizing.add_run("是否是高致敏药物：否")
    run.bold = True

    sibling = doc.add_paragraph("工艺", style=toc)
    _add_outline_level(sibling, 1)
    path = tmp_path / "product-properties.docx"
    doc.save(path)

    structure = parse_docx_structure(path)

    assert structure.find_section("产品的基本性质").level == 1
    product = structure.find_section("产品的结构")
    assert product is not None
    assert product.level == 2
    assert product.paras == [
        "性状：类白色到白色粉末",
        "是否是高致敏药物：否",
    ]
    assert not any(heading.startswith("性状：") for heading in structure.headings)


def test_multirow_table_header_is_canonical_and_rows_keep_raw_indices(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=4)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "毒理标识"
    table.cell(0, 2).merge(table.cell(0, 3)).text = "安全参数"
    headers = ["活性成分(API)", "试验项目", "NOAEL", "PDE (mg/天)"]
    for idx, value in enumerate(headers):
        table.cell(1, idx).text = value
    values = ["HRS-1597", "28天重复给药", "30", "7.5"]
    for idx, value in enumerate(values):
        table.cell(2, idx).text = value

    path = tmp_path / "toxicology.docx"
    doc.save(path)
    structure = parse_docx_structure(path)
    parsed = structure.tables[0]

    assert parsed.header_row_count == 2
    assert parsed.headers[0].endswith("活性成分(API)")
    assert parsed.headers[2].endswith("NOAEL")
    assert parsed.rows[0][parsed.headers[3]] == "7.5"
    assert parsed.table_index == 0
    assert parsed.row_indices == [2]
    assert parsed.cells[2] == values


def test_heading_inference_rejects_plain_short_body():
    doc = Document()
    para = doc.add_paragraph("这是正文")
    assert infer_heading_level(para) == 0
