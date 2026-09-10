"""Regression tests for citation repair, branch scheduling and honest progress."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.schemas.evidence import TaskBudget
from app.services.extraction.checkpoint_journal import CheckpointJournal, read_checkpoint
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_extraction_tasks import accept_fixture_types
from tests.test_extraction.test_hierarchical_context import TestTokenizer


def source(tmp_path, *paragraphs):
    doc = Document()
    for value in paragraphs:
        doc.add_paragraph(value)
    path = tmp_path / "source.docx"
    doc.save(path)
    return analyze_word_core(path).ir


def worker(model, *, schema=None, repair=True, priority=False, max_tasks=64):
    return GenericExtractionRunner(
        schema or {"urn:Product": {}}, TestTokenizer(), model, model_identity="fixture",
        citation_repair=repair, relationship_priority=priority,
        budget=TaskBudget(max_input_tokens=60000, max_regions_per_task=32, max_tasks=max_tasks),
    )


def parts(user):
    request = json.loads(user)
    context = json.loads(request["context"])
    return request, context["task"], [f for f in context["fragments"] if f["purpose"] == "target"]


def quote(fragment, value=None):
    return {"evidence_id": fragment["anchor"]["evidence_id"],
            "text": value if value is not None else fragment["text"]}


def test_wrong_source_id_is_reproposed_by_model_then_independently_verified(tmp_path):
    ir = source(tmp_path, "产品：试验品甲。", "是否是青霉素：否")
    requests = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        requests.append(request)
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        product, unrelated = targets
        mention = quote(unrelated, "试验品甲")
        if "citation_feedback" in request:
            error = request["citation_feedback"]["errors"][0]
            assert error["evidence_id"] == unrelated["anchor"]["evidence_id"]
            assert error["exact_matches"] == [quote(product, "试验品甲")]
            mention = error["exact_matches"][0]
        return {"entities": [{"class_iri": "urn:Product", "mention": mention}]}

    run = worker(model).run(ir)
    assert run.completion == "complete", run.diagnostics
    assert [r["stage"] for r in requests] == ["recall", "recall", "verify_entity_types"]
    assert len(run.candidates) == 1 and run.candidates[0].positive_eligible
    assert ir.resolve(run.candidates[0].provenance[0].anchors[0]) == "试验品甲"
    assert run.checkpoint["attempt_count"] == 2 and run.checkpoint["model_calls"] == 3
    assert run.checkpoint["citation_repairs"]  # Original errors remain replayable.
    call_count = len(requests)
    restored = worker(model).run(ir, checkpoint=run.checkpoint)
    assert len(requests) == call_count and restored.candidates == run.candidates


@pytest.mark.parametrize("bad", ["改写的产品", "重复产品"])
def test_repair_does_not_accept_rewritten_or_ambiguous_source_and_is_bounded(tmp_path, bad):
    ir = source(tmp_path, "重复产品；重复产品")
    calls = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append(request)
        return {"entities": [{"class_iri": "urn:Product", "mention": quote(targets[0], bad)}]}

    first = worker(model).run(ir)
    assert len(calls) == 2 and not first.candidates and first.completion == "incomplete"
    for _ in range(2):
        first = worker(model).run(ir, checkpoint=first.checkpoint)
    assert len(calls) == 2


def test_pending_repair_survives_pause_and_journal_without_losing_valid_sibling(tmp_path):
    ir = source(tmp_path, "试验品甲；试验品乙")
    calls = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append(request)
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        names = ["试验品甲", "试验品乙" if "citation_feedback" in request else "不存在"]
        return {"entities": [{"class_iri": "urn:Product", "mention": quote(targets[0], name)}
                             for name in names]}

    path = tmp_path / "checkpoint.json"
    journal = CheckpointJournal(path)
    first = worker(model).run(ir, pause_after=1, checkpoint_fn=journal.save)
    assert first.completion == "incomplete" and len(first.candidates) == 1
    assert len(calls) == 2
    saved = read_checkpoint(path)
    assert next(iter(saved["citation_repairs"].values()))["_pending"] is True
    # A pending v1 reservation may contain binding-only suggestions. Rebuild
    # it without resetting its attempt/model-call budget on upgrade.
    next(iter(saved["citation_repairs"].values()))["version"] = "citation-feedback-v1"
    second = worker(model).run(ir, checkpoint=saved)
    assert second.completion == "complete", second.diagnostics
    assert {c.text for c in second.candidates} == {"试验品甲", "试验品乙"}
    assert second.checkpoint["attempt_count"] == 2
    assert second.checkpoint["model_calls"] == len(calls) == 4
    assert next(iter(second.checkpoint["citation_repairs"].values()))["version"] == (
        "citation-feedback-v2-domains"
    )


def test_empty_repair_is_not_success_and_total_budget_still_applies(tmp_path):
    ir = source(tmp_path, "试验品甲")
    calls = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append(request)
        return {"entities": [] if "citation_feedback" in request else [
            {"class_iri": "urn:Product", "mention": quote(targets[0], "不存在")}]}

    run = worker(model).run(ir)
    assert run.completion == "incomplete" and "citation_repair_unresolved" in run.diagnostics
    assert len(calls) == 2
    calls.clear()
    bounded = worker(model, max_tasks=1).run(ir)
    assert len(calls) == 1 and bounded.checkpoint["attempt_count"] == 1
    worker(model, max_tasks=1).run(ir, checkpoint=bounded.checkpoint)
    assert len(calls) == 1


RELATIONS = {
    "urn:Report": {"relationships": [
        {"iri": "urn:describes", "label": "描述产品", "range": ["urn:Product"]},
        {"iri": "urn:hasCleaning", "label": "含清洗方法", "range": ["urn:Cleaning"]},
    ]},
    "urn:Product": {"label": "药物产品"},
    "urn:Cleaning": {"label": "清洁过程"},
}


def test_late_cleaning_record_and_product_get_local_menus_before_full_scan(tmp_path):
    doc = Document()
    doc.add_heading("产品介绍", 1)
    doc.add_paragraph("报告描述试验品甲产品。")
    for _ in range(70):
        doc.add_paragraph("无关步骤数据。")
    doc.add_heading("设备清洗方法", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "设备", "清洁方法"
    table.cell(1, 0).text, table.cell(1, 1).text = "反应釜", "用纯化水冲洗设备五分钟。"
    path = tmp_path / "priority.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    calls, snapshots = [], []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append((request["stage"], task, targets))
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        if request["stage"] == "verify_binding":
            assert schema["$defs"]["SpanProposal"]["properties"]["context"] == {"type": "null"}
            assert "跨单元格" in request["source_quote_rules"]
            candidate = request["candidate"]
            return {"supported": task["predicate_iri"] == "urn:describes",
                    "subject_candidate_id": candidate["subject"]["candidate_id"],
                    "object_candidate_id": candidate["object"]["candidate_id"],
                    "assertion_status": "affirmed", "method": "explicit_assertion",
                    "assertion_spans": [quote(targets[-1])], "conditions": []}
        if task["task_kind"] == "entity":
            values = []
            for cls, word in [("urn:Product", "试验品甲"), ("urn:Cleaning", "纯化水冲洗")]:
                if cls in task["predicate_definition"]["classes"]:
                    values += [{"class_iri": cls, "mention": quote(f, word)} for f in targets
                               if word in f["text"]]
            return {"entities": values}
        return {"assertions": [{"object_candidate_id": task["object_candidates"][0]["candidate_id"],
                                "assertion_status": "affirmed",
                                "assertion_spans": [quote(targets[-1])]}]}

    run = worker(model, schema=RELATIONS, priority=True).run(
        ir, effective_class="urn:Report", pause_after=4, snapshot_fn=snapshots.append,
    )
    recall_tasks = [(task, targets) for stage, task, targets in calls if stage == "recall"]
    assert [t["task_kind"] for t, _ in recall_tasks[:4]] == [
        "entity", "relationship", "entity", "relationship",
    ]
    assert list(recall_tasks[0][0]["predicate_definition"]["classes"]) == ["urn:Product"]
    assert list(recall_tasks[2][0]["predicate_definition"]["classes"]) == ["urn:Cleaning"]
    assert any("纯化水冲洗" in f["text"] for f in recall_tasks[2][1])
    assert all("无关步骤数据" not in f["text"] for _, targets in recall_tasks[:4] for f in targets)
    assert run.branch_progress["urn:describes"]["status"] == "identified"
    assert run.branch_progress["urn:hasCleaning"]["status"] == "relation_failed"
    assert not any(c.positive_eligible for c in run.candidates
                   if c.predicate_iri == "urn:hasCleaning")
    assert any(s.branch_progress["urn:hasCleaning"]["status"] == "extracting_relation"
               for s in snapshots)
    prefix = len(calls)
    resumed = worker(model, schema=RELATIONS, priority=True).run(
        ir, effective_class="urn:Report", checkpoint=run.checkpoint,
    )
    assert resumed.input_id == run.input_id
    assert any("无关步骤数据" in f["text"] for _, _, fs in calls[prefix:] for f in fs)


def test_old_checkpoint_keeps_frozen_schedule_but_can_repair_quotes(tmp_path):
    ir = source(tmp_path, "产品说明")
    calls = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append(request)
        return {"entities": []}

    old = worker(model, schema=RELATIONS, repair=False).run(ir, effective_class="urn:Report")
    checkpoint = deepcopy(old.checkpoint)
    checkpoint.pop("execution_policy")
    count = len(calls)
    new = worker(model, schema=RELATIONS, priority=True).run(
        ir, effective_class="urn:Report", checkpoint=checkpoint,
    )
    assert new.input_id == old.input_id and len(calls) == count
    assert new.checkpoint["execution_policy"] is None
    fresh = worker(model, schema=RELATIONS, priority=True).run(ir, effective_class="urn:Report")
    assert fresh.input_id != old.input_id


def test_repair_keeps_valid_sibling_even_when_corrected_response_omits_it(tmp_path):
    ir = source(tmp_path, "试验品甲；试验品乙")

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        names = ["试验品乙"] if "citation_feedback" in request else ["试验品甲", "不存在"]
        return {"entities": [{"class_iri": "urn:Product", "mention": quote(targets[0], name)}
                             for name in names]}

    first = worker(model).run(ir)
    restored = worker(model).run(ir, checkpoint=first.checkpoint)
    assert first.completion == restored.completion == "complete"
    assert {c.text for c in restored.candidates} == {"试验品甲", "试验品乙"}


def test_empty_branch_progress_is_published_without_new_candidates(tmp_path):
    ir = source(tmp_path, "未涉及产品和清洗的普通记录")
    snapshots = []
    run = worker(lambda *args: {"entities": []}, schema=RELATIONS).run(
        ir, effective_class="urn:Report", snapshot_fn=snapshots.append,
    )
    assert run.completion == "complete" and len(run.candidates) == 1
    states = [s.branch_progress["urn:describes"]["status"] for s in snapshots]
    assert states[0] == "queued" and "extracting_entities" in states and states[-1] == "no_match"
    assert all(len(s.candidates) == 1 for s in snapshots)
    assert run.branch_progress["urn:describes"]["coverage_complete"] is True
    paused = worker(
        lambda *args: pytest.fail("paused task must not call a model"), schema=RELATIONS,
    ).run(ir, effective_class="urn:Report", should_pause=lambda: True)
    branch = paused.branch_progress["urn:describes"]
    assert branch["status"] == "queued" and branch["discovery_tasks"] == 0


@pytest.mark.parametrize("status", [None, "negated", "conditional", "hypothetical", "uncertain"])
def test_checked_nonpositive_relation_does_not_return_to_waiting_status(tmp_path, status):
    ir = source(tmp_path, "产品：试验品甲。条件未满足时不适用。", "其他记录。")

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        conditions = [quote(targets[0], "条件未满足时不适用")] if status == "conditional" else []
        if request["stage"] == "verify_binding":
            candidate = request["candidate"]
            return {"supported": True,
                    "subject_candidate_id": candidate["subject"]["candidate_id"],
                    "object_candidate_id": candidate["object"]["candidate_id"],
                    "assertion_status": status, "method": "explicit_assertion",
                    "assertion_spans": [quote(targets[0])], "conditions": conditions}
        if task["task_kind"] == "entity":
            return {"entities": [{"class_iri": "urn:Product",
                                  "mention": quote(targets[0], "试验品甲")}]}
        return {"assertions": [] if status is None else [{
            "object_candidate_id": task["object_candidates"][0]["candidate_id"],
            "assertion_status": status, "assertion_spans": [quote(targets[0])],
            "conditions": conditions,
        }]}

    run = worker(model, schema=RELATIONS, priority=True).run(
        ir, effective_class="urn:Report", pause_after=2,
    )
    branch = run.branch_progress["urn:describes"]
    assert branch["relationship_tasks"] == 1 and branch["failed_tasks"] == 0
    assert branch["positive_count"] == 0 and branch["coverage_complete"] is False
    assert branch["status"] == "relation_checked"
    relations = [c for c in run.candidates if c.kind == "relationship"]
    assert len(relations) == (0 if status is None else 1)
    assert all(c.validation_status == "passed" and not c.positive_eligible for c in relations)


def test_repair_does_not_override_independent_type_rejection(tmp_path):
    ir = source(tmp_path, "试验品甲", "否")

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        if request["stage"] == "verify_entity_types":
            return {"decisions": [{"candidate_id": c["candidate_id"], "supported": False,
                                   "reason": "类型证据不足"}
                                  for c in request["candidate"]["proposed_entities"]]}
        target = targets[0] if "citation_feedback" in request else targets[1]
        return {"entities": [{"class_iri": "urn:Product", "mention": quote(target, "试验品甲")}]}

    run = worker(model, schema=RELATIONS).run(ir, effective_class="urn:Report")
    assert not any(c.positive_eligible for c in run.candidates if c.class_iri == "urn:Product")
    assert not any(c.kind == "relationship" for c in run.candidates)
    assert run.branch_progress["urn:describes"]["status"] == "entities_failed"


def test_binding_stage_receives_repair_feedback_and_cannot_quote_another_cell_as_context(tmp_path):
    ir = source(tmp_path, "100L反应釜", "清洗：纯化水冲洗设备五分钟。")
    schema = {"urn:Report": {"relationships": [RELATIONS["urn:Report"]["relationships"][1]]},
              "urn:Cleaning": RELATIONS["urn:Cleaning"]}
    bindings = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        target = next(f for f in targets if "纯化水" in f["text"])
        if request["stage"] == "verify_binding":
            bindings.append(request)
            candidate = request["candidate"]
            span = quote(target)
            if "citation_feedback" not in request:
                span["context"] = "100L反应釜"  # Exact text, but from another source unit.
            else:
                assert request["citation_feedback"]["errors"][0]["quote"] == "100L反应釜"
            return {"supported": True, "subject_candidate_id": candidate["subject"]["candidate_id"],
                    "object_candidate_id": candidate["object"]["candidate_id"],
                    "assertion_status": "affirmed", "method": "explicit_assertion",
                    "assertion_spans": [span], "conditions": []}
        if task["task_kind"] == "entity":
            return {"entities": [{"class_iri": "urn:Cleaning", "mention": quote(target)}]}
        return {"assertions": [{
            "object_candidate_id": task["object_candidates"][0]["candidate_id"],
            "assertion_status": "affirmed", "assertion_spans": [quote(target)],
        }]}

    run = worker(model, schema=schema, priority=True).run(
        ir, effective_class="urn:Report", pause_after=3,
    )
    assert len(bindings) == 2 and "citation_feedback" in bindings[1]
    assert run.branch_progress["urn:hasCleaning"]["positive_count"] == 1
    assert not run.checkpoint["failures"]
