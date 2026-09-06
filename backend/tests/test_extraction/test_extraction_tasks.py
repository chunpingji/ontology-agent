import json

from docx import Document

from app.schemas.evidence import TaskBudget
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:NovelProduct": {
        "label": "任意新产品类",
        "properties": [
            {"iri": "urn:strength", "datatype": "decimal", "label": "规格"},
        ],
        "relationships": [{"iri": "urn:uses", "range": ["urn:Device"]}],
    },
    "urn:Device": {"label": "任意设备类", "properties": [], "relationships": []},
}


def test_class_menus_use_token_budget_and_cover_all_regions(tmp_path):
    doc = Document()
    for index in range(3):
        doc.add_paragraph(f"record {index}")
    path = tmp_path / "menus.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    schema = {f"urn:Type{i}": {"label": f"Type {i}"} for i in range(48)}
    requests = []

    def model(system, user, response_schema, budget):
        requests.append(json.loads(json.loads(user)["context"]))
        return {"entities": []}

    runner = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        model,
        model_identity="menus",
        budget=TaskBudget(max_input_tokens=60000, max_regions_per_task=2),
    )
    run = runner.run(ir)
    assert run.completion == "complete", run.diagnostics
    assert len(requests) == 2  # No fixed 16-class cross product.
    coverage = {
        (fragment["anchor"]["evidence_id"], iri)
        for context in requests
        for iri in context["task"]["predicate_definition"]["classes"]
        for fragment in context["fragments"]
        if fragment["purpose"] == "target"
    }
    assert coverage == {(u.evidence_id, iri) for u in ir.evidence_units for iri in schema}
    calls = len(requests)
    runner.run(ir, checkpoint=run.checkpoint)
    assert len(requests) == calls


def test_large_class_definitions_are_split_without_silent_omission():
    from app.services.extraction.evidence_identity import canonical_json
    from app.services.extraction.hierarchical_context import model_task

    schema = {f"urn:Type{i}": {"description": "definition " * 20} for i in range(30)}
    runner = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        None,
        model_identity="menus",
        budget=TaskBudget(max_input_tokens=6000),
    )
    menus = list(runner.pack_classes(sorted(schema)))
    assert len(menus) > 1
    assert [iri for menu in menus for iri in menu] == sorted(schema)
    assert all(len(canonical_json(model_task(runner.entity_menu(menu)))) <= 2000 for menu in menus)


def test_compact_request_does_not_accept_invented_source_offsets(tmp_path):
    doc = Document()
    doc.add_paragraph("药品 A")
    path = tmp_path / "offset.docx"
    doc.save(path)

    def wrong_offset(system, user, schema, budget):
        response = scripted_model(system, user, schema, budget)
        response["entities"][0]["mention"]["start"] = 1
        return response

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        wrong_offset,
        model_identity="bad-offset",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(analyze_word_core(path).ir)
    assert run.completion == "incomplete" and not run.candidates
    assert run.diagnostics == ["source_excerpt_mismatch"]


def test_prompt_projection_revision_invalidates_checkpoint(subject_document, monkeypatch):
    import app.services.extraction.extraction_tasks as module

    ir, _ = subject_document
    calls = []

    def model(*args):
        calls.append(args)
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Drug": {}},
        TestTokenizer(),
        model,
        model_identity="cache",
        budget=TaskBudget(max_input_tokens=20000),
    )
    first = runner.run(ir)
    assert runner.input_id(ir) == first.input_id
    count = len(calls)
    monkeypatch.setattr(module, "MODEL_CONTEXT_VERSION", "next-projection")
    second = runner.run(ir, checkpoint=first.checkpoint)
    assert first.input_id != second.input_id and len(calls) == 2 * count


def test_response_contract_is_visible_in_prompt_not_only_decoder(subject_document):
    from app.services.extraction.evidence_identity import canonical_json

    ir, _ = subject_document

    def model(system, user, response_schema, budget):
        assert canonical_json(response_schema) in system
        assert "Unicode" in system
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Drug": {}},
        TestTokenizer(),
        model,
        model_identity="contract",
        budget=TaskBudget(max_input_tokens=20000),
    )
    assert runner.run(ir).completion == "complete"


def test_bounded_pause_resumes_without_changing_full_run_identity(subject_document):
    ir, _ = subject_document
    calls = []

    def model(*args):
        calls.append(args)
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Drug": {}},
        TestTokenizer(),
        model,
        model_identity="pause",
        budget=TaskBudget(max_input_tokens=20000, max_tasks=256),
    )
    first = runner.run(ir, should_pause=lambda: len(calls) >= 2)
    assert first.completion == "incomplete" and len(calls) == 2
    second = runner.run(ir, checkpoint=first.checkpoint)
    assert second.completion == "complete" and second.input_id == first.input_id
    assert len(calls) == len(ir.evidence_units)


def test_failed_tasks_do_not_starve_resume_and_retry_is_explicit(subject_document):
    ir, _ = subject_document
    calls, fail = [], True

    def model(system, user, schema, budget):
        calls.append(user)
        if fail and len(calls) == 1:
            raise ValueError("source_excerpt_mismatch")
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Drug": {}},
        TestTokenizer(),
        model,
        model_identity="failure-resume",
        budget=TaskBudget(max_input_tokens=20000, max_tasks=10),
    )
    first = runner.run(ir, should_pause=lambda: len(calls) >= 2)
    assert len(calls) == 2 and first.completion == "incomplete"
    second = runner.run(ir, checkpoint=first.checkpoint)
    assert len(calls) == 4 and second.completion == "incomplete"
    assert second.diagnostics == ["source_excerpt_mismatch"]
    fail = False
    third = runner.run(ir, checkpoint=second.checkpoint, retry_failed=True)
    assert len(calls) == 5 and third.completion == "complete"
    assert third.checkpoint["attempt_count"] == 5


def test_pause_does_not_reset_total_task_budget(subject_document):
    ir, _ = subject_document
    calls = []

    def model(*args):
        calls.append(args)
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Drug": {}},
        TestTokenizer(),
        model,
        model_identity="total-budget",
        budget=TaskBudget(max_input_tokens=20000, max_tasks=2),
    )
    first = runner.run(ir, pause_after=1)
    second = runner.run(ir, checkpoint=first.checkpoint)
    third = runner.run(ir, checkpoint=second.checkpoint)
    assert len(calls) == 2 and third.completion == "incomplete"


def test_valid_entity_sibling_survives_bad_quote_without_hiding_task_failure(tmp_path):
    doc = Document()
    doc.add_paragraph("药品 A")
    path = tmp_path / "partial.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    calls = []

    def model(system, user, schema, budget):
        calls.append(user)
        fragment = json.loads(json.loads(user)["context"])["fragments"][0]
        return {
            "entities": [
                {
                    "class_iri": "urn:Drug",
                    "mention": {
                        "evidence_id": fragment["anchor"]["evidence_id"],
                        "text": name,
                    },
                }
                for name in ("药品 A", "不存在的药品 B")
            ]
        }

    runner = GenericExtractionRunner(
        {"urn:Drug": {}},
        TestTokenizer(),
        model,
        model_identity="partial",
        budget=TaskBudget(max_input_tokens=20000),
    )
    first = runner.run(ir)
    assert first.completion == "incomplete" and len(first.candidates) == 1
    assert first.candidates[0].text == "药品 A"
    assert first.candidates[0].validation_status == "passed"
    assert first.candidates[0].review_status == "pending"
    second = runner.run(ir, checkpoint=first.checkpoint)
    assert second.candidates == first.candidates and len(calls) == 1
    assert second.completion == "incomplete"


def test_exact_quote_alignment_is_explicit_unique_and_never_repairs_bad_offsets(tmp_path):
    import pytest

    from app.services.extraction.extraction_tasks import SpanProposal

    doc = Document()
    doc.add_paragraph("𠀀 药品 A；药品 A；设备 E")
    path = tmp_path / "quote.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    unit = ir.evidence_units[0]
    region = ir.anchor(unit.evidence_id, 0, len(unit.text))
    quote = SpanProposal(evidence_id=unit.evidence_id, text="设备 E")
    anchor = GenericExtractionRunner._anchor(quote, ir, [region])
    assert ir.resolve(anchor) == "设备 E"
    assert anchor.span_start == unit.text.index("设备 E")
    repeated = SpanProposal(evidence_id=unit.evidence_id, text="药品 A")
    with pytest.raises(ValueError, match="ambiguous_source_quote"):
        GenericExtractionRunner._anchor(repeated, ir, [region])
    # One exact occurrence inside the declared target window is unambiguous.
    narrowed = ir.anchor(unit.evidence_id, 2, 6)
    assert GenericExtractionRunner._anchor(repeated, ir, [narrowed]) == narrowed
    with pytest.raises(ValueError, match="source_excerpt_mismatch"):
        GenericExtractionRunner._anchor(
            SpanProposal(evidence_id=unit.evidence_id, text="药品A"), ir, [region]
        )
    with pytest.raises(ValueError, match="source_excerpt_mismatch"):
        GenericExtractionRunner._anchor(
            SpanProposal(evidence_id=unit.evidence_id, text="设备 E", start=0, end=4), ir, [region]
        )
    with pytest.raises(ValueError, match="source_quote_outside_scope"):
        GenericExtractionRunner._anchor(quote, ir, [narrowed])
    with pytest.raises(ValueError):
        SpanProposal(evidence_id=unit.evidence_id, text="设备 E", start=0)


def test_quote_mode_preserves_negative_binding_and_review_gate(tmp_path):
    doc = Document()
    doc.add_heading("药品 A", 1)
    doc.add_paragraph("本品不使用设备 E")
    path = tmp_path / "negative-quote.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir

    def quote_only(value):
        if isinstance(value, dict):
            return {
                key: quote_only(item) for key, item in value.items() if key not in {"start", "end"}
            }
        if isinstance(value, list):
            return [quote_only(item) for item in value]
        return value

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        lambda *args: quote_only(scripted_model(*args)),
        model_identity="quote-contract",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(ir)
    assert run.completion == "complete", run.diagnostics
    negative = next(c for c in run.candidates if c.kind == "relationship")
    assert negative.validation_status == "passed" and negative.assertion_status == "negated"
    assert not negative.positive_eligible and negative.review_status == "pending"
    assert negative.commit_status == "not_requested"
    assert ir.resolve(negative.bindings[0].anchors[0]) == "本品不使用设备 E"


def test_explicit_document_class_creates_reviewable_root_even_model_disabled(tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("生产资料", 0)
    doc.add_paragraph("文档正文")
    doc.save(source)
    ir = analyze_word_core(source).ir
    runner = GenericExtractionRunner(
        {"urn:Report": {"relationships": [], "properties": []}},
        None,
        None,
        model_identity="disabled",
    )
    run = runner.run(ir, effective_class="urn:Report")
    assert run.completion == "incomplete"
    root = run.candidates[0]
    assert root.class_iri == "urn:Report" and root.identity["document_root"] == ir.document_hash
    assert root.validation_status == "passed" and root.review_status == "pending"
    from app.services.extraction.evidence_scope import build_scope

    assert len(build_scope(ir, root).ranges) == len([u for u in ir.evidence_units if u.text])


def span(fragment, text):
    base = fragment["anchor"]["span_start"] or 0
    start = fragment["text"].index(text) + base
    return {
        "evidence_id": fragment["anchor"]["evidence_id"],
        "start": start,
        "end": start + len(text),
        "text": text,
    }


def scripted_model(system, user, schema, budget):
    """A contract fixture, not a quality oracle or a production model substitute."""
    request = json.loads(user)
    context = json.loads(request["context"])
    task = context["task"]
    target = next(f for f in context["fragments"] if f["purpose"] == "target")
    if request["stage"] == "verify_binding":
        candidate = request["candidate"]
        return {
            "supported": True,
            "subject_candidate_id": candidate["subject"]["candidate_id"],
            "object_candidate_id": (candidate.get("object") or {}).get("candidate_id"),
            "assertion_status": candidate["assertion_status"],
            "method": "explicit_assertion",
            "assertion_spans": [span(target, target["text"])],
            "conditions": [],
        }
    if task["task_kind"] == "entity":
        entities = []
        for text, iri in (
            ("药品 A", "urn:NovelProduct"),
            ("药品 B", "urn:NovelProduct"),
            ("设备 E", "urn:Device"),
        ):
            if text in target["text"]:
                entities.append({"class_iri": iri, "mention": span(target, text)})
        return {"entities": entities}
    if task["task_kind"] == "property":
        subject = context["subjects"][task["subject"]["candidate_id"]]["text"]
        value = "250 mg" if subject == "药品 A" else "500 mg"
        if value in target["text"]:
            return {"assertions": [{"value": span(target, value)}]}
        return {"assertions": []}
    if "不使用" in target["text"] and task["object_candidates"]:
        return {
            "assertions": [
                {
                    "object_candidate_id": task["object_candidates"][0]["candidate_id"],
                    "assertion_status": "negated",
                    "assertion_spans": [span(target, target["text"])],
                }
            ]
        }
    return {"assertions": []}


def test_heading_entities_precede_body_properties_and_checkpoints_do_not_repeat(tmp_path):
    doc = Document()
    doc.add_heading("药品 A", 1)
    doc.add_paragraph("规格：250 mg")
    doc.add_heading("药品 B", 1)
    doc.add_paragraph("规格：500 mg")
    path = tmp_path / "headings.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    calls = []

    def model(*args):
        calls.append(args)
        return scripted_model(*args)

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        model,
        model_identity="fixture-v1",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(ir)
    assert run.completion == "complete", run.diagnostics
    entities = {c.candidate_id: c.text for c in run.candidates if c.kind == "entity"}
    pairs = {
        (entities[c.subject.candidate_id], c.literal.normalized_value)
        for c in run.candidates
        if c.kind == "property" and c.validation_status == "passed"
    }
    assert pairs == {("药品 A", "250"), ("药品 B", "500")}
    assert all(
        c.review_status == "pending" and c.commit_status == "not_requested" for c in run.candidates
    )
    count = len(calls)
    replay = runner.run(ir, checkpoint=run.checkpoint)
    assert replay.candidates == run.candidates
    assert len(calls) == count
    runner.schema["urn:NewReport"] = {
        "relationships": [{"iri": "urn:describes", "range": ["urn:NovelProduct"]}],
        "properties": [],
    }
    runner.run(ir, effective_class="urn:NewReport", checkpoint=run.checkpoint)
    assert len(calls) > count


def test_explicit_negative_relationship_is_valid_but_not_positive(tmp_path):
    doc = Document()
    doc.add_heading("药品 A", 1)
    doc.add_paragraph("本品不使用设备 E")
    path = tmp_path / "negative.docx"
    doc.save(path)
    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        scripted_model,
        model_identity="fixture-v1",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(analyze_word_core(path).ir)
    assert run.completion == "complete", run.diagnostics
    relationships = [c for c in run.candidates if c.kind == "relationship"]
    assert len(relationships) == 1
    assert relationships[0].validation_status == "passed"
    assert relationships[0].assertion_status == "negated"
    assert not relationships[0].positive_eligible


def test_offline_and_budget_stop_without_legacy_fallback(tmp_path):
    doc = Document()
    doc.add_paragraph("药品 A")
    path = tmp_path / "offline.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    runner = GenericExtractionRunner(SCHEMA, TestTokenizer(), None, model_identity="disabled")
    run = runner.run(ir)
    assert run.completion == "incomplete" and not run.degraded and not run.candidates
    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        scripted_model,
        model_identity="m",
        budget=TaskBudget(max_input_tokens=1),
    )
    run = runner.run(ir)
    assert run.completion == "incomplete" and not run.candidates
    assert any("budget_exceeded" in reason for reason in run.diagnostics)


def test_shared_value_requires_explicit_joint_binding_and_restores(tmp_path):
    doc = Document()
    doc.add_paragraph("药品 A、药品 B 的规格为 250 mg")
    path = tmp_path / "shared.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    joint = []

    def model(system, user, schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        if request["stage"] == "verify_shared_binding":
            joint.append(request)
            return {"supported_subject_ids": [], "refusal_reason": "缺少共享证据"}
        if request["stage"] == "recall" and context["task"]["task_kind"] == "property":
            target = next(f for f in context["fragments"] if f["purpose"] == "target")
            return {"assertions": [{"value": span(target, "250 mg")}]}
        return scripted_model(system, user, schema, budget)

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        model,
        model_identity="shared-fixture",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(ir)
    properties = [c for c in run.candidates if c.kind == "property"]
    assert len(properties) == 2 and all(c.validation_status == "conflict" for c in properties)
    assert len(joint) == 1
    assert runner.run(ir, checkpoint=run.checkpoint).candidates == run.candidates
    assert len(joint) == 1


def test_positive_paths_use_real_parent_edges_and_stop_cycles(tmp_path):
    schema = {
        "urn:Node": {
            "label": "任意节点",
            "properties": [],
            "relationships": [{"iri": "urn:next", "range": ["urn:Node"]}],
        },
    }
    doc = Document()
    doc.add_paragraph("A → B → C → A")
    path = tmp_path / "path.docx"
    doc.save(path)

    def model(system, user, response_schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        target = next(f for f in context["fragments"] if f["purpose"] == "target")
        task = context["task"]
        if task["task_kind"] == "entity":
            return {
                "entities": [
                    {"class_iri": "urn:Node", "mention": span(target, name)} for name in "ABC"
                ]
            }
        if request["stage"] == "verify_binding":
            candidate = request["candidate"]
            return {
                "supported": True,
                "subject_candidate_id": candidate["subject"]["candidate_id"],
                "object_candidate_id": candidate["object"]["candidate_id"],
                "assertion_status": "affirmed",
                "method": "explicit_assertion",
                "assertion_spans": [span(target, target["text"])],
            }
        name = context["subjects"][task["subject"]["candidate_id"]]["text"]
        next_name = {"A": "B", "B": "C", "C": "A"}[name]
        target_id = next(
            ref["candidate_id"]
            for ref in task["object_candidates"]
            if context["subjects"][ref["candidate_id"]]["text"] == next_name
        )
        return {
            "assertions": [
                {
                    "object_candidate_id": target_id,
                    "assertion_spans": [span(target, target["text"])],
                }
            ]
        }

    runner = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        model,
        model_identity="path-fixture",
        budget=TaskBudget(max_input_tokens=60000, max_hops=3),
    )
    run = runner.run(analyze_word_core(path).ir)
    candidates = {c.candidate_id: c for c in run.candidates}
    paths = [
        c for c in candidates.values() if c.kind == "relationship" and len(c.relationship_path) > 1
    ]
    assert paths
    for candidate in paths:
        assert candidate.path_root and candidate.dependency_refs
        parents = [candidates[ref.candidate_id] for ref in candidate.dependency_refs]
        assert all(parent.positive_eligible for parent in parents)
        assert any(
            parent.kind == "relationship" and parent.object == candidate.subject
            for parent in parents
        )
        assert len(candidate.relationship_path) <= 3
    assert len(run.tasks) < 50


def test_root_windows_do_not_starve_other_subjects_or_relationships(tmp_path):
    doc = Document()
    doc.add_heading("报告", 0)
    for name in ("设备 A", "设备 B"):
        doc.add_heading(name, 1)
        for index in range(5):
            doc.add_paragraph(f"记录 {index}")
    path = tmp_path / "fair.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    schema = {
        "urn:Report": {
            "properties": [{"iri": f"urn:report{i}"} for i in range(10)],
            "relationships": [{"iri": "urn:uses", "range": ["urn:Device"]}],
        },
        "urn:Device": {"properties": [{"iri": "urn:id"}, {"iri": "urn:serial"}]},
    }
    calls = []

    def model(system, user, response_schema, budget):
        context = json.loads(json.loads(user)["context"])
        task = context["task"]
        target = next(f for f in context["fragments"] if f["purpose"] == "target")
        if task["task_kind"] == "entity":
            return {
                "entities": [
                    {"class_iri": "urn:Device", "mention": span(target, name)}
                    for name in ("设备 A", "设备 B")
                    if name == target["text"]
                ]
            }
        calls.append(
            (
                context["subjects"][task["subject"]["candidate_id"]]["text"],
                task["predicate_iri"],
                task["task_kind"],
            )
        )
        return {"assertions": []}

    runner = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        model,
        model_identity="fair",
        budget=TaskBudget(max_input_tokens=60000, max_tasks=len(ir.evidence_units) + 9),
    )
    first = runner.run(ir, effective_class="urn:Report", pause_after=len(ir.evidence_units) + 3)
    assert {name for name, _, _ in calls} == {"报告", "设备 A", "设备 B"}
    second = runner.run(ir, effective_class="urn:Report", checkpoint=first.checkpoint)
    assert len(calls) == 9 and second.completion == "incomplete"
    root_calls = [call for call in calls if call[0] == "报告"]
    assert [call[1] for call in root_calls] == ["urn:uses", "urn:report0", "urn:report1"]
    assert second.checkpoint["attempt_count"] == len(ir.evidence_units) + 9


def test_relationship_object_batches_cover_all_pairs_and_resume(tmp_path):
    doc = Document()
    doc.add_paragraph("报告；" + "；".join(f"对象{i}" for i in range(5)))
    path = tmp_path / "objects.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    schema = {
        "urn:Report": {"relationships": [{"iri": "urn:uses", "range": ["urn:Device"]}]},
        "urn:Device": {},
    }
    batches = []

    def model(system, user, response_schema, budget):
        context = json.loads(json.loads(user)["context"])
        task = context["task"]
        target = next(f for f in context["fragments"] if f["purpose"] == "target")
        if task["task_kind"] == "entity":
            return {
                "entities": [
                    {"class_iri": "urn:Device", "mention": span(target, f"对象{i}")}
                    for i in range(5)
                ]
            }
        batches.append({ref["candidate_id"] for ref in task["object_candidates"]})
        return {"assertions": []}

    runner = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        model,
        model_identity="objects",
        budget=TaskBudget(max_input_tokens=60000, max_objects_per_task=2),
    )
    first = runner.run(ir, effective_class="urn:Report", pause_after=2)
    assert len(batches) == 1 and len(batches[0]) == 2
    second = runner.run(ir, effective_class="urn:Report", checkpoint=first.checkpoint)
    assert second.completion == "complete", second.diagnostics
    assert sorted(map(len, batches)) == [1, 2, 2]
    assert set.union(*batches) == {
        c.candidate_id for c in second.candidates if c.class_iri == "urn:Device"
    }
    assert sum(map(len, batches)) == 5


def test_scheduler_revision_invalidates_run_identity(subject_document, monkeypatch):
    import app.services.extraction.extraction_tasks as module

    ir, _ = subject_document
    runner = GenericExtractionRunner({}, TestTokenizer(), None, model_identity="scheduler")
    before = runner.input_id(ir)
    monkeypatch.setattr(module, "SCHEDULER_VERSION", "future-scheduler")
    assert runner.input_id(ir) != before


def test_object_split_uses_exact_context_budget_and_keeps_competitors(subject_document):
    from app.schemas.evidence import EvidenceRange, ExtractionTask
    from app.services.extraction.evidence_scope import build_scope, candidate_ref
    from app.services.extraction.extraction_tasks import SYSTEM, AssertionResponse
    from app.services.extraction.hierarchical_context import build_context

    ir, (a, b) = subject_document
    objects = [
        b.model_copy(update={"candidate_id": f"object-{i}", "identity": {"record": str(i) * 1800}})
        for i in range(3)
    ]
    candidates = {c.candidate_id: c for c in [a, b, *objects]}
    task = ExtractionTask(
        task_id="large-object-context",
        task_kind="relationship",
        predicate_iri="urn:uses",
        subject=candidate_ref(a),
        competing_subjects=[candidate_ref(b)],
        object_candidates=[candidate_ref(c) for c in objects],
        scope=build_scope(ir, a, [b]),
        target_evidence_ids=[ir.evidence_units[1].evidence_id],
        target_ranges=[
            EvidenceRange(
                evidence_id=ir.evidence_units[1].evidence_id,
                start=0,
                end=len(ir.evidence_units[1].text),
            )
        ],
        budget=TaskBudget(max_input_tokens=60000),
    )
    runner = GenericExtractionRunner(
        {}, TestTokenizer(), None, model_identity="split", compact_identifiers=True
    )

    def context(value):
        return build_context(
            value,
            ir,
            candidates,
            runner.tokenizer,
            model=runner.model_identity,
            system_prompt=SYSTEM,
            response_schema=AssertionResponse.model_json_schema(),
            compact_identifiers=True,
        )

    one = task.model_copy(update={"object_candidates": task.object_candidates[:1]})
    limit = context(one).input_tokens + 200
    task = task.model_copy(update={"budget": TaskBudget(max_input_tokens=limit)})
    assert context(task).reason == "budget_exceeded"
    batches = list(runner.fit_assertion_task(task, ir, candidates))
    assert len(batches) == 3
    assert {ref.candidate_id for batch in batches for ref in batch.object_candidates} == {
        c.candidate_id for c in objects
    }
    assert all(context(batch).input_tokens <= limit for batch in batches)
    assert all(batch.competing_subjects == task.competing_subjects for batch in batches)
    assert all(batch.target_ranges == task.target_ranges for batch in batches)
    assert len({batch.task_id for batch in batches}) == 3
    assert list(runner.fit_assertion_task(task, ir, candidates)) == batches
    tiny = task.model_copy(update={"budget": TaskBudget(max_input_tokens=1)})
    minimal = list(runner.fit_assertion_task(tiny, ir, candidates))
    assert len(minimal) == 3  # Unfit singleton tasks remain explicit, not silently skipped.
    assert all(context(batch).reason == "budget_exceeded" for batch in minimal)
