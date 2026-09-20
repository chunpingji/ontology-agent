"""TOC entries belong to one section while retaining their source paragraphs."""

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.extraction.document_annotator import annotate_word
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import ParagraphBlock, parse_docx_structure
from app.services.extraction.ontology_guided.records import RecordIndex


@pytest.mark.parametrize("style_name", ["toc 1", "TOC 2", "目录 3", "Heading 2", "Normal"])
@pytest.mark.parametrize("title", ["目 录", "Table of Contents"])
def test_toc_entries_share_one_section_and_preserve_preview_and_evidence(
    tmp_path, style_name, title,
):
    doc = Document()
    if style_name not in doc.styles:
        doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph(title)
    entries = ["1 产品信息\t1", "1.1 产品结构……2", "2 工艺信息......12"]
    for index, text in enumerate(entries):
        paragraph = doc.add_paragraph(text, style=style_name)
        paragraph.runs[0].bold = True
        if index == 1:
            paragraph.paragraph_format.page_break_before = True
    # A document can begin at Heading 2; it must not become a child of the TOC.
    doc.add_heading("1 产品信息", level=2)
    doc.add_paragraph("产品正文。")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "属性"
    table.cell(1, 0).text = "片剂"
    path = tmp_path / "toc.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    assert structure.headings == [title, "1 产品信息"]
    assert structure.title == "1 产品信息"
    toc = structure.sections[0]
    assert toc.heading == title
    assert toc.paras == entries
    assert toc.para_indices == [1, 2, 3]
    toc_node, body_node = structure.section_tree.children
    assert toc_node.children == []
    assert len(toc_node.pages) == 2
    assert toc_node.layer_metadata.direct_paragraph_count == 3
    assert toc_node.source_range.end_block_id == "paragraph:3:0"
    assert body_node.heading == "1 产品信息"
    assert structure.tables[0].section_path == ["1 产品信息"]

    ir = build_document_ir(path, structure)
    units = [unit for unit in ir.evidence_units if unit.text in entries]
    assert [unit.text for unit in units] == entries
    assert all(unit.kind == "paragraph" for unit in units)
    assert all(unit.navigation_role == "toc_entry" for unit in units)
    assert all(unit.section_node_id == toc_node.node_id for unit in units)
    for unit in units:
        assert ir.resolve(ir.anchor(unit.evidence_id)) == unit.text
    assert not {unit.evidence_id for unit in units}.intersection(
        unit.evidence_id for record in RecordIndex(ir).records for unit in record.source_units
    )

    content, _, _, _ = annotate_word(
        path, engine=None, structure_only=True, structure=structure, ir=ir,
    )
    preview = {
        node.get("attrs", {}).get("sourceBlockId"): node for node in content["content"]
    }
    for index in range(1, 4):
        node = preview[f"paragraph:{index}:0"]
        assert node["type"] == "paragraph"
        assert node["attrs"]["sectionNodeId"] == toc_node.node_id


@pytest.mark.parametrize("page_suffix", ["\t1", "…… ii", "    12"])
def test_toc_without_heading_gets_one_section_without_losing_first_entry(tmp_path, page_suffix):
    doc = Document()
    doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    entries = [f"产品信息{page_suffix}", "工艺信息\t2"]
    for text in entries:
        doc.add_paragraph(text, style="TOC 1")
    # Real body headings can reuse the TOC style, without a trailing page number.
    doc.add_paragraph("产品信息", style="TOC 1")
    doc.add_paragraph("产品正文。")
    path = tmp_path / "untitled-toc.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    assert [section.heading for section in structure.sections] == ["目录", "产品信息"]
    assert structure.sections[0].paras == entries
    toc_node = structure.section_tree.children[0]
    assert toc_node.paragraph_indices == [0, 1]
    blocks = [block for block in structure.blocks if isinstance(block, ParagraphBlock)]
    assert [block.text for block in blocks[:2]] == entries
    assert [block.heading_level for block in blocks[:2]] == [0, 0]
    assert all(block.section_node_id == toc_node.node_id for block in blocks[:2])
    ir = build_document_ir(path, structure)
    assert [unit.navigation_role for unit in ir.evidence_units[:2]] == ["toc_entry"] * 2
    assert [record.text for record in RecordIndex(ir).records] == ["产品正文。"]


def test_toc_styles_without_page_numbers_are_grouped_after_toc_title(tmp_path):
    doc = Document()
    doc.styles.add_style("TOC 2", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading("目录", level=1)
    doc.add_paragraph("产品信息", style="TOC 2")
    doc.add_paragraph("")
    doc.add_paragraph("工艺信息", style="TOC 2")
    doc.add_heading("产品信息", level=1)
    doc.add_paragraph("产品正文。")
    path = tmp_path / "toc-without-pages.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    assert structure.headings == ["目录", "产品信息"]
    assert structure.sections[0].paras == ["产品信息", "工艺信息"]


def _field_char(paragraph, kind):
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    paragraph.add_run()._r.append(node)


@pytest.mark.parametrize("page_number", ["\t2", ""])
def test_toc_field_scope_ends_after_nested_page_reference(tmp_path, page_number):
    doc = Document()
    doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    first = doc.add_paragraph()
    _field_char(first, "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = ' TOC \\o "1-3" \\h '
    first.add_run()._r.append(instruction)
    _field_char(first, "separate")
    first.add_run("产品信息")
    second = doc.add_paragraph("工艺信息")
    _field_char(second, "begin")
    reference = OxmlElement("w:instrText")
    reference.text = " PAGEREF _Toc123 \\h "
    second.add_run()._r.append(reference)
    _field_char(second, "separate")
    second.add_run(page_number)
    _field_char(second, "end")
    _field_char(second, "end")
    doc.add_paragraph("产品信息", style="TOC 1")
    doc.add_paragraph("产品正文。")
    path = tmp_path / "toc-field.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    assert [section.heading for section in structure.sections] == ["目录", "产品信息"]
    assert structure.sections[0].paras == ["产品信息", f"工艺信息{page_number}"]
    assert structure.sections[1].paras == ["产品正文。"]
    ir = build_document_ir(path, structure)
    assert [unit.navigation_role for unit in ir.evidence_units[:2]] == ["toc_entry"] * 2
    assert [record.text for record in RecordIndex(ir).records] == ["产品正文。"]


def test_body_numbers_and_toc_style_reuse_do_not_create_a_toc(tmp_path):
    doc = Document()
    doc.styles.add_style("toc 2", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("产品信息", style="toc 2")
    doc.add_paragraph("批次数量\t12")
    doc.add_paragraph("样品数量\t20").runs[0].bold = True
    doc.add_heading("工艺版本  2", level=1)
    doc.add_paragraph("工艺正文。")
    path = tmp_path / "body.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    assert structure.headings == ["产品信息", "样品数量\t20", "工艺版本  2"]
    assert structure.sections[0].paras == ["批次数量\t12"]
