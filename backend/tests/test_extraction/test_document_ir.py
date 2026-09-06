import pytest
from docx import Document

from app.services.extraction.document_ir import DocumentIR, OffsetMap, build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def sample(tmp_path):
    doc = Document()
    doc.add_heading("重复标题", 1)
    doc.add_paragraph("  规格：250 mg𠀀😀  ")
    doc.add_heading("重复标题", 1)
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵向合并"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = " 内层𠀀 "
    deep = nested.cell(0, 0).add_table(rows=1, cols=1)
    deep.cell(0, 0).text = "第三层"
    path = tmp_path / "source.docx"
    doc.save(path)
    return path


def test_ir_is_stable_serializable_and_role_independent_in_structure(tmp_path):
    path = sample(tmp_path)
    structure = parse_docx_structure(path)
    first = build_document_ir(path, structure, role="analysis_source")
    second = build_document_ir(path, parse_docx_structure(path), role="template_sample")
    assert first.structure_hash == second.structure_hash
    assert first.analysis_id != second.analysis_id
    assert first.evidence_units == second.evidence_units
    assert DocumentIR.model_validate_json(first.model_dump_json()) == first
    assert len({u.section_node_id for u in first.evidence_units if u.kind == "heading"}) == 2
    para = next(u for u in first.evidence_units if "规格" in u.text)
    assert para.text == "  规格：250 mg𠀀😀  "
    start = para.text.index("𠀀")
    assert first.resolve(first.anchor(para.evidence_id, start, start + 2)) == "𠀀😀"


def test_nested_tables_keep_paths_and_vertical_merges_have_one_origin(tmp_path):
    path = sample(tmp_path)
    ir = build_document_ir(path, parse_docx_structure(path))
    inner = next(u for u in ir.evidence_units if u.text == " 内层𠀀 ")
    deep = next(u for u in ir.evidence_units if u.text == "第三层")
    assert len(inner.table_path) == 3
    assert len(deep.table_path) == 5
    assert inner.table_path != deep.table_path
    assert sum(u.text == "纵向合并" for u in ir.evidence_units) == 1
    merged = next(c for c in ir.tables[0]["source_cells"] if c["row_span"] == 2)
    assert merged["row_index"] == 1
    assert ir.tables[0]["grid"][1][0] == ir.tables[0]["grid"][2][0] == merged["cell_id"]


def test_anchor_identity_and_offset_map_never_guess_or_cross_separators(tmp_path):
    path = sample(tmp_path)
    ir = build_document_ir(path, parse_docx_structure(path))
    units = [u for u in ir.evidence_units if u.text][:2]
    mapped = OffsetMap.join(ir, [u.evidence_id for u in units], separator=" | ")
    assert ir.resolve(mapped.source_anchor(0, 2)) == units[0].text[:2]
    with pytest.raises(ValueError, match="separator|multiple"):
        mapped.source_anchor(len(units[0].text) - 1, len(units[0].text) + 4)
    bad = ir.anchor(units[0].evidence_id, 0, 1).model_copy(update={"structure_hash": "0" * 64})
    with pytest.raises(ValueError, match="identity"):
        ir.resolve(bad)


def test_parse_failure_cannot_become_valid_empty_evidence(tmp_path):
    with pytest.raises(ValueError, match="parse|unavailable"):
        path = tmp_path / "missing.docx"
        build_document_ir(path, parse_docx_structure(path))
