"""遗留 .doc → .docx 转换（app.services.extraction.doc_converter）。

错误路径用 monkeypatch 隔离 soffice（缺失 / 非零退出 / 超时 / 无产出），不依赖真实
LibreOffice；一条端到端 roundtrip 集成测试在 soffice+Writer 可用时跑真实转换，否则 skip。
"""

from __future__ import annotations

import subprocess

import pytest

from app.services.extraction import doc_converter
from app.services.extraction.doc_converter import DocConversionError, ensure_docx


def _make_real_doc(tmp_path) -> str | None:
    """用 python-docx 造 .docx，再经 soffice 转为遗留 .doc；不可用则返回 None（skip 依据）。"""
    if not doc_converter.soffice_available():
        return None
    import docx

    d = docx.Document()
    d.add_heading("评估对象", level=1)
    d.add_paragraph("本品为 XX 注射液，属注射剂。")
    src_docx = tmp_path / "seed.docx"
    d.save(str(src_docx))

    profile = tmp_path / "profile"
    try:
        proc = subprocess.run(
            [
                doc_converter._SOFFICE, "--headless", "--norestore",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to", "doc", "--outdir", str(tmp_path), str(src_docx),
            ],
            capture_output=True, timeout=120,
        )
    except Exception:
        return None
    out = tmp_path / "seed.doc"
    return str(out) if proc.returncode == 0 and out.is_file() else None


# ── 纯逻辑（无 soffice 依赖）────────────────────────────────────────────────


def test_docx_passthrough_returns_same_path():
    assert ensure_docx("/some/where/report.docx") == "/some/where/report.docx"
    assert ensure_docx("/some/where/REPORT.DOCX") == "/some/where/REPORT.DOCX"


def test_unsupported_suffix_raises():
    with pytest.raises(DocConversionError):
        ensure_docx("/some/where/note.pdf")


def test_missing_soffice_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(doc_converter, "_SOFFICE", None)
    doc = tmp_path / "x.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0")  # OLE 头，仅占位
    with pytest.raises(DocConversionError, match="不可用"):
        ensure_docx(str(doc))


def test_nonzero_exit_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(doc_converter, "_SOFFICE", "/bin/soffice-fake")

    def _fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, returncode=1, stdout=b"", stderr=b"boom")

    monkeypatch.setattr(doc_converter.subprocess, "run", _fake_run)
    doc = tmp_path / "x.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocConversionError, match="转换失败"):
        ensure_docx(str(doc))


def test_timeout_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(doc_converter, "_SOFFICE", "/bin/soffice-fake")

    def _fake_run(cmd, **kwargs):
        raise subprocess.TimeoutExpired(cmd, kwargs.get("timeout"))

    monkeypatch.setattr(doc_converter.subprocess, "run", _fake_run)
    doc = tmp_path / "x.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocConversionError, match="超时"):
        ensure_docx(str(doc))


def test_no_output_file_raises(monkeypatch, tmp_path):
    monkeypatch.setattr(doc_converter, "_SOFFICE", "/bin/soffice-fake")

    def _fake_run(cmd, **kwargs):
        return subprocess.CompletedProcess(cmd, returncode=0, stdout=b"", stderr=b"")

    monkeypatch.setattr(doc_converter.subprocess, "run", _fake_run)
    doc = tmp_path / "x.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocConversionError, match="未产出"):
        ensure_docx(str(doc))


def test_invalid_output_not_zip_raises(monkeypatch, tmp_path):
    """returncode=0 且产物存在、非空，但不是有效 zip/.docx（soffice 偶发损坏产物）。"""
    monkeypatch.setattr(doc_converter, "_SOFFICE", "/bin/soffice-fake")

    def _fake_run(cmd, **kwargs):
        # 模拟 soffice 在 outdir 写出一个「非 zip」的 <stem>.docx。
        (tmp_path / "x.docx").write_bytes(b"not a real docx, just bytes")
        return subprocess.CompletedProcess(cmd, returncode=0, stdout=b"", stderr=b"")

    monkeypatch.setattr(doc_converter.subprocess, "run", _fake_run)
    doc = tmp_path / "x.doc"
    doc.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(DocConversionError, match="zip 校验失败"):
        ensure_docx(str(doc))


# ── 端到端真实转换（需 LibreOffice + Writer）────────────────────────────────


def test_real_doc_to_docx_roundtrip(tmp_path):
    doc_path = _make_real_doc(tmp_path)
    if doc_path is None:
        pytest.skip("LibreOffice(soffice)+Writer 不可用，跳过真实 .doc 转换")

    out = ensure_docx(doc_path)
    assert out.endswith(".docx")

    import docx

    parsed = docx.Document(out)
    text = "\n".join(p.text for p in parsed.paragraphs)
    assert "注射液" in text  # 中文正文经转换保留
