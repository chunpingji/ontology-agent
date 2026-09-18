"""Failure witnesses for counterevidence authority and paid invalid responses."""
from __future__ import annotations

import json
from collections import Counter

from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.dependencies import DependencyIndex
from app.services.extraction.ontology_guided.model_adapter import LocalModelRecognitionAdapter
from tests.test_extraction.test_semantic_graph_closure import (
    DESCRIBES,
    INGREDIENT,
    _analysis,
    _effective,
    _respond,
    _run,
)


def test_counterevidence_same_paragraph_wrong_subspan_does_not_clear_conflict(
    tmp_path, monkeypatch,
):
    analysis = _analysis(tmp_path, negative=True)
    batches, rechecks = [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        result = _respond(request)
        if request["stage"] == "verification" and request["required_counterevidence"]:
            negative = next(
                item for item in request["fragments"] if "活性成分不含" in item["text"]
            )
            # The source id is correct but this quote omits the actual negation.
            for proposal in result["verifications"]:
                proposal["counterevidence_support"] = [
                    {"evidence_id": negative["evidence_id"], "text": "产品甲片"}
                ]
                proposal["counterevidence_verdict"] = "supported"
            rechecks.append(request)
        return result

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = _run(
        analysis, LocalModelRecognitionAdapter(object(), model_identity="subspan-test"),
        batch_hook=batches.append,
    )
    assert rechecks, "the test must actually reach evidence-driven revalidation"
    dependencies = DependencyIndex.from_snapshot(batches[-1].dependency_index)
    effective = _effective(result, dependencies)
    assert [edge.predicate_iri for edge in effective.edges] == [DESCRIBES]
    assert not any(edge.predicate_iri == INGREDIENT for edge in effective.edges)
    assert result.graph.progress.unresolved_claims > 0


def test_invalid_model_schema_consumes_actual_calls_and_lineage_budget(tmp_path, monkeypatch):
    analysis = _analysis(tmp_path)
    calls = Counter()

    def invalid_response(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls[request["target"]["claim_ref"]["id"]] += 1
        return {"proposals": [{"kind": "not-an-allowed-kind"}]}

    monkeypatch.setattr(model_adapter, "chat_with_schema", invalid_response)
    result = _run(
        analysis, LocalModelRecognitionAdapter(object(), model_identity="invalid-schema-test")
    )
    assert calls and sum(calls.values()) > 0
    assert result.graph.progress.model_calls == sum(calls.values())
    assert max(calls.values()) <= 2
    assert result.graph.progress.records_incomplete > 0
    assert result.graph.progress.completion == "incomplete"
    assert result.graph.edges == []


def test_successful_full_counterevidence_recheck_resumes_descendant_properties(
    tmp_path, monkeypatch,
):
    from docx import Document

    from app.services.extraction.evidence_identity import evidence_hash
    from app.services.extraction.ontology_guided.contracts import SlotSpec
    from app.services.extraction.word_analysis import analyze_word_core
    from tests.test_extraction import test_semantic_graph_closure as fixtures

    path = tmp_path / "closure.docx"
    document = Document()
    document.add_heading("产品描述、组成和熔点", 1)
    document.add_paragraph("本报告描述产品甲片。")
    document.add_paragraph("产品甲片的活性成分为成分乙。")
    document.add_paragraph("成分乙的熔点为80℃。")
    for number in range(6):
        document.add_paragraph(f"过程说明记录 {number}。")
    document.add_paragraph("更正：产品甲片的活性成分不含成分乙。")
    document.save(path)
    analysis = analyze_word_core(path)
    ontology = fixtures._ontology()
    melting_point = "urn:closure#meltingPoint"
    ontology.classes[fixtures.API].declared_properties.append(
        SlotSpec(iri=melting_point, label="熔点", declared_by=[fixtures.API])
    )
    ontology = ontology.model_copy(update={"ontology_hash": evidence_hash(ontology.classes)})
    monkeypatch.setattr(fixtures, "_ontology", lambda: ontology)
    batches, ordered = [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        result = _respond(request)
        if request["stage"] == "verification" and request["required_counterevidence"]:
            negative = next(
                item for item in request["fragments"] if "活性成分不含" in item["text"]
            )
            for proposal in result["verifications"]:
                proposal["counterevidence_support"] = [
                    {"evidence_id": negative["evidence_id"], "text": negative["text"]}
                ]
                proposal["counterevidence_verdict"] = "supported"
            ordered.append("full_counterevidence_rechecked")
        elif request["predicate"]["iri"] == melting_point:
            target = next((item for item in request["fragments"]
                           if item["fact_eligible"] and "熔点为80℃" in item["text"]), None)
            if target:
                def quote(text):
                    return {"evidence_id": target["evidence_id"], "text": text}
                if request["stage"] == "discovery":
                    result = {"proposals": [{
                        "kind": "property", "value_quote": quote("80℃"),
                        "bridge_kind": "owned_field_group", "polarity": "affirmed",
                    }]}
                else:
                    for verification in result["verifications"]:
                        verification.update(
                            subject_support=[quote("成分乙")],
                            subject_binding_verdict="supported",
                            predicate_support=[quote(target["text"])],
                            role_verdict="supported", predicate_verdict="supported",
                            applicability_verdict="supported", bridge_verdict="supported",
                            reason="受控原文属性独立核验通过。",
                        )
                    ordered.append("descendant_property")
        return result

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = _run(
        analysis, LocalModelRecognitionAdapter(object(), model_identity="recovery-test"),
        batch_hook=batches.append,
    )
    assert "full_counterevidence_rechecked" in ordered
    assert ordered.index("descendant_property") < ordered.index("full_counterevidence_rechecked")
    effective = _effective(result, DependencyIndex.from_snapshot(batches[-1].dependency_index))
    assert any(edge.predicate_iri == INGREDIENT for edge in effective.edges)
    assert [(item.predicate_iri, item.raw_value) for item in effective.properties] == [
        (melting_point, "80℃")
    ]
    assert ordered.index("full_counterevidence_rechecked") < len(ordered) - 1
    assert ordered[-1] == "descendant_property"
    assert ordered.count("descendant_property") == 2


def test_late_conditional_negation_rechecks_previously_unqualified_positive(tmp_path, monkeypatch):
    from docx import Document

    from app.services.extraction.word_analysis import analyze_word_core

    _analysis(tmp_path, negative=True)
    path = tmp_path / "closure.docx"
    document = Document(path)
    document.paragraphs[-1].text = "更正：在条件C下，产品甲片的活性成分不含成分乙。"
    document.save(path)
    batches = []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        result = _respond(request)
        for proposal in result.get("proposals", []):
            if proposal["polarity"] == "negated":
                fragment = next(item for item in request["fragments"]
                                if item["fact_eligible"] and "在条件C下" in item["text"])
                proposal["condition_support"] = [
                    {"evidence_id": fragment["evidence_id"], "text": "在条件C下"}
                ]
        return result

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = _run(
        analyze_word_core(path),
        LocalModelRecognitionAdapter(object(), model_identity="conditional-conflict-test"),
        batch_hook=batches.append,
    )
    effective = _effective(result, DependencyIndex.from_snapshot(batches[-1].dependency_index))
    assert [edge.predicate_iri for edge in effective.edges] == [DESCRIBES]
    assert result.graph.progress.unresolved_claims > 0
