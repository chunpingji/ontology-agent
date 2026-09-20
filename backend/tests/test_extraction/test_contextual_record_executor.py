"""Field admission, bounded candidate refresh and cold reuse through the real store."""

import pytest
from docx import Document

from app.models.document_analysis import DocumentRunCurrentState
from app.services.document_analysis import current_state
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.contracts import SlotSpec
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.record_discovery import ContextualDiscoveryPolicy
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction import test_record_executor as fixture
from tests.test_extraction.test_ontology_guided_core import _definition
from tests.test_extraction.test_record_executor import ROOT, A, B, record_setup

pytest_plugins = ["tests.test_extraction.test_tool_engine_resume"]


def setup_contextual(
    tmp_path, monkeypatch, current_run, *, texts, root_number=False,
    contextual_version="contextual-discovery-v2",
):
    def transform(view, _payload, *, record):
        payload = dict(entities=[], properties=[], relations=[], external_links=[],
                       observations=[], reference_bindings=[])
        if not record:
            return payload
        units = view["evidence_units"]
        sources = [unit for unit in units if unit["fact_eligible"]]

        def quote(unit, text=None):
            return {"evidence_id": unit["evidence_id"], "text": text or unit["text"],
                    "context_text": None}

        if view["stage"] == "verification":
            return {"verifications": [{
                "target_id": target["target_id"], "content_hash": target["content_hash"],
                "facets": [{"name": facet, "verdict": "supported",
                            "support": [quote(sources[0])], "counterevidence_support": [],
                            "reason": "原文明确呈现实体名称与类别。"}
                           for facet in target["required_facets"]],
            } for target in view["verification_input"]["targets"]]}
        if view.get("attribute_disambiguation"):
            payload["observations"] = [{
                "subject_id": None, "predicate_iri": None, "kind": "ambiguous",
                "quote": quote(sources[0]), "reason": "编号主体仍缺少充分归属证明。",
            }]
            return payload
        permitted = {card["class_iri"] for card in view["schema_card"]["class_cards"]}
        for class_iri, name in ((A, "装置甲"), (B, "部件乙"), (A, "装置丙")):
            source = next((unit for unit in sources if name in unit["text"]), None)
            if class_iri in permitted and source is not None:
                payload["entities"].append({
                    "local_id": name, "class_iri": class_iri, "representation": "mention",
                    "mentions": [quote(source, name)], "record_components": [],
                    "identifier_claims": [],
                })
        return payload

    args, factory, requests, hooks = record_setup(
        tmp_path, monkeypatch, current_run, split=True, empty_relations=True,
        transform=transform, texts=texts,
    )
    runner = factory()
    adapter = runner.adapter
    adapter.record_discovery = adapter.record_discovery.model_copy(update={
        "contextual": ContextualDiscoveryPolicy(version=contextual_version),
    })
    for class_iri in [A, B, *([ROOT] if root_number else [])]:
        definition = adapter.ontology.classes[class_iri]
        adapter.ontology.classes[class_iri] = _definition(
            class_iri, definition.label, properties=[SlotSpec(
                iri=class_iri + "/number", label="编号", declared_by=[class_iri],
                datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
            )], relationships=definition.declared_relationships,
        )
    adapter.ontology.ontology_hash = evidence_hash(adapter.ontology.classes)
    return args, factory, requests, hooks


def field_rows(current_run):
    store, run, _ = current_run
    rows = current_state.read_rows(store, run, DocumentRunCurrentState,
                                   prefix="work:record_discovery")
    return [row["value"] for row in rows.get("work:record_discovery", {}).values()
            if row["value"]["task"].get("purpose") == "property_disambiguation"]


@pytest.mark.parametrize("text,reason", [
    ("编号：A-001", "attribute_subject_missing"),
    ("编号：", "attribute_value_missing"),
    ("文档编号：DOC-001", "ontology_property_missing"),
])
def test_exclusive_field_has_one_work_item_before_class_cards_and_no_unneeded_call(
    tmp_path, monkeypatch, current_run, text, reason,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=[text],
    )
    result = factory().run(**args, **hooks)
    rows = field_rows(current_run)
    assert len(rows) == 1, result.events
    assert rows[0]["reason_code"] == reason, result.events
    assert rows[0]["disambiguation_attempts"] == 0
    assert all("members" in view for view in requests)
    assert result.graph.properties == []
    store, run, _ = current_run
    calls = current_state.restore_calls(store, run, run.run_fingerprint)
    assert len([p for p in calls["protocols"].values()
                if p.get("task", {}).get("purpose") == "property_disambiguation"]) == 1
    assert rows[0]["task"]["claim_lineage_id"] not in calls["lineage_calls"]
    count = len(requests)
    restored = vars(current_state.restore_work(store, run, run.run_fingerprint))
    factory().run(**args, **hooks, resume_state=restored, model_call_state=calls)
    assert len(requests) == count and len(field_rows(current_run)) == 1


def test_late_candidates_wake_the_same_field_once_and_cold_continue_does_not_repeat(
    tmp_path, monkeypatch, current_run,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run,
        texts=["编号：A-001", "装置甲。", "部件乙。", "装置丙。"], root_number=True,
    )
    result = factory().run(**args, **hooks)
    fields = [view for view in requests if view.get("attribute_disambiguation")]
    assert len(fields) == 2, {
        "diagnostics": result.diagnostics,
        "outcomes": [value for kind, value in result.events if kind == "task_outcome"],
    }
    first_field = next(index for index, view in enumerate(requests)
                       if view.get("attribute_disambiguation"))
    assert any(view["stage"] == "discovery" and not view.get("attribute_disambiguation")
               for view in requests[:first_field]), requests
    assert {node.label for node in result.graph.nodes if not node.root} >= {
        "装置甲", "部件乙", "装置丙",
    }, result.events
    assert len(fields[0]["attribute_disambiguation"]["options"]) == 2
    # Calibration runs after normal discovery and sees both later entities together.
    assert len(fields[1]["attribute_disambiguation"]["options"]) == 4
    assert next(i for i, view in enumerate(requests) if view == fields[1]) > max(
        i for i, view in enumerate(requests)
        if view.get("stage") == "discovery" and not view.get("attribute_disambiguation")
        and "members" not in view
    )
    rows = field_rows(current_run)
    assert len(rows) == 1 and rows[0]["disambiguation_attempts"] == 2
    assert rows[0]["attribute_status"] == "unresolved"
    store, run, _ = current_run
    calls = current_state.restore_calls(store, run, run.run_fingerprint)
    lineage = rows[0]["task"]["claim_lineage_id"]
    assert calls["protocols"][lineage]["assertion_generation"] == 2
    assert calls["lineage_calls"][lineage] == 2
    count = len(requests)
    restored = vars(current_state.restore_work(store, run, run.run_fingerprint))
    factory().run(**args, **hooks, resume_state=restored, model_call_state=calls)
    assert len(requests) == count
    assert len(field_rows(current_run)) == 1


def test_field_is_released_after_relevant_entity_discovery(
    tmp_path, monkeypatch, current_run,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["编号：A-001", "装置甲。"],
        root_number=True,
    )
    factory().run(**args, **hooks)
    assert not requests[0].get("attribute_disambiguation"), requests
    assert any(request.get("attribute_disambiguation") for request in requests[1:])
    assert [unit["text"] for unit in requests[0]["evidence_units"]
            if unit["fact_eligible"]] == ["装置甲。"]


def test_pause_after_paid_field_answer_reuses_it_on_cold_resume(
    tmp_path, monkeypatch, current_run,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["编号：A-001"], root_number=True,
    )
    stopped = False
    persist = hooks["model_call_hook"]

    def calls(state):
        nonlocal stopped
        persist(state)
        if (requests and requests[-1].get("attribute_disambiguation")
                and any(row["field"] == "model_turn"
                        and row["value"]["stage"] == "discovery"
                        for row in state.get("result_changes", {}).values())):
            stopped = True

    first = factory(progress_hook=lambda _: not stopped).run(
        **args, **{**hooks, "model_call_hook": calls},
    )
    assert "execution_pause_requested" in first.diagnostics
    store, run, _ = current_run
    restored = vars(current_state.restore_work(store, run, run.run_fingerprint))
    result = factory().run(
        **args, **hooks, resume_state=restored,
        model_call_state=current_state.restore_calls(store, run, run.run_fingerprint),
    )
    fields = [view for view in requests if view.get("attribute_disambiguation")]
    assert len(fields) == 1, result.events
    assert field_rows(current_run)[0]["disambiguation_attempts"] == 1


def test_zero_call_missing_owner_field_wakes_once_when_an_owner_is_registered(
    tmp_path, monkeypatch, current_run,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["编号：A-001", "装置甲。"],
    )
    result = factory().run(**args, **hooks)
    fields = [view for view in requests if view.get("attribute_disambiguation")]
    assert len(fields) == 1, result.events
    options = fields[0]["attribute_disambiguation"]["options"]
    assert [option["class_iri"] for option in options] == [A]
    rows = field_rows(current_run)
    assert len(rows) == 1 and rows[0]["disambiguation_attempts"] == 1
    store, run, _ = current_run
    calls = current_state.restore_calls(store, run, run.run_fingerprint)
    lineage = rows[0]["task"]["claim_lineage_id"]
    assert calls["protocols"][lineage]["assertion_generation"] == 1
    assert calls["lineage_calls"][lineage] == 1


def test_unrelated_entity_does_not_repeat_the_unchanged_field_candidate_set(
    tmp_path, monkeypatch, current_run,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["编号：A-001", "部件乙。"], root_number=True,
    )
    ontology = factory().adapter.ontology
    definition = ontology.classes[B]
    ontology.classes[B] = _definition(B, definition.label,
                                     relationships=definition.declared_relationships)
    ontology.ontology_hash = evidence_hash(ontology.classes)
    result = factory().run(**args, **hooks)
    assert any(node.class_iri == B for node in result.graph.nodes), result.events
    fields = [view for view in requests if view.get("attribute_disambiguation")]
    assert len(fields) == 1, result.events
    rows = field_rows(current_run)
    assert len(rows) == 1 and rows[0]["disambiguation_attempts"] == 1
    assert rows[0]["attribute_signature"] == fields[0]["attribute_disambiguation"]["signature"]


def test_merged_value_is_deferred_in_each_logical_row_but_has_one_field_task(
    tmp_path, monkeypatch, current_run,
):
    original_arguments = fixture.arguments

    def arguments(path, lines):
        values = original_arguments(path, lines)
        document = Document()
        document.add_heading("设备明细", 1)
        table = document.add_table(rows=3, cols=2)
        for row, cells in enumerate([
            ["对象", "编号"], ["设备一", "A-001"], ["设备二", ""],
        ]):
            for column, text in enumerate(cells):
                table.cell(row, column).text = text
        table.cell(1, 1).merge(table.cell(2, 1))
        source = path / "merged-field.docx"
        document.save(source)
        analysis = analyze_word_core(source)
        values.update(ir=analysis.ir, metadata=prepare_metadata(
            analysis.ir, section_tree=analysis.structure.section_tree.to_dict(),
            summary_version="merged-field-test",
        ), filename=source.name)
        return values

    monkeypatch.setattr(fixture, "arguments", arguments)
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["replaced by real merged table"],
    )
    index = RecordIndex(args["ir"])
    shared = next(unit for unit in args["ir"].evidence_units if unit.text == "A-001")
    owners = index.records_by_evidence[shared.evidence_id]
    assert len(index.records) == len(owners) == 2
    result = factory().run(**args, **hooks)
    ordinary = [view for view in requests if view["stage"] == "discovery"
                and "members" not in view and not view.get("attribute_disambiguation")]
    assert len(ordinary) == 4, result.diagnostics  # Two candidate cards per logical row.
    assert {view["record_id"] for view in ordinary} == {record.record_id for record in owners}
    rows = field_rows(current_run)
    assert len(rows) == 1, result.events
    field_id = rows[0]["task"]["field_id"]
    for view in ordinary:
        deferred = view["deferred_property_fields"]
        assert [field["field_id"] for field in deferred] == [field_id]
        assert deferred[0]["value"] == "A-001"
        assert {ref["evidence_id"] for ref in deferred[0]["value_refs"]} == {shared.evidence_id}
        assert shared.evidence_id in {unit["evidence_id"] for unit in view["evidence_units"]
                                      if unit["fact_eligible"]}
