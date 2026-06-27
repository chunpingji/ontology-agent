"""集成测试：air-gap 结构化抽取端到端（零联网 / 零回归 / 零误降级）。

断网 + 无 Key 下跑结构化 Excel / 含表格 Word，断言：作业成功（status=reviewing）、
结构化候选与黄金基线逐字一致（O7 零回归）、所有 `_emit(..., degraded=False)`（O5）、
全程不触达 anthropic（O9/O10 零外发）、Word 中文表头按映射转 IRI（T014）。

直接驱动 `run_extraction_pipeline`（对齐 test_doc_repo_pipeline.py 的注入式集成做法），
以便监听 `_emit` 与注入「外发即违例」的 anthropic 哨兵模块。
覆盖 quickstart 场景 1/5（SC-001/002/003）。
"""

from __future__ import annotations

import asyncio
import sys
import types

import docx
import openpyxl
import pytest

from app.config import settings
from app.models.extraction import ExtractionCandidate, ExtractionConfig, ExtractionJob
from app.services.extraction import pipeline as pipeline_module
from app.services.extraction.pipeline import run_extraction_pipeline

EQUIP = "http://slpra.org/equipment#Equipment"
MAPPING = {"设备编号": "equipmentID", "设备名称": "equipmentName", "材质": "material"}

# 黄金基线：结构化列经 column_mapping 落 IRI 键（剔除受控词表派生的 _controlled_vocab）。
GOLD_ROWS = [
    {"equipmentID": "CT64201", "equipmentName": "压片机A", "material": "316L不锈钢"},
    {"equipmentID": "DE64203", "equipmentName": "包衣机B", "material": "304不锈钢"},
]


def _xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["设备编号", "设备名称", "材质"])
    ws.append(["CT64201", "压片机A", "316L不锈钢"])
    ws.append(["DE64203", "包衣机B", "304不锈钢"])
    wb.save(path)
    return path


def _docx(path):
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "设备编号", "设备名称", "材质"
    body = table.rows[1].cells
    body[0].text, body[1].text, body[2].text = "CT64201", "压片机A", "316L不锈钢"
    doc.save(path)
    return path


def _config(db, source_type):
    cfg = ExtractionConfig(name="设备台账", target_class_iri=EQUIP,
                           source_type=source_type, column_mapping=dict(MAPPING))
    db.add(cfg)
    db.commit()
    db.refresh(cfg)
    return cfg


def _job(db, source_type, filename):
    job = ExtractionJob(source_type=source_type, source_filename=filename,
                        source_config={}, status="pending")
    db.add(job)
    db.commit()
    db.refresh(job)
    return job


def _instance_props(db, job):
    cands = (db.query(ExtractionCandidate)
             .filter(ExtractionCandidate.job_id == job.id,
                     ExtractionCandidate.candidate_kind == "instance")
             .all())
    out = []
    for c in cands:
        out.append({k: v for k, v in c.extracted_properties.items()
                    if k != "_controlled_vocab"})
    return cands, out


@pytest.fixture()
def offline(monkeypatch):
    """断网 + 无 Key + anthropic 外发哨兵 + 监听 _emit 的 degraded 标志。"""
    monkeypatch.setattr(settings, "llm_cloud_enabled", False)
    monkeypatch.setattr(settings, "anthropic_api_key", "")

    class _Boom(types.ModuleType):
        def __getattr__(self, name):
            raise AssertionError(f"air-gap 违例：抽取触达 anthropic.{name}（应零外发）")

    monkeypatch.setitem(sys.modules, "anthropic", _Boom("anthropic"))

    emits: list = []
    real_emit = pipeline_module._emit

    def _spy_emit(job, stage, pct, status, degraded=False):
        emits.append((stage, degraded))
        return real_emit(job, stage, pct, status, degraded=degraded)

    monkeypatch.setattr(pipeline_module, "_emit", _spy_emit)
    return emits


def test_excel_offline_zero_regression(db, fake_engine, offline, tmp_path):
    """场景 1：断网结构化 Excel→作业成功、候选逐字黄金基线、离线非降级、进度全 degraded=False。"""
    cfg = _config(db, "excel")
    job = _job(db, "excel", "台账.xlsx")
    asyncio.run(run_extraction_pipeline(
        job, cfg, _xlsx(tmp_path / "台账.xlsx"), fake_engine, db))

    db.refresh(job)
    assert job.status == "reviewing"

    cands, props = _instance_props(db, job)
    assert len(cands) == 2
    assert props == GOLD_ROWS                                  # O7 逐字零回归
    assert all(c.degraded_reason is None for c in cands)       # O3 离线非降级
    assert emits_all_not_degraded(offline)                     # O5 进度一致


def test_word_headers_mapped_to_iri_offline(db, fake_engine, offline, tmp_path):
    """场景 5：断网含表格 Word→表头经 column_mapping 走确定性 IRI 键（替代云端），非降级。"""
    cfg = _config(db, "word")
    job = _job(db, "word", "SOP.docx")
    asyncio.run(run_extraction_pipeline(
        job, cfg, _docx(tmp_path / "SOP.docx"), fake_engine, db))

    db.refresh(job)
    assert job.status == "reviewing"

    cands, props = _instance_props(db, job)
    assert len(cands) == 1
    assert props[0] == {
        "equipmentID": "CT64201", "equipmentName": "压片机A", "material": "316L不锈钢",
    }
    assert cands[0].degraded_reason is None
    assert emits_all_not_degraded(offline)


def emits_all_not_degraded(emits) -> bool:
    """进度事件须至少有一条，且全部 degraded=False（离线不广播 degraded=True）。"""
    return bool(emits) and all(d is False for _, d in emits)
