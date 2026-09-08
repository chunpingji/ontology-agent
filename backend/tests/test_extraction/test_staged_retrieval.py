"""Chapter ranking never removes fallback records or expands fact permissions."""

from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.record_retrieval import RecordIndex
from app.evaluation.staged_retrieval import context_ranges, plan_relation_chapters
from app.schemas.evidence import Candidate, DocumentProvenance
from app.services.extraction.word_analysis import analyze_word_core

SCHEMA = {
    "urn:Report": {"label": "报告", "relationships": [
        {"iri": "urn:describes", "label": "描述", "range": ["urn:Item"]}], "properties": []},
    "urn:Item": {"label": "产品", "properties": [
        {"iri": "urn:name", "label": "项目名称"},
        {"iri": "urn:form", "label": "剂型", "aliases": ["制品形式"]},
        {"iri": "urn:temperature", "label": "温度"}], "relationships": [
            {"iri": "urn:contains", "label": "成分", "range": ["urn:Component"]}]},
    "urn:Component": {"label": "远端类秘密", "properties": [
        {"iri": "urn:remote", "label": "远端属性秘密"}], "relationships": []},
    "urn:Unrelated": {"label": "完全无关秘密", "properties": [], "relationships": []},
}


def build(document, tmp_path, schema=None, metadata=None):
    path = tmp_path / "staged.docx"
    document.save(path)
    ir = analyze_word_core(path).ir
    source = next(unit for unit in ir.evidence_units if unit.text)
    subject = Candidate(
        candidate_id="report", kind="entity", class_iri="urn:Report", text=source.text,
        validation_status="passed", identity={"document_root": ir.document_hash},
        provenance=[DocumentProvenance(anchors=[ir.anchor(source.evidence_id)])],
    )
    return RecordIndex(ir, schema or SCHEMA, metadata=metadata), subject


def plan(index, subject, **kwargs):
    return plan_relation_chapters(index, subject, SCHEMA["urn:Report"]["relationships"][0],
                                  index.schema, **kwargs)


def record(index, text):
    return next(item for item in index.records if any(
        index.ir.unit(region.evidence_id).text == text for region in item.source_ranges))


def text(index, ranges):
    return [index.ir.unit(region.evidence_id).text[region.start:region.end] for region in ranges]


def multichapter_document():
    document = Document()
    document.add_heading("产品资料", 1)
    document.add_paragraph("项目名称：项目甲")
    document.add_paragraph("剂型：胶囊")
    document.add_paragraph("温度：20")
    document.add_heading("补充资料", 1)
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [["产品", "温度"], ["样品乙", "25"], ["样品丙", "30"]],
                           strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_heading("复核资料", 1)
    document.add_paragraph("成分：组分丁")
    document.add_heading("原始记事", 1)
    document.add_paragraph("完全无关秘密，远端类秘密，远端属性秘密")
    return document


def test_two_phases_partition_all_nonheadings_and_round_robin_chapters(tmp_path):
    index, subject = build(multichapter_document(), tmp_path)
    result = plan(index, subject)
    primary, fallback = [phase["records"] for phase in result["phases"]]
    expected = {item.record_id for item in index.records if item.kind != "heading"}
    primary_ids = {item["record_id"] for item in primary}
    fallback_ids = {item["record_id"] for item in fallback}
    assert primary_ids.isdisjoint(fallback_ids)
    assert primary_ids | fallback_ids == expected
    assert len(result["records"]) == len(expected)
    assert len({item["section_node_id"] for item in primary[:3]}) == 3
    remote = record(index, "完全无关秘密，远端类秘密，远端属性秘密")
    assert remote.record_id in fallback_ids
    assert next(item for item in fallback if item["record_id"] == remote.record_id)[
        "rank_score"] == 0
    assert result["range_class_iris"] == ["urn:Item"]  # Never follow range's outgoing edges.
    assert set(result["excluded_heading_record_ids"]) == {
        item.record_id for item in index.records if item.kind == "heading"}
    assert result["coverage"]["unsearched_records"] == len(expected)
    assert not result["coverage"]["complete"]


def test_rejection_cannot_erase_later_chapters_and_plan_is_deterministic_detached(tmp_path):
    index, subject = build(multichapter_document(), tmp_path)
    before_records = [item.model_dump(mode="json") for item in index.records]
    before_schema = deepcopy(index.schema)
    result = plan(index, subject, phase1_section_limit=1)
    original = deepcopy(result)
    result["phases"][0]["records"][0]["outcome"] = "type_rejected"
    assert len(result["phases"][1]["records"]) > 0
    assert plan(index, subject, phase1_section_limit=1) == original
    assert [item.model_dump(mode="json") for item in index.records] == before_records
    assert index.schema == before_schema


def test_no_match_yields_empty_phase_one_and_all_records_in_phase_two(tmp_path):
    document = Document()
    document.add_heading("纯粹附录", 1)
    document.add_paragraph("没有任何线索")
    index, subject = build(document, tmp_path)
    result = plan(index, subject)
    assert result["phases"][0]["records"] == []
    assert len(result["phases"][1]["records"]) == 1
    assert result["phases"][1]["records"][0]["rationale"] == [
        "unmatched_record_retained_for_fallback"]


def test_local_aliases_properties_and_direct_relation_names_are_ranking_clues(tmp_path):
    document = Document()
    document.add_heading("附件", 1)
    document.add_paragraph("制品形式：注射液")
    document.add_heading("另一附件", 1)
    document.add_paragraph("成分：组分甲")
    index, subject = build(document, tmp_path)
    result = plan(index, subject)
    assert all(item["rank_score"] > 0 for item in result["records"])
    assert len(result["phases"][0]["section_node_ids"]) == 2


def test_range_subclasses_use_transitive_parent_closure_with_cycle_guard(tmp_path):
    schema = deepcopy(SCHEMA)
    schema["urn:Child"] = {"label": "子类", "parents": ["urn:Item", "urn:Grandchild"]}
    schema["urn:Grandchild"] = {"label": "孙类", "parents": ["urn:Child"]}
    index, subject = build(multichapter_document(), tmp_path, schema=schema)
    assert plan(index, subject)["range_class_iris"] == ["urn:Child", "urn:Grandchild", "urn:Item"]


def test_only_current_local_canonical_predicate_and_index_schema_are_used(tmp_path):
    index, subject = build(multichapter_document(), tmp_path)
    malicious = {"iri": "urn:describes", "label": "完全无关秘密", "range": ["urn:Unrelated"]}
    assert plan_relation_chapters(index, subject, malicious, index.schema) == plan(index, subject)
    with pytest.raises(ValueError, match="local menu"):
        plan_relation_chapters(index, subject, {"iri": "urn:contains"}, index.schema)
    wrong_schema = deepcopy(index.schema)
    wrong_schema["urn:Item"]["label"] = "changed"
    with pytest.raises(ValueError, match="schema does not match"):
        plan_relation_chapters(index, subject, malicious, wrong_schema)
    with pytest.raises(ValueError, match="validated affirmative"):
        plan(index, subject.model_copy(update={"validation_status": "rejected"}))
    with pytest.raises(ValueError, match="document root"):
        plan(index, subject.model_copy(update={"identity": {"document_root": "other-document"}}))


def test_usable_summary_can_rank_but_stale_summary_does_not_and_never_becomes_source(tmp_path):
    document = Document()
    document.add_heading("普通章节", 1)
    document.add_paragraph("无词汇交集")
    index, subject = build(document, tmp_path)
    target = record(index, "无词汇交集")
    metadata = {target.section_node_id: {"content_summary": "产品剂型与项目名称",
                                       "analysis_id": index.ir.analysis_id,
                                       "structure_hash": index.ir.structure_hash,
                                       "summary_status": "complete"}}
    summarized = RecordIndex(index.ir, index.schema, metadata=metadata)
    ranked = plan(summarized, subject)
    assert ranked["phases"][0]["records"][0]["score_components"]["summaries"] > 0
    assert text(summarized, target.source_ranges) == ["无词汇交集"]
    metadata[target.section_node_id]["summary_status"] = "stale"
    stale = RecordIndex(index.ir, index.schema, metadata=metadata)
    assert plan(stale, subject)["phases"][0]["records"] == []


def test_context_default_retains_whole_field_block_beyond_nine_paragraphs(tmp_path):
    document = Document()
    document.add_heading("产品资料", 1)
    document.add_paragraph("项目名称：项目甲")
    for number in range(9):
        document.add_paragraph(f"结构字段{number}：" + "原始内容" * 80)
    document.add_paragraph("剂型：胶囊")
    index, subject = build(document, tmp_path)
    target = record(index, "项目名称：项目甲")
    before = target.model_dump(mode="json")
    ranges = context_ranges(index, target)
    assert len(ranges) == 10
    assert "剂型：胶囊" in text(index, ranges)
    assert sum(len(value) for value in text(index, ranges)) > 1200
    assert "项目名称：项目甲" not in text(index, ranges)
    assert target.model_dump(mode="json") == before
    authorized = index.plan(subject, SCHEMA["urn:Report"]["relationships"][0], kind="relationship")
    planned = next(item for item in authorized.records if item.record_id == target.record_id)
    context = index.record_context(planned)
    assert {item["text"] for item in context["fragments"] if item["fact_eligible"]} == {
        "项目名称：项目甲"}
    assert len(context_ranges(index, target, radius=2)) == 2


@pytest.mark.parametrize("header_only", [False, True])
def test_context_stops_at_tables_including_header_only_tables(tmp_path, header_only):
    document = Document()
    document.add_heading("产品资料", 1)
    document.add_paragraph("项目名称：项目甲")
    table = document.add_table(rows=1 if header_only else 2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "产品", "温度"
    if not header_only:
        table.cell(1, 0).text, table.cell(1, 1).text = "样品乙", "25"
    document.add_paragraph("剂型：胶囊")
    index, subject = build(document, tmp_path)
    assert context_ranges(index, record(index, "项目名称：项目甲")) == []
    assert context_ranges(index, record(index, "剂型：胶囊")) == []


def test_context_stops_at_heading_section_boundaries_and_leaves_table_rows_unchanged(tmp_path):
    index, subject = build(multichapter_document(), tmp_path)
    target = record(index, "温度：20")
    assert set(text(index, context_ranges(index, target))) == {"项目名称：项目甲", "剂型：胶囊"}
    row = record(index, "25")
    before = row.model_dump(mode="json")
    assert context_ranges(index, row) == []
    assert text(index, row.source_ranges) == ["样品乙", "25"]
    assert set(text(index, row.binding_ranges)) == {"产品", "温度"}
    assert row.model_dump(mode="json") == before
    heading = next(item for item in index.records if item.kind == "heading")
    assert context_ranges(index, heading) == []


def test_context_is_detached_and_rejects_foreign_record_or_invalid_radius(tmp_path):
    index, subject = build(multichapter_document(), tmp_path)
    target = record(index, "项目名称：项目甲")
    ranges = context_ranges(index, target)
    ranges[0].end = 1
    assert context_ranges(index, target)[0].end > 1
    with pytest.raises(ValueError, match="does not belong"):
        context_ranges(index, target.model_copy(update={"record_id": "foreign"}))
    with pytest.raises(ValueError, match="radius"):
        context_ranges(index, target, radius=-1)


def test_merged_and_nested_table_records_keep_all_original_memberships_and_headers(tmp_path):
    document = Document()
    document.add_heading("产品资料", 1)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "产品", "温度"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "共享实例"
    table.cell(1, 1).text, table.cell(2, 1).text = "25", "30"
    outer = document.add_table(rows=2, cols=2)
    outer.cell(0, 0).text, outer.cell(0, 1).text = "外层名称", "外层内容"
    outer.cell(1, 0).text = "容器甲"
    nested = outer.cell(1, 1).add_table(rows=2, cols=2)
    nested.cell(0, 0).text, nested.cell(0, 1).text = "产品", "温度"
    nested.cell(1, 0).text, nested.cell(1, 1).text = "内层实例", "40"
    index, subject = build(document, tmp_path)
    before = [item.model_dump(mode="json") for item in index.records]
    result = plan(index, subject)
    assert {item["record_id"] for item in result["records"]} == {
        item.record_id for item in index.records if item.kind != "heading"}
    merged = record(index, "30")
    assert text(index, merged.source_ranges) == ["共享实例", "30"]
    assert any(item["data_row_indices"] == [1, 2] for item in merged.cell_memberships)
    inner = record(index, "40")
    assert len(inner.table_path) == 3
    assert set(text(index, inner.binding_ranges)) == {"产品", "温度"}
    assert inner.parent_links[0]["relationship"] == "physical_containment_only"
    assert context_ranges(index, inner) == []
    assert [item.model_dump(mode="json") for item in index.records] == before
