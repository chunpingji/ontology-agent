from __future__ import annotations

from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.schema_card_mock_retrieval import (
    CMC,
    MAX_RECORDS,
    MAX_SOURCE_CHARACTERS,
    PROCESS_EQUIPMENT,
    USES_EQUIPMENT,
    prepare_retrieval,
)
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core


def fixtures(tmp_path, document, summaries=None):
    path = tmp_path / "mock-retrieval.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    tree = analysis.structure.section_tree.to_dict()

    def fill(node):
        if summary := (summaries or {}).get(node.get("heading")):
            node.setdefault("layer_metadata", {}).update(
                content_summary=summary, summary_status="completed", summary_source="llm",
            )
        for child in node.get("children", []):
            fill(child)

    fill(tree)
    metadata = prepare_metadata(analysis.ir, section_tree=tree, summary_version="test-frozen")
    catalog = {
        CMC: {"class_iri": CMC, "label": "CMC 报告", "properties": [], "allowed_relations": [{
            "iri": USES_EQUIPMENT, "kind": "relationship", "label": "使用设备",
            "range_class_iris": [PROCESS_EQUIPMENT],
        }]},
        PROCESS_EQUIPMENT: {"class_iri": PROCESS_EQUIPMENT, "label": "工艺设备", "properties": [{
            "iri": "urn:equipmentID", "label": "设备编号", "kind": "property",
            "datatype_iris": ["http://www.w3.org/2001/XMLSchema#string"],
        }]},
    }
    return analysis.ir, metadata, catalog


def run(tmp_path, paragraphs, mock_records=()):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    return prepare_retrieval(*fixtures(tmp_path, document), list(mock_records))


def test_mock_names_ids_change_selection_without_entering_original_sources(tmp_path):
    paragraphs = ["使用设备一", "使用设备二", "使用设备三", "XX091 离心机"]
    result = run(tmp_path, paragraphs, [{"equipment_id": "XX091", "name": "离心机",
                                        "material": "档案独有材质"}])
    baseline, enabled = [result[arm]["cases"] for arm in ("baseline", "mock_retrieval")]
    assert [case["content"]["text"] for case in baseline] == paragraphs[:3]
    assert enabled[0]["content"]["text"] == paragraphs[3]
    assert "档案独有材质" not in str(enabled)
    assert result["ablation"]["added_record_ids"] == [enabled[0]["record_id"]]
    assert result["baseline"]["ranking"]["query_count"] == 1
    assert len(baseline) == len(enabled) == MAX_RECORDS == 3
    assert result["query"]["predicate"]["iri"] == USES_EQUIPMENT
    assert all(case["root_class_iri"] == CMC for case in enabled)


def test_exact_identifier_boundaries_distinct_names_and_repetitions(tmp_path):
    result = run(tmp_path, ["XAB10 AB100 AB10_ ab10 MOCK-AB10 AB10-OLD",
                           "AB10 AB10或CD20 离心机 离心机"], [
        {"equipment_id": "AB10", "name": "离心机"},
        {"equipment_id": "CD20", "name": "离心机"},
    ])
    rows = sorted(result["mock_retrieval"]["ranking"]["records"],
                  key=lambda row: row["source_position"])
    assert rows[0]["mock_score"] == 0
    assert rows[0]["selection_status"] == "not_selected_nonpositive"
    assert rows[1]["mock_score"] == 9
    assert rows[1]["score_components"]["distinct_exact_ids"] == 2
    assert rows[1]["score_components"]["distinct_exact_names"] == 1
    assert [hit["term"] for hit in rows[1]["mock_hits"]["exact_ids"]] == ["AB10", "CD20"]
    assert len(rows[1]["mock_hits"]["exact_ids"][0]["anchors"]) == 2
    assert len(result["mock_retrieval"]["cases"]) == 1


def test_mock_scoring_does_not_join_units_or_use_context_and_summaries(tmp_path):
    document = Document()
    document.add_heading("AB10 离心机", 1)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("AB10", "离心机"), ("AB", "10")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    ir, metadata, catalog = fixtures(tmp_path, document, {"AB10 离心机": "AB10 离心机 工艺设备"})
    result = prepare_retrieval(ir, metadata, catalog, [{"equipment_id": "AB10", "name": "离心机"}])
    assert all(row["mock_score"] == 0 for row in result["mock_retrieval"]["ranking"]["records"])
    case = result["mock_retrieval"]["cases"][0]
    assert "AB10" in case["content"]["text"]
    assert "工艺设备" not in case["content"]["text"]


def test_only_single_cmc_predicate_contributes_despite_other_classes_and_relations(tmp_path):
    document = Document()
    document.add_paragraph("使用设备一")
    document.add_paragraph("无关独有术语")
    ir, metadata, catalog = fixtures(tmp_path, document)
    baseline = prepare_retrieval(ir, metadata, catalog, [])
    expanded = deepcopy(catalog)
    expanded[CMC]["allowed_relations"].append({
        "iri": "urn:unrelated", "kind": "relationship", "label": "无关独有术语",
        "range_class_iris": [PROCESS_EQUIPMENT],
    })
    expanded["urn:other"] = dict(catalog[CMC])
    result = prepare_retrieval(ir, metadata, expanded, [])
    assert result["baseline"]["ranking"] == baseline["baseline"]["ranking"]
    assert result["baseline"]["cases"] == baseline["baseline"]["cases"]
    assert result["query"]["subject_class_iri"] == CMC


def test_full_context_roles_and_ir_order_without_neighboring_field_groups(tmp_path):
    document = Document()
    document.add_heading("父标题", 1)
    document.add_heading("子标题", 2)
    document.add_paragraph("无关字段：不应扩展")
    document.add_paragraph("设备编号：AB10")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("设备编号", "设备名称"), ("CD20", "离心机")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("备注：仅限研发用途。")
    ir, metadata, catalog = fixtures(tmp_path, document)
    result = prepare_retrieval(ir, metadata, catalog, [{"equipment_id": "CD20", "name": "离心机"}])
    cases = result["mock_retrieval"]["cases"]
    assert cases
    table_case = next(case for case in cases if "CD20" in case["content"]["text"])
    for text in ("父标题", "子标题", "设备编号", "设备名称", "CD20", "离心机",
                 "备注：仅限研发用途。"):
        assert text in table_case["content"]["text"]
    paragraph_case = next(case for case in cases if "设备编号：AB10" in case["content"]["text"])
    assert "无关字段" not in paragraph_case["content"]["text"]
    index = RecordIndex(ir)
    for case in cases:
        assert "field_group" not in case["roles"]["refs"]
        assert case["roles"]["primary_refs"]
        assert set(case["roles"]["primary_refs"]).isdisjoint(case["roles"]["context_refs"])
        ids = [unit["evidence_id"] for unit in case["source_units"]]
        assert ids == sorted(set(ids), key=index.positions.__getitem__)
        assert case["source_characters"] == sum(len(ir.unit(eid).text) for eid in ids)
        assert list(case["sources"]) == [f"u{i}" for i in range(len(ids))]
        assert all(unit == ir.unit(unit["evidence_id"]).model_dump(mode="json")
                   for unit in case["source_units"])


def test_nested_parent_is_context_and_never_primary(tmp_path):
    document = Document()
    outer = document.add_table(rows=2, cols=2)
    for row, values in zip(outer.rows, [("对象", "细节"), ("母记录甲", "背景")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    inner = outer.cell(1, 1).add_table(rows=2, cols=2)
    for row, values in zip(inner.rows, [("项目", "参数"), ("使用设备", "AB10")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    ir, metadata, catalog = fixtures(tmp_path, document)
    result = prepare_retrieval(ir, metadata, catalog, [])
    case = next(case for case in result["baseline"]["cases"] if case["roles"]["refs"]["parent"])
    assert "母记录甲" in case["content"]["text"]
    parent_ids = {ref["evidence_id"] for ref in case["roles"]["refs"]["parent"]}
    assert all(case["sources"][ref]["evidence_id"] not in parent_ids
               for ref in case["roles"]["primary_refs"])


def test_oversized_record_is_deferred_then_three_complete_later_cases_selected(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    long = "使用设备 " + "长" * MAX_SOURCE_CHARACTERS
    for row, values in zip(table.rows, [("设备编号", long), ("AB10", "离心机")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    for i in range(4):
        document.add_paragraph(f"使用设备{i}")
    ir, metadata, catalog = fixtures(tmp_path, document)
    result = prepare_retrieval(ir, metadata, catalog, [{"equipment_id": "AB10", "name": "离心机"}])
    for arm in ("baseline", "mock_retrieval"):
        assert [case["content"]["text"] for case in result[arm]["cases"]] == [
            "使用设备0", "使用设备1", "使用设备2"]
        ranking = result[arm]["ranking"]
        assert ranking["budgets"]["max_source_characters_per_case"] == 3000
        assert {row["reason"] for row in ranking["deferred_records"]} == {
            "source_character_budget", "record_count_budget"}
        assert len(ranking["records"]) == len(RecordIndex(ir).records)
    assert next(unit for unit in ir.evidence_units if unit.text == long).text == long


def test_per_case_character_budget_does_not_accumulate_across_selected_records(tmp_path):
    texts = ["使用设备" + str(i) + "长" * 2000 for i in range(3)]
    result = run(tmp_path, texts)
    assert len(result["baseline"]["cases"]) == 3
    assert sum(case["source_characters"] for case in result["baseline"]["cases"]) > 3000
    assert all(case["source_characters"] <= 3000 for case in result["baseline"]["cases"])


def test_empty_mock_and_stable_ties_leave_sources_identical_without_mutation(tmp_path):
    document = Document()
    for text in ("使用设备甲", "使用设备乙", "使用设备丙", "使用设备丁"):
        document.add_paragraph(text)
    ir, metadata, catalog = fixtures(tmp_path, document)
    originals = (ir.model_dump(mode="json"), metadata.model_dump(mode="json"), deepcopy(catalog))
    result = prepare_retrieval(ir, metadata, catalog, [])
    assert result["baseline"]["cases"] == result["mock_retrieval"]["cases"]
    assert result["ablation"]["selected_order_changed"] is False
    assert result["ablation"]["new_model_calls"] == 0
    assert [row["source_position"] for row in result["baseline"]["ranking"]["records"]] == sorted(
        row["source_position"] for row in result["baseline"]["ranking"]["records"])
    assert originals == (ir.model_dump(mode="json"), metadata.model_dump(mode="json"), catalog)


@pytest.mark.parametrize("field,value", [("analysis_id", "wrong"),
                                        ("document_hash", "1" * 64),
                                        ("structure_hash", "2" * 64)])
def test_rejects_cross_source_metadata(tmp_path, field, value):
    ir, metadata, catalog = fixtures(tmp_path, Document())
    with pytest.raises(ValueError, match="retrieval_metadata_source_mismatch"):
        prepare_retrieval(ir, metadata.model_copy(update={field: value}), catalog, [])


def test_rejects_stale_metadata_and_ir(tmp_path):
    document = Document()
    document.add_paragraph("使用设备")
    ir, metadata, catalog = fixtures(tmp_path, document)
    changed = metadata.model_copy(deep=True)
    changed.node_summaries[0].summary = "changed"
    with pytest.raises(ValueError, match="metadata_identity_mismatch"):
        prepare_retrieval(ir, changed, catalog, [])
    with pytest.raises(ValueError, match="record_universe_mismatch"):
        prepare_retrieval(ir, metadata.model_copy(update={"source_record_refs": []}), catalog, [])
    changed_ir = ir.model_copy(deep=True)
    changed_ir.evidence_units[0].text = "篡改"
    with pytest.raises(ValueError, match="structure identity"):
        prepare_retrieval(changed_ir, metadata, catalog, [])


@pytest.mark.parametrize("change,error", [
    ("missing_root", "cmc_card_missing"), ("missing_relation", "relation_invalid"),
    ("duplicate_relation", "relation_invalid"), ("wrong_range", "relation_invalid"),
    ("missing_range_card", "range_class_outside_catalog"),
])
def test_rejects_invalid_equipment_query_contract(tmp_path, change, error):
    ir, metadata, catalog = fixtures(tmp_path, Document())
    if change == "missing_root":
        del catalog[CMC]
    elif change == "missing_relation":
        catalog[CMC]["allowed_relations"] = []
    elif change == "duplicate_relation":
        catalog[CMC]["allowed_relations"] *= 2
    elif change == "wrong_range":
        catalog[CMC]["allowed_relations"][0]["range_class_iris"] = ["urn:wrong"]
    else:
        del catalog[PROCESS_EQUIPMENT]
    with pytest.raises(ValueError, match=error):
        prepare_retrieval(ir, metadata, catalog, [])


@pytest.mark.parametrize("records", [[{"equipment_id": "", "name": "名称"}],
                                    [{"equipment_id": "ID", "name": " "}],
                                    [{"equipment_id": " ID", "name": "名称"}]])
def test_rejects_invalid_mock_terms(tmp_path, records):
    with pytest.raises(ValueError, match="term_invalid"):
        run(tmp_path, ["材料"], records)


def test_rejects_duplicate_mock_keys_even_if_name_matches(tmp_path):
    with pytest.raises(ValueError, match="duplicate_equipment_id"):
        run(tmp_path, ["材料"], [{"equipment_id": "ID", "name": "名称"}] * 2)
