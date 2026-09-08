"""Scheduler contracts with production validation and deterministic model fixtures."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.root_guided_variant import build_root_guided_variant
from app.schemas.evidence import TaskBudget
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:Report": {
        "label": "Report",
        "properties": [{"iri": "urn:code", "label": "code", "datatype": "string"}],
        "relationships": [
            {"iri": "urn:describes", "label": "describes", "range": ["urn:Product"]},
            {"iri": "urn:optional", "label": "optional", "range": ["urn:Unused"]},
        ],
    },
    "urn:Product": {
        "label": "Product",
        "properties": [
            {"iri": "urn:strength", "label": "strength", "datatype": "decimal"},
        ],
        "relationships": [{"iri": "urn:uses", "label": "uses", "range": ["urn:Device"]}],
    },
    "urn:Device": {"label": "Device", "properties": [], "relationships": []},
    "urn:Unused": {"label": "Unused", "properties": [], "relationships": []},
    "urn:Unreachable": {"label": "Unreachable", "properties": [], "relationships": []},
}


def source(tmp_path, *, chapters=False):
    document = Document()
    if chapters:
        document.add_heading("Product", level=1)
    document.add_paragraph(
        "Report code R1 describes Product A and Product B. "
        "Product A uses Device D. Strength A=5 and B=7."
    )
    if chapters:
        document.add_heading("Unrelated records", level=1)
        document.add_paragraph("Unrelated administrative material.")
    path = tmp_path / "root-guided.docx"
    document.save(path)
    return analyze_word_core(path).ir


def fixture_model(calls, *, reject_root=False, hallucinate=False):
    def model(system, user, response_schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        calls.append((request, context))
        task = context["task"]
        targets = [f for f in context["fragments"] if f["purpose"] == "target"]

        def quote(fragment, text):
            return {"evidence_id": fragment["anchor"]["evidence_id"], "text": text}

        if request["stage"] == "verify_entity_types":
            assert "discovery_relation" not in task["predicate_definition"]
            return {
                "decisions": [
                    {"candidate_id": c["candidate_id"], "supported": True, "reason": "fixture"}
                    for c in request["candidate"]["proposed_entities"]
                ]
            }
        if request["stage"] == "verify_binding":
            candidate = request["candidate"]
            fragment = next(f for f in targets if "Report code" in f["text"])
            return {
                "supported": not (reject_root and task["predicate_iri"] == "urn:describes"),
                "subject_candidate_id": candidate["subject"]["candidate_id"],
                "object_candidate_id": (candidate.get("object") or {}).get("candidate_id"),
                "assertion_status": candidate["assertion_status"],
                "method": "explicit_assertion",
                "assertion_spans": [quote(fragment, fragment["text"])],
            }
        if task["task_kind"] == "entity":
            assert "discovery_relation" in task["predicate_definition"]
            entities = []
            for fragment in targets:
                for name, iri in (
                    ("Product A", "urn:Product"),
                    ("Product B", "urn:Product"),
                    ("Device D", "urn:Device"),
                ):
                    classes = task["predicate_definition"]["classes"]
                    class_key = next(
                        (
                            key
                            for key, value in classes.items()
                            if key == iri or value.get("name") == iri
                        ),
                        None,
                    )
                    if class_key is not None and name in fragment["text"]:
                        # Full assertion gives unique source ownership even if names repeat.
                        text = "Invented from summary" if hallucinate else name
                        start = fragment["text"].index(name) + (
                            fragment["anchor"]["span_start"] or 0
                        )
                        mention = quote(fragment, text)
                        if not hallucinate:
                            if "output_contract" in request:
                                mention["context"] = next(
                                    sentence + "."
                                    for sentence in fragment["text"].split(".")
                                    if name in sentence
                                )
                            else:
                                mention.update(start=start, end=start + len(name))
                        entities.append({"class_iri": class_key, "mention": mention})
            return {"entities": entities}
        assertions = []
        for fragment in targets:
            if "Report code" not in fragment["text"]:
                continue
            if task["task_kind"] == "relationship":
                owner = context["subjects"][task["subject"]["candidate_id"]]
                for ref in task["object_candidates"]:
                    obj = context["subjects"][ref["candidate_id"]]
                    if (
                        task["predicate_iri"] == "urn:describes"
                        and obj["class_iri"] == "urn:Product"
                        or task["predicate_iri"] == "urn:uses"
                        and owner["text"] == "Product A"
                        and obj["text"] == "Device D"
                    ):
                        assertions.append(
                            {
                                "object_candidate_id": ref["candidate_id"],
                                "assertion_spans": [quote(fragment, fragment["text"])],
                            }
                        )
            else:
                owner = context["subjects"][task["subject"]["candidate_id"]]
                value = (
                    "R1"
                    if task["predicate_iri"] == "urn:code"
                    else ("5" if owner["text"] == "Product A" else "7")
                )
                assertions.append(
                    {
                        "value": quote(fragment, value),
                        "assertion_spans": [quote(fragment, fragment["text"])],
                    }
                )
        return {"assertions": [{"assertion_status": "affirmed", **a} for a in assertions]}

    return model


def runner_for(ir, model, *, schema=None, max_sections=3, max_hops=4, metadata=None):
    base = GenericExtractionRunner(
        deepcopy(schema or SCHEMA),
        TestTokenizer(),
        model,
        model_identity="root-guided-fixture",
        budget=TaskBudget(
            max_input_tokens=100000,
            max_output_tokens=4096,
            max_regions_per_task=8,
            max_tasks=100,
            max_hops=max_hops,
        ),
    )
    return base, build_root_guided_variant(
        base,
        ir,
        metadata=metadata,
        max_sections_per_predicate=max_sections,
    )


def test_direct_range_menus_and_verified_edges_gate_recursion(tmp_path):
    ir, calls = source(tmp_path), []
    base, runner = runner_for(ir, fixture_model(calls))
    before = evidence_hash([base.schema, ir])
    result = runner.run(ir, effective_class="urn:Report")
    assert result.completion == "complete", result.diagnostics
    recalls = [(r, c) for r, c in calls if r["stage"] == "recall"]
    menus = [
        set(c["task"]["predicate_definition"]["classes"])
        for _, c in recalls
        if c["task"]["task_kind"] == "entity"
    ]
    assert menus[0] == {"urn:Product"}
    assert all(len(menu) == 1 and "urn:Unreachable" not in menu for menu in menus)
    first_device = next(
        i
        for i, (r, c) in enumerate(calls)
        if r["stage"] == "recall"
        and set(c["task"]["predicate_definition"].get("classes", {})) == {"urn:Device"}
    )
    root_binding = next(
        i
        for i, (r, c) in enumerate(calls)
        if r["stage"] == "verify_binding" and c["task"]["predicate_iri"] == "urn:describes"
    )
    assert root_binding < first_device
    first_other_root = next(
        i
        for i, (r, c) in enumerate(calls)
        if r["stage"] == "recall"
        and set(c["task"]["predicate_definition"].get("classes", {})) == {"urn:Unused"}
    )
    assert first_device < first_other_root
    assert any(
        c.kind == "relationship" and len(c.relationship_path) == 2 and c.positive_eligible
        for c in result.candidates
    )
    properties = [c for c in result.candidates if c.kind == "property" and c.positive_eligible]
    assert {c.literal.raw_value for c in properties} == {"R1", "5", "7"}
    assert evidence_hash([base.schema, ir]) == before
    assert runner.plan["coverage"] == result.checkpoint["coverage"]


def test_unsupported_parent_binding_prevents_child_discovery_and_properties(tmp_path):
    ir, calls = source(tmp_path), []
    _, runner = runner_for(ir, fixture_model(calls, reject_root=True))
    result = runner.run(ir, effective_class="urn:Report")
    entity_menus = [
        list(c["task"]["predicate_definition"]["classes"])
        for r, c in calls
        if r["stage"] == "recall" and c["task"]["task_kind"] == "entity"
    ]
    assert ["urn:Device"] not in entity_menus
    assert not any(
        c.kind == "property" and c.predicate_iri == "urn:strength" for c in result.candidates
    )
    assert not any(c.kind == "relationship" and c.positive_eligible for c in result.candidates)


def test_pause_resume_restores_exact_work_without_repeat_calls(tmp_path):
    ir, calls = source(tmp_path), []
    _, runner = runner_for(ir, fixture_model(calls))
    first = runner.run(ir, effective_class="urn:Report", pause_after=1)
    assert first.completion == "incomplete"
    second = runner.run(ir, effective_class="urn:Report", checkpoint=first.checkpoint)
    assert second.completion == "complete", second.diagnostics
    count = len(calls)
    third = runner.run(ir, effective_class="urn:Report", checkpoint=second.checkpoint)
    assert len(calls) == count
    assert third.candidates == second.candidates
    fresh_calls = []
    _, fresh = runner_for(ir, fixture_model(fresh_calls))
    fresh.run(ir, effective_class="urn:Report")
    assert len(calls) == len(fresh_calls)


def test_routing_uses_metadata_but_reports_unsearched_sections(tmp_path):
    ir = source(tmp_path, chapters=True)
    metadata = {
        node["node_id"]: {"summary_status": "completed", "content_summary": ""} for node in ir.nodes
    }
    product_node = next(n for n in ir.nodes if n["heading"] == "Product")
    metadata[product_node["node_id"]]["content_summary"] = "Report code describes Product"
    _, runner = runner_for(
        ir, lambda *args: {"entities": [], "assertions": []}, metadata=metadata, max_sections=1
    )
    # A task-appropriate empty model preserves the real response schemas.
    runner.model_call = lambda s, u, r, b: (
        {"entities": []}
        if json.loads(json.loads(u)["context"])["task"]["task_kind"] == "entity"
        else {"assertions": []}
    )
    result = runner.run(ir, effective_class="urn:Report")
    assert result.completion == "incomplete"
    assert "root_guided_unsearched_sections" in result.diagnostics
    assert all(c["selected_node_ids"] == [product_node["node_id"]] for c in runner.coverage)
    assert all(c["deferred_node_ids"] for c in runner.coverage)
    changed = deepcopy(metadata)
    changed[product_node["node_id"]]["content_summary"] = "Different route"
    _, other = runner_for(ir, lambda *args: None, metadata=changed, max_sections=1)
    assert runner.input_id(ir, "urn:Report") != other.input_id(ir, "urn:Report")


def test_production_replay_rejects_summary_hallucinations(tmp_path):
    ir, calls = source(tmp_path), []
    _, runner = runner_for(ir, fixture_model(calls, hallucinate=True))
    result = runner.run(ir, effective_class="urn:Report")
    assert result.completion == "incomplete"
    assert "source_excerpt_mismatch" in result.diagnostics
    assert len(result.candidates) == 2  # Root and its independently valid code property.
    assert not any(c.class_iri == "urn:Product" for c in result.candidates)


def test_depth_boundary_keeps_endpoint_properties_and_marks_unexpanded_edges(tmp_path):
    ir, calls = source(tmp_path), []
    _, runner = runner_for(ir, fixture_model(calls), max_hops=1)
    result = runner.run(ir, effective_class="urn:Report")
    assert any(
        c.kind == "property" and c.predicate_iri == "urn:strength" and c.positive_eligible
        for c in result.candidates
    )
    assert not any(c.class_iri == "urn:Device" for c in result.candidates)
    assert any(c["status"] == "depth_limited" for c in runner.coverage)
    assert result.completion == "incomplete"


def test_ranges_include_declared_subclasses_without_relation_closure(tmp_path):
    ir = source(tmp_path)
    schema = deepcopy(SCHEMA)
    schema["urn:SpecialProduct"] = {"parents": ["urn:Product"], "relationships": []}
    _, runner = runner_for(ir, lambda *args: {"entities": []}, schema=schema)
    assert runner.direct_range_classes(schema["urn:Report"]["relationships"][0]) == [
        "urn:Product",
        "urn:SpecialProduct",
    ]


def test_explicit_root_required_and_stale_summaries_ignored(tmp_path):
    ir = source(tmp_path)
    metadata = {
        n["node_id"]: {
            "analysis_id": "wrong",
            "summary_status": "completed",
            "content_summary": "Invented",
        }
        for n in ir.nodes
    }
    _, runner = runner_for(ir, lambda *args: pytest.fail("model must not run"), metadata=metadata)
    assert runner.plan["summary_node_count"] == 0
    result = runner.run(ir)
    assert result.diagnostics == ["root_guided_requires_explicit_document_class"]


def test_compact_protocol_retains_relation_context_and_independent_verification(tmp_path):
    ir, calls, traces = source(tmp_path), [], []
    _, runner = runner_for(ir, fixture_model(calls))
    runner.compact_identifiers = True
    runner.trace_fn = traces.append
    result = runner.run(ir, effective_class="urn:Report")
    assert result.completion == "complete", result.diagnostics
    assert any(
        c.kind == "relationship" and len(c.relationship_path) == 2 and c.positive_eligible
        for c in result.candidates
    )
    assert all(json.loads(t["user"])["output_contract"] == t["schema"] for t in traces)
    for request, context in calls:
        task = context["task"]
        if request["stage"] == "recall" and task["task_kind"] == "entity":
            classes = task["predicate_definition"]["classes"]
            assert all(key.startswith("t") for key in classes)
            assert task["predicate_definition"]["discovery_relation"]["iri"].startswith("urn:")
            assert task["subject"]["candidate_id"].startswith("c")
        if request["stage"] == "verify_entity_types":
            assert "discovery_relation" not in task["predicate_definition"]
            assert "property_roles" in json.dumps(context)


def test_late_competitor_invalidates_relationship_and_existing_dependents(tmp_path):
    ir, calls = source(tmp_path), []
    schema = deepcopy(SCHEMA)
    schema["urn:Report"]["relationships"].append(
        {"iri": "urn:announces", "label": "announces", "range": ["urn:Product"]}
    )
    schema["urn:Device"]["properties"] = [
        {"iri": "urn:serial", "label": "serial", "datatype": "string"}
    ]
    delegate = fixture_model(calls)

    def model(system, user, response_schema, budget):
        response = delegate(system, user, response_schema, budget)
        request = json.loads(user)
        context = json.loads(request["context"])
        task = context["task"]
        if request["stage"] != "recall":
            return response
        if task["task_kind"] == "entity":
            predicate = task["predicate_definition"]["discovery_relation"]["iri"]
            if predicate in {"urn:describes", "urn:announces"}:
                wanted = "Product A" if predicate == "urn:describes" else "Product B"
                response["entities"] = [
                    c for c in response["entities"] if c["mention"]["text"] == wanted
                ]
        elif task.get("predicate_iri") == "urn:announces":
            target = next(f for f in context["fragments"] if f["purpose"] == "target")
            response["assertions"] = [
                {
                    "object_candidate_id": ref["candidate_id"],
                    "assertion_status": "affirmed",
                    "assertion_spans": [
                        {"evidence_id": target["anchor"]["evidence_id"], "text": target["text"]}
                    ],
                }
                for ref in task["object_candidates"]
                if context["subjects"][ref["candidate_id"]]["text"] == "Product B"
            ]
        return response

    _, runner = runner_for(ir, model, schema=schema)
    result = runner.run(ir, effective_class="urn:Report")
    old_edges = [
        c for c in result.candidates if c.kind == "relationship" and c.predicate_iri == "urn:uses"
    ]
    assert old_edges and all(c.validation_status == "conflict" for c in old_edges)
    descendant_properties = [c for c in result.candidates if c.predicate_iri == "urn:serial"]
    assert descendant_properties
    assert all(c.validation_status == "conflict" for c in descendant_properties)
    assert result.completion == "incomplete"
    assert {"root_guided_late_competing_subject", "root_guided_invalidated_dependency"} <= set(
        result.diagnostics
    )
    assert result.checkpoint["invalidations"] == runner.plan["invalidations"]
    assert result.stage_statistics["scheduler"]["invalidated"] >= 2
