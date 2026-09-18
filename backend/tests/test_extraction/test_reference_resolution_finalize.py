"""Frozen references are independently checked before canonical graph projection."""

from __future__ import annotations

import copy
from types import SimpleNamespace as NS

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.claim_freeze import freeze_proposal
from app.services.extraction.ontology_guided.claim_protocol import (
    ClaimCheckResult,
    DiscoveryEnvelope,
    EntityDependencyView,
    ExternalCandidate,
    IdentityKeySpec,
    SchemaCard,
    VerificationEnvelope,
    build_verification_input,
    finalize_claims,
    validate_verification,
)
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    EdgeSpec,
    ScopeMember,
    SubjectRef,
    TraversalScope,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.mentions import MentionRegistry
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def reference_document(tmp_path):
    document = Document()
    document.add_paragraph("批次A工艺产出 HRS-9267 粗品。")
    for header, value in (("得量", "10 kg"), ("存放", "密封")):
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "产品"
        table.cell(0, 1).text = header
        table.cell(1, 0).text = "HRS-9267粗品"
        table.cell(1, 1).text = value
    document.add_paragraph("批次B工艺产出 HRS-9267粗品，与批次A为不同生产对象。")
    document.add_paragraph("前述批次A粗品精制后得到 HRS-9267 成品。")
    path = tmp_path / "reference-resolution.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    entries = []
    seen = set()
    for record in index.records:
        for unit in record.source_units:
            if "HRS-9267" not in unit.text or unit.evidence_id in seen:
                continue
            seen.add(unit.evidence_id)
            label = ("HRS-9267 成品" if "成品" in unit.text else "HRS-9267 粗品"
                     if "HRS-9267 粗品" in unit.text else "HRS-9267粗品")
            entries.append(NS(record=record, unit=unit, label=label))
    assert len(entries) == 5
    return NS(index=index, entries=entries, registry=MentionRegistry(index.ir, index))


def run_pipeline(
    document, entry_index, *, priors=(), bind=True, binding_verdict="supported",
    relation=False, predicate_verdict="supported", generation=1, class_iri="urn:CrudeProduct",
    known_resolutions=None, declare_relation_bindings=True, external_identity=False,
    external_subject="current", duplicate_entity=False, binding_prior_indices=None,
):
    entry = document.entries[entry_index]
    root = SubjectRef(entity_id="document", revision=1, class_iri="urn:Report",
                      is_document_root=True)
    root_ref = VersionedRef(id=root.entity_id, revision=root.revision)
    predicate = EdgeSpec(iri="urn:describes", label="记载", declared_by=[root.class_iri],
                         range_class_iris=["urn:CrudeProduct", "urn:FinishedProduct"])
    task = RecognitionTask.create(
        subject=root, predicate_iri=predicate.iri, predicate_kind="relationship",
        record_id=entry.record.record_id, phase=1, hop=0, dependency_hash="frozen-input",
    )
    seed = VerificationTarget.create(
        run_fingerprint="reference-test",
        claim_ref=VersionedRef(id=task.claim_lineage_id, revision=1),
        task_id=task.task_id, check_kind="predicate_entailment",
        document_context=DocumentContext(document_hash=document.index.ir.document_hash,
                                         document_class_iri=root.class_iri, root_ref=root_ref),
        subject_ref=root, predicate_iri=predicate.iri, ontology_hash="ontology",
        source_scope_hash=evidence_hash(entry.record.record_id), context_hash=evidence_hash("seed"),
    )
    dependencies, proofs = [], {}
    for prior in priors:
        node = prior.result.nodes[0]
        entity = prior.frozen.entities[0]
        sources = prior.result.reference_resolutions[0]["source_refs"]
        values = dict(
            entity_ref=VersionedRef(id=node.entity_id, revision=node.revision),
            class_iri=node.class_iri, grounding_kind="mention", root_origin=None,
            proposal=entity, source_refs=sources, dependency_refs=node.dependency_refs,
        )
        # Hash the validated representation, including decoded source anchors.
        values["source_refs"] = [document.index.ir.anchor(
            source["evidence_id"], source["span_start"], source["span_end"],
        ) for source in sources]
        dependencies.append(EntityDependencyView(**values, content_hash=evidence_hash(values)))
        proofs[node.entity_id, node.revision] = node
    context = assemble_context(
        seed, entry.record.record_id, document.index, predicate=predicate,
        required_context_refs=[anchor for dependency in dependencies
                               for anchor in dependency.source_refs],
    )
    root_values = dict(entity_ref=root_ref, class_iri=root.class_iri,
                       grounding_kind="document_root", root_origin="user_specified",
                       proposal=None, source_refs=[], dependency_refs=[])
    dependencies.insert(0, EntityDependencyView(
        **root_values, content_hash=evidence_hash(root_values),
    ))
    card = SchemaCard(
        schema_card_id="reference-card", subject_ref=root_ref, class_iris=[root.class_iri],
        ontology_snapshot_id="ontology", menu_id="menu", predicates=[predicate],
        quantity_policies=[], identity_keys=[], unsupported_constraints=[],
    )
    quote = dict(evidence_id=entry.unit.evidence_id, text=entry.label, context_text=None)
    entity = dict(local_id="current", class_iri=class_iri, representation="mention",
                  mentions=[quote], record_components=[], identifier_claims=[])
    external_candidates, external_links = [], []
    if external_identity:
        # The test profile explicitly declares this source field as a key;
        # production name candidates alone never become identity keys.
        card.identity_keys = [IdentityKeySpec(
            class_iri=class_iri, property_iris=["urn:recordKey"], namespace="test-archive",
            scope="dataset", declaration_ref="frozen-test-profile",
        )]
        candidate = ExternalCandidate(
            candidate_id="archive-candidate", source_id="test-archive", system="mock",
            dataset="products", record_key="CP-A001", record_version="v1", class_iri=class_iri,
            matches=[dict(predicate_iri="urn:recordKey", document_quote=quote,
                          record_field="recordKey", record_value=quote["text"],
                          match_kind="exact_key")],
            mapped_fields=[], metadata=[],
        )
        external_candidates.append(candidate)
        external_links = [dict(
            local_id="archive-link",
            subject_id=(dependencies[1].entity_ref.id if external_subject == "existing"
                        else external_subject),
            external_candidate_id=candidate.candidate_id, identity_support=[quote],
        )]
    bindings = [dict(
        local_id=f"binding-{number}", source_id="current", target_id=dependency.entity_ref.id,
        binding_kind="coreference", support=[quote, dependency.proposal.mentions[0].model_dump()],
    ) for number, dependency in enumerate(dependencies[1:])
        if bind and (binding_prior_indices is None or number in binding_prior_indices)]
    relations = []
    if relation:
        assertion = dict(evidence_id=entry.unit.evidence_id, text=entry.unit.text,
                         context_text=None)
        relations = [dict(
            local_id="relation", subject_id=root.entity_id, predicate_iri=predicate.iri,
            object_ids=["current"], selection="all", bridge_support=[assertion],
            selection_support=[], qualifiers=dict(polarity="affirmed", modality="asserted",
                                                  condition_support=[], scope_qualifiers=[]),
            bridge_kind="document_subject_description", bridge_ref_ids=[],
            source_assertion=dict(subject_support=[],
                                  object_support=[dict(object_id="current", support=[quote])],
                                  predicate_support=[assertion],
                                  binding_ids=[binding["local_id"] for binding in bindings]
                                  if declare_relation_bindings else []),
        )]
    proposal = DiscoveryEnvelope.model_validate(dict(
        entities=[entity, *([{**entity, "local_id": "duplicate"}] if duplicate_entity else [])],
        reference_bindings=bindings, relations=relations, properties=[],
        external_links=external_links, observations=[],
    ))
    frozen = freeze_proposal(
        proposal, task=task, context=context, card=card, index=document.index,
        generation=generation, entity_dependencies=dependencies, reference_resolution=True,
        external_candidates=external_candidates,
    )
    scope = TraversalScope.create()
    verification = build_verification_input(
        frozen, discovery_ref="frozen-discovery", context=context, card=card, scope=scope,
        entity_dependencies=dependencies, external_candidates=external_candidates,
        bridge_dependencies=[],
        scope_resolutions=[],
    )
    responses = []
    for target in verification.targets:
        support = (target.payload.mentions if target.target_kind == "entity" else
                   target.payload.support if target.target_kind == "reference_binding" else
                   target.payload.identity_support if target.target_kind == "external_link" else
                   target.payload.source_assertion.predicate_support)
        responses.append(dict(
            target_id=target.target_id, content_hash=target.content_hash,
            facets=[dict(
                name=name, verdict=(binding_verdict if target.target_kind == "reference_binding"
                                    else predicate_verdict if name == "predicate" else "supported"),
                support=[item.model_dump() for item in support], counterevidence_support=[],
                reason="独立核验当前冻结对象和原文角色。",
            ) for name in target.required_facets],
        ))
    verified = validate_verification(
        VerificationEnvelope.model_validate(dict(verifications=responses)),
        targets=verification.targets, context=context,
    )
    result = finalize_claims(
        frozen, verified, scope=scope, context=context, card=card, verification_input=verification,
        registry=document.registry, deterministic_results={
            "relation": ClaimCheckResult(checks={"binding": True, "relation_graph": True}),
        }, entity_proofs=proofs, reference_resolution=True, known_resolutions=known_resolutions,
    )
    return NS(result=result, frozen=frozen, verification=verification, context=context)


def test_three_record_mentions_bind_to_one_entity_with_original_sources(reference_document):
    first = run_pipeline(reference_document, 0)
    yield_row = run_pipeline(reference_document, 1, priors=[first])
    storage_row = run_pipeline(reference_document, 2, priors=[first])
    canonical = first.result.nodes[0].entity_id
    assert first.result.complete and yield_row.result.complete and storage_row.result.complete
    assert yield_row.result.nodes == [] and storage_row.result.nodes == []
    resolutions = [item.result.reference_resolutions[0] for item in (first, yield_row, storage_row)]
    assert {resolution["entity_ref"]["id"] for resolution in resolutions} == {canonical}
    assert len({resolution["source_refs"][0]["evidence_id"] for resolution in resolutions}) == 3
    assert all(resolution["binding_ref"] for resolution in resolutions[1:])
    assert len(reference_document.registry.mentions) == 3
    assert {mention.text for mention in reference_document.registry.mentions} == {
        "HRS-9267 粗品", "HRS-9267粗品",
    }


def test_scoped_tasks_can_recall_same_physical_entity_without_name_only_scope_leak(
    reference_document,
):
    from app.services.extraction.ontology_guided.reference_context import select_reference_entities

    first = run_pipeline(reference_document, 0)
    node = first.result.nodes[0]
    qualified = TraversalScope.create([ScopeMember(
        relation_ref=VersionedRef(id="conditional-relation", revision=1),
        member_ref=VersionedRef(id=node.entity_id, revision=node.revision),
    )])
    origins = {node.entity_id: {"scope": TraversalScope.create().model_dump(mode="json")}}
    for position, expected in ((0, [node.entity_id]), (1, [])):
        task = NS(record_id=reference_document.entries[position].record.record_id, scope=qualified)
        assert select_reference_entities(
            task, reference_document.index, {node.entity_id: node}, origins, {}, {node.class_iri},
        ) == expected


@pytest.mark.parametrize("verdict", ["unsupported", "undetermined"])
def test_unproven_reference_keeps_independent_entity_and_blocks_dependent_relation(
    reference_document, verdict,
):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 1, priors=[first],
                          binding_verdict=verdict, relation=True)
    assert len(second.result.nodes) == 1
    assert second.result.nodes[0].entity_id != first.result.nodes[0].entity_id
    assert second.result.reference_resolutions[0]["binding_ref"] is None
    assert not second.result.edges


def test_successful_binding_does_not_supply_missing_predicate_entailment(reference_document):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 1, priors=[first], relation=True,
                          predicate_verdict="unsupported")
    resolved = second.result.reference_resolutions[0]["entity_ref"]["id"]
    assert resolved == first.result.nodes[0].entity_id
    assert second.result.reference_resolutions[0]["binding_ref"]
    assert not second.result.edges


def test_accepted_relation_keeps_canonical_endpoint_and_binding_dependency(reference_document):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 1, priors=[first], relation=True)
    assert second.result.complete, second.result.reason
    edge = second.result.edges[0]
    assert edge.object_ref.id == first.result.nodes[0].entity_id
    binding_ref = VersionedRef.model_validate(second.result.reference_resolutions[0]["binding_ref"])
    assert binding_ref in edge.dependency_refs
    assert edge.object_evidence_refs and edge.proof_ref


def test_same_physical_mention_has_stable_entity_across_frozen_generations(reference_document):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 0, generation=2)
    assert first.frozen.local_ref_map["current"] != second.frozen.local_ref_map["current"]
    assert first.result.nodes[0].entity_id == second.result.nodes[0].entity_id
    assert first.result.reference_resolutions[0]["entity_ref"] == (
        second.result.reference_resolutions[0]["entity_ref"]
    )


def test_whitespace_and_same_name_only_do_not_merge_different_mentions(reference_document):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 1, priors=[first], bind=False)
    third = run_pipeline(reference_document, 2, priors=[second], bind=False)
    assert len({result.result.nodes[0].entity_id for result in (first, second, third)}) == 3
    assert not second.result.reference_resolutions[0]["binding_ref"]
    assert not third.result.reference_resolutions[0]["binding_ref"]


def test_conflicting_batch_and_finished_product_do_not_merge(reference_document):
    first = run_pipeline(reference_document, 0)
    other_batch = run_pipeline(reference_document, 3, priors=[first], binding_verdict="unsupported")
    finished = run_pipeline(reference_document, 4, priors=[first], class_iri="urn:FinishedProduct")
    assert len({result.result.nodes[0].entity_id for result in (first, other_batch, finished)}) == 3
    assert not other_batch.result.reference_resolutions[0]["binding_ref"]
    assert "reference_binding_class_mismatch" in finished.frozen.claim_issues["binding-0"]


def test_ambiguous_reference_candidates_are_retained_without_choosing_a_target(reference_document):
    first = run_pipeline(reference_document, 0)
    other = run_pipeline(reference_document, 3)
    current = run_pipeline(reference_document, 1, priors=[first, other])
    assert set(current.frozen.claim_issues) == {"binding-0", "binding-1"}
    assert not current.result.reference_resolutions[0]["binding_ref"]
    assert len(current.result.nodes) == 1
    assert current.result.nodes[0].entity_id not in {
        first.result.nodes[0].entity_id, other.result.nodes[0].entity_id,
    }


@pytest.mark.parametrize("same_scope", [True, False])
def test_saved_reference_mapping_reuse_requires_verified_scope(reference_document, same_scope):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 1, priors=[first])
    saved = copy.deepcopy(second.result.reference_resolutions)
    if not same_scope:
        saved[0]["scope"]["scope_id"] = "foreign-scope"
    repeated = run_pipeline(reference_document, 1, generation=2, priors=[first], bind=False,
                            known_resolutions=saved)
    resolution = repeated.result.reference_resolutions[0]
    if same_scope:
        assert resolution["entity_ref"]["id"] == first.result.nodes[0].entity_id
        assert resolution["binding_ref"] == saved[0]["binding_ref"]
        assert repeated.result.nodes == []
    else:
        assert resolution["entity_ref"]["id"] != first.result.nodes[0].entity_id
        assert resolution["binding_ref"] is None
        assert saved[0]["binding_ref"] not in resolution["dependency_refs"]


def test_relation_cannot_redirect_through_undeclared_same_round_binding(reference_document):
    first = run_pipeline(reference_document, 0)
    second = run_pipeline(reference_document, 1, priors=[first], relation=True,
                          declare_relation_bindings=False)
    assert second.result.reference_resolutions[0]["entity_ref"]["id"] == (
        first.result.nodes[0].entity_id
    )
    assert not second.result.complete
    assert not second.result.edges
    assert "reference_binding_dependency_missing" in second.result.reason


@pytest.mark.parametrize("existing", [False, True])
def test_external_identity_enrichment_survives_canonical_projection(reference_document, existing):
    first = run_pipeline(reference_document, 0) if existing else None
    current = run_pipeline(reference_document, 1, priors=[first] if first else [], bind=False,
                           external_identity=True,
                           external_subject="existing" if existing else "current")
    assert current.result.complete, current.result.reason
    expected = first.result.nodes[0].entity_id if first else (
        current.result.reference_resolutions[0]["entity_ref"]["id"]
    )
    enriched = next((node for node in current.result.nodes if node.entity_id == expected), None)
    assert enriched is not None, "accepted external identity must return the enriched node"
    assert enriched.identity_status == "verified"
    assert enriched.external_provenance[0].record_key == "CP-A001"
    assert len(enriched.identity_decision_refs) == 3
    if existing:
        assert first.result.nodes[0].identity_status == "document_local"


def test_external_identity_on_new_coreferent_preserves_canonical_node_fields(reference_document):
    first = run_pipeline(reference_document, 0)
    current = run_pipeline(reference_document, 1, priors=[first], external_identity=True)
    assert current.result.complete, current.result.reason
    enriched = next((node for node in current.result.nodes
                     if node.entity_id == first.result.nodes[0].entity_id), None)
    assert enriched is not None, "binding must not silently discard accepted external identity"
    fields = {"identity_status", "identity_decision_refs", "external_provenance"}
    assert enriched.model_dump(exclude=fields) == first.result.nodes[0].model_dump(exclude=fields)
    assert enriched.identity_status == "verified"
    assert enriched.external_provenance[0].record_key == "CP-A001"


def test_duplicate_physical_claims_preserve_external_identity_from_either_claim(reference_document):
    current = run_pipeline(reference_document, 0, duplicate_entity=True,
                           external_identity=True, external_subject="duplicate")
    assert current.result.complete, current.result.reason
    assert len(current.result.nodes) == 1
    node = current.result.nodes[0]
    assert node.identity_status == "verified"
    assert node.external_provenance[0].record_key == "CP-A001"


@pytest.mark.parametrize("prior_alias", [False, True])
def test_new_binding_cannot_override_canonical_identity(reference_document, prior_alias):
    first = run_pipeline(reference_document, 0)
    other = run_pipeline(reference_document, 3)
    known = (run_pipeline(reference_document, 1, priors=[first]).result.reference_resolutions
             if prior_alias else None)
    current = run_pipeline(reference_document, 1 if prior_alias else 0, generation=2,
                           priors=[first, other], binding_prior_indices=[1],
                           known_resolutions=known)
    assert not current.result.complete
    assert "reference_binding_conflicts_with_registered_identity" in current.result.reason
    assert current.result.reference_resolutions[0]["entity_ref"]["id"] == (
        first.result.nodes[0].entity_id
    )
    assert current.result.reference_resolutions[0]["binding_ref"] == (
        known[0]["binding_ref"] if prior_alias else None
    )
