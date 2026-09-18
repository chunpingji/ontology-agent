"""Rebuild confirmed authorization against a real parsed, frozen Word IR."""

from __future__ import annotations

from dataclasses import FrozenInstanceError

import pytest
from docx import Document
from pydantic import ValidationError

from app.services.extraction.document_ir import DocumentIR
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided import context as module
from app.services.extraction.ontology_guided.claim_protocol import (
    EntityDependencyView,
    ExtractionProfile,
    SchemaCard,
)
from app.services.extraction.ontology_guided.context import (
    AuthorizedFragment,
    ContextAuthorization,
    ContextBindingRefs,
    ContextFeedback,
    IRIdentity,
    SourceCatalogEntry,
    assemble_context,
    authorization_policy_hash,
    build_authorized_context,
    restore_authorized_context,
)
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    GraphNode,
    SlotSpec,
    SubjectRef,
    TraversalScope,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def authorized_context(tmp_path):
    document = Document()
    document.add_heading("记录", level=1)
    table = document.add_table(rows=3, cols=3)
    for column, label in enumerate(("名称", "数量", "单位")):
        table.cell(0, column).text = label
    for column, value in enumerate(("主体甲", "5", "mg")):
        table.cell(1, column).text = value
    for column, value in enumerate(("主体乙", "6", "mg")):
        table.cell(2, column).text = value
    document.add_heading("条件", level=1)
    document.add_paragraph("条件C不满足时，上述数值不得使用。")
    path = tmp_path / "source.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    record = next(value for value in index.records if value.kind == "table_row")
    supplemental = next(value for value in index.records if "条件C不满足" in value.text)
    owner = next(unit for unit in record.source_units if unit.text == "主体甲")
    owner_ref = index.ir.anchor(owner.evidence_id, 0, len(owner.text))
    counter = supplemental.source_units[0]
    counter_ref = index.ir.anchor(counter.evidence_id, 0, len(counter.text))
    subject = SubjectRef(
        entity_id="root", revision=1, class_iri="urn:Record", is_document_root=True
    )
    node = GraphNode(
        entity_id=subject.entity_id, revision=1, class_iri=subject.class_iri,
        class_label="记录", label="主体甲", root=True, root_origin="user_specified",
        grounding_kind="document_root", evidence_refs=[owner_ref],
    )
    task = RecognitionTask.create(
        subject=subject, predicate_iri="urn:quantity", predicate_kind="property",
        record_id=record.record_id, phase=1, hop=0, dependency_hash="frozen-dependency",
    )
    scope_hash = evidence_hash([
        record.record_id, [unit.evidence_id for unit in record.source_units],
        subject.model_dump(mode="json"), task.predicate_iri,
    ])
    run_fingerprint = "frozen-run-1"
    seed = VerificationTarget.create(
        run_fingerprint=run_fingerprint,
        claim_ref=VersionedRef(id=task.claim_lineage_id, revision=1), task_id=task.task_id,
        check_kind="predicate_entailment",
        document_context=DocumentContext(
            document_hash=index.ir.document_hash, document_class_iri=subject.class_iri,
            root_ref=VersionedRef(id=subject.entity_id, revision=subject.revision),
        ),
        subject_ref=subject, predicate_iri=task.predicate_iri,
        ontology_hash=evidence_hash("frozen-ontology"), source_scope_hash=scope_hash,
        context_hash=evidence_hash([
            index.source_text(record.record_id, include_context=True), scope_hash,
        ]),
    )
    predicate = SlotSpec(
        iri=task.predicate_iri, label="数量", declared_by=[subject.class_iri],
        datatype_iris=["http://www.w3.org/2001/XMLSchema#decimal"],
    )
    card = SchemaCard(
        schema_card_id="card-1", subject_ref=VersionedRef(id="root", revision=1),
        class_iris=[subject.class_iri], ontology_snapshot_id="ontology-1", menu_id="menu-1",
        predicates=[predicate], quantity_policies=[], identity_keys=[], unsupported_constraints=[],
    )
    base = assemble_context(
        seed, record.record_id, index, subject_evidence_refs=[owner_ref],
        subject_label=node.label, predicate=predicate, repair_enabled=True,
    )
    dependency = {
        "entity_ref": VersionedRef(id="root", revision=1), "class_iri": subject.class_iri,
        "grounding_kind": "document_root", "root_origin": "user_specified", "proposal": None,
        "source_refs": [owner_ref], "dependency_refs": [],
    }
    dependencies = [EntityDependencyView(**dependency, content_hash=evidence_hash(dependency))]
    trusted = dict(
        target_seed=seed, run_fingerprint=run_fingerprint, card=card, profile=ExtractionProfile(),
        entity_dependencies=dependencies, scope=TraversalScope.create(), subject_node=node,
    )
    policy = authorization_policy_hash(
        task, base, **{name: value for name, value in trusted.items() if name != "run_fingerprint"}
    )
    refs = {name: list(getattr(base, name)) for name in ContextBindingRefs.model_fields}
    refs["counterevidence_refs"].append(counter_ref)
    authorization = ContextAuthorization(
        task_id=task.task_id,
        ir_identity=IRIdentity(
            document_hash=index.ir.document_hash, parser_version=index.ir.parser_version,
            structure_hash=index.ir.structure_hash,
        ),
        context_policy_hash=policy, record_ids=[record.record_id, supplemental.record_id],
        fragments=[
            AuthorizedFragment(
                record_id=record.record_id, evidence_id=value.anchor.evidence_id,
                span_start=value.anchor.span_start, span_end=value.anchor.span_end,
                role=value.purpose, fact_eligible=value.fact_eligible,
            ) for value in base.fragments
        ] + [AuthorizedFragment(
            record_id=supplemental.record_id, evidence_id=counter.evidence_id,
            span_start=0, span_end=len(counter.text), role="counterevidence", fact_eligible=False,
        )],
        bindings=ContextBindingRefs(**refs),
    )
    assembled, source_hash = build_authorized_context(
        task, base, authorization=authorization, index=index, evidence_revision=2, **trusted
    )
    return dict(
        task=task, base=base, authorization=authorization, index=index, trusted=trusted,
        assembled=assembled, evidence_hash=source_hash, path=path, counter_ref=counter_ref,
    )


def _restore(fixture, **changes):
    arguments = dict(
        task=fixture["task"], base=fixture["base"], authorization=fixture["authorization"],
        index=fixture["index"], expected_evidence_revision=2,
        expected_evidence_hash=fixture["evidence_hash"],
        expected_context_hash=fixture["assembled"].context_hash, **fixture["trusted"],
    )
    arguments.update(changes)
    return restore_authorized_context(**arguments)


def test_heading_and_shared_header_can_be_read_and_restored_without_invented_owner(
    authorized_context,
):
    value = authorized_context
    index, task, trusted = value["index"], value["task"], value["trusted"]
    heading = next(unit for unit in index.ir.evidence_units if unit.text == "记录")
    anchor = index.ir.anchor(heading.evidence_id, 0, len(heading.text))
    base = assemble_context(
        trusted["target_seed"], task.record_id, index, subject_evidence_refs=[anchor],
        subject_label=trusted["subject_node"].label,
        predicate=trusted["card"].predicates[0], repair_enabled=True,
    )
    shared = next(fragment for fragment in base.fragments
                  if fragment.purpose == "record_heading_or_header")
    assert module.fragment_record_id(shared, index, task.record_id) == task.record_id
    assert module.fragment_record_id(shared, index, "not-an-owner") is None
    with pytest.raises(ValueError, match="evidence_record_ambiguous"):
        module.fragment_record_id(
            shared.model_copy(update={"fact_eligible": True}), index, "not-an-owner",
        )
    authorization = module.build_retrieval_authorization(
        task, base, [], index=index,
        **{key: item for key, item in trusted.items() if key != "run_fingerprint"},
    )
    unassigned = [f for f in authorization.fragments if f.record_id is None]
    assert any(f.evidence_id == heading.evidence_id for f in unassigned)
    assert all(not f.fact_eligible for f in unassigned)
    assembled, source_hash = build_authorized_context(
        task, base, authorization=authorization, index=index, evidence_revision=2, **trusted,
    )
    restored = restore_authorized_context(
        task, base, authorization=authorization, index=index, expected_evidence_revision=2,
        expected_evidence_hash=source_hash, expected_context_hash=assembled.context_hash,
        **trusted,
    )
    assert restored.fragments == assembled.fragments
    with pytest.raises(ValidationError, match="cannot grant fact eligibility"):
        unassigned[0].fact_eligible = True
    with pytest.raises(ValidationError, match="cannot grant fact eligibility"):
        build_authorized_context(
            task, base, authorization=authorization, index=index, evidence_revision=2, **trusted,
        )


def test_cold_restore_reads_source_and_field_bindings_from_frozen_ir(authorized_context):
    value = authorized_context
    value["path"].unlink()
    cold_ir = DocumentIR.model_validate_json(value["index"].ir.model_dump_json())
    cold_auth = ContextAuthorization.model_validate_json(value["authorization"].model_dump_json())
    restored = _restore(value, index=RecordIndex(cold_ir), authorization=cold_auth)
    assert [fragment.text for fragment in restored.fragments] == [
        cold_ir.resolve(fragment.anchor) for fragment in restored.fragments
    ]
    counter = next(
        fragment for fragment in restored.fragments if fragment.purpose == "counterevidence"
    )
    assert counter.text == "条件C不满足时，上述数值不得使用。"
    assert counter.fact_eligible is False
    assert restored.counterevidence_refs == [value["counter_ref"]]
    assert len(restored.field_bindings) == 3
    assert restored.field_bindings == value["base"].field_bindings
    assert restored.subject_label == value["trusted"]["subject_node"].label
    assert restored.context_hash == value["assembled"].context_hash
    assert restored.target == value["assembled"].target
    assert restored.target.target_id != value["base"].target.target_id
    assert restored.token_count is None
    assert restored.protocol_state == {}


def test_restore_does_not_copy_untrusted_text_labels_or_field_bindings(authorized_context):
    base = authorized_context["base"].model_copy(deep=True)
    base.fragments[0].text = "模型伪造的表头"
    base.subject_label = "错误主体"
    base.field_bindings = []
    base.protocol_state = {"active_input_items": [{"evidence_id": "invented"}]}
    restored = _restore(authorized_context, base=base)
    assert restored.subject_label == "主体甲"
    assert restored.field_bindings
    assert "模型伪造" not in "".join(fragment.text for fragment in restored.fragments)
    assert all(fragment.anchor.evidence_id != "invented" for fragment in restored.fragments)


def test_existing_assembly_hash_payload_remains_unchanged(authorized_context):
    base = authorized_context["base"]
    payload = {
        "target": authorized_context["trusted"]["target_seed"].model_dump(mode="json"),
        "record_id": base.record_id,
        "fragments": [item.model_dump(mode="json") for item in base.fragments],
        **{name: getattr(base, name) for name in ContextBindingRefs.model_fields},
        "subject_label": base.subject_label,
        "context_records_version": module.CONTEXT_RECORDS_VERSION,
        "field_bindings": [item.model_dump(mode="json") for item in base.field_bindings],
    }
    assert base.context_hash == evidence_hash(payload)


@pytest.mark.parametrize("field", ["expected_evidence_hash", "expected_context_hash"])
def test_unconfirmed_hash_cannot_authorize_context(authorized_context, field):
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(authorized_context, **{field: "0" * 64})


@pytest.mark.parametrize("revision", [1, 0, True])
def test_evidence_revision_belongs_to_outer_confirmed_state(authorized_context, revision):
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(authorized_context, expected_evidence_revision=revision)


@pytest.mark.parametrize("mutation", ["fact", "role", "record", "span", "omit_counter"])
def test_source_permissions_and_membership_cannot_change_on_restore(authorized_context, mutation):
    authorization = authorized_context["authorization"].model_copy(deep=True)
    supplement = authorization.fragments[-1]
    if mutation == "fact":
        supplement.fact_eligible = True
    elif mutation == "role":
        supplement.role = "target"
    elif mutation == "record":
        supplement.record_id = authorization.record_ids[0]
    elif mutation == "span":
        supplement.span_end -= 1
    else:
        authorization.fragments.pop()
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(authorized_context, authorization=authorization)


def test_bound_target_cannot_replace_original_seed(authorized_context):
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(authorized_context, target_seed=authorized_context["base"].target)
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(authorized_context, run_fingerprint="another-run")


@pytest.mark.parametrize("change", ["task", "node_revision", "node_label", "class", "card"])
def test_exact_subject_task_and_schema_dependencies_are_required(authorized_context, change):
    value = authorized_context
    arguments = {}
    if change == "task":
        arguments["task"] = value["task"].model_copy(update={"task_id": "another-task"})
    elif change == "card":
        arguments["card"] = value["trusted"]["card"].model_copy(update={"menu_id": "changed-menu"})
    else:
        field, replacement = {
            "node_revision": ("revision", 2), "node_label": ("label", "主体乙"),
            "class": ("class_iri", "urn:Other"),
        }[change]
        arguments["subject_node"] = value["trusted"]["subject_node"].model_copy(
            update={field: replacement}
        )
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(value, **arguments)


def test_stale_ir_and_assembly_version_cannot_restore(authorized_context, monkeypatch):
    value = authorized_context
    stale = value["authorization"].model_copy(deep=True)
    stale.ir_identity.structure_hash = "0" * 64
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(value, authorization=stale)
    monkeypatch.setattr(module, "FIELD_BINDING_VERSION", "changed-assembly")
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(value)


def test_changed_frozen_source_is_detected_even_when_ids_are_unchanged(authorized_context):
    value = authorized_context
    value["index"].ir.unit(value["counter_ref"].evidence_id).text = "条件C满足时允许使用。"
    with pytest.raises(ValueError, match="context_authorization_mismatch"):
        _restore(value)


def test_authorization_does_not_store_duplicate_hashes_or_source_text(authorized_context):
    payload = authorized_context["authorization"].model_dump(mode="json")
    assert not {"evidence_revision", "evidence_hash", "context_hash"}.intersection(payload)
    assert all("text" not in fragment for fragment in payload["fragments"])
    for key, value in (("evidence_revision", 2), ("context_hash", "0" * 64)):
        with pytest.raises(ValidationError):
            ContextAuthorization.model_validate({**payload, key: value}, strict=True)


def test_display_metadata_and_feedback_are_frozen_derived_values():
    entry = SourceCatalogEntry("record-1", None, "摘要只是定位", "metadata", ["ev-1"])
    feedback = ContextFeedback(None, None, "missing_evidence", None, "读取允许的原文。", [])
    with pytest.raises(FrozenInstanceError):
        entry.summary = "新的事实"
    with pytest.raises(FrozenInstanceError):
        feedback.code = "accepted"
