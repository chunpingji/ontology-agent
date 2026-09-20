"""A shared model input never turns one physical row's values into another's."""

import json
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.claim_protocol import ExtractionProfile
from app.services.extraction.ontology_guided.context import assemble_record_discovery_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
    VersionedRef,
)
from app.services.extraction.ontology_guided.record_discovery import (
    RECORD_PIPELINE,
    RecordDiscoveryTask,
    compile_record_schema_card,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.table_reading import table_reading_groups
from app.services.extraction.ontology_guided.tool_model_adapter import ToolModelRecognitionAdapter
from app.services.extraction.ontology_guided.tool_runtime import ToolLimits
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.local_client import ResponseTurn


@pytest.mark.parametrize("cross_row", [False, True])
def test_one_batch_materializes_multiple_owners_and_rejects_cross_row_value(
    tmp_path, monkeypatch, cross_row,
):
    document = Document()
    document.add_heading("装置清单", 1)
    table = document.add_table(rows=3, cols=3)
    for row, values in enumerate([
        ("装置", "型号", "状态"), ("装置甲", "型号甲", "运行中"),
        ("装置乙", "型号乙", "已停机"),
    ]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    path = tmp_path / "batch.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    sources, = table_reading_groups(index, max_group_chars=1200, max_records=8)
    assert len(sources) == 2
    ontology = OntologySnapshot(
        snapshot_id="batch", ontology_hash=evidence_hash("batch"),
        classes={"urn:Device": OntologyClassDefinition(
            iri="urn:Device", label="装置", source_hash="batch",
            declared_properties=[SlotSpec(
                iri=f"urn:{name}", label=label, declared_by=["urn:Device"],
                datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
            ) for name, label in (("model", "型号"), ("status", "状态"))],
        )},
    )
    card = compile_record_schema_card(
        ontology, class_iris=["urn:Device"], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    task = RecordDiscoveryTask.create(
        run_fingerprint="batch", record_id=sources[0], source_record_ids=sources,
        schema_card_id=card.schema_card_id, analysis_scope_ref=card.analysis_scope_ref,
        dependency_hash=evidence_hash(sources),
    )
    context = assemble_record_discovery_context(
        task, card, index, ontology_hash=ontology.ontology_hash,
        document_context=DocumentContext(
            document_hash=index.ir.document_hash, document_class_iri="urn:Document",
            root_ref=VersionedRef(id="root", revision=1),
        ),
    )
    context.tool_inputs = {
        "target_seed": context.target.model_dump(mode="json"),
        "schema_card": card.model_dump(mode="json"), "run_fingerprint": "batch",
        "entity_dependencies": [], "entity_nodes": [], "reference_resolutions": [],
    }
    context.remaining_model_calls = 4
    context.bind_protocol_hook(lambda state: None)
    context.bind_model_call_hook(lambda stage, ordinal: None)

    def quote(text):
        unit = next(unit for unit in index.ir.evidence_units if unit.text == text)
        return {"evidence_id": unit.evidence_id, "text": text, "context_text": None}

    proposal = {name: [] for name in (
        "entities", "properties", "relations", "external_links", "observations",
        "reference_bindings",
    )}
    for identity, name, model, status in (("a", "装置甲", "型号甲", "运行中"),
                                          ("b", "装置乙", "型号乙", "已停机")):
        proposal["entities"].append({
            "local_id": identity, "class_iri": "urn:Device", "representation": "mention",
            "mentions": [quote(name)], "record_components": [], "identifier_claims": [],
        })
        for predicate, label, value in (("model", "型号", model), ("status", "状态", status)):
            proposal["properties"].append({
                "local_id": f"{identity}-{predicate}", "subject_id": identity,
                "predicate_iri": f"urn:{predicate}",
                "value_quote": quote("型号乙" if cross_row and identity == "a"
                                     and predicate == "model" else value),
                "field_support": [quote(label)], "unit_support": [],
                "qualifiers": {"polarity": "affirmed", "modality": "asserted",
                               "condition_support": [], "scope_qualifiers": []},
                "bridge_kind": "role_mapped_table", "bridge_ref_ids": [],
            })
    requests = []

    def transport(_client, **kwargs):
        view = json.loads(kwargs["input_items"][0]["content"][0]["text"])
        requests.append(view)
        answer = deepcopy(proposal) if view["stage"] == "discovery" else {
            "verifications": [{
                "target_id": target["target_id"], "content_hash": target["content_hash"],
                "facets": [{
                    "name": facet, "verdict": "supported",
                    "support": [quote(unit["text"]) for unit in view["evidence_units"]],
                    "counterevidence_support": [], "reason": "各物理行及字段原文可核验。",
                } for facet in target["required_facets"]],
            } for target in view["verification_input"]["targets"]],
        }
        return ResponseTurn(
            response_id=f"batch-{len(requests)}", response_status="completed",
            output_items=[{"type": "message", "role": "assistant", "status": "completed",
                           "content": [{"type": "output_text", "text": json.dumps(answer)}]}],
            incomplete_details=None, error=None, usage={"input_tokens": 10, "output_tokens": 5},
        )

    monkeypatch.setattr("app.services.llm.local_client.responses_create", transport)
    adapter = ToolModelRecognitionAdapter(
        object(), index=index, ontology=ontology, metadata=SimpleNamespace(node_summaries=[]),
        profile=ExtractionProfile(), token_counter=len, model_identity="controlled-table-model",
        max_input_tokens=1000000, max_output_tokens=20000, tool_limits=ToolLimits(20000),
        recognition_pipeline=RECORD_PIPELINE, reference_resolution=True,
    )
    outcome = adapter.inspect_record(task, context, card)
    assert len(requests) == outcome.model_calls == 2
    assert {node.label for node in outcome.nodes} == {"装置甲", "装置乙"}
    labels = {node.entity_id: node.label for node in outcome.nodes}
    actual = {(labels[prop.subject_ref.id], prop.predicate_iri, prop.raw_value)
              for prop in outcome.properties}
    expected = {("装置甲", "urn:status", "运行中"), ("装置乙", "urn:status", "已停机"),
                ("装置乙", "urn:model", "型号乙")}
    if not cross_row:
        expected.add(("装置甲", "urn:model", "型号甲"))
        assert outcome.complete, outcome.reason
    else:
        assert not outcome.complete
        assert "owner_row_mismatch" in outcome.reason
    assert actual == expected, outcome.reason
