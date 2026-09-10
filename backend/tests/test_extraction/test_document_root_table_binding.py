"""A document title can support a root binding without becoming a table row."""

import pytest
from docx import Document

from app.schemas.evidence import BindingEvidence, Candidate, DocumentProvenance
from app.services.extraction.evidence_scope import build_scope, candidate_ref
from app.services.extraction.extraction_tasks import GenericExtractionRunner, binding_record_mapping
from app.services.extraction.semantic_binding import validate_document_candidate
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def record(tmp_path):
    doc = Document()
    doc.add_paragraph("产品生产报告")
    doc.add_paragraph("另一个工艺的备注")
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("设备", "清洁方法"), ("釜甲", "冲洗五分钟"),
                                      ("釜乙", "冲洗十分钟")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    path = tmp_path / "records.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    schema = {"urn:Report": {"relationships": [{"iri": "urn:cleaning", "range": ["urn:Clean"]}]},
              "urn:Clean": {}}
    root = GenericExtractionRunner(schema, None, None, model_identity="fixture").run(
        ir, effective_class="urn:Report",
    ).candidates[0]
    anchors = {u.text: ir.anchor(u.evidence_id) for u in ir.evidence_units if u.text}
    obj = Candidate(candidate_id="clean", kind="entity", class_iri="urn:Clean", text="冲洗五分钟",
                    validation_status="passed", provenance=[DocumentProvenance(
                        anchors=[anchors["冲洗五分钟"]], excerpts=["冲洗五分钟"],
                    )])
    return ir, root, obj, anchors


def validate(record, binding_anchors, *, supported=True):
    ir, root, obj, anchors = record
    mapping = {"table_path": anchors["冲洗五分钟"].table_path, "row_index": 1}
    candidate = Candidate(
        candidate_id="edge", kind="relationship", predicate_iri="urn:cleaning",
        subject=candidate_ref(root), object=candidate_ref(obj), scope=build_scope(ir, root),
        provenance=[DocumentProvenance(anchors=[anchors["冲洗五分钟"]], excerpts=["冲洗五分钟"])],
        bindings=[BindingEvidence(
            method="table_record", subject_candidate_id=root.candidate_id,
            object_candidate_id=obj.candidate_id, predicate_iri="urn:cleaning",
            anchors=binding_anchors, record_mapping=mapping,
        )],
    )
    return validate_document_candidate(
        candidate, ir, {root.candidate_id: root, obj.candidate_id: obj},
        semantic_supported=supported, allowed_predicates={"urn:cleaning"},
    )


def test_title_and_exact_row_pass_without_dropping_title_from_proof(record):
    ir, root, obj, anchors = record
    binding = [anchors[k] for k in ("产品生产报告", "清洁方法", "冲洗五分钟")]
    mapping = binding_record_mapping("table_record", binding, {}, ir=ir,
                                     fact_anchors=[anchors["冲洗五分钟"]], subject=root)
    assert mapping == {"table_path": anchors["冲洗五分钟"].table_path, "row_index": 1}
    checked = validate(record, binding)
    assert checked.positive_eligible and checked.bindings[0].anchors == binding
    denied = validate(record, binding, supported=False)
    assert not denied.positive_eligible
    assert "cooccurrence_only" in {i.code for i in denied.validation_issues}


@pytest.mark.parametrize("extra", ["另一个工艺的备注", "冲洗十分钟"])
def test_other_prose_and_other_rows_cannot_be_ignored(record, extra):
    ir, root, obj, anchors = record
    binding = [anchors[k] for k in ("产品生产报告", "冲洗五分钟", extra)]
    assert binding_record_mapping("table_record", binding, {}, ir=ir,
                                  fact_anchors=[anchors["冲洗五分钟"]], subject=root) == {}
    checked = validate(record, binding)
    assert "record_mapping_mismatch" in {i.code for i in checked.validation_issues}
    assert not checked.positive_eligible


def test_root_title_cannot_replace_table_fact_or_be_supplied_by_ordinary_entity(record):
    ir, root, obj, anchors = record
    binding = [anchors[k] for k in ("产品生产报告", "冲洗五分钟")]
    assert binding_record_mapping("table_record", binding, {}, ir=ir,
                                  fact_anchors=[anchors["产品生产报告"]], subject=root) == {}
    ordinary = root.model_copy(update={"extractor_version": "other-extractor"})
    assert binding_record_mapping("table_record", binding, {}, ir=ir,
                                  fact_anchors=[anchors["冲洗五分钟"]], subject=ordinary) == {}
