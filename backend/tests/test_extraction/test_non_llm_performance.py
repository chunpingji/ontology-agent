"""Performance invariants on real execution paths, with bounded scripted models."""

import json
from copy import deepcopy

import pytest
from docx import Document
from sqlalchemy import event, select

from app.models.evidence import EvidenceCandidateRecord, EvidenceCandidateRevision, EvidenceReview
from app.models.extraction import ExtractionJob
from app.schemas.evidence import ExtractionTask, TaskBudget
from app.services.extraction.candidate_store import CandidateStore
from app.services.extraction.checkpoint_journal import (
    CheckpointJournal,
    journal_path,
    read_checkpoint,
    write_snapshot,
)
from app.services.extraction.evidence_scope import build_scope, candidate_ref
from app.services.extraction.extraction_tasks import AssertionResponse, GenericExtractionRunner
from app.services.extraction.hierarchical_context import split_windows
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_extraction_tasks import accept_fixture_types
from tests.test_extraction.test_fact_commit import entity
from tests.test_extraction.test_hierarchical_context import TestTokenizer


def test_unit_index_tracks_replacement_reordering_deepcopy_and_rejects_mutated_run(tmp_path):
    doc = Document()
    doc.add_paragraph("first")
    doc.add_paragraph("second")
    path = tmp_path / "index.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    first, second = ir.evidence_units
    assert ir.unit(first.evidence_id) is first
    copied = ir.model_copy(deep=True)
    assert copied.unit(first.evidence_id) is copied.evidence_units[0]
    assert copied.unit(first.evidence_id) is not first
    ir.evidence_units.reverse()
    assert ir.unit(first.evidence_id) is first
    replacement = first.model_copy(update={"text": "edited"})
    ir.evidence_units[1] = replacement
    assert ir.unit(first.evidence_id) is replacement
    with pytest.raises(ValueError, match="structure identity"):
        ir.evidence_units = [second]
    with pytest.raises(ValueError, match="unknown evidence"):
        ir.unit(first.evidence_id)
    runner = GenericExtractionRunner({}, TestTokenizer(), None, model_identity="test")
    with pytest.raises(ValueError, match="structure identity"):
        runner.run(ir)


def test_preflight_reuses_envelope_but_mutated_source_and_candidate_invalidate(
    subject_document, monkeypatch
):
    from app.services.extraction import extraction_tasks as module

    ir, (a, b) = subject_document
    scope = build_scope(ir, a, [b])
    task = ExtractionTask(
        task_id="prepared",
        task_kind="property",
        predicate_iri="urn:strength",
        subject=candidate_ref(a),
        scope=scope,
        target_evidence_ids=[ir.evidence_units[1].evidence_id],
        budget=TaskBudget(max_input_tokens=60000),
    )
    runner = GenericExtractionRunner(
        {a.class_iri: {}},
        TestTokenizer(),
        lambda *args: {"assertions": []},
        model_identity="test",
        budget=task.budget,
    )
    original, calls = module.build_context, []

    def build(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(module, "build_context", build)
    candidates = {a.candidate_id: a}
    prepared = list(runner.fit_assertion_task(task, ir, candidates))
    runner.execute_task(prepared[0], ir, candidates)
    assert len(calls) == 1
    a.text += " amended"
    runner._context(task, ir, candidates, "", AssertionResponse)
    assert len(calls) == 2
    ir.title += " amended"
    runner._context(task, ir, candidates, "", AssertionResponse)
    assert len(calls) == 3


def test_whole_menu_and_short_window_use_one_exact_count():
    class Counter(TestTokenizer):
        def __init__(self):
            self.calls = []

        def count(self, text):
            self.calls.append(text)
            return super().count(text)

    counter = Counter()
    schema = {f"urn:Class{i}": {"label": str(i)} for i in range(136)}
    runner = GenericExtractionRunner(
        schema, counter, None, model_identity="test", budget=TaskBudget(max_input_tokens=60000)
    )
    assert list(runner.pack_classes(sorted(schema))) == [sorted(schema)]
    assert len(counter.calls) == 1
    counter.calls.clear()
    assert split_windows("中文𠀀🙂", counter, 20) == [(0, 4)]
    assert len(counter.calls) == 1


def test_checkpoint_delta_replay_removals_torn_tail_and_new_snapshot(tmp_path):
    path = tmp_path / "checkpoint.json"
    writer = CheckpointJournal(path, every=100, seconds=1000)
    value = {
        "input_id": "run",
        "completed": {},
        "failures": {},
        "candidate_ids": [],
        "attempt_count": 0,
        "task_outcomes": {},
        "joint_completed": {},
    }
    writer.save(value)
    value["failures"]["retry"] = {"reason": "timeout"}
    value["attempt_count"] = 1
    writer.save(value)
    del value["failures"]["retry"]
    value["completed"]["retry"] = [{"candidate_id": "one"}]
    value["candidate_ids"] = ["one"]
    value["attempt_count"] = 2
    writer.save(value)
    assert read_checkpoint(path) == value
    with journal_path(path).open("ab") as stream:
        stream.write(b'{"base": "torn')
    assert read_checkpoint(path) == value
    # A restart loads and folds the successful prefix before writing anything.
    restarted = CheckpointJournal(path)
    restarted.save(read_checkpoint(path))
    assert not journal_path(path).exists()
    value["completed"]["retry"][0]["text"] = "immutable source edit detected"
    restarted.save(value)
    assert read_checkpoint(path) == value
    old_log = journal_path(path).read_bytes()
    newer = {**value, "input_id": "new-run", "attempt_count": 0}
    write_snapshot(path, newer)
    journal_path(path).write_bytes(old_log)
    assert read_checkpoint(path) == newer


def test_batch_store_bounds_selects_preserves_review_and_outer_rollback(db):
    job = ExtractionJob(source_type="word", status="annotating")
    db.add(job)
    db.commit()
    job_id = job.id
    values = [entity(f"entity-{i}") for i in range(83)]
    store = CandidateStore(db)
    saved = store.persist_validated(job_id, values, actor="test")
    reviewed = store.review(saved[0].candidate_id, 1, "rejected", "review retained", "analyst")
    queries = []

    def record(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            queries.append(statement)

    db.expire_on_commit = True
    event.listen(db.bind, "before_cursor_execute", record)
    try:
        again = store.persist_validated(job_id, values, actor="test")
    finally:
        event.remove(db.bind, "before_cursor_execute", record)
    assert len(queries) <= 4
    assert again[0] == reviewed
    assert {c.candidate_id for c in again} == {c.candidate_id for c in saved}
    new = entity("rollback")
    staged = store.persist_validated(job_id, [new], actor="test", commit=False)
    db.rollback()
    assert db.get(EvidenceCandidateRecord, staged[0].candidate_id) is None
    assert not db.scalars(
        select(EvidenceCandidateRevision).where(
            EvidenceCandidateRevision.candidate_id == staged[0].candidate_id
        )
    ).all()
    assert not db.scalars(
        select(EvidenceReview).where(EvidenceReview.candidate_id == staged[0].candidate_id)
    ).all()


def test_root_relationships_are_incremental_resume_without_recall_and_keep_properties_deferred(
    tmp_path,
):
    doc = Document()
    doc.add_paragraph("report")
    doc.add_paragraph("device A")
    for i in range(4):
        doc.add_paragraph(f"unrelated {i}")
    path = tmp_path / "incremental.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    schema = {
        "urn:Report": {"relationships": [{"iri": "urn:uses", "range": ["urn:Device"]}]},
        "urn:Device": {"properties": [{"iri": "urn:id"}]},
    }
    calls, snapshots = [], []

    def model(system, user, response_schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        task = context["task"]
        if request["stage"] == "verify_entity_types":
            return accept_fixture_types(request)
        calls.append(
            (
                task["task_kind"],
                [
                    f["anchor"]["evidence_id"]
                    for f in context["fragments"]
                    if f["purpose"] == "target"
                ],
            )
        )
        if task["task_kind"] == "entity":
            target = next(f for f in context["fragments"] if f["purpose"] == "target")
            return {
                "entities": [
                    {
                        "class_iri": "urn:Device",
                        "mention": {
                            "evidence_id": target["anchor"]["evidence_id"],
                            "text": "device A",
                        },
                    }
                ]
                if target["text"] == "device A"
                else []
            }
        return {"assertions": []}

    runner = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        model,
        model_identity="test",
        budget=TaskBudget(max_input_tokens=60000, max_tasks=100),
    )
    first = runner.run(
        ir,
        effective_class="urn:Report",
        pause_after=3,
        snapshot_fn=lambda run: snapshots.append(deepcopy(run)),
    )
    assert [kind for kind, _ in calls] == ["entity", "entity", "relationship"]
    assert snapshots and snapshots[-1].completion == "incomplete"
    assert {c.kind for snapshot in snapshots for c in snapshot.candidates} == {"entity"}
    before = list(calls)
    second = runner.run(ir, effective_class="urn:Report", checkpoint=deepcopy(first.checkpoint))
    assert second.completion == "complete", second.diagnostics
    assert len(calls) == len(set((kind, tuple(ids)) for kind, ids in calls))
    assert calls[: len(before)] == before
    replay_count = len(calls)
    runner.run(ir, effective_class="urn:Report", checkpoint=deepcopy(second.checkpoint))
    assert len(calls) == replay_count
    assert second.performance["stages"]["context"]["calls"] > 0


def test_legacy_scheduler_checkpoint_is_resumed_without_repeating_model_tasks(
    tmp_path, monkeypatch
):
    from app.services.extraction import extraction_tasks as module

    doc = Document()
    for text in ["one", "two", "three"]:
        doc.add_paragraph(text)
    path = tmp_path / "legacy.docx"
    doc.save(path)
    ir, calls = analyze_word_core(path).ir, []

    def model(*args):
        calls.append(1)
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Item": {}},
        TestTokenizer(),
        model,
        model_identity="test",
        budget=TaskBudget(max_input_tokens=60000),
    )
    with monkeypatch.context() as patch:
        patch.setattr(module, "SCHEDULER_VERSION", "document-path-priority-v3")
        old = runner.run(ir, pause_after=1)
    resumed = runner.run(ir, checkpoint=old.checkpoint)
    assert resumed.completion == "complete"
    assert resumed.scheduler_version == "document-path-priority-v3"
    assert len(calls) == len(ir.evidence_units)
    assert any(t["status"] == "restored" for t in resumed.tasks)
    fresh = runner.run(ir)
    assert fresh.scheduler_version == "document-batch-priority-v4"
    assert fresh.input_id != resumed.input_id


def test_overlapping_windows_cannot_publish_a_verdict_that_a_later_window_rejects(
    tmp_path,
    monkeypatch,
):
    from app.services.extraction import extraction_tasks as module

    doc = Document()
    doc.add_paragraph("report")
    doc.add_paragraph("x Device y")
    path = tmp_path / "overlap.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    snapshots, verdicts = [], []
    monkeypatch.setattr(
        module,
        "split_windows",
        lambda text, *args: (
            [(0, len(text) - 1), (1, len(text))] if "Device" in text else [(0, len(text))]
        ),
    )

    def model(system, user, schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        if request["stage"] == "verify_entity_types":
            verdicts.append(1)
            return {
                "decisions": [
                    {
                        "candidate_id": c["candidate_id"],
                        "supported": len(verdicts) == 1,
                        "reason": "fixture verdict",
                    }
                    for c in request["candidate"]["proposed_entities"]
                ]
            }
        if context["task"]["task_kind"] == "entity":
            target = next(f for f in context["fragments"] if f["purpose"] == "target")
            return {
                "entities": [
                    {
                        "class_iri": "urn:Device",
                        "mention": {
                            "evidence_id": target["anchor"]["evidence_id"],
                            "text": "Device",
                        },
                    }
                ]
                if "Device" in target["text"]
                else []
            }
        return {"assertions": []}

    runner = GenericExtractionRunner(
        {
            "urn:Report": {"relationships": [{"iri": "urn:uses", "range": ["urn:Device"]}]},
            "urn:Device": {},
        },
        TestTokenizer(),
        model,
        model_identity="overlap",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(
        ir,
        effective_class="urn:Report",
        snapshot_fn=lambda value: snapshots.append(deepcopy(value)),
    )
    assert len(verdicts) == 2
    assert all(c.text != "Device" for snapshot in snapshots for c in snapshot.candidates)
    candidate = next(c for c in run.candidates if c.text == "Device")
    assert candidate.validation_status != "passed"
