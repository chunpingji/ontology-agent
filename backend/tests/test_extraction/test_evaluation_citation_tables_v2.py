"""Real DOCX grids retain all column headers without broadening source permissions."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.citation_protocol import PROTOCOL_VERSION, CitationProtocol
from app.schemas.evidence import (
    Candidate,
    DocumentProvenance,
    EvidenceRange,
    EvidenceScope,
    ExtractionTask,
    TaskBudget,
)
from app.services.extraction.evidence_scope import candidate_ref
from app.services.extraction.extraction_tasks import (
    SYSTEM,
    BindingDecision,
    GenericExtractionRunner,
)
from app.services.extraction.table_records import table_records
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_evaluation_citation_protocol import (
    SCHEMA,
    candidate,
    envelope,
    source_id,
    subject_id,
)
from tests.test_extraction.test_evaluation_citation_protocol import (
    source as citation_source,
)
from tests.test_extraction.test_hierarchical_context import TestTokenizer

source = citation_source


def parsed(document, tmp_path):
    path = tmp_path / "column-context.docx"
    document.save(path)
    return analyze_word_core(path).ir


def record_envelope(ir, value, row):
    records = table_records(ir)
    value_unit = next(u for u in ir.evidence_units if u.text == value)
    identity_unit = next(u for u in ir.evidence_units if u.text == "Equipment DE64603.")
    anchor = ir.anchor(identity_unit.evidence_id)
    subject = Candidate(
        candidate_id="equipment",
        kind="entity",
        class_iri="urn:Product",
        text="DE64603",
        validation_status="passed",
        provenance=[
            DocumentProvenance(
                document_role=ir.document_role,
                anchors=[anchor],
                excerpts=[identity_unit.text],
            )
        ],
    )
    task = ExtractionTask(
        task_id="column-context",
        task_kind="property",
        subject=candidate_ref(subject),
        scope=EvidenceScope(
            scope_id="source-scope",
            document_hash=ir.document_hash,
            subject=candidate_ref(subject),
            ranges=[EvidenceRange(evidence_id=u.evidence_id) for u in ir.evidence_units],
            construction_evidence=[anchor],
        ),
        predicate_iri="urn:dose",
        predicate_definition=SCHEMA["urn:Product"]["properties"][0],
        target_evidence_ids=[u.evidence_id for u in records.row_units(value_unit.table_path, row)],
        budget=TaskBudget(max_input_tokens=100000),
    )
    runner = GenericExtractionRunner(
        SCHEMA, TestTokenizer(), lambda *args: None, model_identity="fixture"
    )
    original = runner._context(
        task, ir, {subject.candidate_id: subject}, "urn:Product", BindingDecision
    )
    proposed = {
        "candidate_id": "proposed",
        "kind": "property",
        "subject": candidate_ref(subject).model_dump(),
        "predicate_iri": "urn:dose",
        "assertion_status": "affirmed",
        "provenance": [
            {
                "kind": "document",
                "anchors": [ir.anchor(value_unit.evidence_id).model_dump(mode="json")],
                "excerpts": [value],
            }
        ],
    }
    return original, proposed


def context(protocol):
    return json.loads(json.loads(protocol.user)["context"])


def bounds(anchors, ir):
    return {
        (
            a.evidence_id,
            a.span_start or 0,
            a.span_end if a.span_end is not None else len(ir.unit(a.evidence_id).text),
        )
        for a in anchors
    }


@pytest.mark.parametrize("compact", [False, True])
def test_fourteen_column_record_keeps_every_header_and_never_moves_na(tmp_path, compact):
    document = Document()
    document.add_paragraph("Equipment DE64603.")
    table = document.add_table(rows=3, cols=14)
    headers = ["设备编号", "规格型号", *[f"其他列{index}" for index in range(2, 14)]]
    for row, values in zip(
        table.rows,
        [headers, ["DE64603", "24盘", *["N/A"] * 12], ["DE64604", "12盘", *["other"] * 12]],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    ir = parsed(document, tmp_path)
    original, proposed = record_envelope(ir, "24盘", 1)
    before = original.model_dump()
    protocol = CitationProtocol(
        SYSTEM,
        original,
        BindingDecision,
        ir=ir,
        stage="verify_binding",
        candidate=proposed,
        compact_identifiers=compact,
    )
    fragments = context(protocol)["fragments"]
    sources = {f["anchor"]["evidence_id"]: f for f in fragments}
    assert PROTOCOL_VERSION == "atomic-citations-v8-condition-review"
    assert set(headers) <= {f["text"] for f in fragments}
    assert len([f for f in fragments if f["text"] == "N/A"]) == 12
    assert not any(f["text"] in {"DE64604", "12盘", "other"} for f in fragments)
    for fragment in fragments:
        layout = fragment.get("table_context")
        if not layout or layout["is_header"]:
            continue
        column = layout["logical_columns"][0]
        assert layout["logical_rows"] == [1]
        assert [sources[h["evidence_id"]]["text"] for h in layout["column_headers"]] == [
            headers[column]
        ]
        assert layout["column_headers"][0]["covered_columns"] == [column]
        if fragment["text"] == "N/A":
            assert column != 1
    assert bounds(protocol.envelope.allowed_fact_regions, ir) == bounds(
        original.allowed_fact_regions, ir
    )
    assert bounds(protocol.envelope.allowed_binding_regions, ir) <= bounds(
        original.allowed_binding_regions, ir
    )
    assert original.model_dump() == before


@pytest.mark.parametrize("value,row", [("24盘", 2), ("N/A", 3)])
def test_merged_headers_and_cells_keep_logical_grid_without_pulling_other_rows(
    tmp_path, value, row
):
    document = Document()
    document.add_paragraph("Equipment DE64603.")
    table = document.add_table(rows=4, cols=4)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Identity"
    table.cell(0, 2).merge(table.cell(0, 3)).text = "Capacity"
    for index, text in enumerate(["设备编号", "规格型号", "最低值", "最高值"]):
        table.cell(1, index).text = text
    for index, text in enumerate(["DE64603", "24盘", "10", "20"]):
        table.cell(2, index).text = text
    table.cell(2, 0).merge(table.cell(3, 0)).text = "DE64603"
    table.cell(3, 1).text = "36盘"
    table.cell(3, 2).merge(table.cell(3, 3)).text = "N/A"
    ir = parsed(document, tmp_path)
    assert ir.tables[0]["header_row_count"] == 2
    original, proposed = record_envelope(ir, value, row)
    protocol = CitationProtocol(
        SYSTEM, original, BindingDecision, ir=ir, stage="verify_binding", candidate=proposed
    )
    fragments = context(protocol)["fragments"]
    sources = {f["anchor"]["evidence_id"]: f for f in fragments}
    assert {"Identity", "Capacity", "设备编号", "规格型号", "最低值", "最高值"} <= {
        f["text"] for f in fragments
    }
    shared = next(f for f in fragments if f["text"] == "DE64603")["table_context"]
    assert shared["logical_rows"] == [2, 3]
    assert shared["logical_columns"] == [0]
    if row == 2:
        assert not any(f["text"] in {"36盘", "N/A"} for f in fragments)
    else:
        assert not any(f["text"] in {"24盘", "10", "20"} for f in fragments)
        merged = next(f for f in fragments if f["text"] == "N/A")["table_context"]
        assert merged["logical_rows"] == [3]
        assert merged["logical_columns"] == [2, 3]
        assert {
            sources[h["evidence_id"]]["text"]: h["covered_columns"]
            for h in merged["column_headers"]
        } == {"Capacity": [2, 3], "最低值": [2], "最高值": [3]}


def test_missing_header_is_display_only_and_cannot_expand_permissions(source):
    original = envelope(source)
    payload = json.loads(original.serialized_input)
    missing = next(f for f in payload["fragments"] if f["text"] == "Approval")
    identity = missing["anchor"]["evidence_id"]
    payload["fragments"].remove(missing)
    original.allowed_binding_regions = [
        a for a in original.allowed_binding_regions if a.evidence_id != identity
    ]
    original.serialized_input = json.dumps(payload)
    protocol = CitationProtocol(
        SYSTEM,
        original,
        BindingDecision,
        ir=source,
        stage="verify_binding",
        candidate=candidate(source),
    )
    header = next(f for f in protocol.envelope.fragments if f["text"] == "Approval")
    assert not header["fact_eligible"] and not header["binding_eligible"]
    assert identity not in {a.evidence_id for a in protocol.envelope.allowed_binding_regions}
    with pytest.raises(ValueError, match="disallowed_citation"):
        protocol.decode(
            {
                "supported": True,
                "subject_candidate_id": subject_id(protocol),
                "assertion_status": "affirmed",
                "method": "table_record",
                "assertion_spans": [
                    {"evidence_id": source_id(protocol, source, "Approval"), "text": "Approval"}
                ],
            }
        )


@pytest.mark.parametrize("compact", [False, True])
def test_duplicate_source_roles_keep_all_subjects_and_per_role_permissions(source, compact):
    original = envelope(source)
    payload = json.loads(original.serialized_input)
    target = next(f for f in payload["fragments"] if f["text"] == "A")
    for subject in ("entity-A", "entity-B"):
        role = deepcopy(target)
        role.update(
            purpose="subject_evidence",
            subject_id=subject,
            fact_eligible=False,
            binding_permitted=False,
        )
        role["anchor"].update(span_start=0, span_end=1)
        payload["fragments"].append(role)
    original.serialized_input = json.dumps(payload)
    protocol = CitationProtocol(
        SYSTEM,
        original,
        BindingDecision,
        ir=source,
        stage="verify_binding",
        candidate=candidate(source),
        compact_identifiers=compact,
    )
    fragments = [f for f in context(protocol)["fragments"] if f["text"] == "A"]
    assert len(fragments) == 1
    fragment = fragments[0]
    assert fragment["fact_eligible"] and fragment["binding_eligible"]
    assert fragment["purpose"] == "target"
    assert {"target", "subject_evidence", "source"} <= set(fragment["roles"])
    assert {s["subject_id"] for s in fragment["subject_roles"]} == {
        subject_id(protocol, "A"),
        subject_id(protocol, "B"),
    }
    assert len(fragment["source_roles"]) == 3
    assert sum(s.get("binding_permitted") is False for s in fragment["source_roles"]) == 2
    assert all("text" not in role and "anchor" not in role for role in fragment["source_roles"])
    assert "subject_id" not in fragment
