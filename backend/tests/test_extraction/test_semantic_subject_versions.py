"""An obsolete intermediate endpoint cannot authorize a current descendant."""

from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.contracts import (
    EdgeSpec,
    GraphEdge,
    GraphNode,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
    VersionedRef,
)
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor, TaskOutcome
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.word_analysis import analyze_word_core


def test_stale_intermediate_endpoint_does_not_authorize_unchanged_descendant(tmp_path):
    root_class, first_class, second_class = "urn:versions:Root", "urn:versions:A", "urn:versions:B"
    root_predicate, first_predicate, leaf_property = (
        "urn:versions:r",
        "urn:versions:s",
        "urn:versions:p",
    )
    document = Document()
    document.add_paragraph("第一条原文。")
    document.add_paragraph("第二条原文。")
    path = tmp_path / "versions.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="versions-test",
    )

    def definition(iri, relationships=(), properties=()):
        return OntologyClassDefinition(
            iri=iri,
            label=iri,
            source_hash="frozen",
            declared_relationships=list(relationships),
            declared_properties=list(properties),
        )

    classes = {
        root_class: definition(
            root_class,
            [
                EdgeSpec(
                    iri=root_predicate,
                    label="关系甲",
                    range_class_iris=[first_class],
                )
            ],
        ),
        first_class: definition(
            first_class,
            [
                EdgeSpec(
                    iri=first_predicate,
                    label="关系乙",
                    range_class_iris=[second_class],
                )
            ],
        ),
        second_class: definition(
            second_class,
            properties=[
                SlotSpec(
                    iri=leaf_property,
                    label="属性丙",
                )
            ],
        ),
    }
    ontology = OntologySnapshot(
        snapshot_id="versions",
        ontology_hash=evidence_hash(classes),
        classes=classes,
        created_from="frozen_fixture",
    )
    first_anchor = analysis.ir.anchor(analysis.ir.evidence_units[0].evidence_id, 0, 6)
    first_node = GraphNode(
        entity_id="A",
        revision=1,
        class_iri=first_class,
        class_label="中间节点",
        label="对象甲",
        decision_status="undetermined",
        evidence_refs=[first_anchor],
    )
    second_node = GraphNode(
        entity_id="B", revision=1, class_iri=second_class, class_label="下游节点", label="对象乙"
    )
    calls = []

    class Adapter:
        model_identity = "version-fixture"

        def inspect(self, task, context, predicate, menu):
            calls.append((task.subject.entity_id, predicate.iri, context.record_id))
            first_record = any(
                "第一条" in fragment.text
                for fragment in context.fragments
                if fragment.fact_eligible
            )
            if predicate.iri == root_predicate and not first_record:
                return TaskOutcome(
                    semantic_outcome="supported",
                    reason_code="node_observation_advanced",
                    reason="同一物理提及的类型观察升版，尚无指向新版本的入边。",
                    model_calls=1,
                    nodes=[
                        first_node.model_copy(
                            update={"revision": 2, "decision_status": "supported"}
                        )
                    ],
                )
            if first_record and predicate.iri in (root_predicate, first_predicate):
                node = first_node if predicate.iri == root_predicate else second_node
                edge = GraphEdge(
                    candidate_id=f"edge:{node.entity_id}",
                    revision=1,
                    subject_ref=VersionedRef(
                        id=task.subject.entity_id, revision=task.subject.revision
                    ),
                    object_ref=VersionedRef(id=node.entity_id, revision=node.revision),
                    predicate_iri=predicate.iri,
                    predicate_label=predicate.label,
                    decision_status="supported",
                    structural_valid=True,
                    model_supported=True,
                    policy_eligible=True,
                    proof_ref=VersionedRef(id=f"proof:{node.entity_id}", revision=1),
                    decision_refs=[VersionedRef(id=f"decision:{node.entity_id}", revision=1)],
                )
                return TaskOutcome(
                    semantic_outcome="supported",
                    reason_code="supported",
                    reason="本测试的已核验边。",
                    model_calls=1,
                    nodes=[node],
                    edges=[edge],
                )
            return TaskOutcome(
                semantic_outcome="not_checked",
                reason_code="no_candidate_observed",
                reason="当前记录没有候选。",
                model_calls=1,
            )

    result = OntologyGuidedExecutor(ontology=ontology, engine=object(), adapter=Adapter()).run(
        recognition_run_id="versions-run",
        run_fingerprint="versions-fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=root_class,
        root_class_label="文档根",
        filename=path.name,
    )
    assert [predicate for _, predicate, _ in calls[:3]] == [
        root_predicate,
        first_predicate,
        root_predicate,
    ]
    assert not any(subject == "B" for subject, _, _ in calls)
    blocked_leaf = [
        value
        for kind, value in result.events
        if kind == "task_outcome" and value["task"]["subject"]["entity_id"] == "B"
    ]
    assert blocked_leaf
    assert all(
        value["outcome"]["reason_code"] == "subject_dependency_invalidated"
        for value in blocked_leaf
    )
