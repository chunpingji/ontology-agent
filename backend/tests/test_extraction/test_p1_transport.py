import json
from copy import deepcopy

from docx import Document

from app.schemas.evidence import TaskBudget
from app.services.extraction.extraction_tasks import (
    SYSTEM,
    EntityResponse,
    GenericExtractionRunner,
)
from app.services.extraction.hierarchical_context import model_request
from app.services.extraction.model_protocol import ModelProtocol, planning_schema
from app.services.extraction.verification_context import verification_payload
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.model_runtime import ModelCancelled
from tests.test_extraction.test_extraction_tasks import SCHEMA, scripted_model, span
from tests.test_extraction.test_hierarchical_context import TestTokenizer


def test_verification_keeps_source_roles_competitors_identity_and_ancestor_closure():
    schema = {
        "root": {"label": "root"},
        "product": {"parents": ["root"], "properties": [{"iri": "id", "identity_key": True}]},
        "drug": {"parents": ["root", "product"]},
        "sibling": {"parents": ["root", "product"]},
        "competitor": {"label": "different interpretation"},
        "irrelevant": {"label": "unrelated"},
    }
    payload = {
        "task": {"predicate_definition": {"classes": deepcopy(schema)}},
        "subjects": {"c1": {"class_iri": "competitor"}},
        "effective_class": "root",
        "fragments": [
            {
                "text": "same original",
                "purpose": "target",
                "fact_eligible": True,
                "anchor": {"evidence_id": "e", "table_path": "t", "row_index": 2},
            }
        ],
    }
    original = deepcopy(payload)
    result = verification_payload(payload, [{"class_iri": "drug"}], schema)
    assert payload == original
    assert result["fragments"] == payload["fragments"]
    assert result["subjects"] == payload["subjects"]
    classes = result["task"]["predicate_definition"]["classes"]
    assert set(classes) == {"root", "product", "drug", "sibling", "competitor"}
    assert classes["product"] == schema["product"]


def test_dynamic_constraints_do_not_change_system_or_ontology_prefix():
    payload = {
        "task": {
            "task_kind": "entity",
            "predicate_definition": {
                "classes": {"urn:Product": {"label": "product"}},
            },
        },
        "subjects": {},
        "effective_class": "",
        "fragments": [
            {"anchor": {"evidence_id": "source-one"}, "text": "A", "purpose": "target"},
        ],
    }
    schema = EntityResponse.model_json_schema()
    first = ModelProtocol(SYSTEM, model_request(payload, "recall"), schema)
    payload["fragments"].append({"anchor": {"evidence_id": "source-two"}, "text": "B"})
    second = ModelProtocol(SYSTEM, model_request(payload, "recall"), schema)
    assert first.schema != second.schema  # Concrete evidence enum changes.
    assert first.system == second.system
    assert first.user.split('\\"fragments\\"')[0] == second.user.split('\\"fragments\\"')[0]
    assert list(json.loads(first.user)) == ["stage", "context", "candidate", "output_contract"]
    assert json.loads(first.user)["output_contract"] == first.schema


def test_root_type_keeps_attribute_roles_subtypes_and_relation_neighbors():
    schema = {
        "operation": {"label": "操作", "properties": [
            {"iri": "condition", "label": "执行条件", "description": "压力与介质"},
        ], "relationships": [{"iri": "uses", "range": ["tool"]}]},
        "specific": {"parents": ["operation"]},
        "owner": {"relationships": [{"iri": "contains", "range": ["operation"]}]},
        "tool": {"label": "工具"}, "unrelated": {"label": "不相关"},
    }
    payload = {"task": {"predicate_definition": {"classes": {
        "operation": {"label": "操作", "parents": []},
    }}}, "subjects": {}, "fragments": [], "effective_class": ""}
    result = verification_payload(payload, [{"class_iri": "operation"}], schema)
    definitions = result["task"]["predicate_definition"]["classes"]
    assert set(definitions) == {"operation", "specific", "owner", "tool"}
    assert definitions["operation"]["property_roles"] == schema["operation"]["properties"]
    assert definitions["operation"]["relationship_roles"] == schema["operation"]["relationships"]
    assert "property_roles" not in payload["task"]["predicate_definition"]["classes"]["operation"]


def test_legacy_planning_is_frozen_despite_new_optional_proposal_fields():
    payload = {
        "task": {"task_kind": "entity", "predicate_definition": {"classes": {}}},
        "subjects": {},
        "fragments": [],
        "effective_class": "",
    }
    schema = EntityResponse.model_json_schema()
    legacy = planning_schema(schema)
    assert "supported" in schema["$defs"]["EntityProposal"]["properties"]
    assert "supported" not in legacy["$defs"]["EntityProposal"]["properties"]
    user = model_request(payload, "recall", planning=True)
    old = ModelProtocol(SYSTEM, user, legacy, planning=True)
    current = ModelProtocol(SYSTEM, user, schema, planning=True)
    assert (current.system, current.user, current.schema) == (old.system, old.user, old.schema)


def source(tmp_path):
    doc = Document()
    doc.add_paragraph("药品 A；规格 2mg")
    path = tmp_path / "p1.docx"
    doc.save(path)
    return analyze_word_core(path).ir


def test_partial_refusal_keeps_valid_entity_and_bounded_wrong_quote_diagnostics(tmp_path):
    ir = source(tmp_path)
    evidence = ir.evidence_units[0].evidence_id

    def model(system, user, schema, budget):
        request = json.loads(user)
        if request["stage"] == "recall":
            return {
                "refusal_reason": "规格不是药品，其余有支持",
                "entities": [
                    {
                        "class_iri": "urn:Product",
                        "mention": {"evidence_id": evidence, "text": "药品 A"},
                    },
                    {
                        "class_iri": "urn:Product",
                        "supported": False,
                        "reason": "数量不能作为药品",
                        "mention": {"evidence_id": evidence, "text": "2mg"},
                    },
                    {
                        "class_iri": "urn:Product",
                        "mention": {"evidence_id": evidence, "text": "错" * 500},
                    },
                ],
            }
        return {
            "decisions": [
                {
                    "candidate_id": c["candidate_id"],
                    "supported": True,
                    "reason": "原文明示药品 A",
                    "identity_supported": False,
                }
                for c in request["candidate"]["proposed_entities"]
            ]
        }

    runner = GenericExtractionRunner(
        {"urn:Product": {}},
        TestTokenizer(),
        model,
        model_identity="test",
        budget=TaskBudget(max_input_tokens=60000),
    )
    result = runner.run(ir)
    assert [(c.text, c.validation_status) for c in result.candidates] == [("药品 A", "passed")]
    assert result.completion == "incomplete"
    outcomes = next(iter(result.checkpoint["task_outcomes"].values()))
    issue = next(item for item in outcomes if item["code"] == "source_excerpt_mismatch")
    assert issue["stage"] == "entity_recall" and issue["task_id"]
    assert issue["evidence_id"] == evidence and issue["proposal_index"] == 2
    assert issue["quote"] == "错" * 240 and issue["quote_truncated"]
    assert "unsupported_type" in result.diagnostics


def test_cancelled_verification_is_automatically_resumable_without_resetting_calls(tmp_path):
    ir = source(tmp_path)
    cancelled = True

    def model(system, user, schema, budget):
        request = json.loads(user)
        if request["stage"] == "recall":
            return {
                "entities": [
                    {
                        "class_iri": "urn:Product",
                        "mention": {
                            "evidence_id": ir.evidence_units[0].evidence_id,
                            "text": "药品 A",
                        },
                    }
                ]
            }
        if cancelled:
            raise ModelCancelled()
        return {
            "decisions": [
                {
                    "candidate_id": c["candidate_id"],
                    "supported": True,
                    "reason": "independent verdict",
                }
                for c in request["candidate"]["proposed_entities"]
            ]
        }

    runner = GenericExtractionRunner(
        {"urn:Product": {}},
        TestTokenizer(),
        model,
        model_identity="test",
        budget=TaskBudget(max_input_tokens=60000),
    )
    paused = runner.run(ir)
    assert paused.checkpoint["attempt_count"] == 0
    assert paused.checkpoint["model_calls"] == 2
    assert next(iter(paused.checkpoint["failures"].values()))["interrupted"]
    cancelled = False
    resumed = runner.run(ir, checkpoint=paused.checkpoint)
    assert resumed.completion == "complete"
    assert resumed.input_id == paused.input_id
    assert resumed.checkpoint["model_calls"] == 4
    assert len(resumed.candidates) == 1


def test_cancelled_joint_binding_stays_unpublishable_and_resumes_only_joint_call(tmp_path):
    doc = Document()
    doc.add_paragraph("药品 A、药品 B 的规格为 250 mg")
    path = tmp_path / "shared.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    calls, cancelled = [], True

    def model(system, user, schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        calls.append(request["stage"])
        if request["stage"] == "verify_shared_binding":
            if cancelled:
                raise ModelCancelled()
            return {"supported_subject_ids": [], "refusal_reason": "缺少共享证据"}
        if request["stage"] == "recall" and context["task"]["task_kind"] == "property":
            target = next(f for f in context["fragments"] if f["purpose"] == "target")
            return {"assertions": [{"value": span(target, "250 mg")}]}
        return scripted_model(system, user, schema, budget)

    runner = GenericExtractionRunner(SCHEMA, TestTokenizer(), model, model_identity="test",
                                     budget=TaskBudget(max_input_tokens=60000))
    paused = runner.run(ir)
    assert paused.completion == "incomplete" and not paused.checkpoint["joint_completed"]
    assert all(not c.positive_eligible for c in paused.candidates if c.kind == "property")
    before = len(calls)
    cancelled = False
    resumed = runner.run(ir, checkpoint=paused.checkpoint)
    assert calls[before:] == ["verify_shared_binding"]
    assert resumed.checkpoint["joint_completed"]
