"""Rich preview should resolve a paragraph style once, preserving run overrides."""

from types import SimpleNamespace

from docx import Document
from docx.shared import Pt

from app.services.extraction.document_annotator import _para_runs_and_text


def test_resolves_paragraph_style_once_and_preserves_direct_run_formatting():
    document = Document()
    style = document.styles["Normal"]
    style.font.bold = True
    style.font.italic = True
    style.font.size = Pt(11)
    style.font.name = "Example Font"
    paragraph = document.add_paragraph()
    paragraph.add_run("继承")
    direct = paragraph.add_run("直接")
    direct.bold = False
    direct.italic = False
    direct.underline = True
    direct.font.size = Pt(14)

    class CountedParagraph:
        reads = 0
        runs = paragraph.runs

        @property
        def style(self):
            self.reads += 1
            return paragraph.style

    counted = CountedParagraph()
    text, runs = _para_runs_and_text(counted, rich=True)
    assert counted.reads == 1
    assert text == "继承直接"
    assert runs == [
        (0, 2, [
            {"type": "bold"}, {"type": "italic"},
            {"type": "textStyle", "attrs": {"fontSize": "11pt", "fontFamily": "Example Font"}},
        ]),
        (2, 4, [
            {"type": "underline"},
            {"type": "textStyle", "attrs": {"fontSize": "14pt", "fontFamily": "Example Font"}},
        ]),
    ]


def test_missing_style_still_preserves_text():
    paragraph = Document().add_paragraph("正文")
    text, runs = _para_runs_and_text(SimpleNamespace(style=None, runs=paragraph.runs), rich=True)
    assert text == "正文"
    assert runs == [(0, 2, [])]
