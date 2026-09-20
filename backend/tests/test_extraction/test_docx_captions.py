"""Table captions stay with their tables without absorbing later sections."""

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.services.extraction.document_annotator import annotate_word
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import infer_heading_level, parse_docx_structure


@pytest.mark.parametrize("font_size", [10.5, 18])
@pytest.mark.parametrize("caption", ["表1得量收率范围", "表 1 得量收率范围", "Table 1 Yield range"])
def test_table_caption_preserves_sibling_sections_and_source_membership(
    tmp_path, font_size, caption,
):
    doc = Document()
    doc.add_heading("工艺", level=2)
    doc.add_heading("得量收率范围", level=3)
    paragraph = doc.add_paragraph(caption)
    paragraph.runs[0].bold = True
    paragraph.runs[0].font.size = Pt(font_size)
    # Blank spacing paragraphs do not separate a caption from its table.
    doc.add_paragraph("")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "收率"
    table.cell(1, 0).text = "50%~85%"
    doc.add_heading("存放条件", level=3)
    doc.add_paragraph("避光保存。")
    doc.add_paragraph("表2中间体及成品的存放条件").runs[0].bold = True
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "条件"
    table.cell(1, 0).text = "密封"
    doc.add_heading("设备需求", level=3)
    path = tmp_path / "captions.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    assert structure.headings == ["工艺", "得量收率范围", "存放条件", "设备需求"]
    process, = structure.section_tree.children
    yields, storage, equipment = process.children
    assert [node.heading for node in process.children] == [
        "得量收率范围", "存放条件", "设备需求",
    ]
    assert yields.table_indices == [0]
    assert storage.table_indices == [1]
    assert equipment.table_indices == []
    assert structure.tables[0].section_path == ["工艺", "得量收率范围"]
    assert structure.tables[1].section_path == ["工艺", "存放条件"]
    assert structure.find_section("得量收率范围").paras == [caption]
    assert yields.source_range.end_block_id == "table:0"
    assert storage.source_range.start_block_id == "paragraph:4:0"
    assert storage.source_range.end_block_id == "table:1"
    assert yields.pages[0].table_indices == [0]
    assert storage.pages[0].table_indices == [1]

    ir = build_document_ir(path, structure)
    for text, owner in [(caption, yields), ("50%~85%", yields), ("密封", storage)]:
        unit = next(unit for unit in ir.evidence_units if unit.text == text)
        assert unit.section_node_id == owner.node_id
        assert ir.resolve(ir.anchor(unit.evidence_id)) == text
        if text == caption:
            assert unit.kind == "paragraph"

    content, _, _, _ = annotate_word(
        path, engine=None, structure_only=True, rich_style=True, structure=structure, ir=ir,
    )
    preview = {
        node.get("attrs", {}).get("sourceBlockId"): node for node in content["content"]
    }
    caption_node = preview["paragraph:2:0"]
    assert caption_node["type"] == "paragraph"
    assert caption_node["attrs"]["sectionNodeId"] == yields.node_id
    assert any(
        mark["type"] == "bold"
        for run in caption_node["content"] for mark in run.get("marks", [])
    )
    assert preview["table:0"]["attrs"]["sectionNodeId"] == yields.node_id
    assert preview["paragraph:4:0"]["type"] == "heading"
    assert preview["paragraph:4:0"]["attrs"]["sectionNodeId"] == storage.node_id


@pytest.mark.parametrize("style_name", ["Caption", "题注", "表题"])
def test_explicit_caption_style_does_not_need_number_or_adjacent_table(style_name):
    doc = Document()
    if style_name not in doc.styles:
        doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph = doc.add_paragraph("得量收率范围", style=style_name)
    paragraph.runs[0].bold = True
    paragraph.runs[0].font.size = Pt(18)
    assert infer_heading_level(paragraph) == 0


@pytest.mark.parametrize("caption", ["表一 得量收率范围", "表１ 得量收率范围", "Table 1.2 Yield"])
@pytest.mark.parametrize("position", ["before", "after"])
def test_numbered_caption_can_be_above_or_below_table(caption, position):
    doc = Document()
    if position == "after":
        doc.add_table(rows=1, cols=1)
        doc.add_paragraph("")
    paragraph = doc.add_paragraph(caption)
    paragraph.runs[0].bold = True
    if position == "before":
        doc.add_paragraph("")
        doc.add_table(rows=1, cols=1)
    assert infer_heading_level(paragraph) == 0


@pytest.mark.parametrize("heading_source", ["style", "outline"])
def test_explicit_heading_takes_precedence_over_caption_like_text(heading_source):
    doc = Document()
    paragraph = doc.add_paragraph("表1结果分析")
    if heading_source == "style":
        paragraph.style = "Heading 3"
    else:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "2")
        paragraph._p.get_or_add_pPr().append(outline)
    doc.add_table(rows=1, cols=1)
    assert infer_heading_level(paragraph) == 3


def test_caption_detection_does_not_demote_unrelated_bold_headings():
    doc = Document()
    numbered = doc.add_paragraph("表1结果分析")
    numbered.runs[0].bold = True
    doc.add_paragraph("这是章节正文，后面才是表格。")
    doc.add_table(rows=1, cols=1)
    semantic = doc.add_paragraph("存放条件")
    semantic.runs[0].bold = True
    doc.add_table(rows=1, cols=1)
    assert infer_heading_level(numbered) == 2
    assert infer_heading_level(semantic) == 2
