"""Regression barriers for the isolated joint-evidence diagnostic driver.

Synthetic DOCX input and controlled transport keep these checks independent of
the benchmark artifacts, model server, and production database.
"""

from __future__ import annotations

import importlib.util
import json
import time
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import GraphNode
from app.services.extraction.ontology_guided.model_adapter import LocalModelRecognitionAdapter
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_joint_evidence_validation import _fixture, _response
from tests.test_extraction.test_semantic_graph_closure import DESCRIBES, ROOT, _ontology


@pytest.fixture
def runner(monkeypatch):
    scripts = Path(__file__).resolve().parents[2] / "scripts"
    monkeypatch.syspath_prepend(str(scripts))
    spec = importlib.util.spec_from_file_location(
        "joint_evidence_runner_test", scripts / "benchmark_joint_evidence.py",
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def inspection(tmp_path, runner):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    output = tmp_path / "experiment"
    output.mkdir()
    experiment = runner.Experiment(output, {
        "run_id": "joint-runner-test",
        "limits": {
            "max_tasks": 8, "max_calls": 16, "deadline_seconds": 1200,
            "max_model_calls_per_record": 4,
        },
    })
    experiment.recognition_started = time.perf_counter()
    experiment.active = "product-current"
    experiment.active_lineage = task.claim_lineage_id
    context = assemble_context(target, task.record_id, index)
    # This is a scheduling-time snapshot; reserve must still check live limits.
    context.remaining_model_calls = 2
    context.bind_model_call_hook(experiment.reserve)
    adapter = LocalModelRecognitionAdapter(object(), model_identity="joint-runner-test")
    predicate = _ontology().classes[ROOT].declared_relationships[0]
    return SimpleNamespace(
        experiment=experiment, adapter=adapter, task=task, context=context,
        predicate=predicate, index=index, binding=binding,
    )


def _forbid_dispatch(monkeypatch):
    dispatched = []

    def unexpected(*_args, **_kwargs):
        dispatched.append(True)
        raise AssertionError("model dispatched past a failed reservation barrier")

    monkeypatch.setattr(model_adapter, "chat_with_schema", unexpected)
    return dispatched


def test_merged_table_reference_resolves_the_requested_logical_row(tmp_path, runner):
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for column, text in enumerate(("产品", "属性", "值")):
        table.cell(0, column).text = text
    table.cell(1, 0).merge(table.cell(2, 0)).text = "产品甲"
    table.cell(1, 1).text, table.cell(1, 2).text = "外观", "白色固体"
    table.cell(2, 1).text, table.cell(2, 2).text = "批号", "批次乙"
    path = tmp_path / "merged-rows.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)

    first = runner.selected_record(index, runner.table_row(0, 1))
    second = runner.selected_record(index, runner.table_row(0, 2))
    shared = {unit.evidence_id for unit in first.source_units} & {
        unit.evidence_id for unit in second.source_units
    }
    assert shared, "fixture must reproduce overlapping physical evidence across rows"
    assert first.record_id != second.record_id
    assert second.table_path == ("table:0",) and second.row_index == 2
    assert second.text == "产品甲\n批号\n批次乙"

    support = runner.selected_units(index, runner.table_row(0, 2))
    assert {unit.evidence_id for unit in support} == {
        unit.evidence_id for unit in second.source_units
    }
    assert [unit.text for unit in support] == ["产品甲", "批号", "批次乙"]
    assert shared <= {unit.evidence_id for unit in support}
    with pytest.raises(ValueError, match="ambiguous target table row"):
        runner.selected_record(index, runner.table_row(0, 99))


@pytest.mark.parametrize("exhausted", ["request", "lineage", "deadline"])
def test_live_budget_failure_prevents_dispatch(inspection, runner, monkeypatch, exhausted):
    experiment = inspection.experiment
    limits = experiment.manifest["limits"]
    if exhausted == "request":
        experiment.calls = limits["max_calls"]
    elif exhausted == "lineage":
        experiment.lineage_calls[experiment.active_lineage] = limits["max_model_calls_per_record"]
    else:
        experiment.recognition_started -= limits["deadline_seconds"] + 1
    before = (experiment.calls, dict(experiment.lineage_calls))
    dispatched = _forbid_dispatch(monkeypatch)
    reason = "deadline_exhausted" if exhausted == "deadline" else f"{exhausted}_budget_exhausted"

    with pytest.raises(runner.BudgetStop, match=reason):
        inspection.adapter.inspect(
            inspection.task, inspection.context, inspection.predicate, None,
        )

    assert dispatched == []
    assert (experiment.calls, experiment.lineage_calls) == before
    assert experiment.reservations == []
    assert not (experiment.output / "reservations.json").exists()


@pytest.mark.parametrize("failed_write", ["event", "reservations"])
def test_reservation_persistence_failure_prevents_dispatch(
    inspection, runner, monkeypatch, failed_write,
):
    experiment = inspection.experiment
    dispatched = _forbid_dispatch(monkeypatch)

    def failed(*_args, **_kwargs):
        raise OSError("isolated reservation write failure")

    if failed_write == "event":
        monkeypatch.setattr(experiment, "event", failed)
    else:
        monkeypatch.setattr(runner, "write_json", failed)

    with pytest.raises(OSError, match="isolated reservation write failure"):
        inspection.adapter.inspect(
            inspection.task, inspection.context, inspection.predicate, None,
        )

    assert dispatched == []
    # A failed durability barrier must not silently reclaim a possible request.
    assert experiment.calls == experiment.lineage_calls[experiment.active_lineage] == 1
    assert not (experiment.output / "reservations.json").exists()


def test_comparison_arms_share_a_persisted_lineage_budget(inspection, runner, monkeypatch):
    experiment = inspection.experiment
    experiment.ontology = _ontology()
    experiment.engine = None
    experiment.bind = None
    experiment.index = inspection.index
    experiment.fingerprint = "joint-runner-fixture"
    experiment.settings = SimpleNamespace(evidence_max_input_tokens=8192)
    experiment.adapter = inspection.adapter
    experiment.root = GraphNode(
        entity_id=inspection.task.subject.entity_id, revision=1, class_iri=ROOT,
        class_label="报告", label="测试报告", root=True, root_origin="user_specified",
    )
    monkeypatch.setattr(runner, "CMC_ROOT", ROOT)
    requests = []

    def respond(_client, *, user, **_kwargs):
        # Observe the durable reservation at the actual transport boundary.
        reservations = json.loads((experiment.output / "reservations.json").read_text())
        assert len(reservations) == experiment.calls == len(requests) + 1
        assert reservations[-1]["lineage"] == experiment.active_lineage
        request = json.loads(user)
        requests.append(request)
        return _response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    arguments = {
        "subject": inspection.task.subject, "node": experiment.root, "predicate": DESCRIBES,
        "record": experiment.index.by_id[inspection.task.record_id],
    }
    current, before = experiment.inspect(
        case_id="product", arm="current", support=[], **arguments,
    )
    joint, after = experiment.inspect(
        case_id="product", arm="joint", support=[inspection.binding], **arguments,
    )
    assert before.semantic_outcome == "undetermined"
    assert after.semantic_outcome == "supported"
    assert current["lineage"] == joint["lineage"]
    assert current["context_hash"] != joint["context_hash"]
    assert current["model_calls"] == joint["model_calls"] == 2
    assert experiment.lineage_calls == {current["lineage"]: 4}

    # Renaming the case cannot reset the subject/predicate/record request budget.
    blocked, outcome = experiment.inspect(
        case_id="renamed-product", arm="joint", support=[inspection.binding], **arguments,
    )
    assert blocked["lineage"] == current["lineage"]
    assert blocked["status"] == outcome.semantic_outcome == "not_checked"
    assert blocked["reason_code"] == "record_model_call_budget_exhausted"
    assert blocked["model_calls"] == 0 and not outcome.complete
    assert experiment.calls == len(requests) == 4
