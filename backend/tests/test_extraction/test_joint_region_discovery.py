"""Joint discovery keeps source bundles, legal predicates and per-claim decisions."""

import copy
import json

import pytest
from docx import Document

from app.services.extraction.ontology_guided.claim_protocol import ExtractionProfile
from app.services.extraction.ontology_guided.context import assemble_record_discovery_context
from app.services.extraction.ontology_guided.contracts import EdgeSpec
from app.services.extraction.ontology_guided.record_discovery import (
    RecordDiscoveryPolicy,
    compile_record_schema_card,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_record_model_adapter import setup_record


def joint_setup(tmp_path, monkeypatch, *, change=None, stop_at=None, conditional=True):
    document = Document()
    document.add_paragraph("对象乙的数量为5 mg，对象丙的名称为样品丙。")
    document.add_paragraph(
        "上述对象乙连接对象丙" + ("；仅在试验阶段连接。" if conditional else "。")
    )
    path = tmp_path / "joint.docx"
    if not path.exists():
        document.save(path)  # Cold resume must reuse identical source bytes, including ZIP dates.
    index = RecordIndex(analyze_word_core(path).ir)
    records = list(index.records)
    units = [unit for record in records for unit in record.source_units]

    def quote(text):
        unit = next(unit for unit in units if text in unit.text)
        return dict(evidence_id=unit.evidence_id, text=text, context_text=None)

    source = dict(
        index=index,
        record_id=records[0].record_id,
        source_unit=units[0],
        quote=quote,
        proposal=dict(
            entities=[
                dict(
                    local_id=identity,
                    class_iri=cls,
                    representation="mention",
                    mentions=[quote(name)],
                    record_components=[],
                    identifier_claims=[],
                )
                for identity, cls, name in [("b", "urn:T", "对象乙"), ("c", "urn:Other", "对象丙")]
            ],
            properties=[],
            relations=[],
            external_links=[],
            observations=[],
        ),
    )
    bundle = [quote(unit.text) for unit in units]

    def transform(answer, view, ordinal):
        if view["stage"] == "discovery":
            answer["relations"] = [
                dict(
                    local_id="link",
                    subject_id="b",
                    predicate_iri="urn:connects",
                    object_ids=["c"],
                    selection="all",
                    bridge_kind="explicit_assertion",
                    bridge_ref_ids=[],
                    bridge_support=[quote("连接对象丙")],
                    selection_support=[],
                    qualifiers=dict(
                        polarity="affirmed",
                        modality="asserted",
                        condition_support=([quote("仅在试验阶段连接")] if conditional else []),
                        scope_qualifiers=[],
                    ),
                    source_assertion=dict(
                        subject_support=[quote("对象乙")],
                        object_support=[dict(object_id="c", support=[quote("对象丙")])],
                        predicate_support=[quote("连接对象丙")],
                        binding_ids=[],
                    ),
                )
            ]
        else:
            for result in answer["verifications"]:
                target = next(
                    t
                    for t in view["verification_input"]["targets"]
                    if t["target_id"] == result["target_id"]
                )
                for facet in result["facets"]:
                    facet["support"] = copy.deepcopy(bundle)
                    if facet["name"] == "counterevidence" or (
                        facet["name"] == "qualifiers"
                        and (target["target_kind"] == "property" or not conditional)
                    ):
                        facet["support"] = []
                        facet["reason"] = "已核对当前两段原文，无额外限定或反证。"
        if change:
            change(answer, view, ordinal)

    adapter, task, original, _, storage, requests, proposal = setup_record(
        source,
        monkeypatch,
        transform=transform,
        stop_at=stop_at,
    )
    adapter.ontology.classes["urn:T"].declared_relationships = [
        EdgeSpec(
            iri="urn:connects",
            label="连接",
            description="主体与对象之间的连接，须保留适用条件。",
            declared_by=["urn:T"],
            range_class_iris=["urn:Other"],
        )
    ]
    card = compile_record_schema_card(
        adapter.ontology,
        class_iris=["urn:T", "urn:Other"],
        analysis_scope_ref=task.analysis_scope_ref,
        profile=ExtractionProfile(),
    )
    task = task.model_copy(
        update={
            "schema_card_id": card.schema_card_id,
            "source_record_ids": [r.record_id for r in records],
        }
    )
    context = assemble_record_discovery_context(
        task,
        card,
        index,
        ontology_hash=adapter.ontology.ontology_hash,
        document_context=original.target.document_context,
    )
    context.tool_inputs = {
        **original.tool_inputs,
        "schema_card": card.model_dump(mode="json"),
        "target_seed": context.target.model_dump(mode="json"),
        "graph_phase": "evidence_review",
    }
    context.remaining_model_calls = 6
    context.bind_protocol_hook(original._protocol_hook)
    context.bind_model_call_hook(original.before_model_call)
    adapter.record_discovery = RecordDiscoveryPolicy(graph_phase="evidence_review")
    return adapter, task, context, card, storage, requests, proposal


def test_same_discovery_creates_entities_properties_and_combined_source_relation(
    tmp_path, monkeypatch
):
    adapter, task, context, card, storage, requests, _ = joint_setup(tmp_path, monkeypatch)
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete, outcome.reason
    assert len(outcome.nodes) == len(outcome.properties) == 2
    (edge,) = outcome.edges
    assert edge.policy_eligible and edge.proof_ref
    assert edge.conditions == ["仅在试验阶段连接"]
    assert len({ref.evidence_id for ref in edge.predicate_evidence_refs}) == 2
    assert all(adapter.index.ir.resolve(ref) for ref in edge.evidence_refs)
    assert len(requests) == 2  # one joint discovery and one independent verification
    first_schema = requests[0]["text_format"]["schema"]
    assert first_schema["properties"]["properties"].get("maxItems") != 0
    assert first_schema["$defs"]["RelationProposal"]["properties"]["predicate_iri"]["enum"] == [
        "urn:connects",
    ]
    view = json.loads(requests[1]["input_items"][0]["content"][0]["text"])
    assert len(view["verification_input"]["attribute_candidates"]) == 2
    assert {t["target_kind"] for t in view["verification_input"]["targets"]} == {
        "entity",
        "property",
        "relation",
    }
    context.protocol_state = copy.deepcopy(storage["protocol"])
    context.protocol_results = copy.deepcopy(storage["results"])
    restored = copy.copy(adapter).inspect_record(task, context, card)
    assert len(requests) == 2 and restored.edges == outcome.edges


def test_unconditional_relation_records_checked_scope_without_invented_quotes(
    tmp_path, monkeypatch
):
    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch, conditional=False)
    outcome = adapter.inspect_record(task, context, card)
    (edge,) = outcome.edges
    assert edge.policy_eligible and not edge.conditions
    decisions = [
        d for d in outcome.decision_payloads if d["reason_code"].startswith("scope_checked_no_")
    ]
    assert {d["check_kind"] for d in decisions} == {"qualifiers", "counterevidence"}
    assert all(
        d["searched_context_refs"]
        and not d["support_refs"]
        and not d["counterevidence_refs"]
        and d["reason"]
        for d in decisions
    )


def test_explicit_counterevidence_cannot_be_skipped_as_absent(tmp_path, monkeypatch):
    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch)
    context.counterevidence_refs = [context.fragments[-1].bounded_anchor()]
    outcome = adapter.inspect_record(task, context, card)
    (edge,) = outcome.edges
    assert not edge.policy_eligible and edge.proof_ref is None
    assert "counterevidence" in edge.reason and edge.evidence_refs


def test_changed_checked_scope_cannot_obtain_relation_proof(tmp_path, monkeypatch):
    from app.services.extraction.ontology_guided.verification import ProofGate

    evaluate = ProofGate.evaluate_frozen_claim
    changed = []

    def changed_scope(gate, claim, decisions, **kwargs):
        if claim.target_kind == "relation":
            decisions = [
                d.model_copy(update={"searched_context_refs": d.searched_context_refs[:1]})
                if d.check_kind == "counterevidence" else d
                for d in decisions
            ]
            changed.append(claim.target_id)
        return evaluate(gate, claim, decisions, **kwargs)

    monkeypatch.setattr(ProofGate, "evaluate_frozen_claim", changed_scope)
    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch)
    outcome = adapter.inspect_record(task, context, card)
    (edge,) = outcome.edges
    assert changed and len(context.fragments) > 1
    assert not edge.policy_eligible and edge.proof_ref is None
    assert "counterevidence_checked_scope_mismatch" in edge.reason


def test_joint_discovery_cold_resume_only_finishes_verification(tmp_path, monkeypatch):
    adapter, task, context, card, storage, requests, _ = joint_setup(
        tmp_path, monkeypatch, stop_at="discovery"
    )
    with pytest.raises(RuntimeError, match="pause after durable commit"):
        adapter.inspect_record(task, context, card)
    saved = copy.deepcopy(storage)
    assert len(requests) == 1
    adapter, task, context, card, storage, resumed, _ = joint_setup(tmp_path, monkeypatch)
    storage.update(saved)
    context.protocol_state = saved["protocol"]
    context.protocol_results = saved["results"]
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete and len(outcome.nodes) == len(outcome.properties) == 2
    assert len(outcome.edges) == 1 and outcome.edges[0].policy_eligible
    assert len(resumed) == 1
    view = json.loads(resumed[0]["input_items"][0]["content"][0]["text"])
    assert view["stage"] == "verification"
    assert len(view["verification_input"]["attribute_candidates"]) == 2
    assert storage["results"][saved["protocol"]["discovery_ref"]] == (
        saved["results"][saved["protocol"]["discovery_ref"]]
    )


def test_relation_uses_property_sources_without_waiting_for_property_acceptance(
    tmp_path, monkeypatch
):
    def change(answer, view, _ordinal):
        if view["stage"] == "verification":
            assert view["verification_input"]["attribute_candidates"]
            for target, result in zip(
                view["verification_input"]["targets"], answer["verifications"]
            ):
                if target["target_kind"] == "property":
                    next(f for f in result["facets"] if f["name"] == "predicate").update(
                        verdict="undetermined",
                        reason="属性映射未确定，但保留原文字段供关系核对。",
                    )

    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch, change=change)
    outcome = adapter.inspect_record(task, context, card)
    assert all(not p.policy_eligible for p in outcome.properties)
    (edge,) = outcome.edges
    assert edge.policy_eligible and edge.proof_ref


def test_incomplete_source_assertion_keeps_cited_relation_pending(tmp_path, monkeypatch):
    def change(answer, view, _ordinal):
        if view["stage"] == "discovery":
            answer["relations"][0]["source_assertion"] = None

    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch, change=change)
    outcome = adapter.inspect_record(task, context, card)
    (edge,) = outcome.edges
    assert edge.decision_status == "undetermined"
    assert not edge.policy_eligible and edge.proof_ref is None
    assert "source_assertion_required" in edge.reason and edge.evidence_refs


@pytest.mark.parametrize("mode", ["predicate", "subject_binding", "counterevidence", "qualifiers"])
def test_failed_semantics_or_missing_real_condition_keeps_dashed_candidate(
    tmp_path, monkeypatch, mode
):
    def change(answer, view, _ordinal):
        if view["stage"] != "verification":
            return
        for target, result in zip(view["verification_input"]["targets"], answer["verifications"]):
            if target["target_kind"] != "relation":
                continue
            facet = next(f for f in result["facets"] if f["name"] == mode)
            if mode == "qualifiers":
                facet["support"] = []
            else:
                facet.update(verdict="undetermined", reason="仅有共现或归属/反证尚未解决。")

    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch, change=change)
    outcome = adapter.inspect_record(task, context, card)
    (edge,) = outcome.edges
    assert (
        not edge.policy_eligible and not edge.proof_ref and edge.decision_status == "undetermined"
    )
    assert mode in edge.reason and edge.evidence_refs and edge.conditions
    assert len(outcome.properties) == 2  # relation failure does not discard independent attributes


@pytest.mark.parametrize(
    "mode", ["illegal_predicate", "wrong_owner", "invalid_quote", "bad_entity"]
)
def test_joint_discovery_cannot_escape_card_identity_or_source(tmp_path, monkeypatch, mode):
    def change(answer, view, _ordinal):
        if view["stage"] == "discovery":
            relation = answer["relations"][0]
            if mode == "illegal_predicate":
                relation["predicate_iri"] = "urn:invented"
            elif mode == "wrong_owner":
                relation.update(subject_id="c", object_ids=["b"])
            elif mode == "invalid_quote":
                relation["bridge_support"][0]["text"] = "原文不存在的证明"
        elif mode == "bad_entity":
            for target, result in zip(
                view["verification_input"]["targets"], answer["verifications"]
            ):
                if target["target_kind"] == "entity" and target["payload"]["local_id"] == "b":
                    next(f for f in result["facets"] if f["name"] == "type")["verdict"] = (
                        "unsupported"
                    )

    adapter, task, context, card, _, _, _ = joint_setup(tmp_path, monkeypatch, change=change)
    outcome = adapter.inspect_record(task, context, card)
    assert not any(edge.policy_eligible for edge in outcome.edges), outcome.reason
    assert outcome.reason_code
