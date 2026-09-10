"""Production atomic citation and logical-record task regressions."""

import json

import pytest
from docx import Document

from app.schemas.evidence import (
    Candidate,
    DocumentProvenance,
    EvidenceRange,
    ExtractionTask,
    TaskBudget,
)
from app.services.extraction.evidence_scope import build_scope, candidate_ref
from app.services.extraction.extraction_tasks import GenericExtractionRunner, PartialTaskFailure
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.record_targets import field_group_scope, record_targets
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_extraction_tasks import accept_fixture_types
from tests.test_extraction.test_hierarchical_context import TestTokenizer
from tests.test_extraction.test_template_relation_recovery import parts, quote, source


def runner(model, **options):
    return GenericExtractionRunner(
        {'urn:Product': {'label': '产品'}}, TestTokenizer(), model,
        model_identity='atomic-record-test', atomic_citations=True,
        budget=TaskBudget(max_input_tokens=60000, max_regions_per_task=1), **options,
    )


def test_atomic_id_only_quote_replays_original_without_losing_a_valid_sibling(tmp_path):
    ir = source(tmp_path, 'SM5592- A14', '不存在的另一个提议')
    calls = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append(request['stage'])
        if request['stage'] == 'verify_entity_types':
            return accept_fixture_types(request)
        assert set(schema['$defs']['FactSpan']['properties']) == {'evidence_id', 'text'}
        return {'entities': [
            {'class_iri': 'urn:Product',
             'mention': {'evidence_id': targets[0]['anchor']['evidence_id']}},
            {'class_iri': 'urn:Product', 'mention': quote(targets[1], 'SM5592-A14')},
        ]}

    worker = runner(model)
    regions = [EvidenceRange(evidence_id=unit.evidence_id, start=0, end=len(unit.text))
               for unit in ir.evidence_units if unit.text]
    task = ExtractionTask(task_id='atomic', task_kind='entity',
                          target_class_iris=['urn:Product'], target_ranges=regions,
                          target_evidence_ids=[value.evidence_id for value in regions],
                          predicate_definition=worker.entity_menu(['urn:Product'])['predicate_definition'],
                          budget=worker.budget)
    with pytest.raises(PartialTaskFailure) as failure:
        worker.execute_task(task, ir, {})
    assert calls == ['recall', 'verify_entity_types']
    assert [item.text for item in failure.value.candidates] == ['SM5592- A14']
    assert failure.value.candidates[0].positive_eligible
    assert any(event.get('field_path') == 'entities[1].mention' for event in worker._task_events)
    assert ir.resolve(failure.value.candidates[0].provenance[0].anchors[0]) == 'SM5592- A14'


def test_atomic_names_and_values_are_not_reclipped_to_separate_binding_quotes(tmp_path):
    ir = source(tmp_path, '产品甲的分子量为537.18。')
    calls = []

    def model(system, user, response_schema, budget):
        request, task, targets = parts(user)
        calls.append(request['stage'])
        if request['stage'] == 'verify_entity_types':
            return accept_fixture_types(request)
        if request['stage'] == 'verify_binding':
            return {'supported': True,
                    'subject_candidate_id': request['candidate']['subject']['candidate_id'],
                    'assertion_status': 'affirmed', 'method': 'explicit_assertion',
                    'assertion_spans': [quote(targets[0])], 'conditions': []}
        if task['task_kind'] == 'entity':
            return {'entities': [{'class_iri': 'urn:Product',
                                  'mention': quote(targets[0], '产品甲'),
                                  'assertion_spans': [quote(targets[0], '分子量为537.18')]}]}
        return {'assertions': [{'value': quote(targets[0], '537.18'),
                                'assertion_spans': [quote(targets[0], '分子量')],
                                'assertion_status': 'affirmed'}]}

    worker = GenericExtractionRunner(
        {'urn:Product': {'label': '产品', 'properties': [
            {'iri': 'urn:weight', 'label': '分子量', 'datatype': 'decimal'},
        ]}}, TestTokenizer(), model, model_identity='separate-quotes',
        atomic_citations=True, record_level_targets=True,
        budget=TaskBudget(max_input_tokens=60000),
    )
    result = worker.run(ir)
    assert result.completion == 'complete', result.diagnostics
    assert calls == ['recall', 'verify_entity_types', 'recall', 'verify_binding']
    assert [(candidate.kind, candidate.positive_eligible) for candidate in result.candidates] == [
        ('entity', True), ('property', True),
    ]
    assert result.candidates[1].literal.raw_value == '537.18'


@pytest.mark.parametrize('compact', [False, True])
@pytest.mark.parametrize('verdict', ['metadata', 'real_condition', 'missing', 'duplicate',
                                    'wrong_index', 'conflicting', 'blank_reason'])
def test_recalled_conditions_need_explicit_complete_independent_review(tmp_path, compact, verdict):
    condition_text = '温度低于20℃时' if verdict == 'real_condition' else '记录标题'
    text = f'产品甲，状态：已完成；{condition_text}。'
    ir = source(tmp_path, text)
    unit = ir.evidence_units[0]
    anchor = ir.anchor(unit.evidence_id, 0, 3)
    subject = Candidate(candidate_id='product', kind='entity', class_iri='urn:Product',
                        text='产品甲', validation_status='passed', provenance=[DocumentProvenance(
                            anchors=[anchor], excerpts=['产品甲'],
                        )])
    seen = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        seen.append(request['stage'])
        condition = quote(targets[0], condition_text)
        proof = {'evidence_id': targets[0]['anchor']['evidence_id']}
        if request['stage'] == 'recall':
            return {'assertions': [{'value': quote(targets[0], '已完成'),
                                    'assertion_spans': [proof], 'conditions': [condition],
                                    'assertion_status': 'affirmed'}]}
        assert request['stage'] == 'verify_binding'
        assert 'condition_reviews' in schema['required']
        assert schema['properties']['condition_reviews']['minItems'] == 1
        assert schema['properties']['condition_reviews']['maxItems'] == 1
        assert schema['$defs']['ConditionReview']['properties']['condition_index']['enum'] == [0]
        assert len(request['candidate']['condition_anchors']) == 1
        reviews = [{'condition_index': 0, 'is_condition': verdict == 'real_condition',
                    'reason': '原文明示温度限制，必须保留。' if verdict == 'real_condition'
                              else '记录标题是表单元数据，不是断言成立的前提。'}]
        if verdict == 'missing':
            reviews = []
        elif verdict == 'duplicate':
            reviews *= 2
        elif verdict == 'wrong_index':
            reviews[0]['condition_index'] = 1
        elif verdict == 'blank_reason':
            reviews[0]['reason'] = '   '
        return {'supported': True,
                'subject_candidate_id': request['candidate']['subject']['candidate_id'],
                'assertion_status': 'affirmed', 'method': 'explicit_assertion',
                'assertion_spans': [proof], 'condition_reviews': reviews,
                # Even an explicit empty list cannot discard a real condition.
                'conditions': [condition] if verdict == 'conflicting' else []}

    worker = runner(model, compact_identifiers=compact)
    task = ExtractionTask(task_id='condition-review', task_kind='property',
                          subject=candidate_ref(subject), predicate_iri='urn:state',
                          scope=build_scope(ir, subject), budget=worker.budget,
                          target_evidence_ids=[unit.evidence_id], target_ranges=[
                              EvidenceRange(evidence_id=unit.evidence_id, start=0, end=len(text)),
                          ])
    if verdict in {'metadata', 'real_condition'}:
        candidate = worker.execute_task(task, ir, {'product': subject})[0]
        assert candidate.validation_status == 'passed'
        assert candidate.positive_eligible is (verdict == 'metadata')
        assert [ir.resolve(a) for a in candidate.condition_anchors] == (
            [condition_text] if verdict == 'real_condition' else []
        )
        review = candidate.bindings[0].record_mapping['condition_reviews'][0]
        assert ir.resolve(review['anchor']) == condition_text
        assert review['is_condition'] is (verdict == 'real_condition') and review['reason']
    else:
        with pytest.raises(PartialTaskFailure) as failure:
            worker.execute_task(task, ir, {'product': subject})
        assert not failure.value.candidates
        assert any('condition_review_' in reason for reason in failure.value.issues)
    assert seen == ['recall', 'verify_binding']


def table_ir(tmp_path):
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, [
        ('设备', '用途', '清洁方法'), ('釜甲', '反应', '水洗五分钟'), ('', '结晶', '乙醇洗十分钟'),
    ], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    table.cell(1, 0).merge(table.cell(2, 0))
    path = tmp_path / 'records.docx'
    document.save(path)
    return analyze_word_core(path).ir


def test_complete_logical_rows_preserve_merged_mentions_and_header_permissions(tmp_path):
    ir = table_ir(tmp_path)
    seen = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        assert request['stage'] == 'recall'
        seen.append([fragment['text'] for fragment in targets])
        return {'entities': []}

    result = runner(model, record_level_targets=True).run(ir)
    assert result.completion == 'complete', result.diagnostics
    assert seen == [['釜甲', '反应', '水洗五分钟'], ['釜甲', '结晶', '乙醇洗十分钟']]
    tasks = [event['task'] for event in result.tasks if event['status'] == 'complete']
    assert len({task['target_record_id'] for task in tasks}) == 2
    assert all(len(task['target_ranges']) == 3 for task in tasks)
    assert tasks[0]['target_ranges'][0] == tasks[1]['target_ranges'][0]


def test_oversized_record_is_incomplete_and_never_split_into_partial_rows(tmp_path):
    ir = table_ir(tmp_path)
    worker = runner(lambda *_: pytest.fail('oversized record must not dispatch'))
    index = RecordIndex(ir)
    record_id, regions = record_targets(index, [
        EvidenceRange(evidence_id=unit.evidence_id) for unit in ir.evidence_units if unit.text
    ])[0]
    task = ExtractionTask(task_id='record', task_kind='entity', target_record_id=record_id,
                          target_class_iris=['urn:Product'], target_ranges=regions,
                          target_evidence_ids=[value.evidence_id for value in regions],
                          predicate_definition=worker.entity_menu(['urn:Product'])['predicate_definition'],
                          budget=TaskBudget(max_input_tokens=1))
    assert worker.split_output_task(task) == []
    with pytest.raises(ValueError, match='context_budget_exceeded'):
        worker.execute_task(task, ir, {})
    partial = task.model_copy(update={'target_ranges': regions[:1],
                                      'target_evidence_ids': [regions[0].evidence_id]})
    with pytest.raises(ValueError, match='incomplete_record_target'):
        worker.execute_task(partial, ir, {})


def test_new_protocol_and_record_targets_invalidate_the_recognition_identity(tmp_path):
    ir = source(tmp_path, '产品甲')
    worker = runner(None)
    original = worker.input_id(ir)
    worker.atomic_citations = False
    assert worker.input_id(ir) != original
    worker.atomic_citations = True
    worker.record_level_targets = True
    assert worker.input_id(ir) != original


@pytest.mark.parametrize('escape_target', [False, True])
def test_route_caption_keeps_sibling_structure_as_binding_only(tmp_path, escape_target):
    document = Document()
    document.add_heading('工艺', level=1)
    document.add_heading('工艺描述', level=2)
    document.add_paragraph('3.1.1合成路线图')
    document.add_heading('步骤1：中间体甲的合成', level=2)
    document.add_paragraph('原料经反应得到中间体甲。')
    document.add_heading('步骤2：成品的合成', level=2)
    document.add_paragraph('中间体甲经精制得到成品。')
    document.add_heading('其他项目', level=1)
    document.add_heading('步骤1：其他产品的合成', level=2)
    path = tmp_path / 'route.docx'
    document.save(path)
    ir = analyze_word_core(path).ir
    seen = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        context = json.loads(request['context'])
        related = [fragment for fragment in context['fragments']
                   if fragment['purpose'] == 'related_section_reference_context']
        assert [fragment['text'] for fragment in related] == [
            '步骤1：中间体甲的合成', '原料经反应得到中间体甲。',
            '步骤2：成品的合成', '中间体甲经精制得到成品。',
        ]
        assert all(not fragment['fact_eligible'] and fragment['binding_eligible']
                   for fragment in related)
        seen.append(request['stage'])
        if request['stage'] == 'verify_entity_types':
            return accept_fixture_types(request)
        return {'entities': [{'class_iri': 'urn:Route',
                              'mention': quote(related[0] if escape_target else targets[0]),
                              'assertion_spans': [quote(fragment) for fragment in related]}]}

    schema = {'urn:Route': {'label': '合成路线', 'relationships': [
        {'iri': 'urn:hasStep', 'range': ['urn:Step']},
    ]}, 'urn:Step': {'label': '合成步骤'}}
    worker = GenericExtractionRunner(schema, TestTokenizer(), model, model_identity='route-test',
                                     atomic_citations=True, record_level_targets=True,
                                     budget=TaskBudget(max_input_tokens=60000))
    unit = next(unit for unit in ir.evidence_units if unit.text == '3.1.1合成路线图')
    identity, regions = record_targets(RecordIndex(ir), [
        EvidenceRange(evidence_id=unit.evidence_id),
    ])[0]
    task = ExtractionTask(task_id='route', task_kind='entity', target_record_id=identity,
                          target_ranges=regions, target_evidence_ids=[unit.evidence_id],
                          budget=worker.budget, **worker.entity_menu(['urn:Route']))
    if escape_target:
        with pytest.raises(PartialTaskFailure, match='unknown_or_disallowed_citation_source'):
            worker.execute_task(task, ir, {})
        assert seen == ['recall']
    else:
        candidates = worker.execute_task(task, ir, {})
        assert seen == ['recall', 'verify_entity_types']
        assert len(candidates) == 1 and candidates[0].positive_eligible
        assert candidates[0].text == '3.1.1合成路线图'


def test_field_group_lookup_requires_independent_cross_record_owner_verification(tmp_path):
    ir = source(
        tmp_path, '本报告描述产品甲。', '分子量：537.18', '项目名称：产品甲', '性状：白色固体',
    )
    units = [unit for unit in ir.evidence_units if unit.text]
    anchor = ir.anchor(units[0].evidence_id, 5, 8)
    assert ir.resolve(anchor) == '产品甲'
    subject = Candidate(candidate_id='product', kind='entity', class_iri='urn:Product',
                        text='产品甲', validation_status='passed', provenance=[DocumentProvenance(
                            anchors=[anchor], excerpts=['产品甲'],
                        )])
    scope = field_group_scope(build_scope(ir, subject), subject, RecordIndex(ir))
    assert {region.evidence_id for region in scope.reference_ranges} == {
        unit.evidence_id for unit in units[1:]
    }
    calls = []

    def model(system, user, schema, budget):
        request, task, targets = parts(user)
        calls.append(request['stage'])
        context = json.loads(request['context'])
        owner = next(fragment for fragment in context['fragments']
                     if fragment['text'] == '项目名称：产品甲')
        assert owner['fact_eligible'] is False
        assert owner['binding_eligible'] is True
        if request['stage'] == 'recall':
            return {'assertions': [{'value': quote(targets[0], '537.18'),
                                    'assertion_spans': [quote(targets[0])],
                                    'assertion_status': 'affirmed'}]}
        candidate = request['candidate']
        if request['stage'] == 'verify_binding':
            return {'supported': True, 'subject_candidate_id': candidate['subject']['candidate_id'],
                    'assertion_status': 'affirmed', 'method': 'explicit_assertion',
                    'assertion_spans': [quote(targets[0])], 'conditions': []}
        # Same-name retrieval is insufficient: failing either-end reference
        # verification must reject the otherwise valid numeric property.
        assert request['stage'] == 'verify_reference'
        return {'supported': False, 'subject_candidate_id': 'product',
                'assertion_spans': [quote(targets[0])], 'reason': '归属尚未证实。'}

    worker = runner(model)
    regions = [EvidenceRange(evidence_id=units[1].evidence_id, start=0, end=len(units[1].text))]
    task = ExtractionTask(task_id='weight', task_kind='property', subject=candidate_ref(subject),
                          predicate_iri='urn:weight', scope=scope, target_ranges=regions,
                          target_evidence_ids=[units[1].evidence_id], budget=worker.budget)
    with pytest.raises(PartialTaskFailure) as failure:
        worker.execute_task(task, ir, {'product': subject})
    assert calls == ['recall', 'verify_binding', 'verify_reference']
    assert not failure.value.candidates
    assert 'reference_not_supported' in failure.value.issues


@pytest.mark.parametrize('field_name,expected', [('HRS-X12', True), ('HRS-X1', False)])
def test_owner_lookup_accepts_role_suffix_but_not_code_prefixes(tmp_path, field_name, expected):
    ir = source(tmp_path, 'HRS-X12项目', '分子量：537.18', '项目名称：' + field_name)
    anchor = ir.anchor(ir.evidence_units[0].evidence_id)
    subject = Candidate(candidate_id='product', kind='entity', class_iri='urn:Product',
                        text=ir.resolve(anchor), validation_status='passed',
                        provenance=[DocumentProvenance(anchors=[anchor],
                                                       excerpts=[ir.resolve(anchor)])])
    scope = field_group_scope(build_scope(ir, subject), subject, RecordIndex(ir))
    assert bool(scope.reference_ranges) is expected
    assert not subject.identity  # A lookup never becomes a global identity decision.


def test_verified_first_level_relation_gets_properties_before_broad_entity_scan(tmp_path):
    ir = source(tmp_path, '报告描述试验品甲产品，分子量：537.18。', '无关附录记录')
    calls = []
    schema = {
        'urn:Report': {'relationships': [
            {'iri': 'urn:describes', 'label': '描述产品', 'range': ['urn:Product']},
        ]},
        'urn:Product': {'label': '产品', 'properties': [
            {'iri': 'urn:weight', 'label': '分子量', 'datatype': 'string'},
        ]},
    }

    def model(system, user, response_schema, budget):
        request, task, targets = parts(user)
        calls.append((request['stage'], task['task_kind'], [item['text'] for item in targets]))
        if request['stage'] == 'verify_entity_types':
            return accept_fixture_types(request)
        if request['stage'] == 'verify_binding':
            candidate = request['candidate']
            return {'supported': True, 'subject_candidate_id': candidate['subject']['candidate_id'],
                    'object_candidate_id': (candidate.get('object') or {}).get('candidate_id'),
                    'assertion_status': 'affirmed', 'method': 'explicit_assertion',
                    'assertion_spans': [quote(targets[0])], 'conditions': []}
        target = next((item for item in targets if '试验品甲' in item['text']), None)
        if task['task_kind'] == 'entity':
            return {'entities': ([{
                'class_iri': 'urn:Product', 'mention': quote(target, '试验品甲'),
            }] if target else [])}
        if task['task_kind'] == 'relationship':
            return {'assertions': ([{
                'object_candidate_id': task['object_candidates'][0]['candidate_id'],
                'assertion_status': 'affirmed', 'assertion_spans': [quote(target)],
            }] if target else [])}
        return {'assertions': ([{'value': quote(target, '537.18'), 'assertion_status': 'affirmed',
                                'assertion_spans': [quote(target)]}] if target else [])}

    worker = GenericExtractionRunner(
        schema, TestTokenizer(), model, model_identity='priority-test',
        atomic_citations=True, relationship_priority=True, record_level_targets=True,
        budget=TaskBudget(max_input_tokens=60000, max_tasks=64),
    )
    result = worker.run(ir, effective_class='urn:Report')
    property_turn = next(index for index, call in enumerate(calls)
                         if call[:2] == ('recall', 'property'))
    broad_turn = next(index for index, call in enumerate(calls)
                      if call[:2] == ('recall', 'entity') and '无关附录记录' in call[2])
    assert property_turn < broad_turn
    assert any(candidate.kind == 'property' and candidate.positive_eligible
               and candidate.literal.raw_value == '537.18' for candidate in result.candidates)
    assert result.branch_progress['urn:describes']['positive_property_count'] > 0
