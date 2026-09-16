from docx import Document

from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure
from app.services.extraction.template_structure_builder import (
    build_template_structure,
    origin_status,
    template_structure_titles,
)


def test_offline_skeleton_has_stable_origins_and_separate_label_value_anchors(tmp_path):
    doc = Document()
    doc.add_heading("产品", 1)
    doc.add_paragraph("  规格：250 mg ")
    doc.add_paragraph("备注：")
    path = tmp_path / "sample.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path), role="template_sample")
    first = build_template_structure(ir)
    assert first == build_template_structure(ir)
    section = first[0]
    assert section["title"] == "产品"
    field = section["groups"][0]["candidates"][0]
    assert field["label"] == "规格"
    assert ir.resolve(field["origin"]["label_anchor"]) == "规格"
    assert ir.resolve(field["origin"]["value_anchor"]) == "250 mg"
    empty = section["groups"][0]["candidates"][1]
    assert empty["origin"]["value_anchor"] is None
    assert origin_status(field["origin"], ir) == "valid"
    assert origin_status({**field["origin"], "document_hash": "0" * 64}, ir) == "invalid"


def test_table_label_and_value_in_different_cells_and_prose_have_origins(tmp_path):
    doc = Document()
    doc.add_heading("产品", 1)
    doc.add_paragraph("此段是待作者编写的叙述。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "字段", "取值"
    table.cell(1, 0).text, table.cell(1, 1).text = "规格", "250 mg"
    path = tmp_path / "fields.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path), role="template_sample")
    sections = build_template_structure(ir)
    fields = [c for s in sections for g in s["groups"] for c in g["candidates"]]
    field = next(c for c in fields if c["label"] == "规格")
    assert ir.resolve(field["origin"]["label_anchor"]) == "规格"
    assert ir.resolve(field["origin"]["value_anchor"]) == "250 mg"
    assert (
        field["origin"]["label_anchor"]["evidence_id"]
        != (field["origin"]["value_anchor"]["evidence_id"])
    )
    assert any(c["label"] == "段落行文" for c in fields)


def test_multicolumn_header_does_not_pair_first_two_columns_as_key_value(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for col, label in enumerate(("对象", "规格", "备注")):
        table.cell(0, col).text = label
    path = tmp_path / "columns.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path), role="template_sample")
    fields = [c for s in build_template_structure(ir) for g in s["groups"] for c in g["candidates"]]
    assert all(c["origin"]["value_anchor"] is None for c in fields)


def test_business_titles_use_merged_heading_and_each_nested_tables_own_caption(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "SECTIONⅠ 部分1"
    cell = table.cell(1, 0).merge(table.cell(1, 1))
    cell.text = "Describing the subject of risk evaluation 风险评估对象基本描述"
    for code in ("642", "646"):
        cell.add_paragraph(f"HRS-1234 使用 {code} 车间设备见下表：")
        nested = cell.add_table(rows=2, cols=2)
        nested.cell(0, 0).text = "设备编号"
        nested.cell(0, 1).text = "设备名称"
        nested.cell(1, 0).text = "EQ-1"
        nested.cell(1, 1).text = "反应釜"
    path = tmp_path / "nested.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path), role="template_sample")
    before = ir.model_dump(mode="json")
    sections = build_template_structure(ir)
    groups = sections[0]["groups"]
    assert [g["title"] for g in groups] == [
        "风险评估对象基本描述", "642 车间设备表", "646 车间设备表",
    ]
    assert all(origin_status(g["origin"], ir) == "valid" for g in groups)
    assert ir.model_dump(mode="json") == before
    assert template_structure_titles({"schema_version": 2}, before) == {
        g["id"]: g["title"] for g in groups
    }
    assert template_structure_titles({"schema_version": 2}, {"title": "无有效锚点"}) == {}


def test_table_caption_cannot_leak_across_cells_or_reuse_for_next_table(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "999 车间设备见下表："
    cell = table.cell(0, 1)
    for index in range(3):
        if index == 1:
            cell.add_paragraph("642 车间设备见下表：")
        nested = cell.add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "编号"
        nested.cell(0, 1).text = "名称"
    path = tmp_path / "separate.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path), role="template_sample")
    groups = build_template_structure(ir)[0]["groups"]
    nested = [g for g in groups if len(g["origin"]["label_anchor"]["table_path"]) > 1]
    assert [g["title"] for g in nested] == ["编号、名称", "642 车间设备表", "编号、名称"]
