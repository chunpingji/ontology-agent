"""Current-run referent recall returns evidence, never an identity conclusion."""

import json
from dataclasses import replace

import pytest
from docx import Document

from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided import tool_runtime as runtime
from app.services.extraction.ontology_guided.claim_protocol import (
    EntityDependencyView,
    EntityProposal,
    ExtractionProfile,
    IdentifierProposal,
    Quote,
)
from app.services.extraction.ontology_guided.context import ContextFragment, assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    LocalMenu,
    SlotSpec,
    SubjectRef,
    TraversalScope,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.mentions import MentionRegistry
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.ontology_guided.tool_contracts import ToolCall


@pytest.fixture
def referent_context(tmp_path):
    document = Document()
    for text in ["批次甲：HRS-9267 粗品", "批次乙：HRS-9267粗品", "本品", "未知材料"]:
        document.add_paragraph(text)
    for number in range(9):
        document.add_paragraph(f"批次{number}：HRS-9267粗品")
    path = tmp_path / "referents.docx"
    document.save(path)
    index = RecordIndex(build_document_ir(path, parse_docx_structure(path)))
    units = list(index.ir.evidence_units)
    record = index.records_by_evidence[units[0].evidence_id][0]
    subject = SubjectRef(entity_id="entity-0", revision=1, class_iri="urn:CrudeProduct")
    slot = SlotSpec(iri="urn:storage", label="Storage")
    menu = LocalMenu(
        menu_id="menu", ontology_snapshot_id="ontology", subject=subject, properties=[slot],
    )
    task = RecognitionTask.create(
        subject=subject, predicate_iri=slot.iri, predicate_kind="property",
        record_id=record.record_id, phase=1, hop=0, dependency_hash="dependencies",
    )
    target = VerificationTarget.create(
        run_fingerprint="run", claim_ref=VersionedRef(id="placeholder", revision=1),
        task_id=task.task_id, check_kind="predicate_entailment",
        document_context=DocumentContext(
            document_hash=index.ir.document_hash, document_class_iri="urn:Document",
            root_ref=VersionedRef(id="root", revision=1),
        ),
        subject_ref=subject, predicate_iri=slot.iri, ontology_hash="ontology",
        source_scope_hash="scope", context_hash="initial",
    )
    context = assemble_context(target, record.record_id, index, predicate=slot)
    context = context.model_copy(update={"fragments": [
        ContextFragment(
            anchor=index.ir.anchor(unit.evidence_id, 0, len(unit.text)),
            text=unit.text, purpose="binding", fact_eligible=False,
        ) for unit in units
    ]})
    registry = MentionRegistry(index.ir, index)
    dependencies = {}
    for number, unit in enumerate(units):
        name = unit.text.split("：")[-1]
        start = unit.text.index(name)
        registry.register(
            evidence_id=unit.evidence_id, start=start, end=start + len(name), text=name,
            record_view_ref=index.record_views_by_id[
                index.records_by_evidence[unit.evidence_id][0].record_id
            ].record_view_id,
        )
        if "粗品" not in name:
            continue
        proposal = EntityProposal(
            local_id=f"local-{number}", class_iri="urn:CrudeProduct", representation="mention",
            mentions=[Quote(evidence_id=unit.evidence_id, text=name, context_text=unit.text)],
            record_components=[], identifier_claims=[IdentifierProposal(
                predicate_iri="urn:batch",
                value_quote=Quote(evidence_id=unit.evidence_id,
                                  text=unit.text.split("：")[0], context_text=None),
            )],
        )
        values = dict(
            entity_ref=VersionedRef(id=f"entity-{number}", revision=1),
            class_iri="urn:CrudeProduct", grounding_kind="mention", root_origin=None,
            proposal=proposal,
            source_refs=[index.ir.anchor(unit.evidence_id, start, start + len(name))],
            dependency_refs=[],
        )
        dependencies[proposal.local_id] = EntityDependencyView(
            **values, content_hash=evidence_hash(values),
        )
    return runtime.ToolContext(
        task=task, context=context, index=index, menu=menu, profile=ExtractionProfile(),
        scope=TraversalScope.create([]), stage="discovery", recovery_kind="none",
        frozen_claims={}, local_ref_map={}, entity_dependencies=dependencies,
        registered_mentions={item.mention_id: item for item in registry.mentions},
        limits=runtime.ToolLimits(max_result_tokens=100_000),
        measure_result_tokens=len, reference_resolution=True,
    )


def recall(ctx, text="HRS-9267 粗品", **overrides):
    mention = next(mention for mention in ctx.registered_mentions.values() if mention.text == text)
    arguments = dict(mention_refs=[mention.mention_id], class_iris=["urn:CrudeProduct"],
                     scope_id=ctx.scope.scope_id)
    arguments.update(overrides)
    return runtime.dispatch_tool(ToolCall(
        call_id="recall", name="find_referent_candidates", arguments_json=json.dumps(arguments),
    ), ctx)


def test_physical_match_precedes_whitespace_candidates_without_merging(referent_context):
    ctx = referent_context
    before = {key: value.model_dump() for key, value in ctx.entity_dependencies.items()}
    result = recall(ctx)
    assert result.status == "ok", result
    assert result.data.candidates[0].candidate_entity_ref.id == "entity-0"
    assert result.data.candidates[0].reason == "physical_mention"
    assert result.data.candidates[1].reason == "normalized_name"
    assert result.data.candidates[1].candidate_entity_ref.id == "entity-1"
    assert result.data.identity_status == "not_checked"
    assert result.data.candidates[0].distinctions == ["批次甲"]
    assert result.data.candidates[1].distinctions == ["批次乙"]
    assert before == {key: value.model_dump() for key, value in ctx.entity_dependencies.items()}
    for candidate in result.data.candidates:
        assert candidate.role == "binding" and candidate.fact_eligible is False
        assert candidate.source_texts == [
            ctx.index.ir.resolve(ref) for ref in candidate.source_refs
        ]
        assert candidate.mention_refs


def test_anaphora_returns_competing_candidates_instead_of_choosing_nearest(referent_context):
    result = recall(referent_context, text="本品")
    assert result.status == "ok"
    assert len(result.data.candidates) == 8
    assert all(item.reason == "anaphora_context" for item in result.data.candidates)
    assert result.data.truncated and result.data.excluded_count == 3
    assert result.data.identity_status == "not_checked"


def test_same_name_truncation_and_no_match_are_separate(referent_context):
    result = recall(referent_context)
    assert len(result.data.candidates) == 8
    assert result.data.truncated and result.data.excluded_count == 3
    empty = recall(referent_context, text="未知材料")
    assert empty.status == "no_match"
    assert empty.data.candidates == [] and not empty.data.truncated
    assert empty.data.identity_status == "not_checked"


@pytest.mark.parametrize("arguments,path", [
    ({"mention_refs": ["made-up-mention"]}, "/mention_refs/0"),
    ({"scope_id": "other-run"}, "/scope_id"),
    ({"class_iris": ["urn:FinishedProduct"]}, "/class_iris"),
    ({"class_iris": ["urn:CrudeProduct", "urn:CrudeProduct"]}, "/class_iris"),
])
def test_unknown_and_out_of_scope_arguments_are_blocked(referent_context, arguments, path):
    result = recall(referent_context, **arguments)
    assert result.status == "blocked" and result.data is None
    assert result.issues[0].field_path == path
    assert not result.evidence_refs


def test_other_run_mention_and_changed_dependency_are_invalid(referent_context):
    ctx = referent_context
    mention = next(iter(ctx.registered_mentions.values()))
    changed = mention.model_copy(update={"analysis_id": "other-run"})
    result = recall(replace(ctx, registered_mentions={mention.mention_id: changed}))
    assert result.status == "blocked"
    assert result.issues[0].code == "reference_version_mismatch"
    dependency = next(iter(ctx.entity_dependencies.values()))
    result = recall(replace(ctx, entity_dependencies={
        "changed": dependency.model_copy(update={"content_hash": "stale"}),
    }))
    assert result.status == "blocked"
    assert result.issues[0].code == "reference_version_mismatch"


def test_dependency_sources_do_not_grant_implicit_evidence_permission(referent_context):
    ctx = referent_context
    reduced = ctx.context.model_copy(update={"fragments": ctx.context.fragments[:1]})
    result = recall(replace(ctx, context=reduced))
    assert result.status == "blocked"
    assert result.issues[0].code == "reference_outside_scope"
    assert result.data is None and not result.evidence_refs


def test_result_budget_remains_enforced(referent_context):
    ctx = replace(referent_context, limits=runtime.ToolLimits(max_result_tokens=1))
    result = recall(ctx)
    assert result.status == "blocked" and result.data is None
    assert result.issues[0].code == "result_budget_exceeded"
    assert not result.evidence_refs


def test_same_spelling_with_other_type_is_not_a_candidate(referent_context):
    ctx = referent_context
    dependency = next(iter(ctx.entity_dependencies.values()))
    values = dependency.model_dump(mode="json", exclude={"content_hash"})
    values["class_iri"] = values["proposal"]["class_iri"] = "urn:FinishedProduct"
    other_type = EntityDependencyView(**values, content_hash=evidence_hash(values))
    result = recall(replace(ctx, entity_dependencies={"finished": other_type}))
    assert result.status == "no_match" and result.data.candidates == []


def test_conflicting_entity_revisions_are_not_silently_selected(referent_context):
    ctx = referent_context
    dependency = next(iter(ctx.entity_dependencies.values()))
    values = dependency.model_dump(mode="json", exclude={"content_hash"})
    values["entity_ref"]["revision"] = 2
    changed = EntityDependencyView(**values, content_hash=evidence_hash(values))
    result = recall(replace(ctx, entity_dependencies={"v1": dependency, "v2": changed}))
    assert result.status == "blocked"
    assert result.issues[0].code == "reference_version_mismatch"
    assert result.data is None


def test_new_policy_gate_keeps_old_runs_catalog_unchanged(referent_context):
    enabled = referent_context
    name = "find_referent_candidates"
    assert name in {item["name"] for item in runtime.build_tool_definitions(enabled, "discovery")}
    disabled = replace(enabled, reference_resolution=False)
    assert name not in {
        item["name"] for item in runtime.build_tool_definitions(disabled, "discovery")
    }
    assert recall(disabled).issues[0].code == "tool_not_allowed"
    verification = replace(enabled, stage="verification")
    assert name in {
        item["name"] for item in runtime.build_tool_definitions(verification, "verification")
    }
    assert recall(replace(enabled, stage="finalize")).issues[0].code == "tool_not_allowed"
