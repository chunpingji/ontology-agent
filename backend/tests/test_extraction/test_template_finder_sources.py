"""Actual Finder and anchor behavior on synthetic Word originals."""

import json
from pathlib import Path

import pytest
from docx import Document

from app.services.extraction.word_analysis import analyze_word_core
from app.services.ontology_engine import OntologyEngine
from app.services.template_finder.policy import BUILTINS, Binding
from app.services.template_finder.runner import freeze_ontology, run
from app.services.template_finder.sources import source_for


@pytest.mark.parametrize(("class_name", "headers", "values"), [
    ("SafetyRiskAssessment", ["风险环节", "风险描述", "控制措施"],
     ["投料操作", "粉尘暴露", "密闭投料"]),
    ("QualityRiskAssessment", ["风险环节", "质量风险", "控制措施"],
     ["干燥操作", "水分超标", "监控干燥终点"]),
    ("ProductionRiskAssessment",
     ["风险类型", "风险因素", "控制前风险水平", "控制后风险水平", "可追溯性",
      "风险状态", "控制措施"],
     ["HazID-1", "粉尘暴露", "高", "低", "规程A", "已控制", "密闭投料"]),
])
def test_risk_tables_keep_declared_properties_without_domain(
    tmp_path, engine, class_name, headers, values,
):
    profile = json.loads(BUILTINS["cmc_baseline_v1"].read_text())
    binding = Binding(profile["root_classes"][0], profile)
    doc = Document()
    doc.add_heading("HRS-1234 CMC 报告", 0)
    table = doc.add_table(rows=2, cols=len(headers))
    for row, texts in zip(table.rows, [headers, values]):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    path = tmp_path / "risk.docx"
    doc.save(path)
    analysis = analyze_word_core(path)
    graph = run(analysis, binding, freeze_ontology(engine, binding))
    risk = next(n for n in graph["relationships"]
                if n["object_class_iri"].endswith("/" + class_name))
    props = {p["iri"].rsplit("/", 1)[-1]: p for p in risk["object_data_properties"] if p["iri"]}
    assert props["riskCategory"]["value"] == values[0]
    assert props["controlMeasure"]["value"] == values[-1]
    for name, column in (("riskCategory", 0), ("controlMeasure", len(headers) - 1)):
        anchors = props[name]["source"]["anchors"]
        assert len(anchors) == 1
        assert anchors[0]["row_index"] == 1
        assert anchors[0]["column_index"] == column
        assert analysis.ir.resolve(anchors[0]) == values[column]
    if class_name == "ProductionRiskAssessment":
        raw = [p for p in risk["object_data_properties"] if p["iri"] is None]
        assert [p["value"] for p in raw] == values[1:-1]
        for prop, column in zip(raw, range(1, len(headers) - 1)):
            anchor = prop["source"]["anchors"][0]
            assert anchor["column_index"] == column
            assert analysis.ir.resolve(anchor) == prop["value"]


@pytest.mark.parametrize("invalid_property", ["removedRiskCategory", "hasStep"])
def test_missing_or_object_property_cannot_be_used_as_risk_data(
    tmp_path, engine, monkeypatch, invalid_property,
):
    from app.services.template_finder.finders import DP

    invalid_iri = "https://ontology.pharma-gmp.cn/slpra/drug-development/" + invalid_property
    monkeypatch.setitem(DP, "riskCategory", invalid_iri)
    profile = json.loads(BUILTINS["cmc_baseline_v1"].read_text())
    binding = Binding(profile["root_classes"][0], profile)
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for row, texts in zip(table.rows, [["风险环节", "风险描述"], ["投料", "粉尘暴露"]]):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    path = tmp_path / "invalid-risk.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="Finder 属性不在当前本体菜单中"):
        run(analyze_word_core(path), binding, freeze_ontology(engine, binding))


@pytest.fixture
def engine(tmp_path):
    engine = OntologyEngine(
        ontology_dir=Path(__file__).resolve().parents[3] / "ontology/slpra",
        store_path=tmp_path / "ontology.sqlite3",
    )
    engine.load()
    yield engine
    engine.close()


def test_process_and_yield_have_independent_physical_sources(tmp_path, engine):
    profile = json.loads(BUILTINS["cmc_baseline_v1"].read_text())
    binding = Binding(profile["root_classes"][0], profile)
    doc = Document()
    doc.add_heading("HRS-1234 CMC 报告", 0)
    doc.add_heading("合成路线图", 1)
    doc.add_heading("中间体A制备", 1)
    table = doc.add_table(rows=3, cols=2)
    for row, cells in zip(
        table.rows,
        [
            ["工序", "具体操作"],
            ["投料", "加入起始物料并缓慢搅拌，保持温度在二十摄氏度以下直至完全溶解。"],
            ["反应", "升温至五十摄氏度保温反应两小时，然后降温过滤收集得到的固体。"],
        ],
    ):
        for cell, text in zip(row.cells, cells):
            cell.text = text
    doc.add_heading("收率汇总", 1)
    table = doc.add_table(rows=2, cols=3)
    for row, cells in zip(
        table.rows, [["名称", "参考得量范围", "参考收率范围"], ["中间体A", "10～12", "80～90"]]
    ):
        for cell, text in zip(row.cells, cells):
            cell.text = text
    path = tmp_path / "synthesis.docx"
    doc.save(path)
    analysis = analyze_word_core(path)
    graph = run(analysis, binding, freeze_ontology(engine, binding))
    route = next(
        n for n in graph["relationships"] if n["object_class_iri"].endswith("/SynthesisRoute")
    )
    step = route["sub_relationships"][0]
    props = {p["iri"].rsplit("/", 1)[-1]: p for p in step["object_data_properties"]}
    condition = props["reactionConditions"]["source"]
    mass = props["outputMassRange_kg"]["source"]
    yield_source = props["yieldRangePercent"]["source"]
    assert condition["kind"] == "computed"
    assert {a["table_path"][0] for a in condition["anchors"]} == {"table:0"}
    assert {a["table_path"][0] for a in mass["anchors"]} == {"table:1"}
    assert mass["anchors"][0]["column_index"] == 1
    assert yield_source["anchors"][0]["column_index"] == 2
    assert analysis.ir.resolve(mass["anchors"][0]) == "10～12"
    assert analysis.ir.resolve(yield_source["anchors"][0]) == "80～90"
    assert not props["stepOrder"]["source"]["anchors"]
    assert props["stepOrder"]["source"]["kind"] == "computed"

    def walk(nodes):
        for node in nodes:
            assert "conflict" not in node
            for prop in node["object_data_properties"]:
                for anchor in prop["source"]["anchors"]:
                    analysis.ir.resolve(anchor)
            walk(node["sub_relationships"])

    walk(graph["relationships"])


def test_merged_nested_cells_and_fragment_offsets_never_search_values(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("中文😀前半")
    paragraph.add_run().add_break()
    paragraph.add_run("后半😀中文")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "相同值😀"
    table.cell(0, 1).text = "相同值😀"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套😀"
    path = tmp_path / "cells.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    merged = source_for(
        ir, {"kind": "table_cell", "table_path": ["table:0"], "row": 1, "column": 0}
    )
    assert len(merged["anchors"]) == 1
    assert merged["anchors"][0]["row_index"] == 0
    assert merged["anchors"][0]["column_index"] == 0
    assert ir.resolve(merged["anchors"][0]) == "相同值😀"
    unit = next(u for u in ir.evidence_units if u.text == "嵌套😀")
    located = source_for(
        ir, {"kind": "table_cell", "table_path": unit.table_path, "row": 0, "column": 0}
    )
    assert ir.resolve(located["anchors"][0]) == "嵌套😀"
    assert source_for(ir, "§ 不明确")["anchors"] == []
    assert (
        source_for(ir, {"kind": "table_cell", "table_path": ["table:0"], "row": 99, "column": 0})[
            "anchors"
        ]
        == []
    )
