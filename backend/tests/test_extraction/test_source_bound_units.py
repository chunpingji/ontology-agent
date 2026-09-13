"""Unit evidence through actual discovery/review, proof, and public projection paths."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.services.extraction.literal_normalizer import (
    UNIT_REGISTRY_VERSION,
    LiteralNormalizationError,
    normalize_literal,
)
from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import SlotSpec
from app.services.extraction.ontology_guided.ontology_plan import compile_local_menu
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.repair_adapter import EvidenceRepairAdapter
from app.services.extraction.ontology_guided.value_constraints import (
    UNIT_NORMALIZATION_VERSION,
    XSD,
)
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_evidence_repair import _ontology
from tests.test_extraction.test_joint_evidence_validation import _fixture


def quantity_run(
    tmp_path, monkeypatch, source, *, raw="3.8", unit="kg", header=None,
    unit_source=None, target_unit="kg", unit_verdict="supported", role_verdict="supported",
    unit_support=True, datatype="decimal", unit_context=None, owner=None,
):
    _, _, task, target, _ = _fixture(tmp_path)
    document = Document()
    if header:
        table = document.add_table(rows=2, cols=2)
        for row, texts in zip(table.rows, [header, [source, "7"]], strict=True):
            for cell, text in zip(row.cells, texts, strict=True):
                cell.text = text
    else:
        document.add_paragraph(source)
    path = tmp_path / "quantity.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    record = next(r for r in index.records if any(u.text == source for u in r.source_units))
    predicate = SlotSpec(
        iri="urn:plannedBatchSizeMin", label="预计批量下限", canonical_unit=target_unit,
        datatype_iris=[XSD + datatype],
    )
    task = task.model_copy(update={
        "record_id": record.record_id, "predicate_iri": predicate.iri, "predicate_kind": "property",
    })
    target = target.model_copy(update={
        "predicate_iri": predicate.iri,
        "document_context": target.document_context.model_copy(update={
            "document_hash": analysis.ir.document_hash,
        }),
    })
    owner_refs = []
    if owner:
        subject = task.subject.model_copy(update={"entity_id": "plan", "is_document_root": False})
        task = task.model_copy(update={"subject": subject})
        target = target.model_copy(update={"subject_ref": subject})
        value_unit = next(u for u in record.source_units if u.text == source)
        start = source.index(owner)
        owner_refs = [analysis.ir.anchor(value_unit.evidence_id, start, start + len(owner))]
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    calls = []

    def respond(_client, *, user, schema, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        assert request["unit_normalization_version"] == UNIT_NORMALIZATION_VERSION
        value = next(f for f in request["fragments"] if f["text"] == source)
        quote = {"evidence_id": value["evidence_id"]}
        binding = next((b for b in request["field_bindings"] if any(
            r["evidence_id"] == value["evidence_id"] for r in b["target_value_refs"]
        )), None)
        if request["stage"] == "discovery":
            return {"proposals": [{
                "kind": "property", "value_quote": {**quote, "text": raw},
                "bridge_kind": "role_mapped_table" if header else "explicit_assertion",
                **({"field_binding_id": binding["field_binding_id"]} if binding else {}),
            }]}
        chosen = next(f for f in request["fragments"] if f["text"] == (unit_source or source))
        candidate = request["candidates"][0]
        # Verify the actual transport accepts units from headers, with precise
        # text and only authorized source IDs (not a fact-only citation menu).
        if unit:
            assert chosen["evidence_id"] in schema["$defs"]["UnitQuote"]["properties"][
                "evidence_id"
            ]["enum"]
        return {"verifications": [{
            "candidate_id": candidate["candidate_id"], "target_id": candidate["target_id"],
            "role_verdict": role_verdict,
            "subject_binding": {
                "verdict": "supported",
                "support": [{"evidence_id": quote["evidence_id"]}] if owner else [],
                "local_support": [quote] if owner else [],
            },
            **{key + "_verdict": "supported" for key in (
                "predicate", "applicability", "counterevidence", "bridge",
            )},
            "field_role_support": ([{"evidence_id": r["evidence_id"]}
                                    for r in binding["label_refs"]] if binding else [quote]),
            "bridge_support": [quote], "predicate_support": [quote],
            "condition_support": [], "counterevidence_support": [],
            "unit_verdict": unit_verdict,
            "source_unit_quote": {"evidence_id": chosen["evidence_id"], "text": unit,
                                  "context_text": unit_context}
            if unit else None,
            "unit_binding_support": [quote, {"evidence_id": chosen["evidence_id"]}]
            if unit_support else [],
            "reason": "原文支持当前计划的批量下限。",
        }]}

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    context = assemble_context(
        target, record.record_id, index, repair_enabled=True,
        subject_evidence_refs=owner_refs, subject_label=owner or "",
    )
    adapter = EvidenceRepairAdapter(object(), model_identity="unit-fixture")
    result = adapter.inspect(task, context, predicate, menu)
    return result, context, adapter, task, predicate, menu, analysis, index, calls


@pytest.mark.parametrize("unit", ["kg", "KG", "Kg", "千克", "公斤", "ｋｇ", "ＫＧ"])
def test_range_lower_endpoint_inherits_proven_unit(tmp_path, monkeypatch, unit):
    result, context, *_ = quantity_run(
        tmp_path, monkeypatch, f"本次备样用于Ⅰ期临床，批量3.8- 6.6{unit}。", unit=unit,
    )
    prop = result.properties[0]
    assert prop.policy_eligible and prop.structural_valid and prop.model_supported
    assert prop.raw_value == "3.8" and prop.normalized_value == "3.8"
    assert prop.normalization_record["raw_unit"] == unit
    assert prop.normalization_record["to"] == "kg"
    assert prop.normalization_record["factor"] == "1"
    assert prop.normalization_record["registry_version"] == UNIT_REGISTRY_VERSION
    assert result.proof_payloads[0]["normalization_record"] == prop.normalization_record
    assert prop.unit_evidence_refs[0].span_start == context.fragments[0].text.index(unit)


def test_header_unit_and_exact_conversion(tmp_path, monkeypatch):
    result, *_ = quantity_run(
        tmp_path, monkeypatch, "3800", raw="3800", unit="克",
        header=["预计批量下限（克）", "其他数值（kg）"], unit_source="预计批量下限（克）",
    )
    prop = result.properties[0]
    assert prop.policy_eligible and prop.normalized_value == "3.8" and prop.raw_value == "3800"
    assert prop.normalization_record["factor"] == "0.001"
    assert prop.normalization_record["mapping_kind"] == "conversion"


def test_field_group_label_unit_is_bound_to_its_value(tmp_path, monkeypatch):
    result, *_ = quantity_run(tmp_path, monkeypatch, "预计批量下限（公斤）：3.8", unit="公斤")
    assert result.properties[0].policy_eligible
    assert result.properties[0].normalized_value == "3.8"


def test_plan_owner_and_repeated_units_have_exact_source_bindings(tmp_path, monkeypatch):
    result, *_ = quantity_run(
        tmp_path, monkeypatch, "本次备样批量3.8-6.6 kg，另次批量5 kg。",
        unit_context="本次备样批量3.8-6.6 kg", owner="本次备样",
    )
    assert result.properties[0].policy_eligible
    assert result.properties[0].subject_ref.id == "plan"
    decisions = {d["check_kind"]: d for d in result.decision_payloads}
    assert decisions["local_coreference"]["verdict"] == "supported"
    assert decisions["unit_binding"]["verdict"] == "supported"


@pytest.mark.parametrize("source,kwargs,issue", [
    ("批量3.8。", {"unit": None}, "unit_source_missing"),
    ("批量3.8 kg。", {"unit_support": False}, "unit_binding_source_missing"),
    ("批量3.8 kg。", {"unit_verdict": "undetermined"}, "unit_binding_not_supported"),
    ("批量3.8。设备质量7 kg。", {}, "unit_record_mismatch"),
    ("批量3.8 kg/day。", {}, "unit_quote_partial"),
    ("批量3.8 kg/day。", {"unit": "kg/day"}, "unit_missing_or_incompatible"),
    ("批量3.8 kgf。", {}, "unit_quote_partial"),
    ("批量3.8 Mg。", {"unit": "Mg"}, "unit_unknown"),
    ("批量13.8 kg。", {}, "quantity_value_partial"),
    ("批量3.8 kg。", {"datatype": "integer"}, "datatype_mismatch"),
    ("周期1 s。", {"raw": "1", "unit": "s", "target_unit": "min"},
     "unit_conversion_not_exact"),
    ("3.8", {"header": ["预计批量下限（g）", "其他数值（kg）"],
             "unit_source": "其他数值（kg）"}, "unit_record_mismatch"),
    ("3.8 g", {"header": ["预计批量下限（kg）", "其他数值"],
               "unit_source": "预计批量下限（kg）"}, "source_unit_conflict"),
    ("3.8㎎", {"header": ["预计批量下限（kg）", "其他数值"],
               "unit_source": "预计批量下限（kg）"}, "source_unit_conflict"),
    ("3.8斤", {"header": ["预计批量下限（kg）", "其他数值"],
               "unit_source": "预计批量下限（kg）"}, "source_unit_conflict"),
    ("3.8 kg / day", {"header": ["预计批量下限（kg）", "其他数值"],
                       "unit_source": "预计批量下限（kg）"}, "source_unit_conflict"),
])
def test_unit_failures_keep_semantic_role_and_explain_actual_gate(
    tmp_path, monkeypatch, source, kwargs, issue,
):
    result, context, *_ = quantity_run(tmp_path, monkeypatch, source, **kwargs)
    prop = result.properties[0]
    assert not prop.policy_eligible and not prop.structural_valid
    assert prop.normalized_value is None and prop.normalization_record == {}
    assert prop.reason_code == issue
    assert "数值/单位核验未通过" in prop.reason
    role = next(d for d in result.decision_payloads if d["check_kind"] == "field_role")
    assert role["verdict"] == "supported" and role["support_refs"]
    issues = next(iter(context.protocol_state["gate_issues"].values()))
    assert issue in issues and "field_role_not_supported" not in issues


def test_correct_unit_cannot_override_rejected_field_role(tmp_path, monkeypatch):
    result, *_ = quantity_run(tmp_path, monkeypatch, "预计上限3.8 kg。", role_verdict="unsupported")
    prop = result.properties[0]
    assert not prop.policy_eligible and prop.decision_status == "unsupported"


def test_unit_policy_is_frozen_and_completed_review_replays_without_calls(tmp_path, monkeypatch):
    result, context, adapter, task, predicate, menu, _, _, calls = quantity_run(
        tmp_path, monkeypatch, "批量3.8公斤。", unit="公斤",
    )
    state = deepcopy(context.protocol_state)
    again = adapter.verify_existing(task, context, predicate, menu, state)
    assert len(calls) == 2 and again.model_calls == 0
    assert again.properties == result.properties
    for version in (None, "obsolete"):
        old = deepcopy(state)
        old["unit_normalization_version"] = version
        with pytest.raises(ValueError, match="unit_policy_mismatch"):
            adapter.verify_existing(task, context, predicate, menu, old)
    assert len(calls) == 2


def test_public_property_preserves_units_and_replayable_source_selections(tmp_path, monkeypatch):
    from app.schemas.document_analysis import GraphProperty as PublicProperty
    from app.services.document_analysis.public_projection import _property, build_selection_registry
    from app.services.extraction.ontology_guided.contracts import (
        GraphNode,
        RunProgress,
        VersionedRef,
    )
    from app.services.extraction.ontology_guided.projection import project_graph

    result, _, _, task, _, _, analysis, index, _ = quantity_run(
        tmp_path, monkeypatch, "批量3800-6600 克。", raw="3800", unit="克",
    )
    prop = result.properties[0]
    graph = project_graph(
        recognition_run_id="unit-run", run_revision=1, event_head=1,
        metadata_snapshot_id="metadata",
        root_ref=VersionedRef(id=task.subject.entity_id, revision=1),
        nodes=[GraphNode(
            entity_id=task.subject.entity_id, revision=1, class_iri=task.subject.class_iri,
            class_label="报告", label="报告", root=True,
        )],
        edges=[], properties=result.properties, coverage=[], progress=RunProgress(),
        projection="all", artifact_status="ready",
    )
    registry = build_selection_registry(
        recognition_run_id=graph.recognition_run_id, analysis_id=analysis.ir.analysis_id,
        graph=graph, index=index,
    )
    public = PublicProperty.model_validate(_property(prop, registry, set()))
    assert public.raw_value == "3800" and public.normalized_value == "3.8" and public.unit == "kg"
    assert public.datatype_iri == XSD + "decimal"
    assert public.normalization_record == prop.normalization_record
    refs = public.source_selection_refs.unit
    assert refs and all(registry[r]["selection_role"] == "unit" for r in refs)
    assert "evidence_id" not in public.model_dump_json()
    from app.schemas.evidence import EvidenceAnchor

    assert any(analysis.ir.resolve(EvidenceAnchor.model_validate(registry[r]["anchors"][0])) == "克"
               for r in refs)


def test_online_resume_rejects_old_unit_policy_before_model_dispatch(monkeypatch):
    from app.services.document_analysis import execution
    from app.services.extraction.ontology_guided.evidence_groups import (
        LITERAL_QUOTE_VERSION,
        SCOPE_PROTOCOL_VERSION,
    )
    from app.services.extraction.ontology_guided.evidence_work import EvidenceWorkQueue
    from app.services.extraction.ontology_guided.field_bindings import OWNER_BINDING_VERSION

    policy = {
        "evidence_repair": "evidence-repair-v1", "owner_binding": OWNER_BINDING_VERSION,
        "scope_protocol": SCOPE_PROTOCOL_VERSION, "literal_quotes": LITERAL_QUOTE_VERSION,
        "evidence_work": EvidenceWorkQueue.version,
        "unit_normalization": UNIT_NORMALIZATION_VERSION,
    }
    calls = []
    monkeypatch.setattr(execution, "configured_model_adapter", lambda **kw: calls.append(kw))
    execution._configured_recognition_adapter(policy)
    assert len(calls) == 1
    for value in (None, "old"):
        with pytest.raises(execution.CheckpointMismatch):
            execution._configured_recognition_adapter({**policy, "unit_normalization": value})
    assert len(calls) == 1


@pytest.mark.parametrize("raw", ["3.8 KG", "3.8 公斤", "3.8 千克", "３．８ ｋｇ"])
def test_shared_normalizer_also_accepts_registered_aliases(raw):
    value = normalize_literal(raw, datatype="decimal", target_unit="kg")
    assert value.normalized_value == "3.8" and value.raw_value == raw


def test_aliases_preserve_case_and_compound_dimensions():
    value = normalize_literal("3.8 毫克/公斤", datatype="decimal", target_unit="mg/kg")
    assert value.normalized_value == "3.8"
    for raw in ("3.8 Mg", "3.8 mg/kg"):
        with pytest.raises(LiteralNormalizationError):
            normalize_literal(raw, datatype="decimal", target_unit="kg")
