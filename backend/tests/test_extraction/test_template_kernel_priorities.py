"""Template priorities affect dispatch, while proof inputs and resume stay exact."""

import json

from docx import Document

from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import (
    EdgeSpec,
    OntologyClassDefinition,
    SlotSpec,
)
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.model_adapter import LocalModelRecognitionAdapter
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core

from .test_independent_semantic_verification import _inputs
from .test_semantic_graph_closure import (
    APPEARANCE,
    DESCRIBES,
    PRODUCT,
    ROOT,
    _analysis,
    _ontology,
    _respond,
)


def test_priority_paths_propagate_to_entity_properties_and_replay_dispatch(tmp_path, monkeypatch):
    ontology = _ontology()
    ontology.classes[ROOT].declared_relationships.insert(0, EdgeSpec(
        iri="urn:other", label="其他关系", range_class_iris=[PRODUCT]))
    ontology.classes[PRODUCT].declared_properties.insert(0, SlotSpec(
        iri="urn:other-value", label="其他属性"))
    analysis = _analysis(tmp_path)
    arguments = dict(recognition_run_id="priority-run", run_fingerprint="priority-test",
                     ir=analysis.ir, metadata=prepare_metadata(analysis.ir,
                         section_tree=analysis.structure.section_tree.to_dict(),
                         summary_version="priority-test"),
                     root_class_iri=ROOT, root_class_label="报告", filename="source.docx")
    calls = []

    def respond(_client, *, user, **_kw):
        request = json.loads(user)
        if request["stage"] == "discovery":
            calls.append((request["subject"]["is_document_root"], request["predicate"]["iri"],
                          tuple(f["text"] for f in request["fragments"] if f["fact_eligible"])))
        return _respond(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)

    def executor(progress_hook=None):
        return OntologyGuidedExecutor(
            ontology=ontology, engine=object(), max_tasks=200, progress_hook=progress_hook,
            adapter=LocalModelRecognitionAdapter(object(), model_identity="priorities"),
            priority_paths=[(DESCRIBES, APPEARANCE), (DESCRIBES,)],
        )

    baseline = executor().run(**arguments)
    expected = list(calls)
    assert expected[0][1] == DESCRIBES
    assert next(predicate for root, predicate, _ in expected
                if not root and predicate in {APPEARANCE, "urn:other-value"}) == APPEARANCE
    assert any(predicate == "urn:other" for _, predicate, _ in expected)
    calls.clear()
    batches = []
    executor(lambda stage: stage != "after_model" or len(calls) < 5).run(
        **arguments, batch_hook=batches.append)
    assert len(calls) < len(expected)
    result = executor().run(**arguments, resume_state=batches[-1].model_dump(mode="json"))
    assert calls == expected
    assert result.graph.progress.records_examined == baseline.graph.progress.records_examined


def test_named_overall_object_reads_related_complete_records_only_as_binding(tmp_path):
    task, target_context, _, _ = _inputs(tmp_path)
    document = Document()
    document.add_heading("制造工艺", 1)
    document.add_heading("路线总览", 2)
    document.add_paragraph("合成路线图")
    document.add_heading("合成步骤", 2)
    document.add_paragraph("原料经反应、纯化后得到产品；这是完整操作说明。")
    document.add_heading("另一产品", 1)
    document.add_heading("合成步骤", 2)
    document.add_paragraph("其他产品的工艺不能作为本路线背景。")
    path = tmp_path / "route.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    record = next(item for item in index.records if item.text == "合成路线图")
    ontology = _ontology()
    ontology.classes["urn:Route"] = OntologyClassDefinition(
        iri="urn:Route", label="合成路线", source_hash="route",
        declared_relationships=[EdgeSpec(iri="urn:steps", label="步骤",
                                         range_class_iris=["urn:Step"])])
    ontology.classes["urn:Step"] = OntologyClassDefinition(
        iri="urn:Step", label="合成步骤", source_hash="step")
    predicate = EdgeSpec(iri="urn:route", label="合成路线", range_class_iris=["urn:Route"])
    context = assemble_context(target_context.target, record.record_id, index,
                               predicate=predicate, ontology=ontology)
    assert [item.text for item in context.fragments if item.fact_eligible] == ["合成路线图"]
    assert any(item.text.startswith("原料经反应") and not item.fact_eligible
               for item in context.fragments)
    assert all("其他产品的工艺" not in item.text for item in context.fragments)
