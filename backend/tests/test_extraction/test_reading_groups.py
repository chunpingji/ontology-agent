"""Joint reading retains Word anchors and requires positive continuity evidence."""

import pytest
from docx import Document

from app.services.extraction.ontology_guided.reading_groups import (
    build_reading_groups,
    build_section_paragraph_groups,
    heading_context,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core


def index_for(tmp_path, sections, *, parent="产品基本性质"):
    document = Document()
    document.add_heading(parent, 1)
    for heading, values in sections:
        document.add_heading(heading, 2)
        for value in values:
            document.add_paragraph(value)
    path = tmp_path / "reading.docx"
    document.save(path)
    return RecordIndex(analyze_word_core(path).ir)


def texts(index, anchors):
    return [index.ir.resolve(anchor) for anchor in anchors]


def test_ontology_labeled_siblings_share_reading_but_keep_all_records_and_anchors(tmp_path):
    index = index_for(tmp_path, [("1.1 分子量", ["537.18"]), ("1.2 分子式", ["C23H21N7O3"])])
    before = index.ir.model_dump_json()
    groups = build_reading_groups(index, field_labels=["分子量", "分子式"])

    assert len(groups) == 1
    group = groups[0]
    assert group.record_ids == tuple(record.record_id for record in index.records)
    assert group.section_node_ids == tuple(record.section_node_id for record in index.records)
    assert texts(index, group.binding_refs) == ["产品基本性质", "1.1 分子量", "1.2 分子式"]
    assert [link.reason_code for link in group.reasons] == ["field_label_values"]
    assert group == build_reading_groups(index, field_labels=["分子式", "分子量"])[0]
    assert index.ir.model_dump_json() == before
    assert [texts(index, view.source_refs) for view in index.record_views] == [
        ["537.18"], ["C23H21N7O3"],
    ]


def test_short_adjacent_topics_do_not_create_a_reading_group(tmp_path):
    index = index_for(tmp_path, [("简介", ["产品甲"]), ("研究结论", ["产品乙"])])
    assert build_reading_groups(index) == ()


def test_consecutive_paragraphs_in_one_section_form_bounded_region_batches(tmp_path):
    index = index_for(tmp_path, [("工艺描述", ["第一步。", "第二步。", "第三步。"])])
    groups = build_section_paragraph_groups(index, max_records=2)
    assert [tuple(index.by_id[identity].text for identity in group.record_ids)
            for group in groups] == [("第一步。", "第二步。")]
    assert len({identity for group in groups for identity in group.record_ids}) == 2


def test_section_paragraph_batch_does_not_cross_excluded_table_lead_or_section(tmp_path):
    index = index_for(tmp_path, [
        ("工艺描述", ["第一步。", "结果如下表。", "不连续的第三步。"]),
        ("质量控制", ["检查甲。", "检查乙。"]),
    ])
    excluded = next(record.record_id for record in index.records
                    if record.text == "结果如下表。")
    groups = build_section_paragraph_groups(index, excluded_record_ids=[excluded])
    assert [tuple(index.by_id[identity].text for identity in group.record_ids)
            for group in groups] == [("检查甲。", "检查乙。")]


def test_explicit_label_and_value_split_across_sections_can_be_read_together(tmp_path):
    index = index_for(tmp_path, [("字段", ["分子量："]), ("数值", ["537.18"])])
    groups = build_reading_groups(index, field_labels=["分子量"])
    assert len(groups) == 1
    assert len(groups[0].record_ids) == 2
    assert groups[0].reasons[0].reason_code == "split_field_value"
    assert "分子量：" in texts(index, groups[0].reasons[0].binding_refs)


def test_empty_label_section_binds_next_value_without_inventing_a_record(tmp_path):
    index = index_for(tmp_path, [("分子量：", []), ("数值", ["537.18"])])
    groups = build_reading_groups(index, field_labels=["分子量"])
    assert len(groups) == 1
    assert groups[0].record_ids == (index.records[0].record_id,)
    assert "分子量：" in texts(index, groups[0].binding_refs)


@pytest.mark.parametrize("heading,text", [
    ("工艺（续）", "随后完成过滤。"),
    ("续上文", "随后完成过滤。"),
    ("补充", "接上段：随后完成过滤。"),
])
def test_explicit_continuation_keeps_independent_source_records(tmp_path, heading, text):
    index = index_for(tmp_path, [("工艺", ["先进行溶解。"]), (heading, [text])])
    groups = build_reading_groups(index)
    assert len(groups) == 1
    assert groups[0].reasons[0].reason_code == "explicit_continuation"
    assert len(groups[0].record_ids) == 2
    assert all(index.ir.resolve(anchor) for anchor in groups[0].reasons[0].binding_refs)


@pytest.mark.parametrize("value", [
    "样品B，537.18", "批次B，537.18", "批号：B-01", "物料编号：B-01",
    "未检出", "无", "未批准", "物料乙，537.18", "20℃下测试",
    "仅在20℃条件下成立", "若温度大于20℃", "sample B: 537.18",
    "not applicable", "if heated to 20°C",
])
def test_owner_negation_and_condition_boundaries_block_grouping(tmp_path, value):
    index = index_for(tmp_path, [("分子量", ["537.18"]), ("分子式", [value])])
    assert build_reading_groups(index, field_labels=["分子量", "分子式"]) == ()


def test_distinct_parents_and_intervening_subsections_are_not_crossed(tmp_path):
    document = Document()
    document.add_heading("产品甲", 1)
    document.add_heading("分子量", 2)
    document.add_paragraph("537.18")
    document.add_heading("产品乙", 1)
    document.add_heading("分子式", 2)
    document.add_paragraph("C23H21N7O3")
    document.add_heading("分子量", 2)
    document.add_heading("详细数值", 3)
    document.add_paragraph("621.22")
    path = tmp_path / "owners.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    assert build_reading_groups(index, field_labels=["分子量", "分子式"]) == ()


def test_a_table_is_not_absorbed_as_a_short_field_value(tmp_path):
    document = Document()
    document.add_heading("产品基本性质", 1)
    document.add_heading("分子量", 2)
    document.add_paragraph("537.18")
    document.add_heading("分子式", 2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "名称", "数值"
    table.cell(1, 0).text, table.cell(1, 1).text = "产品乙", "C23H21N7O3"
    path = tmp_path / "table.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    assert any(record.kind == "table_row" for record in index.records)
    assert build_reading_groups(index, field_labels=["分子量", "分子式"]) == ()


def test_capacity_bounds_do_not_overlap_or_silently_truncate_records(tmp_path):
    sections = [("分子量" if number % 2 == 0 else "分子式", [str(number)]) for number in range(5)]
    index = index_for(tmp_path, sections)
    labels = ["分子量", "分子式"]
    groups = build_reading_groups(index, field_labels=labels, max_records=2)
    assert [len(group.record_ids) for group in groups] == [2, 2]
    grouped = [record for group in groups for record in group.record_ids]
    assert len(grouped) == len(set(grouped)) == 4
    assert index.records[-1].record_id not in grouped
    assert build_reading_groups(index, field_labels=labels, max_group_chars=1) == ()
    assert build_reading_groups(index, field_labels=labels, max_section_chars=1) == ()


def test_repeated_field_headings_alone_do_not_prove_continuity(tmp_path):
    index = index_for(tmp_path, [("编号", ["A-01"]), ("编号", ["B-01"])])
    assert build_reading_groups(index, field_labels=["编号"]) == ()
    index = index_for(tmp_path, [("编号", ["A-01"]), ("型号", ["M-01"]), ("编号", ["B-01"])])
    groups = build_reading_groups(index, field_labels=["编号", "型号"])
    assert len(groups) == 1
    assert groups[0].record_ids == tuple(record.record_id for record in index.records[:2])


def test_heading_context_keeps_actual_ancestors_and_never_creates_filename_evidence(tmp_path):
    index = index_for(tmp_path, [("分子量", ["537.18"])])
    record = index.records[0]
    assert texts(index, heading_context(index, record.record_id)) == ["产品基本性质", "分子量"]
    assert texts(index, heading_context(index, record.record_id, max_ancestors=0)) == ["分子量"]
    document = Document()
    document.add_paragraph("孤立正文")
    path = tmp_path / "not-source-evidence.docx"
    document.save(path)
    no_heading = RecordIndex(analyze_word_core(path).ir)
    assert heading_context(no_heading, no_heading.records[0].record_id) == ()


def test_navigation_heading_is_not_binding_evidence_or_a_bridge_between_fields(tmp_path):
    document = Document()
    document.add_heading("产品基本性质", 1)
    document.add_heading("分子量", 2)
    document.add_paragraph("537.18")
    document.add_heading("目录", 2)
    document.add_paragraph("1 研究结论........1")
    document.add_paragraph("这里是真实正文。")
    document.add_heading("分子式", 2)
    document.add_paragraph("C23H21N7O3")
    path = tmp_path / "navigation.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    body = next(record for record in index.records if record.text == "这里是真实正文。")
    assert "目录" not in texts(index, heading_context(index, body.record_id))
    assert build_reading_groups(index, field_labels=["分子量", "分子式"]) == ()


@pytest.mark.parametrize("limits", [
    {"max_records": 0}, {"max_group_chars": 0}, {"max_section_chars": 0},
])
def test_invalid_capacity_is_rejected(tmp_path, limits):
    index = index_for(tmp_path, [("字段", ["值"])])
    with pytest.raises(ValueError, match="reading_group_limit_invalid"):
        build_reading_groups(index, **limits)
