"""Source parsing survives rejected mappings without becoming an accepted fact."""

from __future__ import annotations

import copy
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.extraction.ontology_guided.claim_freeze import freeze_record_proposal
from app.services.extraction.ontology_guided.claim_protocol import (
    DiscoveryEnvelope,
    ExtractionProfile,
    PropertyProposal,
)
from app.services.extraction.ontology_guided.context import assemble_record_discovery_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
    VersionedRef,
)
from app.services.extraction.ontology_guided.record_discovery import (
    RecordDiscoveryTask,
    compile_record_schema_card,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.tool_contracts import (
    ToolIssue,
    ValidateGraphArgs,
    ValidateMetricArgs,
)
from app.services.extraction.ontology_guided.value_constraints import XSD
from app.services.extraction.ontology_guided.value_observation import (
    normalize_field_support,
    parse_attribute_value,
)
from app.services.extraction.tool_validation.metric import normalize_metric
from app.services.extraction.tool_validation.shacl import PROFILE_VERSION, validate_metric_result
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_tool_engine_quantity import metric_case


@pytest.mark.parametrize("raw,expected,value,datatype,precision", [
    ("2026年02月", None, "2026-02", "gYearMonth", "month"),
    ("2026年2月3日", None, "2026-02-03", "date", "day"),
    ("否", None, False, "boolean", None),
    ("0", XSD + "boolean", False, "boolean", None),
    ("0012", XSD + "string", "0012", "string", None),
    ("1e400", XSD + "decimal", "1" + "0" * 400, "decimal", None),
])
def test_source_syntax_is_deterministic_and_does_not_round(
    raw, expected, value, datatype, precision,
):
    result = parse_attribute_value(raw, expected_datatype=expected)
    assert result.value == value
    assert result.datatype_iri == XSD + datatype
    assert result.precision == precision
    assert not result.issues


def test_expected_date_never_supplies_day_precision_and_unknown_boolean_is_not_false():
    result = parse_attribute_value("2026年02月", expected_datatype=XSD + "date")
    assert (result.value, result.precision) == ("2026-02", "month")
    assert result.datatype_iri == XSD + "gYearMonth"
    assert result.issues == ["datatype_mismatch"]
    unknown = parse_attribute_value("—", expected_datatype=XSD + "boolean")
    assert unknown.value is not False
    assert "datatype_mismatch" in unknown.issues
    assert parse_attribute_value("2026-02-30").value is None


def test_untrusted_interval_and_bound_keep_units_endpoints_and_openness():
    interval = parse_attribute_value("[1.20, 2.5) kg")
    assert interval.value is None
    assert interval.quantity == {
        "kind": "range", "operator": "eq", "scalar": None,
        "lower": "1.2", "upper": "2.5", "lower_inclusive": True,
        "upper_inclusive": False, "source_unit": "kg", "dimension": "mass",
    }
    bound = parse_attribute_value("≤25℃")
    assert bound.value is None
    assert bound.quantity["upper"] == "25"
    assert bound.quantity["upper_inclusive"] is True
    assert bound.quantity["lower"] is None
    approximate = parse_attribute_value("≈2 kg")
    assert approximate.value is None
    assert approximate.quantity["scalar"] == "2"
    assert "quantity_approximate" in approximate.issues


@pytest.mark.parametrize("failure", ["binding", "semantic", "datatype", "unit"])
def test_failed_mapping_keeps_parsing_but_no_trusted_quantity_or_literal(failure):
    case = metric_case("2026年02月", unit=None, datatype="gYearMonth", requirement="not_declared")
    if failure == "binding":
        case["binding"].validation_status = "failed"
        case["binding"].issues = [ToolIssue(
            code="field_column_mismatch", field_path=None, message="wrong label", evidence_ids=[],
        )]
    elif failure == "semantic":
        case["verified"].decisions[0].verdict = "undetermined"
    elif failure == "datatype":
        case["slot"].datatype_iris = [XSD + "date"]
    else:
        case = metric_case("2", unit="kg")
    result = normalize_metric(**case)
    assert result.validation_status != "passed"
    assert result.quantity is None and result.normalized_literal is None
    assert result.parsed_value.value == ("2" if failure == "unit" else "2026-02")
    if failure != "unit":
        shacl = validate_metric_result(
            result, slot=case["slot"], quantity_policy=case["quantity_policy"],
            candidate_ref=case["candidate_ref"], shape_profile_id=PROFILE_VERSION,
            raw=case["raw"],
        )
        assert shacl.evaluated is False and shacl.conforms is None
        assert shacl.blocked_by == [issue.code for issue in result.issues]


def test_llm_proposals_and_metric_arguments_cannot_supply_parsed_values():
    assert "parsed_value" not in PropertyProposal.model_fields
    assert "parsed_value" not in ValidateMetricArgs.model_fields


def test_graph_tool_reports_blocking_cause_instead_of_only_generic_shacl_failure(monkeypatch):
    from app.services.extraction.ontology_guided import tool_runtime

    case = metric_case("2026年02月", unit=None, datatype="gYearMonth", requirement="not_declared")
    case["binding"].validation_status = "failed"
    case["binding"].issues = [ToolIssue(
        code="field_column_mismatch", field_path=None, message="wrong label", evidence_ids=[],
    )]
    metric = normalize_metric(**case)
    monkeypatch.setattr(tool_runtime, "_metric_inputs", lambda *args: (
        case["target"], case["slot"], case["quantity_policy"],
    ))
    context = SimpleNamespace(stage="finalize", frozen_claims={}, metric_result=metric)
    result = tool_runtime._validate_graph(
        ValidateGraphArgs(claim_id=case["candidate_ref"].id, shape_profile_id=PROFILE_VERSION),
        context,
    )
    assert result.status == "ok"
    assert result.data.blocked_by == ["field_column_mismatch"]
    assert result.data.validation_status == "incomplete"
    assert result.data.evaluated is False and result.data.conforms is None
    assert [issue.code for issue in result.issues] == [
        "field_column_mismatch", "shacl_not_evaluated",
    ]
    assert "field_column_mismatch" in result.issues[-1].message


@pytest.fixture
def inline_field(tmp_path, request):
    document = Document()
    raw = getattr(request, "param", "时间：2026年02月")
    document.add_paragraph(raw)
    document.add_paragraph("名称：样品甲")
    path = tmp_path / "inline.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    record = next(r for r in index.records if r.text.startswith("时间"))
    ontology = OntologySnapshot(
        snapshot_id="ontology", ontology_hash="a" * 64,
        classes={"urn:Report": OntologyClassDefinition(
            iri="urn:Report", label="报告", source_hash="fixture",
            declared_properties=[SlotSpec(
                iri="urn:plannedDate", label="计划生产日期", datatype_iris=[XSD + "gYearMonth"],
                declared_by=["urn:Report"],
            )],
        )},
    )
    card = compile_record_schema_card(
        ontology, class_iris=ontology.classes, analysis_scope_ref="b" * 64,
        profile=ExtractionProfile(),
    )
    task = RecordDiscoveryTask.create(
        run_fingerprint="run", record_id=record.record_id, schema_card_id=card.schema_card_id,
        analysis_scope_ref=card.analysis_scope_ref, dependency_hash="c" * 64,
    )
    context = assemble_record_discovery_context(
        task, card, index, ontology_hash=ontology.ontology_hash,
        document_context=DocumentContext(
            document_hash=index.ir.document_hash, document_class_iri="urn:Report",
            root_ref=VersionedRef(id="root", revision=1),
        ),
    )
    def quote(text):
        return dict(evidence_id=record.source_units[0].evidence_id, text=text, context_text=None)
    proposal = DiscoveryEnvelope(
        entities=[], relations=[], external_links=[], observations=[],
        properties=[PropertyProposal(
            local_id="date", subject_id="root", predicate_iri="urn:plannedDate",
            value_quote=quote("2026年02月"), field_support=[quote(raw)],
            unit_support=[], bridge_kind="owned_field_group", bridge_ref_ids=[],
            qualifiers={"polarity": "affirmed", "modality": "asserted",
                        "condition_support": [], "scope_qualifiers": []},
        )],
    )
    return SimpleNamespace(index=index, context=context, task=task, card=card, proposal=proposal)


def test_freeze_narrows_label_without_changing_original_or_business_meaning(inline_field):
    case = inline_field
    original = case.proposal.model_dump(mode="json")
    frozen = freeze_record_proposal(
        case.proposal, task=case.task, context=case.context, index=case.index, card=case.card,
        generation=1,
    )
    assert frozen.properties[0].field_support[0].text == "时间"
    assert frozen.properties[0].value_quote.text == "2026年02月"
    assert frozen.properties[0].predicate_iri == "urn:plannedDate"
    assert case.proposal.model_dump(mode="json") == original
    # Freezing a structural correction does not mark its proposed meaning supported.
    assert "normalized_literal" not in frozen.properties[0].model_dump()


@pytest.mark.parametrize("mutation", ["scope", "duplicate", "wrong_source", "partial_label"])
def test_field_narrowing_does_not_repair_ambiguous_or_unauthorized_sources(inline_field, mutation):
    case = inline_field
    proposal = case.proposal.model_copy(deep=True)
    context = case.context.model_copy(deep=True)
    if mutation == "scope":
        context.fragments = [f.model_copy(update={"fact_eligible": False})
                             for f in context.fragments]
    elif mutation == "duplicate":
        context.field_bindings.append(context.field_bindings[0].model_copy(deep=True))
    elif mutation == "wrong_source":
        proposal.properties[0].field_support[0].evidence_id = "not-in-context"
    else:
        proposal.properties[0].field_support[0].text = "间：2026年02月"
    original = copy.deepcopy(proposal.model_dump())
    result = normalize_field_support(proposal, context=context, index=case.index)
    assert result.model_dump() == original


@pytest.mark.parametrize("inline_field", [
    "时间：2026年02月，名称：样品甲", "时间：2026年02月\n名称：样品甲",
], indirect=True)
def test_field_narrowing_does_not_cross_another_field_or_line(inline_field):
    case = inline_field
    original = copy.deepcopy(case.proposal.model_dump())
    normalize_field_support(case.proposal, context=case.context, index=case.index)
    assert case.proposal.model_dump() == original
