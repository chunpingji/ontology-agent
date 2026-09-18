"""018 — batched bottom-up Word tree summary contracts."""

from __future__ import annotations

import json
from copy import deepcopy
from threading import Barrier, Event, Lock, get_ident

import pytest
from docx import Document

from app.config import settings
from app.services.extraction import word_tree_summarizer
from app.services.extraction.docx_structure import parse_docx_structure
from app.services.llm.model_runtime import ModelCancelled, model_scope, runtime


def _structure(tmp_path):
    doc = Document()
    doc.add_heading("父章节", level=1)
    doc.add_paragraph("父章节直接内容。")
    doc.add_heading("子章节", level=2)
    doc.add_paragraph("子章节正文。")
    doc.add_page_break()
    doc.add_paragraph("子章节第二页正文。")
    doc.add_heading("另一个子章节", level=2)
    doc.add_paragraph("另一个子章节正文。")
    path = tmp_path / "summary.docx"
    doc.save(path)
    return parse_docx_structure(path)


def _enable(monkeypatch):
    monkeypatch.setattr(settings, "local_llm_enabled", True)
    monkeypatch.setattr(settings, "llm_word_tree_summary_enabled", True)
    monkeypatch.setattr(settings, "word_tree_summary_max_nodes_per_batch", 20)
    monkeypatch.setattr(settings, "word_tree_summary_max_batch_chars", 24000)
    monkeypatch.setattr(settings, "word_tree_summary_max_concurrency", 2)
    monkeypatch.setattr(settings, "local_llm_max_concurrency", 2)


def test_summary_order_is_pages_then_deepest_sections_then_root(tmp_path, monkeypatch):
    structure = _structure(tmp_path)
    _enable(monkeypatch)
    calls: list[list[str]] = []
    prompts: dict[str, str] = {}

    def fake_chat(_client, **kwargs):
        payload = json.loads(kwargs["user"])
        ids = [node["node_id"] for node in payload["nodes"]]
        calls.append(ids)
        for node in payload["nodes"]:
            prompts[node["node_id"]] = node["content"]
        return {
            "summaries": [
                {"node_id": node_id, "content_summary": f"摘要-{node_id}"}
                for node_id in ids
            ]
        }

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", fake_chat)
    root = word_tree_summarizer.summarize_word_tree(structure, object())

    assert root is not None
    flat_calls = [node_id for batch in calls for node_id in batch]
    page_index = next(i for i, value in enumerate(flat_calls) if ":page:" in value)
    child = root.children[0].children[0]
    parent = root.children[0]
    assert page_index < flat_calls.index(child.node_id)
    assert flat_calls.index(child.node_id) < flat_calls.index(parent.node_id)
    assert flat_calls.index(parent.node_id) < flat_calls.index("document")
    assert f"摘要-{child.node_id}" in prompts[parent.node_id]
    assert parent.layer_metadata.summary_status == "completed"


def test_unknown_and_missing_ids_fallback_per_node(tmp_path, monkeypatch):
    structure = _structure(tmp_path)
    _enable(monkeypatch)

    def fake_chat(_client, **kwargs):
        payload = json.loads(kwargs["user"])
        first = payload["nodes"][0]["node_id"]
        return {
            "summaries": [
                {"node_id": "unknown", "content_summary": "不得写入"},
                {"node_id": first, "content_summary": "有效摘要"},
                {"node_id": first, "content_summary": "重复摘要"},
            ]
        }

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", fake_chat)
    root = word_tree_summarizer.summarize_word_tree(structure, object())
    assert root is not None
    nodes = [root, *root.children, *root.children[0].children]
    summaries = [node.layer_metadata for node in nodes]
    assert all(metadata.content_summary != "不得写入" for metadata in summaries)
    assert any(
        metadata.summary_source == "extractive_fallback" for metadata in summaries
    )


def test_disabled_summary_keeps_complete_tree_without_call(tmp_path, monkeypatch):
    structure = _structure(tmp_path)
    monkeypatch.setattr(settings, "local_llm_enabled", False)
    monkeypatch.setattr(settings, "llm_word_tree_summary_enabled", True)

    def boom(*_args, **_kwargs):
        raise AssertionError("disabled summary must not call the model")

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", boom)
    root = word_tree_summarizer.summarize_word_tree(structure, None)
    assert root is structure.section_tree
    assert root.layer_metadata.summary_status == "disabled"
    assert root.children[0].children[0].pages
    assert root.children[0].children[0].pages[0].page_metadata.summary_status == "disabled"


def test_enabled_but_unavailable_client_uses_deterministic_fallback(tmp_path, monkeypatch):
    structure = _structure(tmp_path)
    _enable(monkeypatch)
    root = word_tree_summarizer.summarize_word_tree(structure, None)
    assert root is not None
    assert root.layer_metadata.summary_source == "extractive_fallback"
    assert root.layer_metadata.content_summary


@pytest.mark.parametrize("lost", [False, True])
def test_pause_or_lease_loss_stops_remaining_summary_calls(tmp_path, monkeypatch, lost):
    from app.services.extraction.annotation_execution import ExecutionLost

    structure = _structure(tmp_path)
    _enable(monkeypatch)
    calls = []

    def model(*args, **kwargs):
        calls.append(1)
        return {"summaries": []}

    def should_stop():
        if calls and lost:
            raise ExecutionLost()
        return bool(calls)

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", model)
    if lost:
        with pytest.raises(ExecutionLost):
            word_tree_summarizer.summarize_word_tree(structure, object(),
                                                   should_stop_fn=should_stop)
    else:
        word_tree_summarizer.summarize_word_tree(structure, object(), should_stop_fn=should_stop)
    assert calls == [1]


def _reply(nodes):
    return {"summaries": [
        {"node_id": n["node_id"], "content_summary": "HRS-5592：仅限临床备样，未批准商业生产。"}
        for n in nodes
    ]}


def test_single_page_reuses_generation_without_changing_source_or_scope(tmp_path, monkeypatch):
    _enable(monkeypatch)
    doc = Document()
    doc.add_heading("临床备样", level=1)
    doc.add_paragraph("HRS-5592：仅限临床备样，未批准商业生产。")
    path = tmp_path / "single.docx"
    doc.save(path)
    structure = parse_docx_structure(path)
    before = deepcopy((structure.blocks, structure.sections, structure.tables,
                       structure.paragraphs))
    leaf = structure.section_tree.children[0]
    original_hash = leaf.layer_metadata.content_hash
    calls = []

    def chat(_client, **kwargs):
        nodes = json.loads(kwargs["user"])["nodes"]
        calls.extend(n["node_id"] for n in nodes)
        return _reply(nodes)

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", chat)
    word_tree_summarizer.summarize_word_tree(structure, object())
    page = leaf.pages[0]
    assert page.node_id in calls and leaf.node_id not in calls
    assert leaf.layer_metadata.content_summary == page.page_metadata.content_summary
    assert leaf.layer_metadata.generated_at == page.page_metadata.generated_at
    assert leaf.layer_metadata.summary_source == "llm"
    assert leaf.layer_metadata.summary_scope == "subtree"
    assert page.page_metadata.summary_scope == "page_segment"
    assert leaf.layer_metadata.content_hash == original_hash
    assert before == (structure.blocks, structure.sections, structure.tables, structure.paragraphs)


def test_parent_deduplicates_only_successful_pages_and_keeps_failed_source(tmp_path, monkeypatch):
    _enable(monkeypatch)
    structure = _structure(tmp_path)
    leaf = structure.section_tree.children[0].children[0]
    first, second = leaf.pages
    prompts = {}

    def chat(_client, **kwargs):
        nodes = json.loads(kwargs["user"])["nodes"]
        prompts.update({n["node_id"]: n["content"] for n in nodes})
        return _reply([n for n in nodes if n["node_id"] != second.node_id])

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", chat)
    word_tree_summarizer.summarize_word_tree(structure, object())
    assert first.page_metadata.summary_status == "completed"
    assert second.page_metadata.summary_source == "extractive_fallback"
    material = prompts[leaf.node_id]
    assert "子章节正文。" not in material
    assert "子章节第二页正文。" in material
    assert "HRS-5592：仅限临床备样，未批准商业生产。" in material
    assert "父章节直接内容。" in prompts[structure.section_tree.children[0].node_id]
    assert leaf.layer_metadata.summary_status == "partial"
    assert structure.section_tree.layer_metadata.summary_status == "partial"


@pytest.mark.parametrize("truncated", ["input", "output"])
def test_partial_page_cannot_hide_its_original_source(tmp_path, monkeypatch, truncated):
    _enable(monkeypatch)
    structure = _structure(tmp_path)
    leaf = structure.section_tree.children[0].children[0]
    page = leaf.pages[0]
    original = word_tree_summarizer._page_material(structure, page)
    if truncated == "input":
        monkeypatch.setattr(settings, "word_tree_summary_max_input_chars_per_node", 3)
    else:
        monkeypatch.setattr(settings, "word_tree_summary_max_output_chars", 3)
    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", lambda _c, **kw:
                        _reply(json.loads(kw["user"])["nodes"]))
    word_tree_summarizer._apply_batch(object(), [{
        "node_id": page.node_id, "material": original, "metadata": page.page_metadata,
    }])
    assert page.page_metadata.summary_status == "partial"
    assert "子章节正文。" in word_tree_summarizer._chapter_material(structure, leaf)


@pytest.mark.parametrize("capacity", [1, 2])
def test_parallel_siblings_keep_parent_barrier_and_callbacks_on_owner(
    tmp_path, monkeypatch, capacity,
):
    _enable(monkeypatch)
    monkeypatch.setattr(settings, "word_tree_summary_max_nodes_per_batch", 1)
    monkeypatch.setattr(settings, "local_llm_max_concurrency", capacity)
    doc = Document()
    for i in range(4):
        doc.add_heading(f"章节{i}", level=1)
        doc.add_paragraph(f"HRS-5592，批号{i}，未批准商业生产。")
    path = tmp_path / "siblings.docx"
    doc.save(path)
    structure = parse_docx_structure(path)
    owner = get_ident()
    lock, barrier = Lock(), Barrier(capacity)
    state = {"active": 0, "peak": 0, "pages_done": 0}
    events, callbacks = [], []

    def callback(event=None):
        assert get_ident() == owner
        callbacks.append(True)
        if event:
            events.append(event)

    def chat(_client, **kwargs):
        assert get_ident() != owner
        assert runtime.get()["run_id"] == "owned-run"
        assert runtime.get()["job_id"] == "owned-job"
        assert runtime.get()["on_model_wait"] is None
        assert runtime.get()["stage"] == "word_tree_summary"
        nodes = json.loads(kwargs["user"])["nodes"]
        is_page = ":page:" in nodes[0]["node_id"]
        with lock:
            state["active"] += 1
            state["peak"] = max(state["peak"], state["active"])
        try:
            if is_page:
                barrier.wait(timeout=3)
                with lock:
                    state["pages_done"] += 1
            else:
                assert state["pages_done"] == 4
                assert "HRS-5592" in nodes[0]["content"]
            runtime.get()["progress"]({"node_id": nodes[0]["node_id"]})
            return _reply(nodes)
        finally:
            with lock:
                state["active"] -= 1

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", chat)
    with model_scope(run_id="owned-run", job_id="owned-job", progress=callback,
                     on_model_wait=callback):
        word_tree_summarizer.summarize_word_tree(structure, object())
    assert state["peak"] == capacity and state["active"] == 0
    assert len(events) == 5 and callbacks
    assert structure.section_tree.layer_metadata.summary_status == "completed"


def test_cancel_parallel_summary_closes_requests_and_does_not_start_parents(
    tmp_path, monkeypatch, isolated_model_scheduler,
):
    import asyncio

    import httpx

    from app.services.llm import local_client
    from app.services.llm.local_client import LocalModelClient

    from .test_model_scheduler import rows

    _enable(monkeypatch)
    monkeypatch.setattr(settings, "word_tree_summary_max_nodes_per_batch", 1)
    monkeypatch.setattr(local_client, "POLL_SECONDS", 0.01)
    bind = isolated_model_scheduler()
    both_started = Event()
    lock = Lock()
    sent, closed = [], []

    async def transport(_request):
        with lock:
            sent.append(True)
            if len(sent) == 2:
                both_started.set()
        try:
            await asyncio.sleep(3)
            raise AssertionError("cancellation must close the request first")
        finally:
            with lock:
                closed.append(True)

    client = LocalModelClient("http://model.test/v1", "test", httpx.MockTransport(transport))
    with model_scope(bind=bind, run_id="summary-cancel", should_stop=both_started.is_set):
        with pytest.raises(ModelCancelled):
            word_tree_summarizer.summarize_word_tree(_structure(tmp_path), client)
    assert len(sent) == len(closed) == 2
    requests = rows(bind)
    assert len(requests) == 2
    assert {r.status for r in requests} == {"cancelled"}
    assert {r.run_id for r in requests} == {"summary-cancel"}
    assert {r.stage for r in requests} == {"word_tree_summary"}
