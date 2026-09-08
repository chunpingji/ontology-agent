"""Attribution and scheduling regressions; scripted models test contracts, not quality."""

import json

import pytest
from docx import Document

from app.schemas.evidence import (
    BindingEvidence,
    Candidate,
    DocumentProvenance,
    ExtractionTask,
    TaskBudget,
)
from app.services.extraction.evidence_preview import preview_relationships
from app.services.extraction.evidence_scope import (
    _statement_bounds,
    build_scope,
    candidate_ref,
    scope_contains,
)
from app.services.extraction.extraction_tasks import GenericExtractionRunner, PartialTaskFailure
from app.services.extraction.gap_extraction import expanded_gap_scope, gap_subject_dependencies
from app.services.extraction.template_priorities import template_priority_paths
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_extraction_tasks import span
from tests.test_extraction.test_hierarchical_context import TestTokenizer


def source(tmp_path):
    document = Document()
    document.add_heading("报告", 0)
    document.add_heading("计划A", 1)
    document.add_heading("计划B", 1)
    document.add_heading("简介", 1)
    document.add_paragraph("计划A的批量3.8-6.6kg；计划B的批量9-12kg。")
    document.add_paragraph("供应商：甲公司；数量：300kg；存放地点：仓库。")
    path = tmp_path / "plans.docx"
    document.save(path)
    return analyze_word_core(path).ir


def entity(ir, index, identity, class_iri):
    unit = ir.evidence_units[index]
    return Candidate(
        candidate_id=identity,
        kind="entity",
        class_iri=class_iri,
        text=unit.text,
        validation_status="passed",
        provenance=[
            DocumentProvenance(anchors=[ir.anchor(unit.evidence_id)], excerpts=[unit.text])
        ],
    )


def bound_relation(ir, root, target, text):
    unit = ir.evidence_units[4]
    start = unit.text.index(text)
    anchor = ir.anchor(unit.evidence_id, start, start + len(text))
    return Candidate(
        candidate_id="relation-" + target.candidate_id,
        kind="relationship",
        subject=candidate_ref(root),
        object=candidate_ref(target),
        predicate_iri="urn:hasPlan",
        validation_status="passed",
        provenance=[DocumentProvenance(anchors=[anchor], excerpts=[text])],
        bindings=[
            BindingEvidence(
                method="explicit_assertion",
                subject_candidate_id=root.candidate_id,
                object_candidate_id=target.candidate_id,
                predicate_iri="urn:hasPlan",
                anchors=[anchor],
            )
        ],
    )


def test_bound_object_gets_only_its_verified_clause_not_sibling_or_parent_scope(tmp_path):
    ir = source(tmp_path)
    root, a, b = (
        entity(ir, i, name, cls)
        for i, name, cls in (
            (0, "root", "urn:Report"),
            (1, "a", "urn:Plan"),
            (2, "b", "urn:Plan"),
        )
    )
    ra = bound_relation(ir, root, a, "计划A的批量3.8-6.6kg")
    rb = bound_relation(ir, root, b, "计划B的批量9-12kg")
    scope = build_scope(ir, a, [b], bound_relationships=[ra, rb])
    assert scope_contains(scope, ra.bindings[0].anchors[0], ir)
    assert not scope_contains(scope, rb.bindings[0].anchors[0], ir)
    assert not scope_contains(scope, ir.anchor(ir.evidence_units[4].evidence_id), ir)
    assert not scope_contains(scope, ir.anchor(ir.evidence_units[5].evidence_id), ir)


def test_short_relation_quote_recovers_following_attributes_in_same_statement_and_gap(tmp_path):
    ir = source(tmp_path)
    root, a, b = (
        entity(ir, 0, "root", "urn:Report"),
        entity(ir, 1, "a", "urn:Plan"),
        entity(ir, 2, "b", "urn:Plan"),
    )
    ra = bound_relation(ir, root, a, "计划A")
    rb = bound_relation(ir, root, b, "计划B")
    candidates = [root, a, b, ra, rb]
    prefix, refs = gap_subject_dependencies(
        root, a, [{"predicate_iri": "urn:hasPlan", "direction": "forward"}], candidates
    )
    assert prefix == ["urn:hasPlan"]
    assert {ref.candidate_id for ref in refs} == {root.candidate_id, ra.candidate_id}
    scope = expanded_gap_scope(ir, a, candidates, 0, bound_relationships=[ra, rb])
    unit = ir.evidence_units[4]
    anchor = ir.anchor(unit.evidence_id, unit.text.index("3.8"), unit.text.index("kg") + 2)
    assert scope_contains(scope, anchor, ir)
    assert not scope_contains(scope, rb.provenance[0].anchors[0], ir)
    unrelated_root = root.model_copy(update={"candidate_id": "unrelated"})
    assert gap_subject_dependencies(
        unrelated_root, a, [{"predicate_iri": "urn:hasPlan", "direction": "forward"}], candidates
    ) == ([], [])


@pytest.mark.parametrize(
    "text,quote,expected",
    [
        ("Plan A: 3.8-6.6kg. Plan B: 9-12kg.", "Plan A", "Plan A: 3.8-6.6kg"),
        (
            "计划A批量3.8-6.6kg。 计划B批量9-12kg。",
            "计划A批量3.8-6.6kg。 ",
            "计划A批量3.8-6.6kg。 ",
        ),
    ],
)
def test_statement_extension_does_not_cross_sentence_boundary_or_split_decimals(
    text, quote, expected
):
    start = text.index(quote)
    left, right = _statement_bounds(text, start, start + len(quote))
    assert text[left:right] == expected


@pytest.mark.parametrize("change", ["negative", "stale", "rejected", "unverified"])
def test_ineligible_relation_cannot_expand_object_evidence(tmp_path, change):
    ir = source(tmp_path)
    root, a = entity(ir, 0, "root", "urn:Report"), entity(ir, 1, "a", "urn:Plan")
    relation = bound_relation(ir, root, a, "计划A的批量3.8-6.6kg")
    if change == "negative":
        relation.assertion_status = "negated"
    elif change == "stale":
        relation.object = relation.object.model_copy(update={"revision": 2})
    elif change == "rejected":
        relation.review_status = "rejected"
    else:
        relation.validation_status = "pending"
    scope = build_scope(ir, a, bound_relationships=[relation])
    assert not scope_contains(scope, relation.provenance[0].anchors[0], ir)


SCHEMA = {
    "urn:Report": {"relationships": [{"iri": "urn:hasPlan", "range": ["urn:Plan"]}]},
    "urn:Plan": {
        "properties": [
            {"iri": "urn:lower", "datatype": "decimal"},
            {"iri": "urn:upper", "datatype": "decimal"},
        ]
    },
    "urn:Supplier": {"properties": [{"iri": "urn:wrong", "datatype": "decimal"}]},
}


def plan_model(system, user, schema, budget):
    request = json.loads(user)
    context = json.loads(request["context"])
    task = context["task"]
    targets = [f for f in context["fragments"] if f["purpose"] == "target"]
    if request["stage"] == "verify_entity_types":
        return {
            "decisions": [
                {
                    "candidate_id": c["candidate_id"],
                    "supported": c["text"] != "300kg",
                    "reason": "按正文类型和表格角色复核",
                }
                for c in request["candidate"]["proposed_entities"]
            ]
        }
    if task["task_kind"] == "entity":
        return {
            "entities": [
                {"class_iri": cls, "mention": span(f, name)}
                for f in targets
                for name, cls in (
                    ("计划A", "urn:Plan"),
                    ("计划B", "urn:Plan"),
                    ("300kg", "urn:Supplier"),
                )
                if cls in task["predicate_definition"]["classes"]
                and (f["text"] == name or name == "300kg" and name in f["text"])
            ]
        }
    if request["stage"] == "verify_binding":
        candidate = request["candidate"]
        return {
            "supported": True,
            "subject_candidate_id": candidate["subject"]["candidate_id"],
            "object_candidate_id": (candidate.get("object") or {}).get("candidate_id"),
            "assertion_status": "affirmed",
            "method": "explicit_assertion",
            "assertion_spans": candidate["provenance"][0]["anchors"]
            and [span(f, f["text"]) for f in targets if "批量" in f["text"]],
        }
    if task["task_kind"] == "relationship":
        assertions = []
        for ref in task["object_candidates"]:
            name = context["subjects"][ref["candidate_id"]]["text"]
            clause = {"计划A": "计划A的批量3.8-6.6kg", "计划B": "计划B的批量9-12kg"}[name]
            for f in targets:
                if clause in f["text"]:
                    assertions.append(
                        {
                            "object_candidate_id": ref["candidate_id"],
                            "assertion_spans": [span(f, clause)],
                        }
                    )
        return {"assertions": assertions}
    name = context["subjects"][task["subject"]["candidate_id"]]["text"]
    value = {
        ("计划A", "urn:lower"): "3.8",
        ("计划A", "urn:upper"): "6.6",
        ("计划B", "urn:lower"): "9",
        ("计划B", "urn:upper"): "12",
    }.get((name, task["predicate_iri"]))
    return {
        "assertions": [{"value": span(f, value)} for f in targets if value and value in f["text"]]
    }


def test_heading_objects_keep_distinct_batch_bounds_through_relation_paths_and_resume(tmp_path):
    ir = source(tmp_path)
    calls = []

    def model(*args):
        calls.append(json.loads(args[1]))
        return plan_model(*args)

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        model,
        model_identity="contract",
        budget=TaskBudget(max_input_tokens=60000, max_regions_per_task=1),
        priority_paths=[("urn:hasPlan", "urn:lower"), ("urn:hasPlan", "urn:upper")],
    )
    run = runner.run(ir, effective_class="urn:Report")
    assert run.completion == "complete", run.diagnostics
    graph = preview_relationships(run.candidates, SCHEMA)
    values = {
        edge["object_text"]: {p["iri"]: p["value"] for p in edge["object_data_properties"]}
        for edge in graph
    }
    assert values == {
        "计划A": {"urn:lower": "3.8", "urn:upper": "6.6"},
        "计划B": {"urn:lower": "9", "urn:upper": "12"},
    }
    assert all(c.review_status == "pending" for c in run.candidates)
    assert not any(c.class_iri == "urn:Supplier" for c in run.candidates)
    count = len(calls)
    replay = runner.run(ir, effective_class="urn:Report", checkpoint=run.checkpoint)
    assert replay.candidates == run.candidates and len(calls) == count


def test_wrong_entity_type_is_recorded_but_never_scheduled(tmp_path):
    ir = source(tmp_path)
    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        plan_model,
        model_identity="types",
        budget=TaskBudget(max_input_tokens=60000),
    )
    run = runner.run(ir)
    wrong = next(c for c in run.candidates if c.text == "300kg")
    assert wrong.validation_status == "rejected"
    assert wrong.type_verification.supported is False
    assert not any(
        (t["task"].get("subject") or {}).get("candidate_id") == wrong.candidate_id
        for t in run.tasks
    )


@pytest.mark.parametrize("mode", ["missing", "duplicate", "unknown"])
def test_incomplete_or_ambiguous_entity_type_verdict_cannot_authorize_entities(tmp_path, mode):
    ir = source(tmp_path)

    def model(system, user, schema, budget):
        request = json.loads(user)
        if request["stage"] != "verify_entity_types":
            return plan_model(system, user, schema, budget)
        proposed = request["candidate"]["proposed_entities"]
        decisions = [
            {"candidate_id": c["candidate_id"], "supported": True, "reason": "fixture"}
            for c in proposed
        ]
        if mode == "missing":
            decisions.pop()
        elif mode == "duplicate":
            decisions.append(decisions[0])
        else:
            decisions[0]["candidate_id"] = "unknown"
        return {"decisions": decisions}

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        model,
        model_identity="types",
        budget=TaskBudget(max_input_tokens=60000),
    )
    task = ExtractionTask(
        task_id="types",
        task_kind="entity",
        target_evidence_ids=[u.evidence_id for u in ir.evidence_units],
        budget=runner.budget,
        **runner.entity_menu(list(SCHEMA)),
    )
    with pytest.raises(PartialTaskFailure, match="entity_type_decisions_mismatch") as failure:
        runner.execute_task(task, ir, {})
    assert failure.value.candidates == []


def test_priority_paths_are_part_of_resume_identity(tmp_path):
    ir = source(tmp_path)
    runner = GenericExtractionRunner(SCHEMA, TestTokenizer(), plan_model, model_identity="contract")
    previous = runner.input_id(ir, "urn:Report")
    runner.priority_paths = [("urn:hasPlan", "urn:upper")]
    assert runner.input_id(ir, "urn:Report") != previous


def test_bad_assertion_sibling_keeps_verified_value_and_source_record_mapping(tmp_path):
    ir = source(tmp_path)
    subject = entity(ir, 4, "plan", "urn:Plan")
    unit = ir.evidence_units[4]

    def model(system, user, schema, budget):
        request = json.loads(user)
        if request["stage"] == "recall":
            return {
                "assertions": [
                    {"value": {"evidence_id": unit.evidence_id, "text": text}}
                    for text in ("3.8", "absent value")
                ]
            }
        return {
            "supported": True,
            "subject_candidate_id": subject.candidate_id,
            "assertion_status": "affirmed",
            "method": "section_record",
            "assertion_spans": [{"evidence_id": unit.evidence_id, "text": unit.text}],
        }

    runner = GenericExtractionRunner(SCHEMA, TestTokenizer(), model, model_identity="partial")
    task = ExtractionTask(
        task_id="partial",
        task_kind="property",
        subject=candidate_ref(subject),
        predicate_iri="urn:lower",
        predicate_definition={"datatype": "decimal"},
        scope=build_scope(ir, subject),
        target_evidence_ids=[unit.evidence_id],
        budget=TaskBudget(max_input_tokens=20000),
    )
    with pytest.raises(PartialTaskFailure, match="source_excerpt_mismatch") as failure:
        runner.execute_task(task, ir, {subject.candidate_id: subject})
    assert len(failure.value.candidates) == 1
    valid = failure.value.candidates[0]
    assert valid.literal.normalized_value == "3.8" and valid.validation_status == "passed"
    assert valid.bindings[0].record_mapping == {"section_node_id": unit.section_node_id}


def test_terminal_failure_remains_visible_after_bad_sibling():
    failure = PartialTaskFailure([], ["source_excerpt_mismatch", "model_unavailable"])
    assert str(failure) == "model_unavailable"
    assert failure.issues == ["source_excerpt_mismatch", "model_unavailable"]


def test_type_verification_batches_fit_output_budget_without_dropping_entities(tmp_path):
    document = Document()
    document.add_paragraph("A B C D E F")
    path = tmp_path / "type-batches.docx"
    document.save(path)
    ir = analyze_word_core(path).ir
    batches = []

    def model(system, user, schema, budget):
        request = json.loads(user)
        if request["stage"] == "recall":
            return {
                "entities": [
                    {
                        "class_iri": "urn:Token",
                        "mention": {
                            "evidence_id": ir.evidence_units[0].evidence_id,
                            "text": text,
                        },
                    }
                    for text in "ABCDEF"
                ]
            }
        proposed = request["candidate"]["proposed_entities"]
        batches.append(len(proposed))
        return {
            "decisions": [
                {"candidate_id": c["candidate_id"], "supported": True, "reason": "source token"}
                for c in proposed
            ]
        }

    runner = GenericExtractionRunner(
        {"urn:Token": {}},
        TestTokenizer(),
        model,
        model_identity="batch-fixture",
        budget=TaskBudget(max_input_tokens=20000, max_output_tokens=1024),
    )
    run = runner.run(ir)
    assert run.completion == "complete"
    assert batches == [2, 2, 2]
    assert {c.text for c in run.candidates} == set("ABCDEF")


def test_template_projection_priorities_are_full_paths_and_reverse_paths_are_not_reinterpreted():
    schema = {
        "schema_version": 2,
        "definitions": {
            "bindings": {
                "plan": {
                    "kind": "facts",
                    "contract_ref": {"root_class_iri": "urn:Report"},
                    "scope": {
                        "root": {"kind": "source_root"},
                        "predicate_path": [
                            {"predicate_iri": "urn:hasPlan", "direction": "forward"}
                        ],
                    },
                }
            },
            "inputs": {
                "bounds": {
                    "binding_ref": "plan",
                    "projection": {
                        "kind": "record",
                        "fields": {
                            "lower": {"value": {"kind": "property", "property_iri": "urn:lower"}},
                            "upper": {"value": {"kind": "property", "property_iri": "urn:upper"}},
                        },
                    },
                }
            },
        },
    }
    assert template_priority_paths(schema, "urn:Report") == [
        ("urn:hasPlan", "urn:lower"),
        ("urn:hasPlan", "urn:upper"),
        ("urn:hasPlan",),
    ]
    schema["definitions"]["bindings"]["plan"]["scope"]["predicate_path"][0]["direction"] = "inverse"
    assert template_priority_paths(schema, "urn:Report") == []
