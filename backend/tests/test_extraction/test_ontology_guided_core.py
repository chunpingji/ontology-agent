from __future__ import annotations

import json

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    EdgeSpec,
    GraphEdge,
    GraphNode,
    GraphProperty,
    OntologyClassDefinition,
    OntologySnapshot,
    PredicateEvidence,
    RunProgress,
    SemanticDecision,
    SlotSpec,
    SubjectRef,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.dependencies import DependencyIndex
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor, TaskOutcome
from app.services.extraction.ontology_guided.mentions import MentionRegistry
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.model_adapter import LocalModelRecognitionAdapter
from app.services.extraction.ontology_guided.ontology_plan import compile_local_menu
from app.services.extraction.ontology_guided.projection import project_graph
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.retrieval import plan_slot
from app.services.extraction.ontology_guided.scheduler import FrontierScheduler, RecognitionTask
from app.services.extraction.ontology_guided.verification import ProofGate, observe_value
from app.services.extraction.word_analysis import analyze_word_core

REPORT = "urn:Report"
PRODUCT = "urn:Product"
DESCRIBES = "urn:describes"
APPEARANCE = "urn:appearance"


def _definition(iri, label, *, parents=(), properties=(), relationships=()):
    payload = {
        "iri": iri,
        "label": label,
        "parent_iris": list(parents),
        "declared_properties": list(properties),
        "declared_relationships": list(relationships),
    }
    return OntologyClassDefinition(
        **payload,
        source_hash=evidence_hash(payload),
    )


def ontology():
    describes = EdgeSpec(
        iri=DESCRIBES,
        label="描述产品",
        declared_by=[REPORT],
        range_class_iris=[PRODUCT],
    )
    appearance = SlotSpec(
        iri=APPEARANCE,
        label="外观",
        declared_by=[PRODUCT],
        datatype_iris=["string"],
    )
    classes = {
        REPORT: _definition(REPORT, "报告", relationships=[describes]),
        PRODUCT: _definition(PRODUCT, "产品", properties=[appearance]),
    }
    digest = evidence_hash({key: value.model_dump(mode="json") for key, value in classes.items()})
    return OntologySnapshot(
        snapshot_id="ontology-snapshot",
        ontology_hash=digest,
        classes=classes,
        created_from="frozen_fixture",
    )


def sample(tmp_path):
    document = Document()
    document.add_heading("产品候选", 1)
    document.add_paragraph("背景资料提到产品乙，但不是本报告对象。")
    document.add_heading("正文", 1)
    document.add_paragraph("本报告明确描述产品甲，外观为白色片剂。")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "API"
    table.cell(0, 1).text = "来源"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "成分甲"
    table.cell(1, 1).text = "试验来源-1"
    table.cell(2, 1).text = "试验来源-2"
    path = tmp_path / "source.docx"
    document.save(path)
    return analyze_word_core(path)


def test_direct_menu_does_not_preexpand_product_properties(tmp_path):
    analysis = sample(tmp_path)
    snapshot = ontology()
    root = SubjectRef(entity_id="root", revision=1, class_iri=REPORT, is_document_root=True)
    menu = compile_local_menu(snapshot, root)
    assert [item.iri for item in menu.relationships] == [DESCRIBES]
    assert menu.properties == []
    assert APPEARANCE not in menu.model_dump_json()

    product = SubjectRef(entity_id="product", revision=1, class_iri=PRODUCT)
    product_menu = compile_local_menu(snapshot, product)
    assert [item.iri for item in product_menu.properties] == [APPEARANCE]

    index = RecordIndex(analysis.ir)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="test",
    )
    plan = plan_slot(
        root,
        menu.relationships[0],
        index,
        metadata,
        ontology_hash=snapshot.ontology_hash,
        phase1_section_limit=1,
    )
    assert {item.record_id for item in plan.records} == set(plan.ledger)
    assert {item.phase for item in plan.records} == {1, 2}


def test_local_menu_is_frozen_snapshot_only_and_expands_frozen_range_subclasses():
    child = "urn:ProductChild"
    snapshot = ontology()
    snapshot.classes[child] = _definition(child, "子类产品", parents=[PRODUCT])

    class ExplodingEngine:
        def __getattribute__(self, name):
            raise AssertionError(f"live ontology access is forbidden: {name}")

    menu = compile_local_menu(
        snapshot,
        SubjectRef(entity_id="root", revision=1, class_iri=REPORT, is_document_root=True),
        engine=ExplodingEngine(),
    )

    assert menu.relationships[0].range_class_iris == [PRODUCT, child]
    assert [item.label for item in menu.relationships[0].range_classes] == ["产品", "子类产品"]


def test_physical_mention_is_shared_across_logical_merged_rows(tmp_path):
    analysis = sample(tmp_path)
    index = RecordIndex(analysis.ir)
    merged = next(unit for unit in analysis.ir.evidence_units if unit.text == "成分甲")
    views = [view for view in index.record_views if merged.source_cell_id in view.source_cell_ids]
    assert len(views) == 2
    registry = MentionRegistry(analysis.ir, index)
    first = registry.register(
        evidence_id=merged.evidence_id,
        start=0,
        end=len(merged.text),
        text=merged.text,
        record_view_ref=views[0].record_view_id,
    )
    second = registry.register(
        evidence_id=merged.evidence_id,
        start=0,
        end=len(merged.text),
        text=merged.text,
        record_view_ref=views[1].record_view_id,
    )
    assert first is second
    assert len(registry.mentions) == 1
    assert len(first.record_view_refs) == 2


def _target(anchor):
    subject = SubjectRef(entity_id="root", revision=1, class_iri=REPORT, is_document_root=True)
    return VerificationTarget.create(
        run_fingerprint="run-fingerprint",
        claim_ref=VersionedRef(id="claim", revision=1),
        task_id="task",
        check_kind="predicate_entailment",
        document_context=DocumentContext(
            document_hash=anchor.document_hash,
            document_class_iri=REPORT,
            root_ref=VersionedRef(id="root", revision=1),
        ),
        subject_ref=subject,
        predicate_iri=DESCRIBES,
        ontology_hash="ontology",
        source_scope_hash="scope",
        context_hash="context",
    )


def _decision(target, anchor, kind):
    return SemanticDecision(
        decision_id=f"decision-{kind}",
        target_id=target.target_id,
        check_kind=kind,
        verdict="supported",
        reason_code="supported",
        reason="原文支持",
        support_refs=[anchor],
        searched_context_refs=[anchor],
        verifier_version="test",
        attempt_id=f"attempt-{kind}",
    )


def test_proof_gate_rejects_same_name_even_when_model_says_supported(tmp_path):
    analysis = sample(tmp_path)
    unit = next(unit for unit in analysis.ir.evidence_units if "产品甲" in unit.text)
    anchor = analysis.ir.anchor(unit.evidence_id, 0, len(unit.text))
    target = _target(anchor)
    proof = PredicateEvidence(
        proof_id="proof",
        target_id=target.target_id,
        predicate_iri=DESCRIBES,
        subject_role_refs=[VersionedRef(id="subject-role", revision=1)],
        object_role_refs=[VersionedRef(id="object-role", revision=1)],
        predicate_support_refs=[anchor],
        bridge_kind="same_name",
        verdict="supported",
    )
    bundle = ProofGate().evaluate(
        target,
        proof,
        [
            _decision(target, anchor, "predicate_entailment"),
            _decision(target, anchor, "applicability"),
        ],
        is_property=False,
    )
    assert not bundle.policy_eligible
    assert "retrieval_hint_is_not_predicate_proof" in bundle.validation_issues


def test_proof_gate_accepts_controlled_bundle_facets_but_keeps_target_and_scope_bound(
    tmp_path,
):
    analysis = sample(tmp_path)
    unit = next(unit for unit in analysis.ir.evidence_units if "产品甲" in unit.text)
    anchor = analysis.ir.anchor(unit.evidence_id, 0, len(unit.text))
    target = _target(anchor)
    proof = PredicateEvidence(
        proof_id="proof-explicit",
        target_id=target.target_id,
        predicate_iri=DESCRIBES,
        subject_role_refs=[VersionedRef(id="subject-role", revision=1)],
        object_role_refs=[VersionedRef(id="object-role", revision=1)],
        predicate_support_refs=[anchor],
        bridge_kind="explicit_assertion",
        verdict="supported",
    )
    decisions = [
        _decision(target, anchor, kind)
        for kind in ("type", "field_role", "predicate_entailment", "applicability")
    ]

    accepted = ProofGate().evaluate(target, proof, decisions, is_property=False)
    assert accepted.structural_valid
    assert accepted.model_supported
    assert accepted.policy_eligible

    wrong_target = decisions[0].model_copy(update={"target_id": "another-target"})
    rejected_target = ProofGate().evaluate(
        target,
        proof,
        [wrong_target, *decisions[1:]],
        is_property=False,
    )
    assert not rejected_target.structural_valid
    assert not rejected_target.policy_eligible
    assert "decision belongs to another verification target" in rejected_target.validation_issues

    outside_scope = decisions[1].model_copy(update={"searched_context_refs": []})
    rejected_scope = ProofGate().evaluate(
        target,
        proof,
        [decisions[0], outside_scope, *decisions[2:]],
        is_property=False,
    )
    assert not rejected_scope.structural_valid
    assert not rejected_scope.policy_eligible
    assert (
        "decision cites evidence outside its searched context" in rejected_scope.validation_issues
    )


def test_missing_values_and_false_are_not_collapsed():
    ref = VersionedRef(id="urn:slot", revision=1)
    assert observe_value(slot_ref=ref, raw_text="N/A", source_refs=[]).presence == "placeholder"
    false_value = observe_value(slot_ref=ref, raw_text="否", source_refs=[])
    assert false_value.presence == "present"
    assert false_value.normalized_value == "否"
    assert observe_value(slot_ref=ref, raw_text=None, source_refs=[]).presence == "not_mentioned"


def test_dependency_alternatives_and_scheduler_fairness():
    index = DependencyIndex()
    index.add_proof("claim", ["left"])
    index.add_proof("claim", ["right"])
    assert index.invalidate("left") == {"left"}
    assert index.is_valid("claim")
    assert "claim" in index.invalidate("right")

    scheduler = FrontierScheduler(max_tasks=4)
    subject = SubjectRef(entity_id="root", revision=1, class_iri=REPORT, is_document_root=True)
    for number in range(2):
        scheduler.enqueue(
            RecognitionTask.create(
                subject=subject,
                predicate_iri=f"urn:root:{number}",
                predicate_kind="relationship",
                record_id=f"root-{number}",
                phase=1,
                hop=0,
                dependency_hash="d",
            ),
            root_branch=True,
        )
    child = subject.model_copy(update={"entity_id": "child", "is_document_root": False})
    for number in range(2):
        scheduler.enqueue(
            RecognitionTask.create(
                subject=child,
                predicate_iri=f"urn:child:{number}",
                predicate_kind="property",
                record_id=f"child-{number}",
                phase=1,
                hop=1,
                dependency_hash="d",
            )
        )
    assert [scheduler.next_task().record_id for _ in range(3)] == [
        "child-0",
        "child-1",
        "root-0",
    ]


class FakeAdapter:
    model_identity = "fake"

    def inspect(self, task, context, predicate, menu):
        text = "\n".join(fragment.text for fragment in context.fragments if fragment.fact_eligible)
        subject_ref = VersionedRef(id=task.subject.entity_id, revision=task.subject.revision)
        if predicate.iri == DESCRIBES and "明确描述产品甲" in text:
            node = GraphNode(
                entity_id="product",
                revision=1,
                class_iri=PRODUCT,
                class_label="产品",
                label="产品甲",
            )
            edge = GraphEdge(
                candidate_id="edge",
                revision=1,
                subject_ref=subject_ref,
                object_ref=VersionedRef(id="product", revision=1),
                predicate_iri=DESCRIBES,
                predicate_label="描述产品",
                decision_status="supported",
                structural_valid=True,
                model_supported=True,
                policy_eligible=True,
                proof_ref=VersionedRef(id="proof-describes", revision=1),
                decision_refs=[VersionedRef(id="decision-describes", revision=1)],
            )
            return TaskOutcome(
                semantic_outcome="supported",
                reason_code="supported",
                reason="明确描述",
                model_calls=1,
                nodes=[node],
                edges=[edge],
            )
        if predicate.iri == APPEARANCE and "外观为白色片剂" in text:
            prop = GraphProperty(
                candidate_id="appearance",
                revision=1,
                subject_ref=subject_ref,
                predicate_iri=APPEARANCE,
                predicate_label="外观",
                raw_value="白色片剂",
                decision_status="supported",
                structural_valid=True,
                model_supported=True,
                policy_eligible=True,
                proof_ref=VersionedRef(id="proof-appearance", revision=1),
                decision_refs=[VersionedRef(id="decision-appearance", revision=1)],
            )
            return TaskOutcome(
                semantic_outcome="supported",
                reason_code="supported",
                reason="明确属性",
                model_calls=1,
                properties=[prop],
            )
        return TaskOutcome(
            semantic_outcome="unsupported",
            reason_code="record_not_supporting_target",
            reason="当前记录不支持目标",
            model_calls=1,
        )


def test_effective_projection_and_frontier_share_the_complete_proof_gate(tmp_path):
    analysis = sample(tmp_path)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="test",
    )

    class UngatedAdapter(FakeAdapter):
        def inspect(self, task, context, predicate, menu):
            outcome = super().inspect(task, context, predicate, menu)
            return outcome.model_copy(
                update={
                    "edges": [
                        edge.model_copy(update={"policy_eligible": False}) for edge in outcome.edges
                    ]
                }
            )

    result = OntologyGuidedExecutor(
        ontology=ontology(),
        engine=object(),
        adapter=UngatedAdapter(),
        phase1_section_limit=1,
    ).run(
        recognition_run_id="proof-gate-run",
        run_fingerprint="proof-gate-fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=REPORT,
        root_class_label="报告",
        filename="source.docx",
    )
    assert {edge.candidate_id for edge in result.graph.edges} == {"edge"}
    assert not any(item.predicate_iri == APPEARANCE for item in result.graph.coverage)

    effective = project_graph(
        recognition_run_id=result.graph.recognition_run_id,
        run_revision=result.graph.run_revision,
        event_head=result.graph.event_head,
        metadata_snapshot_id=result.graph.metadata_snapshot_id,
        root_ref=result.graph.root_ref,
        nodes=result.graph.nodes,
        edges=result.graph.edges,
        properties=result.graph.properties,
        coverage=result.graph.coverage,
        progress=result.graph.progress,
        projection="effective",
        artifact_status=result.graph.artifact_status,
    )
    assert effective.edges == []
    assert len(effective.nodes) == 1


def test_one_bounded_retry_does_not_starve_fresh_records(tmp_path):
    analysis = sample(tmp_path)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="test",
    )

    class RetryOnceAdapter(FakeAdapter):
        def __init__(self):
            self.calls = {}
            self.failed_key = None

        def inspect(self, task, context, predicate, menu):
            key = (task.subject.entity_id, task.predicate_iri, task.record_id)
            self.calls[key] = self.calls.get(key, 0) + 1
            if self.failed_key is None:
                self.failed_key = key
                raise TimeoutError("transient model timeout")
            return super().inspect(task, context, predicate, menu)

    adapter = RetryOnceAdapter()
    result = OntologyGuidedExecutor(
        ontology=ontology(),
        engine=object(),
        adapter=adapter,
        phase1_section_limit=1,
    ).run(
        recognition_run_id="retry-run",
        run_fingerprint="retry-fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=REPORT,
        root_class_label="报告",
        filename="source.docx",
    )

    assert adapter.calls[adapter.failed_key] == 2
    assert max(adapter.calls.values()) == 2
    assert "bounded_technical_retry_scheduled" in result.diagnostics
    assert result.graph.progress.completion == "in_scope_complete"


@pytest.mark.parametrize("polarity", ["negated", "conditional"])
def test_supported_nonaffirmed_candidate_keeps_support_without_expanding_frontier(
    tmp_path, monkeypatch, polarity
):
    document = Document()
    source_text = (
        "仅针对批次一，本报告描述产品甲。" if polarity == "conditional" else "本报告不描述产品甲。"
    )
    document.add_paragraph(source_text)
    path = tmp_path / "polarity.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="test",
    )

    def supported_nonaffirmed(_client, *, user, **_kwargs):
        request = json.loads(user)
        fragment = next(
            (
                item
                for item in request["fragments"]
                if item["fact_eligible"] and source_text == item["text"]
            ),
            None,
        )
        if request["predicate"]["iri"] != DESCRIBES or fragment is None:
            return {"proposals": []}
        evidence_id = fragment["evidence_id"]
        if request["stage"] == "verification":
            return {"verifications": [
                {
                    "candidate_id": candidate["candidate_id"],
                    "target_id": candidate["target_id"],
                    "type_verdict": "supported",
                    "role_verdict": "supported",
                    "subject_binding_verdict": "supported",
                    "predicate_verdict": "supported",
                    "applicability_verdict": "supported",
                    "counterevidence_verdict": "undetermined",
                    "bridge_verdict": "supported",
                    "type_support": [{"evidence_id": evidence_id, "text": source_text}],
                    "predicate_support": [{"evidence_id": evidence_id, "text": source_text}],
                    "subject_support": [],
                    "condition_support": (
                        [{"evidence_id": evidence_id, "text": "仅针对批次一"}]
                        if polarity == "conditional"
                        else []
                    ),
                    "counterevidence_support": [],
                    "reason": "独立读取原文后确认带极性和条件的断言。",
                }
                for candidate in request["candidates"]
            ]}
        return {
            "proposals": [
                {
                    "kind": "relationship",
                    "object_class_iri": PRODUCT,
                    "object_label": "产品甲",
                    "object_quote": {"evidence_id": evidence_id, "text": "产品甲"},
                    "predicate_support": [{"evidence_id": evidence_id, "text": source_text}],
                    "subject_support": [{"evidence_id": evidence_id, "text": "本报告"}],
                    "bridge_kind": "explicit_assertion",
                    "type_verdict": "supported",
                    "role_verdict": "supported",
                    "predicate_verdict": "supported",
                    "applicability_verdict": "supported",
                    "polarity": polarity,
                    "condition_support": (
                        [{"evidence_id": evidence_id, "text": "仅针对批次一"}]
                        if polarity == "conditional"
                        else []
                    ),
                    "reason": "原文明确支持带极性的断言。",
                }
            ]
        }

    monkeypatch.setattr(model_adapter, "chat_with_schema", supported_nonaffirmed)
    result = OntologyGuidedExecutor(
        ontology=ontology(),
        engine=object(),
        adapter=LocalModelRecognitionAdapter(object(), model_identity="polarity-test-model"),
        phase1_section_limit=1,
    ).run(
        recognition_run_id="polarity-run",
        run_fingerprint="polarity-fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=REPORT,
        root_class_label="报告",
        filename="source.docx",
    )

    assert result.graph.edges[0].polarity == polarity
    if polarity == "conditional":
        assert result.graph.edges[0].conditions == ["仅针对批次一"]
        assert [
            analysis.ir.resolve(ref) for ref in result.graph.edges[0].condition_evidence_refs
        ] == ["仅针对批次一"]
    assert result.graph.edges[0].decision_status == "supported"
    assert result.graph.edges[0].policy_eligible is True
    assert result.graph.nodes[-1].decision_status == "supported"
    assert result.graph.progress.supported == 1
    assert not any(item.predicate_iri == APPEARANCE for item in result.graph.coverage)


def test_executor_continues_into_phase2_then_expands_direct_object_menu(tmp_path):
    analysis = sample(tmp_path)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="test",
    )
    result = OntologyGuidedExecutor(
        ontology=ontology(),
        engine=object(),
        adapter=FakeAdapter(),
        phase1_section_limit=1,
    ).run(
        recognition_run_id="run",
        run_fingerprint="fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=REPORT,
        root_class_label="报告",
        filename="source.docx",
    )
    assert {edge.candidate_id for edge in result.graph.edges} == {"edge"}
    assert {item.candidate_id for item in result.graph.properties} == {"appearance"}
    assert any(item.predicate_iri == APPEARANCE for item in result.graph.coverage)
    describes = next(item for item in result.graph.coverage if item.predicate_iri == DESCRIBES)
    assert describes.phase2 > 0
    assert result.graph.progress == RunProgress(**result.graph.progress.model_dump(mode="json"))
    assert result.graph.progress.completion == "in_scope_complete"
