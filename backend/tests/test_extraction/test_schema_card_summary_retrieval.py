from __future__ import annotations

from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.schema_card_evidence import build_sources
from app.evaluation.schema_card_summary_retrieval import (
    MAX_RECORDS,
    MAX_SOURCE_CHARACTERS,
    retrieve_scope,
)
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core

REPORT = "urn:report"
PRODUCT = "urn:product"
FIELD = {"iri": "urn:metric", "label": "指标", "kind": "property",
         "datatype_iris": ["http://www.w3.org/2001/XMLSchema#string"]}


def fixtures(tmp_path, document, summaries=None):
    path = tmp_path / "summary-retrieval.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    tree = analysis.structure.section_tree.to_dict()

    def fill(node):
        summary = (summaries or {}).get(node.get("heading"))
        if summary:
            node.setdefault("layer_metadata", {}).update(
                content_summary=summary, summary_status="completed", summary_source="llm",
            )
        for child in node.get("children", []):
            fill(child)

    fill(tree)
    metadata = prepare_metadata(analysis.ir, section_tree=tree, summary_version="test-frozen")
    menus = {"document": {"class_iri": REPORT, "label": "报告", "fields": {"metric": FIELD}}}
    catalog = {REPORT: {"class_iri": REPORT, "label": "报告", "properties": [FIELD]}}
    return analysis.ir, metadata, menus, catalog


def ranked_document():
    document = Document()
    for index in range(8):
        document.add_heading(f"章节{index}", 1)
        document.add_paragraph(f"指标原值{index}")
    return document


def test_summary_changes_selection_but_never_enters_evidence(tmp_path):
    args = fixtures(tmp_path, ranked_document(), {"章节7": "指标 摘要独有且不可引用的虚构值"})
    ir, metadata, menus, catalog = args
    before = metadata.model_dump(mode="json")
    result = retrieve_scope(*args, "arbitrary-query-family")
    retrieval = result["retrieval"]
    enabled, masked = retrieval["summary_enabled"], retrieval["summary_masked"]
    assert enabled["selected_record_ids"] != masked["selected_record_ids"]
    assert len(enabled["selected_record_ids"]) == len(masked["selected_record_ids"]) == 6
    assert "指标原值7" in result["case"]["content"]["text"]
    assert "摘要独有" not in result["case"]["content"]["text"]
    assert "虚构值" not in str(result["case"]["source_units"])
    assert enabled["metadata_snapshot_id"] != masked["metadata_snapshot_id"]
    assert all(row["score_components"]["summary"] == 0 for row in masked["records"])
    assert metadata.model_dump(mode="json") == before
    for unit in result["case"]["source_units"]:
        assert unit == ir.unit(unit["evidence_id"]).model_dump(mode="json")
    assert retrieval["budgets"]["max_records"] == MAX_RECORDS == 6
    assert retrieval["budgets"]["max_source_characters"] == MAX_SOURCE_CHARACTERS == 12000
    assert len(retrieval["summary_enabled"]["records"]) == len(RecordIndex(ir).records)
    assert retrieval["ablation"]["source_unit_set_changed"] is True
    assert retrieval["ablation"]["source_input_changed"] is True
    assert result["case"]["document_ref"] == ""


def test_ranking_order_change_is_not_evidence_input_change(tmp_path):
    document = Document()
    for index in range(2):
        document.add_heading(f"章节{index}", 1)
        document.add_paragraph(f"指标{index}")
    result = retrieve_scope(*fixtures(tmp_path, document, {"章节1": "指标"}), "scope")
    ablation = result["retrieval"]["ablation"]
    assert ablation["selected_order_changed"] is True
    assert ablation["source_unit_set_changed"] is False
    assert ablation["source_input_changed"] is False
    assert ablation["added_record_ids"] == ablation["removed_record_ids"] == []


@pytest.mark.parametrize("field,value", [
    ("analysis_id", "another-analysis"),
    ("document_hash", "1" * 64),
    ("structure_hash", "2" * 64),
])
def test_rejects_cross_source_metadata(tmp_path, field, value):
    ir, metadata, menus, catalog = fixtures(tmp_path, ranked_document())
    with pytest.raises(ValueError, match="retrieval_metadata_source_mismatch"):
        retrieve_scope(ir, metadata.model_copy(update={field: value}), menus, catalog, "scope")


def test_rejects_metadata_that_loses_records_or_reuses_a_stale_identity(tmp_path):
    ir, metadata, menus, catalog = fixtures(tmp_path, ranked_document())
    with pytest.raises(ValueError, match="record_universe_mismatch"):
        retrieve_scope(ir, metadata.model_copy(update={"source_record_refs": []}),
                       menus, catalog, "scope")
    changed = metadata.model_copy(deep=True)
    changed.node_summaries[-1].summary = "指标外加的摘要"
    with pytest.raises(ValueError, match="metadata_identity_mismatch"):
        retrieve_scope(ir, changed, menus, catalog, "scope")


def test_full_record_headers_notes_field_group_and_original_headings_are_replayed(tmp_path):
    document = Document()
    document.add_heading("原始上级标题", 1)
    document.add_heading("原始子标题", 2)
    document.add_paragraph("对象：甲")
    document.add_paragraph("指标：5")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("产品", "条件"), ("甲", "指标不超过25℃"),
                                       ("乙", "无关记录")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("备注：仅限本次研究。")
    args = fixtures(tmp_path, document)
    result = retrieve_scope(*args, "not-a-source-filter")
    text = result["case"]["content"]["text"]
    for expected in ("原始上级标题", "原始子标题", "对象：甲", "指标：5", "产品", "条件",
                     "甲", "指标不超过25℃", "备注：仅限本次研究。"):
        assert expected in text
    assert "无关记录" not in text
    roles = result["retrieval"]["record_roles"]
    selected = result["retrieval"]["summary_enabled"]["selected_record_ids"]
    assert any(roles[rid]["refs"]["field_group"] for rid in selected)
    assert any(roles[rid]["refs"]["header"] for rid in selected)
    assert any(roles[rid]["refs"]["note"] for rid in selected)
    assert all(roles[rid]["refs"]["ancestor_heading"] for rid in selected)
    actual_ids = [unit["evidence_id"] for unit in result["case"]["source_units"]]
    assert len(actual_ids) == len(set(actual_ids))
    index = RecordIndex(args[0])
    assert actual_ids == sorted(actual_ids, key=index.positions.__getitem__)


def test_parser_header_context_does_not_become_a_confirmed_column_header(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "投料"
    table.cell(0, 1).text = "步骤一：指标达到标准后继续反应。"
    table.cell(1, 0).text = "下一步"
    table.cell(1, 1).text = "继续操作"
    args = fixtures(tmp_path, document)
    result = retrieve_scope(*args, "scope")
    narrative = next(unit for unit in result["case"]["source_units"] if "步骤一" in unit["text"])
    role_check = result["retrieval"]["unit_header_status"][narrative["evidence_id"]]
    assert role_check["parser_header_candidate"] is True
    assert role_check["is_header"] is False
    assert any(narrative["evidence_id"] == ref["evidence_id"]
               for role in result["retrieval"]["record_roles"].values()
               for ref in role["refs"]["header"])
    sources = build_sources(result["case"], args[0].model_dump(mode="json"))
    assert next(s for s in sources.values() if s["evidence_id"] == narrative["evidence_id"])[
        "is_header"] is False


def test_nested_table_retains_parent_record_as_separate_context(tmp_path):
    document = Document()
    document.add_heading("研究", 1)
    outer = document.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "产品"
    outer.cell(0, 1).text = "研究内容"
    outer.cell(1, 0).text = "母记录甲"
    outer.cell(1, 1).text = "研究背景"
    inner = outer.cell(1, 1).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "项目"
    inner.cell(0, 1).text = "结果"
    inner.cell(1, 0).text = "指标"
    inner.cell(1, 1).text = "5"
    args = fixtures(tmp_path, document)
    result = retrieve_scope(*args, "scope")
    selected = result["retrieval"]["summary_enabled"]["selected_record_ids"]
    roles = result["retrieval"]["record_roles"]
    parent_refs = [a for rid in selected for a in roles[rid]["refs"]["parent"]]
    assert parent_refs
    assert any(args[0].resolve(ref) == "母记录甲" for ref in parent_refs)
    assert "母记录甲" in result["case"]["content"]["text"]
    assert "研究背景" in result["case"]["content"]["text"]
    assert all(ref not in roles[rid]["refs"]["source"]
               for rid in selected for ref in roles[rid]["refs"]["parent"])


def test_oversized_complete_context_is_deferred_without_truncation(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    long_text = "指标：" + "长" * MAX_SOURCE_CHARACTERS
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = long_text
    table.cell(1, 0).text = "甲"
    table.cell(1, 1).text = "值"
    document.add_paragraph("指标：3")
    args = fixtures(tmp_path, document)
    result = retrieve_scope(*args, "scope")
    selected = result["retrieval"]["summary_enabled"]
    assert selected["deferred_records"]
    assert any(row["reason"] == "source_character_budget" for row in selected["deferred_records"])
    assert result["case"]["content"]["text"] == "指标：3"
    assert selected["source_characters"] == len("指标：3")
    long_unit = next(unit for unit in args[0].evidence_units if unit.text == long_text)
    deferred_ids = {row["record_id"] for row in selected["deferred_records"]}
    assert any(long_unit.evidence_id in result["retrieval"]["record_roles"][rid]["unit_ids"]
               for rid in deferred_ids)
    assert long_unit.text == long_text


def test_ties_are_stable_and_duplicate_slots_or_scope_names_cannot_bias_queries(tmp_path):
    ir, metadata, menus, catalog = fixtures(tmp_path, ranked_document())
    baseline = retrieve_scope(ir, metadata, menus, catalog, "first")
    duplicate = deepcopy(menus)
    duplicate["untrusted-slot"] = deepcopy(menus["document"])
    duplicate["untrusted-slot"]["anchor"] = {"quote": "指标原值7", "ref": "old-u7"}
    duplicate = dict(reversed(list(duplicate.items())))
    again = retrieve_scope(ir, metadata, duplicate, catalog, "指标原值7")
    left, right = [item["retrieval"]["summary_enabled"] for item in (baseline, again)]
    assert left == right
    assert left["query_count"] == 1
    assert len(left["selected_record_ids"]) == MAX_RECORDS
    assert all(row["rank_score"] == 4 for row in left["records"])
    assert [row["source_position"] for row in left["records"]] == sorted(
        row["source_position"] for row in left["records"])
    assert len(left["deferred_records"]) == 2


def test_relationship_range_labels_come_from_frozen_catalog(tmp_path):
    document = Document()
    document.add_paragraph("此处记载待识别对象")
    ir, metadata, menus, catalog = fixtures(tmp_path, document)
    menus["document"]["fields"] = {"rel": {
        "kind": "relationship", "iri": "urn:relation", "label": "无匹配谓词",
        "range_class_iris": [PRODUCT],
    }}
    catalog[PRODUCT] = {"class_iri": PRODUCT, "label": "待识别对象", "properties": []}
    result = retrieve_scope(ir, metadata, menus, catalog, "scope")
    assert result["retrieval"]["summary_enabled"]["selected_record_ids"]
    predicate = result["retrieval"]["queries"][0]["predicate"]
    assert predicate["range_classes"][0]["label"] == "待识别对象"


def test_zero_matches_are_unselected_and_coverage_keeps_empty_units(tmp_path):
    document = Document()
    document.add_heading("章节", 1)
    document.add_paragraph("与查询无关的材料")
    document.add_paragraph("")
    args = fixtures(tmp_path, document)
    result = retrieve_scope(*args, "scope")
    assert result["case"]["source_units"] == []
    assert result["retrieval"]["summary_enabled"]["selected_record_ids"] == []
    assert all(row["selection_status"] == "not_selected_nonpositive"
               for row in result["retrieval"]["summary_enabled"]["records"])
    categories = result["coverage"]["unit_categories"]
    assert categories["empty"]["count"] > 0
    assert sum(row["count"] for row in categories.values()) == len(args[0].evidence_units)
    flattened = [eid for row in categories.values() for eid in row["evidence_ids"]]
    assert len(flattened) == len(set(flattened))
