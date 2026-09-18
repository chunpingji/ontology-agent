"""Isolated evidence-group experiments retain the production proof pipeline.

Only the model transport is controlled. These tests do not measure model
quality or timings, and do not add a new production context-selection policy.
"""

from __future__ import annotations

import json

import pytest
from docx import Document

from app.services.extraction.ontology_guided import executor, model_adapter
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    GraphNode,
    RunProgress,
    SubjectRef,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.model_adapter import (
    LocalModelRecognitionAdapter,
    RecognitionModelFailure,
)
from app.services.extraction.ontology_guided.projection import project_graph
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_semantic_graph_closure import (
    APPEARANCE,
    DESCRIBES,
    PRODUCT,
    ROOT,
    _ontology,
    _run,
)

NAME = "项目名称：产品甲片"
BINDING = "本报告记录的药物产品由项目名称字段确定；对照资料另提到产品乙。"
APPEARANCE_TEXT = "产品甲片外观为白色片剂。"


def _fixture(tmp_path, *, extra_text=None):
    document = Document()
    for heading, text in (
        ("字段信息", NAME), ("文档适用范围", BINDING), ("产品外观", APPEARANCE_TEXT),
    ):
        document.add_heading(heading, 1)
        document.add_paragraph(text)
    if extra_text:
        document.add_heading("复核说明", 1)
        document.add_paragraph(extra_text)
    path = tmp_path / "joint-source.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    target_record = next(record for record in index.records if record.text == NAME)
    binding_unit = next(unit for unit in analysis.ir.evidence_units if unit.text == BINDING)
    binding = analysis.ir.anchor(binding_unit.evidence_id, 0, len(binding_unit.text))
    subject = SubjectRef(entity_id="root", revision=1, class_iri=ROOT, is_document_root=True)
    task = RecognitionTask.create(
        subject=subject, predicate_iri=DESCRIBES, predicate_kind="relationship",
        record_id=target_record.record_id, phase=1, hop=0, dependency_hash="joint-evidence",
    )
    target = VerificationTarget.create(
        run_fingerprint="joint-evidence-experiment", claim_ref=VersionedRef(
            id=task.claim_lineage_id, revision=1,
        ),
        task_id=task.task_id, check_kind="predicate_entailment",
        document_context=DocumentContext(
            document_hash=analysis.ir.document_hash, document_class_iri=ROOT,
            root_ref=VersionedRef(id="root", revision=1),
        ),
        subject_ref=subject, predicate_iri=DESCRIBES, ontology_hash=_ontology().ontology_hash,
        source_scope_hash="base-source", context_hash="base-context",
    )
    return analysis, index, task, target, binding


def _response(request):
    predicate = request["predicate"]["iri"]
    fragments = request["fragments"]
    if request["stage"] == "discovery":
        for fragment in fragments:
            if not fragment["fact_eligible"]:
                continue
            if predicate == DESCRIBES and fragment["text"] == NAME:
                return {"proposals": [{
                    "kind": "relationship", "object_class_iri": PRODUCT,
                    "object_label": "产品甲片",
                    "object_quote": {"evidence_id": fragment["evidence_id"], "text": "产品甲片"},
                    "bridge_kind": "document_subject_description",
                }]}
            if predicate == APPEARANCE and fragment["text"] == APPEARANCE_TEXT:
                return {"proposals": [{
                    "kind": "property", "value_quote": {
                        "evidence_id": fragment["evidence_id"], "text": "白色片剂",
                    }, "bridge_kind": "explicit_assertion",
                }]}
        return {"proposals": []}
    result = []
    for candidate in request["candidates"]:
        endpoint_id = candidate["claim"]["endpoint_quote"]["evidence_id"]
        endpoint = next(fragment for fragment in fragments
                        if fragment["evidence_id"] == endpoint_id)
        binding = next((fragment for fragment in fragments if fragment["text"] == BINDING), None)
        root = request["subject"]["is_document_root"]
        supported = bool(binding) if root else endpoint["text"] == APPEARANCE_TEXT
        verdict = "supported" if supported else "undetermined"
        quote = {"evidence_id": endpoint["evidence_id"], "text": endpoint["text"]}
        predicate_support = [quote]
        if root and binding:
            predicate_support.append({"evidence_id": binding["evidence_id"], "text": BINDING})
        result.append({
            "candidate_id": candidate["candidate_id"], "target_id": candidate["target_id"],
            "type_verdict": "supported", "role_verdict": verdict,
            "subject_binding_verdict": "supported" if root else verdict,
            "predicate_verdict": verdict, "applicability_verdict": "supported",
            "counterevidence_verdict": "undetermined", "bridge_verdict": verdict,
            "type_support": [quote] if root else [],
            "predicate_support": predicate_support,
            "subject_support": [] if root else [{
                "evidence_id": endpoint["evidence_id"], "text": "产品甲片",
            }],
            "condition_support": [], "counterevidence_support": [],
            "reason": "受控独立核验：只有完整角色说明可将项目字段绑定为本报告药物产品。",
        })
    return {"verifications": result}


def test_joint_sources_change_proof_scope_while_preserving_the_semantic_task(tmp_path, monkeypatch):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    local = assemble_context(target, task.record_id, index)
    joint = assemble_context(target, task.record_id, index, required_context_refs=[binding])
    assert local.target.target_id != joint.target.target_id
    assert local.target.context_hash != joint.target.context_hash
    assert local.target.source_scope_hash != joint.target.source_scope_hash
    assert local.target.claim_ref == joint.target.claim_ref
    assert local.target.task_id == joint.target.task_id == task.task_id
    assert {item.anchor.evidence_id for item in local.fragments if item.fact_eligible} == {
        item.anchor.evidence_id for item in joint.fragments if item.fact_eligible
    }
    assert all(not item.fact_eligible for item in joint.fragments if item.text == BINDING)
    requests, reservations = [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        requests.append(request)
        return _response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    adapter = LocalModelRecognitionAdapter(object(), model_identity="joint-proof-test")
    predicate = _ontology().classes[ROOT].declared_relationships[0]
    for context in (local, joint):
        context.bind_model_call_hook(lambda stage, ordinal: reservations.append((stage, ordinal)))
    before = adapter.inspect(task, local, predicate, None)
    after = adapter.inspect(task, joint, predicate, None)

    assert [request["stage"] for request in requests] == [
        "discovery", "verification", "discovery", "verification",
    ]
    assert len(reservations) == before.model_calls + after.model_calls == 4
    assert before.complete and before.semantic_outcome == "undetermined"
    assert not before.edges[0].policy_eligible
    assert after.semantic_outcome == "supported" and after.edges[0].policy_eligible
    assert before.edges[0].candidate_id == after.edges[0].candidate_id
    assert before.proof_payloads[0]["target_id"] != after.proof_payloads[0]["target_id"]
    assert binding.evidence_id in {
        ref.evidence_id for ref in after.edges[0].predicate_evidence_refs
    }
    assert after.proof_payloads and after.decision_payloads
    assert requests[0]["fragments"] == requests[1]["fragments"]
    assert requests[2]["fragments"] == requests[3]["fragments"]


def test_binding_group_does_not_authorize_a_new_object_endpoint(tmp_path, monkeypatch):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    context = assemble_context(target, task.record_id, index, required_context_refs=[binding])
    requests = []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        requests.append(request)
        fragment = next(item for item in request["fragments"] if item["text"] == BINDING)
        return {"proposals": [{
            "kind": "relationship", "object_class_iri": PRODUCT, "object_label": "产品乙",
            "object_quote": {"evidence_id": fragment["evidence_id"], "text": "产品乙"},
            "bridge_kind": "document_subject_description",
        }]}

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    with pytest.raises(RecognitionModelFailure, match="discovery_citation_invalid") as caught:
        LocalModelRecognitionAdapter(object(), model_identity="joint-proof-test").inspect(
            task, context, _ontology().classes[ROOT].declared_relationships[0], None,
        )
    assert len(requests) == caught.value.model_calls == 1
    assert caught.value.cause_type == "source_quote_outside_scope"


@pytest.mark.parametrize("failure", ["foreign_source", "budget"])
def test_joint_context_failure_prevents_model_dispatch(tmp_path, monkeypatch, failure):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    if failure == "foreign_source":
        binding = binding.model_copy(update={"document_hash": "0" * 64})

    class Counter:
        def count(self, text):
            return 101

    context = assemble_context(
        target, task.record_id, index, required_context_refs=[binding],
        token_counter=Counter() if failure == "budget" else None, max_input_tokens=100,
    )

    def unexpected(*_args, **_kwargs):
        pytest.fail("invalid or over-budget grouped context must not reach the model")

    monkeypatch.setattr(model_adapter, "chat_with_schema", unexpected)
    outcome = LocalModelRecognitionAdapter(object(), model_identity="joint-proof-test").inspect(
        task, context, _ontology().classes[ROOT].declared_relationships[0], None,
    )
    assert not outcome.complete and outcome.model_calls == 0
    assert outcome.reason_code == (
        "required_context_missing" if failure == "foreign_source" else "context_budget_exceeded"
    )


def test_joint_proof_enables_property_scheduling_but_undetermined_is_not_auto_retried(
    tmp_path, monkeypatch,
):
    analysis, _index, _task, _target, binding = _fixture(tmp_path)
    monkeypatch.setattr(
        model_adapter, "chat_with_schema",
        lambda _client, *, user, **_kwargs: _response(json.loads(user)),
    )
    baseline = _run(
        analysis, LocalModelRecognitionAdapter(object(), model_identity="joint-control"),
    )
    original_assemble = executor.assemble_context

    def with_group(target, record_ref, index, **kwargs):
        if target.predicate_iri == DESCRIBES and index.by_id[record_ref].text == NAME:
            kwargs["required_context_refs"] = [*kwargs.get("required_context_refs", []), binding]
        return original_assemble(target, record_ref, index, **kwargs)

    monkeypatch.setattr(executor, "assemble_context", with_group)
    grouped = _run(analysis, LocalModelRecognitionAdapter(object(), model_identity="joint-control"))

    before = [payload for kind, payload in baseline.events if kind == "task_outcome"]
    assert any(item["outcome"]["semantic_outcome"] == "undetermined" for item in before)
    assert all(item["task"]["retry_kind"] is None for item in before)
    assert not any(item.predicate_iri == APPEARANCE for item in baseline.graph.coverage)
    assert any(item.predicate_iri == APPEARANCE for item in grouped.graph.coverage)
    assert [(item.predicate_iri, item.raw_value) for item in grouped.graph.properties
            if item.policy_eligible] == [(APPEARANCE, "白色片剂")]
    assert any(item.policy_eligible for item in grouped.graph.edges
               if item.predicate_iri == DESCRIBES)


def test_arbitrary_remote_binding_does_not_bypass_property_owner_gate(tmp_path, monkeypatch):
    analysis, _index, original_task, original_target, _binding = _fixture(tmp_path)
    path = tmp_path / "joint-source.docx"
    document = Document(path)
    document.paragraphs[-1].text = "外观：白色片剂"
    document.add_heading("跨段归属说明", 1)
    document.add_paragraph("上节外观字段属于产品甲片。")
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    record = next(item for item in index.records if item.text == "外观：白色片剂")
    name_unit = next(unit for unit in analysis.ir.evidence_units if unit.text == NAME)
    owner_unit = next(unit for unit in analysis.ir.evidence_units if "上节外观字段" in unit.text)
    name_ref = analysis.ir.anchor(name_unit.evidence_id)
    owner_ref = analysis.ir.anchor(owner_unit.evidence_id)
    subject = SubjectRef(entity_id="product", revision=1, class_iri=PRODUCT)
    task = RecognitionTask.create(
        subject=subject, predicate_iri=APPEARANCE, predicate_kind="property",
        record_id=record.record_id, phase=1, hop=1, dependency_hash="joint-owner-test",
    )
    values = original_target.model_dump(mode="python", exclude={"target_id"})
    values.update(
        task_id=task.task_id, subject_ref=subject, predicate_iri=APPEARANCE,
        claim_ref=VersionedRef(id=task.claim_lineage_id, revision=1),
        document_context=DocumentContext(
            document_hash=analysis.ir.document_hash, document_class_iri=ROOT,
            root_ref=VersionedRef(id=original_task.subject.entity_id, revision=1),
        ),
    )
    target = VerificationTarget.create(run_fingerprint="joint-property-owner", **values)
    context = assemble_context(
        target, record.record_id, index, subject_label="产品甲片",
        subject_evidence_refs=[name_ref], required_context_refs=[owner_ref],
    )
    assert context.owner_field_refs == []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        endpoint = next(item for item in request["fragments"] if item["fact_eligible"])
        if request["stage"] == "discovery":
            return {"proposals": [{
                "kind": "property", "bridge_kind": "owned_field_group",
                "value_quote": {"evidence_id": endpoint["evidence_id"], "text": "白色片剂"},
            }]}
        response = _response(request)
        for item in response["verifications"]:
            item.update(
                role_verdict="supported", subject_binding_verdict="supported",
                predicate_verdict="supported", bridge_verdict="supported",
                subject_support=[{"evidence_id": fragment["evidence_id"], "text": "产品甲片"}
                                 for fragment in request["fragments"]
                                 if "产品甲片" in fragment["text"]],
            )
        return response

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    outcome = LocalModelRecognitionAdapter(object(), model_identity="joint-control").inspect(
        task, context, _ontology().classes[PRODUCT].declared_properties[0], None,
    )
    assert outcome.model_calls == 2 and outcome.semantic_outcome == "undetermined"
    assert not outcome.properties[0].policy_eligible
    coreference = next(item for item in outcome.decision_payloads
                       if item["check_kind"] == "local_coreference")
    assert coreference["verdict"] == "undetermined" and coreference["support_refs"] == []


def test_support_record_alone_has_no_target_endpoint_for_the_reported_product(
    tmp_path, monkeypatch,
):
    _analysis, index, original_task, original_target, binding = _fixture(tmp_path)
    support_record = next(record for record in index.records if record.text == BINDING)
    task = RecognitionTask.create(
        subject=original_task.subject, predicate_iri=DESCRIBES, predicate_kind="relationship",
        record_id=support_record.record_id, phase=1, hop=0, dependency_hash="support-only",
    )
    values = original_target.model_dump(mode="python", exclude={"target_id"})
    values.update(
        task_id=task.task_id, claim_ref=VersionedRef(id=task.claim_lineage_id, revision=1),
    )
    target = VerificationTarget.create(run_fingerprint="support-only", **values)
    context = assemble_context(target, support_record.record_id, index)
    assert not any("产品甲片" in fragment.text for fragment in context.fragments)
    assert any(fragment.anchor.evidence_id == binding.evidence_id and fragment.fact_eligible
               for fragment in context.fragments)
    monkeypatch.setattr(
        model_adapter, "chat_with_schema",
        lambda _client, *, user, **_kwargs: _response(json.loads(user)),
    )
    outcome = LocalModelRecognitionAdapter(object(), model_identity="joint-control").inspect(
        task, context, _ontology().classes[ROOT].declared_relationships[0], None,
    )
    assert outcome.complete and outcome.model_calls == 1
    assert outcome.reason_code == "no_candidate_observed"
    assert not outcome.edges and not outcome.nodes


@pytest.mark.parametrize("case,source", [
    ("negated", "更正：本报告不描述产品甲片。"),
    ("conditional", "本报告仅在批次一的范围内描述产品甲片。"),
    ("competing_owner", "项目名称字段的产品甲片属于对照报告，本报告记录产品乙。"),
])
def test_joint_counterevidence_cannot_produce_an_unconditional_effective_parent(
    tmp_path, monkeypatch, case, source,
):
    analysis, index, task, target, binding = _fixture(tmp_path, extra_text=source)
    unit = next(unit for unit in analysis.ir.evidence_units if unit.text == source)
    counter = analysis.ir.anchor(unit.evidence_id, 0, len(unit.text))
    context = assemble_context(
        target, task.record_id, index, required_context_refs=[binding],
        counterevidence_refs=[counter],
    )

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        result = _response(request)
        counter_fragment = next(item for item in request["fragments"] if item["text"] == source)
        quote = {"evidence_id": counter_fragment["evidence_id"], "text": source}
        if request["stage"] == "discovery":
            if case in {"negated", "conditional"}:
                result["proposals"][0]["polarity"] = case
            if case == "conditional":
                result["proposals"][0]["condition_support"] = [{
                    "evidence_id": quote["evidence_id"], "text": "仅在批次一的范围内",
                }]
        else:
            for item, candidate in zip(result["verifications"], request["candidates"], strict=True):
                item.update(counterevidence_verdict="supported", counterevidence_support=[quote])
                item["predicate_support"].append(quote)
                item["condition_support"] = candidate["claim"]["conditions"]
                if case == "competing_owner":
                    item.update(role_verdict="unsupported", predicate_verdict="unsupported")
        return result

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    outcome = LocalModelRecognitionAdapter(object(), model_identity="joint-control").inspect(
        task, context, _ontology().classes[ROOT].declared_relationships[0], None,
    )
    assert outcome.model_calls == 2 and outcome.proof_payloads and outcome.decision_payloads
    assert outcome.edges
    if case in {"negated", "conditional"}:
        assert outcome.edges[0].policy_eligible
        assert outcome.edges[0].polarity == case
    else:
        assert not outcome.edges[0].policy_eligible
    root = GraphNode(
        entity_id=task.subject.entity_id, revision=1, class_iri=ROOT, class_label="报告",
        label="受控测试文档", root=True,
    )
    effective = project_graph(
        recognition_run_id="joint-counter-test", run_revision=1, event_head=0,
        metadata_snapshot_id=None, root_ref=VersionedRef(id=root.entity_id, revision=1),
        nodes=[root, *outcome.nodes], edges=outcome.edges, properties=[], coverage=[],
        progress=RunProgress(), projection="effective", artifact_status="partial",
    )
    assert effective.edges == []
    assert len(effective.nodes) == 1
