"""Record-scoped Responses use complete claims and resume saved paid work."""

from __future__ import annotations

import copy
import json
from dataclasses import replace
from types import SimpleNamespace

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.claim_protocol import (
    ExtractionProfile,
    plan_verification_target_batches,
)
from app.services.extraction.ontology_guided.context import assemble_record_discovery_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
    VersionedRef,
)
from app.services.extraction.ontology_guided.current_work import (
    call_request_key,
    validate_tool_protocol,
)
from app.services.extraction.ontology_guided.model_adapter import RecognitionModelFailure
from app.services.extraction.ontology_guided.record_discovery import (
    RECORD_PIPELINE,
    RecordDiscoveryTask,
    compile_record_schema_card,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.tool_model_adapter import ToolModelRecognitionAdapter
from app.services.extraction.ontology_guided.tool_runtime import ToolLimits
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.local_client import ResponseTurn

XSD = "http://www.w3.org/2001/XMLSchema#"


@pytest.fixture
def source(tmp_path):
    document = Document()
    document.add_heading("正文", level=1)
    document.add_paragraph("对象乙的数量为5 mg，对象丙的名称为样品丙。")
    path = tmp_path / "records.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    record = next(record for record in index.records if "对象乙" in record.text)
    unit = record.source_units[0]

    def quote(text):
        return {"evidence_id": unit.evidence_id, "text": text, "context_text": None}

    return {"index": index, "record_id": record.record_id, "quote": quote,
            "source_unit": unit, "proposal": {
                "entities": [{"local_id": identity, "class_iri": class_iri,
                              "representation": "mention", "mentions": [quote(text)],
                              "record_components": [], "identifier_claims": []}
                             for identity, class_iri, text in (("b", "urn:T", "对象乙"),
                                                              ("c", "urn:Other", "对象丙"))],
                "properties": [], "relations": [], "reference_bindings": [],
                "external_links": [], "observations": [],
            }}


def setup_record(source, monkeypatch, *, stop_at=None, empty=False, tool_first=False,
                 invalid_first=False, budget=6, no_properties=False, transform=None):
    index = source["index"]
    ontology = OntologySnapshot(
        snapshot_id="ontology", ontology_hash=evidence_hash("ontology"),
        classes={iri: OntologyClassDefinition(
            iri=iri, label=iri, source_hash="fixture",
            declared_properties=[] if no_properties else [
                SlotSpec(iri="urn:value", label="数量" if unit else "名称", declared_by=[iri],
                         datatype_iris=[datatype], canonical_unit=unit),
            ],
        ) for iri, datatype, unit in (("urn:T", XSD + "decimal", "mg"),
                                    ("urn:Other", XSD + "string", None))},
    )
    card = compile_record_schema_card(
        ontology, class_iris=ontology.classes, analysis_scope_ref=evidence_hash("scope"),
        profile=ExtractionProfile(),
    )
    task = RecordDiscoveryTask.create(
        run_fingerprint="run", record_id=source["record_id"],
        schema_card_id=card.schema_card_id, analysis_scope_ref=card.analysis_scope_ref,
        dependency_hash=evidence_hash("dependencies"),
    )
    context = assemble_record_discovery_context(
        task, card, index, ontology_hash=ontology.ontology_hash,
        document_context=DocumentContext(document_hash=index.ir.document_hash,
                                         document_class_iri="urn:Document",
                                         root_ref=VersionedRef(id="document-root", revision=1)),
    )
    context.tool_inputs = {
        "target_seed": context.target.model_dump(mode="json"),
        "schema_card": card.model_dump(mode="json"), "run_fingerprint": "run",
        "entity_dependencies": [], "entity_nodes": [], "reference_resolutions": [],
    }
    context.remaining_model_calls = budget
    proposal = copy.deepcopy(source["proposal"])
    proposal.update(relations=[], reference_bindings=[])
    proposal["entities"][1]["class_iri"] = "urn:Other"
    if not no_properties:
        proposal["properties"] = [{
            "local_id": "pb", "subject_id": "b", "predicate_iri": "urn:value",
            "value_quote": source["quote"]("5"),
            "field_support": [source["quote"]("数量")],
            "unit_support": [source["quote"]("mg")],
            "qualifiers": {"polarity": "affirmed", "modality": "asserted",
                           "condition_support": [], "scope_qualifiers": []},
            "bridge_kind": "explicit_assertion", "bridge_ref_ids": [],
        }, {
            "local_id": "pc", "subject_id": "c", "predicate_iri": "urn:value",
            "value_quote": source["quote"]("样品丙"),
            "field_support": [source["quote"]("名称")], "unit_support": [],
            "qualifiers": {"polarity": "affirmed", "modality": "asserted",
                           "condition_support": [], "scope_qualifiers": []},
            "bridge_kind": "explicit_assertion", "bridge_ref_ids": [],
        }]
    if empty:
        proposal = {field: [] for field in proposal}
    storage = {"protocol": {}, "results": {}, "reservations": []}

    def checkpoint(value):
        value = copy.deepcopy(value)
        changes = value.pop("result_changes", {})
        validate_tool_protocol(value)
        storage["protocol"] = value
        storage["results"].update(changes)
        if stop_at and any(row["field"] == stop_at for row in changes.values()):
            raise RuntimeError("pause after durable commit")

    def reserve(stage, ordinal):
        assert storage["protocol"]["pending_request"]["attempt"] == ordinal
        storage["reservations"].append((stage, ordinal))

    context.bind_protocol_hook(checkpoint)
    context.bind_model_call_hook(reserve)
    requests = []

    def transport(client, **kwargs):
        requests.append(copy.deepcopy(kwargs))
        view = json.loads(kwargs["input_items"][0]["content"][0]["text"])
        has_paired_tool_result = any(
            item.get("type") == "function_call_output" for item in kwargs["input_items"]
        )
        if tool_first and kwargs.get("tools") and not has_paired_tool_result:
            output = [{"type": "function_call", "id": "fc-1", "call_id": "anchor-1",
                       "name": "resolve_source_anchor", "arguments": json.dumps({
                           "evidence_id": source["source_unit"].evidence_id,
                           "quote": "对象乙", "context_text": None,
                       })}]
        else:
            answer = copy.deepcopy(proposal) if view["stage"] == "discovery" else {
                "verifications": [{
                    "target_id": target["target_id"], "content_hash": target["content_hash"],
                    "facets": [{"name": facet, "verdict": "supported",
                                "support": [source["quote"](source["source_unit"].text)],
                                "counterevidence_support": [], "reason": "原文支持"}
                               for facet in target["required_facets"]],
                } for target in view["verification_input"]["targets"]],
            }
            if invalid_first and len(requests) == 1:
                answer = {"unknown": "invalid response"}
            if transform is not None:
                transform(answer, view, len(requests))
            output = [{"type": "message", "id": "message", "role": "assistant",
                       "status": "completed", "content": [{"type": "output_text",
                                                              "text": json.dumps(answer)}]}]
        return ResponseTurn(response_id=f"response-{len(requests)}", output_items=output,
                            response_status="completed", incomplete_details=None, error=None,
                            usage={"input_tokens": 10, "output_tokens": 5})

    monkeypatch.setattr("app.services.llm.local_client.responses_create", transport)
    adapter = ToolModelRecognitionAdapter(
        object(), index=index, ontology=ontology, metadata=SimpleNamespace(node_summaries=[]),
        profile=ExtractionProfile(), token_counter=len, model_identity="qwen-test",
        max_input_tokens=1000000, max_output_tokens=2000, tool_limits=ToolLimits(20000),
        recognition_pipeline=RECORD_PIPELINE,
    )
    return adapter, task, context, card, storage, requests, proposal


def test_record_entities_and_each_owner_property_use_complete_verifier(source, monkeypatch):
    adapter, task, context, card, storage, requests, _ = setup_record(source, monkeypatch)
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete
    assert len(outcome.nodes) == 2
    assert len(outcome.properties) == 2
    assert not outcome.relationship_groups
    assert len(requests) == len(storage["reservations"]) == outcome.model_calls == 2
    first = json.loads(requests[0]["input_items"][0]["content"][0]["text"])
    assert "subject_ref" not in first and "predicate_iri" not in first
    assert "subject" not in first["task"]
    assert len(first["schema_card"]["class_cards"]) == 2
    assert not requests[0].get("tools")
    assert requests[0].get("tool_choice") is None
    verified = json.loads(requests[1]["input_items"][0]["content"][0]["text"])
    assert len(verified["verification_input"]["targets"]) == 4
    assert outcome.controller_checks["metric"] == outcome.controller_checks["shacl"] == 2
    for request in requests:
        offered = {tool["name"] for tool in request.get("tools") or []}
        assert "propose_mentions" not in offered
        assert offered <= {
            "inspect_evidence", "resolve_source_anchor",
            "find_referent_candidates", "check_claim_binding",
        }


def test_record_saved_request_cannot_restore_boundary_ner(source, monkeypatch):
    adapter, task, context, card, storage, _, _ = setup_record(source, monkeypatch)
    assert adapter.inspect_record(task, context, card).complete
    protocol = copy.deepcopy(storage["protocol"])
    protocol["stage"] = "discovery"
    protocol["request_attempt"] += 1
    protocol["pending_request"] = {
        "attempt": protocol["request_attempt"],
        "stage": "discovery",
        "request_hash": evidence_hash("saved-boundary-request"),
        "reservation_key": call_request_key({
            "lineage_id": protocol["lineage_id"],
            "protocol_attempt": protocol["request_attempt"],
        }),
        "allowed_tool_names": ["propose_mentions"],
    }
    with pytest.raises(ValueError, match="^record_protocol_tool_outside_scope$"):
        validate_tool_protocol(protocol)


def test_competing_record_types_use_entity_only_first_pass(source, monkeypatch):
    from app.services.llm import local_client

    adapter, task, context, card, _, requests, _ = setup_record(source, monkeypatch)
    adapter.record_discovery = adapter.record_discovery.model_copy(
        update={"max_classes_per_card": 1},
    )
    transport = local_client.responses_create

    def entity_only(client, **kwargs):
        view = json.loads(kwargs["input_items"][0]["content"][0]["text"])
        if view["stage"] == "discovery":
            assert kwargs["text_format"]["schema"]["properties"]["properties"][
                "maxItems"
            ] == 0
            assert "本轮只做实体识别" in kwargs["instructions"]
            turn = transport(client, **kwargs)
            payload = json.loads(turn.output_items[0]["content"][0]["text"])
            payload["properties"] = []
            turn.output_items[0]["content"][0]["text"] = json.dumps(payload)
            return turn
        return transport(client, **kwargs)

    monkeypatch.setattr(local_client, "responses_create", entity_only)
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete
    assert len(outcome.nodes) == 2
    assert not outcome.properties
    assert len(requests) == 2
    assert not requests[0].get("tools")


def test_dependency_aware_verification_batches_bound_twenty_three_targets():
    entities = [VersionedRef(id=f"{'e' * 63}{index}", revision=1) for index in range(6)]
    targets = []
    for index, entity in enumerate(entities):
        targets.append(SimpleNamespace(
            target_id=evidence_hash(["entity", index]), claim_ref=entity,
            dependency_refs=[], required_facets=["type", "referent", "subject_role"],
        ))
        property_count = (5, 3, 3, 2, 2, 2)[index]
        for ordinal in range(property_count):
            targets.append(SimpleNamespace(
                target_id=evidence_hash(["property", index, ordinal]),
                claim_ref=VersionedRef(
                    id=evidence_hash(["property-ref", index, ordinal]), revision=1,
                ),
                dependency_refs=[entity],
                required_facets=[
                    "subject_binding", "field_role", "predicate", "bridge", "value",
                    "qualifiers", "counterevidence",
                ],
            ))
    assert len(targets) == 23
    batches = plan_verification_target_batches(SimpleNamespace(targets=targets))
    assert len(batches) > 1
    assert {target_id for batch in batches for target_id in batch} == {
        target.target_id for target in targets
    }
    assert all(1 <= len(batch) <= 8 for batch in batches)
    by_id = {target.target_id: target for target in targets}
    assert all(
        sum(len(by_id[target_id].required_facets) for target_id in batch) <= 48
        for batch in batches
    )


def test_truncated_completed_json_is_split_without_replaying_fragment(source, monkeypatch):
    from app.services.llm import local_client

    adapter, task, context, card, storage, requests, _ = setup_record(
        source, monkeypatch, budget=10,
    )
    transport = local_client.responses_create
    truncated = False
    marker = "TRUNCATED_FRAGMENT_MUST_NOT_REPLAY"

    def truncate_first_verification(client, **kwargs):
        nonlocal truncated
        view = json.loads(kwargs["input_items"][0]["content"][0]["text"])
        turn = transport(client, **kwargs)
        if view["stage"] == "verification" and not truncated:
            truncated = True
            return replace(
                turn,
                output_items=[{
                    "type": "message", "role": "assistant", "status": "completed",
                    "content": [{"type": "output_text",
                                 "text": '{"verifications":["' + marker}],
                }],
                response_status="completed", incomplete_details=None,
                usage={"input_tokens": 10, "output_tokens": adapter.max_output_tokens},
            )
        return turn

    monkeypatch.setattr(local_client, "responses_create", truncate_first_verification)
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete
    verification_views = [
        json.loads(request["input_items"][0]["content"][0]["text"])
        for request in requests
        if json.loads(request["input_items"][0]["content"][0]["text"])["stage"]
        == "verification"
    ]
    assert [len(view["verification_input"]["targets"]) for view in verification_views] == [4, 2, 2]
    verification_texts = [
        request["input_items"][0]["content"][0]["text"]
        for request in requests
        if json.loads(request["input_items"][0]["content"][0]["text"])["stage"]
        == "verification"
    ]
    stable_prefixes = [text.split(',"verification_input":', 1)[0]
                       for text in verification_texts]
    assert len(set(stable_prefixes)) == 1
    assert all(marker not in json.dumps(request["input_items"], ensure_ascii=False)
               for request in requests[2:])
    batches = storage["protocol"]["verification_batches"]["batches"]
    assert [batch["status"] for batch in batches].count("split") == 1
    assert [batch["status"] for batch in batches].count("completed") == 2


def test_completed_verification_batch_survives_single_target_truncation_and_resume(
    source, monkeypatch,
):
    from app.services.llm import local_client

    adapter, task, context, card, storage, requests, _ = setup_record(
        source, monkeypatch, budget=12,
    )
    transport = local_client.responses_create
    verification_call = 0

    def split_then_fail_one(client, **kwargs):
        nonlocal verification_call
        view = json.loads(kwargs["input_items"][0]["content"][0]["text"])
        turn = transport(client, **kwargs)
        if view["stage"] != "verification":
            return turn
        verification_call += 1
        if verification_call in {1, 3, 4}:
            return replace(
                turn,
                output_items=[{
                    "type": "message", "role": "assistant", "status": "completed",
                    "content": [{"type": "output_text", "text": '{"verifications":['}],
                }],
                usage={"input_tokens": 10, "output_tokens": adapter.max_output_tokens},
            )
        return turn

    monkeypatch.setattr(local_client, "responses_create", split_then_fail_one)
    with pytest.raises(RecognitionModelFailure, match="model_output_truncated"):
        adapter.inspect_record(task, context, card)
    batches = storage["protocol"]["verification_batches"]["batches"]
    completed_ids = {
        target_id for batch in batches if batch["status"] == "completed"
        for target_id in batch["target_ids"]
    }
    assert completed_ids
    assert any(batch["status"] == "failed" for batch in batches)
    paid_before_resume = len(requests)

    context.remaining_model_calls = 12
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete
    resumed_views = [
        json.loads(request["input_items"][0]["content"][0]["text"])
        for request in requests[paid_before_resume:]
    ]
    resumed_target_ids = {
        target["target_id"] for view in resumed_views
        if view["stage"] == "verification"
        for target in view["verification_input"]["targets"]
    }
    assert completed_ids.isdisjoint(resumed_target_ids)


@pytest.mark.parametrize("stop_at", ["model_turn", "tool_result", "discovery", "verification"])
def test_record_cold_resume_reuses_paid_responses_and_verification_tool_results(
    source, monkeypatch, stop_at,
):
    adapter, task, context, card, saved, paid, _ = setup_record(
        source, monkeypatch, stop_at=stop_at, tool_first=True,
    )
    with pytest.raises(RuntimeError, match="pause after durable commit"):
        adapter.inspect_record(task, context, card)
    saved = copy.deepcopy(saved)
    paid_count = len(paid)
    adapter, task, context, card, storage, resumed, _ = setup_record(
        source, monkeypatch, tool_first=True,
    )
    storage.update(copy.deepcopy(saved))
    context.protocol_state, context.protocol_results = saved["protocol"], saved["results"]
    context.remaining_model_calls = 6 - paid_count
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete and len(outcome.properties) == 2
    assert paid_count + len(resumed) == 3
    assert len(storage["reservations"]) == 3
    assert storage["protocol"]["request_attempt"] == 3
    tools = [v for v in storage["results"].values() if v["field"] == "tool_result"]
    assert len(tools) == 1
    assert tools[0]["value"]["result"]["status"] == "ok"
    if resumed and stop_at == "tool_result":
        assert len([item for item in resumed[0]["input_items"]
                    if item.get("type") == "function_call_output"]) == 1


def test_empty_record_and_propertyless_class_never_pay_empty_verification(source, monkeypatch):
    adapter, task, context, card, storage, requests, _ = setup_record(
        source, monkeypatch, empty=True, no_properties=True,
    )
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete and not outcome.nodes and not outcome.properties
    assert outcome.model_calls == len(requests) == 1
    assert storage["protocol"]["verification_ref"]


def test_record_malformed_answer_corrects_without_changing_task_or_schema(source, monkeypatch):
    adapter, task, context, card, _, requests, _ = setup_record(
        source, monkeypatch, invalid_first=True,
    )
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete and len(outcome.properties) == 2
    assert outcome.model_calls == 3
    assert requests[0]["text_format"] == requests[1]["text_format"]
    assert "stage_answer_invalid" in json.dumps(requests[1]["input_items"])
    assert "invalid response" not in json.dumps(requests[1]["input_items"])


def test_record_context_capacity_failure_is_incomplete_without_paid_call(source, monkeypatch):
    adapter, task, context, card, storage, requests, _ = setup_record(source, monkeypatch)
    adapter.max_input_tokens = 1
    with pytest.raises(RecognitionModelFailure, match="context_budget_exceeded"):
        adapter.inspect_record(task, context, card)
    assert not requests and not storage["reservations"]
    assert storage["protocol"]["task"] == task.model_dump(mode="json")
    assert storage["protocol"]["discovery_ref"] is None


def test_record_feedback_reopens_only_published_result_and_keeps_lineage_budget(
    source, monkeypatch,
):
    adapter, task, context, card, storage, requests, _ = setup_record(
        source, monkeypatch, budget=4,
    )
    first = adapter.inspect_record(task, context, card)
    assert first.model_calls == 2 and first.complete
    original = copy.deepcopy(storage["protocol"])
    feedback_hash = evidence_hash("missing candidate evidence")
    context.tool_inputs.update(
        record_feedback_hash=feedback_hash,
        record_feedback={"hash": feedback_hash, "source_refs": [
            context.fragments[0].anchor.model_dump(mode="json"),
        ], "refs": [{"id": "relation-task", "revision": 1}], "reason": "缺少urn:T端点"},
    )
    context.remaining_model_calls = 2
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.model_calls == 2 and outcome.complete
    assert len(requests) == len(storage["reservations"]) == 4
    protocol = storage["protocol"]
    assert protocol["request_attempt"] == 4
    assert protocol["completed_attempts"] == [1, 2, 3, 4]
    assert protocol["task"] == original["task"]
    assert protocol["lineage_id"] == original["lineage_id"]
    assert protocol["assertion_generation"] == protocol["evidence_revision"] == 2
    assert protocol["record_feedback_hash"] == feedback_hash
    first_view = json.loads(requests[0]["input_items"][0]["content"][0]["text"])
    second_view = json.loads(requests[2]["input_items"][0]["content"][0]["text"])
    assert first_view["schema_card"] == second_view["schema_card"]
    assert first_view["evidence_units"] == second_view["evidence_units"]
    assert second_view["feedback"][0]["hash"] == feedback_hash
    assert second_view["feedback"][0]["reason"] == "缺少urn:T端点"
    context.tool_inputs["record_feedback_hash"] = evidence_hash("another feedback")
    context.remaining_model_calls = 0
    with pytest.raises(RecognitionModelFailure, match="model_budget_exhausted"):
        adapter.inspect_record(task, context, card)
    assert storage["protocol"] == protocol
    assert len(requests) == 4


def test_record_feedback_cannot_replace_an_unfinished_paid_attempt(source, monkeypatch):
    adapter, task, context, card, storage, requests, _ = setup_record(
        source, monkeypatch, stop_at="discovery",
    )
    with pytest.raises(RuntimeError, match="pause after durable commit"):
        adapter.inspect_record(task, context, card)
    original = copy.deepcopy(storage["protocol"])
    context.tool_inputs["record_feedback_hash"] = evidence_hash("too early feedback")
    with pytest.raises(ValueError, match="record_feedback_refresh_not_ready"):
        adapter.inspect_record(task, context, card)
    assert storage["protocol"] == original and len(requests) == 1


@pytest.mark.parametrize("partial_composition", [False, True])
def test_record_entity_requires_every_component_while_other_entity_survives(
    source, monkeypatch, partial_composition,
):
    from app.services.llm import local_client

    adapter, task, context, card, _, requests, proposal = setup_record(source, monkeypatch)
    proposal["entities"][0].update(representation="record", mentions=[], record_components=[
        {"role": role, "quote": source["quote"](text)}
        for role, text in (("subject", "对象乙"), ("field", "数量"), ("value", "5"))
    ])
    transport = local_client.responses_create

    def verify_components(client, **kwargs):
        response = transport(client, **kwargs)
        view = json.loads(kwargs["input_items"][0]["content"][0]["text"])
        if partial_composition and view["stage"] == "verification":
            answer = json.loads(response.output_items[0]["content"][0]["text"])
            target_id = next(target["target_id"] for target in view["verification_input"]["targets"]
                             if target["payload"].get("local_id") == "b")
            for target in answer["verifications"]:
                if target["target_id"] == target_id:
                    for facet in target["facets"]:
                        if facet["name"] == "referent":
                            facet["support"] = [source["quote"]("对象乙")]
            response.output_items[0]["content"][0]["text"] = json.dumps(answer)
        return response

    monkeypatch.setattr(local_client, "responses_create", verify_components)
    outcome = adapter.inspect_record(task, context, card)
    assert outcome.complete is (not partial_composition)
    if partial_composition:
        assert outcome.reason_code == "record_composition_source_coverage_missing"
    # Missing composition proof is downgraded locally; it never resends the
    # same source and full verification answer for correction.
    assert len(requests) == 2
    assert any(node.class_iri == "urn:Other" for node in outcome.nodes)
    records = [node for node in outcome.nodes if node.grounding_kind == "record"]
    if partial_composition:
        assert not records
        assert len(outcome.properties) == 1
    else:
        assert len(records) == 1
        assert records[0].composition_decision_ref is not None
        assert len(outcome.properties) == 2


@pytest.mark.parametrize("generation", [2, 5])
def test_record_feedback_cannot_recreate_registered_composition_at_new_generation(
    source, monkeypatch, generation,
):
    from app.services.extraction.ontology_guided.claim_freeze import freeze_record_proposal
    from app.services.extraction.ontology_guided.claim_protocol import (
        DiscoveryEnvelope,
        EntityDependencyView,
    )

    _, task, context, card, _, _, proposal = setup_record(source, monkeypatch)
    proposal["entities"][0].update(representation="record", mentions=[], record_components=[
        {"role": role, "quote": source["quote"](text)}
        for role, text in (("subject", "对象乙"), ("field", "数量"), ("value", "5"))
    ])
    proposal = DiscoveryEnvelope.model_validate(proposal)
    index = source["index"]
    unit = source["source_unit"]
    dependency = dict(
        entity_ref=VersionedRef(id="already-registered", revision=1), class_iri="urn:T",
        grounding_kind="record", root_origin=None, proposal=proposal.entities[0],
        source_refs=[index.ir.anchor(unit.evidence_id, 0, len(unit.text))], dependency_refs=[],
    )
    registered = EntityDependencyView(**dependency, content_hash=evidence_hash(dependency))
    frozen = freeze_record_proposal(
        proposal, task=task, context=context, card=card, index=index,
        generation=generation, entity_dependencies=[registered], reference_resolution=True,
    )
    assert frozen.claim_issues["b"] == ["record_entity_already_registered"]
    assert "c" not in frozen.claim_issues
    assert frozen.local_ref_map["already-registered"] == registered.entity_ref
