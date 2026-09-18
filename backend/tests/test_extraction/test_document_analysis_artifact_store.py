from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from fastapi import UploadFile

from app.services.document_analysis.artifact_store import (
    RunArtifactStorage,
    SourceArtifactError,
    stage_upload,
)


def _docx_bytes(path) -> bytes:
    document = Document()
    document.add_heading("质量概要", level=1)
    document.add_paragraph("产品甲由本报告明确描述。")
    document.save(path)
    return path.read_bytes()


@pytest.mark.asyncio
async def test_upload_is_streamed_hashed_and_confined_to_its_run(tmp_path):
    run_id = uuid4()
    raw = _docx_bytes(tmp_path / "source.docx")
    upload = UploadFile(filename="../../报告.docx", file=BytesIO(raw))

    staged = await stage_upload(
        upload,
        storage_root=tmp_path / "artifacts",
        run_id=run_id,
        max_upload_bytes=len(raw) + 10,
    )

    assert staged.filename == "报告.docx"
    assert staged.path.read_bytes() == raw
    assert staged.storage_uri == f"{run_id}/source/original.docx"
    assert (
        RunArtifactStorage(tmp_path / "artifacts").resolve(run_id, staged.storage_uri)
        == staged.path
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "content", "expected_code"),
    [
        ("empty.docx", b"", "EMPTY_SOURCE"),
        ("notes.txt", b"plain", "UNSUPPORTED_SOURCE_TYPE"),
        ("fake.docx", b"not-a-zip", "UNSUPPORTED_SOURCE_TYPE"),
        ("fake.doc", b"not-an-ole-file", "UNSUPPORTED_SOURCE_TYPE"),
    ],
)
async def test_upload_rejects_empty_extension_and_magic_mismatches(
    tmp_path, filename, content, expected_code
):
    with pytest.raises(SourceArtifactError) as caught:
        await stage_upload(
            UploadFile(filename=filename, file=BytesIO(content)),
            storage_root=tmp_path / "artifacts",
            run_id=uuid4(),
            max_upload_bytes=1024,
        )
    assert caught.value.code == expected_code


@pytest.mark.asyncio
async def test_upload_size_limit_removes_partial_run_directory(tmp_path):
    run_id = uuid4()
    with pytest.raises(SourceArtifactError) as caught:
        await stage_upload(
            UploadFile(filename="large.docx", file=BytesIO(b"x" * 32)),
            storage_root=tmp_path / "artifacts",
            run_id=run_id,
            max_upload_bytes=16,
        )
    assert caught.value.code == "SOURCE_TOO_LARGE"
    assert not (tmp_path / "artifacts" / str(run_id)).exists()


def test_storage_rejects_foreign_and_traversal_references(tmp_path):
    storage = RunArtifactStorage(tmp_path / "artifacts")
    run_id = uuid4()
    for reference in ("../secret", f"{uuid4()}/source/original.docx", "/etc/passwd"):
        with pytest.raises(SourceArtifactError):
            storage.resolve(run_id, reference)
