"""Quality scheduler/routing contracts; extraction verdicts are explicit fixtures."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.quality_guided_variant import build_quality_guided_variant
from app.schemas.evidence import (
    BindingEvidence,
    Candidate,
    DocumentProvenance,
    ExtractionTask,
    LiteralValue,
    TaskBudget,
)
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.evidence_scope import scope_contains
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:Report": {
        "label": "Report",
        "properties": [{"iri": "urn:code", "label": "Document identifier", "datatype": "string"}],
        "relationships": [{"iri": "urn:describes", "label": "describes", "range": ["urn:Item"]}],
    },
    "urn:Item": {
        "label": "Item",
        "properties": [{"iri": "urn:size", "label": "Size", "datatype": "decimal"}],
        "relationships": [{"iri": "urn:uses", "label": "uses", "range": ["urn:Device"]}],
    },
    "urn:Device": {"label": "Device", "properties": [], "relationships": []},
    "urn:Unreachable": {"label": "Unreachable", "properties": [], "relationships": []},
}


@pytest.fixture
def ir(tmp_path):
    doc = Document()
    doc.add_heading("Report", level=1)
    table = doc.add_table(rows=2, cols=3)
    for row, values in zip(
        table.rows, [("Item", "Size", "Device"), ("Item A", "5", "Device D")], strict=True
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    doc.add_paragraph("Document code R1.")
    doc.add_paragraph("Administrative note.")
    path = tmp_path / "quality-routing.docx"
    doc.save(path)
    return analyze_word_core(path).ir


def make_runner(ir, model=None):
    base = GenericExtractionRunner(
        deepcopy(SCHEMA),
        TestTokenizer(),
        model or (lambda *args: pytest.fail("unexpected model call")),
        model_identity="quality-routing-fixture",
        compact_identifiers=False,
        budget=TaskBudget(max_input_tokens=100000, max_output_tokens=4096, max_tasks=100),
    )
    return build_quality_guided_variant(base, ir)


def entity(ir, text, *, cls="urn:Item", cid="item-a"):
    source = next(u for u in ir.evidence_units if u.text == text)
    return Candidate(
        candidate_id=cid,
        kind="entity",
        class_iri=cls,
        text=text,
        validation_status="passed",
        provenance=[DocumentProvenance(anchors=[ir.anchor(source.evidence_id)], excerpts=[text])],
    )


def root(ir):
    return entity(ir, "Report", cls="urn:Report", cid="report").model_copy(
        update={"identity": {"document_root": ir.document_hash}}
    )


def predicates(iri):
    definition = SCHEMA[iri]
    return [
        (kind, p)
        for kind, field in (("property", "properties"), ("relationship", "relationships"))
        for p in definition.get(field, [])
    ]


def routing_task(runner, kind, regions, **kwargs):
    return ExtractionTask(
        task_id="routing:" + evidence_hash([kind, regions, kwargs]),
        task_kind=kind,
        target_evidence_ids=[r.evidence_id for r in regions],
        target_ranges=regions,
        budget=runner.budget,
        **kwargs,
    )


def test_route_requests_only_contain_current_direct_menu_and_complete_records(ir):
    requests = []

    def model(system, user, schema, budget):
        payload = json.loads(user)
        requests.append(payload)
        return {
            "records": [
                {"record_id": record_id, "predicate_ids": []} for record_id in payload["records"]
            ]
        }

    runner = make_runner(ir, model)
    runner.route_batch_size = 1
    result = runner._routes(
        root(ir),
        predicates("urn:Report"),
        lambda kind, regions, **kw: routing_task(runner, kind, regions, **kw),
    )
    expected = {r.record_id for r in runner.record_index.records if r.kind != "heading"}
    assert set(result["routes"]) == expected
    assert all(entry["status"] == "route_negative_not_extracted" for entry in result["ledger"])
    for request in requests:
        assert {p["iri"] for p in request["local_menu"].values()} == {"urn:code", "urn:describes"}
        assert "urn:uses" not in json.dumps(request["local_menu"])
        assert "urn:Device" not in json.dumps(request["local_menu"])
        assert "urn:Unreachable" not in json.dumps(request)
    table = next(
        record
        for request in requests
        for record in request["records"].values()
        if record["kind"] == "table_row"
    )
    assert [cell["fragments"] for cell in table["cells"]] == [["Item A"], ["5"], ["Device D"]]
    assert [cell["column_headers"][0]["text"] for cell in table["cells"]] == [
        "Item",
        "Size",
        "Device",
    ]
    assert [cell["column_indices"] for cell in table["cells"]] == [[0], [1], [2]]
    assert "source" not in table and "column_headers" not in table
    count = len(requests)
    runner._routes(
        root(ir),
        predicates("urn:Report"),
        lambda kind, regions, **kw: routing_task(runner, kind, regions, **kw),
    )
    assert len(requests) == count


def test_routing_cells_pair_multiline_and_merged_values_with_actual_headers(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Item", "Size", "Device"], strict=True):
        cell.text = text
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Item A"
    table.cell(1, 1).text = "5"
    table.cell(1, 1).add_paragraph("millimetres")
    table.cell(1, 2).text = "Device D"
    table.cell(2, 1).text = "6"
    table.cell(2, 2).text = "Device E"
    path = tmp_path / "routing-multiline.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    requests = []

    def model(system, user, schema, budget):
        payload = json.loads(user)
        requests.append(payload)
        return {"records": [{"record_id": key, "predicate_ids": []} for key in payload["records"]]}

    runner = make_runner(ir, model)
    runner._routes(
        entity(ir, "Item A", cls="urn:Report", cid="report"),
        predicates("urn:Report"),
        lambda kind, regions, **kw: routing_task(runner, kind, regions, **kw),
    )
    rows = [r for request in requests for r in request["records"].values()]
    assert len(rows) == 2
    assert rows[0]["cells"][1]["fragments"] == ["5", "millimetres"]
    assert rows[0]["cells"][1]["column_headers"] == [{"column_indices": [1], "text": "Size"}]
    assert rows[0]["cells"][2]["column_headers"] == [{"column_indices": [2], "text": "Device"}]
    assert rows[0]["cells"][0]["data_row_indices"] == [1, 2]
    assert rows[1]["cells"][0]["shared_across_rows"]
    assert rows[1]["cells"][1]["fragments"] == ["6"]


@pytest.mark.parametrize(
    "malformation", ["missing_record", "duplicate_record", "unknown_predicate"]
)
def test_route_response_must_cover_each_record_once_and_use_known_predicates(ir, malformation):
    def model(system, user, schema, budget):
        payload = json.loads(user)
        values = [{"record_id": key, "predicate_ids": []} for key in payload["records"]]
        if malformation == "missing_record":
            values.pop()
        elif malformation == "duplicate_record":
            values.append(values[0])
        else:
            values[0]["predicate_ids"] = ["invented-predicate"]
        return {"records": values}

    runner = make_runner(ir, model)
    with pytest.raises(ValueError, match="record_route_"):
        runner._routes(
            root(ir),
            predicates("urn:Report"),
            lambda kind, regions, **kw: routing_task(runner, kind, regions, **kw),
        )
    assert not runner._route_cache


def test_record_scope_retains_authorized_rows_and_requires_reference_for_external_record(ir):
    runner = make_runner(ir)
    subject = entity(ir, "Item A")
    predicate = SCHEMA["urn:Item"]["properties"][0]
    plan = runner.record_index.plan(subject, predicate, kind="property")
    row = next(r for r in plan.records if r.kind == "table_row")
    paragraph = next(
        r
        for r in plan.records
        if any(ir.unit(p.evidence_id).text == "Document code R1." for p in r.source_ranges)
    )
    original = plan.authorized_scope.model_dump(mode="json")
    assert runner._record_scope(plan, row) == plan.authorized_scope
    external = runner._record_scope(plan, paragraph)
    assert external.reference_ranges == paragraph.retrieved_target_ranges
    assert external.construction_evidence == plan.authorized_scope.construction_evidence
    assert plan.authorized_scope.model_dump(mode="json") == original
    assert all(
        scope_contains(external, ir.anchor(r.evidence_id, r.start, r.end), ir)
        for r in paragraph.retrieved_target_ranges
    )
    assert not paragraph.target_ranges


def install_scheduler_fixtures(monkeypatch, runner, ir, *, reject_root=False):
    events = []

    def routes(subject, menu, make_task):
        events.append(("routes", subject.class_iri, [p["iri"] for _, p in menu]))
        expected = {p["iri"] for _, p in predicates(subject.class_iri)}
        assert {p["iri"] for _, p in menu} == expected
        selected = {}
        for record in runner.record_index.records:
            if record.kind == "heading":
                continue
            text = " ".join(ir.unit(r.evidence_id).text for r in record.source_ranges)
            if subject.class_iri == "urn:Report":
                selected[record.record_id] = (
                    ["urn:describes"]
                    if "Item A" in text
                    else ["urn:code"]
                    if "Document code" in text
                    else []
                )
            elif subject.class_iri == "urn:Item":
                selected[record.record_id] = ["urn:size", "urn:uses"] if "Item A" in text else []
            else:
                selected[record.record_id] = []
        value = {
            "routes": selected,
            "ledger": [
                {
                    "record_id": rid,
                    "predicate_iris": iris,
                    "status": "routed" if iris else "route_negative_not_extracted",
                }
                for rid, iris in selected.items()
            ],
        }
        runner._route_cache[subject.candidate_id] = value
        return value

    def execute_task(task, source, current, effective_class=""):
        runner._task_events = []
        record = runner._record_for_task[task.task_id]
        assert task.target_ranges == record.retrieved_target_ranges
        assert set(task.target_evidence_ids) == {r.evidence_id for r in record.source_ranges}
        for ref in [task.subject, task.path_root, *task.competing_subjects, *task.dependency_refs]:
            if ref is not None:
                assert (
                    ref.candidate_id in current
                    and current[ref.candidate_id].revision == ref.revision
                )
        events.append(
            (task.task_kind, task.predicate_iri, list(task.target_class_iris), task.task_id)
        )
        if task.task_kind == "entity":
            values = []
            for cls, text, cid in (
                ("urn:Item", "Item A", "item-a"),
                ("urn:Device", "Device D", "device-d"),
            ):
                if cls in task.target_class_iris:
                    candidate = entity(ir, text, cls=cls, cid=cid)
                    candidate.task_id = task.task_id
                    values.append(candidate)
            return values
        anchors = [source.anchor(r.evidence_id, r.start, r.end) for r in task.target_ranges]
        targets = task.object_candidates if task.task_kind == "relationship" else [None]
        values = []
        for target in targets:
            literal = None
            if task.task_kind == "property":
                literal = (
                    LiteralValue(kind="text", raw_value="R1", normalized_value="R1")
                    if task.predicate_iri == "urn:code"
                    else LiteralValue(
                        kind="number",
                        raw_value="5",
                        normalized_value="5",
                        datatype_iri="http://www.w3.org/2001/XMLSchema#decimal",
                    )
                )
            values.append(
                Candidate(
                    candidate_id="assertion:" + evidence_hash([task.task_id, target]),
                    kind=task.task_kind,
                    subject=task.subject,
                    object=target,
                    predicate_iri=task.predicate_iri,
                    literal=literal,
                    task_id=task.task_id,
                    scope=task.scope,
                    path_root=task.path_root,
                    relationship_path=task.relationship_path,
                    dependency_refs=task.dependency_refs,
                    validation_status="rejected"
                    if reject_root and task.predicate_iri == "urn:describes"
                    else "passed",
                    provenance=[
                        DocumentProvenance(
                            anchors=anchors, excerpts=[source.resolve(a) for a in anchors]
                        )
                    ],
                    bindings=[
                        BindingEvidence(
                            method="explicit_assertion",
                            subject_candidate_id=task.subject.candidate_id,
                            predicate_iri=task.predicate_iri,
                            object_candidate_id=target.candidate_id if target else None,
                            anchors=anchors,
                        )
                    ],
                )
            )
        return values

    monkeypatch.setattr(runner, "_routes", routes)
    monkeypatch.setattr(runner, "execute_task", execute_task)
    return events


def test_verified_record_edge_expands_child_before_other_root_records(monkeypatch, ir):
    runner = make_runner(ir)
    events = install_scheduler_fixtures(monkeypatch, runner, ir)
    result = runner.run(ir, effective_class="urn:Report")
    recalls = [event for event in events if event[0] == "entity"]
    assert recalls[0][2] == ["urn:Item"]
    assert any(event[2] == ["urn:Device"] for event in recalls)
    root_edge = next(
        i for i, event in enumerate(events) if event[:2] == ("relationship", "urn:describes")
    )
    child_route = next(i for i, event in enumerate(events) if event[:2] == ("routes", "urn:Item"))
    other_root = next(i for i, event in enumerate(events) if event[:2] == ("property", "urn:code"))
    assert root_edge < child_route < other_root
    assert any(
        c.kind == "relationship" and len(c.relationship_path) == 2 and c.positive_eligible
        for c in result.candidates
    )
    assert "route_negative_records_not_exhaustively_extracted" in result.diagnostics
    assert result.completion == "incomplete"
    negatives = [
        item
        for route in runner.plan["routes"].values()
        for item in route["ledger"]
        if item["status"] == "route_negative_not_extracted"
    ]
    assert negatives
    admin = next(
        record.record_id
        for record in runner.record_index.records
        if any(ir.unit(r.evidence_id).text == "Administrative note." for r in record.source_ranges)
    )
    assert not any(c["record_id"] == admin for c in runner.coverage)


def test_rejected_root_edge_does_not_route_or_extract_child(monkeypatch, ir):
    runner = make_runner(ir)
    events = install_scheduler_fixtures(monkeypatch, runner, ir, reject_root=True)
    result = runner.run(ir, effective_class="urn:Report")
    assert not any(event[:2] == ("routes", "urn:Item") for event in events)
    assert not any(event[:2] == ("property", "urn:size") for event in events)
    assert not any(c.kind == "relationship" and c.positive_eligible for c in result.candidates)


def test_focus_path_validates_formal_hops_and_changes_checkpoint_identity(ir):
    runner = make_runner(ir)
    original = runner.input_id(ir, "urn:Report")
    runner.focus_path = ("urn:describes", "urn:uses")
    runner._validate_focus_path("urn:Report")
    assert runner.input_id(ir, "urn:Report") != original
    runner.focus_path = ("urn:uses",)
    with pytest.raises(ValueError, match="focus_path_not_reachable"):
        runner.run(ir, effective_class="urn:Report")
    runner.focus_path = ("urn:describes", "urn:uses")
    runner.budget = runner.budget.model_copy(update={"max_hops": 1})
    with pytest.raises(ValueError, match="focus_path_exceeds_max_hops"):
        runner.run(ir, effective_class="urn:Report")


def test_resume_at_record_boundary_keeps_completed_record_and_remaining_queue(monkeypatch, ir):
    runner = make_runner(ir)
    events = install_scheduler_fixtures(monkeypatch, runner, ir)
    first = runner.run(ir, effective_class="urn:Report", pause_after=2)
    assert first.checkpoint["queue"] and first.completion == "incomplete"
    assert [event[0] for event in events if event[0] != "routes"] == ["entity", "relationship"]
    completed_ids = [task["task"]["task_id"] for task in first.tasks]
    second = runner.run(ir, effective_class="urn:Report", checkpoint=deepcopy(first.checkpoint))
    assert not second.checkpoint["queue"]
    executed = [event[3] for event in events if event[0] != "routes"]
    assert all(executed.count(task_id) == 1 for task_id in completed_ids)
    assert any(c.predicate_iri == "urn:size" and c.positive_eligible for c in second.candidates)
    count = len(events)
    third = runner.run(ir, effective_class="urn:Report", checkpoint=deepcopy(second.checkpoint))
    assert len(events) == count
    assert third.candidates == second.candidates
    wrong = deepcopy(second.checkpoint)
    wrong["input_id"] = "other-input"
    with pytest.raises(ValueError, match="quality_checkpoint_identity_mismatch"):
        runner.run(ir, effective_class="urn:Report", checkpoint=wrong)


def test_many_records_for_one_predicate_do_not_delay_other_predicates(monkeypatch, tmp_path):
    doc = Document()
    doc.add_heading("Report", level=1)
    for index in range(7):
        doc.add_paragraph(f"Item frequent record {index}.")
    doc.add_paragraph("Device singleton record.")
    path = tmp_path / "predicate-fairness.docx"
    doc.save(path)
    source = analyze_word_core(path).ir
    runner = make_runner(source)
    runner.schema["urn:Report"]["relationships"] = [
        {"iri": "urn:busy", "label": "Item", "range": ["urn:Item"]},
        {"iri": "urn:rare", "label": "Device", "range": ["urn:Device"]},
    ]
    from app.evaluation.record_retrieval import RecordIndex

    runner.record_index = RecordIndex(source, runner.schema)
    seen, menus = [], []

    def routes(subject, menu, make_task):
        menus.append([(kind, p["iri"]) for kind, p in menu])
        assignments = {}
        for record in runner.record_index.records:
            if record.kind == "heading":
                continue
            text = " ".join(source.unit(r.evidence_id).text for r in record.source_ranges)
            assignments[record.record_id] = ["urn:busy"] if "frequent" in text else ["urn:rare"]
        result = {"routes": assignments, "ledger": []}
        runner._route_cache[subject.candidate_id] = result
        return result

    def execute(task, ir, current, effective_class=""):
        runner._task_events = []
        record = runner._record_for_task[task.task_id]
        assert task.target_ranges == record.retrieved_target_ranges
        seen.append(task.predicate_definition["discovery_relation"]["iri"])
        return []

    monkeypatch.setattr(runner, "_routes", routes)
    monkeypatch.setattr(runner, "execute_task", execute)
    result = runner.run(source, effective_class="urn:Report")
    assert menus[0] == [
        ("relationship", "urn:busy"),
        ("relationship", "urn:rare"),
        ("property", "urn:code"),
    ]
    assert seen[:2] == ["urn:busy", "urn:rare"]
    assert seen.count("urn:busy") == 7 and seen.count("urn:rare") == 1
    assert len(result.checkpoint["coverage"]) == 8
    assert result.checkpoint["record_turns"] == 8


def install_long_child_branch(monkeypatch, tmp_path):
    doc = Document()
    doc.add_heading("Report", level=1)
    table = doc.add_table(rows=2, cols=3)
    for row, values in zip(
        table.rows, [("Item", "Size", "Device"), ("Item A", "5", "Device D")], strict=True
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    for index in range(8):
        doc.add_paragraph(f"Child detail {index} with size 5.")
    doc.add_paragraph("Document code R1.")
    doc.add_paragraph("Document code R2.")
    path = tmp_path / "root-child-fairness.docx"
    doc.save(path)
    source = analyze_word_core(path).ir
    runner = make_runner(source)
    events = install_scheduler_fixtures(monkeypatch, runner, source)

    def routes(subject, menu, make_task):
        events.append(("routes", subject.class_iri, [p["iri"] for _, p in menu]))
        assignments = {}
        for record in runner.record_index.records:
            if record.kind == "heading":
                continue
            text = " ".join(source.unit(r.evidence_id).text for r in record.source_ranges)
            if subject.class_iri == "urn:Report":
                values = (
                    ["urn:describes"]
                    if "Item A" in text
                    else ["urn:code"]
                    if "Document code" in text
                    else []
                )
            elif subject.class_iri == "urn:Item":
                values = (
                    ["urn:uses", "urn:size"]
                    if "Item A" in text
                    else (["urn:size"] if "Child detail" in text else [])
                )
            else:
                values = []
            assignments[record.record_id] = values
        result = {"routes": assignments, "ledger": []}
        runner._route_cache[subject.candidate_id] = result
        return result

    monkeypatch.setattr(runner, "_routes", routes)
    return source, runner, events


def test_pending_root_records_advance_during_long_child_branch(monkeypatch, tmp_path):
    source, runner, _ = install_long_child_branch(monkeypatch, tmp_path)
    result = runner.run(source, effective_class="urn:Report")
    coverage = result.checkpoint["coverage"]
    root_slots = [index for index, record in enumerate(coverage) if not record["relationship_path"]]
    assert root_slots == [0, 2, 5]
    assert len(coverage) > root_slots[-1] + 3  # Root work progresses before children finish.
    last_root, longest, current = root_slots[-1], 0, 0
    for record in coverage[: last_root + 1]:
        current = current + 1 if record["relationship_path"] else 0
        longest = max(longest, current)
    assert longest <= 2
    assert result.checkpoint["record_turns"] == len(coverage)


def test_resume_preserves_root_fairness_turn_after_two_records(monkeypatch, tmp_path):
    source, runner, events = install_long_child_branch(monkeypatch, tmp_path)
    # The root relation and first child relation each need discovery plus binding.
    first = runner.run(source, effective_class="urn:Report", pause_after=4)
    assert first.checkpoint["record_turns"] == 2
    assert first.checkpoint["queue"]
    previous_tasks = {event[3] for event in events if event[0] != "routes"}
    boundary = len(events)
    second = runner.run(source, effective_class="urn:Report", checkpoint=deepcopy(first.checkpoint))
    resumed = [event for event in events[boundary:] if event[0] != "routes"]
    assert resumed[0][:2] == ("property", "urn:code")
    assert previous_tasks.isdisjoint(event[3] for event in resumed)
    coverage = second.checkpoint["coverage"]
    assert [i for i, record in enumerate(coverage) if not record["relationship_path"]] == [0, 2, 5]
    assert second.checkpoint["record_turns"] == len(coverage)
