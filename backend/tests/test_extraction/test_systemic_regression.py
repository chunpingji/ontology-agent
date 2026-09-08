"""Replay the systemic report's failures with source documents and independent verdicts.

Scripted verdicts isolate contracts; these tests do not measure model recall.
"""

import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document

from app.schemas.evidence import (
    BindingEvidence,
    Candidate,
    DocumentProvenance,
    EvidenceRange,
    ExtractionTask,
    TaskBudget,
)
from app.services.extraction.evidence_preview import preview_relationships
from app.services.extraction.evidence_scope import build_scope, candidate_ref, scope_contains
from app.services.extraction.extraction_tasks import GenericExtractionRunner, PartialTaskFailure
from app.services.extraction.gap_extraction import expanded_gap_scope
from app.services.extraction.hierarchical_context import build_context
from app.services.extraction.literal_normalizer import normalize_literal
from app.services.extraction.table_records import TableRecords
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.local_client import StructuredModelError, chat_with_schema
from app.services.ontology_instance_writer import instance_iri
from tests.test_extraction.test_hierarchical_context import TestTokenizer


def analyze(document, tmp_path):
    path = tmp_path / "systemic.docx"
    document.save(path)
    return analyze_word_core(path).ir


def unit(ir, text):
    return next(u for u in ir.evidence_units if u.text == text)


def quote(unit, text=None, **kwargs):
    return {"evidence_id": unit.evidence_id, "text": unit.text if text is None else text, **kwargs}


def entity(ir, source, name=None, identity="subject", cls="urn:Product"):
    name = name or source.text
    start = source.text.index(name)
    return Candidate(
        candidate_id=identity,
        kind="entity",
        class_iri=cls,
        text=name,
        validation_status="passed",
        provenance=[
            DocumentProvenance(
                anchors=[ir.anchor(source.evidence_id, start, start + len(name))],
                excerpts=[name],
            )
        ],
    )


def property_task(ir, subject, target, *, datatype="string", canonical_unit=None, scope=None):
    return ExtractionTask(
        task_id="property",
        task_kind="property",
        subject=candidate_ref(subject),
        predicate_iri="urn:value",
        predicate_definition={
            "datatype": datatype,
            "canonical_unit": canonical_unit,
        },
        target_evidence_ids=[target.evidence_id],
        scope=scope or build_scope(ir, subject),
        budget=TaskBudget(max_input_tokens=100000),
    )


def execute_property(ir, subject, task, proposal, spans, **verdict):
    stages = []

    def model(system, user, schema, budget):
        request = json.loads(user)
        stages.append(request["stage"])
        if request["stage"] == "recall":
            return {"assertions": [proposal]}
        if request["stage"] == "verify_reference":
            return {
                "supported": verdict.get("reference_supported", True),
                "subject_candidate_id": subject.candidate_id,
                "assertion_spans": spans,
                "reason": "独立核对前文主体与当前指代",
            }
        return {
            "supported": True,
            "subject_candidate_id": subject.candidate_id,
            "assertion_status": "affirmed",
            "method": "explicit_assertion",
            "assertion_spans": spans,
            **{k: v for k, v in verdict.items() if k != "reference_supported"},
        }

    runner = GenericExtractionRunner({}, TestTokenizer(), model, model_identity="systemic")
    return runner.execute_task(task, ir, {subject.candidate_id: subject}), stages


def table_document(tmp_path):
    doc = Document()
    table = doc.add_table(rows=4, cols=3)
    for row, values in zip(
        table.rows,
        [
            ["API", "试验", "包装方式"],
            ["HRS-1597", "大鼠14天", "双层袋"],
            ["", "犬14天", "复合袋"],
            ["HRS-5678", "另一试验", "纸袋"],
        ],
        strict=True,
    ):
        for cell, text in zip(row.cells, values, strict=True):
            cell.text = text
    table.cell(1, 0).merge(table.cell(2, 0))
    return analyze(doc, tmp_path)


def test_header_and_shared_cell_bind_to_specific_data_row(tmp_path):
    ir = table_document(tmp_path)
    shared = next(u for u in ir.evidence_units if "HRS-1597" in u.text)
    subject = entity(ir, shared, "HRS-1597")
    value = unit(ir, "复合袋")
    task = property_task(ir, subject, value)
    assert scope_contains(task.scope, ir.anchor(value.evidence_id), ir)
    assert not scope_contains(task.scope, ir.anchor(unit(ir, "纸袋").evidence_id), ir)
    spans = [quote(unit(ir, "包装方式")), quote(shared), quote(value)]
    (candidate,), _ = execute_property(
        ir,
        subject,
        task,
        {"value": quote(value)},
        spans,
        method="table_record",
    )
    assert candidate.validation_status == "passed", candidate.validation_issues
    assert candidate.bindings[0].record_mapping == {"table_path": value.table_path, "row_index": 2}
    # Explicitly specifying the other covered row cannot move this value there.
    (wrong,), _ = execute_property(
        ir,
        subject,
        task,
        {"value": quote(value)},
        spans,
        method="table_record",
        record_mapping={"table_path": value.table_path, "row_index": 1},
    )
    assert "record_mapping_mismatch" in {i.code for i in wrong.validation_issues}
    tables = TableRecords(ir)
    assert tables.infer_mapping([ir.anchor(shared.evidence_id)]) == {}
    assert not tables.valid_mapping(
        candidate.bindings[0].record_mapping,
        [ir.anchor(unit(ir, "双层袋").evidence_id), ir.anchor(value.evidence_id)],
    )


def test_second_row_entity_context_includes_shared_subject_only_as_metadata(tmp_path):
    ir = table_document(tmp_path)
    target = unit(ir, "犬14天")
    task = ExtractionTask(
        task_id="row2",
        task_kind="entity",
        target_class_iris=["urn:Study"],
        target_evidence_ids=[target.evidence_id],
        budget=TaskBudget(max_input_tokens=100000),
    )
    context = build_context(task, ir, {}, TestTokenizer(), model="test")
    assert any("HRS-1597" in ir.resolve(a) for a in context.allowed_binding_regions)
    assert [ir.resolve(a) for a in context.allowed_fact_regions] == ["犬14天"]
    assert not any(ir.resolve(a) == "大鼠14天" for a in context.allowed_binding_regions)


def test_multilevel_and_horizontal_headers_explain_columns_without_claiming_rows(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "对象"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "批量"
    for col, text in enumerate(["名称", "下限(kg)", "上限(kg)"]):
        table.cell(1, col).text = text
    for col, text in enumerate(["计划A", "3.8", "6.6"]):
        table.cell(2, col).text = text
    ir = analyze(doc, tmp_path)
    subject, value = entity(ir, unit(ir, "计划A")), unit(ir, "6.6")
    spans = [quote(unit(ir, text)) for text in ("批量", "上限(kg)", "计划A", "6.6")]
    (candidate,), _ = execute_property(
        ir,
        subject,
        property_task(ir, subject, value, datatype="decimal", canonical_unit="kg"),
        {"value": quote(value)},
        spans,
        method="table_record",
        source_unit=quote(unit(ir, "上限(kg)"), "kg"),
    )
    assert candidate.validation_status == "passed"
    assert candidate.literal.canonical_unit == "kg"
    assert candidate.literal.conversion_record["evidence"][0]["row_index"] == 1


@pytest.mark.parametrize("name,start", [("计划A", 5), ("计划B", 14)])
def test_repeated_value_is_located_inside_unique_assertion(tmp_path, name, start):
    doc = Document()
    doc.add_paragraph("计划A批量5kg，计划B批量5kg。")
    ir = analyze(doc, tmp_path)
    target = ir.evidence_units[0]
    subject = entity(ir, target, name)
    assertion = quote(target, name + "批量5kg")
    (candidate,), stages = execute_property(
        ir,
        subject,
        property_task(ir, subject, target, datatype="decimal"),
        {"value": quote(target, "5kg"), "assertion_spans": [assertion]},
        [assertion],
    )
    assert candidate.validation_status == "passed" and "verify_binding" in stages
    assert candidate.provenance[0].anchors[0].span_start == start
    with pytest.raises(PartialTaskFailure, match="ambiguous_source_quote"):
        execute_property(
            ir,
            subject,
            property_task(ir, subject, target),
            {"value": quote(target, "5kg")},
            [assertion],
        )
    # A context cannot authorize a value in an excluded part of the source.
    task = property_task(ir, subject, target)
    task.target_ranges = [EvidenceRange(evidence_id=target.evidence_id, start=0, end=8)]
    if name == "计划B":
        with pytest.raises(PartialTaskFailure, match="source_quote_outside_scope"):
            execute_property(
                ir,
                subject,
                task,
                {"value": quote(target, "5kg", context=name + "批量5kg")},
                [assertion],
            )


def relation_document(tmp_path):
    doc = Document()
    doc.add_heading("临床备样生产计划", 0)
    doc.add_heading("报告内容", 1)
    doc.add_paragraph("本报告包含临床备样生产计划。该计划生产用途为一期临床。")
    doc.add_paragraph("该计划批量下限为3.8kg，上限为6.6kg。")
    doc.add_paragraph("计划B用于二期临床。该计划批量为20kg。")
    ir = analyze(doc, tmp_path)
    subject = entity(ir, ir.evidence_units[0], cls="urn:Plan")
    seed_unit = ir.evidence_units[2]
    seed = ir.anchor(seed_unit.evidence_id, 0, seed_unit.text.index("。"))
    relation = Candidate(
        candidate_id="relation",
        kind="relationship",
        subject=candidate_ref(subject),
        object=candidate_ref(subject),
        predicate_iri="urn:hasPlan",
        validation_status="passed",
        provenance=[DocumentProvenance(anchors=[seed], excerpts=[ir.resolve(seed)])],
        bindings=[
            BindingEvidence(
                method="explicit_assertion",
                subject_candidate_id="subject",
                object_candidate_id="subject",
                predicate_iri="urn:hasPlan",
                anchors=[seed],
            )
        ],
    )
    return ir, subject, relation


def test_relation_continuations_need_independent_reference_verification(tmp_path):
    ir, subject, relation = relation_document(tmp_path)
    competitor = entity(ir, ir.evidence_units[4], "计划B", "plan-b", "urn:Plan")
    scope = build_scope(ir, subject, [competitor], bound_relationships=[relation])
    purpose = ir.evidence_units[2]
    assert scope_contains(
        scope,
        ir.anchor(purpose.evidence_id, purpose.text.index("一期"), purpose.text.index("一期") + 4),
        ir,
    )
    target = ir.evidence_units[3]
    assert scope_contains(scope, ir.anchor(target.evidence_id, 0, len(target.text) - 1), ir)
    assert not scope_contains(scope, ir.anchor(ir.evidence_units[4].evidence_id), ir)
    task = property_task(ir, subject, target, datatype="decimal", scope=scope)
    spans = [quote(purpose, "本报告包含临床备样生产计划"), quote(target)]
    (candidate,), stages = execute_property(
        ir,
        subject,
        task,
        {"value": quote(target, "3.8kg")},
        spans,
    )
    assert candidate.validation_status == "passed", candidate.validation_issues
    assert "verify_reference" in stages
    assert candidate.bindings[-1].method == "identity_reference"
    with pytest.raises(PartialTaskFailure, match="reference_not_supported"):
        execute_property(
            ir, subject, task, {"value": quote(target, "3.8kg")}, spans, reference_supported=False
        )


def test_gap_expansion_starts_at_relation_destination_and_marks_reference_requirement(tmp_path):
    ir, subject, relation = relation_document(tmp_path)
    scope = expanded_gap_scope(ir, subject, [], 1, bound_relationships=[relation])
    assert any(r.evidence_id == ir.evidence_units[3].evidence_id for r in scope.ranges)
    assert scope.reference_ranges
    assert len(scope.expansion_history) <= 1


@pytest.mark.parametrize(
    "raw,unit_name,expected",
    [
        ("0.2g", "mg", "200"),
        ("12mg/kg", "mg/kg", "12"),
        ("0.2g/day", "mg/day", "200"),
        ("0.2mg/day", "mg/day", "0.2"),
        ("25℃", "°C", "25"),
        ("25℃", "K", "298.15"),
        ("298.15K", "°C", "25"),
        ("0.1mg/kg/day", "ug/kg/day", "100"),
        ("2μg/m³", "ug/m3", "2"),
    ],
)
def test_compound_units_and_affine_temperatures_are_exact(raw, unit_name, expected):
    value = normalize_literal(raw, datatype="decimal", target_unit=unit_name)
    assert value.normalized_value == expected and value.raw_value == raw
    assert value.canonical_unit == unit_name


def test_chinese_dates_legends_and_unknowns_preserve_meaning():
    assert normalize_literal("2026年6月1日", datatype="date").normalized_value == "2026-06-01"
    for raw in ("2026年2月29日", "2026年6月", "2026年6月32日"):
        with pytest.raises(ValueError, match="invalid date"):
            normalize_literal(raw, datatype="date")
    for symbol, meaning in (("√", True), ("×", False)):
        assert (
            normalize_literal(
                symbol, datatype="boolean", boolean_legend="√：是；×：否；—：研究数据不足"
            ).normalized_value
            is meaning
        )
        with pytest.raises(ValueError, match="boolean"):
            normalize_literal(symbol, datatype="boolean")
    with pytest.raises(ValueError, match="unknown_boolean"):
        normalize_literal("—", datatype="boolean", boolean_legend="—：研究数据不足")
    with pytest.raises(ValueError):
        normalize_literal("XX", datatype="decimal")
    with pytest.raises(ValueError, match="unit_conversion_not_exact"):
        normalize_literal("1s", datatype="decimal", target_unit="min")
    with pytest.raises(ValueError, match="unit_missing_or_incompatible"):
        normalize_literal("5", datatype="decimal", target_unit="mg")


def test_real_document_boolean_legend_is_replayed_from_table_note(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for row, texts in zip(table.rows, [("产品", "毒性"), ("产品A", "√")], strict=True):
        for cell, text in zip(row.cells, texts, strict=True):
            cell.text = text
    legend = "备注：“√”代表有相应毒性；“×”代表没有相应的毒性 ；“—”代表相关的研究数据不充分。"
    doc.add_paragraph(legend)
    ir = analyze(doc, tmp_path)
    subject, value = entity(ir, unit(ir, "产品A")), unit(ir, "√")
    (candidate,), _ = execute_property(
        ir,
        subject,
        property_task(ir, subject, value, datatype="boolean"),
        {"value": quote(value)},
        [quote(unit(ir, "毒性")), quote(value)],
        method="table_record",
        boolean_legend=quote(unit(ir, legend)),
    )
    assert candidate.validation_status == "passed"
    assert candidate.literal.normalized_value is True
    assert candidate.literal.conversion_record["boolean_legend"] == legend
    assert candidate.literal.conversion_record["evidence"]
    assert (
        normalize_literal("×", datatype="boolean", boolean_legend=legend).normalized_value is False
    )
    with pytest.raises(ValueError, match="unknown_boolean"):
        normalize_literal("—", datatype="boolean", boolean_legend=legend)


def test_subject_in_one_study_row_does_not_inherit_other_study_values(tmp_path):
    ir = table_document(tmp_path)
    rat = entity(ir, unit(ir, "大鼠14天"), cls="urn:Study", identity="rat")
    dog = entity(ir, unit(ir, "犬14天"), cls="urn:Study", identity="dog")
    assert instance_iri(rat) != instance_iri(dog)
    rat_scope, dog_scope = build_scope(ir, rat, [dog]), build_scope(ir, dog, [rat])
    assert not scope_contains(rat_scope, ir.anchor(unit(ir, "复合袋").evidence_id), ir)
    assert not scope_contains(dog_scope, ir.anchor(unit(ir, "双层袋").evidence_id), ir)


def test_nested_table_record_cannot_borrow_parent_headers(tmp_path):
    doc = Document()
    outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).text = "外部主体"
    outer.cell(0, 1).text = "外部值"
    outer.cell(1, 0).text = "外部对象"
    nested = outer.cell(1, 1).add_table(rows=2, cols=2)
    for row, texts in zip(nested.rows, [("主体", "值"), ("内部对象", "5")], strict=True):
        for cell, text in zip(row.cells, texts, strict=True):
            cell.text = text
    ir = analyze(doc, tmp_path)
    source, value = unit(ir, "内部对象"), unit(ir, "5")
    records = TableRecords(ir)
    mapping = records.infer_mapping([ir.anchor(source.evidence_id), ir.anchor(value.evidence_id)])
    assert mapping["table_path"] == source.table_path and len(source.table_path) == 3
    assert records.valid_mapping(
        mapping, [ir.anchor(unit(ir, "值").evidence_id), ir.anchor(value.evidence_id)]
    )
    assert not records.valid_mapping(
        mapping, [ir.anchor(unit(ir, "外部值").evidence_id), ir.anchor(value.evidence_id)]
    )


def test_source_unit_in_another_column_cannot_relabel_value(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for row, texts in zip(
        table.rows, [("对象", "剂量(mg)", "质量(g)"), ("A", "5", "0.2")], strict=True
    ):
        for cell, text in zip(row.cells, texts, strict=True):
            cell.text = text
    ir = analyze(doc, tmp_path)
    subject, value = entity(ir, unit(ir, "A")), unit(ir, "5")
    task = property_task(ir, subject, value, datatype="decimal", canonical_unit="mg")
    # Permit both cells as targets to isolate the unit/column contract.
    task.target_evidence_ids.append(unit(ir, "0.2").evidence_id)
    with pytest.raises(PartialTaskFailure, match="unit_record_mismatch"):
        execute_property(
            ir,
            subject,
            task,
            {"value": quote(value)},
            [quote(value), quote(unit(ir, "质量(g)"))],
            source_unit=quote(unit(ir, "质量(g)"), "g"),
        )


@pytest.mark.parametrize(
    "finish,content,error",
    [
        ("length", '{"x":1}', "model_output_truncated"),
        ("stop", "{broken", "model_parse_error"),
        ("stop", "", "model_empty_response"),
    ],
)
def test_model_failure_categories_preserve_request_outcome(finish, content, error):
    client = SimpleNamespace(
        base_url="http://model.test/v1",
        chat=SimpleNamespace(
            completions=SimpleNamespace(
                create=AsyncMock(side_effect=lambda **kwargs: SimpleNamespace(
                    choices=[
                        SimpleNamespace(
                            finish_reason=finish,
                            message=SimpleNamespace(content=content),
                        )
                    ]
                )),
            )
        )
    )
    with pytest.raises(StructuredModelError, match=error):
        chat_with_schema(
            client, system="test", user="test", schema={}, max_attempts=1, raise_on_error=True
        )
    assert chat_with_schema(client, system="test", user="test", schema={}, max_attempts=1) is None


@pytest.mark.parametrize(
    "error,code", [(TimeoutError(), "model_timeout"), (RuntimeError(), "model_request_failed")]
)
def test_transport_errors_are_not_reported_as_model_unavailable(error, code):
    async def create(**kwargs):
        raise error

    client = SimpleNamespace(base_url="http://model.test/v1",
                             chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    with pytest.raises(StructuredModelError, match=code):
        chat_with_schema(
            client, system="test", user="test", schema={}, max_attempts=1, raise_on_error=True
        )


def test_stage_statistics_survive_checkpoint_replay(tmp_path):
    doc = Document()
    doc.add_paragraph("未记载相关事实")
    ir = analyze(doc, tmp_path)
    calls = []

    def model(*args):
        calls.append(1)
        return {"entities": []}

    runner = GenericExtractionRunner(
        {"urn:Product": {}},
        TestTokenizer(),
        model,
        model_identity="counts",
        budget=TaskBudget(max_input_tokens=30000),
    )
    run = runner.run(ir)
    assert run.stage_statistics == {"entity_recall": {"not_found": 1}}
    replay = runner.run(ir, checkpoint=run.checkpoint)
    assert replay.stage_statistics == run.stage_statistics and len(calls) == 1


@pytest.mark.parametrize("same_key,verified", [(True, True), (False, True), (True, False)])
def test_cross_mention_identity_requires_declared_identifier_and_independent_verdict(
    tmp_path,
    same_key,
    verified,
):
    doc = Document()
    doc.add_paragraph("设备：过滤器；唯一编号：EQ-1")
    doc.add_paragraph("设备：过滤器；唯一编号：" + ("EQ-1" if same_key else "EQ-2"))
    ir = analyze(doc, tmp_path)
    schema = {
        "urn:Equipment": {
            "properties": [
                {"iri": "urn:id", "datatype": "string", "identity_key": True},
            ]
        }
    }

    def model(system, user, response_schema, budget):
        request = json.loads(user)
        if request["stage"] == "verify_entity_types":
            return {
                "decisions": [
                    {
                        "candidate_id": c["candidate_id"],
                        "supported": True,
                        "identity_supported": verified,
                        "reason": "编号及具体实例身份独立复核",
                    }
                    for c in request["candidate"]["proposed_entities"]
                ]
            }
        context = json.loads(request["context"])
        target = ir.unit(context["fragments"][0]["anchor"]["evidence_id"])
        return {
            "entities": [
                {
                    "class_iri": "urn:Equipment",
                    "mention": quote(target, "过滤器"),
                    "identifier": {
                        "predicate_iri": "urn:id",
                        "value": quote(target, target.text[-4:]),
                    },
                }
            ]
        }

    runner = GenericExtractionRunner(schema, TestTokenizer(), model, model_identity="identity")
    entities = []
    for i, target in enumerate(ir.evidence_units):
        task = ExtractionTask(
            task_id=f"entity-{i}",
            task_kind="entity",
            target_class_iris=["urn:Equipment"],
            target_evidence_ids=[target.evidence_id],
            budget=TaskBudget(max_input_tokens=30000),
        )
        entities.extend(runner.execute_task(task, ir, {}))
    a, b = entities
    assert a.candidate_id != b.candidate_id  # Every physical mention remains auditable.
    assert (instance_iri(a) == instance_iri(b)) is (same_key and verified)
    assert scope_contains(build_scope(ir, a, [b]), b.provenance[0].anchors[0], ir) is (
        same_key and verified
    )
    if same_key and verified:
        task = property_task(ir, b, ir.evidence_units[1])
        (value,), _ = execute_property(
            ir,
            b,
            task,
            {"value": quote(ir.evidence_units[1], "EQ-1")},
            [quote(ir.evidence_units[1])],
        )
        relation = Candidate(
            candidate_id="link",
            kind="relationship",
            subject=candidate_ref(a),
            object=candidate_ref(a),
            predicate_iri="urn:has",
            validation_status="passed",
            provenance=a.provenance,
            bindings=[
                BindingEvidence(
                    method="explicit_assertion",
                    subject_candidate_id=a.candidate_id,
                    object_candidate_id=a.candidate_id,
                    predicate_iri="urn:has",
                    anchors=a.provenance[0].anchors,
                )
            ],
        )
        assert preview_relationships([a, b, value, relation], schema)[0]["object_data_properties"]
