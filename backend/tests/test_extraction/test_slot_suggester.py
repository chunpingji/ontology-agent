"""019: IR-owned skeletons and ontology-menu-only semantic enrichment."""

from __future__ import annotations

import json
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from app.services.extraction.document_annotator import parse_word_to_tiptap
from app.services.extraction.slot_suggester import suggest_slots, tiptap_to_text
from tests.fixtures.ontology import (
    BIOLOGIC,
    DRUG_NS,
    DRUG_PRODUCT,
    MANUFACTURED_BY,
    MANUFACTURER,
    build_drug_ontology,
)


@pytest.fixture
def sample_content(tmp_path):
    doc = Document()
    doc.add_heading("评估对象", 1)
    doc.add_paragraph("药品名称：XX注射液")
    path = tmp_path / "sample.docx"
    doc.save(path)
    return parse_word_to_tiptap(path)


def _make_client(response):
    client = MagicMock()
    client.base_url = "http://model.test/v1"
    client.chat.completions.create = AsyncMock()

    async def create(**kwargs):
        request = json.loads(kwargs["messages"][1]["content"])
        value = {"section_id": request["section"]["id"], "document_summary": "评估摘要", **response}
        reply = MagicMock()
        reply.choices[0].message.content = json.dumps(value, ensure_ascii=False)
        return reply

    client.chat.completions.create.side_effect = create
    return client


class TestBasic:
    def test_offline_structure_does_not_depend_on_semantic_success(self, sample_content):
        offline = suggest_slots(None, "", content_json=sample_content)
        client = MagicMock()
        client.base_url = "http://model.test/v1"
        client.chat.completions.create = AsyncMock()
        client.chat.completions.create.side_effect = RuntimeError("model unavailable")
        failed = suggest_slots(client, "", content_json=sample_content)
        assert offline["sections"] == failed["sections"]
        assert offline["completion"] == failed["completion"] == "incomplete"
        assert offline["degraded"] is False
        assert failed["diagnostics"]

    def test_unknown_or_invented_structure_is_rejected(self, sample_content):
        offline = suggest_slots(None, "", content_json=sample_content)
        client = _make_client({"fields": [{"id": "invented", "semantic_label": "bad"}]})
        result = suggest_slots(client, "", content_json=sample_content)
        assert result["sections"] == offline["sections"]
        assert result["completion"] == "incomplete"
        assert result["coverage"] == []

    def test_model_can_enrich_only_existing_candidate_id(self, sample_content):
        original = suggest_slots(None, "", content_json=sample_content)
        candidate = original["sections"][0]["groups"][0]["candidates"][0]
        client = _make_client({"fields": [{"id": candidate["id"], "semantic_label": "药品"}]})
        result = suggest_slots(client, "", content_json=sample_content)
        actual = result["sections"][0]["groups"][0]["candidates"][0]
        assert actual["semantic_label"] == "药品"
        assert actual["origin"] == candidate["origin"]
        assert actual["label"] == candidate["label"]
        assert result["completion"] == "complete"

    def test_legacy_text_without_analysis_requires_reparse(self):
        client = MagicMock()
        result = suggest_slots(client, "原文")
        assert result["sections"] == []
        assert result["completion"] == "incomplete"
        assert "analysis_required" in result["diagnostics"][0]
        client.chat.completions.create.assert_not_called()

    def test_oversized_section_is_explicitly_incomplete_not_truncated(self, tmp_path):
        doc = Document()
        doc.add_heading("长文档", 1)
        doc.add_paragraph("字段：" + "值" * 13000)
        path = tmp_path / "long.docx"
        doc.save(path)
        content = parse_word_to_tiptap(path)
        client = MagicMock()
        result = suggest_slots(client, "", content_json=content)
        assert result["sections"]
        assert result["completion"] == "incomplete"
        assert "budget_exceeded" in result["diagnostics"][0]
        client.chat.completions.create.assert_not_called()


class TestOntologyCoverage:
    def test_graph_sourced_section_emits_required_type_coverage(self, sample_content):
        engine = build_drug_ontology()
        client = _make_client(
            {
                "coverage": [
                    {"predicate_iri": MANUFACTURED_BY, "range_class_iri": MANUFACTURER},
                    {"predicate_iri": MANUFACTURED_BY, "range_class_iri": MANUFACTURER},
                    {
                        "predicate_iri": MANUFACTURED_BY,
                        "range_class_iri": DRUG_NS + "individual_001",
                    },
                ]
            }
        )
        result = suggest_slots(
            client,
            "",
            ontology_engine=engine,
            doc_class_iri=DRUG_PRODUCT,
            content_json=sample_content,
        )
        assert len(result["coverage"]) == 1
        cov = result["coverage"][0]
        assert cov["kind"] == "ontology_relation"
        assert cov["doc_class_iri"] == DRUG_PRODUCT
        assert cov["predicate_iri"] == MANUFACTURED_BY
        assert cov["range_class_iri"] == MANUFACTURER
        assert cov["required"] is True

    def test_unexpected_model_fields_do_not_leak_into_contract(self, sample_content):
        client = _make_client({"unresolved_candidates": [{"proposed_label": "虚构字段"}]})
        result = suggest_slots(client, "", content_json=sample_content)
        assert result["coverage"] == []
        assert result["completion"] == "incomplete"
        assert "unresolved_candidates" not in result
        assert result["sections"]

    def test_doc_class_iri_optional_request_still_has_structure(self, sample_content):
        from app.schemas.extraction import SuggestSlotsRequest

        request = SuggestSlotsRequest(document_text="doc")
        assert request.doc_class_iri is None
        result = suggest_slots(_make_client({}), "", content_json=sample_content)
        assert result["coverage"] == []
        assert result["sections"]
        assert result["document_summary"] == "评估摘要"


class TestSupplementalCmcEdges:
    """016 schema-driven: broad-domain object props now have ``rdfs:domain`` declared
    in TTL, so ``get_relation_schema`` BFS discovers them directly. The AI coverage
    menu and extraction pipeline share the same ``get_relation_schema`` single source
    of truth.
    """

    def test_cmc_schema_edges_in_menu(self):
        from app.services.extraction.slot_suggester import (
            _build_ontology_context,
            _extract_coverage,
        )
        from tests.fixtures.report_ontology_constants import (
            CMC_REPORT_IRI,
            DEGRADATION_PATHWAY_IRI,
            EQUIPMENT_IRI,
            HAS_DEGRADATION_PATHWAY_IRI,
            HAS_STORAGE_CONDITION_IRI,
            STORAGE_CONDITION_IRI,
            USES_EQUIPMENT_IRI,
        )

        # CMCReport 现有 domain 声明 → get_relation_schema BFS 自然覆盖 3 条边。
        engine = build_drug_ontology()
        schema_edges, prompt = _build_ontology_context(engine, CMC_REPORT_IRI)

        edge_keys = {(e["predicate_iri"], e["range_class_iri"]) for e in schema_edges}
        assert (USES_EQUIPMENT_IRI, EQUIPMENT_IRI) in edge_keys
        assert (HAS_STORAGE_CONDITION_IRI, STORAGE_CONDITION_IRI) in edge_keys
        assert (HAS_DEGRADATION_PATHWAY_IRI, DEGRADATION_PATHWAY_IRI) in edge_keys

        for label in ("使用设备", "存放条件", "含降解途径"):
            assert label in prompt

        r2 = {
            "coverage": [
                {"predicate_iri": USES_EQUIPMENT_IRI, "range_class_iri": EQUIPMENT_IRI},
            ]
        }
        cov = _extract_coverage(r2, schema_edges, CMC_REPORT_IRI)
        assert len(cov) == 1
        assert cov[0]["predicate_iri"] == USES_EQUIPMENT_IRI
        assert cov[0]["doc_class_iri"] == CMC_REPORT_IRI
        assert cov[0]["required"] is True

    def test_non_cmc_doc_class_gets_own_schema_edges(self):
        """Other doc types get their own schema edges (no cross-type leakage)."""
        from app.services.extraction.slot_suggester import _supplemented_schema_edges

        engine = build_drug_ontology()
        dp_edges = _supplemented_schema_edges(engine, DRUG_PRODUCT)
        assert len(dp_edges) >= 1
        assert all(e["domain_class_iri"] == DRUG_PRODUCT for e in dp_edges if e["hop"] == 1)


class TestCoverageCapable:
    """Bugfix「仅启用已建模类型」(S9/S10): a document type is *capable* iff the ontology
    models ≥1 hop-1 coverage edge for it. ``coverage_capable`` (behind the
    ``/coverage-doc-classes`` probe that gates the authoring dropdown) and
    ``_build_ontology_context`` (the AI menu) MUST agree — both consume the one
    ``_supplemented_schema_edges`` source of truth, so "UI enables it" ⇔ "AI can emit
    coverage". Per contracts/suggest-slots-api.md.
    """

    def test_modeled_type_is_capable(self):
        """DrugProduct has a hop-1 object property (manufacturedBy) → capable."""
        from app.services.extraction.slot_suggester import coverage_capable

        engine = build_drug_ontology()
        assert coverage_capable(engine, DRUG_PRODUCT) is True

    def test_cmc_capable_via_domain_declarations(self):
        """CMCReport has 3 object properties with domain declarations → BFS finds
        them at hop-1, so CMCReport is capable."""
        from app.services.extraction.slot_suggester import coverage_capable
        from tests.fixtures.report_ontology_constants import CMC_REPORT_IRI

        engine = build_drug_ontology()
        assert coverage_capable(engine, CMC_REPORT_IRI) is True

    def test_unmodeled_type_not_capable(self):
        """Leaf/document-classification types with no object property → NOT capable
        (Manufacturer is a leaf; BiologicalDrugProduct is a subclass with no own
        object props — domain-literal matching, mirroring real doc-class subclasses)."""
        from app.services.extraction.slot_suggester import coverage_capable

        engine = build_drug_ontology()
        assert coverage_capable(engine, MANUFACTURER) is False
        assert coverage_capable(engine, BIOLOGIC) is False

    def test_missing_engine_or_iri_not_capable(self):
        """FR-012 graceful degradation: no engine / no iri → not capable (no raise)."""
        from app.services.extraction.slot_suggester import coverage_capable

        engine = build_drug_ontology()
        assert coverage_capable(None, DRUG_PRODUCT) is False
        assert coverage_capable(engine, None) is False

    def test_build_ontology_context_contract_unchanged(self):
        """S9 regression: ``(edges, prompt)`` contract intact — modeled type yields a
        non-empty menu prompt; no-hop1 / no-engine degrades to ``([], "")``."""
        from app.services.extraction.slot_suggester import _build_ontology_context

        engine = build_drug_ontology()
        edges, prompt = _build_ontology_context(engine, DRUG_PRODUCT)
        assert edges and prompt
        assert "生产者" in prompt  # hop-1 predicate label rendered into the menu

        # no hop-1 edge → prompt empty (dropdown/AI both see "unmodeled")
        assert _build_ontology_context(engine, MANUFACTURER) == ([], "")
        # engine/iri absent → fully degraded
        assert _build_ontology_context(None, DRUG_PRODUCT) == ([], "")
        assert _build_ontology_context(engine, None) == ([], "")

    def test_capable_iff_menu_nonempty(self):
        """S9 parity: capability ⇔ a non-empty AI menu, for every fixture type."""
        from app.services.extraction.slot_suggester import (
            _build_ontology_context,
            coverage_capable,
        )
        from tests.fixtures.report_ontology_constants import CMC_REPORT_IRI

        engine = build_drug_ontology()
        for iri in (DRUG_PRODUCT, CMC_REPORT_IRI, MANUFACTURER, BIOLOGIC):
            _, prompt = _build_ontology_context(engine, iri)
            assert coverage_capable(engine, iri) is bool(prompt), iri


# ── 013: tiptap → LLM text (server-side, structure-faithful) ───────────────


class TestTiptapToText:
    def test_empty_returns_empty_string(self):
        assert tiptap_to_text(None) == ""
        assert tiptap_to_text({}) == ""

    def test_heading_prefixed_with_hash(self):
        doc = {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 2},
                    "content": [{"type": "text", "text": "评估对象"}],
                },
            ],
        }
        assert tiptap_to_text(doc) == "## 评估对象"

    def test_heading_level_clamped_to_six(self):
        doc = {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 9},
                    "content": [{"type": "text", "text": "深层标题"}],
                },
            ],
        }
        assert tiptap_to_text(doc) == "###### 深层标题"

    def test_paragraph_and_list_item(self):
        doc = {
            "type": "doc",
            "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "普通段落"}]},
                {
                    "type": "bulletList",
                    "content": [
                        {
                            "type": "listItem",
                            "content": [
                                {
                                    "type": "paragraph",
                                    "content": [{"type": "text", "text": "列表项"}],
                                },
                            ],
                        },
                    ],
                },
            ],
        }
        # listItem 走 _node_text 递归取全部文本，不再重复 emit 内部段落。
        assert tiptap_to_text(doc) == "普通段落\n- 列表项"

    def test_table_rows_joined_with_pipe(self):
        doc = {
            "type": "doc",
            "content": [
                {
                    "type": "table",
                    "content": [
                        {
                            "type": "tableRow",
                            "content": [
                                {
                                    "type": "tableCell",
                                    "content": [
                                        {
                                            "type": "paragraph",
                                            "content": [{"type": "text", "text": "设备编号"}],
                                        }
                                    ],
                                },
                                {
                                    "type": "tableCell",
                                    "content": [
                                        {
                                            "type": "paragraph",
                                            "content": [{"type": "text", "text": "CT64201"}],
                                        }
                                    ],
                                },
                            ],
                        },
                    ],
                },
            ],
        }
        assert tiptap_to_text(doc) == "[表格]\n设备编号 | CT64201"

    def test_round_tripped_block_text_is_substring(self):
        # 忠于原文：LLM 引用的 evidence_span 必是序列化文本的子串（联动前提）。
        doc = {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 1},
                    "content": [{"type": "text", "text": "评估对象"}],
                },
                {
                    "type": "paragraph",
                    "content": [{"type": "text", "text": "本品为XX注射液。"}],
                },
            ],
        }
        text = tiptap_to_text(doc)
        assert "评估对象" in text
        assert "本品为XX注射液。" in text

    def test_serialization_never_silently_truncates_source(self):
        long = "字" * 13000
        doc = {
            "type": "doc",
            "content": [{"type": "paragraph", "content": [{"type": "text", "text": long}]}],
        }
        out = tiptap_to_text(doc)
        assert out == long
