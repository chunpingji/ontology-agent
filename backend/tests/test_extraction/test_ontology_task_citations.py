"""Production adapter citations and independently proved cross-field ownership."""

import json

import pytest
from docx import Document

from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context import ContextFragment, assemble_context
from app.services.extraction.ontology_guided.contracts import DocumentContext, SlotSpec, SubjectRef
from app.services.extraction.ontology_guided.model_adapter import (
    LocalModelRecognitionAdapter,
    RecognitionModelFailure,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.source_citations import resolve_fragment_quote
from app.services.extraction.word_analysis import analyze_word_core

from .test_independent_semantic_verification import _inputs
from .test_semantic_graph_closure import PRODUCT, ROOT, _respond


def test_proof_ids_replay_complete_source_without_replacing_entity_name(tmp_path, monkeypatch):
    args = _inputs(tmp_path)

    def respond(_client, *, user, schema, **_kwargs):
        request = json.loads(user)
        response = _respond(request)
        assert all(fragment['evidence_id'].startswith('E') for fragment in request['fragments'])
        if request['stage'] == 'verification':
            proof = schema['$defs']['BindingProof']
            assert proof['required'] == ['evidence_id']
            assert 'text' not in proof['properties']
            for item in response['verifications']:
                for field in ('type_support', 'predicate_support'):
                    item[field] = [{'evidence_id': quote['evidence_id']} for quote in item[field]]
        return response

    monkeypatch.setattr(model_adapter, 'chat_with_schema', respond)
    result = LocalModelRecognitionAdapter(object(), model_identity='atomic').inspect(*args)
    assert result.edges[0].policy_eligible
    assert result.nodes[0].label == '产品甲片'
    name_anchor = result.nodes[0].evidence_refs[0]
    assert name_anchor.span_end - name_anchor.span_start == 4
    proof = next(item for item in result.decision_payloads if item['check_kind'] == 'type')
    assert proof['support_refs'][0]['span_end'] == len('本报告描述产品甲片。')
    assert result.model_calls == 2


@pytest.mark.parametrize('failure', [
    'background_endpoint', 'ellipsis', 'missing_name', 'coordinates',
])
def test_bad_source_never_reaches_independent_verification(tmp_path, monkeypatch, failure):
    args = _inputs(tmp_path)
    calls = []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request['stage'])
        response = _respond(request)
        quote = response['proposals'][0]['object_quote']
        if failure == 'background_endpoint':
            fragment = next(f for f in request['fragments'] if not f['fact_eligible'])
            quote.update(evidence_id=fragment['evidence_id'], text=fragment['text'])
        elif failure == 'ellipsis':
            quote['text'] = '产品…片'
        elif failure == 'missing_name':
            quote.pop('text')
        else:
            quote.update(start=0, end=4)
        return response

    monkeypatch.setattr(model_adapter, 'chat_with_schema', respond)
    with pytest.raises(RecognitionModelFailure, match='discovery_citation_invalid') as failure:
        LocalModelRecognitionAdapter(object(), model_identity='bad-citation').inspect(*args)
    assert failure.value.model_calls == 1
    assert calls == ['discovery']


def test_repeated_substrings_and_disjoint_permissions_are_not_guessed(tmp_path):
    document = Document()
    document.add_paragraph('设备甲 设备甲')
    path = tmp_path / 'repeated.docx'
    document.save(path)
    ir = analyze_word_core(path).ir
    unit = ir.evidence_units[0]
    full = ContextFragment(anchor=ir.anchor(unit.evidence_id), text=unit.text,
                           purpose='target', fact_eligible=True)
    with pytest.raises(ValueError, match='ambiguous_source_quote'):
        resolve_fragment_quote(unit.evidence_id, '设备甲', [full], fact_required=True)
    pieces = [ContextFragment(anchor=ir.anchor(unit.evidence_id, start, end),
                              text=unit.text[start:end], purpose='target', fact_eligible=True)
              for start, end in [(0, 3), (4, 7)]]
    with pytest.raises(ValueError, match='ambiguous_allowed_source_fragments'):
        resolve_fragment_quote(unit.evidence_id, None, pieces)
    anchor, text = resolve_fragment_quote(unit.evidence_id, '设备甲', [pieces[0], pieces[0]])
    assert ir.resolve(anchor) == text == '设备甲'
    assert anchor.span_start == 0 and anchor.span_end == 3


@pytest.mark.parametrize('support', ['both', 'original_only', 'field_only'])
def test_form_property_requires_original_owner_and_current_role_field(
    tmp_path, monkeypatch, support,
):
    task, root_context, _, _ = _inputs(tmp_path)
    document = Document()
    document.add_paragraph('X12项目是一种原料药。')
    document.add_heading('产品信息', 1)
    document.add_paragraph('分子量：123.45')
    document.add_paragraph('项目名称：X12')
    document.add_paragraph('CAS：N/A')
    path = tmp_path / 'form.docx'
    document.save(path)
    ir = analyze_word_core(path).ir
    index = RecordIndex(ir)
    original = ir.evidence_units[0]
    record = next(r for r in index.records if r.text.startswith('分子量：'))
    subject = SubjectRef(entity_id='product', revision=1, class_iri=PRODUCT)
    predicate = SlotSpec(iri='urn:weight', label='分子量')
    task = task.model_copy(update={'subject': subject, 'predicate_iri': predicate.iri,
                                  'predicate_kind': 'property', 'record_id': record.record_id})
    target = root_context.target.model_copy(update={
        'subject_ref': subject, 'predicate_iri': predicate.iri,
        'document_context': DocumentContext(document_hash=ir.document_hash,
            document_class_iri=ROOT, root_ref=root_context.target.document_context.root_ref),
    })
    context = assemble_context(target, record.record_id, index, subject_label='X12项目',
                               subject_evidence_refs=[ir.anchor(original.evidence_id, 0, 5)])
    assert [f.text for f in context.fragments if f.fact_eligible] == ['分子量：123.45']
    assert [ir.resolve(a) for a in context.owner_field_refs] == ['项目名称：X12']
    assert any(f.text == 'CAS：N/A' and not f.fact_eligible for f in context.fragments)

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        fragments = {f['text']: f['evidence_id'] for f in request['fragments']}
        value_id = fragments['分子量：123.45']
        if request['stage'] == 'discovery':
            return {'proposals': [{'kind': 'property', 'bridge_kind': 'owned_field_group',
                                  'value_quote': {'evidence_id': value_id, 'text': '123.45'}}]}
        claim = request['candidates'][0]
        sources = []
        if support != 'field_only':
            sources.append({'evidence_id': fragments[original.text]})
        if support != 'original_only':
            sources.append({'evidence_id': fragments['项目名称：X12']})
        return {'verifications': [{
            'candidate_id': claim['candidate_id'], 'target_id': claim['target_id'],
            **{field + '_verdict': 'supported' for field in (
                'type', 'role', 'subject_binding', 'predicate', 'applicability',
                'counterevidence', 'bridge')},
            'type_support': [], 'predicate_support': [{'evidence_id': value_id}],
            'subject_support': sources, 'condition_support': [], 'counterevidence_support': [],
            'reason': '分别核对原主体、表单中的主体角色和当前分子量记录。',
        }]}

    monkeypatch.setattr(model_adapter, 'chat_with_schema', respond)
    result = LocalModelRecognitionAdapter(object(), model_identity='form-owner').inspect(
        task, context, predicate, None)
    assert result.properties[0].raw_value == '123.45'
    assert result.properties[0].policy_eligible is (support == 'both')
    assert result.model_calls == 2
