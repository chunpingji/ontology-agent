"""Record boundaries and scope permissions, independent of extraction quality."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.record_retrieval import RecordIndex
from app.schemas.evidence import (
    Candidate,
    DocumentProvenance,
    ExtractionTask,
    TaskBudget,
    TypeVerification,
)
from app.services.extraction.evidence_scope import candidate_ref, scope_contains
from app.services.extraction.extraction_tasks import GenericExtractionRunner, PartialTaskFailure
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:Report": {
        "label": "Report",
        "properties": [{"iri": "urn:temperature", "label": "Temperature", "datatype": "decimal"}],
        "relationships": [{"iri": "urn:describes", "label": "describes", "range": ["urn:Item"]}],
    },
    "urn:Item": {
        "label": "Item",
        "properties": [
            {"iri": "urn:identifier", "label": "Identifier", "identity_key": True},
            {"iri": "urn:temperature", "label": "Temperature", "datatype": "decimal"},
        ],
        "relationships": [{"iri": "urn:uses", "label": "uses", "range": ["urn:Device"]}],
    },
    "urn:Device": {"label": "Device", "properties": [], "relationships": []},
}
TEMPERATURE = SCHEMA["urn:Item"]["properties"][1]


def analyze(document, tmp_path):
    path = tmp_path / "record-retrieval.docx"
    document.save(path)
    return analyze_word_core(path).ir


def unit(ir, text):
    return next(u for u in ir.evidence_units if u.text == text)


def entity(ir, source, *, name=None, identity="subject", verified_identity=False):
    name = name or source.text
    start = source.text.index(name)
    return Candidate(
        candidate_id=identity,
        kind="entity",
        class_iri="urn:Item",
        text=name,
        validation_status="passed",
        identity={
            "instance_iri": "urn:instance:item-1",
            "key_predicate": "urn:identifier",
            "key_value": "K-1",
        }
        if verified_identity
        else {},
        type_verification=TypeVerification(
            supported=True,
            identity_supported=verified_identity,
            reason="independent fixture identity decision",
        ),
        provenance=[
            DocumentProvenance(
                anchors=[ir.anchor(source.evidence_id, start, start + len(name))], excerpts=[name]
            )
        ],
    )


def root(ir):
    source = next(u for u in ir.evidence_units if u.text)
    return entity(ir, source).model_copy(
        update={
            "class_iri": "urn:Report",
            "identity": {"document_root": ir.document_hash},
        }
    )


def contents(ir, ranges):
    return {ir.unit(r.evidence_id).text[r.start : r.end] for r in ranges}


def for_text(ir, plan, text):
    return next(record for record in plan.records if text in contents(ir, record.source_ranges))


def test_logical_rows_preserve_merged_memberships_and_headers_as_context(tmp_path):
    doc = Document()
    doc.add_heading("Measurements", level=1)
    table = doc.add_table(rows=4, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Item", "Temperature", "Period"], strict=True):
        cell.text = text
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Item K-1"
    table.cell(1, 1).text, table.cell(1, 2).text = "30", "Day 1"
    table.cell(2, 1).text, table.cell(2, 2).text = "40", "Day 2"
    for cell, text in zip(table.rows[3].cells, ["Item Z-9", "90", "Day 9"], strict=True):
        cell.text = text
    ir = analyze(doc, tmp_path)
    subject = entity(ir, unit(ir, "Item K-1"))
    index = RecordIndex(ir, SCHEMA)
    plan = index.plan(subject, TEMPERATURE, kind="property")
    first, second, other = (for_text(ir, plan, text) for text in ("30", "40", "90"))
    assert contents(ir, first.target_ranges) == {"Item K-1", "30", "Day 1"}
    assert contents(ir, second.target_ranges) == {"Item K-1", "40", "Day 2"}
    assert other.target_ranges == []
    assert contents(ir, second.binding_ranges) == {"Item", "Temperature", "Period"}
    membership = next(
        c for c in second.cell_memberships if c["evidence_id"] == unit(ir, "Item K-1").evidence_id
    )
    assert membership["data_row_indices"] == [1, 2] and membership["shared_across_rows"]
    context = index.record_context(second)
    assert {f["text"] for f in context["fragments"] if f["fact_eligible"]} == {
        "Item K-1",
        "40",
        "Day 2",
    }
    assert all(not f["fact_eligible"] for f in context["fragments"] if f["text"] == "Temperature")
    assert "30" not in {f["text"] for f in context["fragments"]}


def test_nested_rows_keep_parent_containment_separate_from_local_column_headers(tmp_path):
    doc = Document()
    outer = doc.add_table(rows=2, cols=2)
    outer.cell(0, 0).text, outer.cell(0, 1).text = "Outer item", "Outer values"
    outer.cell(1, 0).text = "Parent P"
    nested = outer.cell(1, 1).add_table(rows=2, cols=2)
    nested.cell(0, 0).text, nested.cell(0, 1).text = "Inner item", "Temperature"
    nested.cell(1, 0).text, nested.cell(1, 1).text = "Child C", "30"
    ir = analyze(doc, tmp_path)
    index = RecordIndex(ir, SCHEMA)
    plan = index.plan(root(ir), TEMPERATURE, kind="property")
    record = for_text(ir, plan, "30")
    assert len(record.table_path) == 3
    assert contents(ir, record.target_ranges) == {"Child C", "30"}
    assert contents(ir, record.binding_ranges) == {"Inner item", "Temperature"}
    assert "Outer values" in contents(ir, record.parent_table_ranges)
    assert record.parent_links[0]["relationship"] == "physical_containment_only"
    context = index.record_context(record)
    outer_fragments = [f for f in context["fragments"] if f["text"] == "Outer values"]
    assert outer_fragments and all(f["purpose"] == "parent_table_context" for f in outer_fragments)
    assert all(not f["fact_eligible"] for f in outer_fragments)


def identity_document(tmp_path):
    doc = Document()
    doc.add_heading("First record", level=1)
    doc.add_paragraph("Item K-1 registered here.")
    doc.add_heading("Second record", level=1)
    doc.add_paragraph("Item K-1 measured elsewhere.")
    doc.add_paragraph("Temperature 30.")
    return analyze(doc, tmp_path)


def test_same_name_does_not_grant_cross_section_scope(tmp_path):
    ir = identity_document(tmp_path)
    subject = entity(ir, unit(ir, "Item K-1 registered here."), name="Item K-1")
    alias = entity(ir, unit(ir, "Item K-1 measured elsewhere."), name="Item K-1", identity="other")
    index = RecordIndex(ir, SCHEMA)
    plan = index.plan(subject, TEMPERATURE, kind="property", candidates=[subject, alias])
    remote = for_text(ir, plan, "Item K-1 measured elsewhere.")
    assert remote.authorization == "needs_identity_verification"
    assert not remote.target_ranges
    assert for_text(ir, plan, "Temperature 30.").authorization == "needs_subject_binding"
    assert not scope_contains(
        plan.authorized_scope, ir.anchor(unit(ir, "Temperature 30.").evidence_id), ir
    )


def test_only_declared_verified_identity_can_join_remote_mentions(tmp_path):
    ir = identity_document(tmp_path)
    subject = entity(
        ir, unit(ir, "Item K-1 registered here."), name="Item K-1", verified_identity=True
    )
    alias = entity(
        ir,
        unit(ir, "Item K-1 measured elsewhere."),
        name="Item K-1",
        identity="other",
        verified_identity=True,
    )
    index = RecordIndex(ir, SCHEMA)
    plan = index.plan(subject, TEMPERATURE, kind="property", candidates=[alias])
    remote = for_text(ir, plan, "Item K-1 measured elsewhere.")
    assert remote.authorization == "verified_identity" and remote.target_ranges
    assert remote.identity_dependency_refs == [candidate_ref(alias)]
    # Identity authorizes that mention's local record, not a neighboring paragraph.
    assert not for_text(ir, plan, "Temperature 30.").target_ranges
    forged = alias.model_copy(
        update={
            "type_verification": TypeVerification(
                supported=True, identity_supported=False, reason="same name alone"
            ),
        }
    )
    unverified = index.plan(subject, TEMPERATURE, kind="property", candidates=[forged])
    assert not for_text(ir, unverified, "Item K-1 measured elsewhere.").target_ranges
    wrong_key = alias.model_copy(
        update={"identity": {**alias.identity, "key_predicate": "urn:temperature"}}
    )
    unverified = index.plan(subject, TEMPERATURE, kind="property", candidates=[wrong_key])
    assert not for_text(ir, unverified, "Item K-1 measured elsewhere.").target_ranges


def test_unmatched_records_are_retained_and_only_execution_updates_coverage(tmp_path):
    ir = identity_document(tmp_path)
    index = RecordIndex(ir, SCHEMA)
    plan = index.plan(root(ir), TEMPERATURE, kind="property")
    assert len(plan.records) == len(index.records)
    assert any(not record.matched for record in plan.records)
    assert plan.coverage["unsearched_records"] == len(index.records)
    for record in plan.records:
        plan.mark(record.record_id, "no_evidence", task_ids=[record.record_id])
    assert plan.coverage["complete"]
    limited = index.plan(root(ir), TEMPERATURE, kind="property", max_records=1)
    assert len(limited.records) == len(plan.records)
    omitted = next(r for r in limited.records if not r.selected)
    assert limited.ledger[omitted.record_id]["reason"] == "rank_limit"
    with pytest.raises(ValueError, match="unselected"):
        limited.mark(omitted.record_id, "searched")


def test_summaries_rank_records_but_never_become_source_fragments(tmp_path):
    ir = identity_document(tmp_path)
    metadata = {
        n["node_id"]: {
            "summary_status": "completed",
            "analysis_id": ir.analysis_id,
            "content_summary": "Temperature Phantom-summary",
        }
        for n in ir.nodes
    }
    index = RecordIndex(ir, SCHEMA, metadata=metadata)
    plan = index.plan(root(ir), TEMPERATURE, kind="property")
    assert any(record.score_components["summaries"] > 0 for record in plan.records)
    assert all(
        "Phantom-summary" not in f["text"]
        for r in plan.records
        for f in index.record_context(r)["fragments"]
    )
    changed = deepcopy(metadata)
    for value in changed.values():
        value["analysis_id"] = "another-analysis"
    stale = RecordIndex(ir, SCHEMA, metadata=changed)
    other = stale.plan(root(ir), TEMPERATURE, kind="property")
    assert all(record.score_components["summaries"] == 0 for record in other.records)
    assert other.plan_id != plan.plan_id


def test_direct_range_does_not_expand_downstream_relationship_types(tmp_path):
    ir = identity_document(tmp_path)
    schema = deepcopy(SCHEMA)
    schema["urn:SpecialItem"] = {"label": "Special Item", "parents": ["urn:Item"]}
    index = RecordIndex(ir, schema)
    plan = index.plan(root(ir), schema["urn:Report"]["relationships"][0], kind="relationship")
    assert set(plan.direct_range_class_iris) == {"urn:Item", "urn:SpecialItem"}
    assert "urn:Device" not in plan.direct_range_class_iris
    with pytest.raises(ValueError, match="not declared"):
        index.plan(root(ir), {"iri": "urn:invented"}, kind="property")


@pytest.mark.parametrize("reference_supported", [False, True])
def test_provisional_record_scope_requires_independent_production_reference_verdict(
    tmp_path, reference_supported
):
    ir = identity_document(tmp_path)
    subject = entity(ir, unit(ir, "Item K-1 registered here."), name="Item K-1")
    index = RecordIndex(ir, SCHEMA)
    plan = index.plan(subject, TEMPERATURE, kind="property")
    record = for_text(ir, plan, "Temperature 30.")
    accepted = plan.authorized_scope.model_dump(mode="json")
    context = index.reference_context(plan, record.record_id)
    scope = context["scope"]
    assert plan.authorized_scope.model_dump(mode="json") == accepted
    assert scope.construction_evidence == plan.authorized_scope.construction_evidence
    assert scope.reference_ranges == record.retrieved_target_ranges
    assert plan.ledger[record.record_id]["status"] == "unsearched"
    assert record.authorization == "needs_subject_binding" and not record.target_ranges
    stages = []

    def model(system, user, schema, budget):
        request = json.loads(user)
        stages.append(request["stage"])
        value_source = unit(ir, "Temperature 30.")
        local = {"evidence_id": value_source.evidence_id, "text": value_source.text}
        if request["stage"] == "recall":
            return {
                "assertions": [
                    {
                        "value": {"evidence_id": value_source.evidence_id, "text": "30"},
                        "assertion_spans": [local],
                    }
                ]
            }
        if request["stage"] == "verify_binding":
            return {
                "supported": True,
                "subject_candidate_id": subject.candidate_id,
                "assertion_status": "affirmed",
                "method": "explicit_assertion",
                "assertion_spans": [local],
            }
        return {
            "supported": reference_supported,
            "subject_candidate_id": subject.candidate_id,
            "reason": "independent reference fixture",
            "assertion_spans": [
                {"evidence_id": document_anchors_for(subject)[0].evidence_id, "text": subject.text},
                local,
            ],
        }

    runner = GenericExtractionRunner(
        SCHEMA,
        TestTokenizer(),
        model,
        model_identity="record-fixture",
        budget=TaskBudget(max_input_tokens=100000),
    )
    task = ExtractionTask(
        task_id="record-property",
        task_kind="property",
        subject=candidate_ref(subject),
        predicate_iri=TEMPERATURE["iri"],
        predicate_definition=TEMPERATURE,
        target_ranges=record.retrieved_target_ranges,
        target_evidence_ids=[r.evidence_id for r in record.retrieved_target_ranges],
        scope=scope,
        budget=runner.budget,
    )
    if reference_supported:
        values = runner.execute_task(task, ir, {subject.candidate_id: subject})
        assert len(values) == 1 and values[0].positive_eligible
        assert any(binding.method == "identity_reference" for binding in values[0].bindings)
    else:
        with pytest.raises(PartialTaskFailure, match="reference_not_supported"):
            runner.execute_task(task, ir, {subject.candidate_id: subject})
    assert stages == ["recall", "verify_binding", "verify_reference"]


def document_anchors_for(candidate):
    return [anchor for provenance in candidate.provenance for anchor in provenance.anchors]
