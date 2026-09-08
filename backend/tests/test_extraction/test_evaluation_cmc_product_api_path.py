"""Synthetic multi-chapter CMC path through real execution and citation replay.

Only LLM outputs are substituted: no silver, frozen artifacts, or execute_task mocks.
"""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.quality_guided_variant import build_quality_guided_variant
from app.schemas.evidence import TaskBudget
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

DEV = "https://ontology.pharma-gmp.cn/slpra/drug-development/"
DRUG = "https://ontology.pharma-gmp.cn/slpra/drug/"
CMC_REPORT = DEV + "CMCReport"
DRUG_PRODUCT = DRUG + "DrugProduct"
API = DRUG + "ActivePharmaceuticalIngredient"
DESCRIBES = DEV + "describes"
HAS_API = DRUG + "hasActiveIngredient"
DOSAGE_FORM = DRUG + "dosageForm"
FOCUS_PATH = [DESCRIBES, HAS_API]
BAD_LABEL = "DrugProduct 制剂产品"
BAD_TEXT = f"目录条目：{BAD_LABEL}。这里的制剂产品只是章节标签，不是可给药药品。"
PRODUCT_NAME = "晨星制剂"
PRODUCT_TEXT = f"本报告描述的制剂产品为{PRODUCT_NAME}，其成品是口服片剂。"
API_NAME = "晨星活性成分"
API_TEXT = f"{PRODUCT_NAME}的活性药物成分（API）为{API_NAME}，该原料药用于生产上述制剂。"
OTHER_API_NAME = "暮光活性成分"
OTHER_API_TEXT = (
    f"暮光制剂的活性药物成分（API）为{OTHER_API_NAME}；本记录仅属于暮光制剂，"
    f"不能归于{PRODUCT_NAME}。"
)

# Exact IRIs/domain/range pairs from repository ontology; invented instance names.
SCHEMA = {
    CMC_REPORT: {
        "iri": CMC_REPORT,
        "label": "CMCReport",
        "description": "药学研究报告",
        "parents": [],
        "properties": [],
        "relationships": [
            {
                "iri": DESCRIBES,
                "label": "describes",
                "description": "报告明确描述的制剂产品",
                "range": [DRUG_PRODUCT],
            }
        ],
    },
    DRUG_PRODUCT: {
        "iri": DRUG_PRODUCT,
        "label": "DrugProduct",
        "description": "可给药的成品制剂药品",
        "parents": [],
        "properties": [{"iri": DOSAGE_FORM, "label": "制剂剂型", "datatype": "string"}],
        "relationships": [
            {
                "iri": HAS_API,
                "label": "含有活性成分",
                "description": "该制剂实际包含的活性药物成分",
                "range": [API],
            }
        ],
    },
    API: {
        "iri": API,
        "label": "ActivePharmaceuticalIngredient",
        "description": "活性药物成分原料药",
        "parents": [],
        "properties": [],
        "relationships": [],
    },
    DRUG + "UnreachableTestType": {
        "iri": DRUG + "UnreachableTestType",
        "label": "UnreachableTestType",
        "parents": [],
        "properties": [],
        "relationships": [],
    },
}


@pytest.fixture
def source(tmp_path):
    document = Document()
    for heading, text in (
        (BAD_LABEL, BAD_TEXT),
        ("一般说明", "本节讨论文档结构，不指认药品实体。"),
        ("成品身份", PRODUCT_TEXT),
        ("其他制剂原料", OTHER_API_TEXT),
        ("活性成分", API_TEXT),
    ):
        document.add_heading(heading, 1)
        document.add_paragraph(text)
    path = tmp_path / "independent-cmc-path.docx"
    document.save(path)
    return analyze_word_core(path).ir


def model_response(calls, *, reference_fault=None):
    def model(system, user, schema, budget):
        request = json.loads(user)
        if "local_menu" in request:
            # Even an older hard router would select only the wrong title for
            # describes, and no API records. Focus relations must bypass it.
            routes = []
            for identity, record in request["records"].items():
                source_text = " ".join(
                    record.get("source", [])
                    + [text for cell in record.get("cells", []) for text in cell["fragments"]]
                )
                selected = [
                    key
                    for key, definition in request["local_menu"].items()
                    if (definition["iri"] == DOSAGE_FORM and PRODUCT_TEXT in source_text)
                    or (definition["iri"] == DESCRIBES and BAD_TEXT in source_text)
                ]
                routes.append({"record_id": identity, "predicate_ids": selected})
            calls.append({"stage": "route_records", "request": request})
            return {"records": routes}
        context = json.loads(request["context"])
        task, stage = context["task"], request["stage"]
        event = {
            "stage": stage,
            "task": task,
            "context": context,
            "candidate": request.get("candidate"),
        }
        calls.append(event)
        targets = [fragment for fragment in context["fragments"] if fragment["fact_eligible"]]

        def respond(value):
            event["response"] = deepcopy(value)
            return value

        def fragment(text, *, target=False):
            return next(
                f for f in (targets if target else context["fragments"]) if f["text"] == text
            )

        def quote(value, text=None):
            reference = {"evidence_id": value["anchor"]["evidence_id"]}
            if text is not None:
                reference["text"] = text
            return reference

        if stage == "verify_entity_types":
            assert "discovery_relation" not in task["predicate_definition"]
            return respond(
                {
                    "decisions": [
                        {
                            "candidate_id": proposed["candidate_id"],
                            "supported": proposed["text"] != BAD_LABEL,
                            "identity_supported": False,
                            "reason": "目录标签不是具体制剂"
                            if proposed["text"] == BAD_LABEL
                            else "原文明示具体制剂或活性成分",
                        }
                        for proposed in request["candidate"]["proposed_entities"]
                    ]
                }
            )
        if stage == "recall" and task["task_kind"] == "entity":
            classes = task["predicate_definition"]["classes"]
            class_names = {definition["name"]: alias for alias, definition in classes.items()}
            assert "UnreachableTestType" not in class_names
            proposals = []
            for type_name, source_text, name in (
                ("DrugProduct", BAD_TEXT, BAD_LABEL),
                ("DrugProduct", PRODUCT_TEXT, PRODUCT_NAME),
                ("ActivePharmaceuticalIngredient", API_TEXT, API_NAME),
                ("ActivePharmaceuticalIngredient", OTHER_API_TEXT, OTHER_API_NAME),
            ):
                if type_name in class_names and any(f["text"] == source_text for f in targets):
                    proposals.append(
                        {
                            "class_iri": class_names[type_name],
                            "mention": quote(fragment(source_text, target=True), name),
                            "assertion_spans": [quote(fragment(source_text, target=True))],
                            "supported": True,
                            "reason": "召回原文提及，交由独立类型复核",
                        }
                    )
            return respond({"entities": proposals})
        if stage == "recall" and task["task_kind"] == "relationship":
            source_text = next(
                (
                    text
                    for text in (PRODUCT_TEXT, API_TEXT, OTHER_API_TEXT)
                    if any(f["text"] == text for f in targets)
                ),
                None,
            )
            return respond(
                {
                    "assertions": [
                        {
                            "object_candidate_id": ref["candidate_id"],
                            "assertion_status": "affirmed",
                            "assertion_spans": [quote(fragment(source_text, target=True))],
                        }
                        for ref in task["object_candidates"]
                    ]
                    if source_text
                    else []
                }
            )
        if stage == "recall":
            assert task["predicate_iri"] == DOSAGE_FORM
            return respond(
                {
                    "assertions": [
                        {
                            "value": quote(fragment(PRODUCT_TEXT, target=True), "片剂"),
                            "assertion_status": "affirmed",
                        }
                    ]
                }
            )
        if stage == "verify_binding":
            proposed = request["candidate"]
            source_text = next(
                text
                for text in (PRODUCT_TEXT, API_TEXT, OTHER_API_TEXT)
                if any(f["text"] == text for f in targets)
            )
            # Deliberately accept wrong-record binding. Independent reference
            # verification must still reject its wrong-product attribution.
            return respond(
                {
                    "supported": True,
                    "subject_candidate_id": proposed["subject"]["candidate_id"],
                    "object_candidate_id": (proposed.get("object") or {}).get("candidate_id"),
                    "assertion_status": proposed["assertion_status"],
                    "method": "explicit_assertion",
                    "record_mapping": {},
                    "assertion_spans": [quote(fragment(source_text, target=True))],
                }
            )
        assert stage == "verify_reference", stage
        assert task["predicate_iri"] == HAS_API
        source_text = next(
            text for text in (API_TEXT, OTHER_API_TEXT) if any(f["text"] == text for f in targets)
        )
        groups = context["reference_verification"]
        assert groups["subject_construction_sources"] and groups["current_fact_sources"]
        bridge = [quote(fragment(PRODUCT_TEXT)), quote(fragment(source_text, target=True))]
        subject = task["subject"]["candidate_id"]
        if reference_fault == "wrong_subject" and source_text == API_TEXT:
            subject = task["path_root"]["candidate_id"]
        if reference_fault == "current_record_only" and source_text == API_TEXT:
            bridge = bridge[1:]
        return respond(
            {
                "supported": source_text == API_TEXT,
                "subject_candidate_id": subject,
                "assertion_spans": bridge,
                "reason": "双端原文明示晨星制剂身份"
                if source_text == API_TEXT
                else "该API仅属于暮光制剂，不能归给晨星制剂",
            }
        )

    return model


def run_path(source, calls, *, reference_fault=None):
    base = GenericExtractionRunner(
        deepcopy(SCHEMA),
        TestTokenizer(),
        model_response(calls, reference_fault=reference_fault),
        model_identity="independent-cmc-path-fixture",
        compact_identifiers=True,
        budget=TaskBudget(
            max_input_tokens=200000,
            max_output_tokens=4096,
            max_tasks=200,
            max_hops=3,
            max_regions_per_task=32,
        ),
    )
    runner = build_quality_guided_variant(base, source, focus_path=FOCUS_PATH)
    assert runner.execute_task.__func__ is GenericExtractionRunner.execute_task
    return runner, runner.run(source, effective_class=CMC_REPORT)


def test_first_title_type_rejection_does_not_block_later_product_api_path_or_product_property(
    source,
):
    calls = []
    before = source.model_dump(mode="json")
    runner, result = run_path(source, calls)
    positive = [candidate for candidate in result.candidates if candidate.positive_eligible]
    by_id = {candidate.candidate_id: candidate for candidate in result.candidates}
    products = [
        candidate
        for candidate in positive
        if candidate.kind == "entity" and candidate.class_iri == DRUG_PRODUCT
    ]
    assert [candidate.text for candidate in products] == [PRODUCT_NAME], result.diagnostics
    rejected = next(
        candidate
        for candidate in result.candidates
        if candidate.text == BAD_LABEL and not candidate.identity.get("document_root")
    )
    assert not rejected.positive_eligible and rejected.type_verification.supported is False
    edges = [candidate for candidate in positive if candidate.kind == "relationship"]
    describes = next(edge for edge in edges if edge.predicate_iri == DESCRIBES)
    ingredient = next(edge for edge in edges if edge.predicate_iri == HAS_API)
    assert len(edges) == 2, result.diagnostics
    assert by_id[describes.subject.candidate_id].identity["document_root"] == source.document_hash
    assert describes.object == ingredient.subject
    assert by_id[ingredient.object.candidate_id].text == API_NAME
    assert by_id[ingredient.object.candidate_id].class_iri == API
    assert ingredient.relationship_path == FOCUS_PATH
    assert describes.candidate_id in {ref.candidate_id for ref in ingredient.dependency_refs}
    assert any(binding.method == "identity_reference" for binding in ingredient.bindings)
    properties = [candidate for candidate in positive if candidate.kind == "property"]
    assert len(properties) == 1 and properties[0].predicate_iri == DOSAGE_FORM
    assert properties[0].subject == describes.object
    assert properties[0].literal.raw_value == "片剂"
    assert properties[0].literal.normalized_value == "片剂"
    assert "reference_not_supported" in result.diagnostics
    assert not any(by_id[edge.object.candidate_id].text == OTHER_API_NAME for edge in edges)
    decisions = [
        (proposed["text"], decision["supported"])
        for call in calls
        if call["stage"] == "verify_entity_types"
        for proposed, decision in zip(
            call["candidate"]["proposed_entities"], call["response"]["decisions"], strict=True
        )
    ]
    assert decisions.index((BAD_LABEL, False)) < decisions.index((PRODUCT_NAME, True))
    assert (API_NAME, True) in decisions
    reference_calls = [call for call in calls if call["stage"] == "verify_reference"]
    assert any(call["response"]["supported"] for call in reference_calls)
    assert any(not call["response"]["supported"] for call in reference_calls)
    assert any(row.get("reference_verification_required") for row in runner.coverage)
    assert all(
        definition["iri"] not in FOCUS_PATH
        for call in calls
        if call["stage"] == "route_records"
        for definition in call["request"]["local_menu"].values()
    )
    assert source.model_dump(mode="json") == before


@pytest.mark.parametrize("reference_fault", ["wrong_subject", "current_record_only"])
def test_api_relationship_needs_correct_product_and_both_reference_ends(source, reference_fault):
    calls = []
    _, result = run_path(source, calls, reference_fault=reference_fault)
    positive = [candidate for candidate in result.candidates if candidate.positive_eligible]
    assert any(
        candidate.kind == "relationship" and candidate.predicate_iri == DESCRIBES
        for candidate in positive
    )
    assert any(
        candidate.kind == "property" and candidate.predicate_iri == DOSAGE_FORM
        for candidate in positive
    )
    assert any(candidate.kind == "entity" and candidate.text == API_NAME for candidate in positive)
    assert not any(
        candidate.kind == "relationship" and candidate.predicate_iri == HAS_API
        for candidate in positive
    )
    assert "reference_not_supported" in result.diagnostics
    assert result.completion == "incomplete"
    assert any(
        call["stage"] == "verify_reference" and call["response"]["supported"] for call in calls
    )
