"""Atomic table citations through the real closed transport and source verifier."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.evaluation.citation_protocol import CitationProtocol
from app.schemas.evidence import (
    Candidate,
    DocumentProvenance,
    EvidenceRange,
    EvidenceScope,
    ExtractionTask,
    TaskBudget,
)
from app.services.extraction.evidence_scope import build_scope, candidate_ref
from app.services.extraction.extraction_tasks import (
    SYSTEM,
    AssertionResponse,
    BindingDecision,
    EntityResponse,
    EntityTypeResponse,
    GenericExtractionRunner,
    PartialTaskFailure,
    ReferenceDecision,
)
from app.services.extraction.hierarchical_context import ContextEnvelope
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:Product": {
        "iri": "urn:Product",
        "label": "Product",
        "description": "A named product",
        "parents": [],
        "properties": [
            {"iri": "urn:dose", "label": "Dose", "datatype": "decimal", "canonical_unit": "mg"}
        ],
        "relationships": [],
    }
}


@pytest.fixture
def source(tmp_path):
    document = Document()
    document.add_heading("Results", 1)
    table = document.add_table(rows=3, cols=3)
    for row, cells in zip(
        table.rows,
        [
            ["Product", "Dose (mg)", "Approval"],
            ["A", "5 mg", "Not approved"],
            ["B", "9 mg", "Approved"],
        ],
        strict=True,
    ):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    document.add_paragraph("Legend: — means not tested.")
    document.add_paragraph("These doses are planned, not actual.")
    document.add_paragraph("Repeated 5 mg; 5 mg.")
    document.add_heading("Unrelated appendix", 1)
    document.add_paragraph("Unrelated material that must not inflate verification context." * 50)
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [["Other item", "Other dose"], ["C", "7 mg"]], strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    path = tmp_path / "citations.docx"
    document.save(path)
    return analyze_word_core(path).ir


def unit(ir, text):
    return next(item for item in ir.evidence_units if item.text == text)


def entity(ir, text):
    anchor = ir.anchor(unit(ir, text).evidence_id)
    return Candidate(
        candidate_id=f"entity-{text}",
        kind="entity",
        class_iri="urn:Product",
        text=text,
        validation_status="passed",
        provenance=[
            DocumentProvenance(
                document_role=ir.document_role,
                anchors=[anchor],
                excerpts=[text],
            )
        ],
    )


def envelope(ir, *, kind="property"):
    a, b, unrelated = [entity(ir, text) for text in ("A", "B", "C")]
    scope = EvidenceScope(
        scope_id="test-explicit-scope",
        document_hash=ir.document_hash,
        subject=candidate_ref(a),
        ranges=[EvidenceRange(evidence_id=u.evidence_id) for u in ir.evidence_units],
    )
    task = ExtractionTask(
        task_id="citation-task",
        task_kind=kind,
        subject=candidate_ref(a),
        scope=scope,
        competing_subjects=[candidate_ref(b)],
        object_candidates=[candidate_ref(unrelated)],
        predicate_iri="urn:dose" if kind != "entity" else None,
        predicate_definition=SCHEMA["urn:Product"]["properties"][0]
        if kind != "entity"
        else {"classes": SCHEMA},
        target_class_iris=["urn:Product"] if kind == "entity" else [],
        target_evidence_ids=[u.evidence_id for u in ir.evidence_units],
    )
    fragments, facts, bindings = [], [], []
    for source_unit in ir.evidence_units:
        anchor = ir.anchor(source_unit.evidence_id)
        header = source_unit.row_index == 0 and bool(source_unit.table_path)
        fact = source_unit.kind != "heading" and not header
        fragments.append(
            {
                "anchor": anchor.model_dump(mode="json"),
                "text": source_unit.text,
                "purpose": "target" if fact else "table_record_metadata",
                "fact_eligible": fact,
                "role": "source",
            }
        )
        bindings.append(anchor)
        if fact:
            facts.append(anchor)
    payload = {
        "task": task.model_dump(mode="json"),
        "subjects": {c.candidate_id: c.model_dump(mode="json") for c in (a, b, unrelated)},
        "fragments": fragments,
        "effective_class": "urn:Product",
    }
    return ContextEnvelope(
        lookup_key="test",
        context_hash="original",
        completion="complete",
        serialized_input=json.dumps(payload),
        fragments=fragments,
        allowed_fact_regions=facts,
        allowed_binding_regions=bindings,
    )


def candidate(ir):
    return {
        "kind": "property",
        "candidate_id": "proposal",
        "subject": candidate_ref(entity(ir, "A")).model_dump(),
        "predicate_iri": "urn:dose",
        "assertion_status": "conditional",
        "literal": {"raw_value": "5"},
        "provenance": [
            {
                "kind": "document",
                "document_role": ir.document_role,
                "anchors": [ir.anchor(unit(ir, "5 mg").evidence_id, 0, 1).model_dump(mode="json")],
                "excerpts": ["5"],
            }
        ],
    }


def source_id(protocol, ir, text):
    canonical = unit(ir, text).evidence_id
    if not protocol.compact_identifiers:
        return canonical
    return next(
        alias for alias, value in protocol.references["evidence"].items() if value == canonical
    )


def subject_id(protocol, name="A"):
    canonical = f"entity-{name}"
    if not protocol.compact_identifiers:
        return canonical
    return next(
        alias for alias, value in protocol.references["candidate"].items() if value == canonical
    )


def quote(protocol, ir, source_text, text=None, **coordinates):
    return {
        "evidence_id": source_id(protocol, ir, source_text),
        "text": text if text is not None else source_text,
        **coordinates,
    }


@pytest.mark.parametrize("compact", [False, True])
def test_atomic_table_quote_list_preserves_polarity_roles_and_source(source, compact):
    original = envelope(source)
    before_ir, before_envelope = source.model_dump(), original.model_dump()
    protocol = CitationProtocol(
        SYSTEM,
        original,
        BindingDecision,
        ir=source,
        stage="verify_binding",
        candidate=candidate(source),
        compact_identifiers=compact,
    )
    response = {
        "supported": False,
        "subject_candidate_id": subject_id(protocol),
        "object_candidate_id": None,
        "assertion_status": "conditional",
        "method": "table_record",
        "record_mapping": {},
        "assertion_spans": [
            quote(protocol, source, text) for text in ("A", "Dose (mg)", "5 mg", "Not approved")
        ],
        "conditions": [quote(protocol, source, "These doses are planned, not actual.")],
    }
    decoded = protocol.decode(response)
    assert decoded["supported"] is False
    assert decoded["assertion_status"] == "conditional"
    assert len({s["evidence_id"] for s in decoded["assertion_spans"]}) == 4
    for span in [*decoded["assertion_spans"], *decoded["conditions"]]:
        assert (
            source.resolve(source.anchor(span["evidence_id"], span["start"], span["end"]))
            == span["text"]
        )
    assert source.model_dump() == before_ir
    assert original.model_dump() == before_envelope
    context = json.loads(json.loads(protocol.user)["context"])
    texts = {f["text"] for f in context["fragments"]}
    assert {
        "A",
        "B",
        "5 mg",
        "9 mg",
        "Not approved",
        "Approval",
        "Dose (mg)",
        "Legend: — means not tested.",
        "These doses are planned, not actual.",
    } <= texts
    assert "C" not in texts and "7 mg" not in texts
    assert not any("Unrelated material" in text for text in texts)
    assert context["task"]["object_candidates"] == []
    assert context["task"]["competing_subjects"][0]["candidate_id"] == subject_id(protocol, "B")
    assert all(f["scope_id"] == "test-explicit-scope" and "role" in f for f in context["fragments"])
    assert all(f["binding_eligible"] for f in context["fragments"])
    assert "context" not in protocol.schema["$defs"]["BindingSpan"]["properties"]
    assert protocol.schema["$defs"]["BindingSpan"]["additionalProperties"] is False
    assert json.loads(protocol.user)["output_contract"] == protocol.schema


@pytest.mark.parametrize("compact", [False, True])
@pytest.mark.parametrize(
    "failure", ["context", "stitched_text", "unknown", "bounds", "wrong_offset"]
)
def test_invalid_atomic_citations_fail_closed(source, compact, failure):
    protocol = CitationProtocol(
        SYSTEM,
        envelope(source),
        AssertionResponse,
        ir=source,
        stage="recall",
        compact_identifiers=compact,
    )
    value = quote(protocol, source, "5 mg", "5")
    if failure == "context":
        value["context"] = "A | Dose (mg) | 5 mg"
    elif failure == "stitched_text":
        value["text"] = "A | 5 mg"
    elif failure == "unknown":
        value["evidence_id"] = "made-up"
    elif failure == "bounds":
        value.update(start=90, end=91)
    else:
        value.update(start=2, end=3)
    with pytest.raises(ValueError):
        protocol.decode({"assertions": [{"value": value, "assertion_status": "affirmed"}]})


def test_header_is_binding_context_and_cannot_become_a_value(source):
    protocol = CitationProtocol(
        SYSTEM, envelope(source), AssertionResponse, ir=source, stage="recall"
    )
    with pytest.raises(ValueError, match="disallowed_citation"):
        protocol.decode(
            {
                "assertions": [
                    {"assertion_status": "affirmed", "value": quote(protocol, source, "Dose (mg)")}
                ]
            }
        )


def test_duplicate_quote_rejects_even_with_correct_model_coordinates(source):
    protocol = CitationProtocol(
        SYSTEM, envelope(source), AssertionResponse, ir=source, stage="recall"
    )
    repeated = "Repeated 5 mg; 5 mg."
    assertion = {"assertion_status": "affirmed", "value": quote(protocol, source, repeated, "5 mg")}
    with pytest.raises(ValueError, match="ambiguous_source_quote"):
        protocol.decode({"assertions": [assertion]})
    start = repeated.rfind("5 mg")
    assertion["value"].update(start=start, end=start + 4)
    with pytest.raises(ValueError, match="non_atomic_source_citation"):
        protocol.decode({"assertions": [assertion]})


def test_explicit_counterevidence_is_never_removed_by_local_projection(source):
    original = envelope(source)
    payload = json.loads(original.serialized_input)
    extra = next(f for f in payload["fragments"] if "Unrelated material" in f["text"])
    extra["role"] = "counterevidence"
    original.serialized_input = json.dumps(payload)
    protocol = CitationProtocol(
        SYSTEM,
        original,
        BindingDecision,
        ir=source,
        stage="verify_binding",
        candidate=candidate(source),
    )
    assert any(f["role"] == "counterevidence" for f in protocol.envelope.fragments)


def test_unpermitted_retrieval_context_never_becomes_a_binding_citation(source):
    original = envelope(source)
    payload = json.loads(original.serialized_input)
    fragment = next(f for f in payload["fragments"] if f["text"] == "7 mg")
    fragment.update(
        purpose="retrieval_candidate",
        record_role="parent_context",
        binding_permitted=False,
        fact_eligible=False,
    )
    original.serialized_input = json.dumps(payload)
    protocol = CitationProtocol(
        SYSTEM,
        original,
        BindingDecision,
        ir=source,
        stage="verify_binding",
        candidate=candidate(source),
    )
    retained = next(f for f in protocol.envelope.fragments if f["text"] == "7 mg")
    assert retained["record_role"] == "parent_context"
    assert retained["binding_eligible"] is False
    with pytest.raises(ValueError, match="disallowed_citation"):
        protocol.decode(
            {
                "supported": True,
                "subject_candidate_id": subject_id(protocol),
                "object_candidate_id": None,
                "assertion_status": "affirmed",
                "method": "explicit_assertion",
                "assertion_spans": [quote(protocol, source, "7 mg")],
            }
        )


def test_citation_region_intersections_never_widen_fact_permissions(source):
    original = envelope(source)
    identity = unit(source, "5 mg").evidence_id
    payload = json.loads(original.serialized_input)
    full = next(f for f in payload["fragments"] if f["anchor"]["evidence_id"] == identity)
    full["fact_eligible"], full["purpose"] = False, "table_record_metadata"
    target = deepcopy(full)
    target.update(
        anchor=source.anchor(identity, 0, 1).model_dump(mode="json"),
        text="5",
        fact_eligible=True,
        purpose="target",
    )
    payload["fragments"].append(target)
    original.serialized_input = json.dumps(payload)
    original.allowed_fact_regions = [source.anchor(identity, 0, 1)]
    for f in payload["fragments"]:
        if f is not target:
            f["fact_eligible"] = False
    original.serialized_input = json.dumps(payload)
    protocol = CitationProtocol(SYSTEM, original, AssertionResponse, ir=source, stage="recall")
    with pytest.raises(ValueError, match="scope|outside"):
        protocol.decode(
            {
                "assertions": [
                    {"assertion_status": "affirmed", "value": quote(protocol, source, "5 mg", "mg")}
                ]
            }
        )
    value = quote(protocol, source, "5 mg", "5")
    assert protocol.decode({"assertions": [{"assertion_status": "affirmed", "value": value}]})


def test_mismatched_source_fragment_rejected_before_request(source):
    original = envelope(source)
    payload = json.loads(original.serialized_input)
    payload["fragments"][0]["text"] += "invented"
    original.serialized_input = json.dumps(payload)
    with pytest.raises(ValueError, match="fragment_source_mismatch"):
        CitationProtocol(SYSTEM, original, EntityResponse, ir=source, stage="recall")


def test_entity_verification_keeps_class_competition_but_removes_discovery_hint(source):
    original = envelope(source, kind="entity")
    payload = json.loads(original.serialized_input)
    payload["task"]["predicate_definition"]["discovery_relation"] = {"label": "recall-only hint"}
    original.serialized_input = json.dumps(payload)
    proposal = {"proposed_entities": [entity(source, "A").model_dump(mode="json")]}
    protocol = CitationProtocol(
        SYSTEM,
        original,
        EntityTypeResponse,
        ir=source,
        stage="verify_entity_types",
        candidate=proposal,
        ontology_schema=SCHEMA,
    )
    context = json.loads(json.loads(protocol.user)["context"])
    definition = context["task"]["predicate_definition"]
    assert "discovery_relation" not in definition
    assert any("property_roles" in item for item in definition["classes"].values())
    decoded = protocol.decode(
        {
            "decisions": [
                {
                    "candidate_id": subject_id(protocol),
                    "supported": False,
                    "identity_supported": False,
                    "reason": "Not an entity in this role",
                }
            ]
        }
    )
    assert decoded["decisions"][0]["candidate_id"] == "entity-A"
    assert decoded["decisions"][0]["supported"] is False


def test_reference_projection_keeps_explicit_parent_record_bridge(source):
    original = envelope(source)
    payload = json.loads(original.serialized_input)
    parent = next(f for f in payload["fragments"] if f["text"] == "7 mg")
    parent["purpose"] = "parent_record_reference_context"
    parent["fact_eligible"] = False
    original.serialized_input = json.dumps(payload)
    protocol = CitationProtocol(
        SYSTEM,
        original,
        ReferenceDecision,
        ir=source,
        stage="verify_reference",
        candidate=candidate(source),
    )
    assert any(f["text"] == "7 mg" and not f["fact_eligible"] for f in protocol.envelope.fragments)
    decision = protocol.decode(
        {
            "supported": False,
            "subject_candidate_id": subject_id(protocol),
            "assertion_spans": [quote(protocol, source, text) for text in ("A", "7 mg", "5 mg")],
            "reason": "Parent proximity does not prove identity",
        }
    )
    assert decision["supported"] is False
    assert len(decision["assertion_spans"]) == 3


@pytest.mark.parametrize("reference_mode", ["unsupported", "value_only", "identity_only"])
def test_real_reference_verifier_rejects_retrieved_record_without_identity_support(
    source, reference_mode
):
    seen = []

    class AtomicRunner(GenericExtractionRunner):
        def _invoke(self, task, original, response_type, stage, candidate=None):
            seen.append(stage)
            protocol = CitationProtocol(
                SYSTEM,
                original,
                response_type,
                ir=source,
                stage=stage,
                candidate=candidate,
                ontology_schema=self.schema,
            )
            if stage == "recall":
                raw = {
                    "assertions": [
                        {
                            "assertion_status": "affirmed",
                            "value": quote(protocol, source, "9 mg", "9"),
                        }
                    ]
                }
            elif stage == "verify_binding":
                raw = {
                    "supported": True,
                    "subject_candidate_id": subject_id(protocol),
                    "object_candidate_id": None,
                    "assertion_status": "affirmed",
                    "method": "explicit_assertion",
                    "record_mapping": {},
                    "assertion_spans": [quote(protocol, source, text) for text in ("A", "9 mg")],
                    "source_unit": quote(protocol, source, "9 mg", "mg"),
                }
            else:
                context = json.loads(json.loads(protocol.user)["context"])
                groups = context["reference_verification"]
                assert groups["subject_construction_sources"]
                assert groups["current_fact_sources"][0]["evidence_id"] == source_id(
                    protocol, source, "9 mg"
                )
                assert groups["current_record_binding_sources_to_check"]
                assert "仅引用属性值不足以通过" in protocol.system
                sources = {
                    "unsupported": ("A", "9 mg"),
                    "value_only": ("9 mg",),
                    "identity_only": ("A",),
                }
                raw = {
                    "supported": reference_mode != "unsupported",
                    "subject_candidate_id": subject_id(protocol),
                    "assertion_spans": [
                        quote(protocol, source, text) for text in sources[reference_mode]
                    ],
                    "reason": "The retrieved row belongs to B, not A",
                }
            return response_type.model_validate(protocol.decode(raw), strict=True)

    a, b = entity(source, "A"), entity(source, "B")
    scope = build_scope(source, a, [b])
    retrieved = EvidenceRange(evidence_id=unit(source, "9 mg").evidence_id)
    scope = scope.model_copy(
        update={"ranges": [*scope.ranges, retrieved], "reference_ranges": [retrieved]}
    )
    task = ExtractionTask(
        task_id="unproven-retrieved-record",
        task_kind="property",
        subject=candidate_ref(a),
        scope=scope,
        competing_subjects=[candidate_ref(b)],
        predicate_iri="urn:dose",
        predicate_definition=SCHEMA["urn:Product"]["properties"][0],
        target_evidence_ids=[retrieved.evidence_id],
        budget=TaskBudget(max_input_tokens=60000),
    )
    runner = AtomicRunner(
        SCHEMA,
        TestTokenizer(),
        lambda *args: None,
        model_identity="fixture",
        budget=task.budget,
        compact_identifiers=True,
    )
    with pytest.raises(PartialTaskFailure, match="reference_not_supported") as failure:
        runner.execute_task(task, source, {a.candidate_id: a, b.candidate_id: b})
    assert failure.value.candidates == []
    assert seen == ["recall", "verify_binding", "verify_reference"]


@pytest.mark.parametrize("compact", [False, True])
def test_adapter_runs_real_table_property_binding_and_unit_validation(source, compact):
    requests = []

    class AtomicRunner(GenericExtractionRunner):
        def _invoke(self, task, original, response_type, stage, candidate=None):
            protocol = CitationProtocol(
                SYSTEM,
                original,
                response_type,
                ir=source,
                stage=stage,
                candidate=candidate,
                ontology_schema=self.schema,
                compact_identifiers=compact,
            )
            requests.append(protocol)
            if stage == "recall":
                raw = {
                    "assertions": [
                        {
                            "assertion_status": "affirmed",
                            "value": quote(protocol, source, "5 mg", "5"),
                        }
                    ]
                }
            else:
                raw = {
                    "supported": True,
                    "subject_candidate_id": subject_id(protocol),
                    "object_candidate_id": None,
                    "assertion_status": "affirmed",
                    "method": "table_record",
                    "assertion_spans": [
                        quote(protocol, source, text) for text in ("A", "Dose (mg)", "5 mg")
                    ],
                    "source_unit": quote(protocol, source, "5 mg", "mg"),
                    "record_mapping": {},
                }
            return response_type.model_validate(protocol.decode(raw), strict=True)

    a, b = entity(source, "A"), entity(source, "B")
    scope = build_scope(source, a, [b])
    task = ExtractionTask(
        task_id="real-table-property",
        task_kind="property",
        subject=candidate_ref(a),
        scope=scope,
        competing_subjects=[candidate_ref(b)],
        predicate_iri="urn:dose",
        predicate_definition=SCHEMA["urn:Product"]["properties"][0],
        target_evidence_ids=[unit(source, "5 mg").evidence_id],
        budget=TaskBudget(max_input_tokens=60000),
    )
    runner = AtomicRunner(
        SCHEMA,
        TestTokenizer(),
        lambda *args: None,
        model_identity="fixture",
        budget=task.budget,
        compact_identifiers=compact,
    )
    result = runner.execute_task(task, source, {a.candidate_id: a, b.candidate_id: b})
    assert len(result) == 1
    assert result[0].positive_eligible, result[0].validation_issues
    assert result[0].literal.raw_value == "5"
    assert result[0].literal.canonical_unit == "mg"
    assert result[0].bindings[0].record_mapping["row_index"] == 1
    assert len(requests) == 2
