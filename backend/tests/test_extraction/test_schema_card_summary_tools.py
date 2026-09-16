"""D runner isolation, frozen inputs and the unchanged bounded Qwen protocol."""

from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document

from app.evaluation import schema_card_summary_tools as runner
from app.evaluation.schema_card_probe import CMC, digest, read, write
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.word_analysis import analyze_word_core

ITEM = "urn:test:item"


@pytest.fixture
def experiment(tmp_path, monkeypatch):
    baseline, metadata_dir = tmp_path / "C-baseline", tmp_path / "metadata"
    baseline.mkdir()
    metadata_dir.mkdir()
    document = Document()
    document.add_heading("测试章节", 1)
    document.add_paragraph("旧局部范围")
    document.add_paragraph("全IR新命中")
    document.save(baseline / "source.docx")
    analysis = analyze_word_core(baseline / "source.docx")
    ir = analysis.ir
    catalog = {iri: {"class_iri": iri, "label": "报告" if iri == CMC else "对象",
                     "properties": [], "allowed_relations": []} for iri in (CMC, ITEM)}
    write(baseline / "ir.json", ir.model_dump(mode="json"))
    write(baseline / "cards.json", catalog)
    tree = analysis.structure.section_tree.to_dict()
    tree.setdefault("layer_metadata", {}).update(
        content_summary="摘要独有虚构值", summary_status="completed", summary_source="llm",
    )
    metadata = prepare_metadata(ir, section_tree=tree, summary_version="test-summary-v1")
    write(metadata_dir / "metadata.json", {"metadata_snapshot": metadata.model_dump(mode="json")})
    write(metadata_dir / "section-tree.json", tree)
    write(metadata_dir / "summaries.json", {"content_summary": "摘要独有虚构值"})
    metadata_manifest = {
        "document_ref": "upload-test", "source_sha256": ir.document_hash,
        "analysis_id": ir.analysis_id, "structure_hash": ir.structure_hash,
        "metadata_snapshot_id": metadata.snapshot_id,
        "metadata_dependency_hash": metadata.dependency_hash,
        "input_hashes": {"ir": digest(baseline / "ir.json")},
        "output_hashes": {name: digest(metadata_dir / name) for name in runner.METADATA_FILES[:-1]},
    }
    write(metadata_dir / "manifest.json", metadata_manifest)
    frozen_code = tmp_path / "frozen.py"
    frozen_code.write_text("# Existing C code remains frozen.\n")
    (baseline / "code").mkdir()
    (baseline / "code/frozen.py").write_bytes(frozen_code.read_bytes())
    vocabulary = {
        "profile": "test-v1", "missing": [],
        "groups": {"entities": ["对象名称"], "values": [], "units": []},
        "entries": {"对象名称": {"iri": ITEM, "role": "entity_name", "description": "对象原名"}},
    }
    old_plan = {
        "get_schema_card": [{"class_iri": ITEM}, {"class_iri": ITEM}],
        "inspect_evidence": [{"ref": "OLD-REF-MUST-NOT-SURVIVE"}],
        "propose_mentions": [{"group": "entities"}],
    }
    _, _, _, menus = runner.c_runner.candidate_schema(old_plan, {}, catalog, {})
    for scope in runner.SCOPES:
        folder = baseline / "C" / scope
        (folder / "plan").mkdir(parents=True)
        write(folder / "plan/proposal.json", old_plan)
        write(folder / "subject-cards.json", menus)
        write(folder / "vocabulary.json", vocabulary)
        # If an implementation consults historical source scopes or silver it fails JSON decoding.
        (folder / "sources.json").write_text("DO NOT READ OLD SOURCES")
    for name in ("cases.json", "reference.json", "acceptance.json"):
        (baseline / name).write_text("DO NOT READ OLD SELECTORS OR SILVER")
    old = {
        "version": runner.c_runner.VERSION, "status": "completed", "run_id": "frozen-C",
        "source_sha256": ir.document_hash, "ontology_hash": "ontology",
        "model": "Qwen-fixed", "declared_model_revision": "qwen-fixed-revision",
        "gliner2_model": {"repo": "fixed", "revision": "fixed"},
        "packages": {name: "fixed" for name in runner.c_runner.PACKAGES},
        "ner_device": "cuda:0", "labels_per_batch": 1,
        "backend": runner.c_runner.BACKEND, "architecture": runner.c_runner.ARCHITECTURE,
        "input_hashes": {name: digest(baseline / name) for name in ("ir.json", "cards.json")},
        "code_hashes": {"frozen.py": digest(frozen_code)},
        "prepared_tool_hashes": {scope: {
            name: digest(baseline / "C" / scope / name)
            for name in ("plan/proposal.json", "vocabulary.json")
        } for scope in runner.SCOPES},
    }
    write(baseline / "result.json", old)
    monkeypatch.setattr(runner, "SOURCE_SHA256", ir.document_hash)
    monkeypatch.setattr(runner, "code_paths", lambda: {"frozen.py": frozen_code})
    monkeypatch.setattr(runner.c_runner, "verify_model_files", lambda _: old["gliner2_model"])
    monkeypatch.setattr(runner.c_runner, "package_versions", lambda: old["packages"])
    monkeypatch.setattr(runner.c_runner, "build_extraction_vocabulary",
                        lambda *a, **kw: deepcopy(vocabulary))
    selected_unit = ir.evidence_units[-1].model_dump(mode="json")
    retrieval_calls = []

    def retrieve(source_ir, snapshot, demand_menus, cards, scope):
        # DocumentIR keeps a private mutable unit lookup cache; compare frozen public content.
        assert source_ir.model_dump(mode="json") == ir.model_dump(mode="json")
        assert snapshot == metadata and cards == catalog
        assert all("anchor" not in menu for menu in demand_menus.values())
        retrieval_calls.append(scope)
        return {
            "case": {"scope_id": scope, "document_ref": "", "source_units": [selected_unit],
                     "content": {"text": selected_unit["text"]}},
            "retrieval": {"policy": "test-fixed-policy",
                          "ablation": {"source_input_changed": False}},
            "coverage": {"ir_unit_count": len(ir.evidence_units)},
        }

    monkeypatch.setattr(runner, "retrieve_scope", retrieve)
    args = SimpleNamespace(
        baseline=baseline, metadata_dir=metadata_dir, output=tmp_path / "prepared",
        ontology_dir=tmp_path / "ontology", model_path=tmp_path / "model",
        device="cuda:0", prepare_only=True, ner_only=False, prepared_tools=None,
    )
    return SimpleNamespace(args=args, ir=ir, metadata=metadata, catalog=catalog,
                           vocabulary=vocabulary, old=old, retrieval_calls=retrieval_calls)


def fake_ner(monkeypatch):
    configuration = []

    def extractor(path, **kwargs):
        configuration.append(kwargs)
        return object()

    def mentions(sources, *, groups, extractor, labels_per_batch):
        assert labels_per_batch == 1 and groups == {"entities": ["对象名称"]}
        return {
            "execution_status": "completed", "semantic_status": "not_checked",
            "fact_eligible": False,
            "spans": [{"ref": ref, "start": 0, "end": len(unit["text"]),
                       "text": unit["text"], "label": "对象名称", "score": 0.9,
                       "id": "mention-" + ref, "group": "entities"}
                      for ref, unit in sources.items()],
        }

    monkeypatch.setattr(runner.c_runner, "Gliner2Extractor", extractor)
    monkeypatch.setattr(runner.c_runner, "propose_mentions", mentions)
    return configuration


def prepare_ner_run(experiment, monkeypatch):
    configuration = fake_ner(monkeypatch)
    args = deepcopy(experiment.args)
    args.prepare_only, args.ner_only = False, True
    manifest = runner.execute(args)
    assert manifest["status"] == "ner_completed"
    assert configuration[0]["max_len"] == 160 and configuration[0]["word_splitter"] == "char"
    return args, manifest


def test_prepare_only_rebuilds_refs_from_full_ir_without_models_or_silver(experiment, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Model loading or inference during preparation")

    monkeypatch.setattr(runner.c_runner, "Gliner2Extractor", forbidden)
    monkeypatch.setattr(runner.c_runner, "invoke", forbidden)
    result = runner.execute(experiment.args)
    assert result["status"] == "prepared" and result["results"] == []
    assert result["max_model_calls"] == 16
    assert experiment.retrieval_calls == list(runner.SCOPES)
    for scope in runner.SCOPES:
        directory = experiment.args.output / "D" / scope
        assert read(directory / "sources.json")["u0"]["text"] == "全IR新命中"
        plan = read(directory / "plan/proposal.json")
        assert plan["inspect_evidence"] == [{"ref": "u0"}]
        assert plan["get_schema_card"] == [{"class_iri": ITEM}, {"class_iri": ITEM}]
        assert plan["propose_mentions"] == [{"group": "entities"}]
        assert read(directory / "case.json")["document_ref"] == "upload-test"
        assert "摘要独有" not in str(read(directory / "sources.json"))
    assert not (experiment.args.output / "reference.json").exists()
    assert not (experiment.args.output / "acceptance.json").exists()


@pytest.mark.parametrize("key,value", [
    ("source_sha256", "0" * 64), ("analysis_id", "another-analysis"),
    ("structure_hash", "1" * 64), ("input_hashes", {"ir": "other-ir-hash"}),
])
def test_cross_source_summary_fails_before_retrieval_and_is_recorded(experiment, key, value):
    directory = experiment.args.metadata_dir
    manifest = read(directory / "manifest.json")
    manifest[key] = value
    write(directory / "manifest.json", manifest)
    with pytest.raises(ValueError, match="summary_metadata_source_mismatch"):
        runner.execute(experiment.args)
    saved = read(experiment.args.output / "result.json")
    assert saved["status"] == "failed" and saved["failure_stage"] == "metadata_identity"
    assert experiment.retrieval_calls == []


def test_hash_updated_snapshot_still_must_match_section_tree(experiment):
    directory = experiment.args.metadata_dir
    snapshot = read(directory / "metadata.json")
    snapshot["metadata_snapshot"]["node_summaries"][0]["summary"] = "篡改摘要"
    write(directory / "metadata.json", snapshot)
    manifest = read(directory / "manifest.json")
    manifest["output_hashes"]["metadata.json"] = digest(directory / "metadata.json")
    write(directory / "manifest.json", manifest)
    with pytest.raises(ValueError, match="summary_snapshot_content_mismatch"):
        runner.load_metadata(directory, experiment.ir, digest(experiment.args.baseline / "ir.json"))


def test_baseline_class_plan_and_current_code_are_frozen(experiment):
    path = experiment.args.baseline / "C" / runner.SCOPES[0] / "plan/proposal.json"
    changed = read(path)
    changed["get_schema_card"].pop()
    write(path, changed)
    with pytest.raises(ValueError, match="baseline_demand_changed"):
        runner.execute(experiment.args)


def test_prepared_reuse_runs_no_ner_and_keeps_new_source_identity(experiment, monkeypatch):
    first, before = prepare_ner_run(experiment, monkeypatch)
    args = deepcopy(first)
    args.output = first.output.parent / "reused"
    args.prepared_tools = first.output
    monkeypatch.setattr(runner.c_runner, "Gliner2Extractor",
                        lambda *a, **kw: pytest.fail("NER repeated during prepared reuse"))
    after = runner.execute(args)
    assert after["status"] == "ner_completed" and after["ner_reused_from"] == before["run_id"]
    assert after["prepared_tool_hashes"] == before["prepared_tool_hashes"]
    assert after["scope_input_hashes"] == before["scope_input_hashes"]


@pytest.mark.parametrize("file", [
    "D/introduction/tools.json", "D/introduction/sources.json", "metadata/metadata.json",
    "code/frozen.py", "cases.json",
])
def test_prepared_actual_files_are_rechecked_not_only_claimed_hashes(experiment, monkeypatch, file):
    first, _ = prepare_ner_run(experiment, monkeypatch)
    path = first.output / file
    path.write_text(path.read_text() + "\n ")
    args = deepcopy(first)
    args.output, args.prepared_tools = first.output.parent / "reused", first.output
    with pytest.raises(ValueError, match="prepared_D_(actual_inputs_changed|tool_digest_mismatch)"):
        runner.execute(args)
    assert read(args.output / "result.json")["failure_stage"] == "prepared_tools_identity"
    assert not (args.output / "D/introduction/tools.json").exists()


@pytest.mark.parametrize("mutation", ["old-ref", "old-text"])
def test_old_source_spans_cannot_be_reused_even_if_tool_hash_is_updated(
    experiment, monkeypatch, mutation,
):
    first, _ = prepare_ner_run(experiment, monkeypatch)
    path = first.output / "D/introduction/tools.json"
    tools = read(path)
    span = tools["propose_mentions"]["spans"][0]
    span.update({"ref": "OLD-REF-MUST-NOT-SURVIVE"} if mutation == "old-ref"
                else {"text": "旧局部范围", "end": len("旧局部范围")})
    write(path, tools)
    manifest = read(first.output / "result.json")
    manifest["prepared_tool_hashes"]["introduction"] = digest(path)
    write(first.output / "result.json", manifest)
    args = deepcopy(first)
    args.output, args.prepared_tools = first.output.parent / "reused", first.output
    with pytest.raises(ValueError, match="prepared_mention_not_in_new_sources"):
        runner.execute(args)


def test_changed_vocabulary_fails_before_any_ner(experiment, monkeypatch):
    vocabulary = deepcopy(experiment.vocabulary)
    vocabulary["entries"]["对象名称"]["description"] = "changed description"
    monkeypatch.setattr(runner.c_runner, "build_extraction_vocabulary", lambda *a, **kw: vocabulary)
    with pytest.raises(ValueError, match="frozen_C_vocabulary_changed"):
        runner.execute(experiment.args)


def model_reply(calls, fail_stage=None):
    def invoke(client, directory, run_id, scope, stage, prompt, payload, schema, records):
        calls.append({"stage": stage, "payload": payload, "prompt": prompt})
        records.append({"stage": stage, "seconds": 0.1})
        if stage == fail_stage:
            raise ValueError("HTTP_400_no_automatic_truncation")
        if stage == "candidates":
            return {"subjects": {sid: {
                "anchor": {"ref": "", "quote": "", "span_id": ""},
                "attributes": {}, "relations": {},
            } for sid in ("document", "s1", "s2")}, "observations": []}
        return {"types": {}, "claims": {}, "observations": {}}

    return invoke


@pytest.mark.parametrize("failure,expected", [
    (None, ["candidates", "verification"]), ("candidates", ["candidates"]),
    ("verification", ["candidates", "verification"]),
])
def test_D_reuses_two_C_stages_without_retry_or_summary_in_prompt(
    experiment, monkeypatch, failure, expected,
):
    args, _ = prepare_ner_run(experiment, monkeypatch)
    calls = []
    monkeypatch.setattr(runner.c_runner, "invoke", model_reply(calls, failure))
    directory = args.output / "D/introduction"
    row = runner.run_candidate_case(None, directory, "test-D", experiment.catalog)
    assert [call["stage"] for call in calls] == expected
    assert row["arm"] == read(directory / "result.json")["arm"] == "D"
    assert row["status"] == ("complete" if failure is None else "failed")
    assert len(row["calls"]) == len(expected) <= 2
    assert calls[0]["prompt"] == runner.c_runner.CANDIDATE_PROMPT
    assert "全IR新命中" in str(calls[0]["payload"])
    assert "摘要独有" not in str(calls[0]["payload"])
    assert "rank_score" not in str(calls[0]["payload"])


def test_whole_experiment_has_exact_sixteen_request_budget(experiment, monkeypatch):
    args, manifest = prepare_ner_run(experiment, monkeypatch)
    calls = []
    monkeypatch.setattr(runner.c_runner, "invoke", model_reply(calls))
    runner.run_qwen(None, args.output, manifest, experiment.catalog)
    assert len(calls) == sum(len(row["calls"]) for row in manifest["results"]) == 16
    summary = runner.summarize(manifest["results"])
    assert summary["completed"] == 8
    assert summary["quality_scoring"] == "not_performed_on_old_scope_silver"
    assert "local_checklists_passed" not in summary
    with pytest.raises(ValueError, match="already_has_model_results"):
        runner.run_qwen(None, args.output, manifest, experiment.catalog)


def test_global_budget_gate_stops_before_another_scope_request(experiment, monkeypatch):
    args, manifest = prepare_ner_run(experiment, monkeypatch)
    calls = []
    monkeypatch.setattr(runner.c_runner, "invoke", model_reply(calls))
    monkeypatch.setattr(runner, "MAX_MODEL_CALLS", 3)
    with pytest.raises(ValueError, match="summary_experiment_budget_exhausted"):
        runner.run_qwen(None, args.output, manifest, experiment.catalog)
    assert len(calls) == 2


def test_failure_is_saved_without_overwriting_existing_output(experiment, monkeypatch):
    fake_ner(monkeypatch)
    monkeypatch.setattr(runner.c_runner, "Gliner2Extractor",
                        lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("load failed")))
    args = deepcopy(experiment.args)
    args.prepare_only = False
    with pytest.raises(RuntimeError, match="load failed"):
        runner.execute(args)
    path = args.output / "result.json"
    result = read(path)
    assert result["status"] == "failed" and result["failure_stage"] == "ner"
    previous = path.read_bytes()
    with pytest.raises(FileExistsError):
        runner.execute(args)
    assert path.read_bytes() == previous
