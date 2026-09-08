"""The recorded GLiNER load is the cold call, not a second availability check."""

from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document

from app.evaluation import cmc_benchmark, hierarchy_variant
from app.schemas.evidence import TaskBudget
from app.services.extraction import gliner_extractor, local_semantic_model, word_analysis
from app.services.extraction.extraction_tasks import ExtractionRun, GenericExtractionRunner
from tests.test_extraction.test_hierarchical_context import TestTokenizer


@pytest.fixture
def experiment(tmp_path, monkeypatch):
    prepared, output = tmp_path / "prepared", tmp_path / "run"
    prepared.mkdir()
    output.mkdir()
    doc = Document()
    doc.add_heading("设备信息", level=1)
    doc.add_paragraph("设备 R-101。")
    source = prepared / "source.docx"
    doc.save(source)
    analysis = word_analysis.analyze_word_core(source)
    schema = {cmc_benchmark.CMC_CLASS: {"label": "CMCReport"}}
    cmc_benchmark.write_json(prepared / "schema.json", schema)
    cmc_benchmark.write_json(
        prepared / "summaries.json",
        {
            "document_hash": analysis.ir.document_hash,
            "metadata": {},
            "generation_seconds": 5.0,
        },
    )
    manifest = {
        "source_filename": "source.docx",
        "analysis_id": analysis.ir.analysis_id,
        "document_hash": analysis.ir.document_hash,
        "runtime_hash": "fixture-runtime",
        "ontology_hash": "fixture-ontology",
        "priority_paths": [],
        "settings": {},
    }
    base = GenericExtractionRunner(
        schema,
        TestTokenizer(),
        Mock(side_effect=AssertionError("no live model calls")),
        model_identity="fixture",
        budget=TaskBudget(max_input_tokens=60000),
    )
    monkeypatch.setattr(local_semantic_model, "configured_generic_runner", lambda engine: base)
    monkeypatch.setattr(word_analysis, "analyze_word_core", lambda *args, **kwargs: analysis)
    monkeypatch.setattr(cmc_benchmark, "slot_snapshot", lambda: [])
    monkeypatch.setattr(cmc_benchmark, "emit", lambda *args, **kwargs: None)
    clock = [100.0]
    monkeypatch.setattr(cmc_benchmark.time, "perf_counter", lambda: clock[0])
    monkeypatch.setattr(hierarchy_variant, "perf_counter", lambda: clock[0])

    def extract_without_model(self, *args, **kwargs):
        clock[0] += 2.0
        return ExtractionRun(input_id="fixture", completion="complete")

    run = Mock(side_effect=extract_without_model)
    monkeypatch.setattr(
        hierarchy_variant.HierarchyHintRunner,
        "run",
        lambda self, *args, **kwargs: run(self, *args, **kwargs),
    )
    args = SimpleNamespace(
        mode="structure_summary_gliner",
        timeout=600,
        timeout_retries=0,
        resume=False,
        deadline_seconds=None,
        pause_after=24,
    )
    return SimpleNamespace(
        args=args,
        prepared=prepared,
        output=output,
        manifest=manifest,
        clock=clock,
        run=run,
    )


def test_first_gliner_load_is_measured_once_and_already_in_total_time(experiment, monkeypatch):
    env = experiment

    class LazyExtractor:
        calls = 0

        def is_available(self):
            self.calls += 1
            if self.calls == 1:
                env.clock[0] += 7.0
            return True

    extractor = LazyExtractor()
    monkeypatch.setattr(gliner_extractor, "get_gliner_extractor", lambda: extractor)
    cmc_benchmark._run_segment(env.args, env.prepared, env.manifest, env.output, [])
    result = cmc_benchmark.read_json(env.output / "result.json")
    assert extractor.calls == 1
    assert result["variant_statistics"]["gliner_load_wall_seconds"] == 7.0
    assert result["planner_seconds_this_segment"] == 7.0
    assert result["elapsed_seconds_this_segment"] == 9.0  # load 7 + extraction 2
    assert result["cold_metadata_accounted_seconds"] == 14.0  # + separate summaries 5


@pytest.mark.parametrize("disabled", [False, True])
def test_unavailable_gliner_never_produces_a_successful_run(experiment, monkeypatch, disabled):
    env = experiment
    available = Mock(return_value=False)
    extractor = None if disabled else SimpleNamespace(is_available=available)
    monkeypatch.setattr(gliner_extractor, "get_gliner_extractor", lambda: extractor)
    with pytest.raises(RuntimeError, match="GLiNER.*unavailable"):
        cmc_benchmark._run_segment(env.args, env.prepared, env.manifest, env.output, [])
    if not disabled:
        available.assert_called_once_with()
    env.run.assert_not_called()
    assert not (env.output / "result.json").exists()
