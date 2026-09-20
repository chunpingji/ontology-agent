"""Grouped record and attribute work retain exact source and paid-result ownership."""

import json
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document

from app.models.document_analysis import DocumentRunCandidate, DocumentRunResult
from app.services.document_analysis import current_state
from app.services.document_analysis.run_store import HeadConflict, RunNotFound
from app.services.extraction.ontology_guided.attribute_disambiguation import (
    extract_attribute_fields,
    field_payload,
)
from app.services.extraction.ontology_guided.context import ContextBindingRefs
from app.services.extraction.ontology_guided.contracts import GraphNode
from app.services.extraction.ontology_guided.current_work import protocol_result_ref
from app.services.extraction.ontology_guided.executor import TaskOutcome
from app.services.extraction.ontology_guided.record_discovery import RecordDiscoveryTask
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_record_discovery_current_state import (
    confirm,
    persist,
    publication,
    publish,
    record_protocol,
    reservation,
)

pytest_plugins = ["tests.test_extraction.test_tool_engine_resume"]


def grouped_protocol(current_run, tmp_path, *, attribute=False):
    store, run, token = current_run
    document = Document()
    document.add_heading("设备标识", 1)
    for text in ["装置甲。", "编号：A-001", "组外装置乙。"]:
        document.add_paragraph(text)
    path = tmp_path / "contextual.docx"
    document.save(path)
    ir = analyze_word_core(path).ir
    index = RecordIndex(ir)
    records = index.records
    field = extract_attribute_fields(index, property_labels=["编号"])[0]
    protocol = record_protocol()
    task = RecordDiscoveryTask.create(
        run_fingerprint=run.run_fingerprint, record_id=records[0].record_id,
        schema_card_id=protocol["task"]["schema_card_id"],
        analysis_scope_ref=protocol["task"]["analysis_scope_ref"], dependency_hash="f" * 64,
        source_record_ids=[record.record_id for record in records[:2]],
        reading_section_ids=[records[0].section_node_id],
        field_id=field.field_id if attribute else None,
    )
    protocol.update(task=task.model_dump(mode="json"), lineage_id=task.claim_lineage_id)
    protocol["base_target"]["task_id"] = task.task_id
    protocol["base_target"]["document_context"]["document_hash"] = ir.document_hash
    run.document_hash = ir.document_hash
    payload = {"analysis": ir.model_dump(mode="json")}
    store.update_artifact(
        run.recognition_run_id, run.owner_id, token, artifact_kind="metadata",
        expected_revision=0, artifact_hash=current_state.content_hash(payload),
        status="ready", payload=payload,
    )
    store.db.commit()
    return protocol, ir, records


def field_row(protocol, ir, *, outcome_ref=None):
    field = extract_attribute_fields(RecordIndex(ir), property_labels=["编号"])[0]
    return {"key": protocol["lineage_id"], "position": 0, "value": {
        "task": deepcopy(protocol["task"]), "status": "examined", "outcome_ref": outcome_ref,
        "attribute_field": field_payload(field),
        "attribute_options": [{"subject_ref": {"id": "root", "revision": 1},
                               "class_iri": "urn:Report", "predicate_iri": "urn:identifier"}],
        "attribute_signature": "d" * 64, "attribute_status": "unresolved",
        "disambiguation_attempts": 1,
    }}


def finish(current_run, protocol, monkeypatch, *, outcome=None, row=None):
    batch, outcome_ref = publication(current_run, protocol, monkeypatch)
    if outcome is not None:
        batch.outcome = outcome
        value = outcome.model_dump(mode="json")
        outcome_ref = protocol_result_ref(protocol["lineage_id"], "outcome", value)
        batch.protocol_state_changes[protocol["lineage_id"]]["outcome_ref"] = outcome_ref
        batch.protocol_result_changes = {outcome_ref: {
            "lineage_id": protocol["lineage_id"], "field": "outcome", "value": value,
        }}
        batch.work_changes.pop("nodes", None)
        batch.work_changes.pop("scheduler:record_relations", None)
        batch.graph = batch.graph.model_copy(update={"nodes": []})
    if row is not None:
        row = deepcopy(row)
        row["value"]["outcome_ref"] = outcome_ref
        batch.work_changes["record_discovery"] = {json.dumps(protocol["lineage_id"]): row}
    publish(current_run, batch)
    return deepcopy(batch.protocol_state_changes[protocol["lineage_id"]]), row


def reopen(current_run, protocol, saved, ir, unit, *, refs=()):
    store, run, token = current_run
    key = json.dumps(protocol["lineage_id"])
    pending = deepcopy(saved)
    feedback = "9" * 64
    anchor = ir.anchor(unit.evidence_id, 0, len(unit.text)).model_dump(mode="json")
    pending["value"].update(status="pending", feedback_hash=feedback, feedback_refs=list(refs),
                            feedback_source_refs=[anchor])
    current_state.persist_boundary(
        store, run, token, changes={"record_discovery": {key: pending}},
        fingerprint=run.run_fingerprint, expected_version=run.work_version,
    )
    updated = deepcopy(protocol)
    updated.update(
        assertion_generation=protocol["assertion_generation"] + 1,
        evidence_revision=protocol["evidence_revision"] + 1,
        context_hash="c" * 64, evidence_hash="c" * 64, stage="discovery",
        active_instructions="Resolve only the current field with its changed candidates.",
        turn_refs=[], completed_tool_results=[], materialized_refs={}, stage_input_items=[],
        discovery_ref=None, verification_ref=None, outcome_ref=None, record_feedback_hash=feedback,
    )
    updated["base_target"]["context_hash"] = updated["context_hash"]
    active = deepcopy(pending)
    active["value"].update(status="active", consumed_feedback_hash=feedback,
                           generation=updated["assertion_generation"], disambiguation_attempts=2)
    return updated, {"record_discovery": {key: active}}


def test_grouped_field_reopens_from_second_record_without_refunding_calls(
    current_run, tmp_path, monkeypatch,
):
    protocol, ir, records = grouped_protocol(current_run, tmp_path, attribute=True)
    persist(current_run, protocol)
    for _ in range(2):
        receipt = reservation(current_run, protocol)
        persist(current_run, protocol, receipts=[receipt])
        confirm(current_run, protocol)
    protocol, saved = finish(current_run, protocol, monkeypatch, row=field_row(protocol, ir))
    updated, changes = reopen(current_run, protocol, saved, ir, records[1].source_units[0])
    persist(current_run, updated, work_changes=changes)
    store, run, _ = current_run
    store.db.expire_all()
    restored = current_state.restore_calls(store, run, run.run_fingerprint)
    assert restored["protocols"][protocol["lineage_id"]] == updated
    assert restored["lineage_calls"] == {protocol["lineage_id"]: 2}
    assert restored["reservation_sequence"] == 2
    assert current_state.get_row(store, run, "work:record_discovery",
                                 json.dumps(protocol["lineage_id"])) == next(iter(
        changes["record_discovery"].values()
    ))


@pytest.mark.parametrize("authorized", [True, False])
def test_field_reopening_replaces_old_revision_only_with_owned_feedback(
    current_run, tmp_path, monkeypatch, authorized,
):
    protocol, ir, records = grouped_protocol(current_run, tmp_path, attribute=True)
    protocol["reference_context"] = {
        "version": 1, "entity_refs": [{"id": "record-entity", "revision": 1}],
    }
    persist(current_run, protocol)
    protocol, saved = finish(current_run, protocol, monkeypatch, row=field_row(protocol, ir))
    store, run, token = current_run
    previous = store.db.get(DocumentRunCandidate, (run.recognition_run_id, "record-entity", 1))
    changed = {**previous.payload, "revision": 2}
    store.put_candidate(
        run.recognition_run_id, run.owner_id, token, candidate_id="record-entity", revision=2,
        kind="entity", payload=changed, payload_hash=current_state.content_hash(changed),
        proof_refs=[], expected_head_revision=1, event_sequence=run.event_head,
    )
    store.db.commit()
    new_ref = {"id": "record-entity", "revision": 2}
    updated, changes = reopen(current_run, protocol, saved, ir, records[1].source_units[0],
                              refs=[new_ref] if authorized else [])
    updated["reference_context"] = {"version": 1, "entity_refs": [new_ref]}
    if not authorized:
        with pytest.raises(HeadConflict, match="unauthorized entity references"):
            persist(current_run, updated, work_changes=changes)
        return
    persist(current_run, updated, work_changes=changes)
    restored = current_state.restore_calls(store, run, run.run_fingerprint)
    assert restored["protocols"][protocol["lineage_id"]]["reference_context"] == {
        "version": 1, "entity_refs": [new_ref],
    }


@pytest.mark.parametrize("fault", [None, "extra", "missing"])
def test_field_reopening_uses_exact_published_candidate_set_when_it_shrinks(
    current_run, tmp_path, monkeypatch, fault,
):
    protocol, ir, records = grouped_protocol(current_run, tmp_path, attribute=True)
    store, run, token = current_run
    removed = GraphNode(entity_id="unrelated-now", revision=1, class_iri="urn:Entity",
                        class_label="Entity", label="B").model_dump(mode="json")
    store.put_candidate(
        run.recognition_run_id, run.owner_id, token, candidate_id="unrelated-now", revision=1,
        kind="entity", payload=removed, payload_hash=current_state.content_hash(removed),
        proof_refs=[], expected_head_revision=0, event_sequence=run.event_head or None,
    )
    store.db.commit()
    retained_ref = {"id": "record-entity", "revision": 1}
    removed_ref = {"id": "unrelated-now", "revision": 1}
    protocol["reference_context"] = {"version": 1, "entity_refs": [retained_ref, removed_ref]}
    persist(current_run, protocol)
    protocol, saved = finish(current_run, protocol, monkeypatch, row=field_row(protocol, ir))
    updated, changes = reopen(current_run, protocol, saved, ir, records[1].source_units[0],
                              refs=[retained_ref])
    after = [retained_ref, removed_ref] if fault == "extra" else [] if fault else [retained_ref]
    updated["reference_context"] = {"version": 1, "entity_refs": after}
    if fault:
        with pytest.raises(HeadConflict, match="unauthorized entity references"):
            persist(current_run, updated, work_changes=changes)
    else:
        persist(current_run, updated, work_changes=changes)
        assert current_state.restore_calls(store, run, run.run_fingerprint)["protocols"][
            protocol["lineage_id"]
        ]["reference_context"]["entity_refs"] == [retained_ref]


@pytest.mark.parametrize("fault", ["outside_record", "heading", "changed_membership"])
def test_reopening_cannot_expand_group_sources(current_run, tmp_path, monkeypatch, fault):
    protocol, ir, records = grouped_protocol(current_run, tmp_path, attribute=True)
    persist(current_run, protocol)
    protocol, saved = finish(current_run, protocol, monkeypatch, row=field_row(protocol, ir))
    unit = (next(unit for unit in ir.evidence_units if unit.kind == "heading")
            if fault == "heading" else records[-1].source_units[0])
    updated, changes = reopen(current_run, protocol, saved, ir, unit)
    if fault == "changed_membership":
        updated["task"]["source_record_ids"].append(records[-1].record_id)
    store, run, _ = current_run
    before = current_state.restore_calls(store, run, run.run_fingerprint)
    version = run.work_version
    with pytest.raises(HeadConflict):
        persist(current_run, updated, work_changes=changes)
    assert run.work_version == version
    assert current_state.restore_calls(store, run, run.run_fingerprint) == before


def authorization(protocol, ir, records):
    heading = next(unit for unit in ir.evidence_units if unit.kind == "heading")
    return {
        "task_id": protocol["task"]["task_id"],
        "ir_identity": {name: getattr(ir, name) for name in (
            "document_hash", "parser_version", "structure_hash",
        )},
        "context_policy_hash": "f" * 64,
        "record_ids": [record.record_id for record in records[:2]],
        "fragments": [{"record_id": record.record_id, "evidence_id": unit.evidence_id,
                       "span_start": 0, "span_end": len(unit.text), "role": "target",
                       "fact_eligible": True}
                      for record in records[:2] for unit in record.source_units] + [
            {"record_id": None, "evidence_id": heading.evidence_id, "span_start": 0,
             "span_end": len(heading.text), "role": "required_context", "fact_eligible": False},
        ],
        "bindings": {name: [] for name in ContextBindingRefs.model_fields},
    }


@pytest.mark.parametrize("fault", [None, "extra_record", "forged_record", "wrong_ir",
                                    "unknown_record", "unknown_section"])
def test_initial_authorization_grants_only_explicit_group_facts(current_run, tmp_path, fault):
    protocol, ir, records = grouped_protocol(current_run, tmp_path)
    allowed = authorization(protocol, ir, records)
    if fault == "extra_record":
        allowed["record_ids"].append(records[-1].record_id)
    elif fault == "forged_record":
        allowed["fragments"][0]["evidence_id"] = records[-1].source_units[0].evidence_id
    elif fault == "wrong_ir":
        allowed["ir_identity"]["structure_hash"] = "8" * 64
    elif fault == "unknown_record":
        protocol["task"]["source_record_ids"].append("unknown-record")
    elif fault == "unknown_section":
        protocol["task"]["reading_section_ids"].append("unknown-section")
    protocol["context_authorization"] = allowed
    if fault:
        with pytest.raises(HeadConflict, match="record.*source"):
            persist(current_run, protocol)
    else:
        persist(current_run, protocol)
        store, run, _ = current_run
        assert current_state.restore_calls(store, run, run.run_fingerprint)["protocols"][
            protocol["lineage_id"]
        ]["context_authorization"] == allowed


@pytest.mark.parametrize("reason", ["attribute_value_missing", "attribute_subject_missing",
                                    "ontology_property_missing"])
def test_zero_call_field_outcome_publishes_and_restores_without_fake_requests(
    current_run, tmp_path, monkeypatch, reason,
):
    protocol, ir, _ = grouped_protocol(current_run, tmp_path, attribute=True)
    persist(current_run, protocol)
    outcome = TaskOutcome(semantic_outcome="undetermined", complete=False,
                          reason_code=reason, reason="字段尚不能形成已验证属性。")
    saved = field_row(protocol, ir)
    saved["value"].update(reason_code=reason, disambiguation_attempts=0, status="incomplete")
    protocol, saved = finish(current_run, protocol, monkeypatch, outcome=outcome, row=saved)
    store, run, _ = current_run
    store.db.expire_all()
    restored = current_state.restore_calls(store, run, run.run_fingerprint)
    assert restored["lineage_calls"] == {}
    assert restored["reservations"] == []
    assert restored["protocols"][protocol["lineage_id"]]["discovery_ref"] is None
    assert restored["protocols"][protocol["lineage_id"]]["verification_ref"] is None
    assert current_state.load_protocol_result(store, run, protocol["lineage_id"],
                                              protocol["outcome_ref"], "outcome") == (
        outcome.model_dump(mode="json")
    )
    assert current_state.get_row(store, run, "work:record_discovery",
                                 json.dumps(protocol["lineage_id"])) == saved
    foreign = SimpleNamespace(recognition_run_id=run.recognition_run_id, owner_id="another-owner")
    with pytest.raises(RunNotFound):
        current_state.load_protocol_result(store, foreign, protocol["lineage_id"],
                                          protocol["outcome_ref"], "outcome")


def test_field_publication_rolls_back_status_and_result_together(
    current_run, tmp_path, monkeypatch,
):
    protocol, ir, _ = grouped_protocol(current_run, tmp_path, attribute=True)
    persist(current_run, protocol)
    batch, outcome_ref = publication(current_run, protocol, monkeypatch)
    saved = field_row(protocol, ir, outcome_ref=outcome_ref)
    batch.work_changes["record_discovery"] = {json.dumps(protocol["lineage_id"]): saved}

    def interrupted(*_args, **_kwargs):
        raise RuntimeError("field publication interrupted")

    monkeypatch.setattr(current_state, "publish_display", interrupted)
    store, run, _ = current_run
    with pytest.raises(RuntimeError, match="field publication interrupted"):
        publish(current_run, batch)
    assert current_state.get_row(store, run, "work:record_discovery",
                                 json.dumps(protocol["lineage_id"])) is None
    assert current_state.get_row(store, run, "calls:results", outcome_ref,
                                 model=DocumentRunResult) is None
    assert current_state.restore_calls(store, run, run.run_fingerprint)["protocols"][
        protocol["lineage_id"]
    ]["outcome_ref"] is None
