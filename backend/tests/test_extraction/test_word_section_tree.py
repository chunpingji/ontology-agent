"""018 — canonical Word blocks, explicit chapter tree and leaf page nodes."""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.extraction.document_annotator import annotate_word
from app.services.extraction.docx_structure import (
    PageBreakBlock,
    ParagraphBlock,
    TableBlock,
    parse_docx_structure,
)


def _manual_break(run) -> None:
    element = OxmlElement("w:br")
    element.set(qn("w:type"), "page")
    run._r.append(element)


def test_explicit_tree_handles_level_jumps_duplicates_and_root_content(tmp_path):
    doc = Document()
    doc.add_paragraph("文档前言正文。")
    doc.add_heading("产品信息", level=1)
    doc.add_paragraph("产品正文。")
    doc.add_heading("重复章节", level=3)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "属性"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "剂型"
    table.cell(1, 1).text = "片剂"
    doc.add_heading("重复章节", level=2)
    doc.add_paragraph("第二个重复章节正文。")
    doc.add_heading("工艺信息", level=1)
    path = tmp_path / "tree.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    root = structure.section_tree
    assert root is not None
    assert [child.heading for child in root.children] == ["产品信息", "工艺信息"]
    assert [child.heading for child in root.children[0].children] == [
        "重复章节", "重复章节",
    ]
    assert root.children[0].children[0].level == 3
    assert root.children[0].children[0].node_id != root.children[0].children[1].node_id
    assert root.paragraph_indices == [0]
    assert root.children[0].children[0].table_indices == [0]

    ordered_sections = []

    def visit(node):
        for child in node.children:
            ordered_sections.append(child.heading)
            visit(child)

    visit(root)
    assert ordered_sections == [section.heading for section in structure.sections if section.level]

    block_ids = [block.block_id for block in structure.blocks]
    assert len(block_ids) == len(set(block_ids))
    assert any(isinstance(block, ParagraphBlock) for block in structure.blocks)
    assert any(isinstance(block, TableBlock) for block in structure.blocks)
    assert root.source_range.start_block_id == "paragraph:0:0"
    assert root.source_range.end_block_id == root.children[-1].source_range.end_block_id


def test_leaf_pages_follow_manual_break_and_do_not_cross_sections(tmp_path):
    doc = Document()
    doc.add_heading("第一章", level=1)
    paragraph = doc.add_paragraph()
    first = paragraph.add_run("第一页内容")
    _manual_break(first)
    paragraph.add_run("第二页内容")
    doc.add_heading("第二章", level=1)
    doc.add_paragraph("同一物理页上的第二章内容。")
    path = tmp_path / "pages.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    root = structure.section_tree
    assert root is not None
    first_chapter, second_chapter = root.children
    assert structure.pagination.mode == "explicit_markers"
    assert [page.physical_page_number for page in first_chapter.pages] == [1, 2]
    assert [page.physical_page_number for page in second_chapter.pages] == [2]
    assert "paragraph:1:1" in first_chapter.pages[1].block_ids
    assert all(
        not block_id.startswith("paragraph:2:")
        for page in first_chapter.pages
        for block_id in page.block_ids
    )


def test_no_markers_create_virtual_leaf_page_without_physical_number(tmp_path):
    doc = Document()
    doc.add_heading("章节", level=1)
    doc.add_paragraph("正文。")
    path = tmp_path / "fallback.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    leaf = structure.section_tree.children[0]
    assert structure.pagination.mode == "single_page_fallback"
    assert structure.pagination.physical_page_numbers_available is False
    assert len(leaf.pages) == 1
    assert leaf.pages[0].physical_page_number is None
    assert leaf.pages[0].node_id.endswith(":page:1")


def test_tiptap_nodes_carry_canonical_source_coordinates(tmp_path):
    doc = Document()
    doc.add_heading("章节", level=1)
    doc.add_paragraph("正文。")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "表头"
    table.cell(1, 0).text = "数据"
    path = tmp_path / "coordinates.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    content, _, _, _ = annotate_word(
        path, engine=None, structure_only=True, structure=structure
    )
    heading, paragraph, table_node = [
        node
        for node in content["content"]
        if node["type"] in {"heading", "paragraph", "table"}
    ]
    assert heading["attrs"]["sourceBlockId"] == "paragraph:0:0"
    assert heading["attrs"]["sectionNodeId"] == "section:0"
    assert paragraph["attrs"]["sourceBlockId"] == "paragraph:1:0"
    assert table_node["attrs"]["sourceBlockId"] == "table:0"
    assert table_node["attrs"]["sourceTableIndex"] == 0


def test_consecutive_breaks_and_line_break_keep_preview_fragment_ids(tmp_path):
    doc = Document()
    doc.add_heading("章节", level=1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("第一行")
    run.add_break()
    _manual_break(run)
    _manual_break(run)
    paragraph.add_run("第三页内容")
    path = tmp_path / "consecutive-breaks.docx"
    doc.save(path)

    structure = parse_docx_structure(path)
    paragraph_blocks = [
        block
        for block in structure.blocks
        if isinstance(block, ParagraphBlock) and block.paragraph_index == 1
    ]
    page_breaks = [
        block for block in structure.blocks if isinstance(block, PageBreakBlock)
    ]
    assert [block.block_id for block in paragraph_blocks] == [
        "paragraph:1:0",
        "paragraph:1:2",
    ]
    assert paragraph_blocks[0].text == "第一行\n"
    assert paragraph_blocks[1].text == "第三页内容"
    assert len(page_breaks) == 2

    content, _, _, _ = annotate_word(
        path, engine=None, structure_only=True, structure=structure
    )
    preview_ids = [
        node.get("attrs", {}).get("sourceBlockId")
        for node in content["content"]
        if node["type"] in {"heading", "paragraph"}
    ]
    assert "paragraph:1:0" in preview_ids
    assert "paragraph:1:2" in preview_ids
    assert [node["type"] for node in content["content"]].count("pageBreak") == 2
