"""No-model equivalence, complete logical coverage and bounded task creation."""

from __future__ import annotations

import copy
import json

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context_records import (
    field_group_context,
    named_object_context,
)
from app.services.extraction.ontology_guided.contracts import (
    PlannedRecord,
    RecallEntry,
    RetrievalPlan,
    SubjectRef,
)
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.model_adapter import LocalModelRecognitionAdapter
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import FrontierScheduler, RecognitionTask
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_ontology_guided_core import DESCRIBES, REPORT, ontology
from tests.test_extraction.test_semantic_graph_closure import (
    APPEARANCE as TEMPLATE_APPEARANCE,
)
from tests.test_extraction.test_semantic_graph_closure import (
    DESCRIBES as TEMPLATE_DESCRIBES,
)
from tests.test_extraction.test_semantic_graph_closure import (
    ROOT,
    _analysis,
    _ontology,
    _respond,
)
from tests.test_extraction.test_semantic_scheduler import task


def make_plan(subject="root", predicate="urn:p", count=30, kind="relationship"):
    records = [PlannedRecord(
        record_id=f"record-{position}", phase=1 if position % 3 else 2,
        section_node_id=f"section-{position % 4}", record_kind="paragraph",
    ) for position in range(count)]
    # Freeze a non-source traversal order to exercise the independent source index.
    records = records[::2] + records[1::2]
    source_ids = [f"record-{position}" for position in range(count)]
    return RetrievalPlan(
        plan_id=f"plan:{subject}:{predicate}",
        subject=SubjectRef(entity_id=subject, revision=1, class_iri="urn:Class",
                           is_document_root=subject == "root"),
        predicate_iri=predicate, predicate_kind=kind, metadata_snapshot_id="metadata",
        ontology_hash="ontology", records=records,
        ledger={record.record_id: RecallEntry(record_id=record.record_id, phase=record.phase)
                for record in records},
        frozen_record_ids=source_ids, frozen_record_hash=evidence_hash(source_ids),
    )


def enqueue_plan(scheduler, plan, *, lazy=True, template_priority=False, hop=None):
    positions = {identity: index for index, identity in enumerate(plan.frozen_record_ids)}
    hop = (0 if plan.subject.is_document_root else 1) if hop is None else hop
    options = dict(hop=hop, dependency_hash="dependency",
                   root_branch=plan.subject.is_document_root)
    if lazy:
        return scheduler.enqueue_plan(
            plan, source_positions=positions, template_priority=template_priority, **options,
        )
    for record in plan.records:
        scheduler.enqueue(RecognitionTask.create(
            subject=plan.subject, predicate_iri=plan.predicate_iri,
            predicate_kind=plan.predicate_kind, record_id=record.record_id,
            phase=record.phase, section_node_id=record.section_node_id,
            source_position=positions[record.record_id],
            **{key: value for key, value in options.items() if key != "root_branch"},
        ), root_branch=options["root_branch"])


def test_peek_never_copies_scheduler_and_detaches_only_the_selected_task():
    class CopyForbidden:
        def __deepcopy__(self, memo):
            raise AssertionError("peek copied unrelated scheduler state")

    scheduler = FrontierScheduler()
    scheduler.unrelated_state = CopyForbidden()
    original = task()
    scheduler.enqueue(original, root_branch=True)
    before = scheduler.snapshot()
    upcoming = scheduler.peek_task()
    assert upcoming == original
    upcoming.subject.class_iri = "urn:changed-by-caller"
    assert scheduler.snapshot() == before
    assert scheduler.next_task() == original


@pytest.mark.parametrize("phase_interleaving", [True, False])
def test_logical_frontier_dispatch_and_restoration_equal_eager_with_ranks_retries_and_exclusions(
    phase_interleaving,
):
    eager = FrontierScheduler(max_tasks=1000, phase_interleaving=phase_interleaving)
    lazy = FrontierScheduler(max_tasks=1000, phase_interleaving=phase_interleaving)
    plans = [make_plan(subject, predicate, kind=kind)
             for subject in ("root", "child-a", "child-b")
             for predicate, kind in (("urn:p", "relationship"), ("urn:a", "property"))]
    for plan in plans:
        enqueue_plan(eager, plan, lazy=False)
        enqueue_plan(lazy, plan)
    for scheduler in (eager, lazy):
        for plan in plans:
            scheduler.reorder_slot(plan.subject, plan.predicate_iri,
                                   list(reversed(plan.frozen_record_ids[:12])), epoch_seq=1)
            scheduler.reorder_slot(plan.subject, plan.predicate_iri,
                                   list(reversed(plan.frozen_record_ids[12:])), epoch_seq=2)
        scheduler.enqueue(task(record="recheck", retry="evidence_recheck"), root_branch=True)
    logical_before = lazy.snapshot()
    assert all(entry.coverage_state == "unattempted"
               for plan in plans for entry in plan.ledger.values())
    position = 0
    dispatched = []
    while eager.pending and position < 1000:
        excluded = {("child-a", 1, "urn:p")} if position % 3 else set()
        assert lazy.peek_task(excluded) == eager.peek_task(excluded)
        actual = lazy.next_task(excluded)
        assert actual == eager.next_task(excluded)
        if actual is not None:
            dispatched.append(actual.task_id)
        assert lazy.pending == eager.pending
        if position % 17 == 0:
            restored = FrontierScheduler.from_snapshot(lazy.snapshot(), max_tasks=1000)
            assert restored.snapshot() == lazy.snapshot()
            lazy = restored
        position += 1
    assert lazy.pending == eager.pending == 0
    assert len(dispatched) == len(set(dispatched)) == 181
    assert logical_before["logical_frontiers"][0]["consumed"] == []


def test_ten_thousand_logical_opportunities_create_only_selected_task_models(monkeypatch):
    plan = make_plan(count=10000)
    created = 0
    original = RecognitionTask.create.__func__

    def counted(cls, **values):
        nonlocal created
        created += 1
        return original(cls, **values)

    monkeypatch.setattr(RecognitionTask, "create", classmethod(counted))
    scheduler = FrontierScheduler(max_tasks=20000)
    assert enqueue_plan(scheduler, plan) == 10000
    assert created == 0
    assert scheduler.pending == 10000
    assert scheduler.pending_slots == {("root", 1, "urn:p")}
    snapshot = scheduler.snapshot()
    assert snapshot["root"] == snapshot["child"] == []
    restored = FrontierScheduler.from_snapshot(snapshot, max_tasks=20000)
    assert created == 0
    assert restored.pending == 10000
    before = restored.snapshot()
    assert restored.peek_slot() == (plan.subject, plan.predicate_iri)
    assert created == 0
    assert restored.peek_task() is not None
    assert created == 1
    assert restored.snapshot() == before
    assert restored.next_task() is not None
    assert created == 2
    assert restored.pending == 9999
    assert sum(entry.coverage_state == "unattempted" for entry in plan.ledger.values()) == 10000


def test_compact_frontier_reduces_serialized_queue_without_reducing_opportunities():
    plan = make_plan(count=1000)
    eager, lazy = FrontierScheduler(), FrontierScheduler()
    enqueue_plan(eager, plan, lazy=False)
    enqueue_plan(lazy, plan)
    compact_bytes = len(json.dumps(lazy.snapshot()).encode())
    eager_bytes = len(json.dumps(eager.snapshot()).encode())
    assert lazy.pending == eager.pending == 1000
    assert compact_bytes < eager_bytes * 0.2


def test_logical_budget_and_hop_stop_keep_complete_unattempted_frontier():
    plan = make_plan(count=100)
    scheduler = FrontierScheduler(max_tasks=1)
    enqueue_plan(scheduler, plan)
    assert scheduler.next_task() is not None
    assert scheduler.peek_task() is None
    assert scheduler.pending == 99
    assert scheduler.next_task() is None
    assert scheduler.unexplored_frontier == [{
        "slot": ["root", 1, "urn:p"], "reason": "max_tasks", "logical_records": 99,
    }]
    restored = FrontierScheduler.from_snapshot(scheduler.snapshot(), max_tasks=1)
    assert restored.snapshot() == scheduler.snapshot()
    assert all(entry.coverage_state == "unattempted" for entry in plan.ledger.values())
    exceeded = FrontierScheduler(max_hops=1)
    assert enqueue_plan(exceeded, plan, hop=2) == 0
    assert exceeded.unexplored_frontier[0]["logical_records"] == 100
    assert exceeded.next_task() is None


@pytest.mark.parametrize("corrupt", ["record", "hash", "consumed", "arrival", "version"])
def test_corrupt_logical_frontier_is_rejected(corrupt):
    scheduler = FrontierScheduler()
    enqueue_plan(scheduler, make_plan())
    raw = copy.deepcopy(scheduler.snapshot())
    frontier = raw["logical_frontiers"][0]
    if corrupt == "record":
        frontier["records"][0][0] = "different-record"
    elif corrupt == "hash":
        frontier["content_hash"] = "wrong-hash"
    elif corrupt == "consumed":
        frontier["consumed"] = [1000]
    elif corrupt == "arrival":
        raw["next_arrival"] = 1
    else:
        raw["schema_version"] = 999
    with pytest.raises(ValueError):
        FrontierScheduler.from_snapshot(raw)


def test_template_priority_is_separate_frozen_strategy_with_regular_and_tail_opportunities():
    scheduler = FrontierScheduler(template_interleaving=True)
    priority, ordinary = make_plan(predicate="urn:template"), make_plan(predicate="urn:ordinary")
    enqueue_plan(scheduler, priority)  # A plan alone confers no priority.
    enqueue_plan(scheduler, ordinary)
    scheduler.prioritize_slot(priority.subject, priority.predicate_iri)
    for plan in (priority, ordinary):
        scheduler.reorder_slot(plan.subject, plan.predicate_iri,
                               list(reversed(plan.frozen_record_ids)), epoch_seq=1)
    first = [scheduler.next_task() for _ in range(12)]
    assert [item.predicate_iri for item in first] == ["urn:template", "urn:template",
                                                    "urn:ordinary"] * 4
    restored = FrontierScheduler.from_snapshot(scheduler.snapshot())
    remaining = []
    while scheduler.pending:
        assert restored.peek_task() == scheduler.peek_task()
        actual = scheduler.next_task()
        assert restored.next_task() == actual
        remaining.append(actual)
    assert {item.record_id for item in [*first, *remaining]} == set(priority.frozen_record_ids)
    assert scheduler.snapshot() == restored.snapshot()


def test_logical_subject_discard_preserves_dedup_and_other_frontiers():
    scheduler = FrontierScheduler()
    plan = make_plan()
    enqueue_plan(scheduler, plan)
    enqueue_plan(scheduler, make_plan("child"))
    scheduler.next_task()
    assert len(scheduler.discard_subject(plan.subject)) == 30
    assert scheduler.pending == 29
    assert enqueue_plan(scheduler, plan) == 0
    restored = FrontierScheduler.from_snapshot(scheduler.snapshot())
    assert restored.pending == scheduler.pending
    assert restored.next_task() == scheduler.next_task()


def test_record_indices_preserve_field_owners_context_and_source_order(tmp_path):
    document = Document()
    document.add_heading("产品概况", 1)
    document.add_heading("产品", 2)
    document.add_paragraph("产品甲")
    document.add_heading("外观", 2)
    document.add_paragraph("颜色：白色")
    document.add_paragraph("产品名称：产品甲")
    path = tmp_path / "indexed-records.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    first = index.records[0]
    name = index.records[-1]
    expected_units, expected_owners = field_group_context(index, name.record_id, "产品甲")
    snapshot = ontology()
    predicate = next(item for item in snapshot.classes[REPORT].declared_relationships
                     if item.iri == DESCRIBES)
    expected_named = named_object_context(index, first.record_id, predicate, snapshot)
    assert expected_named
    assert [index.ir.resolve(anchor) for anchor in expected_owners] == ["产品名称：产品甲"]
    assert index.record_positions == {record.record_id: position
                                      for position, record in enumerate(index.records)}

    class NoScan:
        def __iter__(self):
            raise AssertionError("context lookup scanned the entire record/group collection")

    index.records = index.field_groups = NoScan()
    assert field_group_context(index, name.record_id, "产品甲") == (expected_units, expected_owners)
    assert named_object_context(index, first.record_id, predicate, snapshot) == expected_named


@pytest.mark.parametrize(
    "initial_lazy,initial_template", [(False, False), (True, False), (True, True)],
)
def test_executor_resume_keeps_frozen_frontier_and_template_strategy_when_defaults_change(
    tmp_path, monkeypatch, initial_lazy, initial_template,
):
    analysis, snapshot = _analysis(tmp_path), _ontology()
    calls = []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        if request["stage"] == "discovery":
            calls.append((request["subject"], request["predicate"]["iri"], request["fragments"]))
        return _respond(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    arguments = dict(
        recognition_run_id="frozen-frontier-strategy", run_fingerprint="frozen-strategy",
        ir=analysis.ir,
        metadata=prepare_metadata(analysis.ir,
                                  section_tree=analysis.structure.section_tree.to_dict(),
                                  summary_version="frozen-strategy"),
        root_class_iri=ROOT, root_class_label="报告", filename="source.docx",
    )

    def executor(*, lazy, template, progress_hook=None):
        return OntologyGuidedExecutor(
            ontology=snapshot, engine=object(), max_tasks=200,
            lazy_frontier=lazy, template_interleaving=template, progress_hook=progress_hook,
            adapter=LocalModelRecognitionAdapter(object(), model_identity="frozen-strategy"),
            priority_paths=[(TEMPLATE_DESCRIBES, TEMPLATE_APPEARANCE)],
        )

    baseline = executor(lazy=initial_lazy, template=initial_template).run(**arguments)
    expected = list(calls)
    calls.clear()
    batches = []
    executor(
        lazy=initial_lazy, template=initial_template,
        progress_hook=lambda stage: stage != "after_model" or len(calls) < 2,
    ).run(**arguments, batch_hook=batches.append)
    assert 0 < len(calls) < len(expected)
    checkpoint = batches[-1].model_dump(mode="json")
    assert checkpoint["frontier"].get("schema_version", 1) == (2 if initial_lazy else 1)
    resumed = executor(lazy=not initial_lazy, template=not initial_template).run(
        **arguments, resume_state=checkpoint,
    )
    assert calls == expected
    assert resumed.graph.progress.records_examined == baseline.graph.progress.records_examined
