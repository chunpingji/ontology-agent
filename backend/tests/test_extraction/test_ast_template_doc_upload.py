"""AST 模板上传端点对遗留 .doc 的受理（后端转 .docx）。

覆盖 parse-sample / {id}/sample / {id}/default-source 三个上传面：
  - 真实 .doc（需 LibreOffice+Writer）经后端转换后成功解析/落盘为 .docx；
  - 不支持的后缀（.txt）仍 422（回归保护）。

真实 .doc 相关用例在 soffice/Writer 不可用时 skip；422 用例无 soffice 依赖，恒跑。
Fixtures(client/db/analyst_headers) 来自 tests/conftest.py。
"""

from __future__ import annotations

import io
import subprocess

import pytest

from app.services.extraction import doc_converter
from app.services.reporting.ast_template import load_default_template


def _seed_template(db, *, name="DocUp", version="v1"):
    from app.models.extraction import AstTemplate

    row = AstTemplate(
        name=name, version=version, doc_no="DOC-1",
        schema_json=load_default_template().model_dump(),
        status="draft", created_by="creator",
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _real_doc_bytes(tmp_path) -> bytes | None:
    """造一个真实遗留 .doc 的字节；soffice/Writer 不可用则 None（skip 依据）。"""
    if not doc_converter.soffice_available():
        return None
    import docx

    d = docx.Document()
    d.add_heading("评估对象", level=1)
    d.add_paragraph("本品为 XX 注射液，属注射剂。")
    src = tmp_path / "seed.docx"
    d.save(str(src))
    profile = tmp_path / "prof"
    try:
        proc = subprocess.run(
            [
                doc_converter._SOFFICE, "--headless", "--norestore",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to", "doc", "--outdir", str(tmp_path), str(src),
            ],
            capture_output=True, timeout=120,
        )
    except Exception:
        return None
    out = tmp_path / "seed.doc"
    if proc.returncode != 0 or not out.is_file():
        return None
    return out.read_bytes()


_WORD_DOC_MIME = "application/msword"


class TestParseSampleDoc:
    def test_txt_rejected(self, client, analyst_headers):
        resp = client.post(
            "/api/ast-templates/parse-sample",
            headers=analyst_headers,
            files={"file": ("x.txt", b"hello", "text/plain")},
        )
        assert resp.status_code == 422

    def test_real_doc_parsed(self, client, analyst_headers, tmp_path):
        data = _real_doc_bytes(tmp_path)
        if data is None:
            pytest.skip("LibreOffice+Writer 不可用")
        resp = client.post(
            "/api/ast-templates/parse-sample",
            headers=analyst_headers,
            files={"file": ("legacy.doc", data, _WORD_DOC_MIME)},
        )
        assert resp.status_code == 200
        assert "注射液" in resp.json()["plain_text"]


class TestReplaceSampleDoc:
    def test_txt_rejected(self, client, db, analyst_headers):
        row = _seed_template(db)
        resp = client.post(
            f"/api/ast-templates/{row.id}/sample",
            headers=analyst_headers,
            files={"file": ("x.txt", b"hello", "text/plain")},
        )
        assert resp.status_code == 422

    def test_real_doc_converted_and_persisted(self, client, db, analyst_headers, tmp_path):
        data = _real_doc_bytes(tmp_path)
        if data is None:
            pytest.skip("LibreOffice+Writer 不可用")
        row = _seed_template(db)
        resp = client.post(
            f"/api/ast-templates/{row.id}/sample",
            headers=analyst_headers,
            files={"file": ("legacy.doc", data, _WORD_DOC_MIME)},
        )
        assert resp.status_code == 200
        assert "注射液" in resp.json()["plain_text"]
        db.refresh(row)
        # 持久化的样例始终是 .docx（遗留 .doc 已转换）。
        assert (row.sample_docx_path or "").endswith(".docx")
        assert "注射液" in (row.sample_text or "")


class TestDefaultSourceCopyOnWrite:
    def test_convert_failure_preserves_old_source(self, client, db, analyst_headers, monkeypatch):
        """转换失败（422）时，既有默认源文件与 DB 引用必须原样保留（copy-on-write 回归保护）。

        无需真实 soffice：直接令转换抛 DocConversionError。旧实现「先删后写」会在此丢失旧源。
        """
        from pathlib import Path as _P

        from app.api import ast_templates
        from app.services.extraction import doc_converter

        row = _seed_template(db)
        # 预置一份既有默认源（.docx）落盘并挂到行上。
        old_path = _P(ast_templates._UPLOADS) / f"tpl_{row.id}_source_existing.docx"
        old_path.parent.mkdir(parents=True, exist_ok=True)
        old_path.write_bytes(b"OLD-SOURCE-BYTES")
        row.default_source_path = str(old_path)
        row.default_source_filename = "old.docx"
        db.commit()

        def _boom(_src):
            raise doc_converter.DocConversionError("模拟转换失败")

        monkeypatch.setattr(doc_converter, "ensure_docx", _boom)

        resp = client.post(
            f"/api/ast-templates/{row.id}/default-source",
            headers=analyst_headers,
            files={"file": ("new.doc", b"\xd0\xcf\x11\xe0new", _WORD_DOC_MIME)},
        )
        assert resp.status_code == 422
        # 旧源文件与 DB 引用未受影响。
        assert old_path.is_file()
        assert old_path.read_bytes() == b"OLD-SOURCE-BYTES"
        db.refresh(row)
        assert row.default_source_path == str(old_path)
        assert row.default_source_filename == "old.docx"
        # 转换失败不得在 data/uploads 里留下本次上传的孤儿文件（仅既有那份存在）。
        leftovers = list(old_path.parent.glob(f"tpl_{row.id}_source_*"))
        assert leftovers == [old_path]

    def test_unsupported_suffix_rejected(self, client, db, analyst_headers):
        row = _seed_template(db)
        resp = client.post(
            f"/api/ast-templates/{row.id}/default-source",
            headers=analyst_headers,
            files={"file": ("x.pdf", b"%PDF-1.4", "application/pdf")},
        )
        assert resp.status_code == 422


class TestDefaultSourceDoc:
    def test_real_doc_stored_as_docx(self, client, db, analyst_headers, tmp_path):
        data = _real_doc_bytes(tmp_path)
        if data is None:
            pytest.skip("LibreOffice+Writer 不可用")
        row = _seed_template(db)
        resp = client.post(
            f"/api/ast-templates/{row.id}/default-source",
            headers=analyst_headers,
            files={"file": ("origin.doc", data, _WORD_DOC_MIME)},
        )
        assert resp.status_code == 200
        # 展示名保留用户原始 .doc；落盘路径为转换后的 .docx（下游标注链只认 .docx）。
        assert resp.json()["default_source_filename"] == "origin.doc"
        db.refresh(row)
        assert (row.default_source_path or "").endswith(".docx")
