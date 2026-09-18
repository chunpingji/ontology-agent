"""Style lookup must scale without changing preview or evidence identities."""

from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.styles import CT_Styles
from docx.shared import Pt

from app.services.extraction import docx_reader
from app.services.extraction.document_annotator import parse_word_to_tiptap


def test_cached_read_matches_uncached_structure_styles_and_evidence(tmp_path, monkeypatch):
    document = Document()
    document.styles["Normal"].font.bold = True
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("评估报告", 1)
    document.add_paragraph("产品信息")
    custom = document.styles.add_style("ReportDetails", WD_STYLE_TYPE.PARAGRAPH)
    custom.font.italic = True
    document.add_paragraph("评估内容", style=custom)
    # Missing style and wrong-type style both retain python-docx's default fallback.
    document.add_paragraph("缺失样式")._p.get_or_add_pPr().get_or_add_pStyle().val = "Missing"
    document.styles.add_style("CharacterOnly", WD_STYLE_TYPE.CHARACTER)
    mismatch = document.add_paragraph("类型不匹配")
    mismatch._p.get_or_add_pPr().get_or_add_pStyle().val = "CharacterOnly"
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "合并单元格"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套段落"
    path = tmp_path / "sample.docx"
    document.save(path)

    cached = parse_word_to_tiptap(path)
    monkeypatch.setattr(docx_reader, "load_docx_for_reading", lambda path: Document(str(path)))
    uncached = parse_word_to_tiptap(path)
    assert cached == uncached


def test_large_style_catalog_is_not_rescanned_for_every_paragraph(tmp_path, monkeypatch):
    document = Document()
    for index in range(300):
        style = OxmlElement("w:style")
        style.set(qn("w:type"), "paragraph")
        style.set(qn("w:styleId"), f"ImportedStyle{index}")
        document.styles.element.append(style)
    paragraph = document.add_paragraph("正文内容")._p
    for _ in range(99):
        paragraph.addnext(deepcopy(paragraph))
    path = tmp_path / "many-styles.docx"
    document.save(path)

    scans = []
    original = CT_Styles.default_for

    def counted(styles, style_type):
        scans.append(style_type)
        return original(styles, style_type)

    monkeypatch.setattr(CT_Styles, "default_for", counted)
    parsed = parse_word_to_tiptap(path)
    assert len(parsed["analysis"]["evidence_units"]) == 100
    # Two fresh reads (canonical structure and styled preview), not 900 scans.
    assert len(scans) <= 2


def test_style_cache_is_scoped_to_each_read_not_filename_or_user(tmp_path):
    path = tmp_path / "same-name.docx"
    first = Document()
    first.styles["Normal"].font.bold = True
    first.add_paragraph("正文")
    first.save(path)
    loaded_first = docx_reader.load_docx_for_reading(path)
    assert loaded_first.paragraphs[0].style.font.bold is True

    second = Document()
    second.styles["Normal"].font.bold = False
    second.add_paragraph("正文")
    second.save(path)
    loaded_second = docx_reader.load_docx_for_reading(path)
    assert loaded_second.paragraphs[0].style.font.bold is False
    assert loaded_first.paragraphs[0].style.font.bold is True
