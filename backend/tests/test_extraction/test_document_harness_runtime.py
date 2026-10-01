"""Independent harness persistence, ownership and API routing regressions."""

from __future__ import annotations

import builtins
import json
from uuid import uuid4

import pytest
from docx import Document
from rdflib import OWL, RDF, Graph, URIRef
from sqlalchemy import select

from app.api import document_analysis
from app.config import settings
from app.models.document_analysis import (
    DocumentRunCurrentState,
    DocumentRunRequest,
    DocumentRunResult,
)
from app.services.document_analysis.application import (
    DocumentAnalysisApplication,
    DocumentAnalysisError,
)
from app.services.document_analysis.dispatcher import _execute
from app.services.document_analysis.run_store import DocumentAnalysisRunStore, content_hash
from app.services.document_harness.application import ENGINE
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.projection import graph_response, source_response
from app.services.document_harness.ranking import default_policy
from app.services.document_harness.runtime import (
    HarnessCallFailed,
    Repository,
    initialize_state,
    read_rows,
)
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.local_client import StructuredModelError
from app.services.llm.model_runtime import ModelCancelled


def _catalog():
    graph = Graph()
    graph.add((URIRef("urn:example:Report"), RDF.type, OWL.Class))
    return catalog_from_graph(graph, "urn:example:Report").model_dump(mode="json")


def _run(db, *, owner="analyst", engine=ENGINE, document_hash="d" * 64, catalog=None):
    key = uuid4().hex
    catalog = catalog or _catalog()
    store = DocumentAnalysisRunStore(db)
    run, _ = store.create_run_with_source(
        owner_id=owner, request_key=key, filename="report.docx", document_hash=document_hash,
        root_class_iri="urn:example:Report", root_class_label="Report",
        ontology_snapshot_hash=content_hash(catalog), source_artifact_id="source:" + key,
        source_storage_uri="original.docx", source_media_type="application/docx",
        source_size_bytes=1, source_payload={"engine": engine, "policy": {
            "max_input_tokens": 32768, "max_context_tokens": 65536,
            "card_ranking": default_policy(), "protocol": "document-harness-v6",
            "execution_policy": {"flow": "local_reading", "reading_concurrency": 2},
        }},
        ontology_artifact_id="schema:" + key, ontology_payload=catalog,
        progress={"engine": engine},
    )
    initialize_state(db, run, catalog)
    token = store.claim(run.recognition_run_id, owner, actor="test", worker_id="test")
    db.commit()
    return run, token


def _discovery():
    return {"entities": [], "document_field_ids": [], "document_source_fields": [],
            "unowned_fields": [], "relation_hints": [], "reference_cues": [], "complete": True}


def _requests(db, run):
    repo = Repository(db, run, "")
    return {row.payload["key"]: repo.get_call_record(row.payload["key"])
            for row in db.scalars(select(DocumentRunRequest).where(
                DocumentRunRequest.recognition_run_id == run.recognition_run_id,
                DocumentRunRequest.domain == "harness:calls",
            ))}


def test_paid_response_is_saved_once_and_reused_before_application(db, monkeypatch):
    from app.services.document_harness import model
    from app.services.llm.model_runtime import runtime

    run, token = _run(db)
    calls = []
    answer = {"output": _discovery(), "usage": {"input_tokens": 9, "output_tokens": 2},
              "seconds": 1.5, "raw_response": {"id": "paid"}, "error": None}
    def call(*args):
        assert runtime.get()["stage"] == "discover"
        calls.append(args)
        return answer

    monkeypatch.setattr(model, "call_model", call)
    repository = Repository(db, run, token)
    assert repository.invoke("discover", {"source": "original"}, {}) == answer["output"]
    # A fresh repository simulates pause/restart before the controller applies it.
    assert Repository(db, run, token).invoke("discover", {"source": "original"}, {}) == _discovery()
    assert len(calls) == 1
    results = read_rows(db, run, model=DocumentRunResult)["calls"]
    assert next(iter(results.values()))["raw_response"] == {"id": "paid"}
    assert next(iter(_requests(db, run).values()))[
        "status"
    ] == "completed"
    request = next(iter(_requests(db, run).values()))
    assert request["payload"] == {"source": "original"}
    assert request["schema"] == {}
    cost = graph_response(db, run)["progress"]["stage_costs"][0]
    assert cost["stage"] == "discover" and cost["calls"] == 1
    assert cost["input_tokens"] == 9 and cost["output_tokens"] == 2
    assert cost["seconds"] >= 0 and cost["unmeasured_attempts"] == 0


def test_span_extension_uses_separate_paid_answer_and_reuses_it_on_resume(db, monkeypatch):
    from app.services.document_harness import model

    run, token = _run(db)
    calls = []
    original = {"sources": [{"source_id": "S1", "text": "A1 unit and A2 unit"}],
                "evidence_spans": [], "new_spans_allowed": True}
    expanded = {**original, "evidence_spans": [{"span_id": "P1", "text": "A1"},
                                             {"span_id": "P2", "text": "A2"}],
                "new_spans_allowed": False}

    def expected(payload):
        return {"new_spans": [], "partitions": [], "expressions": [
            {"id": ref["span_id"], "span_id": ref["span_id"], "property_iri": "urn:test:id"}
            for ref in payload["evidence_spans"]
        ]}

    def call(stage, payload, schema, policy):
        calls.append(payload)
        return {"output": expected(payload), "usage": {"output_tokens": 8},
                "seconds": 0.2, "raw_response": {}, "error": None}

    monkeypatch.setattr(model, "call_model", call)
    for _ in range(2):
        repository = Repository(db, run, token)
        assert repository.invoke("referent_candidates", original, {}) == expected(original)
        assert repository.invoke("referent_candidates", expanded, {}) == expected(expanded)
    assert calls == [original, expanded]
    requests = _requests(db, run)
    assert len(requests) == 2
    assert all(r["attempts"] == 1 and r["status"] == "completed" for r in requests.values())


def test_ranking_results_are_separate_reused_and_scoped_to_the_run(db, monkeypatch):
    from app.services.document_harness.ranking import CardRanker

    run, token = _run(db)
    catalog = catalog_from_graph(
        Graph().parse(data='<urn:example:Report> a <http://www.w3.org/2002/07/owl#Class> .',
                      format="turtle"), "urn:example:Report",
    )
    calls = []

    def rank(_self, cards, payload, budget):
        calls.append((cards.snapshot_id, payload, budget))
        return {"snapshot_id": cards.snapshot_id, "selected_iris": [], "seconds": 0.25,
                "operations": [{"operation": "score_pairs", "tokens": 5}]}

    monkeypatch.setattr(CardRanker, "rank", rank)
    first = Repository(db, run, token).rank(catalog, {"sources": []}, 10000)
    assert Repository(db, run, token).rank(catalog, {"sources": []}, 10000) == first
    assert len(calls) == 1
    assert set(read_rows(db, run, model=DocumentRunResult)) == {"rankings"}
    assert "rankings" not in read_rows(db, run)
    other, other_token = _run(db)
    Repository(db, other, other_token).rank(catalog, {"sources": []}, 10000)
    assert len(calls) == 2


def test_ranking_failure_is_recorded_and_never_cached_as_a_success(db, monkeypatch):
    from app.services.document_harness.ontology import SchemaCatalog
    from app.services.document_harness.ranking import CardRanker

    run, token = _run(db)

    def unavailable(*_args):
        raise RuntimeError("ranking_model_unavailable")

    monkeypatch.setattr(CardRanker, "rank", unavailable)
    repository = Repository(db, run, token)
    with pytest.raises(RuntimeError, match="ranking_model_unavailable"):
        repository.rank(SchemaCatalog.model_validate(_catalog()), {"sources": []}, 10000)
    assert not read_rows(db, run, model=DocumentRunResult)
    request = next(iter(read_rows(db, run, model=DocumentRunRequest)["rankings"].values()))
    assert request["status"] == "failed"
    assert request["error"] == "ranking_model_unavailable"


def test_invalid_paid_answer_persists_and_is_not_silently_retried(db, monkeypatch):
    from app.services.document_harness import model

    run, token = _run(db)
    calls = []
    monkeypatch.setattr(model, "call_model", lambda *args: calls.append(1) or {
        "output": None, "raw_response": {"text": "{"}, "usage": {"output_tokens": 16384},
        "seconds": 12, "error": "harness_model_incomplete",
    })
    for _ in range(2):
        with pytest.raises(HarnessCallFailed, match="harness_model_incomplete"):
            Repository(db, run, token).invoke("discover", {}, {})
    assert calls == [1]
    row = next(iter(read_rows(db, run, model=DocumentRunResult)["calls"].values()))
    assert row["raw_response"] == {"text": "{"}
    assert graph_response(db, run)["progress"]["completed_calls"] == 0
    assert graph_response(db, run)["observations"][0]["reason"] == "harness_model_incomplete"


@pytest.mark.parametrize("error_code", ["model_stream_incomplete", "model_request_failed"])
def test_retryable_model_error_saves_only_complete_answer(db, monkeypatch, error_code):
    from app.services.document_harness import model

    run, token = _run(db)
    calls = []

    def call(*_args):
        calls.append(1)
        if len(calls) == 1:
            raise StructuredModelError(error_code)
        return {"output": _discovery(), "raw_response": {"id": "complete"},
                "usage": {"input_tokens": 9}, "seconds": 1.5, "error": None}

    monkeypatch.setattr(model, "call_model", call)
    assert Repository(db, run, token).invoke("discover", {}, {}) == _discovery()
    assert Repository(db, run, token).invoke("discover", {}, {}) == _discovery()
    assert len(calls) == 2
    request = next(iter(_requests(db, run).values()))
    assert request["status"] == "completed"
    assert request["attempts"] == 2
    assert request["unmeasured_attempts"] == 0
    result = next(iter(read_rows(db, run, model=DocumentRunResult)["calls"].values()))
    assert result["raw_response"] == {"id": "complete"}
    costs = graph_response(db, run)["progress"]["stage_costs"]
    assert costs[0]["calls"] == 2
    assert costs[0]["input_tokens"] is None


@pytest.mark.parametrize("error_code", ["model_stream_incomplete", "model_request_failed"])
def test_retryable_model_error_stops_after_bounded_retries(db, monkeypatch, error_code):
    from app.services.document_harness import model

    run, token = _run(db)
    calls = []

    def call(*_args):
        calls.append(1)
        raise StructuredModelError(error_code)

    monkeypatch.setattr(model, "call_model", call)
    with pytest.raises(StructuredModelError, match=error_code):
        Repository(db, run, token).invoke("discover", {}, {})
    assert len(calls) == 3
    assert not read_rows(db, run, model=DocumentRunResult)
    request = next(iter(_requests(db, run).values()))
    assert request["status"] == "failed"
    assert request["attempts"] == 3
    assert request["unmeasured_attempts"] == 0
    assert request["error"] == error_code


@pytest.mark.parametrize("error_code", ["model_stream_incomplete", "model_request_failed"])
def test_retryable_model_error_does_not_retry_after_pause(db, monkeypatch, error_code):
    from app.services.document_harness import model

    run, token = _run(db)
    calls = []

    def call(*_args):
        calls.append(1)
        store = DocumentAnalysisRunStore(db)
        current = store.get_owned(run.recognition_run_id, run.owner_id)
        store.request_control(
            run.recognition_run_id, run.owner_id, action="pause",
            expected_revision=current.revision, request_key="pause-stream-retry",
        )
        db.commit()
        raise StructuredModelError(error_code)

    monkeypatch.setattr(model, "call_model", call)
    with pytest.raises(ModelCancelled):
        Repository(db, run, token).invoke("discover", {}, {})
    assert len(calls) == 1
    request = next(iter(_requests(db, run).values()))
    assert request["attempts"] == 1
    assert request["error"] == error_code


def test_business_updates_preserve_other_rows_and_root_is_not_a_fact(db):
    run, token = _run(db)
    repository = Repository(db, run, token)
    root = {"id": "document", "role": "document_root", "state": "accepted"}
    a = {"id": "a", "label": "A", "state": "candidate"}
    repository.save({"entities": {"document": root, "a": a},
                     "cursor": {"main": {"stage": "discover"}}})
    before = db.scalar(select(DocumentRunCurrentState).where(
        DocumentRunCurrentState.domain == "harness:entities",
        DocumentRunCurrentState.business_key == content_hash("document"),
    )).work_version
    repository.save({"entities": {"a": {**a, "state": "accepted"}}})
    assert repository.load()["entities"]["document"] == root
    assert db.scalar(select(DocumentRunCurrentState).where(
        DocumentRunCurrentState.domain == "harness:entities",
        DocumentRunCurrentState.business_key == content_hash("document"),
    )).work_version == before
    projected = graph_response(db, run)
    assert projected["progress"]["fact_count"] == 1
    assert len(projected["entities"]) == 2


def test_source_is_exact_and_run_owned_without_old_record_protocol(db, tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("项目名称：示例产品")
    doc.save(path)
    ir = analyze_word_core(path).ir
    run, token = _run(db)
    Repository(db, run, token).save({"input": {"document": ir.model_dump(mode="json")}})
    unit = ir.evidence_units[0]
    assert source_response(db, run, unit.evidence_id)["text"] == unit.text
    foreign, _ = _run(db, owner="someone-else")
    with pytest.raises(Exception, match="原文尚未解析"):
        source_response(db, foreign, unit.evidence_id)


def test_new_get_apis_are_read_only_and_old_endpoints_explicitly_reject(
    client, db, analyst_headers, monkeypatch,
):
    run, _token = _run(db)
    calls = []
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *a, **k: calls.append(1))
    base = f"/api/document-analysis/runs/{run.recognition_run_id}"
    status = client.get(base, headers=analyst_headers)
    assert status.status_code == 200, status.text
    assert status.json()["extraction_protocol"] == ENGINE
    response = client.get(base + "/harness-graph", headers=analyst_headers)
    assert response.status_code == 200, response.text
    assert response.json()["protocol"] == ENGINE
    for suffix in ("/metadata", "/graph", "/target-graph", "/source", "/harness",
                   "/harness/context?call_id=x", "/ranking-summary", "/reviews", "/repairs"):
        response = client.get(base + suffix, headers=analyst_headers)
        assert response.status_code == 409, (suffix, response.text)
        assert response.json()["error"]["code"] == "ENGINE_NOT_APPLICABLE"
    assert not calls


def test_retired_non_template_cannot_resume_or_reenter_old_engine(db):
    run, token = _run(db, engine="retired")
    _execute(db, DocumentAnalysisRunStore(db), run, token)
    db.refresh(run)
    assert run.execution_status == "blocked_dependency"
    assert run.error["code"] == "ENGINE_RETIRED"
    with pytest.raises(DocumentAnalysisError, match="旧文档引擎已隔离"):
        DocumentAnalysisApplication(db, ontology_engine=None).control(
            run, action="resume", expected_revision=run.revision, request_key="resume",
            reason=None, role="senior_analyst",
        )


def test_create_api_defaults_to_new_engine_without_old_policy(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    from app.services.document_harness import ontology
    from app.services.document_harness.ontology import SchemaCatalog

    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "runs")
    monkeypatch.setattr(
        ontology, "freeze_catalog", lambda *_: SchemaCatalog.model_validate(_catalog()),
    )
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *a, **k: None)
    doc = Document()
    doc.add_paragraph("对象名称：A")
    path = tmp_path / "report.docx"
    doc.save(path)
    response = client.post(
        "/api/document-analysis/runs", headers=analyst_headers,
        files={"file": ("report.docx", path.read_bytes())},
        data={"root_class_iri": "urn:example:Report", "request_key": "new-engine",
              "metadata_mode": "structure_only"},
    )
    assert response.status_code == 202, response.text
    assert response.json()["stage"] == "accepted"
    run_id = response.json()["recognition_run_id"]
    links = response.json()["links"]
    assert links["metadata"] is None
    assert links["graph"] == f"/api/document-analysis/runs/{run_id}/harness-graph"
    assert links["source"] == f"/api/document-analysis/runs/{run_id}/source?format=original"
    original = client.get(links["source"], headers=analyst_headers)
    assert original.status_code == 200, original.text
    assert original.content == path.read_bytes()
    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers)
    assert status.json()["extraction_protocol"] == ENGINE
    from app.models.document_analysis import DocumentAnalysisArtifact, DocumentAnalysisRun

    run = db.get(DocumentAnalysisRun, run_id)
    source = db.get(DocumentAnalysisArtifact, run.source_artifact_ref).payload
    assert "performance_policy" not in source
    assert source["policy"]["max_output_tokens"] == 16384
    assert set(read_rows(db, run)) == {"metrics", "display"}
    assert graph_response(db, run)["progress"]["completed_calls"] == 0
    graph_url = links["graph"]
    graph = client.get(graph_url, headers=analyst_headers)
    assert graph.status_code == 200, graph.text
    assert graph.json()["stage"] == "ingest"
    assert graph.json()["status"] == "queued"
    assert graph.json()["progress"]["completed_calls"] == 0
    assert set(read_rows(db, run)) == {"metrics", "display"}
    assert graph_response(db, run)["progress"]["completed_calls"] == 0


def test_worker_pause_commits_paid_answer_and_resume_uses_current_business_state(
    db, tmp_path, monkeypatch,
):
    from app.services.document_harness import model
    from app.services.document_harness.runtime import execute_claimed

    doc = Document()
    doc.add_paragraph("本页没有可发现的实体。")
    path = tmp_path / "pause.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    run, token = _run(db, document_hash=ir.original_document_hash)
    repository = Repository(db, run, token)
    repository.save({"input": {"document": ir.model_dump(mode="json")}})
    calls = []

    def answer(stage, _payload, _schema, _policy):
        assert stage == "discover"
        calls.append(stage)
        store = DocumentAnalysisRunStore(db)
        current = store.get_owned(run.recognition_run_id, run.owner_id)
        store.request_control(
            run.recognition_run_id, run.owner_id, action="pause",
            expected_revision=current.revision, request_key="pause-test",
        )
        db.commit()
        return {"output": {"entities": [], "relation_hints": [], "complete": True,
                           "document_field_ids": [], "document_source_fields": [],
                           "unowned_fields": [], "reference_cues": []},
                "raw_response": {}, "error": None, "usage": {}, "seconds": 0.1}

    monkeypatch.setattr(model, "call_model", answer)
    execute_claimed(db, run, token)
    db.refresh(run)
    assert run.execution_status == "paused"
    events = DocumentAnalysisApplication(db, ontology_engine=None).events_response(
        run, after_sequence=0,
    )
    assert events[-1][2]["status"] == "paused"
    assert events[-1][2]["run_revision"] == run.revision
    assert read_rows(db, run)["cursor"]["main"]["phase"] == "reading"
    assert len(read_rows(db, run)["cursor"]["main"]["active_batches"]) == 1
    assert len(read_rows(db, run, model=DocumentRunResult)["calls"]) == 1
    store = DocumentAnalysisRunStore(db)
    store.request_control(
        run.recognition_run_id, run.owner_id, action="resume", expected_revision=run.revision,
        request_key="resume-test",
    )
    token = store.claim(run.recognition_run_id, run.owner_id, actor="test", worker_id="test")
    db.commit()
    execute_claimed(db, run, token)
    db.refresh(run)
    assert run.execution_status == "finished", run.error
    assert graph_response(db, run)["progress"]["scope_complete"] is True
    assert graph_response(db, run)["progress"]["fact_count"] == 0
    assert calls == ["discover"]


def test_new_default_claim_status_and_sse_never_load_old_work_protocols(
    client, db, analyst_headers, monkeypatch,
):
    original_import = builtins.__import__

    def protected(name, *args, **kwargs):
        if name.startswith("app.services.extraction.ontology_guided") or name in {
            "app.services.document_analysis.current_state",
            "app.services.document_analysis.state_artifacts",
            "app.services.document_analysis.reviews",
            "app.services.document_analysis.execution",
        }:
            raise AssertionError("new run loaded old engine: " + name)
        if name == "app.services.document_analysis" and "current_state" in kwargs.get(
            "fromlist", args[2] if len(args) > 2 else (),
        ):
            raise AssertionError("new run loaded old current_state")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", protected)
    run, token = _run(db)
    repository = Repository(db, run, token)
    repository.save({"cursor": {"main": {"stage": "complete", "scope_complete": True}}})
    repository.current()
    repository.progress(stage="complete", status="finished")
    db.refresh(run)
    base = f"/api/document-analysis/runs/{run.recognition_run_id}"
    status = client.get(base, headers=analyst_headers)
    assert status.status_code == 200, status.text
    assert status.json()["stage"] == "complete"
    assert status.json()["status"] == "finished"
    response = client.get(base + "/events", headers=analyst_headers)
    assert response.status_code == 200, response.text
    frames = [json.loads(line.removeprefix("data: ")) for line in response.text.splitlines()
              if line.startswith("data: ")]
    assert frames[-1]["status"] == "finished"
    assert frames[-1]["stage"] == "complete"
    assert frames[-1]["event_head"] == run.event_head
    assert frames[-1]["run_revision"] == run.revision


def test_registered_report_uses_new_engine_and_only_original_document(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    from app.models.document_analysis import DocumentAnalysisArtifact, DocumentAnalysisRun
    from app.models.entity_shadow import EntityShadow
    from app.models.extraction import ExtractionJob
    from app.services.document_harness import ontology
    from app.services.document_harness.ontology import SchemaCatalog

    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "runs")
    monkeypatch.setattr(
        ontology, "freeze_catalog", lambda *_: SchemaCatalog.model_validate(_catalog()),
    )
    dispatched = []
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *a, **k: dispatched.append(a))
    doc = Document()
    doc.add_paragraph("原件中的对象信息。")
    path = tmp_path / "registered.docx"
    doc.save(path)
    iri = "urn:example:uploaded-report"
    job = ExtractionJob(
        source_type="word", source_filename=path.name, document_path=str(path),
        source_config={"doc_class_iri": "urn:example:Report", "doc_ref": iri}, status="paused",
    )
    db.add(job)
    db.flush()
    db.add(EntityShadow(
        iri=iri, class_iri="urn:example:Report", module="document", label_zh="报告",
        properties_json={"job_id": str(job.id)},
    ))
    db.commit()
    response = client.post(
        "/api/document-analysis/documents/runs", params={"document_iri": iri},
        headers=analyst_headers, json={"request_key": "registered-source"},
    )
    assert response.status_code == 202, response.text
    run = db.get(DocumentAnalysisRun, response.json()["recognition_run_id"])
    source = db.get(DocumentAnalysisArtifact, run.source_artifact_ref).payload
    assert source["engine"] == ENGINE
    assert source["origin"]["source_job_id"] == str(job.id)
    assert "template_id" not in source["origin"]
    assert len(dispatched) == 1
    latest = client.get(
        "/api/document-analysis/documents/runs", params={"document_iri": iri},
        headers=analyst_headers,
    )
    assert latest.json()["run"]["extraction_protocol"] == ENGINE
    assert len(dispatched) == 1
    assert job.status == "paused"


def test_public_graph_exposes_review_quotes_and_excludes_unknown_schema_constraints(
    client, db, analyst_headers, tmp_path,
):
    from app.schemas.document_harness import HarnessGraph, HarnessSource
    from app.services.document_harness.ontology import RelationCard
    from app.services.document_harness.source import reference

    document = Document()
    document.add_paragraph("报告描述产品甲。产品甲由工厂乙生产。")
    path = tmp_path / "evidence.docx"
    document.save(path)
    ir = analyze_word_core(path).ir
    unit = ir.evidence_units[0]
    candidate_quote = reference(ir, unit.evidence_id, 0, 8)
    review_quote = reference(ir, unit.evidence_id, 8, len(unit.text))
    catalog = _catalog()
    catalog["classes"]["urn:example:Report"]["relations"] = [
        RelationCard(
            iri="urn:example:describes", label="描述", description="描述对象",
            range_class_iris=("urn:example:Report",),
        ).model_dump(mode="json"),
        RelationCard(
            iri="urn:example:unknown", label="未知边界", description="无法解析约束",
            constraint_status="unresolved",
        ).model_dump(mode="json"),
    ]
    run, token = _run(db, catalog=catalog, document_hash=ir.original_document_hash)
    Repository(db, run, token).save({
        "input": {"document": ir.model_dump(mode="json")},
        "entities": {"document": {
            "id": "document", "label": "报告", "role": "document_root",
            "class_iri": "urn:example:Report", "class_label": "报告", "state": "accepted",
            "reason": "用户选择类型", "evidence": [],
        }},
        "relations": {"r": {
            "id": "r", "subject_id": "document", "object_id": "product",
            "predicate_iri": "urn:example:describes", "label": "描述", "state": "accepted",
            "reason": "原文支持", "evidence": [candidate_quote],
            "review_evidence": [candidate_quote, review_quote],
            "polarity": "positive", "conditions": [],
        }},
    })
    base = f"/api/document-analysis/runs/{run.recognition_run_id}"
    response = client.get(base + "/harness-graph", headers=analyst_headers)
    assert response.status_code == 200, response.text
    graph = HarnessGraph.model_validate(response.json())
    assert [ref.model_dump() for ref in graph.relations[0].evidence] == [
        candidate_quote, review_quote,
    ]
    assert "review_evidence" not in response.json()["relations"][0]
    assert [target.predicate_iri for target in graph.targets] == ["urn:example:describes"]
    source = client.get(base + f"/harness-source/{unit.evidence_id}", headers=analyst_headers)
    assert source.status_code == 200, source.text
    assert HarnessSource.model_validate(source.json()).text == unit.text


@pytest.mark.parametrize("entity_state", ["candidate", "unresolved", "rejected", "accepted"])
def test_raw_observation_keeps_value_and_candidate_owner_in_public_graph(
    db, tmp_path, entity_state,
):
    from app.schemas.document_harness import HarnessGraph
    from app.services.document_harness.source import reference

    doc = Document()
    doc.add_paragraph("负载范围3.8–6.6 kg")
    path = tmp_path / "range.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    unit = ir.evidence_units[0]
    ref = reference(ir, unit.evidence_id, 4, len(unit.text))
    run, token = _run(db)
    Repository(db, run, token).save({
        "entities": {"e": {
            "id": "e", "label": "本次安排", "role": "计划", "class_iri": None,
            "class_label": None, "state": entity_state, "field_ids": ["f"],
            "reason": "类型待核对", "evidence": [ref],
        }},
        "fields": {"f": {"value": "3.8–6.6 kg"}},
        "observations": {"o": {
            "id": "o", "field_id": "f", "label": "负载范围", "reason": "原文观察",
            "evidence": [ref],
        }},
    })
    graph = HarnessGraph.model_validate(graph_response(db, run))
    observation = graph.observations[0]
    assert observation.value == "3.8–6.6 kg"
    assert observation.candidate_subject_ids == ["e"]
    assert graph.progress.fact_count == int(entity_state == "accepted")
    assert observation.evidence[0].text == ir.unit(ref["source_id"]).text[ref["start"]:ref["end"]]


def test_observation_context_uses_saved_cards_and_inherited_predicate_definitions(
    client, db, analyst_headers, monkeypatch,
):
    from copy import deepcopy

    from app.schemas.document_harness import HarnessGraph

    catalog = catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:example:> .
        @prefix doc: <https://example.test/document/> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class ; rdfs:subClassOf :Document ; rdfs:label "报告" .
        :Document a owl:Class ; rdfs:label "文档" .
        :Object a owl:Class ; rdfs:label "对象" .
        :Both a owl:Class ; rdfs:subClassOf :Document, :Object .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :Object .
        doc:number a owl:DatatypeProperty ; rdfs:label "编号" ;
            rdfs:domain :Document ; rdfs:range xsd:string .
        doc:audience a owl:DatatypeProperty ;
            rdfs:domain [owl:unionOf (:Document :Object)] ; rdfs:range xsd:string .
        doc:joint a owl:DatatypeProperty ; rdfs:domain :Document, :Object ;
            rdfs:range xsd:string .
    ''', format="turtle"), "urn:example:Report").model_dump(mode="json")
    run, token = _run(db, catalog=catalog)
    entity = {"id": "document", "label": "报告甲", "role": "document_root",
              "class_iri": "urn:example:Report", "class_label": "报告", "state": "accepted",
              "reason": "用户指定", "evidence": [], "field_ids": ["f"]}
    prop = {"id": "p", "subject_id": "document", "alignment_class_iri": "urn:example:Report",
            "predicate_iri": "https://example.test/document/number", "label": "编号",
            "field_id": "f", "value": "ABC", "source_value": "ABC", "source_unit": None,
            "value_component": "whole", "value_evidence": [], "state": "accepted",
            "reason": "编号有原文依据", "evidence": []}
    repository = Repository(db, run, token)
    repository.save({
        "entities": {"document": entity, "object": {**entity, "id": "object",
            "role": "object", "label": "对象甲", "class_iri": "urn:example:Object"}},
        "fields": {"f": {"value": "ABC"}}, "properties": {"p": prop},
        "windows": {"w": {"guidance_class_iris": ["urn:example:Object"]}},
        "observations": {"o": {"id": "o", "kind": "field", "field_id": "f",
            "window_id": "w", "label": "原编号", "reason": "原文观察", "evidence": [],
            "alignment_outcomes": {"attempt": {"subject_id": "object",
                "class_iri": "urn:example:Object", "state": "unmatched",
                "predicate_iris": ["https://example.test/document/audience"],
                "reason": "可阅对象属性不匹配编号"}}}},
    })
    before = deepcopy(read_rows(db, run))
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *a, **k: pytest.fail("GET ran"))
    response = client.get(
        f"/api/document-analysis/runs/{run.recognition_run_id}/harness-graph",
        headers=analyst_headers,
    )
    assert response.status_code == 200, response.text
    graph = HarnessGraph.model_validate(response.json())
    observation, property_ = graph.observations[0], graph.properties[0]
    assert set(observation.candidate_subject_ids) == {"document", "object"}
    assert [(c.iri, c.role) for c in observation.discovery_cards] == [
        ("urn:example:Report", "document_properties"), ("urn:example:Object", "reading"),
    ]
    groups = {a.subject_id: a for a in observation.alignments}
    assert groups["document"].property_ids == ["p"]
    assert groups["document"].attempts == []
    assert groups["object"].property_ids == []
    attempt = groups["object"].attempts[0]
    assert attempt.state == "unmatched" and attempt.reason == "可阅对象属性不匹配编号"
    assert " 或 " in attempt.predicates[0].domain_text
    assert property_.card.iri == "urn:example:Report"
    assert property_.predicate.namespace == "https://example.test/document/"
    assert property_.predicate.domain_text == "文档 <urn:example:Document>"
    assert property_.state == "accepted"
    assert read_rows(db, run) == before
    from app.services.document_harness.projection import predicate_ref

    conjunction = predicate_ref(catalog["classes"], "https://example.test/document/joint")
    assert " 且 " in conjunction["domain_text"]
    assert " 或 " not in conjunction["domain_text"]

    # Losing recorded provenance does not authorize guessing it from current entity type.
    prop.pop("alignment_class_iri")
    row = {**before["observations"]["o"], "discovery_window_id": None}
    repository.save({"properties": {"p": prop}, "observations": {"o": row}})
    projected = graph_response(db, run)
    assert projected["properties"][0]["card"] is None
    assert projected["observations"][0]["discovery_cards"] == []


def test_public_graph_normalizes_missing_legacy_value_evidence_to_empty_list(db):
    from app.schemas.document_harness import HarnessGraph

    run, token = _run(db)
    Repository(db, run, token).save({
        "entities": {"document": {
            "id": "document", "label": "报告", "role": "document_root",
            "class_iri": "urn:example:Report", "class_label": "报告", "state": "accepted",
            "reason": "用户选择类型", "evidence": [], "field_ids": ["f"],
        }},
        "properties": {"legacy": {
            "id": "legacy", "subject_id": "document", "predicate_iri": None,
            "label": "旧属性", "value": "原值", "source_value": "原值",
            "source_unit": None, "value_component": "whole", "field_id": "f",
            "state": "unresolved", "reason": "旧运行未保存值级引用", "evidence": [],
        }},
    })
    payload = graph_response(db, run)
    assert payload["properties"][0]["value_evidence"] == []
    assert HarnessGraph.model_validate(payload).properties[0].value_evidence == []


def test_failed_transport_records_duration_and_preserves_unknown_token_counts(db, monkeypatch):
    from app.schemas.document_harness import HarnessGraph
    from app.services.document_harness import model, runtime

    run, token = _run(db)
    times = iter((10.0, 13.75, 20.0, 22.0))
    monkeypatch.setattr(runtime, "monotonic", lambda: next(times))

    def failed(*_args):
        raise TimeoutError("provider unavailable")

    monkeypatch.setattr(model, "call_model", failed)
    with pytest.raises(TimeoutError):
        Repository(db, run, token).invoke("discover", {}, {})
    request = next(iter(_requests(db, run).values()))
    assert request["finished_at"] and request["started_at"]
    assert request["duration_us"] == 3_750_000
    assert request["unknown_input"] == request["unknown_output"] == 1
    graph = HarnessGraph.model_validate(graph_response(db, run))
    assert graph.progress.stage_costs[0].seconds == 3.75
    assert graph.progress.stage_costs[0].input_tokens is None
    assert graph.progress.stage_costs[0].output_tokens is None
    assert graph.observations[0].reason == "TimeoutError"
    monkeypatch.setattr(model, "call_model", lambda *_: {
        "output": _discovery(), "raw_response": {},
        "usage": {"input_tokens": 9, "output_tokens": 2},
        "seconds": 2.0, "error": None,
    })
    Repository(db, run, token).invoke("discover", {}, {})
    cost = HarnessGraph.model_validate(graph_response(db, run)).progress.stage_costs[0]
    assert cost.calls == 2 and cost.seconds == 5.75
    # The successful retry cannot invent usage for the earlier network failure.
    assert cost.input_tokens is None and cost.output_tokens is None


def test_independent_public_contract_rejects_unknown_fields_and_invalid_quote_spans():
    from pydantic import ValidationError

    from app.schemas.document_harness import HarnessSourceRef

    quote = {"source_id": "source", "text": "原文", "start": 0, "end": 2, "page": None,
             "section_id": "section", "block_id": "block"}
    assert HarnessSourceRef.model_validate(quote).text == "原文"
    with pytest.raises(ValidationError, match="source_span_length_mismatch"):
        HarnessSourceRef.model_validate({**quote, "end": 1})
    with pytest.raises(ValidationError, match="Extra inputs are not permitted"):
        HarnessSourceRef.model_validate({**quote, "proof_ref": "old-protocol"})


def test_retired_run_offers_only_cleanup_and_cannot_schedule_property_repair(
    client, db, analyst_headers, monkeypatch,
):
    run, token = _run(db, engine="retired")
    repository = Repository(db, run, token)
    repository.current()
    repository.progress(status="failed")
    db.refresh(run)
    calls = []
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *a, **k: calls.append(1))
    base = f"/api/document-analysis/runs/{run.recognition_run_id}"
    response = client.get(base, headers=analyst_headers)
    assert response.status_code == 200, response.text
    assert response.json()["available_actions"] == ["cancel", "delete"]
    history = client.get(base + "/repairs", headers=analyst_headers)
    assert history.status_code == 200, history.text
    assert history.json()["can_repair"] is False
    response = client.post(base + "/repairs", headers=analyst_headers, json={
        "review_id": str(uuid4()), "request_key": "cannot-repair",
        "expected_run_revision": run.revision,
    })
    assert response.status_code == 409, response.text
    assert response.json()["error"]["code"] == "ENGINE_RETIRED"
    assert calls == []
