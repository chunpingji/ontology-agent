"""Reading groups reduce actual coordinator work while retaining physical sources."""

import pytest
from docx import Document

from app.models.document_analysis import DocumentRunCurrentState
from app.services.document_analysis import current_state
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.model_reference_projection import (
    project_reference_payload,
)
from app.services.extraction.ontology_guided.record_discovery import ContextualDiscoveryPolicy
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction import test_record_executor as fixture

pytest_plugins = ["tests.test_extraction.test_tool_engine_resume"]


@pytest.mark.parametrize("enabled,expected_tasks", [(True, 1), (False, 2)])
def test_continuous_sections_share_discovery_and_restore_the_same_work(
    tmp_path, monkeypatch, current_run, enabled, expected_tasks,
):
    original_arguments = fixture.arguments

    def arguments(path, lines):
        values = original_arguments(path, lines)
        document = Document()
        document.add_heading("工艺描述", 1)
        document.add_heading("工艺", 2)
        document.add_paragraph("装置甲先进行溶解。")
        document.add_heading("工艺（续）", 2)
        document.add_paragraph("装置甲随后完成过滤。")
        source = path / "joint-reading.docx"
        document.save(source)
        analysis = analyze_word_core(source)
        values.update(ir=analysis.ir, metadata=prepare_metadata(
            analysis.ir, section_tree=analysis.structure.section_tree.to_dict(),
            summary_version="reading-test",
        ))
        return values

    monkeypatch.setattr(fixture, "arguments", arguments)
    args, factory, requests, hooks = fixture.record_setup(
        tmp_path, monkeypatch, current_run, split=True, empty_relations=True,
        transform=lambda view, payload, **options: {name: [] for name in payload},
    )
    adapter = factory().adapter
    if enabled:
        adapter.record_discovery = adapter.record_discovery.model_copy(update={
            "contextual": ContextualDiscoveryPolicy(),
        })
    result = factory().run(**args, **hooks)
    discovery = [view for view in requests if "members" not in view]
    assert len(discovery) == expected_tasks, result.diagnostics
    index = RecordIndex(args["ir"])
    assert len(index.records) == 2
    sources = {unit.evidence_id for record in index.records for unit in record.source_units}
    wire_sources = set(project_reference_payload({"evidence_ids": sorted(sources)})[
        "evidence_ids"
    ])
    for view in discovery:
        actual = {unit["evidence_id"] for unit in view["evidence_units"]
                  if unit["fact_eligible"]}
        assert actual == wire_sources if enabled else len(actual) == 1
    store, run, _ = current_run
    rows = current_state.read_rows(store, run, DocumentRunCurrentState,
                                   prefix="work:record_discovery")["work:record_discovery"]
    assert len(rows) == expected_tasks
    assert all(row["value"]["status"] == "examined" for row in rows.values())
    count = len(requests)
    factory().run(
        **args, **hooks,
        resume_state=vars(current_state.restore_work(store, run, run.run_fingerprint)),
        model_call_state=current_state.restore_calls(store, run, run.run_fingerprint),
    )
    assert len(requests) == count
