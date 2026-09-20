"""Navigation stays addressable without becoming facts or summary input."""

import json
from copy import deepcopy

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from app.config import settings
from app.services.extraction import word_tree_summarizer
from app.services.extraction.document_ir import DocumentIR, build_document_ir
from app.services.extraction.docx_structure import PARSER_VERSION, parse_docx_structure
from app.services.extraction.evidence_identity import evidence_hash, stable_id
from app.services.extraction.ontology_guided.records import RecordIndex


def _parse(tmp_path, *, body=None, body_style=None, pages=True):
    doc = Document()
    doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading("目录", 1)
    doc.add_paragraph("导航产品条目\t1" if pages else "导航产品条目", style="TOC 1")
    if body:
        doc.add_paragraph(body, style=body_style)
    path = tmp_path / "navigation.docx"
    doc.save(path)
    structure = parse_docx_structure(path)
    return structure, build_document_ir(path, structure)


@pytest.mark.parametrize("body_style,pages", [(None, True), ("TOC 1", True), ("TOC 1", False)])
@pytest.mark.parametrize("body", ["这是目录之后的真实正文。", "文档编号：DOC-001"])
def test_navigation_filtered_but_unheaded_and_toc_styled_body_remains(
    tmp_path, body_style, pages, body,
):
    structure, ir = _parse(tmp_path, body=body, body_style=body_style, pages=pages)
    index = RecordIndex(ir)
    assert ir.parser_version == str(PARSER_VERSION) == "7"
    assert [unit.navigation_role for unit in ir.evidence_units] == [
        "toc_heading", "toc_entry", None,
    ]
    assert [record.text for record in index.records] == [body]
    # Membership alone does not turn this whole section into navigation.
    assert ir.evidence_units[1].section_node_id == ir.evidence_units[2].section_node_id
    assert all(unit.navigation_role is None for values in index.headings_by_parent.values()
               for unit in values)
    assert structure.sections[0].paras[-1] == body
    for unit in ir.evidence_units:
        assert ir.resolve(ir.anchor(unit.evidence_id)) == unit.text
    assert DocumentIR.model_validate_json(ir.model_dump_json()).model_dump(mode="json") == (
        ir.model_dump(mode="json")
    )


def _enable(monkeypatch):
    monkeypatch.setattr(settings, "local_llm_enabled", True)
    monkeypatch.setattr(settings, "llm_word_tree_summary_enabled", True)


def test_navigation_only_does_not_call_summary_model_and_keeps_tree(tmp_path, monkeypatch):
    structure, ir = _parse(tmp_path)
    _enable(monkeypatch)
    before = deepcopy((structure.blocks, structure.sections, structure.paragraphs))

    def unexpected(*_args, **_kwargs):
        pytest.fail("navigation-only material must not call the model")

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", unexpected)
    root = word_tree_summarizer.summarize_word_tree(structure, object())
    assert root.children[0].pages
    for node in word_tree_summarizer._chapters(root):
        assert node.layer_metadata.summary_source == "empty"
        assert node.layer_metadata.content_summary is None
        assert all(page.page_metadata.summary_source == "empty" for page in node.pages)
    assert before == (structure.blocks, structure.sections, structure.paragraphs)
    assert RecordIndex(ir).records == []


@pytest.mark.parametrize("fallback", [False, True])
def test_summaries_exclude_navigation_but_keep_body_in_the_same_section(
    tmp_path, monkeypatch, fallback,
):
    body = "真实正文载明药品甲用于临床备样。"
    structure, _ = _parse(tmp_path, body=body)
    _enable(monkeypatch)
    prompts = []

    def reply(_client, **kwargs):
        nodes = json.loads(kwargs["user"])["nodes"]
        prompts.extend(node["content"] for node in nodes)
        return {"summaries": [{"node_id": node["node_id"], "content_summary": body}
                              for node in nodes]}

    monkeypatch.setattr(word_tree_summarizer, "chat_with_schema", reply)
    if fallback:
        root = word_tree_summarizer.fallback_word_tree_summaries(structure)
        prompts = [node.layer_metadata.content_summary or ""
                   for node in word_tree_summarizer._chapters(root)]
    else:
        root = word_tree_summarizer.summarize_word_tree(structure, object())
    assert prompts and any(body in text for text in prompts)
    assert all("导航产品条目" not in text and "目录" not in text for text in prompts)
    assert root.children[0].paragraph_indices == [1, 2]


def test_legacy_ir_keeps_exact_payload_hash_after_default_role_is_added(tmp_path):
    _, ir = _parse(tmp_path)
    legacy = ir.model_dump(mode="json")
    legacy["parser_version"] = "6"
    for block in legacy["blocks"]:
        block.pop("navigation_role", None)
    for unit in legacy["evidence_units"]:
        unit.pop("navigation_role", None)
    legacy["structure_hash"] = evidence_hash({key: legacy[key] for key in (
        "parser_version", "ir_version", "structure_policy_version", "nodes", "blocks",
        "tables", "evidence_units", "pagination",
    )})
    legacy["analysis_id"] = stable_id("analysis", {
        "document_hash": legacy["document_hash"], "structure_hash": legacy["structure_hash"],
        "role": legacy["document_role"], "original_document_hash": legacy["original_document_hash"],
    })
    loaded = DocumentIR.model_validate(legacy)
    assert all(unit.navigation_role is None for unit in loaded.evidence_units)
    assert loaded.model_dump(mode="json") == legacy
    assert DocumentIR.model_validate_json(loaded.model_dump_json()).structure_hash == (
        legacy["structure_hash"]
    )
    # Existing frozen runs retain their original interpretation and record set.
    assert len(RecordIndex(loaded).records) == 1


def test_navigation_role_changes_invalidate_frozen_structure_identity(tmp_path):
    _, ir = _parse(tmp_path)
    modified = ir.model_dump(mode="json")
    modified["evidence_units"][1]["navigation_role"] = None
    with pytest.raises(ValueError, match="structure identity"):
        DocumentIR.model_validate(modified)
    modified = ir.model_dump(mode="json")
    modified["blocks"][1]["navigation_role"] = None
    with pytest.raises(ValueError, match="structure identity"):
        DocumentIR.model_validate(modified)
