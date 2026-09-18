"""Freeze failures remain local and record referents require actual composition proof."""

from __future__ import annotations

import copy

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.claim_freeze import freeze_proposal
from app.services.extraction.ontology_guided.claim_protocol import (
    DiscoveryEnvelope,
    EntityDependencyView,
    SchemaCard,
    build_verification_input,
)
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    EdgeSpec,
    SemanticDecision,
    SlotSpec,
    SubjectRef,
    TraversalScope,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.mentions import MentionRegistry
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def source(tmp_path):
    document = Document()
    document.add_heading("正文", level=1)
    document.add_paragraph("主体甲关联对象乙或对象丙，恰选其一。数量为5 mg。")
    document.add_heading("补充", level=1)
    document.add_paragraph("附录补充：对象丁属于类型T。")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数值"
    for row, value in ((1, "5"), (2, "7")):
        table.cell(row, 0).text = "同名记录"
        table.cell(row, 1).text = value
    path = tmp_path / "freeze.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    record = next(value for value in index.records if "主体甲关联" in value.text)
    supplement = next(value for value in index.records if "附录补充" in value.text)
    source_unit, supplement_unit = record.source_units[0], supplement.source_units[0]
    subject = SubjectRef(
        entity_id="subject-server-id", revision=3, class_iri="urn:Source", is_document_root=True
    )

    def quote(text, *, supplemental=False):
        unit = supplement_unit if supplemental else source_unit
        return {"evidence_id": unit.evidence_id, "text": text, "context_text": None}

    def task_context(predicate):
        task = RecognitionTask.create(
            subject=subject, predicate_iri=predicate.iri, predicate_kind=predicate.kind,
            record_id=record.record_id, phase=1, hop=0, dependency_hash="frozen-dependency",
        )
        root_ref = VersionedRef(id=subject.entity_id, revision=subject.revision)
        seed = VerificationTarget.create(
            run_fingerprint="test-run",
            claim_ref=VersionedRef(id=task.claim_lineage_id, revision=1),
            task_id=task.task_id, check_kind="predicate_entailment",
            document_context=DocumentContext(
                document_hash=index.ir.document_hash, document_class_iri=subject.class_iri,
                root_ref=root_ref,
            ),
            subject_ref=subject, predicate_iri=predicate.iri,
            ontology_hash=evidence_hash("ontology"),
            source_scope_hash=evidence_hash(record.record_id), context_hash=evidence_hash("seed"),
        )
        owner = index.ir.anchor(source_unit.evidence_id, 0, len("主体甲"))
        context = assemble_context(
            seed, record.record_id, index, subject_evidence_refs=[owner], subject_label="主体甲",
            required_context_refs=[index.ir.anchor(
                supplement_unit.evidence_id, 0, len(supplement_unit.text)
            )], predicate=predicate,
        )
        card = SchemaCard(
            schema_card_id="card-1", subject_ref=root_ref, class_iris=[subject.class_iri],
            ontology_snapshot_id="ontology-1", menu_id="menu-1", predicates=[predicate],
            quantity_policies=[], identity_keys=[], unsupported_constraints=[],
        )
        return dict(task=task, context=context, card=card, index=index, generation=1)

    predicate = EdgeSpec(
        iri="urn:links", label="关联", declared_by=[subject.class_iri], range_class_iris=["urn:T"]
    )
    entities = [dict(
        local_id=identity, class_iri="urn:T", representation="mention", mentions=[quote(label)],
        record_components=[], identifier_claims=[],
    ) for identity, label in (("b", "对象乙"), ("c", "对象丙"))]
    qualifiers = {
        "polarity": "affirmed", "modality": "asserted",
        "condition_support": [], "scope_qualifiers": [],
    }
    relation = dict(
        local_id="r", subject_id=subject.entity_id, predicate_iri=predicate.iri,
        object_ids=["b", "c"], selection="one_of", bridge_support=[quote(source_unit.text)],
        selection_support=[quote("对象乙或对象丙，恰选其一")], qualifiers=qualifiers,
        bridge_kind="explicit_assertion", bridge_ref_ids=[],
    )
    proposal = dict(entities=entities, properties=[], relations=[relation], external_links=[],
                    observations=[])
    return dict(
        index=index, options=task_context(predicate), task_context=task_context,
        proposal=proposal, quote=quote, source_unit=source_unit, predicate=predicate,
    )


def _freeze(source, proposal=None, **options):
    return freeze_proposal(
        DiscoveryEnvelope.model_validate(proposal or source["proposal"], strict=True),
        **{**source["options"], **options},
    )


def _root_dependency(options):
    task, context = options["task"], options["context"]
    payload = {
        "entity_ref": VersionedRef(id=task.subject.entity_id, revision=task.subject.revision),
        "class_iri": task.subject.class_iri, "grounding_kind": "document_root",
        "root_origin": "user_specified", "proposal": None,
        "source_refs": context.subject_evidence_refs, "dependency_refs": [],
    }
    return EntityDependencyView(**payload, content_hash=evidence_hash(payload))


def test_freeze_keeps_original_selection_and_server_assigned_identity(source):
    original = copy.deepcopy(source["proposal"])
    frozen = _freeze(source)
    assert frozen.claim_issues == {}
    assert frozen.relations[0].object_ids == ["b", "c"]
    assert frozen.relations[0].selection == "one_of"
    assert source["proposal"] == original
    assert frozen.local_ref_map["b"].id != "b"
    assert frozen.local_ref_map["subject-server-id"].revision == 3
    assert frozen.model_dump() == _freeze(source).model_dump()
    next_generation = _freeze(source, generation=2)
    assert next_generation.local_ref_map["b"] != frozen.local_ref_map["b"]
    assert next_generation.assertion_generation == 2
    assert frozen.content_hash == evidence_hash(frozen.model_dump(exclude={"content_hash"}))


def test_registered_dependency_in_whole_fragment_keeps_frozen_context_identity(source):
    options = source["options"]
    context = options["context"]
    fragment = next(f for f in context.fragments
                    if f.anchor.evidence_id == source["source_unit"].evidence_id)
    fragment.anchor = fragment.anchor.model_copy(update={"span_start": None, "span_end": None})
    before = context.model_dump()
    frozen = _freeze(source)
    verification = build_verification_input(
        frozen, discovery_ref="discovery", context=context, card=options["card"],
        scope=TraversalScope.create(), entity_dependencies=[_root_dependency(options)],
        external_candidates=[], bridge_dependencies=[], scope_resolutions=[],
    )
    assert {target.payload.local_id for target in verification.targets} == {"b", "c", "r"}
    assert context.model_dump() == before


def test_bad_entity_blocks_whole_dependent_group_but_not_independent_claim(source):
    proposal = copy.deepcopy(source["proposal"])
    proposal["entities"][0]["mentions"][0]["text"] = "不存在的对象"
    sibling = copy.deepcopy(proposal["relations"][0])
    sibling.update(local_id="independent", object_ids=["c"], selection="all", selection_support=[])
    proposal["relations"].append(sibling)
    frozen = _freeze(source, proposal)
    assert set(frozen.claim_issues) == {"b", "r"}
    assert "source_excerpt_mismatch" in frozen.claim_issues["b"]
    assert frozen.claim_issues["r"] == ["entity_dependency_invalid"]
    assert frozen.relations[0].object_ids == ["b", "c"]
    assert frozen.relations[0].selection == "one_of"
    options = source["options"]
    verifier_input = build_verification_input(
        frozen, discovery_ref="discovery-1", context=options["context"], card=options["card"],
        scope=TraversalScope.create(), entity_dependencies=[_root_dependency(options)],
        external_candidates=[], bridge_dependencies=[], scope_resolutions=[],
    )
    assert {target.payload.local_id for target in verifier_input.targets} == {"c", "independent"}


@pytest.mark.parametrize("field", ["class", "predicate", "endpoint", "selection", "bridge"])
def test_local_mechanical_failures_are_retained_as_stable_claim_issues(source, field):
    proposal = copy.deepcopy(source["proposal"])
    expected = {
        "class": ("b", "class_outside_menu"), "predicate": ("r", "predicate_outside_menu"),
        "endpoint": ("r", "entity_reference_missing"),
        "selection": ("r", "selection_evidence_missing"),
        "bridge": ("r", "bridge_reference_missing"),
    }
    if field == "class":
        proposal["entities"][0]["class_iri"] = "urn:NotAuthorized"
    elif field == "predicate":
        proposal["relations"][0]["predicate_iri"] = "urn:NotAuthorized"
    elif field == "endpoint":
        proposal["relations"][0]["object_ids"][0] = "not-registered"
    elif field == "selection":
        proposal["relations"][0]["selection_support"] = []
    else:
        proposal["relations"][0]["bridge_kind"] = "resolved_reference_chain"
    frozen = _freeze(source, proposal)
    local_id, code = expected[field]
    assert code in frozen.claim_issues[local_id]
    assert len(frozen.entities) == 2 and len(frozen.relations) == 1


def test_supplemental_context_can_support_but_cannot_create_entity_fact(source):
    proposal = copy.deepcopy(source["proposal"])
    proposal["entities"][0]["mentions"] = [source["quote"]("对象丁", supplemental=True)]
    proposal["relations"][0]["qualifiers"]["condition_support"] = [
        source["quote"]("对象丁属于类型T", supplemental=True)
    ]
    frozen = _freeze(source, proposal)
    assert frozen.claim_issues["b"] == ["fact_source_outside_scope"]
    assert frozen.claim_issues["r"] == ["entity_dependency_invalid"]


def test_repeated_quote_needs_an_exact_same_source_context(source):
    proposal = copy.deepcopy(source["proposal"])
    mention = proposal["entities"][0]["mentions"][0]
    mention["text"] = "对象"
    assert "ambiguous_source_quote" in _freeze(source, proposal).claim_issues["b"]
    mention["context_text"] = "关联对象乙"
    assert _freeze(source, proposal).claim_issues == {}


def test_record_proposal_requires_a_fact_component_but_accepts_context_components(source):
    proposal = copy.deepcopy(source["proposal"])
    entity = proposal["entities"][0]
    entity.update(representation="record", mentions=[], record_components=[
        {"role": "subject", "quote": source["quote"]("对象乙")},
        {"role": "context", "quote": source["quote"]("对象丁", supplemental=True)},
    ])
    assert _freeze(source, proposal).claim_issues == {}
    entity["record_components"].pop(0)
    assert "record_fact_components_missing" in _freeze(source, proposal).claim_issues["b"]


def test_external_link_must_reference_an_explicit_registered_candidate(source):
    proposal = copy.deepcopy(source["proposal"])
    proposal["external_links"] = [{
        "local_id": "link-1", "subject_id": "b", "external_candidate_id": "not-registered",
        "identity_support": [source["quote"]("对象乙")],
    }]
    frozen = _freeze(source, proposal)
    assert frozen.claim_issues == {"link-1": ["external_candidate_missing"]}
    assert frozen.entities and frozen.relations


def test_freeze_rejects_corrupt_server_text_and_stale_registered_subject_version(source):
    poisoned = source["options"]["context"].model_copy(deep=True)
    poisoned.fragments[0].text = "伪造原文"
    with pytest.raises(ValueError, match="no longer matches source"):
        _freeze(source, context=poisoned)
    dependency = _root_dependency(source["options"]).model_dump(mode="json")
    dependency["entity_ref"]["revision"] += 1
    dependency["content_hash"] = evidence_hash({
        key: value for key, value in dependency.items() if key != "content_hash"
    })
    with pytest.raises(ValueError, match="registered_entity_identity_conflict"):
        _freeze(source, entity_dependencies=[EntityDependencyView.model_validate(dependency)])


def test_property_value_is_fact_bound_but_separate_unit_quote_remains_intact(source):
    predicate = SlotSpec(
        iri="urn:quantity", label="数量", declared_by=["urn:Source"], datatype_iris=["decimal"]
    )
    options = source["task_context"](predicate)
    value = dict(
        local_id="p", subject_id="subject-server-id", predicate_iri=predicate.iri,
        value_quote=source["quote"]("5"), field_support=[source["quote"]("数量")],
        unit_support=[source["quote"]("mg")],
        qualifiers=copy.deepcopy(source["proposal"]["relations"][0]["qualifiers"]),
        bridge_kind="explicit_assertion", bridge_ref_ids=[],
    )
    proposal = dict(
        entities=[], properties=[value], relations=[], external_links=[], observations=[]
    )
    frozen = _freeze(source, proposal, **options)
    assert frozen.claim_issues == {}
    assert frozen.properties[0].value_quote.text == "5"
    assert frozen.properties[0].unit_support[0].text == "mg"
    proposal["properties"][0]["value_quote"] = source["quote"]("类型T", supplemental=True)
    assert _freeze(source, proposal, **options).claim_issues["p"] == ["fact_source_outside_scope"]


def test_registered_identity_collision_is_not_resolved_by_overwriting_the_root(source):
    proposal = copy.deepcopy(source["proposal"])
    proposal["entities"][0]["local_id"] = "subject-server-id"
    with pytest.raises(ValueError, match="local_id_conflicts"):
        _freeze(source, proposal)


def test_allowed_self_relationship_is_not_rejected_by_a_hardcoded_root_rule(source):
    predicate = source["predicate"].model_copy(update={"range_class_iris": ["urn:Source", "urn:T"]})
    proposal = copy.deepcopy(source["proposal"])
    proposal["relations"][0].update(
        object_ids=["subject-server-id"], selection="all", selection_support=[]
    )
    frozen = _freeze(source, proposal, **source["task_context"](predicate))
    assert frozen.claim_issues == {}


def test_observations_are_preserved_as_untrusted_hints(source):
    proposal = copy.deepcopy(source["proposal"])
    proposal["observations"] = [{
        "subject_id": None, "predicate_iri": None,
        "quote": source["quote"]("模型声称存在但原文没有"),
        "kind": "unknown", "reason": "仅为待检查提示",
    }]
    frozen = _freeze(source, proposal)
    assert frozen.observations[0].quote.text == "模型声称存在但原文没有"
    assert frozen.claim_issues == {}
    assert "模型声称存在但原文没有" not in {
        fragment.text for fragment in source["options"]["context"].fragments
    }


def _decision(kind, refs, *, identity="decision", target="entity-target", verdict="supported"):
    return SemanticDecision(
        decision_id=f"{identity}-{kind}", target_id=target, check_kind=kind, verdict=verdict,
        reason_code="test", reason="独立原文核验", support_refs=refs,
        verifier_version="test", attempt_id="attempt-1",
    )


def _record_parts(source):
    index = source["index"]
    views = [view for view in index.record_views if view.table_path]
    components = [view.source_refs for view in views]
    return views, components


def test_record_referent_requires_composition_and_creates_no_fake_mention(source):
    registry = MentionRegistry(source["index"].ir, source["index"])
    views, components = _record_parts(source)
    kwargs = dict(
        component_refs=components[0],
        composition_decision=_decision("referent", components[0]),
        subject_role_decision=_decision("subject_role", components[0]),
    )
    record = registry.create_record_referent([views[0].record_view_id], **kwargs)
    assert record.kind == "record"
    assert record.mention_refs == []
    assert record.record_view_refs == [views[0].record_view_id]
    assert registry.mentions == []
    assert registry.create_record_referent([views[0].record_view_id], **kwargs) is record


def test_same_label_in_distinct_records_and_distinct_components_never_merge(source):
    registry = MentionRegistry(source["index"].ir, source["index"])
    views, components = _record_parts(source)
    identities = []
    for view, parts in ((views[0], components[0]), (views[1], components[1]),
                        (views[0], components[0][:1])):
        record = registry.create_record_referent(
            [view.record_view_id], component_refs=parts,
            composition_decision=_decision("referent", parts),
            subject_role_decision=_decision("subject_role", parts),
        )
        identities.append(record.referent_id)
    assert len(set(identities)) == 3
    assert registry.mentions == []


@pytest.mark.parametrize(
    "failure", ["unverified", "different_target", "missing_part", "wrong_view"]
)
def test_record_referent_cannot_use_unverified_or_unbound_composition(source, failure):
    registry = MentionRegistry(source["index"].ir, source["index"])
    views, components = _record_parts(source)
    composition = _decision("referent", components[0])
    role = _decision("subject_role", components[0])
    selected = [views[0].record_view_id]
    if failure == "unverified":
        role = role.model_copy(update={"verdict": "undetermined"})
    elif failure == "different_target":
        role = role.model_copy(update={"target_id": "another-entity"})
    elif failure == "missing_part":
        composition = composition.model_copy(update={"support_refs": components[0][:1]})
    else:
        selected = [views[1].record_view_id]
    with pytest.raises(ValueError):
        registry.create_record_referent(
            selected, component_refs=components[0], composition_decision=composition,
            subject_role_decision=role,
        )
    assert registry.referents == [] and registry.mentions == []
