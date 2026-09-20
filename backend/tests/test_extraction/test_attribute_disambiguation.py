"""Physical fields retain source ownership while candidate pairs remain unproved."""

from dataclasses import replace

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.attribute_disambiguation import (
    attribute_options,
    compile_attribute_card,
    extract_attribute_fields,
    field_payload,
)
from app.services.extraction.ontology_guided.claim_protocol import (
    EntityDependencyView,
    ExtractionProfile,
)
from app.services.extraction.ontology_guided.contracts import (
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
)
from app.services.extraction.ontology_guided.reading_groups import build_reading_groups
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core

REPORT, EQUIPMENT, BATCH = "urn:Report", "urn:Equipment", "urn:Batch"
STRING = "http://www.w3.org/2001/XMLSchema#string"


def index_for(tmp_path, paragraphs=(), *, heading="正文", table=None):
    doc = Document()
    doc.add_heading(heading, level=1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for row, cells in enumerate(table):
            for col, value in enumerate(cells):
                grid.cell(row, col).text = value
    doc.add_heading("其他章节", level=1)
    doc.add_paragraph("无关设备乙")
    path = tmp_path / "fields.docx"
    doc.save(path)
    return RecordIndex(analyze_word_core(path).ir)


def ontology():
    definitions = {}
    for owner, label, iri in (
        (REPORT, "文档编号", "urn:documentNumber"),
        (EQUIPMENT, "设备编号", "urn:equipmentCode"),
        (BATCH, "批号", "urn:batchNumber"),
    ):
        value = dict(iri=owner, label=owner, declared_properties=[SlotSpec(
            iri=iri, label=label, datatype_iris=[STRING], declared_by=[owner],
        )])
        definitions[owner] = OntologyClassDefinition(**value, source_hash=evidence_hash(value))
    return OntologySnapshot(snapshot_id="ontology", ontology_hash=evidence_hash(definitions),
                            classes=definitions)


def compile_card(field, classes=(REPORT, EQUIPMENT, BATCH)):
    return compile_attribute_card(field, ontology(), classes, "scope", ExtractionProfile())


def entity(index, identity, class_iri, record=None, *, revision=1, root=False):
    unit = (record or index.records[0]).source_units[0]
    value = dict(
        entity_ref={"id": identity, "revision": revision}, class_iri=class_iri,
        grounding_kind="document_root" if root else "mention",
        root_origin="user_specified" if root else None,
        proposal=None if root else dict(
            local_id=identity, class_iri=class_iri, representation="mention",
            mentions=[{"evidence_id": unit.evidence_id, "text": unit.text,
                       "context_text": None}], record_components=[], identifier_claims=[],
        ),
        source_refs=[index.ir.anchor(unit.evidence_id, 0, len(unit.text))], dependency_refs=[],
    )
    return EntityDependencyView(**value, content_hash=evidence_hash(value))


@pytest.mark.parametrize("text,label,value,exclusive", [
    ("文档编号： 0007-A ", "文档编号", "0007-A", True),
    ("编号：A-1", "编号", "A-1", True),
    ("密级：企业内部", "密级", "企业内部", True),
    ("产品名称：物料甲", "产品名称", "物料甲", False),
    ("数量：7 mg", "数量", "7 mg", False),
    ("编号：A-1。生产使用设备甲。", "编号", "A-1", False),
    ("编号：A-1用于另一产品", "编号", "A-1用于另一产品", False),
    ("编号：", "编号", "", True),
    ("编号", "编号", "", True),
])
def test_fields_keep_exact_source_spans_and_only_exclusive_scalar_records_are_skipped(
    tmp_path, text, label, value, exclusive,
):
    index = index_for(tmp_path, [text])
    fields = extract_attribute_fields(index, property_labels=["数量"])
    assert len(fields) == 1
    field = fields[0]
    assert (field.label, field.value, field.exclusive_record) == (label, value, exclusive)
    assert "".join(index.ir.resolve(ref) for ref in field.label_refs) == label
    assert "".join(index.ir.resolve(ref) for ref in field.value_refs) == value
    assert field_payload(field)["label_refs"] == [ref.model_dump(mode="json")
                                                   for ref in field.label_refs]
    assert fields == extract_attribute_fields(index, property_labels=["数量"])


def test_bare_drug_code_is_not_a_field_and_mixed_record_keeps_entity_discovery(tmp_path):
    index = index_for(tmp_path, ["ABC-001", "编号：A01；设备名称：搅拌罐"])
    fields = extract_attribute_fields(index, property_labels=[])
    assert [(field.label, field.value) for field in fields] == [
        ("编号", "A01"), ("设备名称", "搅拌罐"),
    ]
    assert not any(field.exclusive_record for field in fields)


def test_numbered_property_heading_anchors_the_first_short_value_only(tmp_path):
    index = index_for(tmp_path, ["内部", "不能当作第二个密级值"], heading="1.2 密级")
    fields = extract_attribute_fields(index, property_labels=[])
    assert len(fields) == 1
    assert (fields[0].label, fields[0].value, fields[0].exclusive_record) == ("密级", "内部", True)
    assert index.ir.resolve(fields[0].label_refs[0]) == "密级"
    assert index.ir.resolve(fields[0].value_refs[0]) == "内部"


@pytest.mark.parametrize("heading_label", [False, True])
def test_explicit_split_group_replaces_missing_field_with_exact_value(tmp_path, heading_label):
    document = Document()
    document.add_heading("报告属性", 1)
    document.add_heading("1.1 文档编号：" if heading_label else "字段", 2)
    if not heading_label:
        document.add_paragraph("文档编号：")
    document.add_heading("数值", 2)
    document.add_paragraph("0007-A")
    path = tmp_path / "split-field.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    groups = build_reading_groups(index, field_labels=["文档编号"])
    assert len(groups) == 1 and groups[0].reasons[0].reason_code == "split_field_value"
    without_group = extract_attribute_fields(index, property_labels=[])
    assert all(not field.value for field in without_group)
    fields = extract_attribute_fields(index, property_labels=[], reading_groups=groups)
    assert len(fields) == 1
    field = fields[0]
    assert (field.label, field.value, field.record_id) == (
        "文档编号", "0007-A", index.records[-1].record_id,
    )
    assert field.exclusive_record
    assert index.ir.resolve(field.label_refs[0]) == "文档编号"
    assert index.ir.resolve(field.value_refs[0]) == "0007-A"
    assert field.label_refs[0].section_node_id != field.value_refs[0].section_node_id
    changed_reason = replace(groups[0], reasons=tuple(
        replace(link, reason_code="explicit_continuation") for link in groups[0].reasons
    ))
    assert extract_attribute_fields(index, property_labels=[], reading_groups=[changed_reason]) == (
        without_group
    )


def test_split_group_does_not_choose_one_of_multiple_value_records(tmp_path):
    document = Document()
    document.add_heading("报告属性", 1)
    document.add_heading("字段", 2)
    document.add_paragraph("文档编号：")
    document.add_heading("数值", 2)
    document.add_paragraph("0007-A")
    document.add_paragraph("0008-B")
    path = tmp_path / "ambiguous-values.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    groups = build_reading_groups(index, field_labels=["文档编号"])
    assert len(groups) == 1 and groups[0].reasons[0].reason_code == "split_field_value"
    fields = extract_attribute_fields(index, property_labels=[], reading_groups=groups)
    assert len(fields) == 1 and fields[0].value == "" and fields[0].value_refs == ()


def test_table_header_columns_do_not_cross_assign_values_or_claim_exclusivity(tmp_path):
    index = index_for(tmp_path, table=[
        ["设备名称", "编号"], ["设备甲", "001-A"], ["设备乙", "001-B"],
    ])
    fields = extract_attribute_fields(index, property_labels=["设备名称"])
    codes = [field for field in fields if field.label == "编号"]
    assert {field.value for field in codes} == {"001-A", "001-B"}
    assert len({field.field_id for field in codes}) == 2
    assert not any(field.exclusive_record for field in fields)
    for field in codes:
        assert field.label_refs[0].column_index == field.value_refs[0].column_index == 1


def test_merged_physical_value_is_one_field_across_logical_rows(tmp_path):
    document = Document()
    document.add_heading("设备", 1)
    table = document.add_table(rows=3, cols=2)
    for row, values in enumerate([
        ["设备名称", "编号"], ["设备甲", "001"], ["设备乙", ""],
    ]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    table.cell(1, 1).merge(table.cell(2, 1))
    path = tmp_path / "merged-fields.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    assert len(index.records) == 2
    fields = extract_attribute_fields(index, property_labels=[])
    codes = [field for field in fields if field.label == "编号"]
    assert len(codes) == 1 and codes[0].value == "001"
    value_ref = codes[0].value_refs[0]
    assert len(index.records_by_evidence[value_ref.evidence_id]) == 2
    assert not codes[0].exclusive_record


def test_property_card_distinguishes_generic_and_document_number_labels(tmp_path):
    index = index_for(tmp_path, ["编号：001", "文档编号：002", "密级：内部"])
    generic, specific, missing = extract_attribute_fields(index, property_labels=[])
    assert set(compile_card(generic).class_iris) == {REPORT, EQUIPMENT, BATCH}
    specific_card = compile_card(specific)
    assert specific_card.class_iris == [REPORT]
    assert [p.iri for p in specific_card.predicates] == ["urn:documentNumber"]
    gap = compile_card(missing)
    assert gap.class_iris == [REPORT] and gap.predicates == []
    assert compile_card(generic, (BATCH, REPORT, EQUIPMENT)) == compile_card(generic)
    relabelled = extract_attribute_fields(index, property_labels=["新增属性"])
    assert generic.field_id == relabelled[0].field_id


def test_candidate_options_preserve_ambiguity_and_filter_only_by_relevant_sources(tmp_path):
    index = index_for(tmp_path, ["编号：001", "设备甲"])
    field = extract_attribute_fields(index, property_labels=[])[0]
    entities = [entity(index, "root", REPORT, root=True),
                entity(index, "near", EQUIPMENT, index.records[1]),
                entity(index, "far", EQUIPMENT, index.records[-1])]
    card = compile_card(field)
    options, reason, signature = attribute_options(field, card, index, entities)
    assert reason is None
    assert {item["subject_ref"]["id"] for item in options} == {"root", "near"}
    reordered = attribute_options(field, card, index, list(reversed(entities)))
    assert reordered == (options, reason, signature)
    assert attribute_options(field, card, index, entities[:2]) == (options, reason, signature)
    contextual = attribute_options(field, card, index, entities,
                                   context_record_ids=[index.records[-1].record_id])
    assert {item["subject_ref"]["id"] for item in contextual[0]} == {"root", "near", "far"}
    assert contextual[2] != signature
    overflow, reason, _ = attribute_options(field, card, index, entities, max_candidates=1)
    assert overflow == options and reason == "attribute_candidate_capacity"
    changed = [entities[0], entity(index, "near", EQUIPMENT, index.records[1], revision=2)]
    assert attribute_options(field, card, index, changed)[2] != signature


def test_signature_changes_only_with_relevant_source_anchors(tmp_path):
    index = index_for(tmp_path, ["编号：001", "设备甲在此；另称设备A"])
    field = extract_attribute_fields(index, property_labels=[])[0]
    card = compile_card(field)
    near = entity(index, "near", EQUIPMENT, index.records[1])

    def with_sources(*anchors):
        payload = near.model_dump(mode="json", exclude={"content_hash"})
        payload["source_refs"] = [anchor.model_dump(mode="json") for anchor in anchors]
        return EntityDependencyView(**payload, content_hash=evidence_hash(payload))

    unit = index.records[1].source_units[0]
    first = index.ir.anchor(unit.evidence_id, 0, 3)
    alias = index.ir.anchor(unit.evidence_id, unit.text.index("设备A"), len(unit.text))
    far_unit = index.records[-1].source_units[0]
    unrelated = index.ir.anchor(far_unit.evidence_id, 0, len(far_unit.text))
    original = attribute_options(field, card, index, [with_sources(first)])
    relevant = attribute_options(field, card, index, [with_sources(first, alias)])
    assert original[:2] == relevant[:2]
    assert original[2] != relevant[2]
    assert attribute_options(field, card, index, [with_sources(first, unrelated)]) == original
    assert attribute_options(field, card, index, [with_sources(alias, first)]) == relevant


def test_local_alias_adds_candidate_and_binding_signature_without_rewriting_entity(tmp_path):
    index = index_for(tmp_path, ["编号：001", "本设备又称设备乙"])
    field = extract_attribute_fields(index, property_labels=[])[0]
    card = compile_card(field)
    owner = entity(index, "far", EQUIPMENT, index.records[-1])
    original_owner = owner.model_dump_json()
    alias_unit = index.records[1].source_units[0]
    alias_ref = index.ir.anchor(alias_unit.evidence_id, 0, len(alias_unit.text))
    resolution = {
        "entity_ref": owner.entity_ref.model_dump(mode="json"),
        "binding_ref": {"id": "binding", "revision": 1},
        "source_refs": [alias_ref.model_dump(mode="json")],
    }
    initial = attribute_options(field, card, index, [owner])
    assert initial[0] == [] and initial[1] == "attribute_subject_missing"
    matched = attribute_options(field, card, index, [owner], reference_resolutions=[resolution])
    assert matched[1] is None and matched[0][0]["subject_ref"]["id"] == "far"
    assert set(matched[0][0]) == {"subject_ref", "class_iri", "predicate_iri"}
    assert matched[2] != initial[2]
    changed = {**resolution, "binding_ref": {"id": "binding", "revision": 2}}
    assert attribute_options(field, card, index, [owner], reference_resolutions=[changed])[2] != (
        matched[2]
    )
    assert owner.model_dump_json() == original_owner

    unrelated = {**resolution, "source_refs": [
        anchor.model_dump(mode="json") for anchor in owner.source_refs
    ]}
    wrong_revision = {**resolution, "entity_ref": {"id": "far", "revision": 2}}
    unproved = {**resolution, "binding_ref": None}
    for ignored in (unrelated, wrong_revision, unproved):
        assert attribute_options(field, card, index, [owner], reference_resolutions=[ignored]) == (
            initial
        )
        assert attribute_options(
            field, card, index, [owner], reference_resolutions=[ignored, resolution],
        ) == matched
        assert attribute_options(
            field, card, index, [owner], reference_resolutions=[resolution, ignored],
        ) == matched


def test_missing_value_ontology_and_subject_are_distinct_and_roots_are_not_auto_selected(tmp_path):
    index = index_for(tmp_path, ["文档编号：001", "密级：内部"])
    field, gap_field = extract_attribute_fields(index, property_labels=[])
    card = compile_card(field)
    assert attribute_options(field, card, index, [entity(index, "device", EQUIPMENT)])[1] == (
        "attribute_subject_missing"
    )
    assert attribute_options(gap_field, compile_card(gap_field), index, [
        entity(index, "root", REPORT, root=True),
    ])[1] == "ontology_property_missing"
    missing = replace(field, value="", value_refs=())
    assert attribute_options(missing, card, index, [])[1] == "attribute_value_missing"
    options, reason, _ = attribute_options(field, card, index, [
        entity(index, "root", REPORT, root=True),
    ])
    assert reason is None and len(options) == 1
    assert set(options[0]) == {"subject_ref", "class_iri", "predicate_iri"}
