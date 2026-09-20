"""Merged Word cells supply every covered row without inventing source evidence."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.extraction.docx_structure import parse_docx_structure
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.parser import parse_word
from app.services.extraction.word_analysis import analyze_word_core

HEADERS = ["名称", "物料酸碱性", "溶剂", "温度", "溶解度"]
VALUES = [
    ["1234-3", "中性", "THF", "25℃±2℃", "易溶"],
    ["1234-3", "中性", "正庚烷", "10~30℃", "微溶"],
    ["", "", "水", "", ""],
]


def _solubility_table(document, *, omitted_edges=False):
    offset = int(omitted_edges)
    table = document.add_table(rows=4, cols=len(HEADERS) + 2 * offset)
    for ri, values in enumerate([HEADERS, *VALUES]):
        for ci, value in enumerate(values, start=offset):
            table.cell(ri, ci).text = value
    for ci, value in enumerate(VALUES[0][:2], start=offset):
        table.cell(1, ci).merge(table.cell(2, ci)).text = value
    if omitted_edges:
        table.cell(0, 0).text = "序号"
        table.cell(0, 6).text = "备注"
        for row in table.rows[1:]:
            tr = row._tr
            tr.remove(tr.tc_lst[-1])
            tr.remove(tr.tc_lst[0])
            for tag in ("w:gridBefore", "w:gridAfter"):
                element = OxmlElement(tag)
                element.set(qn("w:val"), "1")
                tr.get_or_add_trPr().append(element)
    return table


@pytest.mark.parametrize("omitted_edges", [False, True])
def test_vertical_merges_fill_rows_and_keep_column_alignment(tmp_path, omitted_edges):
    document = Document()
    _solubility_table(document, omitted_edges=omitted_edges)
    path = tmp_path / "solubility.docx"
    document.save(path)

    parsed = parse_docx_structure(path).tables[0]
    expected = [dict(zip(HEADERS, values)) for values in VALUES]
    if omitted_edges:
        expected = [{"序号": "", **row, "备注": ""} for row in expected]
        assert all(row[0] is None and row[-1] is None for row in parsed.grid[1:])
    assert parsed.rows == expected
    assert parsed.cells[1:] == [list(row.values()) for row in expected]
    assert parsed.row_indices == [1, 2, 3]
    assert [section["content"] for section in parse_word(path)] == expected
    mapped = parse_word(path, {"名称": "materialName"})
    assert [section["content"]["materialName"] for section in mapped] == ["1234-3", "1234-3", ""]


def test_filled_rows_share_source_evidence_and_keep_solvent_conditions_separate(tmp_path):
    document = Document()
    _solubility_table(document)
    path = tmp_path / "evidence.docx"
    document.save(path)

    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    first, second, blank = [record for record in index.records if record.kind == "table_row"]
    for record, values in zip((first, second, blank), VALUES):
        assert [unit.text for unit in record.source_units] == [value for value in values if value]
    for value in ("1234-3", "中性"):
        units = [unit for unit in analysis.ir.evidence_units if unit.text == value]
        assert len(units) == 1
        unit = units[0]
        assert unit in first.source_units and unit in second.source_units
        assert unit not in blank.source_units
        assert analysis.ir.resolve(analysis.ir.anchor(unit.evidence_id)) == value
        cell = index.tables.cells[unit.source_cell_id]
        assert cell["row_index"] == 1 and cell["row_span"] == 2


def test_combined_row_and_column_merge_stops_at_its_boundary(tmp_path):
    document = Document()
    table = document.add_table(rows=5, cols=3)
    for ci, label in enumerate(("字段甲", "字段乙", "条件")):
        table.cell(0, ci).text = label
    for ri, value in enumerate(("条件一", "条件二", "条件三", "条件四"), start=1):
        table.cell(ri, 2).text = value
    merged = table.cell(1, 0).merge(table.cell(3, 1))
    merged.text = "  共享值  "
    merged.add_paragraph(" 补充说明 ")
    table.cell(4, 0).text = "独立值"
    path = tmp_path / "rectangular-merge.docx"
    document.save(path)

    parsed = parse_docx_structure(path).tables[0]
    assert parsed.cells[1:] == [
        ["共享值 补充说明", "共享值 补充说明", "条件一"],
        ["共享值 补充说明", "共享值 补充说明", "条件二"],
        ["共享值 补充说明", "共享值 补充说明", "条件三"],
        ["独立值", "", "条件四"],
    ]
    origin = next(cell for cell in parsed.source_cells if cell["row_span"] == 3)
    assert origin["column_span"] == 2
    assert all(row[:2] == [origin["cell_id"]] * 2 for row in parsed.grid[1:4])


@pytest.mark.parametrize("omitted_edges", [False, True])
def test_nested_tables_fill_merged_rows_with_their_own_headers(tmp_path, omitted_edges):
    document = Document()
    outer = document.add_table(rows=2, cols=1)
    outer.cell(0, 0).text = "溶解度试验"
    _solubility_table(outer.cell(1, 0), omitted_edges=omitted_edges)
    path = tmp_path / "nested.docx"
    document.save(path)

    parsed = parse_docx_structure(path).tables[0]
    nested = next(
        block["table"] for cell in parsed.source_cells for block in cell["blocks"]
        if block["kind"] == "table"
    )
    assert [[row[header] for header in HEADERS] for row in nested["rows"]] == VALUES
    assert nested["table_path"] == ["table:0", "cell:1:0", "table:0"]
    assert parsed.cells[1] == [""]
