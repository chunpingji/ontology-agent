"""Real frozen Word IR through the four initial local tool handlers."""

import json
from dataclasses import replace

import pytest
from docx import Document

from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided import tool_runtime as runtime
from app.services.extraction.ontology_guided.claim_protocol import (
    EntityDependencyView,
    EntityProposal,
    ExtractionProfile,
    PropertyProposal,
    Quote,
    RelationProposal,
    VerificationTargetSpec,
    claim_content_hash,
)
from app.services.extraction.ontology_guided.context import ContextFragment, assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    EdgeSpec,
    LocalMenu,
    SlotSpec,
    SubjectRef,
    TraversalScope,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.mentions import MentionRegistry
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.ontology_guided.tool_contracts import (
    EvidenceData,
    ToolCall,
    ToolErrorResult,
    ToolObservation,
    ToolResult,
)
from app.services.llm.model_runtime import ModelCancelled


@pytest.fixture
def tool_context(tmp_path):
    doc = Document()
    doc.add_paragraph("First alpha; second alpha.")
    table = doc.add_table(rows=3, cols=2)
    for row, values in enumerate([("Entity", "Mass (mg/kg/day)"), ("A", "3.8–6.6"), ("B", "7.5")]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    path = tmp_path / "source.docx"
    doc.save(path)
    index = RecordIndex(build_document_ir(path, parse_docx_structure(path)))
    units = {unit.text: unit for unit in index.ir.evidence_units}
    value = units["3.8–6.6"]
    record = index.records_by_evidence[value.evidence_id][0]
    subject = SubjectRef(entity_id="entity-a", revision=1, class_iri="urn:Entity")
    slot = SlotSpec(
        iri="urn:mass",
        label="Mass",
        datatype_iris=["http://www.w3.org/2001/XMLSchema#decimal"],
        canonical_unit="mg/kg/day",
    )
    menu = LocalMenu(
        menu_id="menu", ontology_snapshot_id="ontology", subject=subject, properties=[slot]
    )
    task = RecognitionTask.create(
        subject=subject,
        predicate_iri=slot.iri,
        predicate_kind="property",
        record_id=record.record_id,
        phase=1,
        hop=0,
        dependency_hash="dep",
    )
    target = VerificationTarget.create(
        run_fingerprint="run",
        claim_ref=VersionedRef(id="placeholder", revision=1),
        task_id=task.task_id,
        check_kind="predicate_entailment",
        document_context=DocumentContext(
            document_hash=index.ir.document_hash,
            document_class_iri="urn:Document",
            root_ref=VersionedRef(id="root", revision=1),
        ),
        subject_ref=subject,
        predicate_iri=slot.iri,
        ontology_hash="ontology",
        source_scope_hash="scope",
        context_hash="initial",
    )
    owner = index.ir.anchor(units["A"].evidence_id, 0, 1)
    context = assemble_context(
        target,
        record.record_id,
        index,
        subject_evidence_refs=[owner],
        subject_label="A",
        predicate=slot,
        repair_enabled=True,
    )
    proposal = EntityProposal(
        local_id="s",
        class_iri="urn:Entity",
        representation="mention",
        mentions=[Quote(evidence_id=owner.evidence_id, text="A", context_text=None)],
        record_components=[],
        identifier_claims=[],
    )
    dependency = dict(
        entity_ref=VersionedRef(id="entity-a", revision=1),
        class_iri="urn:Entity",
        grounding_kind="mention",
        root_origin=None,
        proposal=proposal,
        source_refs=[owner],
        dependency_refs=[],
    )
    entity = EntityDependencyView(**dependency, content_hash=evidence_hash(dependency))
    return runtime.ToolContext(
        task=task,
        context=context,
        index=index,
        menu=menu,
        profile=ExtractionProfile(),
        scope=TraversalScope.create([]),
        stage="verification",
        recovery_kind="none",
        frozen_claims={},
        local_ref_map={"s": entity.entity_ref},
        entity_dependencies={"s": entity},
        limits=runtime.ToolLimits(max_result_tokens=10000),
        measure_result_tokens=lambda text: len(text),
    )


def call(name, **arguments):
    return ToolCall(call_id="call-1", name=name, arguments_json=json.dumps(arguments))


def add_property(ctx, *, value_text="3.8–6.6", unit_text="mg/kg/day"):
    units = {unit.text: unit for unit in ctx.index.ir.evidence_units}

    def quote(source, text):
        return Quote(evidence_id=units[source].evidence_id, text=text, context_text=None)

    proposal = PropertyProposal(
        local_id="p",
        subject_id="s",
        predicate_iri="urn:mass",
        value_quote=quote(value_text, value_text),
        field_support=[quote("Mass (mg/kg/day)", "Mass")],
        unit_support=[quote("Mass (mg/kg/day)", unit_text)],
        qualifiers={
            "polarity": "affirmed",
            "modality": "asserted",
            "condition_support": [],
            "scope_qualifiers": [],
        },
        bridge_kind="role_mapped_table",
        bridge_ref_ids=[],
    )
    target = VerificationTargetSpec(
        target_id="target-p",
        target_kind="property",
        claim_ref=VersionedRef(id="claim-p", revision=1),
        payload=proposal,
        scope=ctx.scope,
        dependency_refs=[ctx.local_ref_map["s"]],
        required_facets=["value"],
        content_hash=claim_content_hash(
            "property",
            proposal,
            ctx.scope,
            [ctx.local_ref_map["s"]],
        ),
    )
    return replace(ctx, frozen_claims={"claim-p": target})


def test_card_compiles_only_current_subject_and_authorized_predicate(tool_context):
    ctx = tool_context
    result = runtime.dispatch_tool(
        call("get_schema_card", subject_id="entity-a", predicate_iri="urn:mass"), ctx
    )
    assert result.status == "ok" and result.data.cards[0].predicates[0].iri == "urn:mass"
    foreign = runtime.dispatch_tool(
        call("get_schema_card", subject_id="foreign", predicate_iri=None), ctx
    )
    assert foreign.status == "blocked" and foreign.issues[0].code == "reference_outside_scope"


def test_inspection_keeps_table_coordinates_headers_and_evidence_roles(tool_context):
    ctx = tool_context
    result = runtime.dispatch_tool(
        call(
            "inspect_evidence", evidence_ids=[f.anchor.evidence_id for f in ctx.context.fragments]
        ),
        ctx,
    )
    assert result.status == "ok", result
    units = {unit.text: unit for unit in result.data.units}
    assert units["3.8–6.6"].table.logical_rows == [1]
    assert units["3.8–6.6"].table.logical_columns == [1]
    assert units["3.8–6.6"].fact_eligible
    header = units["Mass (mg/kg/day)"]
    assert header.role == "header" and not header.fact_eligible
    assert header.evidence_id in units["3.8–6.6"].table.column_header_refs
    assert "B" not in units


def test_inspect_permission_is_checked_before_reading_any_requested_source(tool_context):
    authorized = tool_context.context.fragments[0].anchor.evidence_id
    result = runtime.dispatch_tool(
        call("inspect_evidence", evidence_ids=[authorized, "foreign"]), tool_context
    )
    assert isinstance(result, ToolResult[EvidenceData]) and result.data is None
    assert result.issues[0].field_path == "/evidence_ids/1"
    assert not result.evidence_refs


def test_anchor_uses_same_physical_identity_as_ner_without_registering(tool_context):
    ctx = tool_context
    unit = next(u for u in ctx.index.ir.evidence_units if u.text == "A")
    result = runtime.dispatch_tool(
        call("resolve_source_anchor", evidence_id=unit.evidence_id, quote="A", context_text=None),
        ctx,
    )
    assert result.status == "ok"
    registry = MentionRegistry(ctx.index.ir, ctx.index)
    mention = registry.register(
        evidence_id=unit.evidence_id,
        start=0,
        end=1,
        text="A",
        record_view_ref=ctx.index.record_views_by_id[ctx.context.record_id].record_view_id,
    )
    assert result.data.mention_ref == mention.mention_id
    assert not ctx.frozen_claims


@pytest.mark.parametrize("context_text", ["", "unrelated text"])
def test_anchor_context_error_identifies_the_field_without_rewriting_it(
    tool_context, context_text,
):
    unit = next(u for u in tool_context.index.ir.evidence_units if u.text == "A")
    request = call(
        "resolve_source_anchor", evidence_id=unit.evidence_id, quote="A",
        context_text=context_text,
    )
    result = runtime.dispatch_tool(request, tool_context)
    assert result.status == "blocked" and result.data is None
    assert result.issues[0].code == "citation_quote_not_in_source"
    assert result.issues[0].field_path == "/context_text"
    assert json.loads(request.arguments_json)["context_text"] == context_text
    assert not tool_context.registered_mentions


def test_binding_keeps_complete_interval_and_does_not_claim_semantic_support(tool_context):
    result = runtime.dispatch_tool(
        call("check_claim_binding", claim_id="claim-p"), add_property(tool_context)
    )
    assert result.status == "ok" and result.data.validation_status == "passed", result
    assert result.data.source_unit == "mg/kg/day"
    assert "normalized_literal" not in result.data.model_dump()
    assert "semantic_status" not in result.data.model_dump()


def test_controller_binding_is_not_a_model_context_return(tool_context):
    ctx = add_property(tool_context)
    ctx = replace(ctx, limits=runtime.ToolLimits(max_result_tokens=1))
    request = call("check_claim_binding", claim_id="claim-p")
    model = runtime.dispatch_tool(request, ctx)
    assert model.status == "blocked" and model.issues[0].code == "result_budget_exceeded"
    controller = runtime.dispatch_tool(request, ctx, caller="controller")
    assert controller.status == "ok" and controller.data.validation_status == "passed"
    refs = [ref.model_dump_json() for ref in controller.data.resolved_role_refs]
    assert len(refs) == len(set(refs))


def test_binding_rejects_truncated_header_unit(tool_context):
    result = runtime.dispatch_tool(
        call("check_claim_binding", claim_id="claim-p"), add_property(tool_context, unit_text="mg")
    )
    assert result.data.validation_status == "failed"
    assert "unit_quote_partial" in {issue.code for issue in result.data.issues}


@pytest.mark.parametrize(
    "name,arguments,code,path",
    [
        ("unknown", "{}", "unknown_tool", None),
        ("inspect_evidence", "{", "invalid_tool_arguments", None),
        ("inspect_evidence", '{"evidence_ids":[1]}', "invalid_tool_arguments", "/evidence_ids/0"),
        (
            "inspect_evidence",
            '{"evidence_ids":[],"owner":true}',
            "invalid_tool_arguments",
            "/owner",
        ),
        ("inspect_evidence", '{"evidence_ids":[]}', "invalid_tool_arguments", "/evidence_ids"),
        ("validate_graph", '{"claim_id":"p","shape_profile_id":"x"}', "tool_not_allowed", None),
    ],
)
def test_raw_call_errors_are_actionable_without_rewriting_input(
    tool_context, name, arguments, code, path
):
    original = ToolCall(call_id="pair-me", name=name, arguments_json=arguments)
    result = runtime.dispatch_tool(original, tool_context)
    assert isinstance(result, ToolErrorResult) and result.data is None
    assert result.issues[0].code == code and result.issues[0].field_path == path
    assert original.name == name and original.arguments_json == arguments
    output = runtime.to_function_call_output(original.call_id, result)
    assert output["call_id"] == "pair-me" and json.loads(output["output"])["data"] is None


def test_model_tools_are_advertised_with_reference_gates(tool_context):
    ctx = replace(tool_context, stage="discovery")
    result = runtime.dispatch_tool(
        call("propose_mentions", evidence_ids=["x"], schema_card_id="card"), ctx
    )
    assert result.status == "blocked" and result.issues[0].code == "reference_outside_scope"
    names = {item["name"] for item in runtime.build_tool_definitions(ctx, ctx.stage)}
    assert names == {"get_schema_card", "inspect_evidence", "resolve_source_anchor"}
    configured = replace(
        ctx, mention_extractor=object(), ontology_snapshot=object(),
        instance_reader=object(), external_source_ids=("explicit-source",),
        metadata=object(), base_context=ctx.context, target_seed=ctx.context.target,
        run_fingerprint="frozen", subject_node=object(),
    )
    names = {item["name"] for item in runtime.build_tool_definitions(configured, ctx.stage)}
    assert names == {
        "get_schema_card", "inspect_evidence", "resolve_source_anchor", "propose_mentions",
        "query_instances", "retrieve_evidence",
    }


def test_result_limit_blocks_instead_of_truncating_a_record(tool_context):
    ctx = replace(tool_context, limits=runtime.ToolLimits(max_result_tokens=1))
    result = runtime.dispatch_tool(
        call(
            "inspect_evidence",
            evidence_ids=[
                ctx.context.fragments[0].anchor.evidence_id,
            ],
        ),
        ctx,
    )
    assert result.status == "blocked" and result.issues[0].code == "result_budget_exceeded"


def test_cancellation_is_not_reported_as_tool_failure(tool_context):
    def cancelled():
        raise ModelCancelled()

    ctx = replace(tool_context, check_cancelled=cancelled)
    with pytest.raises(ModelCancelled):
        runtime.dispatch_tool(
            call("get_schema_card", subject_id="entity-a", predicate_iri=None), ctx
        )


def updated_claim(ctx, payload):
    previous = ctx.frozen_claims["claim-p"]
    target = previous.model_copy(
        update={
            "payload": payload,
            "content_hash": claim_content_hash(
                previous.target_kind, payload, ctx.scope, previous.dependency_refs
            ),
        }
    )
    return replace(ctx, frozen_claims={"claim-p": target})


def test_scalar_excerpt_cannot_silently_discard_the_rest_of_an_interval(tool_context):
    ctx = add_property(tool_context)
    payload = ctx.frozen_claims["claim-p"].payload.model_copy(deep=True)
    payload.value_quote.text = "3.8"
    result = runtime.dispatch_tool(
        call("check_claim_binding", claim_id="claim-p"), updated_claim(ctx, payload)
    )
    assert result.data.validation_status == "failed"
    assert "quantity_quote_incomplete" in {issue.code for issue in result.data.issues}


def test_binding_does_not_infer_a_missing_unit_from_the_card(tool_context):
    ctx = add_property(tool_context)
    payload = ctx.frozen_claims["claim-p"].payload.model_copy(update={"unit_support": []})
    result = runtime.dispatch_tool(
        call("check_claim_binding", claim_id="claim-p"), updated_claim(ctx, payload)
    )
    assert result.data.validation_status == "incomplete" and result.data.source_unit is None


def test_authorized_other_row_still_does_not_belong_to_the_subject(tool_context):
    ctx = tool_context
    fragments = list(ctx.context.fragments)
    for unit in ctx.index.ir.evidence_units:
        if unit.text in ("B", "7.5"):
            fragments.append(
                ContextFragment(
                    anchor=ctx.index.ir.anchor(unit.evidence_id, 0, len(unit.text)),
                    text=unit.text,
                    purpose="target",
                    fact_eligible=True,
                )
            )
    ctx = replace(ctx, context=ctx.context.model_copy(update={"fragments": fragments}))
    result = runtime.dispatch_tool(
        call("check_claim_binding", claim_id="claim-p"), add_property(ctx, value_text="7.5")
    )
    assert result.data.validation_status == "failed"
    assert "owner_row_mismatch" in {issue.code for issue in result.data.issues}


def test_stale_source_and_stale_claim_dependencies_cannot_be_consumed(tool_context):
    ctx = add_property(tool_context)
    stale_map = {"s": ctx.local_ref_map["s"].model_copy(update={"revision": 2})}
    result = runtime.dispatch_tool(
        call("check_claim_binding", claim_id="claim-p"), replace(ctx, local_ref_map=stale_map)
    )
    assert result.issues[0].code == "reference_version_mismatch"
    source = ctx.context.model_copy(deep=True)
    source.fragments[0].text += "forged"
    result = runtime.dispatch_tool(
        call(
            "inspect_evidence",
            evidence_ids=[
                source.fragments[0].anchor.evidence_id,
            ],
        ),
        replace(ctx, context=source),
    )
    assert result.issues[0].code == "reference_version_mismatch" and result.data is None


def test_self_relation_is_mechanically_checked_without_blanket_rejection(tool_context):
    ctx = tool_context
    edge = EdgeSpec(iri="urn:relatesTo", label="relates to", range_class_iris=["urn:Entity"])
    source = ctx.entity_dependencies["s"].proposal.mentions[0]
    payload = RelationProposal(
        local_id="relation",
        subject_id="s",
        predicate_iri=edge.iri,
        object_ids=["s"],
        selection="all",
        bridge_support=[source],
        selection_support=[],
        bridge_kind="explicit_assertion",
        bridge_ref_ids=[],
        qualifiers={
            "polarity": "affirmed",
            "modality": "asserted",
            "condition_support": [source],
            "scope_qualifiers": [],
        },
    )
    target = VerificationTargetSpec(
        target_id="self-target",
        target_kind="relation",
        claim_ref=VersionedRef(id="self", revision=1),
        payload=payload,
        scope=ctx.scope,
        dependency_refs=[ctx.local_ref_map["s"]],
        required_facets=["predicate"],
        content_hash=claim_content_hash(
            "relation",
            payload,
            ctx.scope,
            [ctx.local_ref_map["s"]],
        ),
    )
    ctx = replace(
        ctx,
        menu=ctx.menu.model_copy(update={"relationships": [edge]}),
        frozen_claims={"self": target},
    )
    result = runtime.dispatch_tool(call("check_claim_binding", claim_id="self"), ctx)
    assert result.data.validation_status == "passed"
    # The label "A" is neither a condition regex nor relationship proof. Only
    # quote replay was checked; the independent verifier must assess its role.
    assert not result.data.issues


def test_valid_arguments_keep_their_registered_result_type_after_authorization_failure(
    tool_context,
):
    request = call("inspect_evidence", evidence_ids=["foreign"])
    result = runtime.dispatch_tool(request, tool_context)
    observation = ToolObservation(
        request_attempt=1,
        call=request,
        parsed_arguments=runtime.parse_tool_arguments(request.name, request.arguments_json),
        result_ref="confirmed-error",
        result=result,
    )
    assert observation.parsed_arguments is not None
    assert observation.result.data is None


def test_failed_validation_is_not_displayed_as_success(tool_context):
    from app.services.llm.model_runtime import model_scope

    ctx = add_property(tool_context, unit_text="mg")
    events = []
    with model_scope(on_harness_event=lambda kind, data: events.append((kind, data))):
        result = runtime.dispatch_tool(call("check_claim_binding", claim_id="claim-p"), ctx)
    assert result.status == "ok" and result.data.validation_status == "failed"
    assert events[-1][0] == "operation_end"
    assert events[-1][1]["status"] == "blocked"
    assert events[-1][1]["result"] == result.model_dump(mode="json")
