"""Admission tests cover omitted tails, scope and technical-failure semantics."""

from __future__ import annotations

import pytest
from docx import Document

from app.services.extraction.ontology_guided.contracts import SlotSpec, SubjectRef
from app.services.extraction.ontology_guided.heuristic_search import (
    HeuristicSearchIndex,
    HeuristicSearchPolicy,
    HeuristicSlotSearch,
)
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.retrieval import plan_slot, validate_record_universe
from app.services.extraction.word_analysis import analyze_word_core


def _search(tmp_path, paragraphs, *, predicate=None, policy=None, root=False, mentions=()):
    document = Document()
    for position, text in enumerate(paragraphs):
        document.add_heading(f"节 {position}", 1)
        document.add_paragraph(text)
    path = tmp_path / "filename-is-not-a-proved-product.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    metadata = prepare_metadata(
        analysis.ir, section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="test", generation_source="structure_only",
    )
    predicate = predicate or SlotSpec(iri="urn:molecularWeight", label="分子量")
    subject = SubjectRef(
        entity_id="subject", revision=1, class_iri="urn:Product", is_document_root=root,
    )
    plan = plan_slot(subject, predicate, index, metadata, ontology_hash="test-ontology")
    return HeuristicSlotSearch(
        plan=plan, predicate=predicate, search_index=HeuristicSearchIndex(index, metadata),
        policy=policy or HeuristicSearchPolicy(), subject_mentions=list(mentions),
        run_fingerprint="test-run",
    )


def _observe_page(search, page, *, supported=False):
    for rid in page.record_ids:
        search.observe(
            rid, "supported" if supported else "unsupported", True,
            "source_supported" if supported else "no_entailed_fact",
            supported_count=int(supported),
        )


def test_clear_field_admits_without_semantic_and_preserves_unchecked_universe(tmp_path):
    search = _search(tmp_path, ["分子量：123.45", "现场背景资料", "另附文件待审"])
    before = search.plan.model_dump(mode="json")
    page = search.next_admission()

    assert page.stage == "H0"
    assert page.source == "heuristic"
    assert len(page.record_ids) == 1
    assert not search.needs_semantic
    assert search.next_admission() is None
    _observe_page(search, page, supported=True)
    assert search.next_admission() is None
    assert search.status == "local_results_only"
    assert search.snapshot()["examined_count"] == 1
    assert search.snapshot()["deferred_count"] == 2
    assert search.plan.model_dump(mode="json") == before
    validate_record_universe(search.plan, search.search_index.index)


def test_multivalue_hits_page_past_first_success_before_deferring_tail(tmp_path):
    search = _search(
        tmp_path, [*[f"设备：装置 {n}" for n in range(11)], "质量审批说明"],
        predicate=SlotSpec(iri="urn:equipment", label="设备"),
        policy=HeuristicSearchPolicy(initial_page_size=2, expanded_page_size=3),
    )
    pages = []
    while page := search.next_admission():
        pages.append(page)
        _observe_page(search, page, supported=True)

    admitted = [rid for page in pages for rid in page.record_ids]
    assert len(admitted) == 11
    assert len(set(admitted)) == 11
    assert len(pages) > 2
    assert {page.stage for page in pages} == {"H0", "H1"}
    assert search.status == "local_results_only"
    assert search.snapshot()["deferred_count"] == 1
    assert not search.needs_semantic


def test_no_direct_hit_expands_registered_alias_before_semantic(tmp_path):
    search = _search(
        tmp_path, ["molecular weight: 321.5", "设备维护说明"],
        predicate=SlotSpec(iri="urn:mass", label="分子量"),
    )
    page = search.next_admission()
    assert page.stage == "H1"
    assert search.search_index.index.by_id[page.record_ids[0]].text.startswith("molecular")
    assert not search.needs_semantic


def test_empty_and_rejected_pages_reach_real_semantic_then_distinct_tail(tmp_path):
    search = _search(
        tmp_path, ["分子量：该字段暂不适用", "补充材料甲", "补充材料乙", "尾页证据"],
        policy=HeuristicSearchPolicy(exploration_page_size=1),
    )
    first = search.next_admission()
    _observe_page(search, first)
    assert search.next_admission() is None
    assert search.needs_semantic
    middle = search.search_index.record_ids[1]
    search.accept_semantic([middle], "committed-epoch-1", committed=True)
    semantic = search.next_admission()
    assert semantic.stage == "H2"
    assert semantic.epoch_id == "committed-epoch-1"
    _observe_page(search, semantic)
    tail = search.next_admission()
    assert tail.stage == "H3"
    assert tail.record_ids == [search.search_index.record_ids[-1]]
    assert not set(tail.record_ids).intersection([*first.record_ids, *semantic.record_ids])
    _observe_page(search, tail)
    assert search.next_admission() is None
    assert search.status == "pass_exhausted"
    assert search.snapshot()["examined_count"] == 3
    assert search.snapshot()["deferred_count"] == 1


def test_no_lexical_candidates_does_not_claim_absence(tmp_path):
    search = _search(tmp_path, ["一般背景", "尾页证据"])
    assert search.next_admission() is None
    assert search.needs_semantic
    assert search.snapshot()["examined_count"] == 0
    assert search.snapshot()["deferred_count"] == 2
    search.accept_semantic([], "empty-committed-epoch", committed=True)
    page = search.next_admission()
    assert page.stage == "H3"
    assert set(page.record_ids) == set(search.search_index.record_ids)


@pytest.mark.parametrize("complete,outcome", [(False, "undetermined"), (True, "not_checked")])
def test_technical_failure_blocks_instead_of_triggering_no_result_expansion(
    tmp_path, complete, outcome,
):
    search = _search(tmp_path, ["分子量：123", "背景材料"])
    page = search.next_admission()
    search.observe(page.record_ids[0], outcome, complete, "model_timeout")
    assert search.next_admission() is None
    assert search.status == "technical_blocked"
    assert not search.needs_semantic
    assert search.snapshot()["examined_count"] == 0


def test_semantic_admission_requires_committed_epoch_and_source_membership(tmp_path):
    search = _search(tmp_path, ["一般背景", "尾页证据"])
    search.next_admission()
    with pytest.raises(ValueError, match="complete committed"):
        search.accept_semantic([], "pending-epoch", committed=False)
    with pytest.raises(ValueError, match="another source"):
        search.accept_semantic(["outside-record"], "epoch", committed=True)
    rid = search.search_index.record_ids[0]
    with pytest.raises(ValueError, match="duplicate"):
        search.accept_semantic([rid, rid], "epoch", committed=True)


def test_valid_empty_discovery_continues_search_instead_of_technical_block(tmp_path):
    search = _search(tmp_path, ["分子量：请参见附录", "附录另有描述"])
    page = search.next_admission()
    search.observe(page.record_ids[0], "not_checked", True, "no_candidate_observed")
    assert search.next_admission() is None
    assert search.needs_semantic
    assert search.snapshot()["examined_count"] == 1


def test_explicitly_disabled_semantic_uses_real_exploration_without_fake_epoch(tmp_path):
    search = _search(tmp_path, ["背景资料", "实际证据"])
    search.next_admission()
    with pytest.raises(ValueError, match="explicitly disabled"):
        search.skip_semantic("model_timeout")
    search.skip_semantic("ranking_policy_deterministic")
    page = search.next_admission()
    assert page.stage == "H3"
    assert page.epoch_id is None
    assert search.snapshot()["semantic_epoch_id"] is None
    assert search.snapshot()["semantic_skip_reason"] == "ranking_policy_deterministic"


def test_root_filename_and_unrelated_subject_mentions_cannot_drive_admission(tmp_path):
    search = _search(
        tmp_path, ["filename-is-not-a-proved-product", "普通背景"],
        root=True, mentions=["filename-is-not-a-proved-product"],
    )
    assert search.subject_mentions == []
    assert search.next_admission() is None
    assert search.needs_semantic
    snapshot = search.snapshot()
    assert snapshot["resume_supported"] is False
    assert snapshot["admitted_count"] == 0


def test_incomplete_active_page_never_admits_following_page(tmp_path):
    search = _search(
        tmp_path, [f"分子量：{n}" for n in range(5)],
        policy=HeuristicSearchPolicy(initial_page_size=2),
    )
    first = search.next_admission()
    search.observe(first.record_ids[0], "unsupported", True, "no_fact")
    assert search.next_admission() is None
    assert search.snapshot()["admitted_count"] == 2
    search.observe(first.record_ids[1], "unsupported", True, "no_fact")
    second = search.next_admission()
    assert second is not None
    assert not set(first.record_ids).intersection(second.record_ids)


def test_existing_core_retry_updates_latest_feedback_without_new_admission(tmp_path):
    search = _search(tmp_path, ["分子量：123", "背景材料"])
    page = search.next_admission()
    rid = page.record_ids[0]
    search.observe(rid, "not_checked", False, "model_timeout", attempt_id="first-attempt")
    assert search.status == "technical_blocked"
    assert search.next_admission() is None
    search.observe(
        rid, "supported", True, "source_supported", supported_count=1,
        attempt_id="core-technical-once",
    )
    search.observe(
        rid, "supported", True, "source_supported", supported_count=1,
        attempt_id="core-technical-once",
    )
    assert search.next_admission() is None
    assert search.status == "local_results_only"
    snapshot = search.snapshot()
    assert snapshot["admitted_count"] == 1
    assert snapshot["examined_count"] == 1
    assert snapshot["supported_output_count"] == 1
    assert len(snapshot["attempt_history"]) == 2
    with pytest.raises(ValueError, match="one attempt"):
        search.observe(
            rid, "unsupported", True, "changed_feedback", attempt_id="core-technical-once",
        )
