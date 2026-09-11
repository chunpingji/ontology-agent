"""Actual transmitted protocol, evidence permissions and paid-stage recovery."""

import json
from copy import deepcopy

import pytest

from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.evidence_groups import evidence_content_hash
from app.services.extraction.ontology_guided.ontology_plan import compile_local_menu
from app.services.extraction.ontology_guided.repair_adapter import EvidenceRepairAdapter
from tests.test_extraction.test_joint_evidence_validation import _fixture, _response
from tests.test_extraction.test_semantic_graph_closure import _ontology as base_ontology


def _ontology():
    ontology = base_ontology()
    for cls in ontology.classes.values():
        for slot in cls.declared_properties:
            slot.datatype_iris = ["http://www.w3.org/2001/XMLSchema#string"]
    return ontology


def repaired_response(request):
    raw = _response(request)
    if request["stage"] == "discovery":
        for proposal in raw["proposals"]:
            quote = proposal.get("object_quote") or proposal["value_quote"]
            binding = next(
                (
                    b
                    for b in request["field_bindings"]
                    if any(r["evidence_id"] == quote["evidence_id"] for r in b["target_value_refs"])
                ),
                None,
            )
            if binding:
                proposal["field_binding_id"] = binding["field_binding_id"]
    else:
        for review in raw["verifications"]:
            review["subject_binding"] = {
                "verdict": review.pop("subject_binding_verdict"),
                "support": review.pop("subject_support"),
            }
            review["subject_binding"]["local_support"] = (
                [] if request["subject"]["is_document_root"]
                else deepcopy(review["subject_binding"]["support"])
            )
            if not request["subject"]["is_document_root"]:
                # This controlled successful proof cites both occurrences;
                # the wrong-owner case below deliberately omits the original.
                review["subject_binding"]["support"].extend(
                    {"evidence_id": r["evidence_id"]}
                    for r in request["subject_evidence_refs"]
                )
            review["field_role_support"] = deepcopy(review["predicate_support"])
            review["bridge_support"] = deepcopy(review["predicate_support"])
            for binding in request["field_bindings"]:
                review["field_role_support"].extend(
                    {"evidence_id": r["evidence_id"]} for r in binding["label_refs"]
                )
            if request["predicate"]["kind"] == "property":
                review.pop("type_support")
                review.pop("type_verdict")
    return raw


def test_supplement_reuses_discovery_but_rechecks_all_facets(tmp_path, monkeypatch):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)
    calls, stages = [], []

    def respond(_client, *, user, schema, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        assert "adjacent_only" not in json.dumps(schema)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    adapter = EvidenceRepairAdapter(object(), model_identity="repair-fixture")
    local = assemble_context(target, task.record_id, index, repair_enabled=True)
    local.bind_protocol_hook(lambda state: stages.append(deepcopy(state)))
    first = adapter.inspect(task, local, predicate, menu)
    assert first.semantic_outcome == "undetermined"
    joint = assemble_context(
        target, task.record_id, index, required_context_refs=[binding], repair_enabled=True
    )
    retry = task.model_copy(update={"retry_kind": "positive_evidence:1"})
    result = adapter.verify_existing(retry, joint, predicate, menu, stages[-1])
    assert result.edges[0].policy_eligible
    assert result.model_calls == 1
    assert [r["stage"] for r in calls] == ["discovery", "verification", "verification"]
    assert any((state.get("discovery") or {}).get("proposals") for state in stages)
    assert all(
        not f["fact_eligible"]
        for f in calls[-1]["fragments"]
        if f["text"] == index.ir.resolve(binding)
    )
    repeated = adapter.verify_existing(retry, joint, predicate, menu, joint.protocol_state)
    assert repeated.model_calls == 0
    assert len(calls) == 3


def test_discovery_checkpoint_survives_interruption_before_verification(tmp_path, monkeypatch):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)
    stages, calls = [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request["stage"])
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    adapter = EvidenceRepairAdapter(object(), model_identity="repair-fixture")
    context = assemble_context(
        target, task.record_id, index, required_context_refs=[binding], repair_enabled=True
    )

    def checkpoint(state):
        stages.append(deepcopy(state))
        if state.get("discovery"):
            raise RuntimeError("simulated owner interruption")

    context.bind_protocol_hook(checkpoint)
    with pytest.raises(RuntimeError, match="owner interruption"):
        adapter.inspect(task, context, predicate, menu)
    context.bind_protocol_hook(lambda _state: None)
    context.protocol_state = stages[-1]
    result = adapter.inspect(task, context, predicate, menu)
    assert result.edges[0].policy_eligible
    assert calls == ["discovery", "verification"]


def test_evidence_hash_ignores_task_identity_and_fragment_order(tmp_path):
    _analysis, index, task, target, _binding = _fixture(tmp_path)
    context = assemble_context(target, task.record_id, index, repair_enabled=True)
    other = context.model_copy(deep=True)
    other.target.task_id = "different-task"
    other.fragments.reverse()
    assert evidence_content_hash(other) == evidence_content_hash(context)
    other.fragments[0].fact_eligible = not other.fragments[0].fact_eligible
    assert evidence_content_hash(other) != evidence_content_hash(context)


@pytest.mark.parametrize("template_priority", [False, True])
@pytest.mark.parametrize("incremental", [False, True])
def test_shared_executor_auto_supplement_and_durable_recovery(
    tmp_path, monkeypatch, template_priority, incremental,
):
    from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor
    from app.services.extraction.ontology_guided.heuristic_search import HeuristicSearchPolicy
    from app.services.extraction.ontology_guided.metadata import prepare_metadata
    from tests.test_extraction.test_semantic_graph_closure import APPEARANCE, DESCRIBES, ROOT

    analysis, _index, _task, _target, _binding = _fixture(tmp_path)
    arguments = dict(
        recognition_run_id="repair-run",
        run_fingerprint="repair-fingerprint",
        ir=analysis.ir,
        metadata=prepare_metadata(
            analysis.ir,
            section_tree=analysis.structure.section_tree.to_dict(),
            summary_version="repair-test",
        ),
        root_class_iri=ROOT,
        root_class_label="报告",
        filename="source.docx",
    )
    calls, batches, states = [], [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)

    def execute(**kw):
        return OntologyGuidedExecutor(
            ontology=_ontology(),
            engine=object(),
            adapter=EvidenceRepairAdapter(object(), model_identity="repair-fixture"),
            evidence_repair=True,
            incremental_performance=incremental,
            template_interleaving=template_priority,
            priority_paths=[(DESCRIBES, APPEARANCE)] if template_priority else [],
            heuristic_policy=HeuristicSearchPolicy.durable(
                initial_page_size=1, incremental=incremental,
            ),
            max_model_calls_per_record=8,
            max_tasks=100,
        ).run(**arguments, **kw)

    class Interrupted(Exception):
        pass

    def checkpoint(batch):
        batches.append(batch.model_copy(deep=True))
        if batch.task.retry_kind and batch.task.retry_kind.startswith("positive_evidence"):
            raise Interrupted()

    with pytest.raises(Interrupted):
        execute(batch_hook=checkpoint, model_call_hook=lambda s: states.append(deepcopy(s)))
    last = batches[-1]
    assert last.outcome.edges[0].policy_eligible
    assert last.outcome.model_calls == 1
    count = len(calls)
    resumed = execute(
        resume_state={
            name: getattr(last, name)
            for name in (
                "task_outcomes",
                "frontier",
                "recall_ledger",
                "dependency_index",
                "diagnostics",
            )
        },
        ranking_state=last.ranking_state,
        model_call_state=states[-1],
    )
    assert any(p.policy_eligible for p in resumed.graph.properties)
    assert all(r["target"]["claim_ref"]["id"] != last.task.claim_lineage_id for r in calls[count:])
    assert resumed.graph.progress.records_examined <= resumed.graph.progress.records_planned


@pytest.mark.parametrize(
    "wrong_column,condition,wrong_owner,expected",
    [
        (False, False, False, True),
        (True, False, False, False),
        (False, True, False, False),
        (False, False, True, False),
    ],
)
@pytest.mark.parametrize("reverse_columns", [False, True])
def test_property_wire_schema_and_real_column_roles(
    tmp_path, monkeypatch, wrong_column, condition, wrong_owner, expected, reverse_columns
):
    from docx import Document

    from app.services.extraction.ontology_guided.contracts import (
        SlotSpec,
        SubjectRef,
    )
    from app.services.extraction.ontology_guided.records import RecordIndex
    from app.services.extraction.word_analysis import analyze_word_core

    _analysis, _index, task, target, _binding = _fixture(tmp_path)
    document = Document()
    table = document.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, (
        ("设备名称", "设备编号", "规格型号"),
        ("1000L反应釜", "RE1/RE2", "1000L"),
        ("1000L反应釜", "RE3", "OTHER"),
    )):
        for cell, text in zip(row.cells, values[::-1] if reverse_columns else values):
            cell.text = text
    path = tmp_path / "equipment.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    record = next(r for r in index.records if r.kind == "table_row")
    predicate = SlotSpec(
        iri="urn:test:modelSpecification",
        label="规格型号",
        datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
    )
    subject = SubjectRef(entity_id="equipment", revision=1, class_iri="urn:Equipment")
    task = task.model_copy(
        update={
            "subject": subject,
            "record_id": record.record_id,
            "predicate_iri": predicate.iri,
            "predicate_kind": "property",
        }
    )
    target = target.model_copy(
        update={
            "subject_ref": subject,
            "predicate_iri": predicate.iri,
            "document_context": target.document_context.model_copy(
                update={
                    "document_hash": analysis.ir.document_hash,
                }
            ),
        }
    )
    source = next(u for u in record.source_units if u.text == "1000L反应釜")
    original_owner = source
    if wrong_owner:
        other = next(r for r in index.records if r.kind == "table_row" and r != record)
        original_owner = next(u for u in other.source_units if u.text == source.text)
    context = assemble_context(
        target,
        record.record_id,
        index,
        repair_enabled=True,
        subject_label=source.text,
        subject_evidence_refs=[index.ir.anchor(original_owner.evidence_id)],
    )
    sent = []

    def respond(_client, *, user, system, schema, **_kwargs):
        request = json.loads(user)
        sent.append(request)
        assert "allowed_object_classes" not in user + system
        assert "type_verdict" not in json.dumps(schema)
        assert "object_class_iri" not in json.dumps(schema)
        value = "RE1/RE2" if wrong_column else "1000L"
        fragment = next(f for f in request["fragments"] if f["text"] == value)
        owner = next(f for f in request["fragments"] if f["text"] == "1000L反应釜")
        binding = next(
            b
            for b in request["field_bindings"]
            if any(r["evidence_id"] == fragment["evidence_id"] for r in b["target_value_refs"])
        )
        if request["stage"] == "discovery":
            return {
                "proposals": [
                    {
                        "kind": "property",
                        "value_quote": {"evidence_id": fragment["evidence_id"], "text": value},
                        "bridge_kind": "role_mapped_table",
                        "field_binding_id": binding["field_binding_id"],
                        "condition_support": (
                            [{"evidence_id": owner["evidence_id"], "text": owner["text"]}]
                            if condition
                            else []
                        ),
                    }
                ]
            }
        owner_schema = schema["$defs"]["SupportedSubjectBinding"]
        assert {"support", "local_support"} <= set(owner_schema["required"])
        assert owner_schema["properties"]["support"]["minItems"] == 1
        assert owner_schema["properties"]["local_support"]["minItems"] == 1
        assert set(schema["$defs"]["OriginalOwnerProof"]["properties"]["evidence_id"]["enum"]) == {
            ref["evidence_id"] for ref in request["subject_evidence_refs"]
        }
        candidate = request["candidates"][0]
        return {
            "verifications": [
                {
                    "candidate_id": candidate["candidate_id"],
                    "target_id": candidate["target_id"],
                    **{
                        f"{name}_verdict": "supported"
                        for name in (
                            "role",
                            "predicate",
                            "applicability",
                            "counterevidence",
                            "bridge",
                        )
                    },
                    "field_role_support": [
                        {"evidence_id": r["evidence_id"]} for r in binding["label_refs"]
                    ],
                    "bridge_support": [{"evidence_id": fragment["evidence_id"]}],
                    "predicate_support": [{"evidence_id": fragment["evidence_id"]}],
                    "subject_binding": {
                        "verdict": "supported",
                        "support": [{"evidence_id": owner["evidence_id"]}],
                        "local_support": [{"evidence_id": owner["evidence_id"]}],
                    },
                    "condition_support": candidate["claim"]["conditions"],
                    "counterevidence_support": [],
                    "reason": "受控核验：服务端仍须检查实际表头映射和条件角色。",
                }
            ]
        }

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = EvidenceRepairAdapter(object(), model_identity="repair-fixture").inspect(
        task,
        context,
        predicate,
        compile_local_menu(_ontology(), _task_subject(target), engine=object()),
    )
    assert result.properties[0].policy_eligible is expected, context.protocol_state.get(
        "gate_issues"
    )
    assert len(sent) == 2


def _task_subject(target):
    from tests.test_extraction.test_semantic_graph_closure import ROOT

    return target.subject_ref.model_copy(update={"class_iri": ROOT, "is_document_root": True})


@pytest.mark.parametrize(
    "raw,datatype,expected,issue",
    [
        ("2026年08月", "gYearMonth", "2026-08", None),
        ("2026-13", "gYearMonth", None, "datatype_mismatch"),
        ("2026-02-30", "date", None, "datatype_mismatch"),
        ("1", "integer", 1, None),
        ("1批", "integer", None, "datatype_mismatch"),
        ("否", "boolean", False, None),
        ("RE1/RE2", "string", "RE1/RE2", None),
        ("RE1/RE2", "integer", None, "datatype_mismatch"),
        ("NaN", "decimal", None, "datatype_mismatch"),
        ("5", "unknownType", None, "constraint_unresolved"),
    ],
)
def test_literal_constraints_preserve_source_precision(raw, datatype, expected, issue):
    from app.services.extraction.ontology_guided.contracts import SlotSpec
    from app.services.extraction.ontology_guided.value_constraints import XSD, normalize_literal

    slot = SlotSpec(iri="urn:slot", label="字段", datatype_iris=[XSD + datatype])
    assert normalize_literal(raw, slot) == (expected, issue)
    slot.datatype_iris.append(XSD + "string")
    if datatype != "string":
        assert normalize_literal(raw, slot) == (None, "constraint_unresolved")


@pytest.mark.parametrize("missing", ["field_role_support", "bridge_support"])
def test_supported_verdict_without_actual_facet_source_never_passes(tmp_path, monkeypatch, missing):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        result = repaired_response(request)
        if request["stage"] == "verification":
            for review in result["verifications"]:
                review[missing] = []
        return result

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    context = assemble_context(
        target, task.record_id, index, required_context_refs=[binding], repair_enabled=True
    )
    adapter = EvidenceRepairAdapter(object(), model_identity="repair-fixture")
    if missing == "field_role_support":
        with pytest.raises(model_adapter.RecognitionModelFailure, match="response_invalid"):
            adapter.inspect(task, context, predicate, menu)
    else:
        outcome = adapter.inspect(task, context, predicate, menu)
        assert outcome.edges and not any(e.policy_eligible for e in outcome.edges)


def test_reproposal_gets_fresh_target_and_independent_review(tmp_path, monkeypatch):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)
    calls = []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    adapter = EvidenceRepairAdapter(object(), model_identity="repair-fixture")
    context = assemble_context(
        target, task.record_id, index, required_context_refs=[binding], repair_enabled=True
    )
    adapter.inspect(task, context, predicate, menu)
    old = deepcopy(context.protocol_state)
    retry = task.model_copy(update={"retry_kind": "rediscovery:1"})
    outcome = adapter.verify_existing(retry, context, predicate, menu, old)
    assert outcome.model_calls == 2
    assert [r["stage"] for r in calls] == ["discovery", "verification"] * 2
    assert calls[1]["candidates"][0]["target_id"] != calls[3]["candidates"][0]["target_id"]
    assert context.protocol_state["completed_attempts"] == [1, 2, 3, 4]
    adapter.inspect(retry, context, predicate, menu)
    assert len(calls) == 4


def test_type_coalescing_keeps_the_exact_selected_proof():
    from app.services.extraction.ontology_guided.contracts import (
        EdgeSpec,
        GraphEdge,
        GraphNode,
        RangeClass,
        VersionedRef,
    )
    from app.services.extraction.ontology_guided.executor import TaskOutcome

    predicate = EdgeSpec(
        iri="urn:uses",
        label="使用",
        range_class_iris=["urn:Base", "urn:Leaf"],
        range_classes=[
            RangeClass(iri="urn:Base", label="基类"),
            RangeClass(iri="urn:Mid", label="中间", parent_iris=["urn:Base"]),
            RangeClass(iri="urn:Leaf", label="叶类", parent_iris=["urn:Mid"]),
        ],
    )
    nodes, edges = [], []
    # The least specific type comes last: blindly retaining the last edge is wrong.
    for cls in ("urn:Leaf", "urn:Base"):
        nodes.append(
            GraphNode(
                entity_id="mention",
                revision=1,
                class_iri=cls,
                class_label=cls,
                label="设备",
                decision_status="supported",
            )
        )
        edges.append(
            GraphEdge(
                candidate_id="same-id",
                revision=1,
                subject_ref=VersionedRef(id="root", revision=1),
                object_ref=VersionedRef(id="mention", revision=1),
                predicate_iri=predicate.iri,
                predicate_label=predicate.label,
                policy_eligible=True,
                decision_status="supported",
                proof_ref=VersionedRef(id="proof:" + cls, revision=1),
            )
        )
    result = EvidenceRepairAdapter._coalesce_types(
        TaskOutcome(
            semantic_outcome="supported",
            reason_code="fixture",
            reason="两类型证明",
            nodes=nodes,
            edges=edges,
        ),
        predicate,
    )
    assert len(result.nodes) == len(result.edges) == 1
    assert result.nodes[0].class_iri == "urn:Leaf"
    assert result.edges[0].proof_ref.id == "proof:urn:Leaf"


@pytest.mark.parametrize("incremental", [False, True])
def test_production_stage_receipt_survives_interruption_without_rediscovery(
    db, tmp_path, monkeypatch, incremental,
):
    from app.services.document_analysis import execution
    from app.services.document_analysis.run_store import DocumentAnalysisRunStore
    from app.services.extraction.ontology_guided.executor import (
        ModelCallPersistenceFailure,
        OntologyGuidedExecutor,
    )
    from app.services.extraction.ontology_guided.heuristic_search import HeuristicSearchPolicy
    from app.services.extraction.ontology_guided.metadata import prepare_metadata
    from tests.test_extraction.test_document_state_artifacts import seed
    from tests.test_extraction.test_semantic_graph_closure import ROOT

    analysis, _index, _task, _target, _binding = _fixture(tmp_path)
    store = DocumentAnalysisRunStore(db)
    run, token = seed(store, "repair-durable")
    run.document_hash = analysis.ir.document_hash
    run.run_fingerprint = "repair-durable-fingerprint"
    if incremental:
        from app.services.document_analysis.run_store import content_hash

        payload = {"performance_policy": {"state_storage_version": 3}}
        store.update_artifact(
            run.recognition_run_id, run.owner_id, token, artifact_kind="source",
            expected_revision=0, artifact_id=f"{run.recognition_run_id}:source",
            artifact_hash=content_hash(payload), payload=payload, status="ready",
            event_head=run.event_head,
        )
    db.commit()
    calls, saved = [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)

    def persist(state):
        execution._persist_model_call_state(
            db,
            store,
            run,
            token,
            final_fingerprint=run.run_fingerprint,
            state=state,
        )
        saved.append(deepcopy(state))
        if any((p.get("discovery") or {}).get("proposals") for p in state["protocols"].values()):
            raise RuntimeError("process exited after committed discovery")

    def execute(**kwargs):
        return OntologyGuidedExecutor(
            ontology=_ontology(),
            engine=object(),
            adapter=EvidenceRepairAdapter(object(), model_identity="repair-fixture"),
            heuristic_policy=HeuristicSearchPolicy.durable(
                initial_page_size=1, incremental=incremental,
            ),
            evidence_repair=True,
            incremental_performance=incremental,
            max_tasks=8,
            max_model_calls_per_record=8,
        ).run(
            recognition_run_id=str(run.recognition_run_id),
            run_fingerprint=run.run_fingerprint,
            ir=analysis.ir,
            metadata=prepare_metadata(
                analysis.ir,
                section_tree=analysis.structure.section_tree.to_dict(),
                summary_version="repair-test",
            ),
            root_class_iri=ROOT,
            root_class_label="报告",
            filename="source.docx",
            **kwargs,
        )

    with pytest.raises(ModelCallPersistenceFailure):
        execute(model_call_hook=persist)
    restored = execution._restore_model_call_state(
        db,
        store,
        run,
        final_fingerprint=run.run_fingerprint,
    )
    assert restored == saved[-1]
    first_lineage = calls[0]["target"]["claim_ref"]["id"]
    result = execute(model_call_state=restored)
    assert (
        sum(
            r["stage"] == "discovery" and r["target"]["claim_ref"]["id"] == first_lineage
            for r in calls
        )
        == 1
    )
    assert result.graph.progress.model_calls == len(calls)
    assert result.graph.progress.model_calls_unresolved == 0
    changed = deepcopy(restored)
    changed["protocols"][first_lineage]["completed_attempts"] = []
    with pytest.raises(execution.CheckpointMismatch, match="receipts cannot regress"):
        execution._persist_model_call_state(
            db,
            store,
            run,
            token,
            final_fingerprint=run.run_fingerprint,
            state=changed,
        )


def test_repeated_number_requires_unique_authorized_context(tmp_path, monkeypatch):
    from docx import Document

    from app.services.extraction.ontology_guided.contracts import SlotSpec
    from app.services.extraction.ontology_guided.records import RecordIndex
    from app.services.extraction.word_analysis import analyze_word_core

    _analysis, _index, task, target, _binding = _fixture(tmp_path)
    document = Document()
    source = "车间611。本次计划生产1批，另有1批待审批。"
    document.add_paragraph(source)
    path = tmp_path / "repeat-count.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    record = index.records[0]
    predicate = SlotSpec(
        iri="urn:batchCount",
        label="批次数",
        datatype_iris=["http://www.w3.org/2001/XMLSchema#integer"],
    )
    task = task.model_copy(
        update={
            "record_id": record.record_id,
            "predicate_iri": predicate.iri,
            "predicate_kind": "property",
        }
    )
    target = target.model_copy(
        update={
            "predicate_iri": predicate.iri,
            "document_context": target.document_context.model_copy(
                update={
                    "document_hash": analysis.ir.document_hash,
                }
            ),
        }
    )
    menu = compile_local_menu(_ontology(), task.subject, engine=object())

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        fragment = next(f for f in request["fragments"] if f["fact_eligible"])
        quote = {"evidence_id": fragment["evidence_id"]}
        if request["stage"] == "discovery":
            return {
                "proposals": [
                    {
                        "kind": "property",
                        "bridge_kind": "explicit_assertion",
                        "value_quote": {**quote, "text": "1", "context_text": "本次计划生产1批"},
                    }
                ]
            }
        candidate = request["candidates"][0]
        return {
            "verifications": [
                {
                    "candidate_id": candidate["candidate_id"],
                    "target_id": candidate["target_id"],
                    **{
                        name + "_verdict": "supported"
                        for name in (
                            "role",
                            "predicate",
                            "applicability",
                            "counterevidence",
                            "bridge",
                        )
                    },
                    "subject_binding": {"verdict": "supported", "support": []},
                    "field_role_support": [quote],
                    "bridge_support": [quote],
                    "predicate_support": [quote],
                    "condition_support": [],
                    "counterevidence_support": [],
                    "reason": "当前计划批次数。",
                }
            ]
        }

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    context = assemble_context(target, record.record_id, index, repair_enabled=True)
    # This fixture tests a narrative number, without a competing structural mapping.
    assert not context.field_bindings
    result = EvidenceRepairAdapter(object(), model_identity="repair-fixture").inspect(
        task,
        context,
        predicate,
        menu,
    )
    prop = result.properties[0]
    assert prop.policy_eligible and prop.raw_value == "1" and prop.normalized_value == 1
    assert prop.value_evidence_refs[0].span_start == source.index("1批")


def test_verification_receipt_replays_after_reordering_sources(tmp_path, monkeypatch):
    _analysis, index, task, target, binding = _fixture(tmp_path)
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)
    calls, saved = [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    context = assemble_context(
        target, task.record_id, index, required_context_refs=[binding], repair_enabled=True
    )

    def interrupt(state):
        saved.append(deepcopy(state))
        if state.get("verification"):
            raise RuntimeError("lost completion acknowledgement")

    adapter = EvidenceRepairAdapter(object(), model_identity="repair-fixture")
    context.bind_protocol_hook(interrupt)
    with pytest.raises(RuntimeError, match="lost completion"):
        adapter.inspect(task, context, predicate, menu)
    context.bind_protocol_hook(lambda _state: None)
    context.fragments.reverse()
    result = adapter.verify_existing(task, context, predicate, menu, saved[-1])
    assert result.edges[0].policy_eligible
    assert len(calls) == 2 and result.model_calls == 0


def test_unproven_late_negative_suspends_positive_path_and_rechecks_sources(tmp_path, monkeypatch):
    from app.services.extraction.ontology_guided.dependencies import DependencyIndex
    from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor
    from app.services.extraction.ontology_guided.metadata import prepare_metadata
    from tests.test_extraction.test_semantic_graph_closure import (
        DESCRIBES,
        ROOT,
        _analysis,
        _effective,
        _respond,
    )

    analysis, batches, calls = _analysis(tmp_path, negative=True), [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        result = _respond(request)
        if request["stage"] == "discovery":
            allowed = {
                "kind",
                "object_class_iri",
                "object_label",
                "object_quote",
                "bridge_kind",
                "polarity",
                "condition_support",
                "applicability",
            }
            result["proposals"] = [
                {k: v for k, v in proposal.items() if k in allowed}
                for proposal in result["proposals"]
            ]
        else:
            by_id = {c["candidate_id"]: c for c in request["candidates"]}
            for review in result["verifications"]:
                review["field_role_support"] = deepcopy(review["predicate_support"])
                review["bridge_support"] = deepcopy(review["predicate_support"])
                review["subject_binding"] = {
                    "verdict": review.pop("subject_binding_verdict"),
                    "support": review.pop("subject_support"),
                }
                review["subject_binding"]["local_support"] = (
                    [] if request["subject"]["is_document_root"]
                    else deepcopy(review["subject_binding"]["support"])
                )
                if not request["subject"]["is_document_root"]:
                    review["subject_binding"]["support"].extend(
                        {"evidence_id": r["evidence_id"]}
                        for r in request["subject_evidence_refs"]
                    )
                if by_id[review["candidate_id"]]["claim"]["polarity"] == "negated":
                    review["role_verdict"] = "undetermined"
        return result

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = OntologyGuidedExecutor(
        ontology=_ontology(),
        engine=object(),
        adapter=EvidenceRepairAdapter(object(), model_identity="repair-fixture"),
        evidence_repair=True,
        max_model_calls_per_record=8,
        max_tasks=50,
    ).run(
        recognition_run_id="negative-repair",
        run_fingerprint="negative-fingerprint",
        ir=analysis.ir,
        metadata=prepare_metadata(
            analysis.ir,
            section_tree=analysis.structure.section_tree.to_dict(),
            summary_version="negative-test",
        ),
        root_class_iri=ROOT,
        root_class_label="报告",
        filename="source.docx",
        batch_hook=batches.append,
    )
    state = batches[-1]
    effective = _effective(result, DependencyIndex.from_snapshot(state.dependency_index))
    assert [edge.predicate_iri for edge in effective.edges] == [DESCRIBES]
    assert any(r["required_counterevidence"] for r in calls)
    assert result.graph.progress.invalidated_count > 0
    assert result.graph.progress.completion == "incomplete"
    assert any(e.polarity == "negated" and not e.policy_eligible for e in result.graph.edges)


def test_last_paid_result_can_be_applied_after_restart_with_zero_remaining_budget(
    tmp_path, monkeypatch
):
    from docx import Document

    from app.services.extraction.ontology_guided.executor import (
        ModelCallPersistenceFailure,
        OntologyGuidedExecutor,
    )
    from app.services.extraction.ontology_guided.heuristic_search import HeuristicSearchPolicy
    from app.services.extraction.ontology_guided.metadata import prepare_metadata
    from app.services.extraction.word_analysis import analyze_word_core
    from tests.test_extraction.test_joint_evidence_validation import NAME
    from tests.test_extraction.test_semantic_graph_closure import ROOT

    document = Document()
    document.add_paragraph(NAME)
    source = tmp_path / "last-paid.docx"
    document.save(source)
    analysis = analyze_word_core(source)
    calls, states, batches = [], [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)

    def run(**kwargs):
        return OntologyGuidedExecutor(
            ontology=_ontology(),
            engine=object(),
            adapter=EvidenceRepairAdapter(object(), model_identity="repair-fixture"),
            evidence_repair=True,
            heuristic_policy=HeuristicSearchPolicy.durable(initial_page_size=1),
            max_tasks=1,
            max_model_calls_per_record=2,
        ).run(
            recognition_run_id="paid-limit",
            run_fingerprint="paid-limit-fingerprint",
            ir=analysis.ir,
            metadata=prepare_metadata(
                analysis.ir,
                section_tree=analysis.structure.section_tree.to_dict(),
                summary_version="paid-test",
            ),
            root_class_iri=ROOT,
            root_class_label="报告",
            filename="source.docx",
            **kwargs,
        )

    def crash(state):
        states.append(deepcopy(state))
        if any(p.get("outcome") for p in state["protocols"].values()):
            raise RuntimeError("crash after last paid result")

    with pytest.raises(ModelCallPersistenceFailure):
        run(model_call_hook=crash)
    assert len(calls) == 2
    result = run(model_call_state=states[-1], batch_hook=batches.append)
    assert len(calls) == 2
    assert batches[0].outcome.complete and batches[0].outcome.model_calls == 0
    assert result.graph.progress.records_examined == 1
    assert result.graph.progress.model_calls == 2
    assert result.graph.progress.model_calls_unresolved == 0


@pytest.mark.parametrize(
    "clause", ["目前项目处于I期临床阶段", "属于肿瘤药产品", "并非青霉素类药物"],
)
def test_supported_classification_clause_cannot_become_an_entity(tmp_path, monkeypatch, clause):
    from tests.test_extraction.test_semantic_graph_closure import PRODUCT

    _analysis, index, task, target, binding = _fixture(tmp_path, extra_text=clause)
    record = next(r for r in index.records if r.text == clause)
    task = task.model_copy(update={"record_id": record.record_id})
    context = assemble_context(
        target, record.record_id, index, required_context_refs=[binding], repair_enabled=True,
    )
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        if request["stage"] == "discovery":
            source = next(f for f in request["fragments"] if f["fact_eligible"])
            return {"proposals": [{
                "kind": "relationship", "object_class_iri": PRODUCT, "object_label": clause,
                "object_quote": {"evidence_id": source["evidence_id"], "text": clause},
                "bridge_kind": "document_subject_description",
            }]}
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = EvidenceRepairAdapter(object(), model_identity="repair-fixture").inspect(
        task, context, predicate, menu,
    )
    assert result.model_calls == 2
    assert result.edges and not any(e.policy_eligible for e in result.edges)
    issues = next(iter(context.protocol_state["gate_issues"].values()))
    assert "entity_reference_not_specific" in issues


@pytest.mark.parametrize("owner_case", ["same_row", "other_row", "other_table", "notes"])
def test_two_owner_citations_cannot_turn_cooccurrence_into_identity(
    tmp_path, monkeypatch, owner_case,
):
    from docx import Document

    from app.services.extraction.ontology_guided.contracts import SlotSpec, SubjectRef
    from app.services.extraction.ontology_guided.records import RecordIndex
    from app.services.extraction.word_analysis import analyze_word_core

    _, _, task, target, _ = _fixture(tmp_path)
    document = Document()
    table = document.add_table(rows=3, cols=4)
    for row, values in zip(table.rows, (
        ("设备名称", "设备编号", "规格型号", "备注"),
        ("真空干燥箱", "DE1", "24盘", "N/A"),
        ("真空干燥箱", "DE2", "48盘", "N/A"),
    )):
        for cell, text in zip(row.cells, values):
            cell.text = text
    cleaning = document.add_table(rows=2, cols=2)
    for row, values in zip(cleaning.rows, (
        ("设备名称及编号", "清洁方法"), ("真空干燥箱", "用纯化水擦拭"),
    )):
        for cell, text in zip(row.cells, values):
            cell.text = text
    path = tmp_path / "separate-owners.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    rows = [r for r in index.records if r.kind == "table_row"]
    current = next(r for r in rows if "DE1" in r.text)
    original_record = (
        next(r for r in rows if "DE2" in r.text) if owner_case == "other_row"
        else next(r for r in rows if "用纯化水擦拭" in r.text) if owner_case == "other_table"
        else current
    )
    original = next(u for u in original_record.source_units if u.text == "真空干燥箱")
    local = next(u for u in current.source_units if u.text == (
        "N/A" if owner_case == "notes" else "真空干燥箱"
    ))
    predicate = SlotSpec(iri="urn:test:equipmentID", label="设备编号",
                         datatype_iris=["http://www.w3.org/2001/XMLSchema#string"])
    subject = SubjectRef(entity_id="dryer", revision=1, class_iri="urn:VacuumDryer")
    task = task.model_copy(update={"subject": subject, "record_id": current.record_id,
                                  "predicate_iri": predicate.iri, "predicate_kind": "property"})
    target = target.model_copy(update={
        "subject_ref": subject, "predicate_iri": predicate.iri,
        "document_context": target.document_context.model_copy(update={
            "document_hash": analysis.ir.document_hash,
        }),
    })
    context = assemble_context(
        target, current.record_id, index, repair_enabled=True, subject_label="真空干燥箱",
        subject_evidence_refs=[index.ir.anchor(original.evidence_id)],
    )

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        original_ref = request["subject_evidence_refs"][0]["evidence_id"]
        local_ref = next(f["evidence_id"] for f in request["fragments"]
                         if f["fact_eligible"] and f["text"] == local.text)
        value = next(f for f in request["fragments"] if f["text"] == "DE1")
        binding = next(b for b in request["field_bindings"]
                       if any(r["evidence_id"] == value["evidence_id"]
                              for r in b["target_value_refs"]))
        if request["stage"] == "discovery":
            return {"proposals": [{"kind": "property", "bridge_kind": "role_mapped_table",
                                   "field_binding_id": binding["field_binding_id"],
                                   "value_quote": {"evidence_id": value["evidence_id"],
                                                   "text": "DE1"}}]}
        candidate = request["candidates"][0]
        return {"verifications": [{
            "candidate_id": candidate["candidate_id"], "target_id": candidate["target_id"],
            **{name + "_verdict": "supported" for name in
               ("role", "predicate", "applicability", "counterevidence", "bridge")},
            "subject_binding": {"verdict": "supported",
                                "support": [{"evidence_id": original_ref}],
                                "local_support": [{"evidence_id": local_ref}]},
            "field_role_support": [{"evidence_id": r["evidence_id"]}
                                   for r in binding["label_refs"]],
            "predicate_support": [{"evidence_id": r["evidence_id"]}
                                  for r in binding["label_refs"]],
            "bridge_support": [{"evidence_id": value["evidence_id"]}],
            "condition_support": [], "counterevidence_support": [],
            "reason": "受控全支持；同名或备注列不能单独证明跨记录主体身份。",
        }]}

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = EvidenceRepairAdapter(object(), model_identity="owner-regression").inspect(
        task, context, predicate, None,
    )
    assert result.properties[0].policy_eligible is (owner_case == "same_row")


def test_previous_owner_policy_result_is_not_silently_reused(tmp_path, monkeypatch):
    _, index, task, target, _ = _fixture(tmp_path)
    menu = compile_local_menu(_ontology(), task.subject, engine=object())
    predicate = next(p for p in menu.relationships if p.iri == task.predicate_iri)
    calls = []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request)
        return repaired_response(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    context = assemble_context(target, task.record_id, index, repair_enabled=True)
    adapter = EvidenceRepairAdapter(object(), model_identity="repair-fixture")
    adapter.inspect(task, context, predicate, menu)
    before = len(calls)
    context.protocol_state.pop("owner_binding_version")
    with pytest.raises(ValueError, match="frozen_assertion_owner_policy_mismatch"):
        adapter.inspect(task, context, predicate, menu)
    assert len(calls) == before
