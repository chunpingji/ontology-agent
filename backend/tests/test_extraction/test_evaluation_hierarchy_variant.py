"""Contract tests for the isolated experiment, not model-quality benchmarks."""

import json

import pytest
from docx import Document

from app.evaluation.hierarchy_variant import build_plan, build_variant
from app.schemas.evidence import TaskBudget
from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_hierarchical_context import TestTokenizer

SCHEMA = {
    "urn:Device": {
        "label": "设备",
        "description": "反应设备",
        "properties": [],
        "relationships": [],
    },
    "urn:Drug": {"label": "药品", "properties": [], "relationships": []},
}


@pytest.fixture
def source(tmp_path):
    doc = Document()
    doc.add_heading("药品信息", level=1)
    doc.add_paragraph("制剂 A。")
    doc.add_heading("设备信息", level=1)
    doc.add_paragraph("反应设备 R-101。")
    path = tmp_path / "hierarchy-variant.docx"
    doc.save(path)
    return analyze_word_core(path).ir


def base(model, *, tokenizer=None):
    return GenericExtractionRunner(
        SCHEMA,
        tokenizer or TestTokenizer(),
        model,
        model_identity="fixture-no-live-model",
        budget=TaskBudget(max_input_tokens=60000, max_regions_per_task=2, max_tasks=100),
    )


def request(user):
    outer = json.loads(user)
    return outer, json.loads(outer["context"])


def metadata_for(ir, text):
    return {
        node["node_id"]: {
            "content_summary": text,
            "summary_source": "fixture",
            "summary_status": "completed",
            "analysis_id": ir.analysis_id,
        }
        for node in ir.nodes
    }


def test_every_class_and_source_region_is_still_visited(source):
    seen = []

    def model(system, user, schema, budget):
        seen.append(request(user)[1])
        return {"entities": []}

    original = base(model)
    before = evidence_hash([original.schema, source])
    runner = build_variant(original, source, metadata=metadata_for(source, "设备与药品"))
    result = runner.run(source)
    assert result.completion == "complete", result.diagnostics
    coverage = {
        (fragment["anchor"]["evidence_id"], iri)
        for context in seen
        for fragment in context["fragments"]
        if fragment["purpose"] == "target"
        for iri in context["task"]["predicate_definition"]["classes"]
    }
    assert coverage == {(unit.evidence_id, iri) for unit in source.evidence_units for iri in SCHEMA}
    assert evidence_hash([original.schema, source]) == before
    assert runner.ontology_release == original.ontology_release
    assert runner.plan["all_source_regions_and_classes_preserved"] is True


def test_metadata_is_only_a_recall_hint_and_verification_uses_original_context(source):
    calls = []

    def model(system, user, schema, budget):
        outer, context = request(user)
        calls.append((outer, context))
        if outer["stage"] == "verify_entity_types":
            return {
                "decisions": [
                    {
                        "candidate_id": candidate["candidate_id"],
                        "supported": True,
                        "reason": "设备原文明示",
                        "identity_supported": False,
                    }
                    for candidate in outer["candidate"]["proposed_entities"]
                ]
            }
        if outer["stage"] == "recall" and context["task"]["task_kind"] == "entity":
            fragment = next(
                (
                    fragment
                    for fragment in context["fragments"]
                    if fragment["purpose"] == "target" and "R-101" in fragment["text"]
                ),
                None,
            )
            if fragment:
                return {
                    "entities": [
                        {
                            "class_iri": "urn:Device",
                            "mention": {
                                "evidence_id": fragment["anchor"]["evidence_id"],
                                "text": "R-101",
                            },
                        }
                    ]
                }
            return {"entities": []}
        return {"assertions": []}

    marker = "摘要专属字符串禁止作为事实"
    runner = build_variant(base(model), source, metadata=metadata_for(source, marker))
    result = runner.run(source)
    assert result.completion == "complete", result.diagnostics
    assert any(
        candidate.text == "R-101" and candidate.validation_status == "passed"
        for candidate in result.candidates
    )
    verification = [(outer, context) for outer, context in calls if outer["stage"] != "recall"]
    assert verification
    for outer, context in calls:
        assert all(marker not in fragment["text"] for fragment in context["fragments"])
        if outer["stage"] == "recall":
            hints = context["task"]["predicate_definition"]["extraction_hints"]
            assert hints["fact_eligible"] is False
        else:
            assert "extraction_hints" not in context["task"]["predicate_definition"]
            assert marker not in json.dumps(context, ensure_ascii=False)


def test_hallucinated_summary_quote_is_rejected_by_production_replay(source):
    def model(system, user, schema, budget):
        _, context = request(user)
        target = next(
            fragment for fragment in context["fragments"] if fragment["purpose"] == "target"
        )
        return {
            "entities": [
                {
                    "class_iri": "urn:Device",
                    "mention": {
                        "evidence_id": target["anchor"]["evidence_id"],
                        "text": "不存在的虚构设备",
                    },
                }
            ]
        }

    runner = build_variant(base(model), source, metadata=metadata_for(source, "不存在的虚构设备"))
    result = runner.run(source)
    assert result.completion == "incomplete"
    assert not result.candidates
    assert "source_excerpt_mismatch" in result.diagnostics


def test_plan_rejects_stale_metadata_and_summary_change_invalidates_checkpoint(source):
    original = base(lambda *args: {"entities": []})
    first = build_variant(original, source, metadata=metadata_for(source, "设备"))
    second = build_variant(original, source, metadata=metadata_for(source, "药品"))
    assert first.input_id(source) != second.input_id(source)
    assert first.input_id(source) != original.input_id(source)
    stale = metadata_for(source, "设备")
    for item in stale.values():
        item["analysis_id"] = "another-analysis"
    plan = build_plan(SCHEMA, source, metadata=stale)
    assert plan["summary_node_count"] == 0
    assert all(d["code"] == "stale_or_unavailable_metadata_ignored" for d in plan["diagnostics"])
    structure = build_plan(SCHEMA, source, mode="structure", metadata=metadata_for(source, "设备"))
    assert structure["summary_node_count"] == 0
    assert all(not cue["summary"] for node in structure["nodes"].values() for cue in node["cues"])


def test_hint_budget_fallback_preserves_original_request(source):
    class TinyHintBudget(TestTokenizer):
        def count(self, text):
            return 100000 if "extraction_hints" in text else len(text)

    def model(system, user, schema, budget):
        assert "extraction_hints" not in user
        return {"entities": []}

    runner = build_variant(
        base(model, tokenizer=TinyHintBudget()), source, metadata=metadata_for(source, "设备")
    )
    assert runner.run(source).completion == "complete"
    assert runner.variant_statistics["hints_omitted_for_budget"] > 0
    assert runner.variant_statistics["hinted_recall_calls"] == 0


def test_gliner_spans_are_replayed_and_shifted_to_original_evidence(source):
    class Spans:
        def is_available(self):
            return True

        def extract_batch_with_spans(self, texts, labels):
            return [
                [
                    {"start": 0, "end": 5, "text": "R-101", "label": "设备", "score": 0.9},
                    {"start": -1, "end": 5, "text": "R-101", "label": "设备", "score": 0.9},
                    {"start": 0, "end": 5, "text": "other", "label": "设备", "score": 0.9},
                    {"start": 0, "end": 5, "text": "R-101", "label": "unknown", "score": 0.9},
                ]
                for text in texts
            ]

    runner = build_variant(
        base(lambda *args: {"entities": []}),
        source,
        mode="structure_summary_gliner",
        gliner=Spans(),
    )
    payload = {
        "task": {
            "task_kind": "entity",
            "predicate_definition": {
                "classes": {"urn:Device": {"label": "设备"}},
            },
        }
    }
    fragment = {"text": "R-101", "anchor": {"evidence_id": "source", "span_start": 7}}
    result = runner._span_hints(payload, [fragment], ["urn:Device"])
    assert len(result) == 1
    assert result[0]["start"] == 7 and result[0]["end"] == 12
    assert result[0]["validated_fact"] is False
    assert runner.variant_statistics["gliner_invalid_spans"] == 3
    assert runner._span_hints(payload, [fragment], ["urn:Device"]) == result
    assert runner.variant_statistics["gliner_batch_calls"] == 1


def test_unavailable_gliner_cannot_silently_be_named_as_gliner_experiment(source):
    class Unavailable:
        def is_available(self):
            return False

    with pytest.raises(RuntimeError, match="GLiNER unavailable"):
        build_variant(base(None), source, mode="structure_summary_gliner", gliner=Unavailable())
