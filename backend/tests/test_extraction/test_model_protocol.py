import json

import pytest

from app.services.extraction.evidence_identity import canonical_json
from app.services.extraction.extraction_tasks import SYSTEM, EntityResponse


def request():
    context = {
        "task": {
            "task_kind": "entity",
            "predicate_definition": {"classes": {
                "https://example.org/Drug": {"label": "药品", "parents": []},
                "https://example.org/Device": {"label": "设备", "parents": []},
            }},
        },
        "subjects": {},
        "fragments": [{
            "anchor": {"evidence_id": "a" * 64, "span_start": 5, "span_end": 12},
            "text": "药品 A", "purpose": "target", "fact_eligible": True,
        }],
        "effective_class": "https://example.org/Report",
    }
    return canonical_json({
        "stage": "recall", "context": canonical_json(context), "candidate": None,
    })


def test_short_references_are_closed_and_source_text_is_never_rewritten():
    from app.services.extraction.model_protocol import ModelProtocol

    schema = EntityResponse.model_json_schema()
    original = request()
    wire = ModelProtocol(SYSTEM, original, schema)
    context = json.loads(json.loads(wire.user)["context"])
    classes = context["task"]["predicate_definition"]["classes"]
    drug = next(key for key, value in classes.items() if value["label"] == "药品")
    fragment = context["fragments"][0]
    evidence_id = fragment["anchor"]["evidence_id"]
    assert len(evidence_id) < 8 and len(drug) < 8
    assert fragment["text"] == "药品 A" and fragment["anchor"]["span_start"] == 5
    definitions = wire.schema["$defs"]
    assert set(definitions["EntityProposal"]["properties"]["class_iri"]["enum"]) == set(classes)
    assert definitions["SpanProposal"]["properties"]["evidence_id"]["enum"] == [evidence_id]
    assert json.loads(wire.user)["output_contract"] == wire.schema
    result = wire.decode({"entities": [{
        "class_iri": drug, "mention": {"evidence_id": evidence_id, "text": "药品 A"},
    }]})
    assert result["entities"][0]["class_iri"] == "https://example.org/Drug"
    assert result["entities"][0]["mention"]["evidence_id"] == "a" * 64
    assert result["entities"][0]["mention"]["text"] == "药品 A"
    assert "enum" not in schema["$defs"]["EntityProposal"]["properties"]["class_iri"]
    for key in ("not-registered", "https://example.org/Drug"):
        with pytest.raises(ValueError, match="unknown_model_reference"):
            wire.decode({"entities": [{"class_iri": key}]})
    # Strings which happen to resemble identifiers are still original evidence.
    result = wire.decode({"entities": [{"class_iri": drug, "mention": {
        "evidence_id": evidence_id, "text": drug,
    }}]})
    assert result["entities"][0]["mention"]["text"] == drug
    assert original == request()


def test_candidate_reference_aliases_preserve_roles_revision_and_polarity():
    from app.services.extraction.extraction_tasks import BindingDecision
    from app.services.extraction.model_protocol import ModelProtocol

    user = json.loads(request())
    context = json.loads(user["context"])
    subject = {"candidate_id": "subject-a", "revision": 7}
    obj = {"candidate_id": "object-b", "revision": 2}
    context["task"].update(subject=subject, object_candidates=[obj], competing_subjects=[obj])
    context["subjects"] = {"subject-a": {**subject, "text": "A"}, "object-b": {**obj, "text": "B"}}
    user["context"] = canonical_json(context)
    user["candidate"] = {"subject": subject, "object": obj, "assertion_status": "negated"}
    wire = ModelProtocol(SYSTEM, canonical_json(user), BindingDecision.model_json_schema())
    assert "table_record" not in wire.schema["properties"]["method"]["enum"]
    assert wire.schema["properties"]["record_mapping"] == {
        "type": "object", "properties": {}, "additionalProperties": False,
    }
    encoded = json.loads(wire.user)
    context = json.loads(encoded["context"])
    s = context["task"]["subject"]["candidate_id"]
    o = context["task"]["object_candidates"][0]["candidate_id"]
    assert s != o and set(context["subjects"]) == {s, o}
    assert context["task"]["subject"]["revision"] == 7
    assert context["task"]["competing_subjects"][0]["candidate_id"] == o
    result = wire.decode({"subject_candidate_id": s, "object_candidate_id": o,
                          "assertion_status": "negated", "supported": True,
                          "method": "explicit_assertion", "assertion_spans": []})
    assert result == {"subject_candidate_id": "subject-a", "object_candidate_id": "object-b",
                      "assertion_status": "negated", "supported": True,
                      "method": "explicit_assertion", "assertion_spans": []}


@pytest.mark.parametrize("kind", ["property", "relationship"])
def test_wire_assertions_require_task_specific_values_and_explicit_polarity(kind):
    from app.services.extraction.extraction_tasks import AssertionResponse
    from app.services.extraction.model_protocol import ModelProtocol

    data = json.loads(request())
    context = json.loads(data["context"])
    context["task"].update(task_kind=kind, object_candidates=[{
        "candidate_id": "object-b", "revision": 1,
    }])
    data["context"] = canonical_json(context)
    wire = ModelProtocol(SYSTEM, canonical_json(data), AssertionResponse.model_json_schema())
    proposal = wire.schema["$defs"]["AssertionProposal"]
    assert "assertions" in wire.schema["required"]
    assert "assertion_status" in proposal["required"]
    if kind == "relationship":
        assert {"object_candidate_id", "assertion_spans"} <= set(proposal["required"])
        assert proposal["properties"]["object_candidate_id"] == {
            "type": "string", "enum": list(wire.references["candidate"]),
        }
        assert proposal["properties"]["value"] == {"type": "null"}
        assert proposal["properties"]["assertion_spans"]["minItems"] == 1
    else:
        assert "value" in proposal["required"]
        assert proposal["properties"]["value"] == {"$ref": "#/$defs/SpanProposal"}
        assert proposal["properties"]["object_candidate_id"] == {"type": "null"}
    # No-evidence remains representable; the grammar must not force a false fact.
    assert wire.schema["properties"]["assertions"].get("minItems", 0) == 0
    assert wire.decode({"assertions": []}) == {"assertions": []}
    with pytest.raises(ValueError, match="invalid_model_response"):
        wire.decode({})
    with pytest.raises(ValueError, match="invalid_model_response"):
        wire.decode({"assertions": [{"assertion_status": "affirmed"}]})
    with pytest.raises(ValueError, match="invalid_model_response"):
        wire.decode({"assertions": [{
            "object_candidate_id": next(iter(wire.references["candidate"])),
            "assertion_spans": [], "value": {},
        }]})


def test_relationship_binding_schema_requires_the_exact_proposed_object():
    from app.services.extraction.extraction_tasks import BindingDecision
    from app.services.extraction.model_protocol import ModelProtocol

    data = json.loads(request())
    context = json.loads(data["context"])
    context["task"].update(task_kind="relationship", subject={
        "candidate_id": "subject-a", "revision": 1,
    }, object_candidates=[{"candidate_id": name, "revision": 1} for name in ("b", "c")])
    data["context"] = canonical_json(context)
    data["candidate"] = {"object": {"candidate_id": "b", "revision": 1}}
    wire = ModelProtocol(SYSTEM, canonical_json(data), BindingDecision.model_json_schema())
    alias = next(key for key, value in wire.references["candidate"].items() if value == "b")
    assert wire.schema["properties"]["object_candidate_id"] == {
        "type": "string", "enum": [alias],
    }
    assert "object_candidate_id" in wire.schema["required"]


def test_closed_transport_runner_replays_full_anchors_and_negative_bindings(tmp_path):
    from docx import Document

    from app.schemas.evidence import TaskBudget
    from app.services.extraction.extraction_tasks import GenericExtractionRunner
    from app.services.extraction.word_analysis import analyze_word_core
    from tests.test_extraction.test_extraction_tasks import SCHEMA, scripted_model, span
    from tests.test_extraction.test_hierarchical_context import TestTokenizer

    doc = Document()
    doc.add_heading("药品 A", 1)
    doc.add_paragraph("本品不使用设备 E")
    path = tmp_path / "closed.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    traces = []

    def model(system, user, response_schema, budget):
        request = json.loads(user)
        context = json.loads(request["context"])
        if context["task"]["task_kind"] != "entity" or request["stage"] == "verify_entity_types":
            return scripted_model(system, user, response_schema, budget)
        classes = context["task"]["predicate_definition"]["classes"]
        target = next(f for f in context["fragments"] if f["purpose"] == "target")
        return {"entities": [
            {"class_iri": key, "mention": span(target, text)}
            for text, iri in (("药品 A", "urn:NovelProduct"), ("设备 E", "urn:Device"))
            for key, definition in classes.items()
            if definition["name"] == iri and text in target["text"]
        ]}

    runner = GenericExtractionRunner(
        SCHEMA, TestTokenizer(), model, model_identity="closed",
        budget=TaskBudget(max_input_tokens=60000), compact_identifiers=True,
    )
    runner.trace_fn = traces.append
    result = runner.run(ir)
    assert result.completion == "complete", result.diagnostics
    negative = next(c for c in result.candidates if c.kind == "relationship")
    assert negative.assertion_status == "negated" and negative.validation_status == "passed"
    assert not negative.positive_eligible and negative.review_status == "pending"
    assert ir.resolve(negative.bindings[0].anchors[0]) == "本品不使用设备 E"
    assert all(t["input_tokens"] == len(t["user"] + t["system"]) + 128 for t in traces)
    count = len(traces)
    assert runner.run(ir, checkpoint=result.checkpoint).candidates == result.candidates
    assert len(traces) == count
    for trace in traces:
        proposal = trace["schema"].get("$defs", {}).get("SpanProposal")
        if proposal:
            assert proposal["properties"]["start"] == {"type": "null"}
            assert proposal["properties"]["end"] == {"type": "null"}
