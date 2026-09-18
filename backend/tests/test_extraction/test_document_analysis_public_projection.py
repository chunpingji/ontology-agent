from __future__ import annotations

import pytest
from docx import Document

from app.schemas.document_analysis import GraphArtifactResponse, GraphRanking
from app.services.document_analysis.public_projection import (
    build_selection_registry,
    public_graph_payload,
    public_ranking_payload,
)
from app.services.extraction.ontology_guided.contracts import (
    CoverageSummary,
    GraphEdge,
    GraphNode,
    GraphProperty,
    GraphRelationshipGroup,
    RunProgress,
    ScopeMember,
    TraversalScope,
    VersionedRef,
)
from app.services.extraction.ontology_guided.dependencies import DependencyIndex
from app.services.extraction.ontology_guided.projection import project_graph
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core

ROOT = "https://ontology.example/Report"
PRODUCT = "https://ontology.example/Product"
DESCRIBES = "https://ontology.example/describes"


@pytest.fixture
def tool_graph_payload(tmp_path):
    document = Document()
    document.add_paragraph("报告可选择产品甲或产品乙。仅加热时产品甲的温度为20摄氏度。独立实体丙。")
    path = tmp_path / "group.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    unit = index.ir.evidence_units[0]
    anchor = index.ir.anchor(unit.evidence_id, 0, len(unit.text))
    def ref(name):
        return VersionedRef(id=name, revision=1)

    root = GraphNode(entity_id="root", revision=1, class_iri=ROOT, class_label="报告",
                     label="报告", root=True, grounding_kind="document_root")
    nodes = [root, *[
        GraphNode(entity_id=name, revision=1, class_iri=PRODUCT, class_label="产品",
                  label=name, grounding_kind="mention", referent_ref=ref(f"referent:{name}"),
                  type_decision_ref=ref(f"type:{name}"),
                  referent_decision_ref=ref(f"decision:{name}"), evidence_refs=[anchor])
        for name in ("a", "b", "isolated")
    ]]
    group = GraphRelationshipGroup(
        candidate_id="choice", revision=1, subject_ref=ref("root"),
        object_refs=[ref("a"), ref("b")], predicate_iri=DESCRIBES, predicate_label="描述",
        selection="alternatives", selection_evidence_refs=[anchor], modality="possible",
        scope=TraversalScope.create(), decision_status="supported", structural_valid=True,
        model_supported=True, policy_eligible=True, proof_ref=ref("proof:group"),
        decision_refs=[ref("decision:group")], predicate_evidence_refs=[anchor],
    )
    scope = TraversalScope.create([
        ScopeMember(relation_ref=ref("choice"), member_ref=ref("a")),
    ])
    prop = GraphProperty(
        candidate_id="temperature", revision=1, subject_ref=ref("a"),
        predicate_iri="https://ontology.example/temperature", predicate_label="温度",
        raw_value="20", modality="asserted", scope=scope, conditions=["仅加热时"],
        decision_status="supported", structural_valid=True, model_supported=True,
        policy_eligible=True, proof_ref=ref("proof:temperature"),
        decision_refs=[ref("decision:temperature")], value_evidence_refs=[anchor],
    )
    graph = project_graph(
        recognition_run_id="a4f9885d-1b70-4718-9a97-2eb03c3ef8c4", run_revision=1,
        event_head=1, metadata_snapshot_id="metadata", root_ref=ref("root"), nodes=nodes,
        edges=[], properties=[prop], relationship_groups=[group], coverage=[],
        progress=RunProgress(), projection="all", artifact_status="ready",
        extraction_protocol="ontology-tool-extraction-v1",
    )
    registry = build_selection_registry(recognition_run_id=graph.recognition_run_id,
                                        analysis_id=index.ir.analysis_id, graph=graph, index=index)
    return {
        "snapshot_id": "graph:groups", "analysis_id": index.ir.analysis_id,
        "ontology_snapshot_id": "ontology", "graph": graph.model_dump(mode="json"),
        "selection_registry": registry,
    }


def _read_tool_graph(stored, projection, protocol="ontology-tool-extraction-v1"):
    return GraphArtifactResponse.model_validate(public_graph_payload(
        recognition_run_id=stored["graph"]["recognition_run_id"], run_revision=1,
        event_head=1, artifact_revision=1, availability="ready", projection=projection,
        stored_payload=stored, extraction_protocol=protocol,
    ))


def test_verified_keeps_groups_and_independent_nodes_without_fabricated_edges(tool_graph_payload):
    response = _read_tool_graph(tool_graph_payload, "verified")
    assert {item.entity_id for item in response.entities} == {"root", "a", "b", "isolated"}
    assert response.relationships == []
    assert len(response.relationship_groups) == len(response.properties) == 1
    group = response.relationship_groups[0]
    assert group.selection == "alternatives" and group.modality == "possible"
    assert group.source_selection_refs.selection
    assert all(tool_graph_payload["selection_registry"][ref]["selection_role"] == "selection"
               for ref in group.source_selection_refs.selection)


def test_filtered_parent_keeps_exact_inherited_scope_and_source_selections(tool_graph_payload):
    response = _read_tool_graph(tool_graph_payload, "conditional")
    assert response.relationship_groups == []
    assert len(response.properties) == len(response.scope_resolutions) == 1
    scope = response.scope_resolutions[0]
    assert scope.scope_id == response.properties[0].scope.scope_id
    assert scope.steps[0].relation_ref.id == "choice"
    assert scope.steps[0].member_ref.entity_id == "a"
    assert scope.steps[0].selection == "alternatives"
    assert scope.steps[0].modality == "possible"
    assert scope.steps[0].evidence_selection_ids
    assert all(ref in tool_graph_payload["selection_registry"]
               for ref in scope.steps[0].evidence_selection_ids)
    effective = _read_tool_graph(tool_graph_payload, "effective_affirmed")
    assert not effective.properties and not effective.relationship_groups


def test_scope_parent_revision_mismatch_cannot_become_unqualified_verified_fact(tool_graph_payload):
    tool_graph_payload["graph"]["relationship_groups"][0]["revision"] = 2
    response = _read_tool_graph(tool_graph_payload, "verified")
    assert response.properties == [] and response.scope_resolutions == []


def test_public_verified_contract_rejects_silently_dropped_scope_explanation(tool_graph_payload):
    payload = _read_tool_graph(tool_graph_payload, "verified").model_dump(mode="json")
    payload["scope_resolutions"] = []
    with pytest.raises(ValueError, match="exact parent resolutions"):
        GraphArtifactResponse.model_validate(payload)


def test_legacy_verified_is_explicitly_unsupported(tool_graph_payload):
    with pytest.raises(ValueError, match="unsupported_projection"):
        _read_tool_graph(tool_graph_payload, "verified", protocol=None)


def test_legacy_public_policy_keeps_its_frozen_name(tool_graph_payload):
    tool_graph_payload["graph"]["projection_policy"] = "legacy-frozen-projection"
    tool_graph_payload["graph"]["relationship_groups"] = []
    tool_graph_payload["graph"]["properties"] = []
    response = _read_tool_graph(tool_graph_payload, "all_candidates", protocol=None)
    assert response.graph_snapshot.projection_policy == "legacy-frozen-projection"


def test_external_identity_preserves_record_version_without_adding_graph_facts(tool_graph_payload):
    node = tool_graph_payload["graph"]["nodes"][1]
    provenance = {
        "kind": "external_record", "system": "registry", "dataset": "objects",
        "record_key": "object-a", "record_version": "v3", "field_path": "identifier",
        "value": "object-a", "identity_match_evidence": node["evidence_refs"],
        "record_snapshot": {"name": "外部别名"}, "fetched_at": None, "applicable_at": None,
    }
    node.update(identity_status="verified", external_provenance=[provenance],
                identity_decision_refs=[{"id": "identity:a", "revision": 1}])
    response = _read_tool_graph(tool_graph_payload, "verified")
    entity = next(item for item in response.entities if item.entity_id == node["entity_id"])
    assert entity.identity_state == "verified_external"
    assert entity.external_provenance[0].model_dump(mode="json") == provenance
    assert entity.identity_decision_refs[0].id == "identity:a"
    assert entity.label == node["label"] and entity.revision == node["revision"]
    assert len(response.entities) == 4 and len(response.properties) == 1
    assert response.relationships == [] and len(response.relationship_groups) == 1
    untouched = next(item for item in response.entities if item.entity_id == "b")
    assert "external_provenance" not in untouched.model_dump(mode="json")
    assert "identity_decision_refs" not in untouched.model_dump(mode="json")


@pytest.mark.parametrize("normalization", [
    {"target_unit": "kg", "source_unit": "g", "form": "scalar"},
    {"to": "kg", "from": "g"},
])
def test_public_normalized_unit_supports_quantity_and_legacy_records(
    tool_graph_payload, normalization,
):
    prop = tool_graph_payload["graph"]["properties"][0]
    prop.update(raw_value="1000 g", normalized_value="1", normalization_record=normalization)
    response = _read_tool_graph(tool_graph_payload, "verified")
    assert response.properties[0].unit == "kg"
    assert response.properties[0].normalized_value == "1"
    assert response.properties[0].normalization_record == normalization


def test_public_projection_registers_role_specific_opaque_source_refs(tmp_path):
    document = Document()
    document.add_heading("产品", level=1)
    document.add_paragraph("本报告明确描述产品甲。")
    path = tmp_path / "graph.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    unit = next(item for item in analysis.ir.evidence_units if "产品甲" in item.text)
    anchor = analysis.ir.anchor(unit.evidence_id, unit.text.index("产品甲"), len(unit.text) - 1)
    predicate_anchor = analysis.ir.anchor(
        unit.evidence_id, unit.text.index("描述"), unit.text.index("描述") + 2
    )
    root = GraphNode(
        entity_id="entity:root",
        revision=1,
        class_iri=ROOT,
        class_label="报告",
        label="graph.docx",
        root=True,
        root_origin="user_specified",
    )
    product = GraphNode(
        entity_id="entity:product",
        revision=1,
        class_iri=PRODUCT,
        class_label="产品",
        label="产品甲",
        evidence_refs=[anchor],
    )
    edge = GraphEdge(
        candidate_id="candidate:describes",
        revision=1,
        subject_ref=VersionedRef(id=root.entity_id, revision=1),
        object_ref=VersionedRef(id=product.entity_id, revision=1),
        predicate_iri=DESCRIBES,
        predicate_label="描述",
        decision_status="supported",
        structural_valid=True,
        model_supported=True,
        policy_eligible=True,
        proof_ref=VersionedRef(id="proof:describes", revision=1),
        decision_refs=[VersionedRef(id="decision:describes", revision=1)],
        object_evidence_refs=[anchor],
        predicate_evidence_refs=[predicate_anchor],
        evidence_refs=[anchor, predicate_anchor],
    )
    progress = RunProgress(
        tasks_attempted=1,
        model_calls=1,
        records_planned=8,
        records_examined=1,
        records_incomplete=0,
        records_unattempted=7,
        phase_counts={"phase1": 1, "phase2": 0},
        supported=1,
        completion="incomplete",
    )
    graph = project_graph(
        recognition_run_id="a4f9885d-1b70-4718-9a97-2eb03c3ef8c4",
        run_revision=4,
        event_head=5,
        metadata_snapshot_id="metadata:one",
        root_ref=VersionedRef(id=root.entity_id, revision=1),
        nodes=[root, product],
        edges=[edge],
        properties=[],
        coverage=[
            CoverageSummary(
                subject_ref=VersionedRef(id=root.entity_id, revision=1),
                predicate_iri=DESCRIBES,
                predicate_label="描述",
                phase1=1,
                phase2=7,
                executed_phase_counts={"phase1": 1, "phase2": 0},
                examined=1,
                unattempted=7,
            )
        ],
        progress=progress,
        projection="all",
        artifact_status="ready",
    )
    registry = build_selection_registry(
        recognition_run_id=graph.recognition_run_id,
        analysis_id=analysis.ir.analysis_id,
        graph=graph,
        index=index,
    )
    payload = public_graph_payload(
        recognition_run_id=graph.recognition_run_id,
        run_revision=7,
        event_head=8,
        artifact_revision=3,
        availability="ready",
        projection="effective_affirmed",
        stored_payload={
            "snapshot_id": "graph-snapshot:one",
            "analysis_id": analysis.ir.analysis_id,
            "ontology_snapshot_id": "ontology-snapshot:one",
            "generated_at": "2026-09-08T12:00:00Z",
            "graph": graph.model_dump(mode="json"),
            "selection_registry": registry,
        },
    )

    response = GraphArtifactResponse.model_validate(payload)
    serialized = response.model_dump(mode="json")
    assert not {"extraction_protocol", "relationship_groups", "scope_resolutions"}.intersection(
        serialized
    )
    assert not {"modality", "scope"} & serialized["relationships"][0].keys()
    assert "selection" not in serialized["relationships"][0]["source_selection_refs"]
    assert "scope" not in serialized["coverage"]["subjects"][0]
    assert response.coverage.subjects[0].phase_counts.phase2 == 7
    assert response.coverage.subjects[0].executed_phase_counts.phase2 == 0
    assert not response.coverage.phase2_started
    assert len(response.relationships) == 1
    refs = response.relationships[0].source_selection_refs
    assert refs.object and refs.predicate_bridge
    selection = registry[refs.object[0]]
    assert selection["source_record_ref"]
    assert selection["record_view_ref"]
    assert selection["anchors"][0]["evidence_id"] == unit.evidence_id


def test_public_projection_replays_durable_dependency_invalidations(tmp_path):
    document = Document()
    document.add_paragraph("本报告明确描述产品甲。")
    path = tmp_path / "invalidated-graph.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    root = GraphNode(
        entity_id="entity:root",
        revision=1,
        class_iri=ROOT,
        class_label="报告",
        label=path.name,
        root=True,
        root_origin="user_specified",
    )
    product = GraphNode(
        entity_id="entity:product",
        revision=1,
        class_iri=PRODUCT,
        class_label="产品",
        label="产品甲",
    )
    edge = GraphEdge(
        candidate_id="candidate:describes",
        revision=2,
        subject_ref=VersionedRef(id=root.entity_id, revision=1),
        object_ref=VersionedRef(id=product.entity_id, revision=1),
        predicate_iri=DESCRIBES,
        predicate_label="描述",
        decision_status="supported",
        structural_valid=True,
        model_supported=True,
        policy_eligible=True,
        proof_ref=VersionedRef(id="proof:describes", revision=3),
        decision_refs=[VersionedRef(id="decision:describes", revision=4)],
    )
    progress = RunProgress(completion="in_scope_complete")
    graph = project_graph(
        recognition_run_id="a4f9885d-1b70-4718-9a97-2eb03c3ef8c4",
        run_revision=4,
        event_head=5,
        metadata_snapshot_id="metadata:one",
        root_ref=VersionedRef(id=root.entity_id, revision=1),
        nodes=[root, product],
        edges=[edge],
        properties=[],
        coverage=[],
        progress=progress,
        projection="all",
        artifact_status="ready",
    )
    dependencies = DependencyIndex()
    dependencies.add_proof("proof:describes@3", ["decision:describes@4"])
    dependencies.add_proof("candidate:describes@2", ["proof:describes@3"])
    dependencies.invalidate("decision:describes@4")
    stored = {
        "snapshot_id": "graph-snapshot:invalidated",
        "analysis_id": analysis.ir.analysis_id,
        "ontology_snapshot_id": "ontology-snapshot:one",
        "generated_at": "2026-09-08T12:00:00Z",
        "graph": graph.model_dump(mode="json"),
        "selection_registry": {},
        "dependency_index": dependencies.snapshot(),
    }

    effective = GraphArtifactResponse.model_validate(
        public_graph_payload(
            recognition_run_id=graph.recognition_run_id,
            run_revision=7,
            event_head=8,
            artifact_revision=3,
            availability="ready",
            projection="effective_affirmed",
            stored_payload=stored,
        )
    )
    assert effective.relationships == []
    assert {item.id for item in effective.invalidated_refs} >= {
        "decision:describes",
        "proof:describes",
        "candidate:describes",
    }

    all_candidates = GraphArtifactResponse.model_validate(
        public_graph_payload(
            recognition_run_id=graph.recognition_run_id,
            run_revision=7,
            event_head=8,
            artifact_revision=3,
            availability="ready",
            projection="all_candidates",
            stored_payload=stored,
        )
    )
    assert len(all_candidates.relationships) == 1
    assert all_candidates.relationships[0].invalidated is True


def test_ranking_projection_preserves_negative_scores_and_hides_private_inputs():
    state = {
        "service": {
            "policy": {"mode": "semantic"},
            "cache": {"private-query": "uploaded document text"},
            "costs": {"model_calls": 3, "input_pairs": 4, "tokens": 128,
                      "technical_retries": 1, "elapsed_ms": 1500},
            "model_observations": [
                {"operation": "embed", "input_tokens": 64, "queue_seconds": 0.2},
                {"operation": "score_pairs", "input_tokens": None, "queue_seconds": 0.1},
            ],
            "epochs": [{
                "epoch_id": "epoch:one", "status": "committed",
                "subject_ref": {"entity_id": "subject:one", "revision": 2},
                "plan_id": "plan:one", "actual_ranking_mode": "deterministic",
                "degraded": True, "reason": "ranking_timeout",
                "queries": [{"query_id": "query:one", "predicate_iri": DESCRIBES,
                             "model_text": "private source"}],
                "ordered_record_ids": ["record:low", "record:high"],
                "observations": [{"record_id": "record:low",
                                  "retrieval_intent": "counterevidence", "intent_rank": 1,
                                  "raw_rerank_score": -3.0, "channel_hits": ["dense"]}],
            }, {"epoch_id": "epoch:pending", "status": "ready"}],
        },
    }
    result = GraphRanking.model_validate(public_ranking_payload(state))
    assert result.committed_epochs == 1
    assert result.degraded and result.reasons == ["ranking_timeout"]
    assert result.epochs[0].records[0].raw_scores == {"counterevidence": -3.0}
    assert result.epochs[0].records[0].intent_ranks == {"counterevidence": 1}
    assert result.epochs[0].records[0].channels == ["dense"]
    assert result.cost.model_calls == 3 and result.cost.input_pairs == 4
    assert result.cost.observed_requests == 2
    assert result.cost.input_tokens == 128 and result.cost.retries == 1
    assert result.cost.elapsed_seconds == 1.5
    assert result.cost.reserved_input_tokens == 128
    assert result.cost.measured_input_tokens == 64
    assert result.cost.unknown_request_count == 2
    assert round(result.cost.queue_seconds, 1) == 0.3
    assert "private" not in result.model_dump_json()
    assert state["service"]["epochs"][1]["status"] == "ready"


def test_missing_ranking_is_explicitly_unobserved_deterministic():
    result = GraphRanking.model_validate(public_ranking_payload({}))
    assert result.requested_mode == "deterministic"
    assert result.committed_epochs == 0 and result.actual_modes == []


def test_paused_uncommitted_pool_does_not_claim_a_deterministic_execution():
    state = {"service": {
        "policy": {"mode": "semantic"},
        "epochs": [{
            "epoch_id": "completed", "status": "committed", "actual_ranking_mode": "semantic",
        }],
        "pending_epochs": [{
            "epoch_id": "paused", "status": "paused", "actual_ranking_mode": "deterministic",
            "reason": "ranking_call_budget_exhausted",
        }],
    }}
    result = GraphRanking.model_validate(public_ranking_payload(state))
    assert result.paused and result.reasons == ["ranking_call_budget_exhausted"]
    assert result.actual_modes == ["semantic"] and result.committed_epochs == 1
    state["service"]["epochs"] = []
    result = GraphRanking.model_validate(public_ranking_payload(state))
    assert result.paused and result.actual_modes == [] and result.committed_epochs == 0
    assert not result.degraded and result.cost.model_calls == 0
