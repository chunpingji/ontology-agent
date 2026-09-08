"""Execution ownership across actual connections, entry points and worker deliveries."""

import asyncio
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from threading import Barrier
from time import sleep

import pytest
from docx import Document
from fastapi import BackgroundTasks, HTTPException
from sqlalchemy import create_engine, update
from sqlalchemy.orm import Session

from app.api import evidence, extraction
from app.db import Base
from app.dependencies import Identity
from app.models.evidence import EvidenceJobState
from app.models.extraction import AnnotationExecution, ExtractionJob
from app.services.extraction import annotation_execution as execution
from app.services.extraction.checkpoint_journal import CheckpointJournal
from app.services.extraction.progress import progress_bus


@pytest.fixture
def db(tmp_path):
    # Independent file-backed connections: StaticPool would share transactions.
    engine = create_engine(f"sqlite:///{tmp_path / 'ownership.db'}")
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False, autoflush=False) as session:
        yield session
    engine.dispose()


@pytest.fixture
def source_job(db, tmp_path, monkeypatch):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("执行互斥测试文档")
    doc.save(path)
    job = ExtractionJob(source_type="word", source_filename=path.name,
                        document_path=str(path), status="paused",
                        source_config={"mode": "template_default"})
    db.add(job)
    db.commit()
    monkeypatch.setattr(extraction, "_annotation_cache_path", lambda _: tmp_path / "cache.json")
    monkeypatch.setattr(extraction, "_annotation_checkpoint_path", lambda _: tmp_path / "ckpt.json")
    progress_bus.reset(str(job.id))
    return job.id


def test_different_entrypoints_claim_only_one_worker(db, source_job, fake_engine):
    barrier = Barrier(2)
    checkpoint = {"input_id": "existing", "attempt_count": 2, "completed": {}}
    extraction._write_annotation_checkpoint(source_job, checkpoint)
    identity = Identity("analyst", "senior_analyst")

    def start(which):
        with Session(db.bind, expire_on_commit=False) as other:
            background = BackgroundTasks()
            barrier.wait(timeout=10)
            try:
                if which == "background":
                    result = asyncio.run(extraction.resume_annotation(
                        source_job, background, other, fake_engine, identity,
                    ))
                else:
                    result = evidence.extract_evidence(source_job, background, db=other,
                                                       engine=fake_engine, identity=identity)
                assert len(background.tasks) == 1
                assert background.tasks[0].kwargs["run_id"] == result["run_id"]
                return result["run_id"]
            except HTTPException as exc:
                assert exc.status_code == 409
                assert not background.tasks
                return None

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(start, ["background", "evidence"]))
    assert sum(value is not None for value in results) == 1
    assert extraction._load_annotation_checkpoint(source_job) == checkpoint
    head = db.get(AnnotationExecution, source_job)
    assert head.run_id in results and head.status == "queued"
    # Clearing a process-local bus must not allow an additional HTTP worker.
    progress_bus.reset(str(source_job))
    with pytest.raises(HTTPException) as error:
        extraction._claim_annotation(source_job, db, mode="rerun")
    assert error.value.status_code == 409
    assert extraction._load_annotation_checkpoint(source_job) == checkpoint


def test_duplicate_delivery_of_same_run_starts_once(db, source_job, fake_engine, monkeypatch):
    receipt = extraction._claim_annotation(source_job, db, mode="start")
    barrier = Barrier(2)
    called = []
    monkeypatch.setattr(extraction, "_annotation_worker", lambda *args: called.append(args[0]))

    def deliver(_):
        barrier.wait(timeout=10)
        asyncio.run(extraction._precompute_annotation_bg(
            source_job, fake_engine, db.bind, run_id=receipt["run_id"],
        ))

    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(deliver, range(2)))
    assert called == [source_job]


@pytest.mark.parametrize("mode", ["continue", "resume", "retry"])
def test_entrypoints_restore_journal_instead_of_stale_sql(
    db, source_job, fake_engine, monkeypatch, mode,
):
    stale = {"input_id": "stale-synchronous-run", "attempt_count": 1}
    db.add(EvidenceJobState(job_id=source_job, extraction_run={"checkpoint": stale}))
    db.commit()
    journal = CheckpointJournal(extraction._annotation_checkpoint_path(source_job))
    base = {"input_id": "background-run", "attempt_count": 2, "completed": {"one": []}}
    journal.save(base)
    latest = {**base, "attempt_count": 3, "completed": {"one": [], "two": []}}
    journal.save(latest)
    observed = []

    def compute(*args, **kwargs):
        observed.append((args[4], kwargs["retry_failed"], kwargs["pause_after"]))
        raise RuntimeError("stop after checking restored input")

    monkeypatch.setattr(extraction, "_compute_annotation", compute)
    background = BackgroundTasks()
    if mode == "resume":
        asyncio.run(extraction.resume_annotation(source_job, background, db, fake_engine,
                                                Identity("analyst", "senior_analyst")))
    else:
        evidence.extract_evidence(source_job, background,
            req=evidence.ExtractOptions(retry_failed=mode == "retry", reason="模型已恢复",
                                        pause_after=8),
            db=db, engine=fake_engine, identity=Identity("analyst", "senior_analyst"))
    asyncio.run(background())
    assert observed[0][0]["input_id"] == "background-run"
    assert observed[0][0]["attempt_count"] == 3
    assert observed[0][0]["completed"] == latest["completed"]
    assert observed[0][1] == (mode == "retry")
    assert observed[0][2] == (None if mode == "resume" else 8)


def test_legacy_sql_checkpoint_is_adopted_only_once(db, source_job):
    checkpoint = {"input_id": "legacy", "attempt_count": 1}
    db.add(EvidenceJobState(job_id=source_job, extraction_run={"checkpoint": checkpoint}))
    db.commit()
    first = extraction._claim_annotation(source_job, db, mode="continue")
    assert first["has_checkpoint"]
    assert extraction._load_annotation_checkpoint(source_job) == checkpoint
    with execution.fence(db, source_job, first["run_id"]) as head:
        head.status = "complete"
        extraction._clear_annotation_checkpoint(source_job)
    second = extraction._claim_annotation(source_job, db, mode="continue")
    assert not second["has_checkpoint"]
    assert second["run_id"] != first["run_id"]
    assert extraction._load_annotation_checkpoint(source_job) is None


def test_heartbeat_blocks_reclaim_and_pause_is_cross_connection(db, source_job, monkeypatch):
    monkeypatch.setattr(execution, "LEASE_SECONDS", 1)
    monkeypatch.setattr(execution, "HEARTBEAT_SECONDS", 0.05)
    token = extraction._claim_annotation(source_job, db, mode="start")["run_id"]
    assert execution.begin_worker(db, source_job, token)
    with execution.WorkerLease(db.bind, source_job, token) as lease:
        # Stand in for a model wait longer than the initial lease.
        sleep(1.2)
        with Session(db.bind) as other:
            with pytest.raises(execution.ExecutionBusy):
                execution.claim(other, source_job, actor="competitor", options={})
            other.rollback()
            assert execution.request_pause(other, source_job)
        assert lease.check() is True
        progress_bus.reset(str(source_job))
        with Session(db.bind) as reader:
            value = execution.public_progress(reader, source_job)
            assert value["run_id"] == token and value["pause_requested"]
            assert value["status"] == "running"
    # A dead worker is recoverable, and ownership checks stop its successor writes.
    db.execute(update(AnnotationExecution).where(AnnotationExecution.job_id == source_job)
               .values(lease_expires_at=execution.now() - timedelta(seconds=1)))
    db.commit()
    assert execution.public_progress(db, source_job)["status"] == "interrupted"
    new = extraction._claim_annotation(source_job, db, mode="continue")["run_id"]
    assert new != token
    with pytest.raises(execution.ExecutionLost):
        lease.check()


@pytest.mark.parametrize("callback", ["checkpoint", "snapshot", "failure", "result"])
def test_expired_worker_cannot_write_after_new_run(
    db, source_job, fake_engine, monkeypatch, callback,
):
    old = extraction._claim_annotation(source_job, db, mode="start")["run_id"]
    new_runs = []

    def compute(*args, **kwargs):
        with Session(db.bind) as other:
            other.execute(update(AnnotationExecution)
                          .where(AnnotationExecution.job_id == source_job)
                          .values(lease_expires_at=execution.now() - timedelta(seconds=1)))
            other.commit()
            new = extraction._claim_annotation(source_job, other, mode="rerun")["run_id"]
            new_runs.append(new)
            with execution.fence(other, source_job, new):
                extraction._write_annotation_checkpoint(source_job, {"input_id": "new"})
                extraction._write_annotation_cache(source_job, {"annotation_run_id": new})
        if callback == "checkpoint":
            kwargs["checkpoint_fn"]({"input_id": "old", "attempt_count": 999})
        elif callback == "snapshot":
            kwargs["snapshot_fn"]({})  # fencing must precede even payload validation
        elif callback == "failure":
            raise RuntimeError("late error from old model")
        return {"_checkpoint": {"input_id": "old"}}

    monkeypatch.setattr(extraction, "_compute_annotation", compute)
    asyncio.run(extraction._precompute_annotation_bg(source_job, fake_engine, db.bind, run_id=old))
    assert len(new_runs) == 1
    assert extraction._load_annotation_checkpoint(source_job) == {"input_id": "new"}
    assert json.loads(extraction._annotation_cache_path(source_job).read_text()) == {
        "annotation_run_id": new_runs[0],
    }
    db.expire_all()
    assert db.get(ExtractionJob, source_job).status == "annotating"
    assert db.get(AnnotationExecution, source_job).run_id == new_runs[0]
    assert execution.public_progress(db, source_job)["annotation_stage"] == "queued"
    assert db.get(EvidenceJobState, source_job) is None


def test_invalid_checkpoint_does_not_claim_or_restart(db, source_job):
    path = extraction._annotation_checkpoint_path(source_job)
    path.write_text("broken json")
    with pytest.raises(HTTPException) as error:
        extraction._claim_annotation(source_job, db, mode="continue")
    assert error.value.status_code == 409
    assert db.get(AnnotationExecution, source_job) is None
    assert path.read_text() == "broken json"


def test_excel_get_preview_does_not_start_recognition(db, source_job, fake_engine, monkeypatch):
    from openpyxl import Workbook

    from app.services.extraction import document_annotator

    job = db.get(ExtractionJob, source_job)
    path = extraction._annotation_checkpoint_path(source_job).with_suffix(".xlsx")
    workbook = Workbook()
    workbook.active.append(["名称", "值"])
    workbook.active.append(["产品", "A"])
    workbook.save(path)
    job.source_type, job.document_path = "excel", str(path)
    db.commit()

    def forbidden(*args, **kwargs):
        pytest.fail("a document GET must not start semantic recognition")

    monkeypatch.setattr(document_annotator, "_annotate_texts", forbidden)
    result = extraction.get_annotated_document(source_job, db=db, engine=fake_engine)
    assert result["preview_only"] is True
    assert result["content"]
    assert db.get(AnnotationExecution, source_job) is None
