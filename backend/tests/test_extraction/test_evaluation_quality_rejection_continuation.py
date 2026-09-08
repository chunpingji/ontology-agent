"""Synthetic rejection continuation through real runner/citation execution.

Only model responses are substituted. These are scheduler regressions, not
claims about any real document, reference artifact, or model's accuracy.
"""

import json
from copy import deepcopy
from dataclasses import dataclass, field

from docx import Document

from app.evaluation.legacy_quality_guided_variant import (
    build_legacy_quality_guided_variant as build_quality_guided_variant,
)
from app.schemas.evidence import TaskBudget
from app.services.extraction.extraction_tasks import GenericExtractionRunner
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction import test_evaluation_cmc_product_api_path as path_fixture

APPEARANCE = "urn:rejection-continuation:appearance"


@dataclass
class SourceRecord:
    heading: str
    text: str
    entity_class: str | None = None
    entity_name: str | None = None
    type_supported: bool = True
    edge_binding: bool = True
    edge_reference: bool = True
    properties: dict[str, tuple[str, bool]] = field(default_factory=dict)


def _model(records, calls, product_text):
    by_text = {record.text: record for record in records}
    by_name = {record.entity_name: record for record in records if record.entity_name}

    def model(system, user, schema, budget):
        request = json.loads(user)
        if "local_menu" in request:
            routes = []
            for record_id, record in request["records"].items():
                matching = [by_text[text] for text in record.get("source", []) if text in by_text]
                selected = [
                    alias
                    for alias, predicate in request["local_menu"].items()
                    if any(predicate["iri"] in item.properties for item in matching)
                ]
                routes.append({"record_id": record_id, "predicate_ids": selected})
            calls.append({"stage": "route_records", "request": request})
            return {"records": routes}

        context = json.loads(request["context"])
        task, stage = context["task"], request["stage"]
        targets = [fragment for fragment in context["fragments"] if fragment["fact_eligible"]]
        current = next(
            by_text[fragment["text"]] for fragment in targets if fragment["text"] in by_text
        )
        event = {"stage": stage, "task": task, "text": current.text}
        calls.append(event)

        def quote(text, mention=None):
            fragment = next(item for item in context["fragments"] if item["text"] == text)
            result = {"evidence_id": fragment["anchor"]["evidence_id"]}
            if mention is not None:
                result["text"] = mention
            return result

        def respond(response):
            event["response"] = deepcopy(response)
            return response

        if stage == "verify_entity_types":
            return respond(
                {
                    "decisions": [
                        {
                            "candidate_id": candidate["candidate_id"],
                            "supported": by_name[candidate["text"]].type_supported,
                            "identity_supported": False,
                            "reason": "独立类型判定；目录标签不构成具体对象",
                        }
                        for candidate in request["candidate"]["proposed_entities"]
                    ]
                }
            )
        if stage == "recall" and task["task_kind"] == "entity":
            classes = task["predicate_definition"]["classes"]
            alias = next(
                (
                    alias
                    for alias, definition in classes.items()
                    if current.entity_class
                    and definition["name"] == path_fixture.SCHEMA[current.entity_class]["label"]
                ),
                None,
            )
            return respond(
                {
                    "entities": [
                        {
                            "class_iri": alias,
                            "mention": quote(current.text, current.entity_name),
                            "assertion_spans": [quote(current.text)],
                            "supported": True,
                            "reason": "原文对象提及，仍需独立类型核验",
                        }
                    ]
                    if alias is not None
                    else []
                }
            )
        if stage == "recall" and task["task_kind"] == "relationship":
            return respond(
                {
                    "assertions": [
                        {
                            "object_candidate_id": candidate["candidate_id"],
                            "assertion_status": "affirmed",
                            "assertion_spans": [quote(current.text)],
                        }
                        for candidate in task["object_candidates"]
                    ]
                }
            )
        if stage == "recall":
            value, _ = current.properties[task["predicate_iri"]]
            return respond(
                {
                    "assertions": [
                        {
                            "value": quote(current.text, value),
                            "assertion_status": "affirmed",
                        }
                    ]
                }
            )
        if stage == "verify_binding":
            candidate = request["candidate"]
            supported = (
                current.edge_binding
                if task["task_kind"] == "relationship"
                else current.properties[task["predicate_iri"]][1]
            )
            return respond(
                {
                    "supported": supported,
                    "subject_candidate_id": candidate["subject"]["candidate_id"],
                    "object_candidate_id": (candidate.get("object") or {}).get("candidate_id"),
                    "assertion_status": candidate["assertion_status"],
                    "method": "explicit_assertion",
                    "record_mapping": {},
                    "assertion_spans": [quote(current.text)],
                }
            )
        assert stage == "verify_reference", stage
        groups = context["reference_verification"]
        assert groups["subject_construction_sources"] and groups["current_fact_sources"]
        return respond(
            {
                "supported": current.edge_reference,
                "subject_candidate_id": task["subject"]["candidate_id"],
                "assertion_spans": [quote(product_text), quote(current.text)],
                "reason": "当前记录明确归于该产品" if current.edge_reference else "仅属于其他产品",
            }
        )

    return model


def _run(tmp_path, records, *, product_text, focus_path):
    document = Document()
    for record in records:
        document.add_heading(record.heading, 1)
        document.add_paragraph(record.text)
    source_path = tmp_path / "rejection-continuation.docx"
    document.save(source_path)
    source = analyze_word_core(source_path).ir
    source_before = source.model_dump(mode="json")
    schema = deepcopy(path_fixture.SCHEMA)
    schema[path_fixture.DRUG_PRODUCT]["properties"].append(
        {"iri": APPEARANCE, "label": "外观", "datatype": "string"}
    )
    schema_before = deepcopy(schema)
    calls = []
    base = GenericExtractionRunner(
        schema,
        path_fixture.TestTokenizer(),
        _model(records, calls, product_text),
        model_identity="synthetic-rejection-continuation",
        compact_identifiers=True,
        budget=TaskBudget(
            max_input_tokens=200000,
            max_output_tokens=4096,
            max_tasks=200,
            max_hops=3,
            max_regions_per_task=32,
        ),
    )
    runner = build_quality_guided_variant(base, source, focus_path=focus_path)
    trace = []
    runner.trace_fn = trace.append
    assert runner.execute_task.__func__ is GenericExtractionRunner.execute_task
    result = runner.run(source, effective_class=path_fixture.CMC_REPORT)
    assert len(trace) == len(calls)
    for event, transport in zip(calls, trace, strict=True):
        assert event["stage"] == transport["stage"]
        event["task_id"] = transport["task_id"]
        if event["stage"] != "route_records":
            subject = event["task"]["subject"]
            # Compact candidate aliases are task-local, not stable entity IDs.
            event["subject_ref"] = {
                **subject,
                "candidate_id": transport["references"]["candidate"].get(
                    subject["candidate_id"], subject["candidate_id"]
                ),
            }
    assert source.model_dump(mode="json") == source_before
    assert schema == schema_before
    assert not result.checkpoint["queue"]
    assert not {"task_budget_or_pause", "model_call_budget_exceeded"}.intersection(
        result.diagnostics
    )
    return runner, result, calls


def _assert_staged_records_executed(runner, result, calls, predicate):
    plans = [plan for plan in runner._staged_plans.values() if plan["predicate_iri"] == predicate]
    assert len(plans) == 1, result.diagnostics
    staged = plans[0]
    coverage = [row for row in runner.coverage if row.get("predicate_iri") == predicate]
    expected = {row["record_id"]: row["phase"] for row in staged["records"]}
    assert {row["record_id"]: row["retrieval_phase"] for row in coverage} == expected
    assert len(coverage) == len(expected)
    tasks = {event["task"]["task_id"] for event in result.tasks}
    assert all(row["task_ids"] and set(row["task_ids"]) <= tasks for row in coverage)
    called_task_ids = {event["task_id"] for event in calls if event["stage"] != "route_records"}
    assert all(set(row["task_ids"]) <= called_task_ids for row in coverage)
    source_by_id = {
        record.record_id: runner.record_index.ir.unit(record.source_ranges[0].evidence_id).text
        for record in runner.record_index.records
        if record.record_id in expected
    }
    called_texts = {
        event["text"]
        for event in calls
        if event["stage"] == "recall"
        and event["task"]["task_kind"] == "entity"
        and event["task"]["predicate_definition"]["discovery_relation"]["iri"] == predicate
    }
    assert called_texts == set(source_by_id.values())
    return {source_by_id[record_id]: phase for record_id, phase in expected.items()}


def _event_index(calls, text, stage, *, predicate=None, supported=None):
    return next(
        index
        for index, event in enumerate(calls)
        if event["stage"] == stage
        and event.get("text") == text
        and (predicate is None or event["task"].get("predicate_iri") == predicate)
        and (supported is None or event["response"]["supported"] is supported)
    )


def _assert_current_references(candidate, by_id):
    for ref in [
        candidate.subject,
        candidate.object,
        candidate.path_root,
        *candidate.dependency_refs,
    ]:
        if ref is not None:
            assert by_id[ref.candidate_id].positive_eligible
            assert by_id[ref.candidate_id].revision == ref.revision


def test_all_phase1_type_and_binding_rejections_keep_phase2_product_and_property(tmp_path):
    rejected = [
        SourceRecord(
            "DrugProduct 目录",
            "DrugProduct是目录标签，不是具体制剂。",
            path_fixture.DRUG_PRODUCT,
            "DrugProduct",
            type_supported=False,
        ),
        SourceRecord(
            "DrugProduct 对照甲",
            "DrugProduct 对照甲制剂是口服片剂，但不是本报告描述的产品。",
            path_fixture.DRUG_PRODUCT,
            "对照甲制剂",
            edge_binding=False,
        ),
        SourceRecord(
            "DrugProduct 对照乙",
            "DrugProduct 对照乙制剂是口服片剂，但不是本报告描述的产品。",
            path_fixture.DRUG_PRODUCT,
            "对照乙制剂",
            edge_binding=False,
        ),
    ]
    accepted = SourceRecord(
        "附件甲",
        path_fixture.PRODUCT_TEXT,
        path_fixture.DRUG_PRODUCT,
        path_fixture.PRODUCT_NAME,
        properties={path_fixture.DOSAGE_FORM: ("片剂", True)},
    )
    runner, result, calls = _run(
        tmp_path,
        [*rejected, accepted],
        product_text=accepted.text,
        focus_path=[path_fixture.DESCRIBES],
    )
    phases = _assert_staged_records_executed(runner, result, calls, path_fixture.DESCRIBES)
    assert {text for text, phase in phases.items() if phase == 1} == {r.text for r in rejected}
    assert phases[accepted.text] == 2
    accepted_index = _event_index(calls, accepted.text, "verify_binding", supported=True)
    for record in rejected:
        if not record.type_supported:
            index = _event_index(calls, record.text, "verify_entity_types")
            assert all(not item["supported"] for item in calls[index]["response"]["decisions"])
        else:
            index = _event_index(calls, record.text, "verify_binding", supported=False)
        assert index < accepted_index
    by_id = {candidate.candidate_id: candidate for candidate in result.candidates}
    edges = [
        candidate
        for candidate in result.candidates
        if candidate.kind == "relationship" and candidate.positive_eligible
    ]
    assert len(edges) == 1, result.diagnostics
    edge = edges[0]
    assert edge.predicate_iri == path_fixture.DESCRIBES
    assert by_id[edge.object.candidate_id].text == accepted.entity_name
    assert edge.relationship_path == [path_fixture.DESCRIBES]
    _assert_current_references(edge, by_id)
    properties = [
        candidate
        for candidate in result.candidates
        if candidate.kind == "property" and candidate.positive_eligible
    ]
    assert len(properties) == 1 and properties[0].subject == edge.object
    assert properties[0].literal.normalized_value == "片剂"
    assert edge.candidate_id in {ref.candidate_id for ref in properties[0].dependency_refs}
    assert not any(candidate.assertion_status == "negated" for candidate in result.candidates)


def test_api_reference_rejections_and_property_rejection_keep_later_phase2_success(tmp_path):
    product = SourceRecord(
        "成品身份",
        "本报告描述的制剂产品为晨星制剂；胶囊是对照制剂的剂型，不属于本品。",
        path_fixture.DRUG_PRODUCT,
        path_fixture.PRODUCT_NAME,
        properties={path_fixture.DOSAGE_FORM: ("胶囊", False)},
    )
    rejected_apis = [
        SourceRecord(
            f"ActivePharmaceuticalIngredient 对照{label}",
            f"ActivePharmaceuticalIngredient 对照{label}活性成分仅属于暮光{label}制剂，"
            "不属于晨星制剂。",
            path_fixture.API,
            f"对照{label}活性成分",
            edge_reference=False,
        )
        for label in ("甲", "乙", "丙")
    ]
    appearance = SourceRecord(
        "外观实录",
        "晨星制剂的外观为白色薄膜衣片。",
        properties={APPEARANCE: ("白色薄膜衣片", True)},
    )
    accepted_api = SourceRecord(
        "附件乙", path_fixture.API_TEXT, path_fixture.API, path_fixture.API_NAME
    )
    runner, result, calls = _run(
        tmp_path,
        [product, *rejected_apis, appearance, accepted_api],
        product_text=product.text,
        focus_path=path_fixture.FOCUS_PATH,
    )
    phases = _assert_staged_records_executed(runner, result, calls, path_fixture.HAS_API)
    _assert_staged_records_executed(runner, result, calls, path_fixture.DESCRIBES)
    assert {text for text, phase in phases.items() if phase == 1} == {r.text for r in rejected_apis}
    assert phases[accepted_api.text] == 2
    success = _event_index(
        calls, accepted_api.text, "verify_reference", predicate=path_fixture.HAS_API, supported=True
    )
    for record in rejected_apis:
        failure = _event_index(
            calls, record.text, "verify_reference", predicate=path_fixture.HAS_API, supported=False
        )
        assert failure < success
        assert calls[failure]["subject_ref"] == calls[success]["subject_ref"]
    property_failure = _event_index(
        calls, product.text, "verify_binding", predicate=path_fixture.DOSAGE_FORM, supported=False
    )
    property_success = _event_index(
        calls, appearance.text, "verify_binding", predicate=APPEARANCE, supported=True
    )
    assert property_failure < property_success < success
    by_id = {candidate.candidate_id: candidate for candidate in result.candidates}
    positive = [candidate for candidate in result.candidates if candidate.positive_eligible]
    edges = [candidate for candidate in positive if candidate.kind == "relationship"]
    assert len(edges) == 2, result.diagnostics
    describes = next(edge for edge in edges if edge.predicate_iri == path_fixture.DESCRIBES)
    ingredient = next(edge for edge in edges if edge.predicate_iri == path_fixture.HAS_API)
    assert describes.object == ingredient.subject
    assert by_id[ingredient.object.candidate_id].text == accepted_api.entity_name
    assert ingredient.relationship_path == path_fixture.FOCUS_PATH
    assert describes.candidate_id in {ref.candidate_id for ref in ingredient.dependency_refs}
    assert any(binding.method == "identity_reference" for binding in ingredient.bindings)
    for edge in edges:
        _assert_current_references(edge, by_id)
    assert {record.entity_name for record in rejected_apis} <= {
        candidate.text for candidate in positive if candidate.kind == "entity"
    }
    assert not any(
        candidate.kind == "property" and candidate.predicate_iri == path_fixture.DOSAGE_FORM
        for candidate in positive
    )
    saved_property = next(
        candidate
        for candidate in positive
        if candidate.kind == "property" and candidate.predicate_iri == APPEARANCE
    )
    assert saved_property.subject == describes.object
    assert saved_property.literal.normalized_value == "白色薄膜衣片"
    assert describes.candidate_id in {ref.candidate_id for ref in saved_property.dependency_refs}
    assert "reference_not_supported" in result.diagnostics
    assert result.completion == "incomplete"
    assert not any(candidate.assertion_status == "negated" for candidate in result.candidates)
