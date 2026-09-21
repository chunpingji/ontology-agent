"""Table roles come from explicit headers or strong repeated source structure."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core

PROSE = (
    "Keep the first source sentence with its own row. "
    "Continue the stated operation only after checking the recorded result."
)


def table_document(values):
    document = Document()
    table = document.add_table(rows=len(values), cols=len(values[0]))
    for row, texts in zip(table.rows, values, strict=True):
        for cell, text in zip(row.cells, texts, strict=True):
            cell.text = text
    return document, table


def mark_header(row, value):
    header = OxmlElement("w:tblHeader")
    if value is not None:
        header.set(qn("w:val"), value)
    row._tr.get_or_add_trPr().append(header)


def analyze(document, tmp_path):
    path = tmp_path / "generic-table.docx"
    document.save(path)
    result = analyze_word_core(path)
    return result.structure.tables[0], RecordIndex(result.ir)


@pytest.mark.parametrize("values", [
    [("A", PROSE), ("B", PROSE + " Preserve the next record too.")],
    [("1", "Question alpha", "choice", "answer"),
     ("2", "Question beta", "choice", "answer"),
     ("3", "Question gamma", "choice", "answer")],
])
def test_headerless_rows_are_all_fact_sources_and_never_other_rows_headers(tmp_path, values):
    document, _ = table_document(values)
    parsed, index = analyze(document, tmp_path)
    assert parsed.header_row_count == 0
    assert parsed.row_indices == list(range(len(values)))
    assert parsed.headers == [f"col{column}" for column in range(len(values[0]))]
    assert [record.row_index for record in index.records] == list(range(len(values)))
    for record, expected in zip(index.records, values, strict=True):
        assert [unit.text for unit in record.source_units] == list(expected)
        assert record.header_units == ()
        assert all(not index.tables.is_header(unit) for unit in record.source_units)
        assert all(index.ir.resolve(index.ir.anchor(unit.evidence_id)) == unit.text
                   for unit in record.source_units)


@pytest.mark.parametrize("true_value", [None, "1", "true", "on"])
@pytest.mark.parametrize("false_value", ["0", "false", "off"])
def test_explicit_headers_override_data_shape_and_end_at_false(tmp_path, true_value, false_value):
    document, table = table_document([("A", PROSE), ("B", PROSE), ("C", PROSE)])
    mark_header(table.rows[0], true_value)
    mark_header(table.rows[1], false_value)
    mark_header(table.rows[2], true_value)
    parsed, index = analyze(document, tmp_path)
    assert parsed.header_row_count == 1
    assert parsed.row_indices == [1, 2]
    assert all(unit.row_index == 0 for record in index.records for unit in record.header_units)
    assert all(unit.row_index > 0 for record in index.records for unit in record.source_units)


def test_contiguous_explicit_headers_do_not_require_horizontal_merges(tmp_path):
    document, table = table_document([
        ("Group A", "Group B"), ("Field A", "Field B"), ("value-A", "value-B"),
    ])
    mark_header(table.rows[0], None)
    mark_header(table.rows[1], "true")
    parsed, index = analyze(document, tmp_path)
    assert parsed.header_row_count == 2
    assert parsed.row_indices == [2]
    assert {unit.row_index for unit in index.records[0].header_units} == {0, 1}


@pytest.mark.parametrize("false_value", ["0", "false", "off"])
def test_disabled_repeat_header_does_not_hide_strong_data_shape(tmp_path, false_value):
    document, table = table_document([("A", PROSE), ("B", PROSE)])
    mark_header(table.rows[0], false_value)
    parsed, _ = analyze(document, tmp_path)
    assert parsed.header_row_count == 0


@pytest.mark.parametrize("row_index,value", [(0, "unknown"), (0, "TRUE"), (1, "true")])
def test_invalid_or_nonleading_header_marks_cannot_authorize_headerless_rows(
    tmp_path, row_index, value,
):
    document, table = table_document([("A", PROSE), ("B", PROSE)])
    mark_header(table.rows[row_index], value)
    parsed, _ = analyze(document, tmp_path)
    assert parsed.header_row_count == 1


@pytest.mark.parametrize("values", [
    [("Label", "Description"), ("A", PROSE), ("B", PROSE)],
    [("1", "2"), ("2", "3"), ("3", "4")],
    [("1", "alpha"), ("3", "beta"), ("4", "gamma")],
    [("1", "alpha"), ("2", "beta")],
    [("A", PROSE)],
    [("A", "short sentence."), ("B", "another short sentence.")],
    [("A", "Long label " * 10), ("B", "Other long label " * 10)],
])
def test_uncertain_first_rows_keep_conservative_header_behavior(tmp_path, values):
    document, _ = table_document(values)
    parsed, _ = analyze(document, tmp_path)
    assert parsed.header_row_count == 1


@pytest.mark.parametrize("complexity", ["vertical_merge", "omitted_cell", "nested_table"])
def test_complex_topology_is_not_flattened_into_headerless_evidence(tmp_path, complexity):
    document, table = table_document([("A", PROSE), ("B", PROSE)])
    if complexity == "vertical_merge":
        table.cell(0, 0).merge(table.cell(1, 0))
    elif complexity == "omitted_cell":
        row = table.rows[1]._tr
        row.remove(row.tc_lst[0])
        before = OxmlElement("w:gridBefore")
        before.set(qn("w:val"), "1")
        row.get_or_add_trPr().append(before)
    else:
        table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "nested"
    parsed, _ = analyze(document, tmp_path)
    assert parsed.header_row_count == 1
