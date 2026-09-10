"""Event-controlled async ranking, write barriers and checkpoint replay."""

from __future__ import annotations

import copy
import json
from threading import Event

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.contracts import (
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
)
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor, TaskOutcome
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.semantic_reranker import RankingPolicy, RankingService
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.model_runtime import check_cancelled

ROOT, FIRST, SECOND = "urn:async:Report", "urn:async:first", "urn:async:second"


class SimulatedCrash(RuntimeError):
    pass


def fixture(tmp_path):
    document = Document()
    document.add_heading("原文记录", 1)
    for number in range(12):
        document.add_paragraph(f"独立来源记录 {number}，未明确给出请求属性。")
    path = tmp_path / "async-ranking.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="async-test",
    )
    definition = OntologyClassDefinition(
        iri=ROOT,
        label="报告",
        source_hash="source",
        declared_properties=[
            SlotSpec(iri=iri, label=label, declared_by=[ROOT])
            for iri, label in ((FIRST, "第一属性"), (SECOND, "第二属性"))
        ],
    )
    ontology = OntologySnapshot(
        snapshot_id="async-ontology",
        ontology_hash=evidence_hash(definition),
        classes={ROOT: definition},
        created_from="frozen_fixture",
    )
    return ontology, dict(
        recognition_run_id="async-run",
        run_fingerprint="async-fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=ROOT,
        root_class_label="报告",
        filename=path.name,
    )


class ControlledRanking:
    identity = {"model": "event-controlled-ranking"}

    def __init__(self, *, block_second=False):
        self.block_second = block_second
        self.second_preparing = Event()
        self.second_admitted = Event()
        self.second_started = Event()
        self.release_second = Event()
        self.calls = []
        self.observations = []

    def count_tokens(self, text):
        if f'"iri":"{SECOND}"' in text:
            self.second_preparing.set()
        return max(1, len(text) // 4)

    def embed(self, texts):
        self.calls.append(("embed", len(texts)))
        if self.block_second and any(f'"iri":"{SECOND}"' in text for text in texts):
            self.second_started.set()
            while not self.release_second.wait(0.01):
                check_cancelled()
        self.observations.append({"operation": "embed", "input_tokens": len(texts)})
        return [[1.0, 0.5] for _ in texts]

    def score_pairs(self, pairs):
        self.calls.append(("score_pairs", len(pairs)))
        self.observations.append({"operation": "score_pairs", "input_tokens": len(pairs)})
        return [float(len(text)) for _, text in pairs]


class NoFactsAdapter:
    model_identity = "controlled-recognition"

    def __init__(self, ranking=None):
        self.ranking = ranking
        self.calls = []
        self.ready_ran_while_second_scoring = Event()

    def inspect(self, task, context, predicate, menu):
        self.calls.append(task.task_id)
        if (
            self.ranking is not None
            and predicate.iri == FIRST
            and self.ranking.second_admitted.is_set()
        ):
            # The cost barrier was acknowledged before this semantic dispatch,
            # so the ranking worker can enter its blocked request independently.
            assert self.ranking.second_started.wait(2), "ranking request never started"
            self.ready_ran_while_second_scoring.set()
        return TaskOutcome(
            semantic_outcome="not_checked",
            reason_code="no_candidate_observed",
            reason="当前记录没有提出候选，不代表全文否定。",
            model_calls=1,
        )


def run(ontology, arguments, adapter, ranking, policy, **kwargs):
    return OntologyGuidedExecutor(
        ontology=ontology,
        engine=object(),
        adapter=adapter,
        ranking_service=RankingService(policy, ranking),
        max_tasks=100,
    ).run(**arguments, **kwargs)


def test_ready_slot_runs_during_scoring_and_checkpoint_replays_exclusions_and_costs(tmp_path):
    ontology, arguments = fixture(tmp_path)
    ranking = ControlledRanking(block_second=True)
    adapter = NoFactsAdapter(ranking)
    policy = RankingPolicy(mode="semantic", ranking_timeout=10, technical_retry_limit=1)
    snapshots = []
    batches = []

    def persist(state):
        snapshots.append(copy.deepcopy(state))
        if ranking.second_preparing.is_set() and state["service"]["costs"]["model_calls"]:
            ranking.second_admitted.set()

    def checkpoint(batch):
        batches.append(batch)
        if adapter.ready_ran_while_second_scoring.is_set():
            raise SimulatedCrash("checkpoint committed while another slot was scoring")

    with pytest.raises(SimulatedCrash):
        run(
            ontology,
            arguments,
            adapter,
            ranking,
            policy,
            ranking_hook=persist,
            batch_hook=checkpoint,
        )
    assert ranking.second_started.is_set()
    assert adapter.ready_ran_while_second_scoring.is_set()
    checkpoint = batches[-1]
    saved = checkpoint.task_outcomes
    assert any(item.get("ranking_excluded_slots") for item in saved)
    saved_ids = {item["task"]["task_id"] for item in saved}
    durable_cost = snapshots[-1]["service"]["costs"]["tokens"]
    durable_attempts = snapshots[-1]["service"]["request_attempts"]
    recovered_model = ControlledRanking()
    recovered_adapter = NoFactsAdapter()
    recovered = run(
        ontology,
        arguments,
        recovered_adapter,
        recovered_model,
        policy,
        resume_state={
            "task_outcomes": saved,
            "frontier": checkpoint.frontier,
            "recall_ledger": checkpoint.recall_ledger,
            "dependency_index": checkpoint.dependency_index,
        },
        ranking_state=snapshots[-1],
    )
    outcomes = [value for kind, value in recovered.events if kind == "task_outcome"]
    assert outcomes[: len(saved)] == saved
    assert not saved_ids.intersection(recovered_adapter.calls)
    all_ids = [item["task"]["task_id"] for item in outcomes]
    assert len(all_ids) == len(set(all_ids)) == 24
    assert recovered.graph.progress.completion == "in_scope_complete"
    state = recovered.ranking_state["service"]
    assert state["costs"]["tokens"] >= durable_cost
    assert state["costs"]["technical_retries"] >= 1
    assert all(state["request_attempts"][key] >= value for key, value in durable_attempts.items())


def test_failed_main_thread_cost_commit_prevents_actual_ranking_inference(tmp_path):
    ontology, arguments = fixture(tmp_path)
    ranking = ControlledRanking()
    adapter = NoFactsAdapter()
    observed = []

    def rejected_commit(state):
        observed.append(state)
        assert state["service"]["costs"]["model_calls"] == 1
        raise RuntimeError("durable write rejected")

    with pytest.raises(RuntimeError, match="durable write rejected"):
        run(
            ontology,
            arguments,
            adapter,
            ranking,
            RankingPolicy(mode="semantic", technical_retry_limit=2),
            ranking_hook=rejected_commit,
        )
    assert len(observed) == 1
    assert ranking.calls == []
    assert adapter.calls == []
    assert json.dumps(observed[0]["service"]["request_attempts"])


@pytest.mark.parametrize("pool_size,allowed_first_turns", [(2, 2), (4, 5)])
def test_exhausted_pool_waits_except_for_a_due_ledger_exploration(
    tmp_path,
    pool_size,
    allowed_first_turns,
):
    ontology, arguments = fixture(tmp_path)
    ranking = ControlledRanking(block_second=True)
    first_during_wait = []
    unauthorized = []
    reached_wait_boundary = Event()

    class WatchingAdapter(NoFactsAdapter):
        def inspect(self, task, context, predicate, menu):
            if predicate.iri == FIRST and not ranking.release_second.is_set():
                first_during_wait.append(task)
                if len(first_during_wait) > allowed_first_turns:
                    unauthorized.append(task)
                    # End the controlled request so a regression fails a precise
                    # assertion instead of hanging until the model deadline.
                    ranking.release_second.set()
            return super().inspect(task, context, predicate, menu)

    adapter = WatchingAdapter(ranking)

    def persist(state):
        if ranking.second_preparing.is_set() and state["service"]["costs"]["model_calls"]:
            ranking.second_admitted.set()

    def progress(boundary):
        if (
            boundary == "ranking_wait"
            and ranking.second_started.is_set()
            and len(first_during_wait) == allowed_first_turns
        ):
            reached_wait_boundary.set()
            ranking.release_second.set()
        return True

    result = OntologyGuidedExecutor(
        ontology=ontology,
        engine=object(),
        adapter=adapter,
        progress_hook=progress,
        ranking_service=RankingService(
            RankingPolicy(mode="semantic", pool_size=pool_size, ranking_timeout=10),
            ranking,
        ),
        max_tasks=100,
    ).run(**arguments, ranking_hook=persist)
    assert reached_wait_boundary.is_set()
    assert unauthorized == []
    assert len(first_during_wait) == allowed_first_turns
    if pool_size == 4:
        assert first_during_wait[-1].ranking_epoch_seq is None
        assert first_during_wait[-1].pool_rank is None
    assert result.graph.progress.completion == "in_scope_complete"


def test_non_rerankable_records_still_receive_actual_exploration_with_pause_policy(tmp_path):
    ontology, arguments = fixture(tmp_path)
    ranking, adapter = ControlledRanking(), NoFactsAdapter()
    policy = RankingPolicy(
        mode="semantic", failure_policy="pause", max_tokens_per_pair=1,
    )
    result = run(ontology, arguments, adapter, ranking, policy)
    assert adapter.calls
    assert ranking.calls == []
    assert result.graph.progress.records_unattempted == 0
    assert result.graph.progress.records_incomplete == 0
    epochs = result.ranking_state["service"]["epochs"]
    assert epochs
    assert all(epoch["status"] == "committed" for epoch in epochs)
    assert all(epoch["actual_ranking_mode"] == "deterministic" for epoch in epochs)
    assert all(epoch["reason"] == "no_rerankable_records" for epoch in epochs)
    assert all(
        observation["score_status"] == "not_selected"
        and observation["raw_rerank_score"] is None
        and observation["view_status"] == "not_rerankable"
        for epoch in epochs for observation in epoch["observations"]
    )
