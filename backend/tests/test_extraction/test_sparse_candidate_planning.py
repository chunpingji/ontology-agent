"""Actual candidate obligations, bounded search completion and frozen replay."""

import json
from dataclasses import replace

import pytest
from docx import Document

from app.services.extraction.ontology_guided.adaptive_retrieval import (
    AdaptivePolicy,
    CalibrationProfile,
)
from app.services.extraction.ontology_guided.contracts import (
    CANDIDATE_POLICY_VERSION,
    GraphProperty,
    RunProgress,
    SubjectRef,
    VersionedRef,
)
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor, TaskOutcome
from app.services.extraction.ontology_guided.heuristic_search import HeuristicSearchPolicy
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.retrieval import plan_slot, validate_record_universe
from app.services.extraction.ontology_guided.semantic_reranker import RankingPolicy, RankingService
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_adaptive_retrieval import adaptive_fixture
from tests.test_extraction.test_heuristic_executor import _product_arguments
from tests.test_extraction.test_ontology_guided_core import (
    APPEARANCE,
    FakeAdapter,
    ontology,
)
from tests.test_extraction.test_semantic_ranking_execution import (
    FIRST,
    ControlledRanking,
    NoFactsAdapter,
    fixture,
)


def executor(snapshot, *, adapter=None, **kwargs):
    return OntologyGuidedExecutor(
        ontology=snapshot, engine=object(), adapter=adapter or NoFactsAdapter(),
        evidence_repair=True, incremental_performance=True,
        candidate_policy=CANDIDATE_POLICY_VERSION,
        heuristic_policy=kwargs.pop("heuristic_policy", HeuristicSearchPolicy.durable(
            incremental=True, exploration_page_size=2,
        )),
        max_tasks=kwargs.pop("max_tasks", 100), **kwargs,
    )


def resume_arguments(batch):
    # A process boundary loses object identity, mutable caches and shared Python
    # references. Round-trip JSON rather than replaying the same batch objects.
    saved = json.loads(batch.model_dump_json())
    return {
        "resume_state": {key: saved[key] for key in (
            "task_outcomes", "frontier", "recall_ledger", "dependency_index", "diagnostics",
        )},
        "ranking_state": saved["ranking_state"],
        "model_call_state": saved["model_call_state"],
    }


def test_only_admitted_candidates_become_obligations_and_policy_can_finish(tmp_path):
    snapshot, args = fixture(tmp_path)
    batches = []
    result = executor(snapshot).run(**args, batch_hook=batches.append)
    progress = result.graph.progress
    assert len(RecordIndex(args["ir"]).records) == 12
    assert progress.records_planned == progress.records_examined == 4
    assert progress.records_incomplete == progress.records_unattempted == 0
    assert progress.candidate_policy == CANDIDATE_POLICY_VERSION
    assert progress.completion == "policy_complete"
    assert progress.stop_reason == "candidate_search_exhausted"
    assert not result.graph.properties
    assert all(batch.graph.progress.completion == "incomplete" for batch in batches)
    for plan in result.retrieval_plans:
        assert len(plan["ledger"]) == len(plan["records"]) == 2
        assert plan["frozen_record_ids"] == []
        assert plan["search_scope_ref"]["record_count"] == 12
        assert all(entry["task_ids"] for entry in plan["ledger"].values())
    summary = next(value for kind, value in result.events if kind == "heuristic_complete")
    assert all("deferred_record_ids" not in search for search in summary["slots"])
    assert all(search["deferred_count"] == 10 for search in summary["slots"])


def test_unselected_domain_is_shared_instead_of_repeated_per_plan(tmp_path):
    snapshot, args = fixture(tmp_path)
    document = Document()
    for number in range(500):
        document.add_paragraph(f"无目标字段的独立背景 {number}")
    path = tmp_path / "large-background.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    metadata = prepare_metadata(analysis.ir, section_tree=analysis.structure.section_tree.to_dict(),
                                summary_version="sparse-test")
    subject = SubjectRef(entity_id="root", revision=1, class_iri=args["root_class_iri"],
                         is_document_root=True)
    slots = snapshot.classes[args["root_class_iri"]].declared_properties
    plans = [plan_slot(subject, slot, index, metadata, ontology_hash=snapshot.ontology_hash,
                       sparse_candidates=True) for slot in slots]
    for plan in plans:
        validate_record_universe(plan, index)
        assert not plan.records and not plan.ledger and not plan.frozen_record_ids
        assert plan.search_scope_ref["record_count"] == 500
        assert len(plan.model_dump_json()) < 1100
    assert plans[0].search_scope_ref == plans[1].search_scope_ref
    bad = plans[0].model_copy(update={"search_scope_ref": {
        **plans[0].search_scope_ref, "record_hash": "0" * 64,
    }})
    with pytest.raises(ValueError, match="shared record universe"):
        validate_record_universe(bad, index)


def test_large_search_snapshot_keeps_only_hits_admissions_and_bounded_exploration(tmp_path):
    search, _service, arguments = adaptive_fixture(
        tmp_path, mode="enhanced", paragraphs=[f"独立背景 {i}" for i in range(500)],
    )
    plan = plan_slot(search.plan.subject, search.predicate, arguments["index"],
                     arguments["metadata"], ontology_hash="frozen-ontology", sparse_candidates=True)
    search = replace(search, plan=plan, policy=HeuristicSearchPolicy.durable(
        adaptive=True, exploration_page_size=2,
    ))
    assert search.next_admission() is None and search.needs_semantic
    search.skip_semantic("ranking_policy_deterministic")
    page = search.next_admission()
    assert len(page.record_ids) == 2
    for rid in page.record_ids:
        search.observe(rid, "not_checked", True, "no_candidate_observed")
    assert search.next_admission() is None and search.status == "pass_exhausted"
    state = search.snapshot()
    assert state["universe_count"] == 500 and state["deferred_count"] == 498
    assert "deferred_record_ids" not in state
    assert state["resume_state"]["matches"] == {}
    assert sum(map(len, state["resume_state"]["candidates"].values())) == 2
    assert len(json.dumps(state)) < 6000

    fresh = replace(search, plan=plan)
    fresh.exhaust_dependency()
    exhausted = fresh.snapshot()
    assert exhausted["resume_state"]["dependency_exhausted"] is True
    assert exhausted["resume_state"]["dispositions"] == {}
    assert fresh.diagnostics().records_dependency_exhausted == 500
    assert len(json.dumps(exhausted)) < 2500
    restored = replace(search, plan=plan)
    restored.restore(json.loads(json.dumps(exhausted)))
    assert restored.diagnostics().records_dependency_exhausted == 500
    assert not restored.needs_evaluation_ids()


@pytest.mark.parametrize("failure", ["technical", "undetermined"])
def test_finished_undetermined_is_a_result_but_technical_failure_is_incomplete(tmp_path, failure):
    snapshot, args = fixture(tmp_path)

    class IncompleteAdapter:
        model_identity = "incomplete-test"

        def inspect(self, *_args):
            return TaskOutcome(
                semantic_outcome="undetermined" if failure == "undetermined" else "not_checked",
                complete=failure == "undetermined", reason_code="unresolved_source",
                reason="已完成核验但原文仍不确定" if failure == "undetermined" else "模型请求失败",
            )

    result = executor(snapshot, adapter=IncompleteAdapter()).run(**args)
    assert result.graph.progress.completion == (
        "incomplete" if failure == "technical" else "policy_complete"
    )
    assert result.graph.progress.stop_reason == (
        "attempted_incomplete" if failure == "technical" else "candidate_search_exhausted"
    )
    assert result.graph.progress.records_planned == 4
    assert result.graph.progress.records_incomplete or result.graph.progress.unresolved_claims


def test_dispatch_budget_leaves_actual_admitted_work_incomplete(tmp_path):
    snapshot, args = fixture(tmp_path)
    result = executor(snapshot, max_tasks=1).run(**args)
    assert result.graph.progress.completion == "incomplete"
    assert result.graph.progress.stop_reason == "task_budget_exhausted"
    assert result.graph.progress.records_examined == 1
    assert result.graph.progress.records_planned == 2
    assert result.graph.progress.records_unattempted == 1


def test_required_ranking_failure_cannot_become_empty_candidate_success(tmp_path):
    snapshot, args = fixture(tmp_path)
    result = executor(snapshot, ranking_service=RankingService(
        RankingPolicy(mode="semantic", failure_policy="pause"), None,
    )).run(**args)
    assert result.graph.progress.completion == "incomplete"
    assert result.graph.progress.stop_reason == "ranking_paused"
    assert result.graph.progress.records_planned == 0


def test_frozen_pruning_can_end_without_fabricating_reviewed_records(tmp_path):
    snapshot, args = fixture(tmp_path)
    _, service, _ = adaptive_fixture(tmp_path, dense=2, rerank=0)
    raw = service.adaptive_policy.calibration.model_dump(mode="json", exclude={"profile_hash"})
    raw["predicate_iris"] = [
        p.iri for p in snapshot.classes[args["root_class_iri"]].declared_properties
    ]
    policy = AdaptivePolicy(mode="enforce", evaluation_only=True,
                            calibration=CalibrationProfile.model_validate(raw))
    result = executor(
        snapshot, adaptive_policy=policy, ranking_service=service,
        heuristic_policy=HeuristicSearchPolicy.durable(adaptive=True),
    ).run(**args)
    assert result.graph.progress.completion == "policy_complete"
    assert result.graph.progress.records_planned == result.graph.progress.records_examined == 0
    assert result.graph.progress.retrieval_diagnostics.records_soft_pruned == 24
    assert not result.graph.edges and not result.graph.properties


def test_child_fields_need_proved_reachable_subject(tmp_path):
    args = _product_arguments(tmp_path)
    result = executor(ontology(), adapter=FakeAdapter()).run(**args)
    assert any(plan["predicate_iri"] == APPEARANCE for plan in result.retrieval_plans)

    class Unproved(FakeAdapter):
        def inspect(self, *arguments):
            outcome = super().inspect(*arguments)
            for edge in outcome.edges:
                edge.policy_eligible = False
                edge.model_supported = False
            return outcome

    rejected = executor(ontology(), adapter=Unproved()).run(**args)
    assert all(plan["predicate_iri"] != APPEARANCE for plan in rejected.retrieval_plans)


@pytest.mark.parametrize("recheck_complete", [True, False])
def test_conflicting_values_finish_after_recheck_but_failed_recheck_stays_incomplete(
    tmp_path, recheck_complete,
):
    snapshot, args = fixture(tmp_path)
    snapshot.classes[args["root_class_iri"]].declared_properties[0].max_count = 1
    document = Document()
    document.add_paragraph("第一属性：1 mg/day。")
    document.add_paragraph("第一属性：2 mg/day。")
    path = tmp_path / "conflicting-values.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    args.update(ir=analysis.ir, metadata=prepare_metadata(
        analysis.ir, section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="conflict-test",
    ))
    rechecks = []

    class ConflictAdapter(NoFactsAdapter):
        def inspect(self, task, context, predicate, menu):
            if task.retry_kind:
                rechecks.append(task.task_id)
                return TaskOutcome(
                    semantic_outcome="undetermined", complete=recheck_complete,
                    reason_code=("conflicting_source_values" if recheck_complete
                                 else "model_timeout"),
                    reason="独立对比后仍有冲突" if recheck_complete else "补验模型超时",
                )
            if predicate.iri != FIRST:
                return super().inspect(task, context, predicate, menu)
            anchor = next(fragment.anchor for fragment in context.fragments
                          if fragment.fact_eligible)
            prop = GraphProperty(
                candidate_id=f"value-{task.record_id}", revision=1,
                subject_ref=VersionedRef(id=task.subject.entity_id, revision=task.subject.revision),
                predicate_iri=predicate.iri, predicate_label=predicate.label,
                raw_value=task.record_id, decision_status="supported",
                structural_valid=True, model_supported=True, policy_eligible=True,
                proof_ref=VersionedRef(id=f"proof-{task.record_id}", revision=1),
                decision_refs=[VersionedRef(id=f"decision-{task.record_id}", revision=1)],
                evidence_refs=[anchor], predicate_evidence_refs=[anchor],
            )
            return TaskOutcome(semantic_outcome="supported", reason_code="supported",
                               reason="来源确实给出该数值", properties=[prop])

    result = executor(snapshot, adapter=ConflictAdapter()).run(**args)
    assert rechecks
    assert result.graph.progress.unresolved_claims > 0
    assert result.graph.progress.completion == (
        "policy_complete" if recheck_complete else "incomplete"
    )
    assert bool(result.graph.progress.records_incomplete) is not recheck_complete


@pytest.mark.parametrize("semantic", [False, True, "adaptive"])
def test_cold_recovery_replays_only_committed_candidates_and_preserves_cost(tmp_path, semantic):
    snapshot, args = fixture(tmp_path)
    model = ControlledRanking()
    def ranking():
        return RankingService(
            RankingPolicy(mode="semantic", pool_size=4, ranking_timeout=30), model,
        ) if semantic else None
    options = ({"adaptive_policy": AdaptivePolicy(mode="enhanced"),
                "heuristic_policy": HeuristicSearchPolicy.durable(adaptive=True)}
               if semantic == "adaptive" else {})
    saved = []

    class Interrupted(Exception):
        pass

    def save(batch):
        saved.append(batch)
        if len(saved) == 2:
            raise Interrupted()

    class PaidAdapter(NoFactsAdapter):
        def inspect(self, task, context, predicate, menu):
            context.before_model_call("candidate-check", 1)
            return super().inspect(task, context, predicate, menu)

    adapter = PaidAdapter()
    with pytest.raises(Interrupted):
        executor(snapshot, adapter=adapter, ranking_service=ranking(), **options).run(
            **args, batch_hook=save,
        )
    prior_calls = list(adapter.calls)
    resumed_adapter = PaidAdapter()
    result = executor(snapshot, adapter=resumed_adapter, ranking_service=ranking(), **options).run(
        **args, **resume_arguments(saved[-1]),
    )
    assert not set(prior_calls) & set(resumed_adapter.calls)
    assert result.graph.progress.completion == "policy_complete"
    assert result.graph.progress.model_calls == len(prior_calls) + len(resumed_adapter.calls)
    assert result.graph.progress.model_calls_reserved == result.graph.progress.model_calls
    assert result.graph.progress.model_calls_unresolved == 0
    assert result.graph.progress.records_planned == result.graph.progress.records_examined


def test_candidate_policy_cannot_silently_replay_legacy_checkpoint(tmp_path):
    snapshot, args = fixture(tmp_path)
    old_batches = []
    OntologyGuidedExecutor(ontology=snapshot, adapter=NoFactsAdapter(), engine=object(),
                          max_tasks=1).run(**args, batch_hook=old_batches.append)
    with pytest.raises(ValueError, match="candidate planning policy changed"):
        executor(snapshot).run(**args, **resume_arguments(old_batches[-1]))
    assert "search_scope_ref" not in old_batches[-1].recall_ledger[next(iter(
        old_batches[-1].recall_ledger
    ))]


def test_policy_complete_retains_unresolved_semantics_but_rejects_unfinished_calls():
    progress = RunProgress(completion="policy_complete", candidate_policy=CANDIDATE_POLICY_VERSION,
                           unresolved_claims=1, undetermined=1)
    assert progress.unresolved_claims == 1
    with pytest.raises(ValueError, match="unresolved work"):
        RunProgress(completion="policy_complete", candidate_policy=CANDIDATE_POLICY_VERSION,
                    model_calls_unresolved=1)
