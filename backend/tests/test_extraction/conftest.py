import pytest
from docx import Document

from app.schemas.evidence import Candidate, DocumentProvenance
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def subject_document(tmp_path):
    doc = Document()
    doc.add_heading("药品 A", 1)
    doc.add_paragraph("规格：250 mg")
    doc.add_heading("药品 B", 1)
    doc.add_paragraph("规格：500 mg")
    path = tmp_path / "subjects.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    subjects = []
    for index in (0, 2):
        unit = ir.evidence_units[index]
        subjects.append(
            Candidate(
                candidate_id=f"subject-{index}",
                kind="entity",
                text=unit.text,
                class_iri="urn:Drug",
                validation_status="passed",
                provenance=[
                    DocumentProvenance(
                        anchors=[ir.anchor(unit.evidence_id, 0, len(unit.text))],
                        excerpts=[unit.text],
                    )
                ],
            )
        )
    return ir, subjects
