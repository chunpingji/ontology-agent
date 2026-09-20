"""Bounded table batches share discovery work without borrowing another row's facts."""

from dataclasses import replace
from types import SimpleNamespace

import pytest
from docx import Document

from app.models.document_analysis import DocumentRunCurrentState
from app.services.document_analysis import current_state
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.record_discovery import (
    ContextualDiscoveryPolicy,
    RecordDiscoveryPolicy,
)
from app.services.extraction.ontology_guided.record_search import record_text
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.table_reading import table_reading_groups
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction import test_record_executor as fixture
from tests.test_extraction.test_contextual_record_context import context_for
from tests.test_extraction.test_contextual_record_executor import setup_contextual

pytest_plugins = ["tests.test_extraction.test_tool_engine_resume"]

LEAD = "残留物溶解度结果如下表（酸或碱水溶液的清洁方式优于溶剂清洁，三类溶剂优于二类溶剂）。"
CAPTION = "表3溶解度信息"


def document_index(tmp_path, *, lead=LEAD, caption=CAPTION, separator=None, rows=2,
                   columns=2, table=True, nested=False, names=None):
    names = names or ["1234-3", "1234-2", *(f"物料{row}" for row in range(2, rows))]
    document = Document()
    document.add_heading("试验资料", 1)
    if lead:
        document.add_paragraph(lead)
    if separator == "heading":
        document.add_heading("另一项试验", 1)
    elif separator:
        document.add_paragraph(separator)
    if caption:
        document.add_paragraph(caption)
    if table:
        grid = document.add_table(rows=rows + 1, cols=columns)
        if nested:
            grid = grid.cell(0, 0).add_table(rows=rows + 1, cols=columns)
        for column in range(columns):
            grid.cell(0, column).text = ("名称", "溶解度")[column]
        for row in range(rows):
            grid.cell(row + 1, 0).text = names[row]
            if columns > 1:
                grid.cell(row + 1, 1).text = ("易溶", "微溶")[row % 2]
    path = tmp_path / "table-reading.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    return RecordIndex(analysis.ir), analysis


def groups(index, **overrides):
    return list(table_reading_groups(index, **{
        "max_group_chars": 1200, "max_records": 8, **overrides,
    }))


def context(index, sources):
    return context_for(index, SimpleNamespace(record_ids=sources, section_node_ids=()))[1]


def test_intro_and_caption_join_consecutive_data_rows_with_full_original_sources(tmp_path):
    index, _ = document_index(tmp_path)
    before = index.ir.model_dump_json()
    joint, = groups(index)
    assert index.by_id[joint[0]].kind == "table_row"
    assert [index.by_id[rid].text for rid in joint[1:3]] == [LEAD, CAPTION]
    row2 = next(r for r in index.records if r.kind == "table_row" and r.record_id != joint[0])
    view = context(index, joint)
    assert {f.text for f in view.fragments if f.fact_eligible} == {
        "1234-3", "易溶", "1234-2", "微溶", LEAD, CAPTION,
    }
    assert index.ir.model_dump_json() == before
    assert row2.record_id in joint
    assert {"名称", "溶解度"} <= {f.text for f in view.fragments if not f.fact_eligible}
    assert all(index.ir.resolve(f.anchor) == f.text for f in view.fragments)


def test_other_rows_get_intro_for_retrieval_and_binding_without_its_fact_permission(tmp_path):
    index, _ = document_index(tmp_path)
    row = [r for r in index.records if r.kind == "table_row"][1]
    view = context(index, (row.record_id,))
    assert {f.text for f in view.fragments if f.fact_eligible} == {"1234-2", "微溶"}
    assert {LEAD, CAPTION} <= {f.text for f in view.fragments if not f.fact_eligible}
    assert not {"1234-3", "易溶"} & {f.text for f in view.fragments}
    assert LEAD in record_text([row.record_id], index)
    assert record_text(list(groups(index)[0]), index).count(LEAD) == 1


@pytest.mark.parametrize("options", [
    {"table": False}, {"rows": 0}, {"columns": 1}, {"nested": True},
    {"separator": "heading", "caption": None},
    {"separator": "另一项试验使用独立样品。", "caption": None},
    {"lead": "某物料在规定条件下易溶。", "caption": None},
    {"lead": "仅参照上表，不适用于本次试验。", "caption": None},
    {"lead": "说明" * 130 + "结果如下表。", "caption": None},
])
def test_missing_ambiguous_or_out_of_scope_tables_keep_original_paragraphs(tmp_path, options):
    index, _ = document_index(tmp_path, **options)
    assert all(index.by_id[rid].kind == "table_row" for group in groups(index) for rid in group)
    if options.get("nested") or options.get("table") is False or options.get("rows") == 0:
        assert groups(index) == []
    assert any(record.kind == "paragraph" for record in index.records)


def test_intervening_prose_is_not_silently_absorbed_through_a_caption(tmp_path):
    index, _ = document_index(tmp_path, separator="样品乙在20℃下未检出残留。")
    joint, = groups(index)
    assert [index.by_id[rid].text for rid in joint if index.by_id[rid].kind == "paragraph"] == [
        CAPTION,
    ]
    assert {LEAD, "样品乙在20℃下未检出残留。"} <= {
        record.text for record in index.records if record.record_id not in joint
    }


@pytest.mark.parametrize("options", [{"caption": None}, {"lead": None}, {"separator": ""}])
def test_explicit_intro_or_numbered_caption_can_independently_bind_the_next_table(
    tmp_path, options,
):
    index, _ = document_index(tmp_path, **options)
    assert len(groups(index)) == 1


def test_extra_names_negation_and_conditions_in_intro_are_not_discarded(tmp_path):
    lead = "样品甲未检出残留；样品乙仅在20℃下有效，结果如下表。"
    index, _ = document_index(tmp_path, lead=lead)
    assert lead in {f.text for f in context(index, groups(index)[0]).fragments if f.fact_eligible}


@pytest.mark.parametrize("limits", [{"max_records": 1}, {"max_group_chars": 10}])
def test_capacity_preserves_every_original_record_for_independent_work(tmp_path, limits):
    index, _ = document_index(tmp_path)
    before = [record.record_id for record in index.records]
    assert groups(index, **limits) == []
    assert [record.record_id for record in index.records] == before
    assert len(before) == 4


@pytest.mark.parametrize("max_group_chars,expected_requests", [(1200, 1), (1, 4)])
def test_real_coordinator_avoids_intro_only_calls_and_cold_resume_keeps_paid_results(
    tmp_path, monkeypatch, current_run, max_group_chars, expected_requests,
):
    original_arguments = fixture.arguments

    def arguments(path, lines):
        values = original_arguments(path, lines)
        index, analysis = document_index(
            path, lead="装置的试验结果如下表。", caption="表3装置信息",
        )
        values.update(ir=index.ir, metadata=prepare_metadata(
            index.ir, section_tree=analysis.structure.section_tree.to_dict(),
            summary_version="table-reading-test",
        ))
        return values

    monkeypatch.setattr(fixture, "arguments", arguments)
    args, factory, requests, hooks = fixture.record_setup(
        tmp_path, monkeypatch, current_run, split=True, empty_relations=True,
        transform=lambda view, payload, **options: {name: [] for name in payload},
    )
    adapter = factory().adapter
    adapter.record_discovery = adapter.record_discovery.model_copy(update={
        "contextual": ContextualDiscoveryPolicy(max_group_chars=max_group_chars),
    })
    stopped = False
    persist = hooks["model_call_hook"]

    def calls(state):
        nonlocal stopped
        persist(state)
        if any(value["field"] == "model_turn"
               for value in state.get("result_changes", {}).values()):
            stopped = True

    factory(progress_hook=lambda _: not stopped).run(
        **args, **{**hooks, "model_call_hook": calls},
    )
    store, run, _ = current_run
    restore = dict(
        resume_state=vars(current_state.restore_work(store, run, run.run_fingerprint)),
        model_call_state=current_state.restore_calls(store, run, run.run_fingerprint),
    )
    first_count = len(requests)
    first_request = requests[0]
    result = factory().run(**args, **hooks, **restore)
    discovery = [view for view in requests if "members" not in view]
    assert len(discovery) == expected_requests, result.diagnostics
    assert requests.count(first_request) == 1 and first_count == 1
    if max_group_chars > 1:
        assert all(any(unit["fact_eligible"] and unit["text"] in {"1234-3", "1234-2"}
                       for unit in view["evidence_units"]) for view in discovery)
    assert len({view["record_id"] for view in discovery}) == expected_requests
    assert {unit["text"] for view in discovery for unit in view["evidence_units"]
            if unit["fact_eligible"]} == {
        "装置的试验结果如下表。", "表3装置信息", "1234-3", "1234-2", "易溶", "微溶",
    }
    rows = current_state.read_rows(store, run, DocumentRunCurrentState,
                                   prefix="work:record_discovery")["work:record_discovery"]
    entities = [row for row in rows.values()
                if row["value"]["task"].get("purpose") != "property_disambiguation"]
    fields = [row["value"] for row in rows.values()
              if row["value"]["task"].get("purpose") == "property_disambiguation"]
    assert len(entities) == expected_requests
    assert [row["attribute_field"]["label"] for row in fields] == ["名称", "名称"]
    count = len(requests)
    factory().run(
        **args, **hooks,
        resume_state=vars(current_state.restore_work(store, run, run.run_fingerprint)),
        model_call_state=current_state.restore_calls(store, run, run.run_fingerprint),
    )
    assert len(requests) == count


def test_joint_table_reading_still_produces_two_distinct_row_entities(
    tmp_path, monkeypatch, current_run,
):
    original_arguments = fixture.arguments

    def arguments(path, lines):
        values = original_arguments(path, lines)
        index, analysis = document_index(
            path, lead="装置的试验结果如下表。", caption="表3装置信息",
            names=("装置甲", "装置丙"),
        )
        values.update(ir=index.ir, metadata=prepare_metadata(
            index.ir, section_tree=analysis.structure.section_tree.to_dict(),
            summary_version="table-entities-test",
        ))
        return values

    monkeypatch.setattr(fixture, "arguments", arguments)
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["装置甲。"],
    )
    result = factory().run(**args, **hooks)
    assert {node.label for node in result.graph.nodes if not node.root} == {"装置甲", "装置丙"}
    discovery = [view for view in requests if "members" not in view
                 and view["stage"] == "discovery" and not view.get("attribute_disambiguation")]
    assert len(discovery) == 1
    for view in discovery:
        names = {unit["text"] for unit in view["evidence_units"] if unit["fact_eligible"]}
        assert names & {"装置甲", "装置丙"} == {"装置甲", "装置丙"}


def test_default_batch_limit_and_first_batch_only_intro_facts(tmp_path):
    index, _ = document_index(tmp_path, rows=6)
    first, second = groups(index)
    assert [sum(index.by_id[rid].kind == "table_row" for rid in group)
            for group in (first, second)] == [4, 2]
    assert {LEAD, CAPTION} <= {
        f.text for f in context(index, first).fragments if f.fact_eligible
    }
    later = context(index, second)
    assert {LEAD, CAPTION} <= {f.text for f in later.fragments if not f.fact_eligible}
    assert not {LEAD, CAPTION} & {f.text for f in later.fragments if f.fact_eligible}
    assert len(set(first) & set(second)) == 0


def test_record_and_character_limits_split_without_dropping_original_rows(tmp_path):
    index, _ = document_index(
        tmp_path, rows=6, lead=None, caption=None, names=[f"物料{row:04d}" for row in range(6)],
    )
    rows = [record for record in index.records if record.kind == "table_row"]
    assert [len(group) for group in groups(index, max_records=2)] == [2, 2, 2]
    limit = len(rows[0].text) + len(rows[1].text)
    batches = groups(index, max_group_chars=limit)
    assert all(sum(len(index.by_id[rid].text) for rid in batch) <= limit for batch in batches)
    assert {rid for batch in batches for rid in batch} == {row.record_id for row in rows}
    assert [record.record_id for record in index.records] == [row.record_id for row in rows]


@pytest.mark.parametrize("boundary", ["row_gap", "different_header", "different_table"])
def test_row_gap_header_or_table_boundary_prevents_joint_reading(tmp_path, boundary):
    index, _ = document_index(tmp_path, lead=None, caption=None)
    first, second = index.records
    if boundary == "row_gap":
        second = replace(second, row_index=second.row_index + 1)
    elif boundary == "different_header":
        second = replace(second, header_units=second.header_units[:1])
    else:
        second = replace(second, table_path=("different-table",))
    index.records = [first, second]
    index.by_id[second.record_id] = second
    assert groups(index) == []


def test_frozen_policy_enables_bounded_table_rows():
    policy = RecordDiscoveryPolicy(contextual=ContextualDiscoveryPolicy())
    assert policy.table_reading == "bounded-table-rows-v2"
    assert policy.model_dump(mode="json")["contextual"]["max_table_rows_per_group"] == 4
    with pytest.raises(ValueError):
        ContextualDiscoveryPolicy(max_table_rows_per_group=5)


def test_vertically_merged_sources_keep_independent_row_contexts(tmp_path):
    from app.services.extraction.ontology_guided.context import fragment_record_id

    document = Document()
    document.add_heading("装置", 1)
    table = document.add_table(rows=5, cols=2)
    for row, values in enumerate([
        ("名称", "型号"), ("装置甲", "A"), ("装置乙", "B"),
        ("装置乙", "C"), ("装置丁", "D"),
    ]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    table.cell(2, 0).merge(table.cell(3, 0))
    path = tmp_path / "merged-table.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    assert groups(index) == []
    for record in index.records:
        view = context(index, (record.record_id,))
        assert all(fragment_record_id(fragment, index, view.record_id) == record.record_id
                   for fragment in view.fragments if fragment.fact_eligible)
