"""Generic regressions for source permissions, missing values and identity roles."""

import json
from copy import deepcopy

import pytest
from docx import Document
from pydantic import ValidationError

from app.services.extraction.ontology_guided.claim_freeze import freeze_record_proposal
from app.services.extraction.ontology_guided.claim_protocol import (
    DiscoveryEnvelope,
    ExtractionProfile,
    IdentityKeySpec,
    Quote,
    build_verification_input,
    compile_stage_schema,
)
from app.services.extraction.ontology_guided.context import assemble_record_discovery_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
    VersionedRef,
)
from app.services.extraction.ontology_guided.model_context_projection import compact_model_context
from app.services.extraction.ontology_guided.model_schema_projection import compact_answer_schema
from app.services.extraction.ontology_guided.record_discovery import (
    RecordDiscoveryTask,
    compile_record_schema_card,
    merge_record_schema_cards,
)
from app.services.extraction.ontology_guided.record_model_adapter import (
    RECORD_TYPE_EVIDENCE_INSTRUCTIONS,
)
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction import test_record_model_adapter as record_fixture
from tests.test_extraction.test_record_model_adapter import source as source  # noqa: F401


def fixture(tmp_path, value="N/A", *, standalone_value=False):
    document = Document()
    document.add_heading("装置资料", 1)
    document.add_paragraph("名称：样件甲")
    document.add_paragraph(value if standalone_value else f"登记号：{value}")
    document.add_paragraph("说明：N/A接口适配器")
    path = tmp_path / "source.docx"
    document.save(path)
    index = RecordIndex(analyze_word_core(path).ir)
    records = [row for row in index.records
               if "：" in row.text or standalone_value and row.text == value]
    properties = [SlotSpec(
        iri=iri, label=label, datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
        identity_key=identity, description=description,
    ) for iri, label, identity, description in (
        ("urn:name", "名称", False, "对象名称，不声明唯一身份。"),
        ("urn:serial", "登记号", True, "原文明示的登记标识；名称不代替唯一实例身份。"),
        ("urn:description", "说明", False, "原文说明。"),
    )]
    ontology = OntologySnapshot(
        snapshot_id="snapshot", ontology_hash="frozen", classes={
            "urn:assembly": OntologyClassDefinition(
                iri="urn:assembly", label="装置", source_hash="frozen",
                description="具备独立安装功能的完整装置；单独零部件不属于该类型。",
                parent_iris=["urn:physical"], declared_properties=properties,
            ),
        },
    )
    card = compile_record_schema_card(
        ontology, class_iris=["urn:assembly"], analysis_scope_ref="scope",
        profile=ExtractionProfile(identity_keys=[IdentityKeySpec(
            class_iri="urn:assembly", property_iris=["urn:serial"], namespace=None,
            scope="document", declaration_ref="frozen-identity",
        )]),
    )
    task = RecordDiscoveryTask.create(
        run_fingerprint="run", record_id=records[0].record_id,
        source_record_ids=[row.record_id for row in records], schema_card_id=card.schema_card_id,
        analysis_scope_ref=card.analysis_scope_ref, dependency_hash="source",
    )
    context = assemble_record_discovery_context(
        task, card, index, ontology_hash=ontology.ontology_hash,
        document_context=DocumentContext(
            document_hash=index.ir.document_hash, document_class_iri="urn:Document",
            root_ref=VersionedRef(id="root", revision=1),
        ),
    )

    def quote(text, row=0):
        return {"evidence_id": records[row].source_units[0].evidence_id,
                "text": text, "context_text": None}

    proposal = {
        "entities": [{"local_id": "entity", "class_iri": "urn:assembly",
                      "representation": "mention", "mentions": [quote("样件甲")],
                      "record_components": [], "identifier_claims": []}],
        "properties": [{
            "local_id": "property", "subject_id": "entity", "predicate_iri": "urn:serial",
            "value_quote": quote(value, 1), "field_support": [quote("登记号", 1)],
            "unit_support": [], "qualifiers": {"polarity": "affirmed", "modality": "asserted",
                                                "condition_support": [], "scope_qualifiers": []},
            "bridge_kind": "explicit_assertion", "bridge_ref_ids": [],
        }], "relations": [], "external_links": [], "observations": [],
    }
    options = dict(task=task, context=context, card=card, index=index, generation=1,
                   reference_resolution=True)
    return proposal, options, quote


def freeze(proposal, options):
    return freeze_record_proposal(DiscoveryEnvelope.model_validate(proposal), **options)


def targets(frozen, options):
    return build_verification_input(
        frozen, discovery_ref="saved", context=options["context"], card=options["card"],
        scope=options["task"].scope, entity_dependencies=[], external_candidates=[],
        bridge_dependencies=[], scope_resolutions=[],
    ).targets


@pytest.mark.parametrize("value,kind", [
    ("N/A", "missing"), ("not available", "missing"), ("未提供", "missing"),
    ("未知", "unknown"),
])
def test_complete_missing_field_is_not_a_fact_and_keeps_source_observation(tmp_path, value, kind):
    proposal, options, _ = fixture(tmp_path, value)
    original = deepcopy(proposal)
    frozen = freeze(proposal, options)
    assert frozen.claim_issues["property"] == [f"attribute_value_{kind}"]
    assert [(row.kind, row.quote.text) for row in frozen.observations] == [(kind, value)]
    assert [row.payload.local_id for row in targets(frozen, options)] == ["entity"]
    assert proposal == original  # Source proposal and rejected candidate remain inspectable.
    assert frozen.properties[0].value_quote.text == value


@pytest.mark.parametrize("value", ["0", "false", "否", "无", "NA", "N/A-007"])
def test_real_values_and_identifier_spellings_are_not_missing(tmp_path, value):
    proposal, options, _ = fixture(tmp_path, value)
    frozen = freeze(proposal, options)
    assert frozen.claim_issues == {} and frozen.observations == []
    assert len(targets(frozen, options)) == 2


def test_marker_substring_in_an_ordinary_field_is_not_a_missing_value(tmp_path):
    proposal, options, quote = fixture(tmp_path)
    prop = proposal["properties"][0]
    prop.update(predicate_iri="urn:description", value_quote=quote("N/A", 2),
                field_support=[quote("说明", 2)])
    frozen = freeze(proposal, options)
    assert frozen.observations == []
    assert "attribute_value_missing" not in frozen.claim_issues.get("property", [])
    # This does not accept the truncated value: independent field/value verification remains.


def test_missing_identifier_cannot_create_identity_or_its_dependent_fact(tmp_path):
    proposal, options, quote = fixture(tmp_path)
    proposal["entities"][0]["identifier_claims"] = [{
        "predicate_iri": "urn:serial", "value_quote": quote("N/A", 1),
    }]
    frozen = freeze(proposal, options)
    assert "attribute_value_missing" in frozen.claim_issues["entity"]
    assert targets(frozen, options) == []


def test_name_field_cannot_be_relabelled_as_unique_identifier(tmp_path):
    proposal, options, quote = fixture(tmp_path, "S-001")
    prop = proposal["properties"][0]
    prop.update(value_quote=quote("样件甲"), field_support=[quote("名称")])
    correct_name = {**deepcopy(prop), "local_id": "name", "predicate_iri": "urn:name"}
    proposal["properties"].append(correct_name)
    frozen = freeze(proposal, options)
    assert frozen.claim_issues == {"property": ["identity_field_role_mismatch"]}
    assert {row.payload.local_id for row in targets(frozen, options)} == {"entity", "name"}


def test_entity_identifier_does_not_bypass_source_field_role(tmp_path):
    proposal, options, quote = fixture(tmp_path, "S-001")
    proposal["entities"][0]["identifier_claims"] = [{
        "predicate_iri": "urn:serial", "value_quote": quote("样件甲"),
    }]
    frozen = freeze(proposal, options)
    assert frozen.claim_issues["entity"] == ["identity_field_role_mismatch"]
    assert targets(frozen, options) == []


def test_missing_marker_outside_fact_permission_does_not_grant_missing_observation(tmp_path):
    proposal, options, quote = fixture(tmp_path)
    evidence_id = quote("N/A", 1)["evidence_id"]
    for fragment in options["context"].fragments:
        if fragment.anchor.evidence_id == evidence_id:
            fragment.fact_eligible = False
    frozen = freeze(proposal, options)
    assert "fact_source_outside_scope" in frozen.claim_issues["property"]
    assert frozen.observations == []


def entity_value_fixture(tmp_path, value, representation, *, standalone=False, excerpt=None):
    proposal, options, quote = fixture(tmp_path, value, standalone_value=standalone)
    source = quote(value if excerpt is None else excerpt, 1)
    entity = proposal["entities"][0]
    if representation == "mention":
        entity["mentions"] = [source]
    else:
        entity.update(representation="record", mentions=[], record_components=[
            {"role": representation, "quote": source},
        ])
        if representation == "value":
            entity["record_components"].insert(0, {"role": "field", "quote": quote("名称")})
    # A valid dependent field must be blocked by its subject, not its own value.
    proposal["properties"][0].update(
        predicate_iri="urn:name", value_quote=quote("样件甲"), field_support=[quote("名称")],
    )
    return proposal, options, quote


@pytest.mark.parametrize("representation", ["mention", "subject", "value"])
@pytest.mark.parametrize("value,kind", [("N/A", "missing"), ("未知", "unknown")])
@pytest.mark.parametrize("standalone", [False, True])
def test_missing_entity_referent_blocks_its_dependent_property_without_rewriting_source(
    tmp_path, representation, value, kind, standalone,
):
    proposal, options, quote = entity_value_fixture(
        tmp_path, value, representation, standalone=standalone,
    )
    original = deepcopy(proposal)
    frozen = freeze(proposal, options)

    assert frozen.claim_issues == {
        "entity": [f"entity_referent_{kind}"], "property": ["entity_dependency_invalid"],
    }
    assert targets(frozen, options) == []
    assert len(frozen.observations) == 1
    observation = frozen.observations[0]
    assert observation.kind == kind
    assert observation.subject_id is None and observation.predicate_iri is None
    assert observation.quote.model_dump(mode="json") == quote(value, 1)
    assert frozen.entities[0].model_dump(mode="json") == original["entities"][0]
    assert frozen.properties[0].model_dump(mode="json") == original["properties"][0]
    assert proposal == original


def test_missing_value_in_record_with_valid_subject_does_not_shorten_composition(tmp_path):
    proposal, options, quote = entity_value_fixture(tmp_path, "N/A", "value")
    components = proposal["entities"][0]["record_components"]
    components.insert(0, {"role": "subject", "quote": quote("样件甲")})
    frozen = freeze(proposal, options)
    assert frozen.claim_issues == {
        "entity": ["entity_referent_missing"], "property": ["entity_dependency_invalid"],
    }
    assert targets(frozen, options) == []
    assert [part.model_dump(mode="json") for part in frozen.entities[0].record_components] == (
        components
    )
    assert frozen.observations[0].subject_id is None
    assert frozen.observations[0].predicate_iri is None


@pytest.mark.parametrize("representation", ["mention", "subject", "value"])
@pytest.mark.parametrize("standalone", [False, True])
def test_context_only_missing_entity_quote_does_not_create_observation(
    tmp_path, representation, standalone,
):
    proposal, options, quote = entity_value_fixture(
        tmp_path, "N/A", representation, standalone=standalone,
    )
    evidence_id = quote("N/A", 1)["evidence_id"]
    for fragment in options["context"].fragments:
        if fragment.anchor.evidence_id == evidence_id:
            fragment.fact_eligible = False
    frozen = freeze(proposal, options)
    assert "fact_source_outside_scope" in frozen.claim_issues["entity"]
    assert "entity_referent_missing" not in frozen.claim_issues["entity"]
    assert frozen.claim_issues["property"] == ["entity_dependency_invalid"]
    assert frozen.observations == [] and targets(frozen, options) == []


@pytest.mark.parametrize("representation", ["mention", "subject", "value"])
def test_missing_entity_quote_absent_from_source_does_not_create_observation(
    tmp_path, representation,
):
    proposal, options, _ = entity_value_fixture(tmp_path, "S-001", representation, excerpt="N/A")
    frozen = freeze(proposal, options)
    assert "source_excerpt_mismatch" in frozen.claim_issues["entity"]
    assert "entity_referent_missing" not in frozen.claim_issues["entity"]
    assert frozen.observations == [] and targets(frozen, options) == []


@pytest.mark.parametrize("representation", ["mention", "subject", "value"])
@pytest.mark.parametrize("value", ["NA", "0", "false", "否", "无", "N/A-007"])
def test_real_entity_source_values_are_not_missing(tmp_path, representation, value):
    proposal, options, _ = entity_value_fixture(tmp_path, value, representation)
    frozen = freeze(proposal, options)
    assert frozen.claim_issues == {} and frozen.observations == []
    assert {row.payload.local_id for row in targets(frozen, options)} == {"entity", "property"}


@pytest.mark.parametrize("representation", ["mention", "subject", "value"])
@pytest.mark.parametrize("standalone", [False, True])
def test_entity_marker_substring_is_left_to_semantic_verification(
    tmp_path, representation, standalone,
):
    proposal, options, _ = entity_value_fixture(
        tmp_path, "N/A接口适配器", representation, standalone=standalone, excerpt="N/A",
    )
    frozen = freeze(proposal, options)
    assert frozen.claim_issues == {} and frozen.observations == []
    assert len(targets(frozen, options)) == 2


@pytest.mark.parametrize("role", ["field", "context"])
def test_auxiliary_record_component_is_not_treated_as_missing_referent(tmp_path, role):
    proposal, options, quote = entity_value_fixture(tmp_path, "N/A", "subject")
    proposal["entities"][0]["record_components"] = [
        {"role": "subject", "quote": quote("样件甲")},
        {"role": role, "quote": quote("N/A", 1)},
    ]
    frozen = freeze(proposal, options)
    assert frozen.claim_issues == {} and frozen.observations == []
    assert len(targets(frozen, options)) == 2


@pytest.mark.parametrize("text", ["", " ", "\n\t", "\r\n", "\u00a0", "\u3000"])
def test_quotes_reject_empty_support_instead_of_locating_it_everywhere(text):
    with pytest.raises(ValidationError):
        Quote(evidence_id="source", text=text, context_text=None)
    with pytest.raises(ValidationError):
        Quote.model_validate_json(json.dumps({
            "evidence_id": "source", "text": text, "context_text": None,
        }))


@pytest.mark.parametrize("text", [
    "A", "中文引文", "第一行\n第二行", "first\r\nsecond", " \n中文引文\t ",
])
def test_quotes_preserve_valid_source_text_including_multiline_whitespace(text):
    payload = {"evidence_id": "source", "text": text, "context_text": None}
    assert Quote.model_validate(payload).text == text
    assert Quote.model_validate_json(json.dumps(payload)).text == text


def test_discovery_value_schema_authorizes_facts_but_keeps_auxiliary_field_context(tmp_path):
    _, options, _ = fixture(tmp_path)
    schema = compile_stage_schema(
        "discovery", card=options["card"], evidence_ids=["fact", "binding"],
        fact_evidence_ids=["fact"], targets=[],
    )
    definitions = schema["$defs"]
    assert definitions["FactQuote"]["properties"]["evidence_id"]["enum"] == ["fact"]
    assert definitions["Quote"]["properties"]["evidence_id"]["enum"] == ["binding", "fact"]
    assert definitions["PropertyProposal"]["properties"]["value_quote"] == {
        "$ref": "#/$defs/FactQuote",
    }
    for view in (schema, compact_answer_schema(schema)):
        for name in ("Quote", "FactQuote"):
            quote_text = view["$defs"][name]["properties"]["text"]
            assert quote_text["minLength"] == 1
            # llama.cpp rejects the former unanchored \\S pattern before generation.
            assert "pattern" not in quote_text
    factual, auxiliary = definitions["RecordComponent"]["anyOf"]
    assert factual["properties"]["role"]["enum"] == ["subject", "value"]
    assert factual["properties"]["quote"] == {"$ref": "#/$defs/FactQuote"}
    assert auxiliary["properties"]["role"]["enum"] == ["field", "context"]
    assert auxiliary["properties"]["quote"] == {"$ref": "#/$defs/Quote"}


def test_frozen_type_definitions_survive_compaction_and_affect_card_identity(tmp_path):
    _, options, _ = fixture(tmp_path)
    card = options["card"]
    compact = compact_model_context({"schema_card": card.model_dump(mode="json")})
    rendered = compact["schema_card"]["class_cards"][0]
    assert rendered["description"] == "具备独立安装功能的完整装置；单独零部件不属于该类型。"
    assert rendered["parent_iris"] == ["urn:physical"]
    assert rendered["properties"][1]["description"]
    changed = card.model_copy(deep=True)
    changed.class_cards[0].description = "属于可替换的单独零部件。"
    with pytest.raises(ValueError, match="record_schema_class_definition_conflict"):
        merge_record_schema_cards([card, changed])
    assert merge_record_schema_cards([card]).schema_card_id != (
        merge_record_schema_cards([changed]).schema_card_id
    )


@pytest.mark.parametrize("type_verdict", ["undetermined", "unsupported"])
def test_real_request_carries_type_definition_and_rejected_type_blocks_dependent_property(
    source, monkeypatch, type_verdict,
):
    original_compile = record_fixture.compile_record_schema_card

    def with_definitions(ontology, **kwargs):
        ontology.classes["urn:T"].description = "具备独立安装功能的完整装置。"
        ontology.classes["urn:T"].parent_iris = ["urn:assembly"]
        ontology.classes["urn:Other"].description = "可替换的单独零部件。"
        return original_compile(ontology, **kwargs)

    monkeypatch.setattr(record_fixture, "compile_record_schema_card", with_definitions)
    seen = []

    def verify_type(answer, view, _ordinal):
        cards = {row["class_iri"]: row for row in view["schema_card"]["class_cards"]}
        assert cards["urn:T"]["description"] == "具备独立安装功能的完整装置。"
        assert cards["urn:T"]["parent_iris"] == ["urn:assembly"]
        assert view["fact_evidence_ids"]
        seen.append(view["stage"])
        if view["stage"] == "verification":
            target = next(row for row in view["verification_input"]["targets"]
                          if row["payload"].get("local_id") == "b")
            response = next(row for row in answer["verifications"]
                            if row["target_id"] == target["target_id"])
            facet = next(row for row in response["facets"] if row["name"] == "type")
            facet.update(verdict=type_verdict, reason="原文提及对象，但未证明独立完整装置层级。")

    adapter, task, context, card, _storage, requests, _ = record_fixture.setup_record(
        source, monkeypatch, transform=verify_type,
    )
    outcome = adapter.inspect_record(task, context, card)
    assert seen == ["discovery", "verification"]
    for request in requests:
        assert RECORD_TYPE_EVIDENCE_INSTRUCTIONS in request["instructions"]
    assert "referent或属性facet成立不替代type成立" in requests[1]["instructions"]
    assert {node.class_iri for node in outcome.nodes} == {"urn:Other"}
    assert len(outcome.properties) == 1 and not outcome.complete
