"""Root-schema routing narrows work without granting metadata fact authority."""

from copy import deepcopy

import pytest
from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided.claim_protocol import ExtractionProfile
from app.services.extraction.ontology_guided.contracts import (
    EdgeSpec,
    MetadataNode,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
)
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.record_discovery import (
    RecordDiscoveryPolicy,
    compile_record_schema_card,
    merge_record_schema_cards,
)
from app.services.extraction.ontology_guided.record_search import RecordSearch
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.schema_region_routing import (
    SchemaRegionRoutingPolicy,
    compile_schema_region_routing,
    metadata_region_is_bounded,
    resolve_metadata_region,
    routing_card_text,
    routing_heading_match,
)
from app.services.extraction.ontology_guided.semantic_reranker import (
    RankingPolicy,
    RankingService,
)
from app.services.extraction.word_analysis import analyze_word_core
from tests.test_extraction.test_contextual_record_executor import setup_contextual
from tests.test_extraction.test_record_executor import A as RECORD_DEVICE

pytest_plugins = ["tests.test_extraction.test_tool_engine_resume"]

ROOT = "urn:Report"
PRODUCT = "urn:Product"
DEVICE = "urn:Device"
PART = "urn:Part"
DEGRADATION = "urn:DegradationPathway"
DESCRIBES = "urn:describes"
USES = "urn:uses"
HAS_PART = "urn:hasPart"
HAS_DEGRADATION = "urn:hasDegradationPathway"


def definition(iri, label, *, description="", properties=(), relationships=()):
    payload = {
        "iri": iri,
        "label": label,
        "description": description,
        "parent_iris": [],
        "declared_properties": list(properties),
        "declared_relationships": list(relationships),
    }
    return OntologyClassDefinition(**payload, source_hash=evidence_hash(payload))


def ontology(*, product_degradation=False):
    root_id = SlotSpec(
        iri="urn:documentNumber", label="文件编号", declared_by=[ROOT],
        datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
    )
    describes = EdgeSpec(
        iri=DESCRIBES, label="描述产品", description="报告描述的产品",
        declared_by=[ROOT], range_class_iris=[PRODUCT],
    )
    uses = EdgeSpec(
        iri=USES, label="使用设备", declared_by=[ROOT], range_class_iris=[DEVICE],
    )
    has_part = EdgeSpec(
        iri=HAS_PART, label="包含部件", declared_by=[DEVICE], range_class_iris=[PART],
    )
    has_degradation = EdgeSpec(
        iri=HAS_DEGRADATION, label="含降解途径", description="产品的降解途径",
        declared_by=[PRODUCT], range_class_iris=[DEGRADATION],
    )
    classes = {
        ROOT: definition(
            ROOT, "报告", properties=[root_id],
            relationships=[describes, uses, *([has_degradation] if product_degradation else [])],
        ),
        PRODUCT: definition(
            PRODUCT, "产品", description="报告所述产品的基本性质",
            relationships=[has_degradation] if product_degradation else [],
        ),
        DEVICE: definition(DEVICE, "设备", relationships=[has_part]),
        PART: definition(PART, "部件"),
        **({DEGRADATION: definition(DEGRADATION, "降解途径")}
           if product_degradation else {}),
    }
    return OntologySnapshot(
        snapshot_id="routing-ontology",
        ontology_hash=evidence_hash({key: value.model_dump(mode="json")
                                     for key, value in classes.items()}),
        classes=classes,
        created_from="frozen_fixture",
    )


def document_index(tmp_path):
    document = Document()
    document.add_heading("研究资料", 1)
    document.add_heading("设备信息", 2)
    document.add_paragraph("设备甲用于生产。")
    document.add_heading("合成路线图", 2)
    document.add_heading("工艺描述", 2)
    document.add_paragraph("产品甲经两步反应制得。")
    path = tmp_path / "routing.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="routing-test",
    )
    return index, metadata


def product_document_index(tmp_path):
    document = Document()
    document.add_heading("研究资料", 1)
    document.add_heading("简介", 2)
    document.add_paragraph("产品甲用于临床研究。")
    document.add_heading("产品的基本性质", 2)
    document.add_paragraph("产品结构：")
    path = tmp_path / "product-routing.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="routing-test",
    )
    return index, metadata


def test_root_relationships_create_independent_compact_branches_and_property_card():
    policy = SchemaRegionRoutingPolicy.model_validate({
        "version": "schema-region-routing-v1",
        "max_regions_per_card": 2,
        "minimum_similarity": 0.25,
    })
    assert (policy.max_group_chars, policy.max_group_records) == (1200, 8)
    plan = compile_schema_region_routing(
        ontology(), ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    assert [(item.iri, item.label) for item in plan.root_property_card.properties] == [
        ("urn:documentNumber", "文件编号"),
    ]
    cards = {card.predicate_iri: card for card in plan.routing_cards}
    assert cards[DESCRIBES].class_iris == [PRODUCT]
    assert cards[USES].class_iris == [DEVICE, PART]
    assert cards[USES].range_class_iris == [DEVICE]
    text = routing_card_text(cards[USES])
    assert {"使用设备", "设备"} <= set(text.splitlines())
    assert "部件" not in text.splitlines()
    assert "urn:uses" not in text


def test_root_route_embedding_excludes_descendant_branch_vocabulary():
    plan = compile_schema_region_routing(
        ontology(product_degradation=True), ROOT,
        analysis_scope_ref="scope", max_hops=2,
    )
    card = next(item for item in plan.routing_cards if item.predicate_iri == DESCRIBES)
    degradation = next(
        item for item in plan.routing_cards if item.predicate_iri == HAS_DEGRADATION
    )

    assert DEGRADATION in card.class_iris
    assert "降解途径" in card.class_labels
    assert "降解途径" not in routing_card_text(card)
    assert "产品" in routing_card_text(card)
    assert "报告所述产品的基本性质" in routing_card_text(card)
    assert "降解途径" in routing_card_text(degradation)


def test_explicit_ontology_term_in_heading_has_priority_over_semantic_noise(tmp_path):
    index, metadata = document_index(tmp_path)
    plan = compile_schema_region_routing(
        ontology(), ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    card = next(item for item in plan.routing_cards if item.predicate_iri == USES)
    device = next(item for item in metadata.node_summaries if item.heading == "设备信息")
    process = next(item for item in metadata.node_summaries if item.heading == "工艺描述")
    assert routing_heading_match(card, device) > routing_heading_match(card, process)


def test_generic_relation_label_does_not_match_a_longer_process_heading():
    plan = compile_schema_region_routing(
        ontology(), ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    card = next(item for item in plan.routing_cards if item.predicate_iri == DESCRIBES)
    card = card.model_copy(update={"label": "描述"})
    process = MetadataNode(
        node_id="process-description", heading="工艺描述", path=["报告", "工艺描述"],
        summary=None, summary_status="pending", summary_source="none",
        source_record_refs=["record"],
    )
    product = process.model_copy(update={
        "node_id": "product", "heading": "产品信息", "path": ["报告", "产品信息"],
    })
    assert routing_heading_match(card, process) == 0
    assert routing_heading_match(card, product) == 1


def test_qualified_range_matches_only_bounded_unqualified_heading_forms():
    plan = compile_schema_region_routing(
        ontology(), ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    card = next(item for item in plan.routing_cards if item.predicate_iri == DESCRIBES)
    card = card.model_copy(update={
        "predicate_iri": "urn:anyRootRelation",
        "range_type_labels": ["药物产品"],
    })
    product = MetadataNode(
        node_id="product", heading="产品的基本性质", path=["报告", "产品的基本性质"],
        summary=None, summary_status="pending", summary_source="none",
        source_record_refs=["record"],
    )
    assessment = product.model_copy(update={
        "node_id": "assessment", "heading": "生产风险评估",
        "path": ["报告", "生产风险评估"],
    })
    toxicity = product.model_copy(update={
        "node_id": "toxicity", "heading": "产品毒性信息",
        "path": ["报告", "产品毒性信息"],
    })

    assert routing_heading_match(card, product) == 1
    assert routing_heading_match(card, assessment) == 0
    assert routing_heading_match(card, toxicity) == 0


def test_descendant_label_does_not_hijack_unrelated_root_branch():
    plan = compile_schema_region_routing(
        ontology(), ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    card = next(item for item in plan.routing_cards if item.predicate_iri == USES)
    card = card.model_copy(update={"class_labels": [*card.class_labels, "存放条件"]})
    node = MetadataNode(
        node_id="storage", heading="存放条件", path=["报告", "存放条件"],
        summary=None, summary_status="pending", summary_source="none",
        source_record_refs=["record"],
    )
    assert routing_heading_match(card, node) == 0


def test_focus_path_prunes_other_root_branches_and_nested_relations():
    plan = compile_schema_region_routing(
        ontology(), ROOT, analysis_scope_ref="scope", max_hops=2,
        focus_paths=[(USES, HAS_PART)],
    )
    card, = plan.routing_cards
    assert card.predicate_iri == USES
    assert card.class_iris == [DEVICE, PART]


def test_empty_structure_node_resolves_nearest_same_parent_fact_region(tmp_path):
    index, metadata = document_index(tmp_path)
    node = next(item for item in metadata.node_summaries if item.heading == "合成路线图")
    resolved = resolve_metadata_region(node.node_id, metadata, index)
    assert resolved is not None and resolved[1] == "adjacent"
    assert [index.by_id[identity].text for identity in resolved[0]] == [
        "产品甲经两步反应制得。",
    ]


def test_document_summary_parent_is_not_an_execution_region(tmp_path):
    index, metadata = document_index(tmp_path)
    root = next(item for item in metadata.node_summaries if len(item.path) == 1)
    resolved = resolve_metadata_region(root.node_id, metadata, index)
    assert resolved is not None and resolved[1] == "descendant"
    assert not metadata_region_is_bounded(root, resolved, max_records=32)

    direct = next(item for item in metadata.node_summaries if item.source_record_refs)
    resolved = resolve_metadata_region(direct.node_id, metadata, index)
    assert resolved is not None and resolved[1] == "direct"
    assert metadata_region_is_bounded(direct, resolved, max_records=32)


class Embeddings:
    identity = {"model": "schema-routing-test"}

    def count_tokens_batch(self, texts):
        return [len(text) for text in texts]

    def count_tokens(self, text):
        return len(text)

    def embed(self, texts):
        return [[float("设备" in text), float("设备" not in text)] for text in texts]

    def score_pairs(self, pairs):
        return [1.0] * len(pairs)


def test_explicit_region_reserves_next_slot_for_preceding_sibling_context(tmp_path):
    index, metadata = product_document_index(tmp_path)
    snapshot = ontology()
    product = compile_record_schema_card(
        snapshot, class_iris=[PRODUCT], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    plan = compile_schema_region_routing(
        snapshot, ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    route = next(item for item in plan.routing_cards if item.predicate_iri == DESCRIBES)
    route = route.model_copy(update={
        "predicate_iri": "urn:anyRootRelation",
        "range_type_labels": ["药物产品"],
        "execution_card_ids": [product.schema_card_id],
    })
    search = RecordSearch(
        [([record.record_id], [], None) for record in index.records],
        cards={product.schema_card_id: product},
        ordinary_cards=[product.schema_card_id],
        ontology=snapshot,
        index=index,
        policy=RecordDiscoveryPolicy(
            minimum_similarity=0.25,
            schema_region_routing=SchemaRegionRoutingPolicy(max_regions_per_card=2),
        ),
        routing_cards=[route],
        metadata=metadata,
    )
    search.prepare(
        RankingService(
            RankingPolicy(mode="semantic", enable_reranker=False), Embeddings(),
        ),
        lambda: None,
    )

    selections = next(
        row["route_selections"][route.routing_card_id]
        for row in search.rows.values()
        if row.get("route_selections", {}).get(route.routing_card_id)
    )
    headings = {
        node.node_id: node.heading for node in metadata.node_summaries
    }
    assert [headings[item["node_id"]] for item in selections] == [
        "产品的基本性质",
        "简介",
    ]


def test_metadata_route_limits_execution_cards_and_keeps_other_region_unattempted(tmp_path):
    index, metadata = document_index(tmp_path)
    snapshot = ontology()
    device = compile_record_schema_card(
        snapshot, class_iris=[DEVICE, PART], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    product = compile_record_schema_card(
        snapshot, class_iris=[PRODUCT], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    plan = compile_schema_region_routing(
        snapshot, ROOT, analysis_scope_ref="scope", max_hops=2,
        focus_paths=[(USES, HAS_PART)],
    )
    route = plan.routing_cards[0].model_copy(update={
        "execution_card_ids": [device.schema_card_id],
    })
    cards = {item.schema_card_id: item for item in (device, product)}
    slots = [([record.record_id], [], None) for record in index.records]
    search = RecordSearch(
        slots,
        cards=cards,
        ordinary_cards=list(cards),
        ontology=snapshot,
        index=index,
        policy=RecordDiscoveryPolicy(
            candidate_cards_per_record=2,
            minimum_similarity=0.5,
            schema_region_routing=SchemaRegionRoutingPolicy(
                max_regions_per_card=1, minimum_similarity=0.5,
            ),
        ),
        routing_cards=[route],
        metadata=metadata,
    )
    ranking = RankingService(
        RankingPolicy(mode="semantic", enable_reranker=False), Embeddings(),
    )
    search.prepare(ranking, lambda: None)

    rows = list(search.rows.values())
    routed = [row for row in rows if row.get("routing_status") == "routed"]
    unrouted = [row for row in rows if row.get("routing_status") == "unrouted"]
    assert len(routed) == 1 and len(unrouted) == 1
    assert routed[0]["remaining"] == [{
        "card_id": device.schema_card_id,
        "score": 1.0,
    }]
    assert unrouted[0]["remaining"] == []
    unrouted_record = unrouted[0]["source_record_ids"][0]
    assert search.admit_dependency(unrouted_record, {DEVICE}, set())
    current_unrouted = next(
        row for row in search.rows.values()
        if unrouted_record in row["source_record_ids"]
    )
    assert current_unrouted["remaining"] == []
    assert not search.admit_dependency(unrouted_record, {PRODUCT}, set())
    assert current_unrouted["remaining"] == []
    selected_card_id = search.take()[1]
    assert search.cards[selected_card_id].class_iris == [DEVICE]
    assert search.take() is None
    assert search.coverage_counts([{"status": "examined"}]) == {
        "planned": 2,
        "examined": 1,
        "incomplete": 0,
        "unattempted": 1,
    }
    diagnostics = search.diagnostics("semantic")
    assert diagnostics["unrouted_groups"] == 1
    assert diagnostics["unselected_groups"] == 0


def test_routed_dependency_prefers_identity_evidence_over_exact_first_group(tmp_path):
    index, metadata = document_index(tmp_path)
    snapshot = ontology()
    product = compile_record_schema_card(
        snapshot, class_iris=[PRODUCT], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    plan = compile_schema_region_routing(
        snapshot, ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    route = next(item for item in plan.routing_cards if item.predicate_iri == DESCRIBES)
    route = route.model_copy(update={"execution_card_ids": [product.schema_card_id]})
    records = [record.record_id for record in index.records]
    search = RecordSearch(
        [([record_id], [], None) for record_id in records],
        cards={product.schema_card_id: product},
        ordinary_cards=[product.schema_card_id],
        ontology=snapshot,
        index=index,
        policy=RecordDiscoveryPolicy(
            schema_region_routing=SchemaRegionRoutingPolicy(max_regions_per_card=2),
        ),
        routing_cards=[route],
        metadata=metadata,
    )
    rows = list(search.rows.items())
    low_key, low = rows[0]
    high_key, high = rows[1]
    search.rows[low_key] = {
        **low,
        "ranked": True,
        "routing_status": "routed",
        "route_card_ids": [route.routing_card_id],
        "remaining": [{"card_id": product.schema_card_id, "score": 0.9}],
    }
    search.rows[high_key] = {
        **high,
        "ranked": True,
        "routing_status": "routed",
        "route_card_ids": [route.routing_card_id],
        "remaining": [{"card_id": product.schema_card_id, "score": 0.1}],
    }

    # The relation points at the high-score first group, but its factual source
    # contains no product identity. The lower semantic score has explicit
    # product evidence and must run first.
    assert search.admit_dependency(records[0], {PRODUCT}, set())
    selected_sources, selected_card, _ = search.take()
    assert selected_sources == [records[1]]
    assert selected_card == product.schema_card_id
    assert not search.rows[low_key]["remaining"][0].get("priority")


def test_region_batch_merges_selected_execution_cards_and_cold_restore_rebuilds_it(tmp_path):
    index, metadata = document_index(tmp_path)
    snapshot = ontology()
    device = compile_record_schema_card(
        snapshot, class_iris=[DEVICE], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    serial = SlotSpec(
        iri="urn:serial", label="编号", declared_by=[DEVICE],
        datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
    )
    manufacturer = SlotSpec(
        iri="urn:manufacturer", label="制造商", declared_by=[DEVICE],
        datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
    )
    full_class = device.class_cards[0].model_copy(update={
        "properties": [manufacturer, serial],
    })
    device = device.model_copy(update={
        "schema_card_id": evidence_hash(full_class), "class_cards": [full_class],
    })
    field_class = full_class.model_copy(update={"properties": [serial]})
    field_card = device.model_copy(update={
        "schema_card_id": evidence_hash(field_class), "class_cards": [field_class],
    })
    part = compile_record_schema_card(
        snapshot, class_iris=[PART], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    plan = compile_schema_region_routing(
        snapshot, ROOT, analysis_scope_ref="scope", max_hops=2,
        focus_paths=[(USES, HAS_PART)],
    )
    route = plan.routing_cards[0].model_copy(update={
        "execution_card_ids": sorted([device.schema_card_id, part.schema_card_id]),
    })
    base_cards = {item.schema_card_id: item for item in (device, part, field_card)}
    base_ids = [*sorted([device.schema_card_id, part.schema_card_id]), field_card.schema_card_id]
    slots = [([record.record_id], [field_card.schema_card_id], None)
             for record in index.records]
    policy = RecordDiscoveryPolicy(
        candidate_cards_per_record=2,
        minimum_similarity=0.25,
        schema_region_routing=SchemaRegionRoutingPolicy(
            max_regions_per_card=1,
            minimum_similarity=0.25,
            execution_mode="region_batch",
            property_field_mode="region_batch",
        ),
    )
    search = RecordSearch(
        slots,
        cards=base_cards,
        ordinary_cards=[device.schema_card_id, part.schema_card_id],
        ontology=snapshot,
        index=index,
        policy=policy,
        routing_cards=[route],
        metadata=metadata,
    )
    uniform = Embeddings()
    uniform.embed = lambda texts: [[1.0] for _ in texts]
    search.prepare(
        RankingService(RankingPolicy(mode="semantic", enable_reranker=False), uniform),
        lambda: None,
    )
    row = next(value for value in search.rows.values() if value["remaining"])
    assert row["source_card_ids"] == base_ids
    assert len(row["remaining"]) == 1
    merged_id = row["execution_card_id"]
    assert search.cards[merged_id].class_iris == [DEVICE, PART]
    assert search.cards[merged_id].for_class(DEVICE).properties == [manufacturer, serial]
    saved = search.rows.drain()

    restored_cards = {item.schema_card_id: item for item in (device, part, field_card)}
    restored = RecordSearch(
        slots,
        cards=restored_cards,
        ordinary_cards=[device.schema_card_id, part.schema_card_id],
        ontology=snapshot,
        index=index,
        policy=policy,
        routing_cards=[route],
        metadata=metadata,
    )
    restored.restore(saved)
    assert restored.cards[merged_id].class_iris == [DEVICE, PART]
    assert restored.take()[1] == merged_id


def test_waiting_root_endpoint_splits_region_batch_to_exact_range_and_restores(tmp_path):
    index, metadata = document_index(tmp_path)
    snapshot = ontology()
    product = compile_record_schema_card(
        snapshot, class_iris=[PRODUCT], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    device = compile_record_schema_card(
        snapshot, class_iris=[DEVICE], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    plan = compile_schema_region_routing(
        snapshot, ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    card_ids = {
        DESCRIBES: [product.schema_card_id],
        USES: [device.schema_card_id],
    }
    routes = [card.model_copy(update={
        "execution_card_ids": card_ids[card.predicate_iri],
    }) for card in plan.routing_cards]
    cards = {
        product.schema_card_id: product,
        device.schema_card_id: device,
    }
    slots = [([index.records[0].record_id], [], None)]
    policy = RecordDiscoveryPolicy(
        candidate_cards_per_record=2,
        schema_region_routing=SchemaRegionRoutingPolicy(execution_mode="region_batch"),
    )
    search = RecordSearch(
        slots,
        cards=dict(cards),
        ordinary_cards=list(cards),
        ontology=snapshot,
        index=index,
        policy=policy,
        routing_cards=routes,
        metadata=metadata,
    )
    key, row = next(iter(search.rows.items()))
    source_ids = sorted(cards)
    merged = search._merged_card(source_ids)
    search.rows[key] = {
        **row,
        "ranked": True,
        "routing_status": "routed",
        "route_card_ids": [card.routing_card_id for card in routes],
        "source_card_ids": source_ids,
        "execution_card_id": merged.schema_card_id,
        "remaining": [{"card_id": merged.schema_card_id, "score": 0.9}],
    }

    assert search.admit_dependency(index.records[0].record_id, {PRODUCT}, set())
    saved = search.rows.drain()

    restored = RecordSearch(
        slots,
        cards=dict(cards),
        ordinary_cards=list(cards),
        ontology=snapshot,
        index=index,
        policy=policy,
        routing_cards=routes,
        metadata=metadata,
    )
    restored.restore(saved)
    _, product_card_id, _ = restored.take()
    assert restored.cards[product_card_id].class_iris == [PRODUCT]
    _, remaining_card_id, _ = restored.take()
    assert restored.cards[remaining_card_id].class_iris == [DEVICE]


def projection_search(tmp_path, *, graph_phase="evidence_verification"):
    """Two source menus produce different projections for the same requested type."""
    index, metadata = document_index(tmp_path)
    snapshot = ontology()
    first = compile_record_schema_card(
        snapshot, class_iris=[PRODUCT, DEVICE], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    code = SlotSpec(
        iri="urn:productCode", label="产品编号", declared_by=[PRODUCT],
        datatype_iris=["http://www.w3.org/2001/XMLSchema#string"],
    )
    snapshot.classes[PRODUCT] = definition(PRODUCT, "产品", properties=[code])
    second = compile_record_schema_card(
        snapshot, class_iris=[PRODUCT, PART], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    base_cards = {card.schema_card_id: card for card in (first, second)}
    plan = compile_schema_region_routing(
        snapshot, ROOT, analysis_scope_ref="scope", max_hops=2,
    )
    routes = [card.model_copy(update={"execution_card_ids": list(base_cards)})
              for card in plan.routing_cards]
    slots = [([record.record_id], [], None) for record in index.records[:2]]

    def factory():
        return RecordSearch(
            slots, cards=dict(base_cards), ordinary_cards=list(base_cards),
            ontology=snapshot, index=index, metadata=metadata, routing_cards=routes,
            policy=RecordDiscoveryPolicy(
                graph_phase=graph_phase,
                schema_region_routing=SchemaRegionRoutingPolicy(execution_mode="region_batch"),
            ),
        )

    search = factory()
    for (key, row), source_id in zip(search.rows.items(), base_cards, strict=True):
        merged = search._merged_card([source_id])
        search.rows[key] = {
            **row, "ranked": True, "routing_status": "routed",
            "route_card_ids": [card.routing_card_id for card in routes],
            "source_card_ids": [source_id], "execution_card_id": merged.schema_card_id,
            "remaining": [{"card_id": merged.schema_card_id, "score": 0.9}],
        }
    return search, factory


@pytest.mark.parametrize("graph_phase", ["candidate_graph", "evidence_review"])
def test_candidate_dependency_projection_retains_selected_downstream_types_on_resume(
    tmp_path, graph_phase,
):
    search, factory = projection_search(tmp_path, graph_phase=graph_phase)
    first, second = list(search.rows)
    search.rows[first] = {**search.rows[first], "remaining": []}
    record_id = search.rows[second]["source_record_ids"][0]
    assert search.admit_dependency(record_id, {PRODUCT}, set())
    remaining = search.rows[second]["remaining"]
    assert [search.cards[item["card_id"]].class_iris for item in remaining] == [[PRODUCT], [PART]]
    restored = factory()
    restored.restore(search.rows.drain())
    while search.pending:
        assert restored.take() == search.take()
    assert not restored.pending


def test_dependency_comparison_does_not_register_unselected_projection(tmp_path):
    search, factory = projection_search(tmp_path)
    initial_cards = set(search.cards)
    projections = {
        key: search._projected_card(row["source_card_ids"], {PRODUCT}, register=False)
        for key, row in search.rows.items()
    }
    assert len({card.schema_card_id for card in projections.values()}) == 2
    assert set(search.cards) == initial_cards

    record_id = next(iter(search.rows.values()))["source_record_ids"][0]
    assert search.admit_dependency(record_id, {PRODUCT}, set())
    adopted = [key for key, row in search.rows.items() if row["card_projections"]]
    assert len(adopted) == 1
    for key, card in projections.items():
        assert (card.schema_card_id in search.cards) is (key in adopted)
    restored = factory()
    restored.restore(search.rows.drain())
    assert set(restored.cards) == set(search.cards)
    assert restored.scope_hash == search.scope_hash


@pytest.mark.parametrize("consume", [0, 1, 3])
def test_adopted_projection_survives_consumption_and_restore_does_not_repeat_work(
    tmp_path, consume,
):
    search, factory = projection_search(tmp_path)
    record_id = next(iter(search.rows.values()))["source_record_ids"][0]
    assert search.admit_dependency(record_id, {PRODUCT}, set())
    for _ in range(consume):
        if search.pending:
            search.take()
    restored = factory()
    restored.restore(search.rows.drain())
    assert set(restored.cards) == set(search.cards)
    model = Embeddings()
    model.embed = lambda _texts: pytest.fail("completed embeddings must not run again")
    restored.prepare(
        RankingService(RankingPolicy(mode="semantic", enable_reranker=False), model),
        lambda: pytest.fail("completed region selection must not run again"),
    )
    original_remaining, restored_remaining = [], []
    while search.pending:
        original_remaining.append(search.take())
    while restored.pending:
        restored_remaining.append(restored.take())
    assert restored_remaining == original_remaining
    if consume == 3:
        assert restored_remaining == []


@pytest.mark.parametrize("mutation,reason", [
    ("source", "record_search_projection_source_changed"),
    ("classes", "record_search_projected_card_changed"),
    ("outside_class", "record_schema_projection_outside_source"),
    ("hash", "record_search_projected_card_changed"),
    ("missing", "record_search_scope_changed"),
])
def test_consumed_projection_restore_rejects_changed_recipe(tmp_path, mutation, reason):
    search, factory = projection_search(tmp_path)
    record_id = next(iter(search.rows.values()))["source_record_ids"][0]
    assert search.admit_dependency(record_id, {PRODUCT}, set())
    while search.pending:
        search.take()
    saved = deepcopy(search.rows.drain())
    row = next(item["value"] for item in saved.values() if item["value"]["card_projections"])
    card_id = next(identity for identity, recipe in row["card_projections"].items()
                   if recipe["class_iris"] == [PRODUCT])
    recipe = row["card_projections"][card_id]
    if mutation == "source":
        recipe["source_card_ids"] = []
    elif mutation == "classes":
        source = search.cards[row["source_card_ids"][0]]
        recipe["class_iris"] = sorted(set(source.class_iris) - {PRODUCT})
    elif mutation == "outside_class":
        recipe["class_iris"].append("urn:outsideSource")
    elif mutation == "hash":
        row["card_projections"]["changed-card"] = row["card_projections"].pop(card_id)
    else:
        row["card_projections"].pop(card_id)
    with pytest.raises(ValueError, match=reason):
        factory().restore(saved)


def test_schema_card_merge_rejects_conflicting_property_definitions():
    snapshot = ontology()
    card = compile_record_schema_card(
        snapshot, class_iris=[DEVICE], analysis_scope_ref="scope",
        profile=ExtractionProfile(),
    )
    left_property = SlotSpec(iri="urn:serial", label="编号", declared_by=[DEVICE])
    right_property = left_property.model_copy(update={"label": "设备编号"})
    left = card.model_copy(update={
        "schema_card_id": "left",
        "class_cards": [card.class_cards[0].model_copy(update={
            "properties": [left_property],
        })],
    })
    right = card.model_copy(update={
        "schema_card_id": "right",
        "class_cards": [card.class_cards[0].model_copy(update={
            "properties": [right_property],
        })],
    })

    with pytest.raises(ValueError, match="record_schema_class_definition_conflict"):
        merge_record_schema_cards([left, right])


@pytest.mark.parametrize("phase", ["evidence_verification", "evidence_review"])
def test_real_region_batch_folds_property_fields_into_the_discovery_request(
    tmp_path, monkeypatch, current_run, phase,
):
    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run,
        texts=["编号：A-001", "装置甲。", "部件乙。"],
    )
    adapter = factory().adapter
    adapter.record_discovery = adapter.record_discovery.model_copy(update={
        "graph_phase": phase,
        "schema_region_routing": SchemaRegionRoutingPolicy(
            max_regions_per_card=1,
            minimum_similarity=0.25,
            execution_mode="region_batch",
            property_field_mode="region_batch",
        ),
    })
    nodes = [node.model_copy(update={
        "heading": "装置与部件信息", "path": ["装置与部件信息"],
    }) for node in args["metadata"].node_summaries]
    metadata = args["metadata"].model_copy(update={
        "node_summaries": nodes,
        "dependency_hash": evidence_hash([args["metadata"].dependency_hash, nodes]),
    })
    args["metadata"] = metadata
    adapter.metadata = metadata

    result = factory().run(**args, **hooks)
    discovery = [view for view in requests if view.get("stage") == "discovery"
                 and "members" not in view]
    if phase == "evidence_verification":
        assert len(discovery) == 1
    # Evidence review may also explore newly registered downstream entities;
    # the field remains a reading aid on its owning source region.
    field_reads = [view for view in discovery if view.get("field_reading_groups")]
    assert field_reads
    assert all(view["field_reading_groups"][0]["fields"][0]["value"] == "A-001"
               for view in field_reads)
    assert all(not group["shared_subject_established"] for view in field_reads
               for group in view["field_reading_groups"])
    assert not any(view.get("attribute_disambiguation") for view in requests)
    assert result.graph.progress.record_discovery.execution_mode == "region_batch"
    assert result.graph.progress.record_discovery.property_field_mode == "region_batch"


@pytest.mark.parametrize("execution_mode", ["per_card", "region_batch"])
def test_real_executor_freezes_routes_and_cold_resume_reuses_paid_results(
    tmp_path, monkeypatch, current_run, execution_mode,
):
    from app.services.document_analysis import current_state

    args, factory, requests, hooks = setup_contextual(
        tmp_path, monkeypatch, current_run, texts=["装置甲。", "部件乙。"],
    )
    adapter = factory().adapter
    adapter.record_discovery = adapter.record_discovery.model_copy(update={
        "schema_region_routing": SchemaRegionRoutingPolicy(
            max_regions_per_card=1, minimum_similarity=0.25,
            execution_mode=execution_mode,
        ),
    })
    nodes = [node.model_copy(update={
        "heading": "装置与部件信息", "path": ["装置与部件信息"],
    }) for node in args["metadata"].node_summaries]
    metadata = args["metadata"].model_copy(update={
        "node_summaries": nodes,
        "dependency_hash": evidence_hash([args["metadata"].dependency_hash, nodes]),
    })
    args["metadata"] = metadata
    adapter.metadata = metadata

    runner = factory()
    result = runner.run(**args, **hooks)
    discovery = [view for view in requests if view.get("stage") == "discovery"
                 and "members" not in view and not view.get("attribute_disambiguation")]
    assert len(discovery) == 1
    assert [item["class_iri"] for item in discovery[0]["schema_card"]["class_cards"]] == [
        RECORD_DEVICE,
    ]
    assert all({unit["text"] for unit in view["evidence_units"] if unit["fact_eligible"]}
               == {"装置甲。", "部件乙。"} for view in discovery)
    assert all(
        not unit["fact_eligible"]
        for view in discovery
        for unit in view["evidence_units"]
        if unit["role"] == "binding"
    )
    diagnostics = result.graph.progress.record_discovery
    assert diagnostics.routing_cards == 1
    assert diagnostics.selected_regions == 1
    assert diagnostics.routed_groups == 1
    assert diagnostics.unrouted_groups == 0

    store, run, _ = current_run
    before = len(requests)
    factory().run(
        **args,
        **hooks,
        resume_state=vars(current_state.restore_work(store, run, run.run_fingerprint)),
        model_call_state=current_state.restore_calls(store, run, run.run_fingerprint),
    )
    assert len(requests) == before
