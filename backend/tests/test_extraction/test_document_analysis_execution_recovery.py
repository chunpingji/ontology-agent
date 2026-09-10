"""Deterministic hard-crash and startup recovery checks for document runs."""

from __future__ import annotations

import json
from collections import Counter
from copy import deepcopy
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select, update

from app.api import document_analysis
from app.config import settings
from app.models.document_analysis import (
    DocumentAnalysisArtifact,
    DocumentAnalysisExecution,
    DocumentAnalysisRun,
    DocumentRecognitionEventBatch,
    DocumentRunArtifact,
    DocumentRunCandidateHead,
    DocumentVerificationProofHead,
)
from app.services.document_analysis import execution as execution_service
from app.services.document_analysis.run_store import (
    DocumentAnalysisRunStore,
    FenceViolation,
)
from app.services.extraction.ontology_guided.contracts import GraphNode
from app.services.extraction.ontology_guided.executor import TaskOutcome
from app.services.llm.model_runtime import ModelWaitFailure

ROOT_IRI = "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"


class SimulatedProcessCrash(BaseException):
    """Escape normal worker error handling like an actual process termination."""


class CountingAdapter:
    model_identity = "durable-checkpoint-test-adapter-v1"

    def __init__(self) -> None:
        self.calls: list[str] = []

    def inspect(self, task, _context, _predicate, _menu):
        self.calls.append(task.task_id)
        first_call_payload = {}
        if len(self.calls) == 1:
            target_id = f"target:{task.task_id}"
            first_call_payload = {
                "nodes": [
                    GraphNode(
                        entity_id="checkpoint-only-candidate",
                        revision=1,
                        class_iri=task.subject.class_iri,
                        class_label="检查点候选",
                        label="检查点候选",
                    )
                ],
                "proof_payloads": [
                    {
                        "proof_id": "checkpoint-proof",
                        "proof_revision": 1,
                        "target_id": target_id,
                    }
                ],
                "decision_payloads": [
                    {
                        "decision_id": "checkpoint-decision",
                        "target_id": target_id,
                        "verdict": "unsupported",
                    }
                ],
            }
        return TaskOutcome(
            semantic_outcome="unsupported",
            reason_code="record_not_supporting_target",
            reason="确定性恢复测试",
            model_calls=1,
            **first_call_payload,
        )


def _word_bytes(tmp_path: Path) -> bytes:
    document = Document()
    document.add_heading("第一章 产品概要", level=1)
    document.add_paragraph("本报告包含用于恢复测试的产品和设备描述。")
    document.add_heading("1.1 质量属性", level=2)
    document.add_paragraph("外观为白色片剂，生产使用冻干机。")
    path = tmp_path / "checkpoint-source.docx"
    document.save(path)
    return path.read_bytes()


def _create_pending_run(client, analyst_headers, tmp_path, monkeypatch, *, key: str) -> str:
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *_args, **_kwargs: None)
    response = client.post(
        "/api/document-analysis/runs",
        headers=analyst_headers,
        files={
            "file": (
                "恢复测试.docx",
                _word_bytes(tmp_path),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={
            "root_class_iri": ROOT_IRI,
            "request_key": key,
            "metadata_mode": "structure_only",
        },
    )
    assert response.status_code == 202, response.text
    return response.json()["recognition_run_id"]


def _claim(db, run_id: str) -> tuple[DocumentAnalysisRunStore, str]:
    store = DocumentAnalysisRunStore(db)
    token = store.claim(
        run_id,
        "analyst",
        actor="checkpoint-test",
        worker_id="checkpoint-worker",
        lease_seconds=30,
    )
    db.commit()
    return store, token


def _expire_lease(db, run_id: str) -> None:
    db.execute(
        update(DocumentAnalysisExecution)
        .where(DocumentAnalysisExecution.recognition_run_id == run_id)
        .values(lease_expires_at=datetime.now(UTC) - timedelta(seconds=1))
    )
    db.commit()
    db.expire_all()


@pytest.mark.parametrize("failure_class", [RuntimeError, ModelWaitFailure])
def test_worker_failure_keeps_bounded_diagnostics_private(
    client, db, analyst_headers, tmp_path, monkeypatch, failure_class,
):
    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="private-worker-diagnostics"
    )
    private_value = "synthetic-private-exception-value"

    def raise_deep_failure(depth):
        if depth:
            return raise_deep_failure(depth - 1)
        raise failure_class(private_value)

    def fail_worker(*_args, **_kwargs):
        raise_deep_failure(12)

    monkeypatch.setattr(execution_service, "_execute_claimed", fail_worker)
    store, token = _claim(db, run_id)
    execution_service._execute_dispatched_run(
        db, store, store.get_owned(run_id, "analyst"), token
    )
    run = store.get_owned(run_id, "analyst")
    assert run.execution_status == "failed"
    assert run.error["code"] == "ANALYSIS_FAILED"
    assert run.error["failure_type"] == failure_class.__name__
    frames = run.error["failure_frames"]
    assert len(frames) == 8
    assert all(set(frame) == {"file", "line", "function"} for frame in frames)
    assert frames[-1]["function"] == "raise_deep_failure"
    assert all(isinstance(frame["line"], int) and frame["line"] > 0 for frame in frames)
    assert private_value not in json.dumps(run.error)

    for suffix in ("", "/graph", "/events"):
        response = client.get(
            f"/api/document-analysis/runs/{run_id}{suffix}", headers=analyst_headers
        )
        assert response.status_code == 200, response.text
        for private_detail in (
            "failure_type", "failure_frames", "raise_deep_failure", __file__, private_value,
        ):
            assert private_detail not in response.text
        if not suffix:
            error = response.json()["error"]
            assert set(error) == {"code", "stage", "retryable", "safe_detail", "occurred_at"}
            assert error["code"] == "ANALYSIS_FAILED"
            assert error["retryable"] is True
            assert error["safe_detail"] == "文档分析执行失败"


def _crash_after_first_batch(db, run_id: str, monkeypatch):
    store, old_token = _claim(db, run_id)
    original = execution_service._persist_recognition_batch

    def persist_then_crash(*args, **kwargs):
        original(*args, **kwargs)
        raise SimulatedProcessCrash("after committed recognition batch")

    monkeypatch.setattr(execution_service, "_persist_recognition_batch", persist_then_crash)
    run = store.get_owned(run_id, "analyst")
    with pytest.raises(SimulatedProcessCrash):
        execution_service._execute_claimed(db, store, run, old_token)
    monkeypatch.setattr(execution_service, "_persist_recognition_batch", original)
    db.expire_all()
    return store, old_token


def _checkpoint_payload(db, run_id: str) -> tuple[dict, DocumentRecognitionEventBatch]:
    receipt = db.scalar(
        select(DocumentRecognitionEventBatch).where(
            DocumentRecognitionEventBatch.recognition_run_id == run_id
        )
    )
    assert receipt is not None
    artifact = db.get(DocumentAnalysisArtifact, receipt.checkpoint_artifact_id)
    assert artifact is not None and artifact.payload is not None
    return artifact.payload, receipt


def test_committed_batch_is_secret_free_closed_waterline_and_resumes_without_recall(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="crash-after-first-batch"
    )
    adapter = CountingAdapter()
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)

    store, old_token = _crash_after_first_batch(db, run_id, monkeypatch)
    first_task_id = adapter.calls[0]
    checkpoint, receipt = _checkpoint_payload(db, run_id)
    serialized = json.dumps(checkpoint, ensure_ascii=False, sort_keys=True)

    assert old_token not in serialized
    assert "execution_token" not in serialized
    assert checkpoint["execution_generation"] == 1
    assert checkpoint["event_seq"] == receipt.last_sequence == receipt.first_sequence
    assert checkpoint["dependency_index"]["requirements"] == {
        "checkpoint-proof@1": [
            ["checkpoint-decision@1"],
        ]
    }
    run = store.get_owned(run_id, "analyst")
    assert run.execution_status == "running"
    assert run.event_head == checkpoint["event_seq"]

    candidate_heads = {
        (row.candidate_id, row.revision)
        for row in db.scalars(
            select(DocumentRunCandidateHead).where(
                DocumentRunCandidateHead.recognition_run_id == run_id
            )
        )
    }
    proof_heads = {
        (row.proof_id, row.proof_revision)
        for row in db.scalars(
            select(DocumentVerificationProofHead).where(
                DocumentVerificationProofHead.recognition_run_id == run_id
            )
        )
    }
    assert candidate_heads == {
        (item["id"], item["revision"]) for item in checkpoint["candidate_head_refs"]
    }
    assert proof_heads == {(item["id"], item["revision"]) for item in checkpoint["proof_head_refs"]}
    assert "checkpoint-only-candidate" in {item[0] for item in candidate_heads}
    assert "checkpoint-proof" in {item[0] for item in proof_heads}

    graph_ref = db.scalar(
        select(DocumentRunArtifact).where(
            DocumentRunArtifact.recognition_run_id == run_id,
            DocumentRunArtifact.artifact_kind == "graph",
            DocumentRunArtifact.event_head == checkpoint["event_seq"],
        )
    )
    assert graph_ref is not None
    graph_artifact = db.get(DocumentAnalysisArtifact, graph_ref.artifact_id)
    assert graph_artifact is not None
    assert graph_artifact.payload["graph"] == checkpoint["graph_state"]

    _expire_lease(db, run_id)
    new_token = store.claim(
        run_id,
        "analyst",
        actor="checkpoint-recovery",
        worker_id="replacement-worker",
        lease_seconds=30,
    )
    db.commit()
    assert new_token != old_token
    with pytest.raises(FenceViolation):
        store.event_batch_receipt(
            run_id,
            "analyst",
            old_token,
            batch_id=receipt.batch_id,
            batch_hash=receipt.content_hash,
        )

    execution_service._execute_claimed(
        db,
        store,
        store.get_owned(run_id, "analyst"),
        new_token,
    )
    db.expire_all()
    assert db.get(DocumentAnalysisRun, run_id).execution_status == "finished"
    assert Counter(adapter.calls)[first_task_id] == 1


def test_hard_crash_before_terminal_event_reuses_checkpoint_without_new_graph_revision(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="crash-before-terminal-event"
    )
    adapter = CountingAdapter()
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    store, old_token = _claim(db, run_id)
    original_finish = execution_service._finish

    def crash_before_terminal(*_args, **_kwargs):
        raise SimulatedProcessCrash("before terminal event")

    monkeypatch.setattr(execution_service, "_finish", crash_before_terminal)
    with pytest.raises(SimulatedProcessCrash):
        execution_service._execute_claimed(db, store, store.get_owned(run_id, "analyst"), old_token)
    monkeypatch.setattr(execution_service, "_finish", original_finish)
    db.expire_all()

    run = store.get_owned(run_id, "analyst")
    checkpoint = execution_service._restore_recognition_checkpoint(
        db,
        store,
        run,
        final_fingerprint=run.run_fingerprint,
    )
    assert checkpoint is not None
    graph_head_before = store.get_artifact_head(run_id, "analyst", "graph")
    calls_before = list(adapter.calls)

    _expire_lease(db, run_id)
    replacement = store.claim(
        run_id,
        "analyst",
        actor="terminal-recovery",
        worker_id="replacement-worker",
        lease_seconds=30,
    )
    db.commit()
    execution_service._execute_claimed(
        db,
        store,
        store.get_owned(run_id, "analyst"),
        replacement,
    )
    db.expire_all()

    assert db.get(DocumentAnalysisRun, run_id).execution_status == "finished"
    assert adapter.calls == calls_before
    graph_head_after = store.get_artifact_head(run_id, "analyst", "graph")
    assert graph_head_after.revision == graph_head_before.revision
    assert graph_head_after.artifact_id == graph_head_before.artifact_id


@pytest.mark.parametrize("budget_setting,initial_value,changed_value", [
    ("evidence_max_tasks", None, None), ("semantic_ranking_pool_size", None, None),
    ("semantic_ranking_device", "cpu", "cuda:0"),
    ("semantic_ranking_device", "cuda:0", "cpu"),
    ("semantic_ranking_dtype", "float32", "float16"),
    ("semantic_ranking_dtype", "float16", "float32"),
    ("semantic_ranking_cuda_version", "12.6", "12.8"),
])
def test_changed_fingerprint_blocks_recovery_without_another_model_call(
    client, db, analyst_headers, tmp_path, monkeypatch, budget_setting, initial_value,
    changed_value,
):
    # Freeze an explicit valid baseline: changing application defaults must not
    # accidentally turn a drift test into a no-op. Cover both CPU/GPU directions.
    monkeypatch.setattr(settings, "semantic_ranking_device", "cuda:0")
    monkeypatch.setattr(settings, "semantic_ranking_dtype", "float32")
    if initial_value is not None:
        monkeypatch.setattr(settings, budget_setting, initial_value)
    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="fingerprint-mismatch"
    )
    adapter = CountingAdapter()
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    _store, _old_token = _crash_after_first_batch(db, run_id, monkeypatch)
    calls_before = list(adapter.calls)

    _expire_lease(db, run_id)
    monkeypatch.setattr(
        settings, budget_setting,
        getattr(settings, budget_setting) + 1 if changed_value is None else changed_value,
    )
    execution_service.dispatch_run(run_id, bind=db.get_bind())
    db.expire_all()

    run = db.get(DocumentAnalysisRun, run_id)
    assert run.execution_status == "blocked_dependency"
    assert run.error["code"] == "FINGERPRINT_MISMATCH"
    assert run.error["retryable"] is False
    assert adapter.calls == calls_before


def test_startup_recovery_claims_expired_pausing_run_and_finishes_pause(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="startup-pausing-recovery"
    )
    adapter = CountingAdapter()
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    store, old_token = _claim(db, run_id)
    running = store.get_owned(run_id, "analyst")
    pausing = store.request_control(
        run_id,
        "analyst",
        action="pause",
        expected_revision=running.revision,
        request_key="pause-before-worker-crash",
    )
    db.commit()
    assert pausing.execution_status == "pausing"
    _expire_lease(db, run_id)

    def parse_after_pause(*_args, **_kwargs):
        pytest.fail("replacement worker started parsing after a durable pause request")

    monkeypatch.setattr(execution_service, "ensure_docx", parse_after_pause)

    assert execution_service.recover_document_analysis_runs(bind=db.get_bind()) == 1
    db.expire_all()
    recovered = db.get(DocumentAnalysisRun, run_id)

    assert recovered.execution_status == "paused"
    assert recovered.stop_reason == "operator_pause"
    assert adapter.calls == []
    with pytest.raises(FenceViolation):
        store.assert_fence(run_id, "analyst", old_token)
    assert execution_service.recover_document_analysis_runs(bind=db.get_bind()) == 0


def test_ranking_commit_is_fenced_immutable_and_read_only_before_graph(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    from copy import deepcopy

    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="ranking-boundary"
    )
    store, token = _claim(db, run_id)
    run = store.get_owned(run_id, "analyst")
    run = store.update_stage(
        run_id, "analyst", token, expected_revision=run.revision,
        stage="recognition", run_fingerprint="ranking-fingerprint",
    )
    db.commit()
    state = {
        "recognition_run_id": run_id, "run_fingerprint": "ranking-fingerprint",
        "service": {"policy": {"mode": "semantic"}, "costs": {"model_calls": 2},
                    "epochs": [{"epoch_id": "epoch:one", "status": "committed",
                                "actual_ranking_mode": "deterministic", "degraded": True,
                                "reason": "model_unavailable", "ordered_record_ids": []}]},
    }
    execution_service._persist_ranking_state(
        db, store, run, token, final_fingerprint="ranking-fingerprint", state=state
    )
    current = store.get_owned(run_id, "analyst")
    waterline = (current.revision, current.event_head, current.artifact_revision)
    restored = execution_service._restore_ranking_state(
        db, store, current, final_fingerprint="ranking-fingerprint"
    )
    assert restored == state
    for _ in range(2):
        response = client.get(
            f"/api/document-analysis/runs/{run_id}/graph", headers=analyst_headers
        )
        assert response.status_code == 200
        assert response.json()["graph_snapshot"] is None
        assert response.json()["ranking"]["committed_epochs"] == 1
        assert response.json()["ranking"]["degraded"]
    db.expire_all()
    current = store.get_owned(run_id, "analyst")
    assert (current.revision, current.event_head, current.artifact_revision) == waterline
    changed = deepcopy(state)
    changed["service"]["epochs"][0]["reason"] = "silently_reordered"
    with pytest.raises(execution_service.CheckpointMismatch):
        execution_service._persist_ranking_state(
            db, store, current, token, final_fingerprint="ranking-fingerprint", state=changed
        )
    _expire_lease(db, run_id)
    _claim(db, run_id)
    with pytest.raises(FenceViolation):
        execution_service._persist_ranking_state(
            db, store, current, token, final_fingerprint="ranking-fingerprint", state=state
        )
    assert execution_service._restore_ranking_state(
        db, store, store.get_owned(run_id, "analyst"), final_fingerprint="ranking-fingerprint"
    ) == state


def test_ranking_policy_and_model_configuration_are_frozen(monkeypatch):
    monkeypatch.setattr(settings, "semantic_ranking_enabled", False)
    first, identity = execution_service._configured_ranking()
    assert first.policy.mode == "deterministic"
    monkeypatch.setattr(settings, "semantic_ranking_max_tokens_per_run", 123456)
    second, changed = execution_service._configured_ranking()
    assert second.policy.max_ranking_tokens_per_run == 123456
    assert identity != changed
    assert identity["policy"]["intent_weights"] == {"discover": 0.5, "counterevidence": 0.5}


def test_ranking_factory_inherits_owner_cancellation_before_freezing(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    from app.services.llm.model_runtime import ModelCancelled, check_cancelled, runtime

    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="probe-owner-cancellation",
    )
    adapter = CountingAdapter()
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    stopped = False

    def factory():
        nonlocal stopped
        assert runtime.get()["run_id"] == run_id
        assert runtime.get()["bind"] is db.get_bind()
        assert callable(runtime.get()["should_stop"])
        stopped = True
        check_cancelled()
        pytest.fail("cancelled GPU probe returned to execution")

    monkeypatch.setattr(execution_service, "_configured_ranking", factory)
    store, token = _claim(db, run_id)
    with pytest.raises(ModelCancelled):
        execution_service._execute_claimed(
            db, store, store.get_owned(run_id, "analyst"), token,
            should_stop=lambda: stopped,
        )
    assert adapter.calls == []
    assert store.get_owned(run_id, "analyst").run_fingerprint is None


@pytest.mark.parametrize("failure_policy", ["pause", "deterministic"])
@pytest.mark.parametrize("reason", [
    "ranking_cuda_out_of_memory", "ranking_cuda_probe_timeout", "ranking_cuda_unavailable",
])
@pytest.mark.parametrize("state_boundary", ["checkpoint", "fingerprint_only"])
def test_transient_probe_failure_preserves_frozen_identity_and_can_resume(
    client, db, analyst_headers, tmp_path, monkeypatch, failure_policy, reason, state_boundary,
):
    from app.services.extraction.ontology_guided.semantic_reranker import (
        RankingPolicy,
        RankingService,
    )
    from tests.test_extraction.test_semantic_ranking import RankingModel

    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="transient-probe-recovery",
    )
    adapter, model = CountingAdapter(), RankingModel()
    model.identity = {"device": "cuda:0", "dtype": "float16", "cuda_version": "12.6",
                      "numeric_environment": {"device_uuid": "controlled-original-GPU"}}
    policy = RankingPolicy(mode="semantic", failure_policy=failure_policy)
    available = True

    def factory():
        return RankingService(policy, model if available else None), {
            "configuration": {"device": "cuda:0", "dtype": "float16", "cuda_version": "12.6"},
            "policy": policy.model_dump(mode="json"),
            "model_identity": deepcopy(model.identity) if available else None,
            "unavailable_reason": None if available else reason,
        }

    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    monkeypatch.setattr(execution_service, "_configured_ranking", factory)
    if state_boundary == "checkpoint":
        store, _token = _crash_after_first_batch(db, run_id, monkeypatch)
    else:
        store, token = _claim(db, run_id)
        persist = execution_service._persist_ranking_state

        def crash_before_first_ranking_state(*_args, **_kwargs):
            raise SimulatedProcessCrash("after frozen fingerprint before ranking state")

        monkeypatch.setattr(
            execution_service, "_persist_ranking_state", crash_before_first_ranking_state,
        )
        with pytest.raises(SimulatedProcessCrash):
            execution_service._execute_claimed(
                db, store, store.get_owned(run_id, "analyst"), token,
            )
        monkeypatch.setattr(execution_service, "_persist_ranking_state", persist)
        db.expire_all()
        assert adapter.calls == []
        assert store.get_artifact_head(run_id, "analyst", "ranking_state") is None
    current = store.get_owned(run_id, "analyst")
    fingerprint = current.run_fingerprint
    assert fingerprint is not None

    def artifact_head_id(kind):
        head = store.get_artifact_head(run_id, "analyst", kind)
        return head.artifact_id if head is not None else None

    heads = {
        kind: artifact_head_id(kind)
        for kind in ("ranking_state", "recognition_checkpoint", "graph")
    }
    main_calls, rank_calls = list(adapter.calls), list(model.calls)
    available = False
    _expire_lease(db, run_id)
    execution_service.dispatch_run(run_id, bind=db.get_bind())
    db.expire_all()
    current = store.get_owned(run_id, "analyst")
    assert current.execution_status == ("paused" if failure_policy == "pause" else "failed")
    assert current.error["code"] == (
        "RANKING_PAUSED" if failure_policy == "pause" else "RANKING_UNAVAILABLE"
    )
    assert current.error["retryable"] is True
    assert current.run_fingerprint == fingerprint
    assert adapter.calls == main_calls and model.calls == rank_calls
    assert all(artifact_head_id(kind) == identifier
               for kind, identifier in heads.items())

    available = True
    token = store.resume(
        run_id, "analyst", expected_revision=current.revision,
        actor="gpu-probe-retry", worker_id="gpu-recovered-worker",
    )
    db.commit()
    execution_service._execute_claimed(db, store, store.get_owned(run_id, "analyst"), token)
    assert store.get_owned(run_id, "analyst").execution_status == "finished"
    assert store.get_owned(run_id, "analyst").run_fingerprint == fingerprint
    assert len(adapter.calls) == len(set(adapter.calls))


def test_probe_unavailability_does_not_hide_changed_frozen_configuration(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    from app.services.extraction.ontology_guided.semantic_reranker import (
        RankingPolicy,
        RankingService,
    )
    from tests.test_extraction.test_semantic_ranking import RankingModel

    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="probe-real-drift",
    )
    adapter, model = CountingAdapter(), RankingModel()
    policy = RankingPolicy(mode="semantic", failure_policy="pause")
    available = True
    config = {"device": "cuda:0", "dtype": "float16", "cuda_version": "12.6"}

    def factory():
        return RankingService(policy, model if available else None), {
            "configuration": dict(config), "policy": policy.model_dump(mode="json"),
            "model_identity": model.identity if available else None,
            "unavailable_reason": None if available else "ranking_cuda_probe_timeout",
        }

    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    monkeypatch.setattr(execution_service, "_configured_ranking", factory)
    _crash_after_first_batch(db, run_id, monkeypatch)
    calls = list(adapter.calls)
    config["device"] = "cuda:1"
    available = False
    _expire_lease(db, run_id)
    execution_service.dispatch_run(run_id, bind=db.get_bind())
    db.expire_all()
    current = db.get(DocumentAnalysisRun, run_id)
    assert current.execution_status == "blocked_dependency"
    assert current.error["code"] == "FINGERPRINT_MISMATCH"
    assert current.error["retryable"] is False
    assert adapter.calls == calls


@pytest.mark.parametrize("crash_boundary", ["committed", "reserved"])
def test_crash_after_ranking_commit_restores_order_without_rescoring(
    client, db, analyst_headers, tmp_path, monkeypatch, crash_boundary,
):
    from app.services.extraction.ontology_guided.semantic_reranker import (
        RankingPolicy,
        RankingService,
    )

    class CountingRankingModel:
        identity = {"model": "controlled-local-ranker-v1"}

        def __init__(self):
            self.pairs = Counter()

        def count_tokens(self, text):
            return len(text)

        def embed(self, texts):
            return [[1.0, 0.0] for _ in texts]

        def score_pairs(self, pairs):
            from app.services.llm.model_runtime import runtime

            assert runtime.get()["run_id"] == run_id
            assert runtime.get()["bind"] is db.get_bind()
            self.pairs.update(pairs)
            return [-float(len(text)) for _query, text in pairs]

    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="ranking-crash-before-inspect"
    )
    adapter = CountingAdapter()
    model = CountingRankingModel()
    policy = RankingPolicy(mode="semantic")
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    monkeypatch.setattr(
        execution_service, "_configured_ranking",
        lambda: (RankingService(policy, model), {"policy": policy.model_dump(),
                                               "model": model.identity}),
    )
    store, token = _claim(db, run_id)
    persist = execution_service._persist_ranking_state

    def persist_then_crash(*args, **kwargs):
        persist(*args, **kwargs)
        service_state = kwargs["state"]["service"]
        if service_state["epochs"] or (
            crash_boundary == "reserved" and service_state["costs"]["model_calls"] > 0
        ):
            raise SimulatedProcessCrash("after committed ranking before record verification")

    monkeypatch.setattr(execution_service, "_persist_ranking_state", persist_then_crash)
    with pytest.raises(SimulatedProcessCrash):
        execution_service._execute_claimed(db, store, store.get_owned(run_id, "analyst"), token)
    assert adapter.calls == []
    first_pool_calls = dict(model.pairs)
    assert bool(first_pool_calls) == (crash_boundary == "committed")
    state = execution_service._restore_ranking_state(
        db, store, store.get_owned(run_id, "analyst"),
        final_fingerprint=store.get_owned(run_id, "analyst").run_fingerprint,
    )
    assert len(state["service"]["epochs"]) == (1 if crash_boundary == "committed" else 0)
    assert list(state["committed_at"].values()) == (
        [0] if crash_boundary == "committed" else []
    )
    reserved_calls = state["service"]["costs"]["model_calls"]
    assert reserved_calls > 0
    monkeypatch.setattr(execution_service, "_persist_ranking_state", persist)
    _expire_lease(db, run_id)
    store, replacement = _claim(db, run_id)
    execution_service._execute_claimed(
        db, store, store.get_owned(run_id, "analyst"), replacement
    )
    assert store.get_owned(run_id, "analyst").execution_status == "finished"
    assert all(model.pairs[pair] == count for pair, count in first_pool_calls.items())
    assert adapter.calls and len(set(adapter.calls)) == len(adapter.calls)
    current = store.get_owned(run_id, "analyst")
    recovered_state = execution_service._restore_ranking_state(
        db, store, current, final_fingerprint=current.run_fingerprint
    )
    assert recovered_state["service"]["costs"]["model_calls"] >= reserved_calls


def test_ranking_pause_is_durable_and_does_not_turn_into_empty_success(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    run_id = _create_pending_run(
        client, analyst_headers, tmp_path, monkeypatch, key="ranking-unavailable-pause"
    )
    adapter = CountingAdapter()
    monkeypatch.setattr(execution_service, "configured_model_adapter", lambda: adapter)
    monkeypatch.setattr(settings, "semantic_ranking_enabled", True)
    monkeypatch.setattr(settings, "semantic_ranking_failure_policy", "pause")
    monkeypatch.setattr(settings, "semantic_ranking_embedding_path", "")
    store, token = _claim(db, run_id)
    execution_service._execute_claimed(db, store, store.get_owned(run_id, "analyst"), token)
    paused = store.get_owned(run_id, "analyst")
    assert paused.execution_status == "paused"
    assert paused.stop_reason == "ranking_paused"
    assert adapter.calls == []
    ranking = client.get(
        f"/api/document-analysis/runs/{run_id}/graph", headers=analyst_headers
    ).json()["ranking"]
    assert ranking["paused"] and ranking["committed_epochs"] == 0
    assert ranking["epochs"][0]["status"] == "paused"
    fingerprint = paused.run_fingerprint
    initial = execution_service._restore_ranking_state(
        db, store, paused, final_fingerprint=fingerprint
    )
    token = store.resume(
        run_id, "analyst", expected_revision=paused.revision,
        actor="ranking-pause-test", worker_id="replacement-worker",
    )
    db.commit()
    execution_service._execute_claimed(db, store, store.get_owned(run_id, "analyst"), token)
    current = store.get_owned(run_id, "analyst")
    assert current.execution_status == "paused" and adapter.calls == []
    restored = execution_service._restore_ranking_state(
        db, store, current, final_fingerprint=fingerprint
    )
    assert restored["service"]["costs"] == initial["service"]["costs"]
    assert restored["service"]["pending_epochs"] == initial["service"]["pending_epochs"]
