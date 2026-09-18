"""Source/adapter closure, independent of simulated ranking scores."""

from __future__ import annotations

import json

from docx import Document

from app.services.extraction.evidence_identity import evidence_hash
from app.services.extraction.ontology_guided import model_adapter
from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    EdgeSpec,
    OntologyClassDefinition,
    OntologySnapshot,
    SlotSpec,
    SubjectRef,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.dependencies import DependencyIndex
from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor
from app.services.extraction.ontology_guided.metadata import prepare_metadata
from app.services.extraction.ontology_guided.model_adapter import LocalModelRecognitionAdapter
from app.services.extraction.ontology_guided.projection import project_graph
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.word_analysis import analyze_word_core

ROOT, PRODUCT, API = "urn:closure#Report", "urn:closure#Product", "urn:closure#API"
DESCRIBES, INGREDIENT = "urn:closure#describes", "urn:closure#hasActiveIngredient"
APPEARANCE = "urn:closure#appearance"


def _ontology():
    def definition(iri, label, **kwargs):
        values = {"iri": iri, "label": label, **kwargs}
        return OntologyClassDefinition(**values, source_hash=evidence_hash(values))

    classes = {
        ROOT: definition(
            ROOT,
            "报告",
            declared_relationships=[
                EdgeSpec(
                    iri=DESCRIBES,
                    label="描述",
                    range_class_iris=[PRODUCT],
                    declared_by=[ROOT],
                )
            ],
        ),
        PRODUCT: definition(
            PRODUCT,
            "产品",
            declared_relationships=[
                EdgeSpec(
                    iri=INGREDIENT,
                    label="活性成分",
                    range_class_iris=[API],
                    declared_by=[PRODUCT],
                )
            ],
            declared_properties=[
                SlotSpec(
                    iri=APPEARANCE,
                    label="外观",
                    declared_by=[PRODUCT],
                    max_count=1,
                )
            ],
        ),
        API: definition(API, "原料药"),
    }
    return OntologySnapshot(
        snapshot_id="closure-ontology",
        ontology_hash=evidence_hash(classes),
        classes=classes,
        created_from="frozen_fixture",
    )


def _analysis(tmp_path, *, negative=False):
    document = Document()
    document.add_heading("产品描述", 1)
    document.add_paragraph("本报告描述产品甲片。")
    document.add_heading("组成检验", 1)
    document.add_paragraph("产品甲片的活性成分为成分乙。")
    document.add_paragraph("另有试验来源-1及成分丙，两者仅在报告共现。")
    document.add_heading("其他原文", 1)
    document.add_paragraph("产品甲片外观为白色片剂。")
    if negative:
        document.add_paragraph("更正：产品甲片的活性成分不含成分乙。")
    path = tmp_path / "closure.docx"
    document.save(path)
    return analyze_word_core(path)


def _proposal(fragment, *, object_text, object_class, polarity="affirmed", subject_text="产品甲片"):
    def quote(text):
        return {"evidence_id": fragment["evidence_id"], "text": text}

    return {
        "kind": "relationship",
        "object_class_iri": object_class,
        "object_label": object_text,
        "object_quote": quote(object_text),
        "predicate_support": [quote(fragment["text"])],
        "subject_support": [quote(subject_text)],
        "subject_binding_verdict": "supported",
        "bridge_kind": "explicit_assertion",
        "type_verdict": "supported",
        "role_verdict": "supported",
        "predicate_verdict": "supported",
        "applicability_verdict": "supported",
        "polarity": polarity,
        "reason": "可回放的原文断言及主体角色。",
    }


def _run(analysis, adapter, **kwargs):
    metadata = prepare_metadata(
        analysis.ir,
        section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="closure-test",
    )
    return OntologyGuidedExecutor(
        ontology=_ontology(),
        engine=object(),
        adapter=adapter,
        phase1_section_limit=1,
    ).run(
        recognition_run_id="closure-run",
        run_fingerprint="closure-fingerprint",
        ir=analysis.ir,
        metadata=metadata,
        root_class_iri=ROOT,
        root_class_label="报告",
        filename="not-a-product-name.docx",
        **kwargs,
    )


def _effective(result, dependencies=None):
    graph = result.graph
    return project_graph(
        recognition_run_id=graph.recognition_run_id,
        run_revision=graph.run_revision,
        event_head=graph.event_head,
        metadata_snapshot_id=graph.metadata_snapshot_id,
        root_ref=graph.root_ref,
        nodes=graph.nodes,
        edges=graph.edges,
        properties=graph.properties,
        coverage=graph.coverage,
        progress=graph.progress,
        dependency_index=dependencies,
        projection="effective",
    )


def _respond(request, *, check_binding=False, false_owner=False):
    predicate = request["predicate"]["iri"]
    if request.get("stage") == "verification":
        verifications = []
        for candidate in request["candidates"]:
            claim = candidate["claim"]
            fragment = next(
                item for item in request["fragments"]
                if item["evidence_id"] == claim["endpoint_quote"]["evidence_id"]
            )
            source = fragment["text"]
            if check_binding and predicate == INGREDIENT and "活性成分" in source:
                assert request["subject_evidence_refs"]
                assert any(
                    not item["fact_eligible"] and "本报告描述产品甲片" in item["text"]
                    for item in request["fragments"]
                )
            owner = "本报告" if predicate == DESCRIBES else "产品甲片"
            verdict = "supported" if owner in source else "undetermined"
            verifications.append({
                "candidate_id": candidate["candidate_id"],
                "target_id": candidate["target_id"],
                "type_verdict": "supported",
                "role_verdict": verdict,
                "subject_binding_verdict": verdict,
                "predicate_verdict": verdict,
                "applicability_verdict": "supported",
                "counterevidence_verdict": "undetermined",
                "bridge_verdict": verdict,
                "type_support": [{"evidence_id": fragment["evidence_id"], "text": source}],
                "predicate_support": [{"evidence_id": fragment["evidence_id"], "text": source}],
                "subject_support": (
                    [{"evidence_id": fragment["evidence_id"], "text": owner}]
                    if owner in source and not request["subject"]["is_document_root"] else []
                ),
                "condition_support": claim["conditions"],
                "counterevidence_support": [],
                "reason": "独立读取原文并核验当前主体、谓词和桥接角色。",
            })
        return {"verifications": verifications}
    for fragment in request["fragments"]:
        if not fragment["fact_eligible"]:
            continue
        text = fragment["text"]
        if predicate == DESCRIBES and text == "本报告描述产品甲片。":
            return {
                "proposals": [
                    _proposal(
                        fragment,
                        object_text="产品甲片",
                        object_class=PRODUCT,
                        subject_text="本报告",
                    )
                ]
            }
        if predicate == INGREDIENT and ("活性成分为" in text or "活性成分不含" in text):
            if check_binding:
                assert request["subject_evidence_refs"]
                assert any(
                    not item["fact_eligible"] and "本报告描述产品甲片" in item["text"]
                    for item in request["fragments"]
                )
            return {
                "proposals": [
                    _proposal(
                        fragment,
                        object_text="成分乙",
                        object_class=API,
                        polarity="negated" if "不含" in text else "affirmed",
                    )
                ]
            }
        if false_owner and predicate == INGREDIENT and "试验来源-1" in text:
            proposal = _proposal(fragment, object_text="成分丙", object_class=API)
            proposal["subject_support"] = []
            return {"proposals": [proposal]}
    return {"proposals": []}


def test_real_adapter_two_hops_receives_subject_source_and_rejects_unbound_owner(
    tmp_path,
    monkeypatch,
):
    analysis = _analysis(tmp_path)

    def respond(_client, *, user, **_kwargs):
        return _respond(json.loads(user), check_binding=True, false_owner=True)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    result = _run(analysis, LocalModelRecognitionAdapter(object(), model_identity="controlled"))
    effective = _effective(result)
    assert len(effective.edges) == 2
    assert {edge.predicate_iri for edge in effective.edges} == {DESCRIBES, INGREDIENT}
    assert {node.label for node in effective.nodes if not node.root} == {"产品甲片", "成分乙"}
    assert any(edge.decision_status == "undetermined" for edge in result.graph.edges)
    assert result.graph.progress.completion == "in_scope_complete"
    assert all(item.executed_phase_counts["phase2"] > 0 for item in result.graph.coverage)
    for edge in effective.edges:
        assert edge.predicate_evidence_refs
        for anchor in edge.evidence_refs:
            assert analysis.ir.resolve(anchor)


def test_required_context_bypasses_pool_quota_and_missing_or_budget_is_incomplete(tmp_path):
    document = Document()
    document.add_paragraph("目标记录。")
    for number in range(12):
        document.add_paragraph(f"必要来源 {number}：在对应批次条件下才适用。")
    path = tmp_path / "context.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    subject = SubjectRef(entity_id="root", revision=1, class_iri=ROOT, is_document_root=True)
    target = VerificationTarget.create(
        run_fingerprint="fingerprint",
        claim_ref=VersionedRef(id="claim", revision=1),
        task_id="task",
        check_kind="predicate_entailment",
        document_context=DocumentContext(
            document_hash=analysis.ir.document_hash,
            document_class_iri=ROOT,
            root_ref=VersionedRef(id="root", revision=1),
        ),
        subject_ref=subject,
        predicate_iri=DESCRIBES,
        ontology_hash="ontology",
        source_scope_hash="scope",
        context_hash="context",
    )
    refs = [analysis.ir.anchor(record.source_units[0].evidence_id) for record in index.records[1:]]
    context = assemble_context(
        target, index.records[0].record_id, index, required_context_refs=refs
    )
    assert len(context.fragments) == 13
    assert len(context.required_context_refs) == 12
    assert sum(fragment.fact_eligible for fragment in context.fragments) == 1
    assert context.budget_status == "within_budget"
    forged = refs[0].model_copy(update={"document_hash": "0" * 64})
    incomplete = assemble_context(
        target,
        index.records[0].record_id,
        index,
        required_context_refs=[forged],
    )
    assert incomplete.budget_status == "required_context_missing"
    assert incomplete.omitted_refs == [forged]

    class Counter:
        def count(self, text):
            return len(text)

    oversized = assemble_context(
        target,
        index.records[0].record_id,
        index,
        required_context_refs=refs,
        token_counter=Counter(),
        max_input_tokens=10,
    )
    assert oversized.budget_status == "context_budget_exceeded"
    assert len(oversized.fragments) == 13


def test_competitor_invalidates_transitive_dependencies_and_keeps_alternative():
    index = DependencyIndex()
    index.add_proof("old-edge", ["old-proof"])
    index.add_proof("child", ["old-edge"])
    index.add_proof("alternative-child", ["old-edge"])
    index.add_proof("alternative-child", ["independent-edge"])
    index.subscribe(
        "old-proof",
        record_or_group="subject-predicate",
        owner_role="product",
        applicability="batch-one",
    )
    affected = index.notify_competitor(
        record_or_group="subject-predicate",
        owner_role="product",
        applicability="batch-one",
    )
    assert {"old-proof", "old-edge", "child"} <= affected
    assert index.is_valid("alternative-child")
    assert not index.notify_competitor(
        record_or_group="subject-predicate",
        owner_role="product",
        applicability="batch-two",
    )


def test_late_counterevidence_blocks_path_and_survives_replay(tmp_path, monkeypatch):
    analysis, batches, seen = _analysis(tmp_path, negative=True), [], []

    def respond(_client, *, user, **_kwargs):
        request = json.loads(user)
        seen.append(request)
        return _respond(request)

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    adapter = LocalModelRecognitionAdapter(object(), model_identity="controlled")
    result = _run(analysis, adapter, batch_hook=batches.append)
    state = batches[-1]
    effective = _effective(result, DependencyIndex.from_snapshot(state.dependency_index))
    assert [edge.predicate_iri for edge in effective.edges] == [DESCRIBES]
    assert result.graph.progress.invalidated_count > 0
    assert result.graph.progress.unresolved_claims > 0
    assert any(request["required_counterevidence"] for request in seen)
    calls = len(seen)
    replay = _run(
        analysis,
        adapter,
        resume_state={
            "task_outcomes": state.task_outcomes,
            "frontier": state.frontier,
            "recall_ledger": state.recall_ledger,
            "dependency_index": state.dependency_index,
            "model_call_state": state.model_call_state,
        },
    )
    assert len(seen) == calls
    assert replay.graph.progress == result.graph.progress
