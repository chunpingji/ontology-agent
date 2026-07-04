"""015 基本信息 tab full-stack backend surfaces.

Covers what the AST Editor 基本信息 tab drives:
  - PATCH /api/ast-templates/{id} — name / doc_no / owner in-place edit (+ 409 on
    (name, version) collision)
  - POST /api/ast-templates/{id}/sample — replace default sample DOCX
  - POST/DELETE /api/ast-templates/{id}/default-source — 默认源文件 upload/clear
  - GET/POST/DELETE /api/ast-templates/{id}/training-pairs — 训练数据 CRUD

Fixtures (client/db/analyst_headers/operator_headers) come from tests/conftest.py,
mirroring test_reporting/test_template_meta_and_section_prompt.py.
"""

from __future__ import annotations

import io

from app.services.reporting.ast_template import load_default_template


def _seed_template(db, *, name="Basic", version="v1"):
    from app.models.extraction import AstTemplate

    row = AstTemplate(
        name=name, version=version, doc_no="DOC-1",
        schema_json=load_default_template().model_dump(),
        status="draft", created_by="creator",
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _docx_bytes() -> bytes:
    """Minimal real .docx so parse_word_to_tiptap yields text."""
    import docx

    doc = docx.Document()
    doc.add_heading("评估对象", level=1)
    doc.add_paragraph("本品为 XX 注射液，属注射剂。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── metadata edit ────────────────────────────────────────────────────────────


class TestMetaEditBasicInfo:
    def test_name_doc_no_owner_updated_in_place(self, client, db, analyst_headers):
        row = _seed_template(db)
        resp = client.patch(
            f"/api/ast-templates/{row.id}",
            headers=analyst_headers,
            json={"name": "改名后", "doc_no": "DOC-9", "owner": "张三"},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["name"] == "改名后"
        assert body["doc_no"] == "DOC-9"
        assert body["owner"] == "张三"
        assert body["version"] == "v1"  # no version bump
        db.refresh(row)
        assert row.name == "改名后" and row.owner == "张三"

    def test_rename_collision_returns_409(self, client, db, analyst_headers):
        _seed_template(db, name="占用", version="v1")
        row = _seed_template(db, name="待改", version="v1")
        resp = client.patch(
            f"/api/ast-templates/{row.id}",
            headers=analyst_headers,
            json={"name": "占用"},
        )
        assert resp.status_code == 409

    def test_clearing_owner_sets_null(self, client, db, analyst_headers):
        row = _seed_template(db)
        row.owner = "李四"
        db.commit()
        resp = client.patch(
            f"/api/ast-templates/{row.id}",
            headers=analyst_headers,
            json={"owner": ""},
        )
        assert resp.status_code == 200
        assert resp.json()["owner"] is None

    def test_role_gated_403(self, client, db, operator_headers):
        row = _seed_template(db)
        resp = client.patch(
            f"/api/ast-templates/{row.id}",
            headers=operator_headers,
            json={"name": "x"},
        )
        assert resp.status_code == 403


# ── default sample replacement ───────────────────────────────────────────────


class TestReplaceSample:
    def test_replace_updates_sample_content(self, client, db, analyst_headers):
        row = _seed_template(db)
        resp = client.post(
            f"/api/ast-templates/{row.id}/sample",
            headers=analyst_headers,
            files={"file": ("s.docx", _docx_bytes(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert resp.status_code == 200
        assert "注射液" in resp.json()["plain_text"]
        db.refresh(row)
        assert row.sample_content_json is not None
        assert "注射液" in (row.sample_text or "")

    def test_non_docx_rejected(self, client, db, analyst_headers):
        row = _seed_template(db)
        resp = client.post(
            f"/api/ast-templates/{row.id}/sample",
            headers=analyst_headers,
            files={"file": ("s.txt", b"hello", "text/plain")},
        )
        assert resp.status_code == 422


# ── default source file ──────────────────────────────────────────────────────


class TestDefaultSource:
    def test_upload_then_delete(self, client, db, analyst_headers):
        row = _seed_template(db)
        up = client.post(
            f"/api/ast-templates/{row.id}/default-source",
            headers=analyst_headers,
            files={"file": ("origin.docx", b"raw-bytes", "application/octet-stream")},
        )
        assert up.status_code == 200
        assert up.json()["default_source_filename"] == "origin.docx"
        db.refresh(row)
        assert row.default_source_path

        rm = client.delete(
            f"/api/ast-templates/{row.id}/default-source",
            headers=analyst_headers,
        )
        assert rm.status_code == 200
        assert rm.json()["default_source_filename"] is None
        db.refresh(row)
        assert row.default_source_path is None


# ── training pairs ───────────────────────────────────────────────────────────


class TestTrainingPairs:
    def test_create_list_delete(self, client, db, analyst_headers):
        row = _seed_template(db)
        # create with source only (report optional/absent)
        c1 = client.post(
            f"/api/ast-templates/{row.id}/training-pairs",
            headers=analyst_headers,
            files={"source_file": ("src.docx", b"src", "application/octet-stream")},
        )
        assert c1.status_code == 201
        assert c1.json()["source_filename"] == "src.docx"
        assert c1.json()["report_filename"] is None

        # create with both source + report
        c2 = client.post(
            f"/api/ast-templates/{row.id}/training-pairs",
            headers=analyst_headers,
            files={
                "source_file": ("src2.docx", b"src2", "application/octet-stream"),
                "report_file": ("rep2.docx", b"rep2", "application/octet-stream"),
            },
        )
        assert c2.status_code == 201
        assert c2.json()["report_filename"] == "rep2.docx"
        pair2_id = c2.json()["id"]

        lst = client.get(
            f"/api/ast-templates/{row.id}/training-pairs", headers=analyst_headers
        )
        assert lst.status_code == 200
        assert len(lst.json()) == 2

        rm = client.delete(
            f"/api/ast-templates/{row.id}/training-pairs/{pair2_id}",
            headers=analyst_headers,
        )
        assert rm.status_code == 204
        lst2 = client.get(
            f"/api/ast-templates/{row.id}/training-pairs", headers=analyst_headers
        )
        assert len(lst2.json()) == 1

    def test_pairs_surfaced_in_detail(self, client, db, analyst_headers):
        row = _seed_template(db)
        client.post(
            f"/api/ast-templates/{row.id}/training-pairs",
            headers=analyst_headers,
            files={"source_file": ("only.docx", b"x", "application/octet-stream")},
        )
        detail = client.get(f"/api/ast-templates/{row.id}", headers=analyst_headers)
        assert detail.status_code == 200
        assert len(detail.json()["training_pairs"]) == 1

    def test_delete_wrong_template_404(self, client, db, analyst_headers):
        row = _seed_template(db)
        other = _seed_template(db, name="Other")
        made = client.post(
            f"/api/ast-templates/{row.id}/training-pairs",
            headers=analyst_headers,
            files={"source_file": ("s.docx", b"x", "application/octet-stream")},
        )
        pid = made.json()["id"]
        resp = client.delete(
            f"/api/ast-templates/{other.id}/training-pairs/{pid}",
            headers=analyst_headers,
        )
        assert resp.status_code == 404
