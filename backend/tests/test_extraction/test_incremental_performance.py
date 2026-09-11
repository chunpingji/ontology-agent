"""Bounded semantic work, whole-method proof and source-priority behavior."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.services.extraction.ontology_guided.context import assemble_context
from app.services.extraction.ontology_guided.contracts import (
    DocumentContext,
    EdgeSpec,
    GraphNode,
    LocalMenu,
    RangeClass,
    SlotSpec,
    SubjectRef,
    VerificationTarget,
    VersionedRef,
)
from app.services.extraction.ontology_guided.heuristic_search import (
    HeuristicSearchPolicy,
    HeuristicSlotSearch,
)
from app.services.extraction.ontology_guided.process_granularity import (
    CLEANING_PROCESS,
    method_fields,
    validate_method_scope,
)
from app.services.extraction.ontology_guided.ranking_execution import RankingPreparation
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.repair_adapter import EvidenceRepairAdapter
from app.services.extraction.ontology_guided.scheduler import FrontierScheduler, RecognitionTask
from app.services.extraction.ontology_guided.semantic_reranker import RankingPolicy, RankingService
from app.services.extraction.ontology_guided.source_priorities import explicit_attribute_sources
from app.services.extraction.word_analysis import analyze_word_core
from app.services.llm.model_runtime import ModelCancelled
from tests.test_extraction.test_heuristic_search import _observe_page, _search
from tests.test_extraction.test_semantic_ranking import RankingModel, setup_slot


def test_h2_consumes_batches_and_restores_mid_batch(tmp_path):
    policy = HeuristicSearchPolicy.durable(incremental=True)
    texts = [f"背景段落 {n}" for n in range(160)]
    search = _search(tmp_path, texts, policy=policy)
    assert search.next_admission() is None and search.needs_semantic
    for number, size in enumerate((8, 16, 32, 64), 1):
        assert search.semantic_boundary["pool_limit"] == size
        search.accept_semantic(search.deferred_record_ids[:size], f"epoch-{number}", committed=True)
        page = search.next_admission()
        assert page.stage == "H2" and search.next_admission() is None
        state = search.snapshot()
        restored = HeuristicSlotSearch(
            plan=search.plan,
            predicate=search.predicate,
            search_index=search.search_index,
            policy=policy,
            run_fingerprint=search.run_fingerprint,
        )
        restored.restore(state)
        assert restored.snapshot() == state
        search = restored
        _observe_page(search, page)
        while page := search.next_admission():
            if page.stage == "H3":
                break
            assert not search.needs_semantic
            _observe_page(search, page)
        if number < 4:
            assert search.needs_semantic
    assert page.stage == "H3" and search.snapshot()["deferred_count"] > 0
    assert all(e.coverage_state == "unattempted" for e in search.plan.ledger.values())


def test_h2_success_defers_unchecked_tail(tmp_path):
    search = _search(
        tmp_path,
        [f"背景 {n}" for n in range(40)],
        policy=HeuristicSearchPolicy.durable(incremental=True),
    )
    search.next_admission()
    search.accept_semantic(search.deferred_record_ids[:8], "epoch-1", committed=True)
    _observe_page(search, search.next_admission(), supported=True)
    assert search.next_admission() is None
    assert search.status == "local_results_only" and len(search.deferred_record_ids) == 32


def test_small_pool_scores_both_intents_and_reuses_committed_work(tmp_path):
    search = _search(tmp_path, [f"背景 {n}" for n in range(90)])
    args = setup_slot(tmp_path)
    args.update(
        plan=search.plan,
        index=search.search_index.index,
        metadata=search.search_index.metadata,
        predicate=search.predicate,
        subject_node=GraphNode(
            entity_id="subject",
            revision=1,
            class_iri="urn:Product",
            class_label="产品",
            label="产品",
        ),
    )
    model = RankingModel()
    service = RankingService(
        RankingPolicy(
            mode="semantic", pool_size=64, batch_size=4, policy_version="semantic-ranking-v2"
        ),
        model,
    )
    for number, size in enumerate((8, 16), 1):
        boundary = {"version": "bounded-semantic-v1", "round": number, "pool_limit": size}
        prior = len(model.calls)
        epoch = service.prepare_next_epoch(**args, expansion_boundary=boundary)
        assert epoch.status == "ready" and len(epoch.record_ids) == size
        assert sum(n for kind, n in model.calls[prior:] if kind == "score") == size * 2
        assert {i["retrieval_intent"] for i in epoch.observations} == {
            "discover",
            "counterevidence",
        }
        service.commit_epoch(epoch)
        restored = RankingService(service.policy, model, state=service.snapshot())
        before = list(model.calls)
        restored.validate_epoch(epoch, **args, expansion_boundary=boundary)
        assert model.calls == before
        with pytest.raises(ValueError, match="boundary"):
            restored.validate_epoch(epoch, **args, expansion_boundary={**boundary, "round": 4})
        service = restored
    assert not set(service.epochs[0].record_ids) & set(service.epochs[1].record_ids)


def test_cancel_retains_reservation_before_dispatch(tmp_path):
    from time import monotonic, sleep

    model = RankingModel()
    preparation = RankingPreparation(
        RankingService(RankingPolicy(mode="semantic"), model),
        task_id="cancel-test",
        arguments=setup_slot(tmp_path),
    )
    deadline = monotonic() + 3
    try:
        while preparation._updates.empty() and monotonic() < deadline:
            sleep(0.01)
        assert not preparation._updates.empty()
        preparation.cancel()
        saved = []
        preparation.drain(saved.append)
        with pytest.raises(ModelCancelled):
            preparation.future.result(timeout=3)
        assert saved and preparation.service.costs["model_calls"] > 0 and model.calls == []
    finally:
        preparation.close()


@pytest.mark.parametrize("after_task", [False, True])
def test_executor_resumes_second_semantic_batch_without_scoring_it_again(tmp_path, after_task):
    from app.services.extraction.ontology_guided.executor import OntologyGuidedExecutor
    from tests.test_extraction.test_semantic_ranking_execution import NoFactsAdapter, fixture

    ontology, arguments = fixture(tmp_path)
    model = RankingModel()
    policy = RankingPolicy(
        mode="semantic", pool_size=64, batch_size=4, policy_version="semantic-ranking-v2"
    )
    batches = []
    ranking_states = []

    def execute(**kwargs):
        return OntologyGuidedExecutor(
            ontology=ontology,
            engine=object(),
            adapter=NoFactsAdapter(),
            max_tasks=24,
            evidence_repair=True,
            incremental_performance=True,
            heuristic_policy=HeuristicSearchPolicy.durable(incremental=True),
            ranking_service=RankingService(policy, model),
        ).run(**arguments, **kwargs)

    def interrupted(batch):
        batches.append(batch.model_copy(deep=True))
        if after_task and len(batch.task_outcomes) == 9:
            raise RuntimeError("after first task in second batch")

    def ranking_checkpoint(state):
        ranking_states.append(deepcopy(state))
        if not after_task and len(state["service"]["epochs"]) == 2:
            raise RuntimeError("after ranking commit in second batch")

    with pytest.raises(RuntimeError, match="second batch"):
        execute(batch_hook=interrupted, ranking_hook=ranking_checkpoint)
    last = batches[-1]
    latest = last.ranking_state if after_task else ranking_states[-1]
    epochs = latest["service"]["epochs"]
    assert [e["expansion_boundary"]["round"] for e in epochs] == [1, 2]
    before = len(model.calls)
    result = execute(
        resume_state={
            key: getattr(last, key)
            for key in (
                "task_outcomes",
                "frontier",
                "recall_ledger",
                "dependency_index",
                "diagnostics",
            )
        },
        ranking_state=latest,
        model_call_state=last.model_call_state,
    )
    assert result.ranking_state["service"]["epochs"][:2] == epochs
    assert sum(n for kind, n in model.calls[before:] if kind == "score") == 24


def cleaning_fixture(tmp_path, paragraphs):
    document = Document()
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("设备名称", "设备编号", "清洗方法")):
        cell.text = text
    for number in range(2):
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = "反应釜", f"E-{number}", paragraphs[0]
        for text in paragraphs[1:]:
            cells[2].add_paragraph(text)
    path = tmp_path / "cleaning.docx"
    document.save(path)
    analysis = analyze_word_core(path)
    index = RecordIndex(analysis.ir)
    record = next(r for r in index.records if any(u.text == "E-0" for u in r.source_units))
    subject = SubjectRef(
        entity_id="root", revision=1, class_iri="urn:Report", is_document_root=True
    )
    predicate = EdgeSpec(
        iri="urn:hasCleaningMethod",
        label="含清洗方法",
        range_class_iris=[CLEANING_PROCESS],
        range_classes=[RangeClass(iri=CLEANING_PROCESS, label="清洗过程")],
    )
    task = RecognitionTask.create(
        subject=subject,
        predicate_iri=predicate.iri,
        predicate_kind="relationship",
        record_id=record.record_id,
        phase=1,
        hop=0,
        dependency_hash="test",
    )
    target = VerificationTarget.create(
        run_fingerprint="cleaning-test",
        claim_ref=VersionedRef(id=task.claim_lineage_id, revision=1),
        task_id=task.task_id,
        check_kind="predicate_entailment",
        subject_ref=subject,
        document_context=DocumentContext(
            document_hash=analysis.ir.document_hash,
            document_class_iri=subject.class_iri,
            root_ref=VersionedRef(id="root", revision=1),
        ),
        predicate_iri=predicate.iri,
        ontology_hash="test",
        source_scope_hash="test",
        context_hash="test",
    )
    context = assemble_context(target, task.record_id, index, repair_enabled=True)
    context.incremental_performance = True
    return index, task, predicate, context


@pytest.mark.parametrize(
    "paragraphs", [["加水冲洗；丙酮循环；氮气吹干"], ["加水冲洗", "丙酮循环", "氮气吹干"]]
)
def test_method_scope_rejects_partial_actions_and_keeps_rows(tmp_path, paragraphs):
    index, task, predicate, context = cleaning_fixture(tmp_path, paragraphs)
    groups = method_fields(context)
    assert len(groups) == 1
    refs = [
        index.ir.anchor(u.evidence_id, 0, len(u.text))
        for u in index.by_id[task.record_id].source_units
        if u.text in paragraphs
    ]
    assert validate_method_scope(context, predicate, CLEANING_PROCESS, refs[0], refs) is None
    action = refs[-1] if len(refs) > 1 else refs[0].model_copy(update={"span_end": 4})
    assert validate_method_scope(context, predicate, CLEANING_PROCESS, action, refs) == (
        "partial_cleaning_method"
    )
    if len(refs) > 1:
        assert validate_method_scope(context, predicate, CLEANING_PROCESS, refs[0], refs[:1]) == (
            "cleaning_method_scope_incomplete"
        )
    assert validate_method_scope(context, predicate, "urn:ActionStep", action, refs) is None
    other = next(r for r in index.records if any(u.text == "E-1" for u in r.source_units))
    other_context = assemble_context(context.target, other.record_id, index, repair_enabled=True)
    assert method_fields(other_context)[0]["group_id"] != groups[0]["group_id"]
    context.incremental_performance = False
    assert validate_method_scope(context, predicate, CLEANING_PROCESS, action, refs) is None


@pytest.mark.parametrize("text", ["加水冲洗", "CIP-方法A", "标准清洗规程CP-2"])
def test_single_step_or_named_method_is_not_rejected_by_length(tmp_path, text):
    assert method_fields(cleaning_fixture(tmp_path, [text])[-1]) == []


def test_explicit_method_identifier_is_valid_with_following_action_paragraphs(tmp_path):
    index, task, predicate, context = cleaning_fixture(
        tmp_path, ["清洗方法编号：CP-001", "加水冲洗", "氮气吹干"],
    )
    unit = next(u for u in index.by_id[task.record_id].source_units if "CP-001" in u.text)
    endpoint = index.ir.anchor(unit.evidence_id, unit.text.index("CP-001"), len(unit.text))
    assert validate_method_scope(context, predicate, CLEANING_PROCESS, endpoint, [endpoint]) is None


def test_attribute_priority_requires_exact_owner_and_role(tmp_path):
    index, task, _, _ = cleaning_fixture(tmp_path, ["加水冲洗"])
    unit = next(u for u in index.by_id[task.record_id].source_units if u.text == "反应釜")
    owner = index.ir.anchor(unit.evidence_id, 0, len(unit.text))
    predicate = SlotSpec(iri="urn:equipmentID", label="设备编号")
    assert explicit_attribute_sources(index, task.record_id, predicate, [owner])
    other = next(r for r in index.records if any(u.text == "E-1" for u in r.source_units))
    assert not explicit_attribute_sources(index, other.record_id, predicate, [owner])
    assert not explicit_attribute_sources(index, task.record_id, predicate, [])
    assert not explicit_attribute_sources(
        index, task.record_id, predicate.model_copy(update={"label": "批数"}), [owner]
    )


def test_attribute_priority_fairness_and_restore():
    scheduler = FrontierScheduler(template_interleaving=True)
    for number in range(18):
        for kind in ("relationship", "property"):
            task = RecognitionTask.create(
                subject=SubjectRef(entity_id="local", revision=1, class_iri="urn:Equipment"),
                predicate_iri="urn:" + kind,
                predicate_kind=kind,
                record_id=f"r-{number}",
                phase=1,
                hop=1,
                dependency_hash="test",
            )
            scheduler.enqueue(task, template_priority=kind == "relationship")
            if kind == "property":
                scheduler.prioritize_source(task, [{"field_binding_id": "physical-field"}])
    kinds = []
    for _ in range(9):
        before = deepcopy(scheduler.snapshot())
        scheduler.peek_task()
        assert scheduler.snapshot() == before
        restored = FrontierScheduler.from_snapshot(before)
        assert restored.next_task() == scheduler.peek_task()
        kinds.append(scheduler.next_task().predicate_kind)
    assert kinds == ["property", "property", "relationship"] * 3


@pytest.mark.parametrize("partial", [True, False])
@pytest.mark.parametrize("multiple", [False, True])
def test_repair_gate_rejects_partial_method_despite_supported_model(
    tmp_path,
    monkeypatch,
    partial,
    multiple,
):
    from app.services.extraction.ontology_guided import model_adapter

    paragraphs = (
        ["加水冲洗", "丙酮循环", "氮气吹干"] if multiple else ["加水冲洗；丙酮循环；氮气吹干"]
    )
    _, task, predicate, context = cleaning_fixture(tmp_path, paragraphs)
    calls = []

    def respond(_client, *, user, **kwargs):
        request = json.loads(user)
        calls.append(request)
        method = next(
            f
            for f in request["fragments"]
            if f["text"] == (paragraphs[-1] if partial else paragraphs[0])
        )
        whole_quotes = [
            {"evidence_id": f["evidence_id"], "text": f["text"]}
            for f in request["fragments"]
            if f["text"] in paragraphs
        ]
        binding = next(
            b
            for b in request["field_bindings"]
            if any(r["evidence_id"] == method["evidence_id"] for r in b["target_value_refs"])
        )
        quote = {"evidence_id": method["evidence_id"], "text": method["text"]}
        if request["stage"] == "discovery":
            return {
                "proposals": [
                    {
                        "kind": "relationship",
                        "object_class_iri": CLEANING_PROCESS,
                        "object_label": "方法",
                        "object_quote": {**quote, "text": "氮气吹干"} if partial else quote,
                        "bridge_kind": "role_mapped_table",
                        "field_binding_id": binding["field_binding_id"],
                    }
                ]
            }
        candidate = request["candidates"][0]
        return {
            "verifications": [
                {
                    "candidate_id": candidate["candidate_id"],
                    "target_id": candidate["target_id"],
                    **{
                        f"{key}_verdict": "supported"
                        for key in (
                            "role",
                            "type",
                            "predicate",
                            "applicability",
                            "counterevidence",
                            "bridge",
                        )
                    },
                    "subject_binding": {"verdict": "supported", "support": [], "local_support": []},
                    "type_support": [quote],
                    "predicate_support": [quote],
                    "bridge_support": whole_quotes,
                    "field_role_support": [
                        {"evidence_id": r["evidence_id"]} for r in binding["label_refs"]
                    ],
                    "condition_support": [],
                    "counterevidence_support": [],
                    "reason": "受控正向核验",
                }
            ]
        }

    monkeypatch.setattr(model_adapter, "chat_with_schema", respond)
    menu = LocalMenu(
        menu_id="cleaning-menu",
        subject=task.subject,
        ontology_snapshot_id="test",
        relationships=[predicate],
    )
    result = EvidenceRepairAdapter(object(), model_identity="fixture").inspect(
        task, context, predicate, menu
    )
    assert len(calls) == 2 and calls[0]["whole_method_fields"]
    assert result.edges[0].policy_eligible is not partial
    if partial:
        assert result.edges[0].decision_status == "undetermined"
        assert any(
            "partial_cleaning_method" in issues
            for issues in context.protocol_state["gate_issues"].values()
        )
    else:
        assert result.nodes[0].label == "\n".join(paragraphs)
        assert len(result.nodes[0].evidence_refs) == len(paragraphs)
