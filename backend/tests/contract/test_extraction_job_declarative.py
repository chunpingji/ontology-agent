"""T012 — contract test: declaration-driven DB extraction job (US1).

Contract: ``contracts/extraction-job.md``. Written **before** the pipeline rewire
(T014–T019) — MUST FAIL until the declarative branch is wired.

Covers: declaration-mode job creation (202), unknown ``class_mapping_id`` (404),
non-source-entity binding (422), the unified candidate output (one per row with
declared identifiers + controlled-vocab normalization + resolvable provenance),
and drift → partial job with E6 ``health="drift"``.
"""

from __future__ import annotations

import json
import uuid
from io import BytesIO

from docx import Document

from tests.fixtures.feature014 import (
    ANALYST,
    APPROVAL_NUMBER,
    DRUG_PRODUCT,
    MANUFACTURED_BY,
    MANUFACTURER,
    RISK_LEVEL,
    add_property_binding,
    declare_db_binding,
)

CLASSES = "/api/ontology/classes"
JOBS = "/api/extraction/jobs"


def _declare_full_binding(client, source="DRUG_DB_DSN", target="drug_product"):
    mid = declare_db_binding(client, source=source, target=target)
    add_property_binding(
        client, mid, property_iri=APPROVAL_NUMBER, source_path="approval_no",
        is_identifier=True,
    )
    add_property_binding(
        client, mid, property_iri=RISK_LEVEL, source_path="risk",
        transform_type="controlled_vocab",
        transform_config={"map": {"高": "HighRisk", "中": "MediumRisk", "低": "LowRisk"}},
    )
    add_property_binding(
        client, mid, property_iri=MANUFACTURED_BY, property_kind="object",
        source_path="mfr_code", object_resolution="id_reference",
        target_class_iri=MANUFACTURER, target_id_path="mfr_code",
    )
    return mid


def _run_job(client, mid, source_type="database"):
    return client.post(
        JOBS, data={"source_type": source_type, "class_mapping_id": mid},
        headers=ANALYST,
    )


def _all_candidates(client, job_id) -> list[dict]:
    body = client.get(f"{JOBS}/{job_id}/candidates").json()
    return body["ungrouped"] + [c for g in body["groups"] for c in g["candidates"]]


def _declare_doc_pattern(client):
    profile = {
        "version": 1,
        "sources": [{
            "locator": "section_kv",
            "anchors": {"any_of": ["产品的基本性质"]},
        }],
        "identity": {"pattern": "[A-Z]{2,4}-[0-9]{3,5}"},
    }
    response = client.post(
        f"{CLASSES}/{DRUG_PRODUCT}/mappings",
        json={
            "mapping_type": "doc_pattern",
            "target": json.dumps(profile, ensure_ascii=False),
            "source_system": "cmc-word",
        },
        headers=ANALYST,
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("HRS-1597 CMC 报告", level=1)
    document.add_heading("产品的基本性质", level=2)
    document.add_paragraph("批准文号：HRS-1597")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_declarative_db_job_happy_path(drug_client, source_table):
    source_table()  # publishes DRUG_DB_DSN → drug_product (DEFAULT_ROWS)
    mid = _declare_full_binding(drug_client)

    r = _run_job(drug_client, mid)
    assert r.status_code == 202, r.text
    job_id = r.json()["id"]

    job = drug_client.get(f"{JOBS}/{job_id}").json()
    assert job["status"] in ("reviewing", "done")
    assert job["total_candidates"] >= 3

    instances = [c for c in _all_candidates(drug_client, job_id)
                 if c["candidate_kind"] == "instance"]
    assert len(instances) == 3
    approvals = {c["extracted_properties"].get(APPROVAL_NUMBER) for c in instances}
    assert "国药准字H20003007" in approvals

    # controlled vocab: the 高 row becomes HighRisk (normalized ontology term).
    high = next(c for c in instances
                if c["extracted_properties"].get(APPROVAL_NUMBER) == "国药准字H20003007")
    assert high["extracted_properties"].get(RISK_LEVEL) == "HighRisk"

    # provenance: every candidate carries a resolvable source_ref (SC-003).
    for c in instances:
        ref = json.loads(c["source_ref"])
        assert ref["system"] == "DRUG_DB_DSN"
        assert ref["entity"] == "drug_product"
        assert ref["record"]


def test_unknown_class_mapping_id_404(drug_client, source_table):
    source_table()
    r = _run_job(drug_client, str(uuid.uuid4()))
    assert r.status_code == 404, r.text


def test_non_source_entity_binding_422(drug_client):
    # a legacy bfo mapping is not a source-entity binding → cannot drive a job.
    bfo = drug_client.post(
        f"{CLASSES}/{DRUG_PRODUCT}/mappings",
        json={"mapping_type": "bfo", "target": "BFO_0000040"},
        headers=ANALYST,
    ).json()
    r = _run_job(drug_client, bfo["id"])
    assert r.status_code == 422, r.text


def test_drift_declared_column_absent_partial(drug_client, source_table):
    source_table()
    mid = declare_db_binding(drug_client)
    add_property_binding(
        drug_client, mid, property_iri=APPROVAL_NUMBER, source_path="approval_no",
        is_identifier=True,
    )
    # a declared column absent from the table → drift (skip binding, partial job).
    add_property_binding(
        drug_client, mid, property_iri=RISK_LEVEL, source_path="does_not_exist",
    )

    r = _run_job(drug_client, mid)
    assert r.status_code == 202, r.text
    job_id = r.json()["id"]

    job = drug_client.get(f"{JOBS}/{job_id}").json()
    assert job["status"] in ("reviewing", "done")  # completes, never crashes
    assert job["total_candidates"] >= 3

    mappings = drug_client.get(f"{CLASSES}/{DRUG_PRODUCT}/mappings").json()
    db_map = next(m for m in mappings if m["id"] == mid)
    assert db_map["health"] == "drift"


def test_doc_pattern_requires_docx_upload(drug_client):
    mid = _declare_doc_pattern(drug_client)

    missing = _run_job(drug_client, mid, source_type="word")
    assert missing.status_code == 422, missing.text

    unsupported = drug_client.post(
        JOBS,
        data={"source_type": "word", "class_mapping_id": mid},
        files={"file": ("report.txt", b"not a docx", "text/plain")},
        headers=ANALYST,
    )
    assert unsupported.status_code == 422, unsupported.text


def test_doc_pattern_upload_is_persisted_and_extracted(drug_client):
    mid = _declare_doc_pattern(drug_client)
    response = add_property_binding(
        drug_client,
        mid,
        property_iri=APPROVAL_NUMBER,
        source_path="批准文号",
        is_identifier=True,
    )
    assert response.status_code == 201, response.text

    created = drug_client.post(
        JOBS,
        data={"source_type": "word", "class_mapping_id": mid},
        files={"file": (
            "HRS-1597 CMC报告.docx",
            _docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )},
        headers=ANALYST,
    )
    assert created.status_code == 202, created.text
    job = created.json()
    assert job["source_filename"] == "HRS-1597 CMC报告.docx"
    assert job["document_path"].endswith(".docx")

    current = drug_client.get(f"{JOBS}/{job['id']}").json()
    assert current["status"] in ("reviewing", "done")
    candidates = _all_candidates(drug_client, job["id"])
    assert len(candidates) == 1
    assert candidates[0]["extracted_properties"][APPROVAL_NUMBER] == "HRS-1597"


def test_doc_pattern_corrupt_docx_degrades_without_unhandled_error(drug_client):
    mid = _declare_doc_pattern(drug_client)
    created = drug_client.post(
        JOBS,
        data={"source_type": "word", "class_mapping_id": mid},
        files={"file": (
            "corrupt.docx",
            b"not a zip package",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )},
        headers=ANALYST,
    )
    assert created.status_code == 202, created.text
    job = drug_client.get(f"{JOBS}/{created.json()['id']}").json()
    assert job["status"] == "degraded"
    assert "DOCX parse failed" in job["error_message"]
