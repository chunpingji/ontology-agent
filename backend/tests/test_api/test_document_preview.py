"""Reading Word content must not start or wait for semantic/model extraction."""

import json

import pytest
from docx import Document

from app.api import extraction
from app.config import settings
from app.models.evidence import EvidenceCandidateRecord, EvidenceJobState
from app.models.extraction import ExtractionJob
from app.services.extraction.docx_structure import PARSER_VERSION
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def preview_job(db, tmp_path, monkeypatch):
    document = Document()
    document.add_heading("产品 A", 1)
    document.add_paragraph("正文无需等待模型。")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格原文"
    source = tmp_path / "preview.docx"
    document.save(source)
    job = ExtractionJob(
        source_type="word", source_filename=source.name, document_path=str(source),
        source_config={"mode": "doc_repo_preview", "doc_class_iri": "urn:CMCReport"},
        status="reviewing",
    )
    db.add(job)
    db.commit()
    monkeypatch.setattr(extraction, "_annotation_cache_path", lambda _: tmp_path / "cache.json")

    def forbidden(*args, **kwargs):
        raise AssertionError("document GET must not call a model or run extraction")

    monkeypatch.setattr(
        "app.services.extraction.local_semantic_model.configured_generic_runner", forbidden,
    )
    monkeypatch.setattr("app.services.llm.local_client.get_local_llm", forbidden)
    monkeypatch.setattr(
        "app.services.extraction.word_tree_summarizer.summarize_word_tree", forbidden,
    )
    return job


@pytest.mark.parametrize("cache_state", ["missing", "stale", "corrupt", "refresh"])
def test_word_preview_is_model_free_and_does_not_overwrite_extraction(
    client, db, preview_job, analyst_headers, cache_state,
):
    cache = extraction._annotation_cache_path(preview_job.id)
    if cache_state != "missing":
        cached = {
            "_version": extraction._ANNOTATOR_VERSION if cache_state == "refresh" else 1,
            "_parser_version": PARSER_VERSION,
            "_summary_prompt_version": settings.word_tree_summary_prompt_version,
            "source_type": "word", "content": {"type": "doc", "content": []},
            "relationships": [{"legacy": True}],
        }
        cache.write_text("{broken" if cache_state == "corrupt" else json.dumps(cached))
    before = cache.read_bytes() if cache.exists() else None
    url = f"/api/extraction/jobs/{preview_job.id}/annotated-document"
    if cache_state == "refresh":
        url += "?refresh=true"

    for _ in range(2):
        response = client.get(url, headers=analyst_headers)
        assert response.status_code == 200, response.text
        payload = response.json()
        assert "正文无需等待模型。" in json.dumps(payload["content"], ensure_ascii=False)
        assert "表格原文" in json.dumps(payload["content"], ensure_ascii=False)
        assert payload["analysis"]["analysis_id"] == analyze_word_core(
            preview_job.document_path, source_filename=preview_job.source_filename,
        ).ir.analysis_id
        assert payload["content"]["analysis"] == payload["analysis"]
        assert payload["preview_only"] is True
        assert payload["completion"] == "incomplete"
        assert "evidence_run" not in payload
        assert payload["triples"] == payload["relationships"] == []
        assert payload["section_tree"]["children"][0]["heading"] == "产品 A"
    assert (cache.read_bytes() if cache.exists() else None) == before
    assert db.get(EvidenceJobState, preview_job.id) is None
    assert db.query(EvidenceCandidateRecord).count() == 0
    db.refresh(preview_job)
    assert preview_job.status == "reviewing"


def test_template_default_current_annotation_cache_still_returns_existing_results(
    client, db, preview_job, analyst_headers,
):
    preview_job.source_config = {
        **(preview_job.source_config or {}),
        "mode": "template_default",
    }
    db.commit()
    cached = {
        "_version": extraction._ANNOTATOR_VERSION,
        "_parser_version": PARSER_VERSION,
        "_summary_prompt_version": settings.word_tree_summary_prompt_version,
        "source_type": "word", "content": {"type": "doc", "content": []},
        "completion": "incomplete", "relationships": [{"existing_preview": True}],
    }
    extraction._annotation_cache_path(preview_job.id).write_text(json.dumps(cached))
    response = client.get(
        f"/api/extraction/jobs/{preview_job.id}/annotated-document", headers=analyst_headers,
    )
    assert response.status_code == 200
    assert response.json() == cached


def test_missing_source_without_cache_returns_404(client, db, preview_job, analyst_headers):
    preview_job.document_path += ".missing"
    db.commit()
    response = client.get(
        f"/api/extraction/jobs/{preview_job.id}/annotated-document", headers=analyst_headers,
    )
    assert response.status_code == 404
