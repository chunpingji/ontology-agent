"""Independent sample sentinel: never depends on uploaded production documents."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.table import _Cell

from app.models.extraction import AstTemplate
from app.services.reporting.batch_demo_layout import REFERENCE_ID
from app.services.reporting.batch_demo_operation_forms import load_forms
from app.services.reporting.batch_demo_sample_content import A14_OPERATIONS, A14_TABLES


def seed_layout(db, path):
    doc = Document()
    doc.styles['Normal'].font.size = Pt(11)
    for i in range(17):
        p = doc.add_paragraph('批生产记录' if i == 6 else 'SAMPLE-ONLY-SECRET')
        if i == 6:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(24)
    sizes = {0: (6, 4), 1: (2, 5), 3: (2, 6), 4: (2, 7), 9: (2, 4),
             15: (2, 5), 27: (2, 4), 28: (3, 5), 29: (2, 4)}
    sizes.update({i: (19, 10) for i in range(10, 15)})
    sizes.update({17: (1, 1), 20: (1, 1)})
    fields = load_forms()["fields"]
    operation_tables = {f["table"] for f in fields.values()}
    for f in fields.values():
        sizes[f["table"]] = (max(sizes.get(f["table"], (2, 5))[0], f["row"] + 1), 5)
    for i in range(max(sizes) + 1):
        if i in A14_TABLES:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph('6．操作记录')
        t = doc.add_table(*sizes.get(i, (2, 2)))
        t.style = 'Table Grid'
        t.autofit = False
        for col in t.columns:
            col.width = Cm(2)
        for row in t._tbl.tr_lst:
            for tc in row.tc_lst:
                _Cell(tc, t).text = 'SAMPLE-ONLY-SECRET'
        for j, c in enumerate(t.rows[0].cells):
            c.text = f'字段{j + 1}'
        if i == 0:
            for r in range(6):
                t.cell(r, 0 if r == 4 else 1).merge(t.cell(r, 1 if r == 4 else 2))
        if i == 28:
            t.cell(2, 0).merge(t.cell(2, 2))
        if i in A14_TABLES:
            for row in t._tbl.tr_lst:
                for tc in row.tc_lst:
                    _Cell(tc, t).text = ''
            if i < 15:
                t.cell(0, 0).merge(t.cell(0, 9)).text = '称量记录'
            elif i in (17, 20):
                t.cell(0, 0).text = 'TLC结果处理：样例要求，照片贴于留白处。'
        if i in operation_tables:
            for c, text in zip(t.rows[0].cells, [
                '原材料/操作', '参数', '记录', '操作人/日期', '复核人/日期',
            ]):
                c.text = text
            for f in fields.values():
                if f['table'] != i:
                    continue
                r = f['row']
                cells = t._tbl.tr_lst[r].tc_lst
                if f['split_calculation']:
                    _Cell(cells[1], t).merge(_Cell(cells[2], t)).text = f['sample_parameter']
                else:
                    _Cell(cells[1], t).text = f['sample_parameter']
                    _Cell(cells[2], t).text = f['sample_record']
            if i in A14_OPERATIONS:
                c = t.cell(1, 0).merge(t.cell(len(t.rows) - 1, 0))
                c.text = ''
                p = c.paragraphs[0]
                run = p.add_run(f'TEMPLATE-A14-{i}: ')
                run.font.name = 'Arial'
                run.font.size = Pt(12)
                run = p.add_run('参考模板操作，不能用源报告重排')
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run.italic = True
            else:
                t.rows[1].cells[0]._tc.get_or_add_tcPr().append(OxmlElement('w:vMerge'))
    for _ in range(5):
        doc.add_section()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin, s.right_margin = Cm(3.175), Cm(3.175)
        s.top_margin, s.bottom_margin = Cm(2.54), Cm(2.54)
    for i, title in enumerate(['物料信息表', 'SM5592-A14的制备', 'SM5592-A15的制备',
                               'HRS-5592粗品的制备', 'HRS-5592成品的制备（一般区）'], 1):
        h = doc.sections[i].header
        h.is_linked_to_previous = False
        t = h.add_table(rows=1, cols=4, width=Cm(17))
        for c, value in zip(t.rows[0].cells, ['工序', title, '生产地点', f'模板车间-{i}']):
            c.text = value
            c.paragraphs[0].runs[0].font.name = 'Arial'
            c.paragraphs[0].runs[0].font.size = Pt(13)
    doc.save(path)
    row = AstTemplate(id=REFERENCE_ID, name='批记录', version='v1', status='draft',
                      schema_version=2, schema_json={'sections': []}, sample_docx_path=str(path),
                      sample_content_json={'type': 'doc'})
    db.add(row)
    db.commit()
    return row


def xml_signature(element):
    """Compare content/formatting without irrelevant in-scope namespace declarations."""
    return (element.tag, dict(element.attrib), element.text,
            [xml_signature(child) for child in element])
