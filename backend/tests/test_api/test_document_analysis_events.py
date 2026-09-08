from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from app.api import document_analysis
from app.config import settings

ROOT_IRI = "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"


def _document(path: Path) -> bytes:
    document = Document()
    document.add_heading("事件测试", level=1)
    document.add_paragraph("仅验证持久事件重放，不触发模型。")
    document.save(path)
    return path.read_bytes()


def _queued_run(client, headers, tmp_path, monkeypatch) -> dict:
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *_args, **_kwargs: None)
    response = client.post(
        "/api/document-analysis/runs",
        headers=headers,
        files={"file": ("events.docx", _document(tmp_path / "events.docx"))},
        data={
            "root_class_iri": ROOT_IRI,
            "request_key": "events-run",
            "metadata_mode": "structure_only",
        },
    )
    assert response.status_code == 202, response.text
    return response.json()


def test_sse_replays_durable_numeric_sequence_without_creating_work(
    client, analyst_headers, tmp_path, monkeypatch
):
    accepted = _queued_run(client, analyst_headers, tmp_path, monkeypatch)
    run_id = accepted["recognition_run_id"]
    paused = client.post(
        f"/api/document-analysis/runs/{run_id}/pause",
        headers=analyst_headers,
        json={
            "expected_revision": accepted["run_revision"],
            "request_key": "pause-events",
            "reason": "形成终态以验证重放",
        },
    )
    assert paused.status_code == 202, paused.text

    replay = client.get(f"/api/document-analysis/runs/{run_id}/events", headers=analyst_headers)
    assert replay.status_code == 200, replay.text
    ids = [int(value) for value in re.findall(r"^id: (\d+)$", replay.text, re.MULTILINE)]
    assert ids == [1, 2]
    assert "event: run_state" in replay.text
    assert '"recognition_run_id":"' + run_id + '"' in replay.text

    resumed_cursor = client.get(
        f"/api/document-analysis/runs/{run_id}/events",
        headers={**analyst_headers, "Last-Event-ID": str(ids[-1])},
    )
    assert resumed_cursor.status_code == 200
    assert resumed_cursor.text == ""

    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers).json()
    assert status["status"] == "paused"
    assert status["event_head"] == 2
    assert status["progress"]["model_calls"] == 0


def test_sse_and_every_run_resource_hide_foreign_owner(
    client, analyst_headers, tmp_path, monkeypatch
):
    accepted = _queued_run(client, analyst_headers, tmp_path, monkeypatch)
    run_id = accepted["recognition_run_id"]
    foreign = {"X-User": "foreign", "X-Role": "senior_analyst"}
    for suffix in ("", "/metadata", "/graph", "/source", "/events"):
        response = client.get(f"/api/document-analysis/runs/{run_id}{suffix}", headers=foreign)
        assert response.status_code == 404, (suffix, response.text)
        assert response.json()["error"]["code"] == "RUN_NOT_FOUND"

    controlled = client.post(
        f"/api/document-analysis/runs/{run_id}/cancel",
        headers=foreign,
        json={
            "expected_revision": accepted["run_revision"],
            "request_key": "foreign-control",
            "reason": "不得成功",
        },
    )
    assert controlled.status_code == 404
