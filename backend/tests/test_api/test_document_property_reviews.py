"""Exact human review, immutable replay inputs and bounded repair HTTP contract."""

from copy import deepcopy
from datetime import UTC, datetime, timedelta
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import select

from app.api import document_analysis
from app.models.document_analysis import (
    DocumentAnalysisArtifact,
    DocumentRecognitionEvent,
    DocumentRunCandidate,
)
from app.models.document_analysis_review import DocumentPropertyRepair, DocumentPropertyReview
from app.services.document_analysis.reviews import (
    mark_repair_operation,
    pending_review_operations,
    repair_operations,
    review_operations,
)
from app.services.document_analysis.run_store import (
    DocumentAnalysisRunStore,
    FenceViolation,
    HeadConflict,
    content_hash,
)
from app.services.extraction.ontology_guided.contracts import (
    GraphNode,
    GraphProperty,
    RunProgress,
    SubjectRef,
    VersionedRef,
)
from app.services.extraction.ontology_guided.projection import project_graph
from app.services.extraction.ontology_guided.records import RecordIndex
from app.services.extraction.ontology_guided.scheduler import RecognitionTask
from app.services.extraction.word_analysis import analyze_word_core

ROOT = "urn:test:Report"
PROPERTY = "urn:test:appearance"


@pytest.fixture
def reviewed_run(db, tmp_path, monkeypatch):
    monkeypatch.setattr(document_analysis, "notify_document_analysis_dispatcher", lambda: True)
    word = Document()
    word.add_paragraph("外观为白色片剂。")
    word.add_paragraph("另一个主体的外观为红色，不可自动换绑。")
    path = tmp_path / "source.docx"
    word.save(path)
    ir = analyze_word_core(path).ir
    index = RecordIndex(ir)
    record = index.records[0]
    anchor = ir.anchor(record.source_units[0].evidence_id)
    store = DocumentAnalysisRunStore(db)
    run, _ = store.create_run(
        owner_id="analyst", request_key="run-review", filename=path.name,
        document_hash=ir.document_hash, root_class_iri=ROOT, root_class_label="报告",
        ontology_snapshot_hash="b" * 64,
    )
    token = store.claim(run.recognition_run_id, run.owner_id, actor="test", worker_id="test")
    node = GraphNode(entity_id="root", revision=1, class_iri=ROOT, class_label="报告",
                     label="报告", root=True, root_origin="user_specified")
    prop = GraphProperty(
        candidate_id="appearance", revision=1, subject_ref=VersionedRef(id="root", revision=1),
        predicate_iri=PROPERTY, predicate_label="外观", raw_value="白色片剂",
        decision_status="supported", structural_valid=True, model_supported=True,
        policy_eligible=True, proof_ref=VersionedRef(id="proof", revision=1),
        decision_refs=[VersionedRef(id="decision", revision=1)], evidence_refs=[anchor],
        value_evidence_refs=[anchor],
    )
    task = RecognitionTask.create(
        subject=SubjectRef(entity_id="root", revision=1, class_iri=ROOT, is_document_root=True),
        predicate_iri=PROPERTY, predicate_kind="property", record_id=record.record_id,
        phase=1, hop=0, dependency_hash="frozen-dependency",
    )
    for kind, candidate, identifier in (("entity", node, "root"),
                                         ("property", prop, "appearance")):
        store.put_candidate(run.recognition_run_id, run.owner_id, token,
                            candidate_id=identifier, revision=1, kind=kind,
                            payload=candidate.model_dump(mode="json"))
    graph = project_graph(
        recognition_run_id=str(run.recognition_run_id), run_revision=run.revision,
        event_head=run.event_head, metadata_snapshot_id="metadata-test",
        root_ref=VersionedRef(id="root", revision=1), nodes=[node], edges=[], properties=[prop],
        coverage=[], progress=RunProgress(), projection="all", artifact_status="ready",
    )
    checkpoint = {"run_fingerprint": "frozen-run", "task_outcomes": [{
        "task": task.model_dump(mode="json"),
        "outcome": {"properties": [prop.model_dump(mode="json")]},
    }]}
    payloads = {
        "structure": {"analysis": ir.model_dump(mode="json")},
        "graph": {"snapshot_id": "graph-test", "analysis_id": ir.analysis_id,
                  "ontology_snapshot_id": "ontology-test",
                  "generated_at": datetime.now(UTC).isoformat(),
                  "graph": graph.model_dump(mode="json"), "selection_registry": {}},
        "recognition_checkpoint": checkpoint,
    }
    for kind, payload in payloads.items():
        store.update_artifact(
            run.recognition_run_id, run.owner_id, token, artifact_kind=kind, expected_revision=0,
            artifact_id=f"{kind}-{uuid4()}", artifact_hash=content_hash(payload),
            status="ready", media_type="application/json", payload=payload,
        )
    store.update_stage(
        run.recognition_run_id, run.owner_id, token, expected_revision=run.revision,
        stage="finalize", execution_status="finished", run_fingerprint="frozen-run",
        analysis_id=ir.analysis_id, metadata_snapshot_id="metadata-test",
        graph_snapshot_id="graph-test",
    )
    db.commit()
    return run


def review_request(run, **changes):
    return {
        "request_key": "review-1", "expected_run_revision": run.revision,
        "graph_snapshot_id": "graph-test", "candidate_id": "appearance", "candidate_revision": 1,
        "expected_review_revision": 0, "decision": "rejected", "reason_code": "incorrect_value",
        "reason": "请复核原文中的属性值", **changes,
    }


def base(run):
    return f"/api/document-analysis/runs/{run.recognition_run_id}"


def post_review(client, run, headers, **changes):
    return client.post(base(run) + "/reviews", headers=headers, json=review_request(run, **changes))


def test_review_updates_projection_and_etag_without_rewriting_replay_inputs(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    candidates = [(row.candidate_id, row.revision, deepcopy(row.payload))
                  for row in db.scalars(select(DocumentRunCandidate))]
    artifacts = [(row.artifact_id, row.content_hash, deepcopy(row.payload))
                 for row in db.scalars(select(DocumentAnalysisArtifact))]
    before = client.get(base(run) + "/graph", headers=analyst_headers)
    assert len(before.json()["properties"]) == 1
    response = post_review(client, run, analyst_headers)
    assert response.status_code == 201, response.text
    reviewed = response.json()
    assert reviewed["run"]["run_revision"] > run.revision - 1
    graph = client.get(base(run) + "/graph", headers=analyst_headers)
    assert graph.json()["properties"] == []
    assert graph.headers["etag"] != before.headers["etag"]
    all_graph = client.get(base(run) + "/graph?projection=all_candidates", headers=analyst_headers)
    assert all_graph.json()["properties"][0]["independent_review"] == "rejected"
    assert all_graph.json()["graph_snapshot"]["snapshot_id"] == (
        reviewed["run"]["identities"]["graph_snapshot_id"]
    )
    rejected = client.get(base(run) + "/graph?projection=rejected", headers=analyst_headers)
    assert rejected.json()["properties"][0]["candidate_id"] == "appearance"
    assert candidates == [(row.candidate_id, row.revision, row.payload)
                          for row in db.scalars(select(DocumentRunCandidate))]
    assert artifacts == [(row.artifact_id, row.content_hash, row.payload)
                         for row in db.scalars(select(DocumentAnalysisArtifact))]
    assert review_operations(db, run)[0]["after_outcomes"] == 1


def test_acceptance_does_not_change_system_proof_qualification(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    graph_ref = DocumentAnalysisRunStore(db).get_artifact(
        run.recognition_run_id, run.owner_id, "graph",
    )
    artifact = db.get(DocumentAnalysisArtifact, graph_ref.artifact_id)
    body = deepcopy(artifact.payload)
    body["graph"]["properties"][0]["model_supported"] = False
    candidate = db.get(DocumentRunCandidate, (run.recognition_run_id, "appearance", 1))
    candidate.payload = deepcopy(body["graph"]["properties"][0])
    candidate.payload_hash = content_hash(candidate.payload)
    artifact.payload = body
    artifact.content_hash = graph_ref.content_hash = content_hash(body)
    db.commit()
    response = post_review(client, run, analyst_headers, decision="accepted", reason="")
    assert response.status_code == 201, response.text
    graph = client.get(base(run) + "/graph", headers=analyst_headers).json()
    assert graph["properties"] == []
    all_graph = client.get(
        base(run) + "/graph?projection=all_candidates", headers=analyst_headers,
    ).json()
    assert all_graph["properties"][0]["independent_review"] == "accepted"
    assert all_graph["properties"][0]["model_supported"] is False


@pytest.mark.parametrize("field,value", [
    ("candidate_revision", 2), ("graph_snapshot_id", "stale"),
    ("expected_review_revision", 1), ("expected_run_revision", 0),
])
def test_stale_review_targets_rejected(client, analyst_headers, reviewed_run, field, value):
    response = post_review(client, reviewed_run, analyst_headers, **{field: value})
    assert response.status_code == 409, response.text


@pytest.mark.parametrize("state", ["queued", "running", "pausing", "cancelled"])
def test_nonquiet_run_cannot_be_reviewed(client, db, analyst_headers, reviewed_run, state):
    reviewed_run.execution_status = state
    db.commit()
    assert post_review(client, reviewed_run, analyst_headers).status_code == 409
    listing = client.get(base(reviewed_run) + "/reviews", headers=analyst_headers).json()
    assert listing["can_review"] is False


def test_review_role_scope_and_property_only(client, db, analyst_headers, qa_headers, reviewed_run):
    assert client.get(base(reviewed_run) + "/reviews", headers=qa_headers).status_code == 404
    response = post_review(client, reviewed_run, {"X-User": "analyst", "X-Role": "operator"})
    assert response.status_code == 403
    response = post_review(client, reviewed_run, analyst_headers, candidate_id="root")
    assert response.status_code == 422
    assert post_review(client, reviewed_run, analyst_headers, reason=" ").status_code == 400
    reviewed_run.owner_id = "qa01"
    db.commit()
    assert post_review(client, reviewed_run, qa_headers).status_code == 201


def test_review_replay_is_exact_and_changed_key_conflicts(client, analyst_headers, reviewed_run):
    request = review_request(reviewed_run)
    first = client.post(base(reviewed_run) + "/reviews", headers=analyst_headers, json=request)
    second = client.post(base(reviewed_run) + "/reviews", headers=analyst_headers, json=request)
    assert second.status_code == 200 and second.json() == first.json()
    conflict = client.post(base(reviewed_run) + "/reviews", headers=analyst_headers,
                           json={**request, "reason": "不同内容"})
    assert conflict.status_code == 409


def test_displayed_payload_mismatch_cannot_be_reviewed_as_another_exact_candidate(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    ref = DocumentAnalysisRunStore(db).get_artifact(run.recognition_run_id, run.owner_id, "graph")
    artifact = db.get(DocumentAnalysisArtifact, ref.artifact_id)
    body = deepcopy(artifact.payload)
    body["graph"]["properties"][0]["raw_value"] = "另一个展示值"
    artifact.payload = body
    artifact.content_hash = ref.content_hash = content_hash(body)
    db.commit()
    assert post_review(client, run, analyst_headers).status_code == 409
    assert not list(db.scalars(select(DocumentPropertyReview)))


def test_repair_freezes_scope_requeues_finished_and_replays_without_dispatch(
    client, db, analyst_headers, reviewed_run, monkeypatch,
):
    run = reviewed_run
    wakeups = []
    monkeypatch.setattr(document_analysis, "notify_document_analysis_dispatcher",
                        lambda: wakeups.append(True) or True)
    review = post_review(client, run, analyst_headers).json()
    request = {"request_key": "repair-1", "expected_run_revision": review["run"]["run_revision"],
               "review_id": review["review"]["review_id"]}
    # Ordinary resume still refuses a finished run.
    resumed = client.post(base(run) + "/resume", headers=analyst_headers,
                          json={"request_key": "resume", "expected_revision": run.revision,
                                "reason": "retry"})
    assert resumed.status_code == 409
    response = client.post(base(run) + "/repairs", headers=analyst_headers, json=request)
    assert response.status_code == 202, response.text
    result = response.json()
    assert result["run"]["status"] == "queued"
    assert result["operation"]["status"] == "queued"
    assert result["operation"]["max_tasks"] == 16
    assert result["operation"]["max_model_calls"] == 32
    assert len(result["operation"]["record_ids"]) == 1
    original = pending_review_operations(db, run)[0]
    assert original["after_outcomes"] == 1
    assert original["target"]["subject_ref"] == {"id": "root", "revision": 1}
    assert original["target"]["reason"] == "请复核原文中的属性值"
    assert original["target"]["source_evidence_ids"]
    repeated = client.post(base(run) + "/repairs", headers=analyst_headers, json=request)
    assert repeated.json() == result and wakeups == [True]
    listing = client.get(base(run) + "/repairs", headers=analyst_headers)
    assert listing.json()["items"][0]["operation_id"] == result["operation"]["operation_id"]
    assert wakeups == [True]


def test_repair_only_accepts_latest_rejection(client, analyst_headers, reviewed_run):
    run = reviewed_run
    first = post_review(client, run, analyst_headers).json()
    second = post_review(
        client, run, analyst_headers, request_key="review-2", decision="accepted", reason="",
        expected_run_revision=first["run"]["run_revision"], expected_review_revision=1,
        graph_snapshot_id=first["run"]["identities"]["graph_snapshot_id"],
    )
    assert second.status_code == 201, second.text
    for entry in (first, second.json()):
        response = client.post(base(run) + "/repairs", headers=analyst_headers, json={
            "request_key": str(uuid4()),
            "expected_run_revision": second.json()["run"]["run_revision"],
            "review_id": entry["review"]["review_id"],
        })
        assert response.status_code == 409


def test_repair_worker_requires_fence_and_preserves_initial_replay_payload(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    review = post_review(client, run, analyst_headers).json()
    response = client.post(base(run) + "/repairs", headers=analyst_headers, json={
        "request_key": "repair-1", "expected_run_revision": review["run"]["run_revision"],
        "review_id": review["review"]["review_id"],
    })
    operation_id = response.json()["operation"]["operation_id"]
    initial = deepcopy(db.scalar(select(DocumentPropertyRepair)).payload)
    store = DocumentAnalysisRunStore(db)
    token = store.claim(run.recognition_run_id, run.owner_id, actor="test", worker_id="repair")
    db.commit()
    with pytest.raises(FenceViolation):
        mark_repair_operation(db, run.recognition_run_id, run.owner_id, "stale",
                              operation_id, "running")
    db.rollback()
    result = {"replacement_candidate_refs": [], "reason_code": "evidence_unresolved",
              "model_calls": 2, "tasks_attempted": 1}
    mark_repair_operation(db, run.recognition_run_id, run.owner_id, token, operation_id,
                          "unresolved", result)
    db.commit()
    assert pending_review_operations(db, run) == []
    assert len(repair_operations(db, run)) == 1
    assert db.scalar(select(DocumentPropertyRepair)).payload == initial
    with pytest.raises(HeadConflict):
        mark_repair_operation(db, run.recognition_run_id, run.owner_id, token, operation_id,
                              "completed", result)


def test_expired_run_and_cross_run_review_do_not_queue_repairs(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    run.expires_at = datetime.now(UTC) - timedelta(seconds=1)
    db.commit()
    assert post_review(client, run, analyst_headers).status_code == 409
    assert not list(db.scalars(select(DocumentPropertyReview)))


def queued_repair(client, run, headers):
    review = post_review(client, run, headers).json()
    response = client.post(base(run) + "/repairs", headers=headers, json={
        "request_key": "repair-1", "expected_run_revision": review["run"]["run_revision"],
        "review_id": review["review"]["review_id"],
    })
    assert response.status_code == 202, response.text
    return response.json()


@pytest.mark.parametrize("action", ["cancel", "delete"])
@pytest.mark.parametrize("repair_state", ["queued", "running"])
def test_lifecycle_stops_pending_repairs_and_fences_late_results(
    client, db, analyst_headers, reviewed_run, monkeypatch, action, repair_state,
):
    from app.services.document_analysis import retention

    monkeypatch.setattr(retention, "delete_run", lambda *_args, **_kwargs: None)
    run = reviewed_run
    queued = queued_repair(client, run, analyst_headers)
    operation = db.scalar(select(DocumentPropertyRepair))
    frozen_payload, frozen_receipt = deepcopy(operation.payload), deepcopy(operation.receipt)
    store = DocumentAnalysisRunStore(db)
    token = None
    if repair_state == "running":
        token = store.claim(run.recognition_run_id, run.owner_id, actor="test", worker_id="repair")
        mark_repair_operation(db, run.recognition_run_id, run.owner_id, token,
                              operation.operation_id, "running", {
                                  "model_calls": 2, "tasks_attempted": 1,
                                  "replacement_candidate_refs": [],
                              })
        db.commit()
    prior_cost = deepcopy(operation.result)
    expected_revision = run.revision
    request = {"request_key": f"{action}-repair", "expected_revision": expected_revision}

    def control():
        if action == "delete":
            return client.delete(base(run), headers=analyst_headers, params=request)
        return client.post(base(run) + "/cancel", headers=analyst_headers,
                           json={**request, "reason": "停止本轮局部识别"})

    response = control()
    assert response.status_code == 202, response.text
    db.refresh(operation)
    assert operation.status == "cancelled"
    assert operation.result["reason_code"] == (
        "run_cancelled" if action == "cancel" else "run_deleted"
    )
    assert all(operation.result[key] == value for key, value in prior_cost.items())
    assert operation.payload == frozen_payload and operation.receipt == frozen_receipt
    assert pending_review_operations(db, run) == []
    event = db.scalar(select(DocumentRecognitionEvent).where(
        DocumentRecognitionEvent.recognition_run_id == run.recognition_run_id,
        DocumentRecognitionEvent.event_key == f"control:{action}:{request['request_key']}",
    ))
    assert event.payload["cancelled_repair_operation_ids"] == [queued["operation"]["operation_id"]]
    updated_at = operation.updated_at
    replay = control()
    assert replay.status_code == 202 and replay.json() == response.json()
    db.refresh(operation)
    assert operation.updated_at == updated_at
    listing = client.get(base(run) + "/repairs", headers=analyst_headers)
    if action == "cancel":
        assert listing.status_code == 200
        assert listing.json()["items"][0]["status"] == "cancelled"
        assert listing.json()["can_repair"] is False
    else:
        assert listing.status_code == 410
    if token:
        with pytest.raises(FenceViolation):
            mark_repair_operation(db, run.recognition_run_id, run.owner_id, token,
                                  operation.operation_id, "completed", {})
        db.rollback()
        db.refresh(operation)
        assert operation.status == "cancelled"


@pytest.mark.parametrize("terminal", ["completed", "unresolved", "failed", "cancelled"])
def test_cancel_preserves_committed_repair_terminal_results(
    client, db, analyst_headers, reviewed_run, terminal,
):
    run = reviewed_run
    queued_repair(client, run, analyst_headers)
    operation = db.scalar(select(DocumentPropertyRepair))
    store = DocumentAnalysisRunStore(db)
    token = store.claim(run.recognition_run_id, run.owner_id, actor="test", worker_id="repair")
    result = {"reason_code": "existing_result", "model_calls": 4}
    mark_repair_operation(db, run.recognition_run_id, run.owner_id, token,
                          operation.operation_id, terminal, result)
    db.commit()
    db.refresh(operation)
    updated_at = operation.updated_at
    response = client.post(base(run) + "/cancel", headers=analyst_headers, json={
        "request_key": "cancel-existing", "expected_revision": run.revision, "reason": "停止",
    })
    assert response.status_code == 202, response.text
    db.refresh(operation)
    assert operation.status == terminal and operation.result == result
    assert operation.updated_at == updated_at


def test_stale_or_unauthorized_control_cannot_cancel_pending_repair(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    queued_repair(client, run, analyst_headers)
    operation = db.scalar(select(DocumentPropertyRepair))
    for headers, revision, expected_status in [
        (analyst_headers, run.revision - 1, 409),
        ({"X-User": "analyst", "X-Role": "operator"}, run.revision, 403),
        ({"X-User": "another", "X-Role": "senior_analyst"}, run.revision, 404),
    ]:
        response = client.post(base(run) + "/cancel", headers=headers, json={
            "request_key": "cancel-denied", "expected_revision": revision, "reason": "停止",
        })
        assert response.status_code == expected_status, response.text
        db.refresh(operation)
        assert operation.status == "queued"


def test_pause_resume_keeps_repair_available_for_later_execution(
    client, db, analyst_headers, reviewed_run,
):
    run = reviewed_run
    queued_repair(client, run, analyst_headers)
    operation = db.scalar(select(DocumentPropertyRepair))
    for action in ("pause", "resume"):
        response = client.post(base(run) + f"/{action}", headers=analyst_headers, json={
            "request_key": action, "expected_revision": run.revision, "reason": action,
        })
        assert response.status_code == 202, response.text
        db.refresh(operation)
        assert operation.status == "queued"
        assert len(pending_review_operations(db, run)) == 1
