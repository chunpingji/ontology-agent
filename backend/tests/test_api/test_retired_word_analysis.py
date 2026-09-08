"""021 T077/T080: old Word recognition is unreachable without harming shared flows."""

from __future__ import annotations

import io
import json
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document

from app.api import extraction
from app.config import settings
from app.models.evidence import EvidenceCandidateRecord, EvidenceCommit, EvidenceJobState
from app.models.extraction import (
    AnnotationExecution,
    AstTemplate,
    ExtractionCandidate,
    ExtractionConfig,
    ExtractionJob,
)
from app.services.extraction import docx_structure
from app.services.reporting.ast_template import load_default_template

ROOT_IRI = "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"
RETIRED_CODE = "WORD_RECOGNITION_RETIRED"


def _word_bytes(text: str = "本报告包含只读结构预览。") -> bytes:
    document = Document()
    document.add_heading("产品概述", level=1)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _word_config(db) -> ExtractionConfig:
    config = ExtractionConfig(
        name="retired-word-config",
        target_class_iri=ROOT_IRI,
        source_type="word",
        column_mapping={"名称": "urn:test:name"},
    )
    db.add(config)
    db.commit()
    return config


def _legacy_word_job(db, tmp_path: Path) -> ExtractionJob:
    source = tmp_path / "legacy.docx"
    source.write_bytes(_word_bytes("旧作业不得重新进入识别。"))
    job = ExtractionJob(
        source_type="word",
        source_filename=source.name,
        source_config={"mode": "auto", "doc_class_iri": ROOT_IRI},
        document_path=str(source),
        status="pending",
    )
    db.add(job)
    db.commit()
    return job


def test_old_sync_and_word_job_creation_routes_are_closed(client, db, analyst_headers, tmp_path):
    config = _word_config(db)
    raw = _word_bytes()

    old_sync = client.post(
        "/api/document-analysis/word",
        headers=analyst_headers,
        files={"file": ("old.docx", raw)},
    )
    assert old_sync.status_code == 404

    before = db.query(ExtractionJob).count()
    direct = client.post(
        "/api/extraction/jobs",
        headers=analyst_headers,
        data={"source_type": "word", "config_id": str(config.id)},
        files={"file": ("old.docx", raw)},
    )
    automatic = client.post(
        "/api/extraction/jobs/auto",
        headers=analyst_headers,
        data={"source_type": "word", "doc_class_iri": ROOT_IRI},
        files={"file": ("old.docx", raw)},
    )
    disguised = client.post(
        "/api/extraction/jobs",
        headers=analyst_headers,
        data={"source_type": "excel", "config_id": str(config.id)},
        files={"file": ("old.docx", raw)},
    )

    for response in (direct, automatic, disguised):
        assert response.status_code == 410, response.text
        assert RETIRED_CODE in response.text
    assert db.query(ExtractionJob).count() == before
    assert db.query(AnnotationExecution).count() == 0


def test_legacy_word_result_start_progress_and_recovery_are_closed(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    job = _legacy_word_job(db, tmp_path)
    cache = tmp_path / "legacy-cache.json"
    cache.write_text('{"source_type":"word","relationships":[{"legacy":true}]}')
    monkeypatch.setattr(extraction, "_annotation_cache_path", lambda _job_id: cache)
    dictionary_candidate = ExtractionCandidate(
        job_id=job.id,
        target_class_iri="urn:test:Legacy",
        extracted_properties={"secret": "must-not-be-returned"},
        review_status="pending",
    )
    commit = EvidenceCommit(
        id="retired-word-commit",
        job_id=job.id,
        idempotency_key="retired-word-commit",
        content_hash="0" * 64,
        manifest={"items": []},
        status="failed",
        actor="legacy-worker",
    )
    db.add_all(
        [
            dictionary_candidate,
            commit,
            EvidenceJobState(
                job_id=job.id,
                revision=9,
                extraction_run={
                    "completion": "incomplete",
                    "checkpoint": {"legacy": "must-not-be-returned"},
                },
            ),
        ]
    )
    db.commit()

    requests = [
        client.get(
            f"/api/extraction/jobs/{job.id}/annotated-document",
            headers=analyst_headers,
        ),
        client.get(f"/api/extraction/jobs/{job.id}/progress", headers=analyst_headers),
        client.get(f"/api/extraction/jobs/{job.id}/candidates", headers=analyst_headers),
        client.get(f"/api/extraction/jobs/{job.id}/evidence", headers=analyst_headers),
        client.get(f"/api/extraction/evidence/commits/{commit.id}", headers=analyst_headers),
        client.get(f"/api/extraction/jobs/{job.id}/evidence/coverage", headers=analyst_headers),
        client.get(f"/api/extraction/jobs/{job.id}/ast-coverage", headers=analyst_headers),
        client.get(
            f"/api/extraction/jobs/{job.id}/pde-conflict/decision",
            headers=analyst_headers,
        ),
        client.post(f"/api/extraction/jobs/{job.id}/start", headers=analyst_headers),
        client.post(f"/api/extraction/jobs/{job.id}/annotation/pause", headers=analyst_headers),
        client.post(f"/api/extraction/jobs/{job.id}/annotation/resume", headers=analyst_headers),
        client.post(f"/api/extraction/jobs/{job.id}/annotation/rerun", headers=analyst_headers),
        client.post(f"/api/extraction/jobs/{job.id}/evidence/extract", headers=analyst_headers),
        client.post(
            f"/api/extraction/evidence/commits/{commit.id}/retry",
            headers=analyst_headers,
        ),
        client.post(f"/api/extraction/jobs/{job.id}/risk-report", headers=analyst_headers),
        client.post(
            "/api/report-previews",
            headers=analyst_headers,
            json={
                "mode": "data",
                "idempotency_key": "retired-word-report-source",
                "source_bindings": {"doc": {"job_id": str(job.id)}},
            },
        ),
        client.post(
            "/api/ast-templates/suggest-slots",
            headers=analyst_headers,
            json={"job_id": str(job.id)},
        ),
        client.put(
            f"/api/extraction/candidates/{dictionary_candidate.id}/review",
            headers=analyst_headers,
            json={"status": "rejected"},
        ),
    ]

    for response in requests:
        assert response.status_code == 410, response.text
        assert RETIRED_CODE in response.text
    db.refresh(job)
    assert job.status == "pending"
    assert db.query(AnnotationExecution).count() == 0
    assert db.query(EvidenceCandidateRecord).count() == 0
    db.refresh(commit)
    assert commit.status == "failed"
    assert commit.attempts == 0
    state = db.get(EvidenceJobState, job.id)
    assert state.revision == 9
    assert state.extraction_run["checkpoint"]["legacy"] == "must-not-be-returned"
    db.refresh(dictionary_candidate)
    assert dictionary_candidate.review_status == "pending"


@pytest.mark.asyncio
async def test_late_legacy_worker_delivery_is_a_noop(db, fake_engine, tmp_path, monkeypatch):
    job = _legacy_word_job(db, tmp_path)
    called = []
    monkeypatch.setattr(
        extraction,
        "_annotation_worker",
        lambda *_args, **_kwargs: called.append(True),
    )

    await extraction._precompute_annotation_bg(
        job.id,
        fake_engine,
        db.get_bind(),
        run_id="already-queued-before-retirement",
    )
    await extraction._run_pipeline_bg(job.id, None, job.document_path, fake_engine, db)

    assert called == []
    assert db.query(AnnotationExecution).count() == 0
    assert db.query(EvidenceCandidateRecord).count() == 0


def test_doc_repo_word_upload_is_preview_only(client, db, analyst_headers, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    original_enqueue = extraction._enqueue_annotation

    def forbidden(*_args, **_kwargs):
        raise AssertionError("doc_repo preview must not enqueue the retired runner")

    monkeypatch.setattr(extraction, "_enqueue_annotation", forbidden)
    response = client.post(
        "/api/extraction/jobs/auto",
        headers=analyst_headers,
        data={
            "source_type": "word",
            "doc_class_iri": ROOT_IRI,
            "purpose": "doc_repo_preview",
        },
        files={"file": ("repository.docx", _word_bytes())},
    )

    assert response.status_code == 202, response.text
    job = db.get(ExtractionJob, UUID(response.json()["id"]))
    assert job.status == "completed"
    assert job.source_config == {
        "mode": "doc_repo_preview",
        "doc_class_iri": ROOT_IRI,
    }
    assert db.get(AnnotationExecution, job.id) is None
    assert db.get(EvidenceJobState, job.id) is None
    assert db.query(EvidenceCandidateRecord).count() == 0
    monkeypatch.setattr(extraction, "_enqueue_annotation", original_enqueue)

    cache_path = extraction._annotation_cache_path(job.id)
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(
        json.dumps(
            {
                "_version": extraction._ANNOTATOR_VERSION,
                "_parser_version": docx_structure.PARSER_VERSION,
                "_summary_prompt_version": settings.word_tree_summary_prompt_version,
                "source_type": "word",
                "preview_only": False,
                "relationships": [{"legacy-secret": True}],
            }
        ),
        encoding="utf-8",
    )

    preview = client.get(
        f"/api/extraction/jobs/{job.id}/annotated-document", headers=analyst_headers
    )
    assert preview.status_code == 200, preview.text
    payload = preview.json()
    assert payload["preview_only"] is True
    assert payload["relationships"] == []
    assert "evidence_run" not in payload
    assert "legacy-secret" not in preview.text

    for suffix in ("annotation/resume", "annotation/rerun", "evidence/extract"):
        blocked = client.post(f"/api/extraction/jobs/{job.id}/{suffix}", headers=analyst_headers)
        assert blocked.status_code == 410, blocked.text


def test_template_default_word_keeps_preview_and_explicit_shared_runner(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    source = tmp_path / "template.docx"
    source.write_bytes(_word_bytes("模板默认源仍可读取结构。"))
    template = AstTemplate(
        name="retirement-template",
        version="v1",
        iri_pattern=ROOT_IRI,
        schema_json=load_default_template().model_dump(),
        status="draft",
    )
    db.add(template)
    db.flush()
    job = ExtractionJob(
        source_type="word",
        source_filename=source.name,
        source_config={
            "mode": "template_default",
            "template_id": str(template.id),
            "doc_class_iri": ROOT_IRI,
        },
        document_path=str(source),
        status="reviewing",
    )
    db.add(job)
    db.commit()
    monkeypatch.setattr(
        extraction, "_annotation_cache_path", lambda _job_id: tmp_path / "template-cache.json"
    )
    monkeypatch.setattr(
        extraction,
        "_annotation_checkpoint_path",
        lambda _job_id: tmp_path / "template-checkpoint.json",
    )

    preview = client.get(
        f"/api/extraction/jobs/{job.id}/annotated-document", headers=analyst_headers
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["preview_only"] is True

    receipt = extraction._claim_annotation(
        job.id,
        db,
        mode="start",
        actor="template-maintainer",
    )
    assert receipt["run_id"]
    assert db.get(AnnotationExecution, job.id).status == "queued"


def test_new_run_rejects_legacy_coordinates(client, analyst_headers, tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    response = client.post(
        "/api/document-analysis/runs",
        headers=analyst_headers,
        data={
            "root_class_iri": ROOT_IRI,
            "request_key": "no-legacy-coordinates",
            "metadata_mode": "structure_only",
            "legacy_runner": "generic",
            "checkpoint": "old-checkpoint",
            "candidate_id": "old-candidate",
        },
        files={"file": ("new.docx", _word_bytes())},
    )

    assert response.status_code == 400, response.text
    assert response.json()["error"]["code"] == "INVALID_REQUEST"
    assert "candidate_id" in response.json()["error"]["message"]
