from docx import Document

from app.config import settings
from app.services.extraction.document_ir import DocumentIR
from app.services.extraction.word_analysis import analyze_word_core


def test_analysis_and_template_entrypoints_share_identity_and_offline_skeleton(
    client, analyst_headers, tmp_path, monkeypatch,
):
    monkeypatch.setattr(settings, "llm_suggest_slots_enabled", False)
    monkeypatch.setattr(
        settings, "document_analysis_storage_dir", tmp_path / "document-analysis"
    )
    doc = Document()
    doc.add_heading("产品 A", 1)
    doc.add_paragraph("规格：250 mg𠀀")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "内层原文"
    path = tmp_path / "evidence.docx"
    doc.save(path)
    files = {"file": (path.name, path.read_bytes())}
    created = client.post(
        "/api/document-analysis/runs",
        files=files,
        data={
            "root_class_iri": (
                "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"
            ),
            "request_key": "word-evidence-shared-ir",
            "metadata_mode": "structure_only",
        },
        headers=analyst_headers,
    )
    assert created.status_code == 202, created.text
    analysis = client.get(
        f"/api/document-analysis/runs/{created.json()['recognition_run_id']}/metadata",
        headers=analyst_headers,
    )
    sample = client.post("/api/ast-templates/parse-sample", files=files, headers=analyst_headers)
    assert analysis.status_code == sample.status_code == 200
    left, right = analysis.json()["content"]["analysis"], sample.json()["analysis"]
    assert left["structure_hash"] == right["structure_hash"]
    assert left["structure_hash"] == analyze_word_core(path).ir.structure_hash
    assert right["document_role"] == "template_sample"
    ir = DocumentIR.model_validate(right)
    content = sample.json()["content_json"]
    assert content["analysis"]["analysis_id"] == ir.analysis_id
    suggestion = client.post("/api/ast-templates/suggest-slots", headers=analyst_headers, json={
        "sample_content_json": content,
    })
    assert suggestion.status_code == 200, suggestion.text
    result = suggestion.json()
    assert result["completion"] == "incomplete"
    assert result["degraded"] is False
    candidate = result["sections"][0]["groups"][0]["candidates"][0]
    assert candidate["label"] == "规格"
    assert ir.resolve(candidate["origin"]["value_anchor"]) == "250 mg𠀀"


def test_preview_preserves_every_recursive_source_unit(tmp_path):
    from app.services.extraction.document_annotator import parse_word_to_tiptap

    doc = Document()
    cell = doc.add_table(rows=2, cols=2).cell(0, 0)
    cell.text = " 第一段 "
    cell.add_paragraph("第二段𠀀")
    deep = cell.add_table(rows=1, cols=1).cell(0, 0)
    deep.text = "第二层"
    deep.add_table(rows=1, cols=1).cell(0, 0).text = "第三层"
    path = tmp_path / "nested.docx"
    doc.save(path)
    content = parse_word_to_tiptap(path)
    ir = DocumentIR.model_validate(content["analysis"])
    nodes = {}

    def walk(node):
        identity = node.get("attrs", {}).get("evidenceId")
        if identity:
            nodes[identity] = node
        for child in node.get("content", []):
            walk(child)

    walk(content)
    for unit in ir.evidence_units:
        node = nodes[unit.evidence_id]
        assert "".join(child.get("text", "") for child in node.get("content", [])) == unit.text
