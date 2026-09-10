"""Static demo uses the visible graph and never runs recognition."""

import hashlib
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.table import _Cell
from sqlalchemy import func, select

from app.models.document_analysis import DocumentAnalysisRun
from app.models.entity_shadow import EntityShadow
from app.models.extraction import AnnotationExecution, AstTemplate, ExtractionJob, GeneratedReport
from app.services.reporting import batch_demo
from app.services.reporting.batch_demo_template import specialize
from app.services.reporting.report_run_service import schema_hash
from app.services.reporting.template_v2 import TemplateV2
from tests.test_api.batch_layout_fixture import seed_layout


@pytest.fixture
def source(db, tmp_path, monkeypatch):
    graph, template = batch_demo.definitions()
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Isolated source; hash explicitly replaced only in this fixture.")
    doc.save(path)
    graph["source_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    monkeypatch.setattr(batch_demo, "definitions", lambda: (deepcopy(graph), deepcopy(template)))
    monkeypatch.setattr(batch_demo, "ARTIFACT_ROOT", tmp_path / "reports")
    iri = "urn:test:batch-demo-source"
    job = ExtractionJob(source_type="word", source_filename=graph["source_filename"],
                        document_path=str(path), status="paused", source_config={
                            "doc_ref": iri, "doc_class_iri": template["root_class_iri"],
                        })
    db.add(job)
    db.flush()
    entity = EntityShadow(iri=iri, module="document", class_iri=template["root_class_iri"],
                          label_zh=graph["source_filename"],
                          properties_json={"job_id": str(job.id), "_version": 1})
    db.add(entity)
    schema = TemplateV2(schema_version=2, template_family_id="demo-test-family",
                        template_revision_id=str(batch_demo.TEMPLATE_ID), revision_no=2)
    row = AstTemplate(id=batch_demo.TEMPLATE_ID, name="批记录", version="v2.2",
                      schema_json=schema.model_dump(mode="json"), schema_version=2,
                      iri_pattern=template["root_class_iri"], status="draft", is_default=False)
    db.add(row)
    db.commit()
    specialize(db, expected_hash=schema_hash(row.schema_json), actor="analyst",
               document_iri=iri, apply=True)
    seed_layout(db, tmp_path / "layout.docx")
    return graph, template, entity, job, path


def get(client, headers, source):
    return client.get("/api/reports/batch-demo", params={"document_iri": source[2].iri},
                      headers=headers)


def post(client, headers, source, **changes):
    data = get(client, headers, source).json()
    body = {"graph_hash": data["graph_hash"], "template_hash": data["template_hash"],
            "request_key": "demo-request", **changes}
    return client.post("/api/reports/batch-demo", params={"document_iri": source[2].iri},
                       headers=headers, json=body)


def test_graph_to_saved_word_without_models(client, db, analyst_headers, source, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Demo must not invoke models or recognition")

    monkeypatch.setattr("app.services.llm.local_client.get_local_llm", forbidden)
    monkeypatch.setattr(
        "app.services.extraction.local_semantic_model.configured_generic_runner", forbidden,
    )
    monkeypatch.setattr(
        "app.services.document_analysis.application.DocumentAnalysisApplication.create_run",
        forbidden,
    )
    before = source[4].read_bytes()
    loaded = get(client, analyst_headers, source)
    assert loaded.status_code == 200
    assert loaded.headers["cache-control"] == "private, no-store"
    data = loaded.json()
    assert data["validation"]["passed"]
    checks = {c["id"]: c["count"] for c in data["validation"]["checks"]}
    assert checks["steps"] == 4 and checks["intermediates"] == 3 and checks["final"] == 1
    assert db.scalar(select(func.count()).select_from(GeneratedReport)) == 0
    result = post(client, analyst_headers, source)
    assert result.status_code == 201, result.text
    body = result.json()
    assert len(body["stages"]) == 5
    assert all(s["status"] == "completed" for s in body["stages"])
    assert get(client, analyst_headers, source).json()["latest_report"]["id"] == body["id"]
    saved = db.scalar(select(GeneratedReport))
    assert saved.rules_summary["graph"] == data["graph"]
    assert batch_demo.digest({"source": saved.rules_summary["source"],
                              "graph": saved.rules_summary["graph"]}) == data["graph_hash"]
    assert saved.report_run_id is None and saved.rules_fired_count == 0
    downloaded = client.get(
        f"/api/extraction/jobs/{body['job_id']}/reports/{body['id']}/download",
        headers=analyst_headers,
    )
    assert downloaded.status_code == 200
    assert "%E6%89%B9%E8%AE%B0%E5%BD%95" in downloaded.headers["content-disposition"]
    word = Document(BytesIO(downloaded.content))
    text = "\n".join(p.text for p in word.paragraphs)
    tables = "\n".join(
        _Cell(c, t).text for t in word.tables for r in t._tbl.tr_lst for c in r.tc_lst
    )
    assert "演示草稿" in text and "QA审核/日期：" in text
    assert "待填写" not in text + tables
    assert "SM5592-A14" in tables and "SM5592-A15" in tables
    assert "RE64615/RE64215" in tables and "95~105℃" in tables
    ops = [t for t in word.tables if t.cell(0, 0).text == "原材料/操作"]
    assert len(ops) == 13  # Ten template A14 tables followed by the three source stages.
    assert 'TEMPLATE-A14-15' in tables
    assert [sum(bool(_Cell(r.tc_lst[0], t).text) for r in t._tbl.tr_lst[1:])
            for t in ops[-3:]] == [12, 15, 9]
    assert db.scalar(select(func.count()).select_from(DocumentAnalysisRun)) == 0
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 0
    assert source[4].read_bytes() == before
    db.refresh(source[3])
    assert source[3].status == "paused"


def test_replay_and_explicit_new_result(client, db, analyst_headers, source):
    first = post(client, analyst_headers, source).json()
    second = post(client, analyst_headers, source).json()
    assert first["id"] == second["id"]
    assert len(list(batch_demo.ARTIFACT_ROOT.glob("*.docx"))) == 1
    third = post(client, analyst_headers, source, request_key="new").json()
    assert third["id"] != first["id"]
    assert db.scalar(select(func.count()).select_from(GeneratedReport)) == 2


@pytest.mark.parametrize("changed", ["filename", "class"])
def test_unrelated_documents(client, db, analyst_headers, source, changed):
    if changed == "filename":
        source[2].label_zh = "another HRS-5592.docx"
    else:
        source[2].class_iri = "urn:OtherClass"
    db.commit()
    assert get(client, analyst_headers, source).json() == {"available": False}


@pytest.mark.parametrize("changed", ["hash", "missing", "ownership"])
def test_stale_source_blocks_demo(client, db, analyst_headers, source, changed):
    if changed == "hash":
        source[4].write_bytes(b"changed")
    elif changed == "missing":
        source[4].unlink()
    else:
        source[3].source_config = {"doc_ref": "urn:another-document"}
        db.commit()
    assert get(client, analyst_headers, source).status_code == 409
    assert db.scalar(select(func.count()).select_from(GeneratedReport)) == 0


def test_missing_multihop_and_operations(client, analyst_headers, source):
    route = next(r for r in source[0]["relationships"]
                 if r["predicate_iri"].endswith("/hasSynthesisRoute"))
    route["sub_relationships"][0]["operations"] = []
    route["sub_relationships"][1]["sub_relationships"] = []
    assert not get(client, analyst_headers, source).json()["validation"]["passed"]
    assert post(client, analyst_headers, source).status_code == 422


def test_snapshot_and_template_conflicts(client, analyst_headers, source):
    assert post(client, analyst_headers, source, graph_hash="0" * 64).status_code == 409
    assert post(client, analyst_headers, source, template_hash="0" * 64).status_code == 409


def test_history_does_not_claim_to_match_a_changed_template(client, analyst_headers, source):
    created = post(client, analyst_headers, source).json()
    source[1]["manual_fields"].append("新修订填写项")
    assert get(client, analyst_headers, source).json()["latest_report"] is None
    # The frozen historical report remains downloadable by its explicit ID.
    assert client.get(
        f"/api/extraction/jobs/{created['job_id']}/reports/{created['id']}/download",
        headers=analyst_headers,
    ).status_code == 200


def test_roles_and_report_owner(client, analyst_headers, operator_headers, source):
    assert get(client, operator_headers, source).status_code == 200
    assert post(client, operator_headers, source).status_code == 403
    created = post(client, analyst_headers, source).json()
    other = {**analyst_headers, "X-User": "someone-else"}
    assert get(client, other, source).json()["latest_report"] is None
    base = f"/api/extraction/jobs/{created['job_id']}/reports"
    assert client.get(base, headers=other).json() == []
    for suffix in ["", "/download"]:
        assert client.get(f"{base}/{created['id']}{suffix}", headers=other).status_code == 404
    assert client.delete(f"{base}/{created['id']}", headers=other).status_code == 404
    assert client.get(f"/api/extraction/jobs/{created['job_id']}/risk-report",
                      headers=analyst_headers).status_code == 404


def test_save_failure_cleans_new_file_only(db, source, monkeypatch):
    from app.api.batch_demo import GenerateBatchDemo

    data, _ = batch_demo.context(db, source[2].iri, "analyst")
    batch_demo.ARTIFACT_ROOT.mkdir()
    existing = batch_demo.ARTIFACT_ROOT / "existing.docx"
    existing.write_bytes(b"keep")

    def fail(*args, **kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(batch_demo.audit, "append", fail)
    with pytest.raises(RuntimeError, match="audit unavailable"):
        batch_demo.generate(db, source[2].iri, "analyst", GenerateBatchDemo(
            graph_hash=data["graph_hash"], template_hash=data["template_hash"], request_key="a",
        ))
    assert existing.read_bytes() == b"keep"
    assert list(batch_demo.ARTIFACT_ROOT.iterdir()) == [existing]
    assert db.scalar(select(func.count()).select_from(GeneratedReport)) == 0


def test_all_demo_ontology_iris_exist():
    from rdflib import RDF, Graph, URIRef
    from rdflib.namespace import OWL

    graph, template = batch_demo.definitions()
    ttl = Graph()
    for path in (Path(__file__).resolve().parents[3] / "ontology/slpra").glob("*.ttl"):
        ttl.parse(path, format="turtle")

    def visit(nodes):
        for n in nodes:
            assert (URIRef(n["object_class_iri"]), RDF.type, OWL.Class) in ttl
            assert (URIRef(n["predicate_iri"]), RDF.type, OWL.ObjectProperty) in ttl
            for p in n["object_data_properties"]:
                if p["iri"]:
                    assert (URIRef(p["iri"]), RDF.type, OWL.DatatypeProperty) in ttl
            visit(n["sub_relationships"])

    visit(graph["relationships"])
    assert batch_demo.validate(graph, template)["passed"]
