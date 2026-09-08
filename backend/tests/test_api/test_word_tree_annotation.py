"""018 — annotated-document orchestration and compatibility contract."""

from __future__ import annotations

from types import SimpleNamespace

from docx import Document

from app.api import extraction
from app.config import settings
from app.services.extraction import (
    document_annotator,
    docx_structure,
    relation_extractor,
    word_analysis,
    word_tree_summarizer,
)


def test_compute_annotation_reuses_structure_and_returns_tree(tmp_path, monkeypatch):
    doc = Document()
    doc.add_heading("章节", level=1)
    doc.add_paragraph("正文。")
    path = tmp_path / "api-tree.docx"
    doc.save(path)

    job = SimpleNamespace(
        source_type="word",
        document_path=str(path),
        source_filename="api-tree.docx",
        source_config={},
    )
    real_parse = docx_structure.parse_docx_structure
    parse_calls = []
    captured = {}

    def counting_parse(*args, **kwargs):
        parse_calls.append((args, kwargs))
        return real_parse(*args, **kwargs)

    def fake_annotate(_path, *args, **kwargs):
        captured["annotator"] = kwargs["structure"]
        return {"type": "doc", "content": []}, [], [], None

    def fake_relationships(_engine, _path, _triples, **kwargs):
        captured["relations"] = kwargs["structure"]
        return {"doc_class": None, "relationships": []}

    monkeypatch.setattr(word_analysis, "parse_docx_structure", counting_parse)
    monkeypatch.setattr(document_annotator, "annotate_word", fake_annotate)
    monkeypatch.setattr(relation_extractor, "extract_relationships", fake_relationships)

    def broken_summary(*_args, **_kwargs):
        raise RuntimeError("unexpected summary failure")

    monkeypatch.setattr(word_tree_summarizer, "summarize_word_tree", broken_summary)

    result = extraction._compute_annotation(job, engine=object(), preview_only=True)

    assert len(parse_calls) == 1
    assert captured["annotator"] is not None
    assert "relations" not in captured
    assert result["_version"] == extraction._ANNOTATOR_VERSION
    assert result["section_tree"]["children"][0]["heading"] == "章节"
    assert result["preview_only"] is True
    assert result["completion"] == "incomplete"
    assert result["analysis"]["evidence_units"]
    assert result["pagination"]["mode"] == "single_page_fallback"
    assert result["relationships"] == []


def test_word_annotation_cache_identity_includes_parser_and_prompt(monkeypatch):
    payload = {
        "_version": extraction._ANNOTATOR_VERSION,
        "source_type": "word",
        "_parser_version": docx_structure.PARSER_VERSION,
        "_summary_prompt_version": settings.word_tree_summary_prompt_version,
    }
    assert extraction._annotation_cache_is_current(payload)

    stale_parser = {**payload, "_parser_version": docx_structure.PARSER_VERSION - 1}
    assert not extraction._annotation_cache_is_current(stale_parser)

    monkeypatch.setattr(
        settings,
        "word_tree_summary_prompt_version",
        "word-tree-summary-v-next",
    )
    assert not extraction._annotation_cache_is_current(payload)


def test_non_word_cache_only_uses_annotator_version():
    payload = {
        "_version": extraction._ANNOTATOR_VERSION,
        "source_type": "excel",
    }
    assert extraction._annotation_cache_is_current(payload)
