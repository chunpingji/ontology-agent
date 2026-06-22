"""US1 契约/集成测试：抽取作业接线与进度（FR-001~007, SC-001）。"""

from __future__ import annotations

import io

import openpyxl

from app.services.extraction.llm_extractor import build_extraction_prompt
from app.services.extraction.vocabulary import CONTROLLED_VOCAB


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["设备编号", "设备名称", "材质"])
    ws.append(["CT64201", "压片机A", "316L不锈钢"])
    ws.append(["DE64203", "包衣机B", "304不锈钢"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_config(client, headers, source_type="excel"):
    r = client.post("/api/extraction/configs", json={
        "name": "设备台账",
        "target_class_iri": "http://slpra.org/equipment#Equipment",
        "source_type": source_type,
        "column_mapping": {"设备编号": "equipmentID", "设备名称": "equipmentName",
                           "材质": "material"},
    })
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_extraction_prompt_injects_controlled_vocab():
    """受控词表取值注入抽取提示，在生成阶段约束 LLM（FR-006 / US1-AC3）。"""
    prompt = build_extraction_prompt(
        source_data=[{"材质": "316L不锈钢"}],
        target_class_iri="http://slpra.org/equipment#Equipment",
        property_schema=[],
        controlled_vocab=CONTROLLED_VOCAB,
    )
    assert "受控取值约束" in prompt
    # 提示须含受控词表取值，使 LLM 在生成时归一化而非自由发挥。
    assert "OEB1" in prompt
    assert "316L不锈钢" in prompt
    assert "µg/day" in prompt


def test_create_job_triggers_pipeline_running(client, analyst_headers):
    cfg_id = _make_config(client, analyst_headers)
    r = client.post(
        "/api/extraction/jobs",
        data={"source_type": "excel", "config_id": cfg_id},
        files={"file": ("设备台账.xlsx", _xlsx_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=analyst_headers,
    )
    assert r.status_code == 202, r.text
    job = r.json()
    # 作业不停留 pending：触发后为 running，后台流水线完成后转 reviewing。
    assert job["status"] in ("running", "reviewing", "done")

    job_id = job["id"]
    final = client.get(f"/api/extraction/jobs/{job_id}").json()
    assert final["status"] == "reviewing"
    assert final["total_candidates"] >= 2


def test_progress_sse_emits_stages(client, analyst_headers):
    cfg_id = _make_config(client, analyst_headers)
    job = client.post(
        "/api/extraction/jobs",
        data={"source_type": "excel", "config_id": cfg_id},
        files={"file": ("台账.xlsx", _xlsx_bytes(), "application/octet-stream")},
        headers=analyst_headers,
    ).json()
    r = client.get(f"/api/extraction/jobs/{job['id']}/progress", headers=analyst_headers)
    assert r.status_code == 200
    body = r.text
    for stage in ("parsing", "extracting", "aligning", "reviewing"):
        assert stage in body


def test_excel_instance_candidates_and_llm_degraded(client, analyst_headers):
    """LLM 不可用（测试默认无 key）→ 回退、候选带 degraded_reason、作业完成（FR-007）。"""
    cfg_id = _make_config(client, analyst_headers)
    job = client.post(
        "/api/extraction/jobs",
        data={"source_type": "excel", "config_id": cfg_id},
        files={"file": ("台账.xlsx", _xlsx_bytes(), "application/octet-stream")},
        headers=analyst_headers,
    ).json()
    cands = client.get(f"/api/extraction/jobs/{job['id']}/candidates").json()
    members = cands["ungrouped"] + [c for g in cands["groups"] for c in g["candidates"]]
    assert len(members) >= 2
    assert all(m["candidate_kind"] == "instance" for m in members)
    assert all(m["degraded_reason"] for m in members)
    # 受控词表归一化注入（316L不锈钢 / 304不锈钢）。
    assert any("_controlled_vocab" in m["extracted_properties"] for m in members)


def test_word_body_produces_action_candidate(client, analyst_headers):
    docx = __import__("docx")
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "设备编号"
    table.rows[0].cells[1].text = "设备名称"
    table.rows[1].cells[0].text = "CT64201"
    table.rows[1].cells[1].text = "压片机"
    doc.add_paragraph("若设备用于高致敏药品生产，则必须执行专用化管理。")
    buf = io.BytesIO()
    doc.save(buf)

    cfg_id = _make_config(client, analyst_headers, source_type="word")
    job = client.post(
        "/api/extraction/jobs",
        data={"source_type": "word", "config_id": cfg_id},
        files={"file": ("SOP.docx", buf.getvalue(), "application/octet-stream")},
        headers=analyst_headers,
    ).json()
    cands = client.get(f"/api/extraction/jobs/{job['id']}/candidates").json()
    members = cands["ungrouped"] + [c for g in cands["groups"] for c in g["candidates"]]
    actions = [m for m in members if m["candidate_kind"] == "action"]
    assert len(actions) >= 1
    assert actions[0]["action_conditions"]["precondition"]
    assert "必须" in actions[0]["action_conditions"]["obligation"]
