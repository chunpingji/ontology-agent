"""HRS-1597 inline comparison and execution-scoped manual review, without model calls."""

from copy import deepcopy
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import func, select

from app.config import settings
from app.dependencies import get_ontology_engine
from app.main import app
from app.models.extraction import AnnotationExecution
from app.models.pde_conflict import PdeConflictDecision
from app.services.template_finder.pde_review import TARGET_FILENAME, attach_comparisons
from app.services.template_finder.service import private_job_id
from tests.test_api.test_template_finder import (
    DEV,
    ROOT,
    FinderOntology,
    finder_setup,  # noqa: F401
    route,
)

SHARED = DEV + "SharedLineAssessmentData"
PROPERTIES = {
    "activeIngredient": "活性成分（API）", "studyType": "试验项目",
    "noael_mg_per_kg_per_day": "NOAEL（mg/kg/天）",
    "pde_mg_per_day": "PDE（mg/天）", "noaelSpecies": "动物种属",
    "noaelDuration": "试验周期", **{f"f{i}": f"F{i}" for i in range(1, 6)},
}


class PdeOntology(FinderOntology):
    def get_data_properties_by_domain(self, iri):
        return [
            {"iri": DEV + key, "name": key, "label": label, "datatype": "string"}
            for key, label in PROPERTIES.items()
        ] if iri == SHARED else []

    def get_relation_schema(self, root, max_hops=4):
        return [{
            "domain_class_iri": ROOT, "range_class_iri": SHARED,
            "predicate_iri": DEV + "hasSharedLineData", "predicate_label": "共线评估数据",
            "range_subclasses": [],
        }]


@pytest.fixture
def pde_setup(finder_setup, db):  # noqa: F811
    template, job, path, _ = finder_setup
    doc = Document()
    doc.add_heading("HRS-1597 CMC报告", 0)
    table = doc.add_table(rows=3, cols=9)
    for row, values in zip(table.rows, [
        ["活性成分(API)", "试验项目", "NOAEL", "F1", "F2", "F3", "F4", "F5", "PDE (mg/天)"],
        ["HRS-1597", "大鼠14天亚急毒试验（试验1）", "200", "5", "10", "10", "1", "1", "100"],
        ["HRS-1597", "犬14天亚急毒", "30", "2", "10", "10", "1", "1", "7.5"],
    ]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(path)
    job.source_filename = TARGET_FILENAME
    job.source_config = {**job.source_config, "mode": "doc_repo_preview"}
    db.commit()
    app.dependency_overrides[get_ontology_engine] = lambda: PdeOntology()
    return template, job, path


def start(client, headers, template, job, previous=None):
    url = route(template, job)
    response = client.post(url, headers=headers, json={
        "request_key": str(uuid4()), "expected_execution_id": previous,
    })
    assert response.status_code == 202, response.text
    status = client.get(url, headers=headers).json()
    assert status["status"] == "completed", status
    graph = client.get(url + "/graph", headers=headers, params={
        "execution_id": status["execution_id"],
    })
    assert graph.status_code == 200, graph.text
    return status["execution_id"], graph.json()


def review_url(template, job, execution):
    return route(template, job) + f"/pde-conflict/decision?execution_id={execution}"


def test_actual_row_factors_compare_point_bands_without_fabricating_negative_hazards(
    client, db, analyst_headers, pde_setup,
):
    template, job, _ = pde_setup
    _, graph = start(client, analyst_headers, template, job)
    first, second = graph["relationships"]
    conflict = first["conflict"]
    assert conflict["asserted"]["pde_mg_day"] == 100
    assert conflict["derived"]["pde_ug_day"] == 20000
    assert conflict["pde_ratio"] == 5
    assert conflict["derived"]["band"] == conflict["asserted"]["band"] == 1
    assert conflict["derived"]["provisional"] is True
    provenance = conflict["derived"]["provenance"]
    assert provenance["inputs"]["study_duration_days"] == 14
    assert all(provenance["inputs"][key] is None for key in (
        "genotoxic", "carcinogenic", "reproductive_toxicant",
    ))
    assert provenance["factor_sources"]["F4_severity"] == "document"
    assert provenance["factor_sources"]["BW_kg"] == "method_default"
    assert "conflict" not in second
    assert "pde_review_issues" not in second
    assert db.scalar(select(func.count()).select_from(PdeConflictDecision)) == 0


def test_review_is_read_only_then_cas_persists_on_source_and_rerun_requires_new_review(
    client, db, analyst_headers, operator_headers, pde_setup, monkeypatch,
):
    from app.services.template_finder import service

    template, job, _ = pde_setup
    execution, _ = start(client, analyst_headers, template, job)
    url = review_url(template, job, execution)
    private = private_job_id("analyst", template.id, job.id)
    cache = settings.template_finder_storage_dir / str(private) / "latest.json"
    unchanged = cache.read_bytes()

    def forbidden(*args, **kwargs):
        pytest.fail("Decision GET performed calculation or dispatched recognition")

    with monkeypatch.context() as patch:
        patch.setattr(service, "worker", forbidden)
        patch.setattr(service, "run", forbidden)
        patch.setattr(service, "freeze_ontology", forbidden)
        default = client.get(url, headers=analyst_headers)
    assert default.status_code == 200, default.text
    assert default.json()["chosen"] == "pending"
    assert db.scalar(select(func.count()).select_from(PdeConflictDecision)) == 0
    assert cache.read_bytes() == unchanged
    assert client.post(url, headers=operator_headers, json={"chosen": "derived"}).status_code == 403
    saved = client.post(url, headers=analyst_headers, json={
        "chosen": "derived", "expected_version": 0, "note": "已核对原文因子",
    })
    assert saved.status_code == 200, saved.text
    assert saved.json()["version"] == 1 and saved.json()["actor"] == "analyst"
    assert saved.json()["job_id"] == str(job.id)
    assert client.get(url, headers=analyst_headers).json()["chosen"] == "derived"
    assert client.post(url, headers=analyst_headers, json={
        "chosen": "asserted", "expected_version": 0,
    }).status_code == 409
    next_execution, _ = start(client, analyst_headers, template, job, execution)
    assert client.get(url, headers=analyst_headers).status_code == 409
    assert client.post(url, headers=analyst_headers, json={
        "chosen": "asserted", "expected_version": 1,
    }).status_code == 409
    next_url = review_url(template, job, next_execution)
    pending = client.get(next_url, headers=analyst_headers).json()
    assert pending["chosen"] == "pending" and pending["version"] == 1
    row = db.scalar(select(PdeConflictDecision))
    assert row.chosen == "derived" and row.version == 1
    assert client.post(next_url, headers=analyst_headers, json={
        "chosen": "asserted", "expected_version": 1,
    }).json()["version"] == 2
    assert db.scalar(select(func.count()).select_from(PdeConflictDecision)) == 1
    assert db.get(AnnotationExecution, private, populate_existing=True).options[
        "pde_decision_version"
    ] == 2


def test_review_rejects_other_owner_changed_source_and_retired_entrypoints(
    client, db, analyst_headers, pde_setup,
):
    template, job, path = pde_setup
    execution, _ = start(client, analyst_headers, template, job)
    url = review_url(template, job, execution)
    other = {**analyst_headers, "X-User": "other-owner"}
    assert client.get(url, headers=other).status_code == 409
    assert client.post(url, headers=other, json={"chosen": "derived"}).status_code == 409
    path.write_bytes(path.read_bytes() + b"changed")
    rejected = client.post(url, headers=analyst_headers, json={"chosen": "derived"})
    assert rejected.status_code == 409, rejected.text
    assert rejected.json()["detail"]["code"] == "INPUT_VERSION_CONFLICT"
    assert db.scalar(select(func.count()).select_from(PdeConflictDecision)) == 0
    legacy = f"/api/extraction/jobs/{job.id}/pde-conflict/decision"
    assert client.get(legacy, headers=analyst_headers).status_code == 410
    assert client.post(
        legacy, headers=analyst_headers, json={"chosen": "derived"},
    ).status_code == 410
    private = private_job_id("analyst", template.id, job.id)
    assert client.get(
        f"/api/extraction/jobs/{private}/pde-conflict/decision", headers=analyst_headers,
    ).status_code == 404


def _node(**replacements):
    values = {"noael_mg_per_kg_per_day": "200", "pde_mg_per_day": "100",
              "studyType": "大鼠14天亚急毒", **{f"f{i}": v for i, v in
              enumerate(("5", "10", "10", "1", "1"), 1)}, **replacements}
    return {
        "object_class_iri": SHARED,
        "object_data_properties": [
            {"iri": DEV + name, "label": PROPERTIES[name], "value": value,
             "raw_value": value, "source": {"raw_value": value, "anchors": [{"row": 1}]}}
            for name, value in values.items() if value is not None
        ],
    }


@pytest.mark.parametrize("change,expected", [
    ({"noael_mg_per_kg_per_day": "XX mg/kg/天（28天）"}, "明确的正数"),
    ({"noael_mg_per_kg_per_day": "200～300"}, "剂量单位"),
    ({"studyType": "14天研究"}, "动物种属"),
    ({"studyType": "duration 14 days"}, "动物种属"),
    ({"studyType": "大鼠研究"}, "试验周期"),
    ({"f3": "未知"}, "明确的正数"),
    ({"f3": "1e300", "f4": "1e300"}, "可计算范围"),
    ({"f3": "1e-300", "f4": "1e-300"}, "可计算范围"),
    ({"noael_mg_per_kg_per_day": "1e308"}, "可计算范围"),
])
def test_missing_or_ambiguous_inputs_remain_unreviewable(change, expected):
    node = _node(**change)
    graph = {"doc_class": {"doc_class_iri": ROOT}, "relationships": [node]}
    attach_comparisons(graph, TARGET_FILENAME)
    assert "conflict" not in node
    assert any(expected in issue for issue in node["pde_review_issues"])


def test_missing_unit_and_location_do_not_borrow_from_other_studies():
    node = _node()
    node["object_data_properties"][0]["label"] = "NOAEL"
    node["object_data_properties"][1]["source"]["anchors"] = []
    graph = {"doc_class": {"doc_class_iri": ROOT}, "relationships": [node, _node()]}
    attach_comparisons(graph, TARGET_FILENAME)
    assert "conflict" not in node
    assert len(node["pde_review_issues"]) == 2
    assert "conflict" in graph["relationships"][1]


def test_other_documents_are_unchanged_and_explicit_factors_override_defaults():
    graph = {"doc_class": {"doc_class_iri": ROOT}, "relationships": [_node(f5="5")]}
    original = deepcopy(graph)
    attach_comparisons(graph, "other.docx")
    assert graph == original
    attach_comparisons(graph, TARGET_FILENAME)
    conflict = graph["relationships"][0]["conflict"]
    assert conflict["derived"]["pde_ug_day"] == 4000
    assert conflict["derived"]["provenance"]["factors"]["F5_loael"] == 5


def test_no_conflict_cannot_be_reviewed(client, db, analyst_headers, pde_setup):
    template, job, path = pde_setup
    doc = Document(path)
    doc.tables[0].cell(1, 8).text = "20"
    doc.save(path)
    execution, graph = start(client, analyst_headers, template, job)
    assert all("conflict" not in node for node in graph["relationships"])
    url = review_url(template, job, execution)
    assert client.get(url, headers=analyst_headers).status_code == 409
    assert client.post(url, headers=analyst_headers, json={"chosen": "derived"}).status_code == 409
    assert db.scalar(select(func.count()).select_from(PdeConflictDecision)) == 0
