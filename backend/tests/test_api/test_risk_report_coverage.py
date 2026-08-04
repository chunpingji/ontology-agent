"""AST-7 integration tests for the risk-report endpoint coverage contract (010).

Exercises POST ``/api/extraction/jobs/{job_id}/risk-report`` end-to-end and
asserts the no-omission evidence is persisted both in
``GeneratedReport.rules_summary["coverage"]`` (AST-6) and in the hash-chained
audit trail ``details["coverage"]`` (FR-012). This is the API-level proof that a
missing required material can never silently disappear from a generated report.
"""

from __future__ import annotations

import json
from io import BytesIO

from docx import Document

from app.models.extraction import ExtractionJob, GeneratedReport
from app.models.ontology_meta import OntologyDecisionRule
from app.models.reasoning import AuditLog

CMC = "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"
SLPRA = "https://ontology.pharma-gmp.cn/slpra"


def _drug_edge() -> dict:
    return {
        "predicate_iri": f"{SLPRA}/drug-development/describes",
        "object_class_iri": f"{SLPRA}/drug/DrugProduct",
        "object_text": "HRS-1234",
        "object_data_properties": [
            {"iri": f"{SLPRA}/drug/pde_mg_per_day", "label": "PDE", "value": "1.80"},
            {"iri": None, "label": "分类", "value": "化学药品"},
        ],
        "source_ref": "§ 产品信息",
    }


def _shared_line_edge() -> dict:
    return {
        "predicate_iri": f"{SLPRA}/drug-development/hasSharedLineData",
        "object_class_iri": f"{SLPRA}/drug-development/SharedLineAssessmentData",
        "object_text": "共线评估",
        "object_data_properties": [],
        "source_ref": "§ 共线评估",
    }


def _pde_conflict() -> dict:
    return {
        "conflict_key": "shared_line_pde",
        "asserted": {"pde_mg_day": 1.8, "pde_ug_day": 1800.0, "band": 2},
        "derived": {
            "band": 5,
            "pde_ug_day": 50.0,
            "input_source": "extracted",
            "provenance": {
                "method": {"id": "ADE-OEB/test", "version": "1.0"},
                "formula": "PDE = NOAEL·BW / (F1·F2·F3·F4·F5)",
            },
        },
        "delta_bands": 3,
        "pde_ratio": 36.0,
        "summary": "推导 OEB band 5 与原文 band 2 不一致。",
    }


def _equipment_edge(code: str = "RE001") -> dict:
    return {
        "predicate_iri": f"{SLPRA}/drug-development/usesEquipment",
        "object_class_iri": f"{SLPRA}/equipment/Equipment",
        "object_text": code,
        "object_data_properties": [
            {"iri": None, "label": "设备名称", "value": f"设备-{code}"},
            {"iri": None, "label": "设备规格", "value": "搅拌釜 500L"},
            {"iri": None, "label": "材质", "value": "316L"},
        ],
        "source_ref": "642车间 设备需求",
    }


def _seed_rule(db, *, key: str, category: str) -> None:
    db.add(OntologyDecisionRule(
        slpra_iri=f"{SLPRA}/rules/{key}",
        label=f"Rule {key}",
        rule_key=key,
        rule_group="risk_assessment",
        antecedent={"op": "some_values_from", "property": "hasSharedLineData",
                    "filler_class": "SharedLineAssessmentData"},
        consequent={
            "risk_level": "HighRisk", "category": category,
            "description": f"风险因素：{category}",
            "control_measure": f"控制措施：{category}",
            "traceability_docs": f"追溯文件：{category}",
            "postconditions": {},
        },
        priority=100,
        status="published",
    ))
    db.commit()


def _seed_job(db, edges: list[dict], tmp_path) -> ExtractionJob:
    job = ExtractionJob(
        source_type="document", source_filename="HRS-1234.docx", status="completed",
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    cache_dir = tmp_path / "data" / "uploads"
    cache_dir.mkdir(parents=True, exist_ok=True)
    (cache_dir / f"{job.id}.annotated.json").write_text(
        json.dumps({
            "doc_class": {"doc_class_iri": CMC},
            "relationships": edges,
        }),
        encoding="utf-8",
    )
    return job


def test_report_persists_full_coverage_and_audit(client, db, analyst_headers, monkeypatch, tmp_path):
    """Happy path: every prerequisite present → coverage persisted, no omissions."""
    monkeypatch.chdir(tmp_path)
    _seed_rule(db, key="R-RA-EQ", category="生产设备")
    job = _seed_job(db, [_drug_edge(), _shared_line_edge(), _equipment_edge()], tmp_path)

    r = client.post(f"/api/extraction/jobs/{job.id}/risk-report", headers=analyst_headers)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK"  # .docx ZIP

    gen = db.query(GeneratedReport).filter(GeneratedReport.job_id == job.id).one()
    coverage = gen.rules_summary["coverage"]
    assert coverage["template_id"] == "QS-A-020F05@v1"
    assert coverage["missing_required"] == 0
    assert coverage["total_slots"] == len(coverage["slots"])  # full manifest stored
    # legacy rows still present for backward compatibility
    assert "rows" in gen.rules_summary

    entry = (
        db.query(AuditLog)
        .filter(AuditLog.action == "report.generate")
        .order_by(AuditLog.seq.desc())
        .first()
    )
    assert entry is not None
    assert entry.details["coverage"]["missing_required"] == 0
    assert entry.details["coverage"]["template_id"] == "QS-A-020F05@v1"


def test_report_records_omissions_when_shared_line_missing(client, db, analyst_headers, monkeypatch, tmp_path):
    """G1: missing shared-line surfaces as missing_required in DB + audit, not silently."""
    monkeypatch.chdir(tmp_path)
    _seed_rule(db, key="R-RA-EQ", category="生产设备")
    job = _seed_job(db, [_drug_edge(), _equipment_edge()], tmp_path)  # no shared-line

    r = client.post(f"/api/extraction/jobs/{job.id}/risk-report", headers=analyst_headers)
    assert r.status_code == 200, r.text

    gen = db.query(GeneratedReport).filter(GeneratedReport.job_id == job.id).one()
    coverage = gen.rules_summary["coverage"]
    assert coverage["missing_required"] > 0
    assert "prereq.shared_line" in coverage["missing_slot_ids"]

    entry = (
        db.query(AuditLog)
        .filter(AuditLog.action == "report.generate")
        .order_by(AuditLog.seq.desc())
        .first()
    )
    assert entry.details["coverage"]["missing_required"] == coverage["missing_required"]
    assert "prereq.shared_line" in entry.details["coverage"]["missing_slot_ids"]


def test_audit_chain_stays_valid_after_report(client, db, analyst_headers, monkeypatch, tmp_path):
    """The coverage-bearing audit entry must keep the hash chain verifiable (SC-008)."""
    from app.services import audit

    monkeypatch.chdir(tmp_path)
    _seed_rule(db, key="R-RA-EQ", category="生产设备")
    job = _seed_job(db, [_drug_edge(), _equipment_edge()], tmp_path)

    r = client.post(f"/api/extraction/jobs/{job.id}/risk-report", headers=analyst_headers)
    assert r.status_code == 200, r.text
    assert audit.verify(db)["ok"] is True


def test_report_persists_and_renders_pde_conflict_decision(
    client, db, analyst_headers, monkeypatch, tmp_path,
):
    """The generated artifact pins the human decision while retaining both values."""
    monkeypatch.chdir(tmp_path)
    shared_line = _shared_line_edge()
    shared_line["conflict"] = _pde_conflict()
    job = _seed_job(db, [_drug_edge(), shared_line, _equipment_edge()], tmp_path)

    decided = client.post(
        f"/api/extraction/jobs/{job.id}/pde-conflict/decision",
        headers=analyst_headers,
        json={
            "chosen": "derived",
            "note": "已核对毒理试验原始记录",
            "expected_version": 0,
        },
    )
    assert decided.status_code == 200, decided.text

    response = client.post(
        f"/api/extraction/jobs/{job.id}/risk-report", headers=analyst_headers,
    )
    assert response.status_code == 200, response.text

    generated = db.query(GeneratedReport).filter(GeneratedReport.job_id == job.id).one()
    progress = generated.rules_summary["_progress"]
    assert progress["stage"] == "completed"
    assert progress["percent"] == 100
    assert [(step["key"], step["status"]) for step in progress["substeps"]] == [
        ("facts", "completed"),
        ("rules", "completed"),
        ("coverage", "completed"),
        ("narratives", "completed"),
    ]
    conflicts = generated.rules_summary["pde_conflicts"]
    assert len(conflicts) == 1
    assert conflicts[0]["asserted"]["pde_mg_day"] == 1.8
    assert conflicts[0]["derived"]["pde_ug_day"] == 50.0
    assert conflicts[0]["effective"]["source"] == "derived"
    assert conflicts[0]["decision"]["actor"] == "analyst"
    assert conflicts[0]["decision"]["note"] == "已核对毒理试验原始记录"

    doc = Document(BytesIO(response.content))
    text = "\n".join(
        [paragraph.text for paragraph in doc.paragraphs]
        + [
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert "PDE/OEB 潜能等级冲突及裁决" in text
    assert "已裁决（采纳推导值）" in text
    assert "原文断言" in text and "确定性推导" in text
    assert "analyst" in text
    assert "已核对毒理试验原始记录" in text
