"""Sample formatting, source completeness, and version-bound generation."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Pt
from docx.table import _Cell

from app.models.extraction import AstTemplate
from app.services.reporting.batch_demo_layout import REFERENCE_ID, metadata
from app.services.reporting.batch_demo_operation_forms import load_forms
from app.services.reporting.batch_demo_sample_content import A14_OPERATIONS, A14_TABLES
from app.services.reporting.output_ast import walk
from tests.test_api.batch_layout_fixture import xml_signature
from tests.test_api.test_batch_demo import get, post
from tests.test_api.test_batch_demo import source as source


def test_sample_format_and_source_only_content(client, db, analyst_headers, source):
    reference = db.get(AstTemplate, REFERENCE_ID)
    before = Path(reference.sample_docx_path).read_bytes()
    report = post(client, analyst_headers, source).json()
    assert '待填写' not in str(report['body_ast'])
    downloaded = client.get(
        f"/api/extraction/jobs/{report['job_id']}/reports/{report['id']}/download",
        headers=analyst_headers,
    ).content
    document = Document(BytesIO(downloaded))
    sample = Document(BytesIO(before))
    title = next(p for p in document.paragraphs if p.text == '批生产记录')
    assert title.runs[0].font.size == Pt(24)  # Comes from fixture, not hardcoded production size.
    assert title.alignment == sample.paragraphs[6].alignment
    assert len(document.sections) == 7
    for section in document.sections:
        assert section.page_width == sample.sections[0].page_width
        assert section.left_margin == sample.sections[0].left_margin
    assert [tc.grid_span for tc in document.tables[0].rows[0]._tr.tc_lst] == [1, 2, 1]
    equipment = document.tables[2]
    assert equipment.cell(1, 1).text == "RE64615/RE64215"
    operations = [t for t in document.tables if t.cell(0, 0).text == '原材料/操作']
    assert [len(t.rows) - 1 for t in operations[-3:]] == [150, 180, 190]
    for actual, ti in zip(operations[:10], A14_OPERATIONS):
        assert xml_signature(actual._tbl) == xml_signature(sample.tables[ti]._tbl)
    # The whole explicitly selected record section is copied, including weighing/TLC forms.
    a14_start = next(i for i, t in enumerate(document.tables)
                     if _Cell(t._tbl.tr_lst[0].tc_lst[0], t).text == '称量记录')
    for actual, ti in zip(document.tables[a14_start:a14_start + 17], A14_TABLES):
        assert xml_signature(actual._tbl) == xml_signature(sample.tables[ti]._tbl)
        # Preserve the original heading and page-break paragraph between forms.
        assert xml_signature(actual._tbl.getprevious()) == xml_signature(
            sample.tables[ti]._tbl.getprevious(),
        )
        assert xml_signature(actual._tbl.getprevious().getprevious()) == xml_signature(
            sample.tables[ti]._tbl.getprevious().getprevious(),
        )
    route = next(n for n in source[0]['relationships']
                 if n['predicate_iri'].endswith('/hasSynthesisRoute'))
    forms = load_forms()
    for table, step, stage in zip(operations[-3:], route['sub_relationships'][1:],
                                  forms['stages'][1:]):
        assert table._tbl.tblGrid.xml == sample.tables[15]._tbl.tblGrid.xml
        physical_rows = table._tbl.tr_lst[1:]
        starts = [i for i, r in enumerate(physical_rows) if _Cell(r.tc_lst[0], table).text]
        assert len(starts) == len(step['operations'])
        for i, (op, binding) in enumerate(zip(step['operations'], stage['operations'])):
            start = starts[i]
            end = starts[i + 1] if i + 1 < len(starts) else len(physical_rows)
            group = physical_rows[start:end]
            assert op['instruction'] in _Cell(group[0].tc_lst[0], table).text
            assert len(group) == (len(binding['fields']) or 1)
            for j, row in enumerate(group):
                assert len(row.tc_lst) == 5
                texts = [_Cell(c, table).text for c in row.tc_lst]
                assert texts[3:] == ['', '']
                for c in (1, 2):
                    assert row.tc_lst[c].vMerge is None and row.tc_lst[c].grid_span == 1
                if len(group) > 1:
                    for c in (0, 3, 4):
                        assert row.tc_lst[c].vMerge == ('restart' if j == 0 else 'continue')
                if binding['fields']:
                    f = forms['fields'][binding['fields'][j]]
                    assert texts[1] == f['parameter_name']
                    assert texts[2].strip() == f['record_format'].strip()
    for i in range(1, 6):
        assert xml_signature(document.sections[i].header._element) == xml_signature(
            sample.sections[i].header._element,
        )
    assert not document.sections[6].header.tables  # No invented process in the appendix.
    copied = operations[0].cell(1, 0).paragraphs[0].runs
    assert [(r.font.name, r.font.size.pt, r.italic) for r in copied if r.text] == [
        ('Arial', 12, None), ('Times New Roman', 10, True),
    ]
    preview = next(n for n in walk(report['body_ast'])
                   if n.kind == 'table' and metadata(n).get('sample_table') == 15)
    runs = metadata(preview.children[1].children[0])['paragraphs'][0]['runs']
    assert [(r['style']['latin_font_family'], r['style']['font_size_pt'], r['style']['italic'])
            for r in runs] == [('Arial', 12, False), ('Times New Roman', 10, True)]
    with ZipFile(BytesIO(downloaded)) as archive:
        assert not any('/media/' in n or '/embeddings/' in n for n in archive.namelist())
        xml = '\n'.join(archive.read(n).decode() for n in archive.namelist() if n.endswith('.xml'))
        assert 'SAMPLE-ONLY-SECRET' not in xml
        assert '待填写' not in xml
        assert '□是' in xml and '☑' not in xml
    assert Path(reference.sample_docx_path).read_bytes() == before


@pytest.mark.parametrize('change', ['sample', 'schema', 'content', 'header', 'a14_instruction'])
def test_layout_version_conflict_preserves_history(client, db, analyst_headers, source, change):
    original = get(client, analyst_headers, source).json()
    report = post(client, analyst_headers, source).json()
    reference = db.get(AstTemplate, REFERENCE_ID)
    if change == 'sample':
        doc = Document(reference.sample_docx_path)
        doc.paragraphs[6].runs[0].font.size = Pt(26)
        doc.save(reference.sample_docx_path)
    elif change in ('header', 'a14_instruction'):
        doc = Document(reference.sample_docx_path)
        cell = doc.sections[2].header.tables[0].cell(0, 3) if change == 'header' else (
            doc.tables[15].cell(1, 0)
        )
        cell.text = '用户修订的参考内容'
        doc.save(reference.sample_docx_path)
    elif change == 'schema':
        reference.schema_json = {**reference.schema_json, 'style_ref': 'updated'}
    else:
        reference.sample_content_json = {'type': 'doc', 'updated': True}
    db.commit()
    current = get(client, analyst_headers, source).json()
    assert original['graph_hash'] == current['graph_hash']
    assert original['template_hash'] != current['template_hash']
    assert current['latest_report'] is None
    stale = post(client, analyst_headers, source, template_hash=original['template_hash'])
    assert stale.status_code == 409
    assert client.get(
        f"/api/extraction/jobs/{report['job_id']}/reports/{report['id']}/download",
        headers=analyst_headers,
    ).status_code == 200


@pytest.mark.parametrize('change', [
    'missing', 'invalid_zip', 'structure', 'new_sections', 'parameter_moved', 'actual_value',
])
def test_unusable_reference_fails_without_fallback(client, db, analyst_headers, source, change):
    reference = db.get(AstTemplate, REFERENCE_ID)
    path = Path(reference.sample_docx_path)
    if change == 'missing':
        path.unlink()
    elif change == 'invalid_zip':
        path.write_bytes(b'not Word')
    elif change == 'structure':
        doc = Document(path)
        table = doc.tables[0]._tbl
        table.remove(table.tr_lst[-1])
        doc.save(path)
    elif change in ('parameter_moved', 'actual_value'):
        doc = Document(path)
        doc.tables[15].cell(8, 1 if change == 'parameter_moved' else 2).text = (
            'unrelated parameter' if change == 'parameter_moved' else '230 kg'
        )
        doc.save(path)
    else:
        reference.schema_json = {'sections': [{'id': 'new-definition'}]}
        db.commit()
    result = get(client, analyst_headers, source)
    assert result.status_code == 409
    assert '样例' in result.text
