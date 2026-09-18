"""Real Word/Finder API tests; all sources and artifacts live under tmp_path."""

import json
from contextlib import nullcontext
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import func, select

from app.config import settings
from app.dependencies import get_ontology_engine
from app.main import app
from app.models.document_analysis import DocumentAnalysisRun
from app.models.extraction import AnnotationExecution, AstTemplate, ExtractionJob
from app.services.extraction.document_ir import DocumentIR
from app.services.template_finder.service import private_job_id

DEV = "https://ontology.pharma-gmp.cn/slpra/drug-development/"
DRUG = "https://ontology.pharma-gmp.cn/slpra/drug/"
ROOT = DEV + "CMCReport"

FINDER_ENGINE = {"recognition_mode": "finder_legacy", "finder_profile_id": "cmc_baseline_v1"}
NORMAL_ENGINE = {"recognition_mode": "ontology_guided", "finder_profile_id": None}


def switch_engine(client, headers, template, desired, expected):
    return client.patch(
        f"/api/ast-templates/{template.id}/recognition-engine",
        json={**desired, **{f"expected_{key}": value for key, value in expected.items()}},
        headers=headers,
    )


@pytest.mark.parametrize("schema_version", [1, 2])
def test_switch_engine_is_shared_without_execution_or_schema_changes(
    client, db, analyst_headers, finder_setup, schema_version
):
    from copy import deepcopy

    from app.models.entity_shadow import EntityShadow

    template, job, _, config_path = finder_setup
    if schema_version == 1:
        template.schema_json = {"sections": []}
    template.status = "published"
    db.add(
        EntityShadow(
            iri="urn:engine-switch",
            class_iri=ROOT,
            module="document",
            properties_json={"source_job_id": str(job.id)},
        )
    )
    db.commit()
    schema = deepcopy(template.schema_json)
    detail_url = f"/api/ast-templates/{template.id}"
    initial = client.get(detail_url, headers=analyst_headers).json()
    endpoint = detail_url + "/recognition-engine"
    assert client.get(endpoint, headers=analyst_headers).json()["finder_profiles"] == [
        {"id": "cmc_baseline_v1", "label": "CMC（本体指引1.0）"}
    ]
    for desired, expected in (
        (NORMAL_ENGINE, FINDER_ENGINE),
        (FINDER_ENGINE, NORMAL_ENGINE),
        (NORMAL_ENGINE, FINDER_ENGINE),
    ):
        saved = switch_engine(client, analyst_headers, template, desired, expected)
        assert saved.status_code == 200, saved.text
        detail = client.get(detail_url, headers=analyst_headers).json()
        context = client.get(
            "/api/ast-templates/recognition-context",
            params={"document_iri": "urn:engine-switch", "template_id": str(template.id)},
            headers=analyst_headers,
        )
        assert context.status_code == 200, context.text
        for key, value in desired.items():
            assert detail[key] == value
            assert context.json()["selected"][key] == value
        assert detail["schema_json"] == schema
        assert detail["schema_hash"] == initial["schema_hash"]
        assert detail["revision_no"] == initial["revision_no"]
        assert detail["status"] == "published"
        finder = client.get(route(template, job), headers=analyst_headers)
        assert finder.status_code == (200 if desired == FINDER_ENGINE else 409)
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 0
    assert db.scalar(select(func.count()).select_from(DocumentAnalysisRun)) == 0
    assert db.scalar(select(func.count()).select_from(ExtractionJob)) == 1
    assert "cmc_baseline_v1" in config_path.read_text()


def test_engine_permission_and_stale_editor_do_not_overwrite_choice(
    client, db, analyst_headers, operator_headers, qa_headers, finder_setup
):
    from app.services.template_finder.policy import resolve

    template, *_ = finder_setup
    for headers in (operator_headers, qa_headers):
        assert (
            client.get(
                f"/api/ast-templates/{template.id}/recognition-engine", headers=headers
            ).status_code
            == 200
        )
        assert (
            switch_engine(client, headers, template, NORMAL_ENGINE, FINDER_ENGINE).status_code
            == 403
        )
    assert (
        switch_engine(client, analyst_headers, template, NORMAL_ENGINE, FINDER_ENGINE).status_code
        == 200
    )
    stale = switch_engine(client, analyst_headers, template, FINDER_ENGINE, FINDER_ENGINE)
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "RECOGNITION_CONFIG_CHANGED"
    db.refresh(template)
    assert template.recognition_mode == "ontology_guided"
    assert resolve(template) is None


@pytest.mark.parametrize(
    "desired",
    [
        {**FINDER_ENGINE, "finder_profile_id": "../../arbitrary-script"},
        {**FINDER_ENGINE, "finder_profile_id": None},
        {**NORMAL_ENGINE, "finder_profile_id": "cmc_baseline_v1"},
        {**NORMAL_ENGINE, "recognition_mode": "unknown"},
    ],
)
def test_invalid_engine_configuration_is_rejected(
    client, db, analyst_headers, finder_setup, desired
):
    from app.services.template_finder.policy import resolve

    template, *_ = finder_setup
    assert (
        switch_engine(client, analyst_headers, template, desired, FINDER_ENGINE).status_code == 422
    )
    db.refresh(template)
    assert template.recognition_mode is None
    assert resolve(template).profile_id == "cmc_baseline_v1"


def test_unsupported_roots_and_static_demo_cannot_enable_finder(
    client, db, analyst_headers, finder_setup
):
    template, _, _, config = finder_setup
    config.write_text('{"bindings": []}')
    template.iri_pattern = "urn:unsupported"
    db.commit()
    endpoint = f"/api/ast-templates/{template.id}/recognition-engine"
    assert client.get(endpoint, headers=analyst_headers).json()["finder_profiles"] == []
    assert (
        switch_engine(client, analyst_headers, template, FINDER_ENGINE, NORMAL_ENGINE).status_code
        == 422
    )
    template.iri_pattern = ROOT
    template.schema_json = {"demo_profile": {"fixture_id": "static"}}
    db.commit()
    assert (
        switch_engine(client, analyst_headers, template, FINDER_ENGINE, NORMAL_ENGINE).status_code
        == 409
    )


def test_saved_engine_is_not_inherited_by_new_template_revision(
    client, db, analyst_headers, finder_setup
):
    template, _, _, config = finder_setup
    config.write_text('{"bindings": []}')
    template.schema_json = {"sections": []}
    db.commit()
    assert (
        switch_engine(client, analyst_headers, template, FINDER_ENGINE, NORMAL_ENGINE).status_code
        == 200
    )
    created = client.put(
        f"/api/ast-templates/{template.id}",
        json={"schema_json": {"template_id": "engine-revision", "sections": []}},
        headers=analyst_headers,
    )
    assert created.status_code == 201, created.text
    assert created.json()["id"] != str(template.id)
    assert created.json()["recognition_mode"] == "ontology_guided"
    assert created.json()["finder_profile_id"] is None


def test_completed_finder_cache_survives_engine_roundtrip(
    client, db, analyst_headers, finder_setup
):
    template, job, *_ = finder_setup
    endpoint = route(template, job)
    assert (
        client.post(
            endpoint, json={"request_key": "before-switch"}, headers=analyst_headers
        ).status_code
        == 202
    )
    original = client.get(endpoint, headers=analyst_headers).json()
    assert original["status"] == "completed"
    for desired, expected in ((NORMAL_ENGINE, FINDER_ENGINE), (FINDER_ENGINE, NORMAL_ENGINE)):
        assert (
            switch_engine(client, analyst_headers, template, desired, expected).status_code == 200
        )
    current = client.get(endpoint, headers=analyst_headers).json()
    assert current["execution_id"] == original["execution_id"]
    assert current["has_result"] and not current["stale"]
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 1


def test_root_metadata_change_cannot_leave_a_broken_finder_binding(
    client, db, analyst_headers, finder_setup
):
    template, *_ = finder_setup
    template.schema_json = {"sections": []}
    db.commit()
    rejected = client.patch(
        f"/api/ast-templates/{template.id}",
        json={"iri_pattern": "urn:unsupported"},
        headers=analyst_headers,
    )
    assert rejected.status_code == 409
    db.rollback()
    db.refresh(template)
    assert template.iri_pattern == ROOT


def test_engine_migration_preserves_existing_template_data(tmp_path):
    from runpy import run_path

    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import Column, MetaData, String, Table, create_engine, inspect

    migration = run_path("alembic/versions/0041_template_engine.py")
    engine = create_engine(f"sqlite:///{tmp_path / 'migration.db'}")
    metadata = MetaData()
    Table("ast_templates", metadata, Column("id", String, primary_key=True))
    metadata.create_all(engine)
    with engine.begin() as connection, Operations.context(MigrationContext.configure(connection)):
        connection.execute(metadata.tables["ast_templates"].insert().values(id="existing"))
        migration["upgrade"]()
        table = Table("ast_templates", MetaData(), autoload_with=connection)
        assert dict(connection.execute(table.select()).one()._mapping) == {
            "id": "existing",
            "recognition_mode": None,
            "finder_profile_id": None,
        }
        migration["downgrade"]()
        assert [c["name"] for c in inspect(connection).get_columns("ast_templates")] == ["id"]
    engine.dispose()


class FinderOntology:
    def sparql_query(self, query):
        return []

    def lexical_read_scope(self):
        return nullcontext()

    def get_class_detail(self, iri):
        return {"iri": iri}

    def get_class_label(self, iri):
        return iri.rsplit("/", 1)[-1]

    def get_subclass_synonyms(self, iri):
        return {}

    def get_data_properties_by_domain(self, iri):
        return (
            [{"iri": DRUG + "appearance", "name": "appearance", "label": "性状"}]
            if iri == DRUG + "DrugProduct"
            else []
        )

    def get_relation_schema(self, root, max_hops=4):
        return [
            {
                "domain_class_iri": ROOT,
                "range_class_iri": DRUG + "DrugProduct",
                "predicate_iri": DEV + "describes",
                "predicate_label": "描述",
                "range_subclasses": [],
            }
        ]


@pytest.fixture
def finder_setup(client, db, tmp_path, monkeypatch):
    config = tmp_path / "bindings.json"
    monkeypatch.setattr(settings, "template_finder_config_path", config)
    monkeypatch.setattr(settings, "template_finder_storage_dir", tmp_path / "finder")
    doc = Document()
    doc.add_heading("HRS-1234 CMC报告", 0)
    doc.add_paragraph("性状：重复值😀")
    doc.add_heading("产品基本性质", 1)
    doc.add_paragraph("性状：重复值😀")
    path = tmp_path / "source.docx"
    doc.save(path)
    job = ExtractionJob(
        source_type="word",
        source_filename="source.docx",
        document_path=str(path),
        source_config={"mode": "template_default", "doc_class_iri": ROOT},
    )
    template = AstTemplate(
        name="Finder V2",
        version="1",
        iri_pattern=ROOT,
        schema_json={"schema_version": 2, "sections": []},
    )
    db.add_all([job, template])
    db.flush()
    template.default_source_job_id = job.id
    template.schema_json = {
        **template.schema_json,
        "template_family_id": str(template.id),
        "template_revision_id": str(template.id),
        "revision_no": 1,
    }
    job.source_config = {**job.source_config, "template_id": str(template.id)}
    db.commit()
    config.write_text(
        json.dumps(
            {
                "bindings": [
                    {
                        "template_id": str(template.id),
                        "root_class_iri": ROOT,
                        "finder_profile": "cmc_baseline_v1",
                    }
                ]
            }
        )
    )
    app.dependency_overrides[get_ontology_engine] = lambda: FinderOntology()
    return template, job, path, config


def route(template, job):
    return f"/api/ast-templates/{template.id}/sources/{job.id}/finder"


def test_finder_upload_only_registers_and_rejects_non_word(
    client,
    db,
    analyst_headers,
    finder_setup,
    monkeypatch,
    tmp_path,
):
    from app.api import ast_templates, extraction

    template, job, path, _ = finder_setup
    uploads = tmp_path / "uploads"
    monkeypatch.setattr(ast_templates, "_UPLOADS", uploads)

    def forbidden(*args, **kwargs):
        pytest.fail("Finder default upload dispatched the generic annotation worker")

    monkeypatch.setattr(extraction, "_enqueue_annotation", forbidden)
    url = f"/api/ast-templates/{template.id}/default-source"
    rejected = client.post(
        url, files={"file": ("test.xlsx", b"not a Word")}, headers=analyst_headers
    )
    assert rejected.status_code == 422
    assert not uploads.exists()
    accepted = client.post(
        url, files={"file": ("new.docx", path.read_bytes())}, headers=analyst_headers
    )
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["recognition_mode"] == "finder_legacy"
    new_id = accepted.json()["default_source_job_id"]
    assert new_id != str(job.id)
    from uuid import UUID

    new_job = db.get(ExtractionJob, UUID(new_id))
    assert new_job.status == "pending"
    assert new_job.source_config["mode"] == "template_default"
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 0


def test_shared_source_explicit_normal_template_and_existing_generic_resume(
    db,
    finder_setup,
    monkeypatch,
    tmp_path,
):
    from app.api import extraction

    template, job, _, config_path = finder_setup
    normal = AstTemplate(
        name="普通模板", version="1", iri_pattern=ROOT, schema_json={"sections": []}
    )
    db.add(normal)
    db.commit()
    checkpoint = tmp_path / "generic.json"
    monkeypatch.setattr(extraction, "_annotation_checkpoint_path", lambda _: checkpoint)
    monkeypatch.setattr(
        extraction, "_annotation_cache_path", lambda _: tmp_path / "generic-cache.json"
    )
    first = extraction._claim_annotation(job.id, db, mode="start", template_id=normal.id)
    assert first["status"] == "queued"
    assert job.source_config["recognition_template_id"] == str(normal.id)
    head = db.get(AnnotationExecution, job.id, populate_existing=True)
    head.status = "completed"
    db.commit()
    configured = config_path.read_text()
    config_path.write_text('{"bindings": []}')
    extraction._claim_annotation(job.id, db, mode="start", template_id=template.id)
    config_path.write_text(configured)
    head = db.get(AnnotationExecution, job.id, populate_existing=True)
    assert "recognition_mode" not in head.options
    head.status = "paused"
    db.commit()
    extraction._write_annotation_checkpoint(
        job.id, {"input_id": "existing", "attempt_count": 0, "completed": {}}
    )
    continued = extraction._claim_annotation(job.id, db, mode="continue")
    assert continued["status"] == "queued"
    assert db.get(AnnotationExecution, job.id, populate_existing=True).options["mode"] == "continue"


def test_real_finder_same_run_original_idempotence_and_isolation(
    client, db, analyst_headers, finder_setup
):
    template, job, path, _ = finder_setup
    url = route(template, job)
    assert client.get(url, headers=analyst_headers).json()["status"] == "not_started"
    preview = client.get(url + "/source", headers=analyst_headers)
    assert preview.status_code == 200, preview.text
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 0
    assert not settings.template_finder_storage_dir.exists()
    request = {"request_key": "first", "expected_execution_id": None}
    created = client.post(url, json=request, headers=analyst_headers)
    assert created.status_code == 202, created.text
    status = client.get(url, headers=analyst_headers).json()
    assert status["status"] == "completed", status
    assert status["has_result"] and not status["stale"]
    eid = status["execution_id"]
    assert client.post(url, json=request, headers=analyst_headers).json()["execution_id"] == eid
    assert (
        client.post(
            url, json={**request, "request_key": "other"}, headers=analyst_headers
        ).status_code
        == 409
    )
    graph = client.get(url + "/graph", params={"execution_id": eid}, headers=analyst_headers).json()
    original = client.get(
        url + "/source", params={"execution_id": eid}, headers=analyst_headers
    ).json()
    for key in ("execution_id", "document_hash", "structure_hash", "parser_version", "analysis_id"):
        assert graph[key] == original[key]
    prop = graph["relationships"][0]["object_data_properties"][0]
    assert prop["value"] == "重复值😀"
    anchor = prop["source"]["anchors"][0]
    assert anchor["paragraph_index"] == 3
    ir = DocumentIR.model_validate(original["content"]["analysis"])
    assert ir.resolve(anchor) == "重复值😀"
    private_id = private_job_id(analyst_headers["X-User"], template.id, job.id)
    head = db.get(AnnotationExecution, private_id, populate_existing=True)
    assert head.run_id != eid
    assert db.scalar(select(func.count()).select_from(DocumentAnalysisRun)) == 0
    for suffix in ("", "/progress", "/annotated-document", "/candidates", "/reports"):
        response = client.get(f"/api/extraction/jobs/{private_id}{suffix}", headers=analyst_headers)
        assert response.status_code == 404, (suffix, response.text)
    assert str(private_id) not in client.get("/api/extraction/jobs", headers=analyst_headers).text
    other = {**analyst_headers, "X-User": "different-owner"}
    assert client.get(url, headers=other).json()["status"] == "not_started"
    assert (
        client.get(url + "/graph", params={"execution_id": eid}, headers=other).status_code == 409
    )
    assert job.source_config["mode"] == "template_default"
    assert template.default_source_job_id == job.id
    assert [f.name for f in settings.template_finder_storage_dir.rglob("*") if f.is_file()] == [
        "latest.json"
    ]
    path.unlink()
    assert client.get(url, headers=analyst_headers).json()["stale"]
    assert (
        client.get(url + "/source", params={"execution_id": eid}, headers=analyst_headers).json()[
            "content"
        ]
        == original["content"]
    )


def test_mode_guards_config_errors_and_rerun(client, db, analyst_headers, finder_setup):
    from fastapi import HTTPException

    from app.api.extraction import _claim_annotation

    template, job, path, config = finder_setup
    url = route(template, job)
    endpoint = f"/api/document-analysis/templates/{template.id}/sources/{job.id}/runs"
    response = client.post(endpoint, json={"request_key": "wrong-domain"}, headers=analyst_headers)
    assert response.status_code == 409, response.text
    assert response.json()["error"]["code"] == "RECOGNITION_MODE_MISMATCH"
    with pytest.raises(HTTPException, match="本体指引1[.]0"):
        _claim_annotation(job.id, db, mode="start", template_id=template.id)
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 0
    response = client.post(url, json={"request_key": "one"}, headers=analyst_headers)
    eid = response.json()["execution_id"]
    second = client.post(
        url, json={"request_key": "two", "expected_execution_id": eid}, headers=analyst_headers
    )
    assert second.status_code == 202, second.text
    assert second.json()["execution_id"] != eid
    assert (
        client.get(
            url + "/graph", params={"execution_id": eid}, headers=analyst_headers
        ).status_code
        == 409
    )
    assert client.post(url, json={"request_key": "one"}, headers=analyst_headers).status_code == 409
    payload = json.loads(config.read_text())
    payload["bindings"][0]["finder_profile"] = "../bad"
    config.write_text(json.dumps(payload))
    failed = client.post(url, json={"request_key": "bad"}, headers=analyst_headers)
    assert failed.status_code == 409
    assert failed.json()["detail"]["code"] == "FINDER_CONFIG_INVALID"
    assert (
        client.get(
            url + "/source",
            params={"execution_id": second.json()["execution_id"]},
            headers=analyst_headers,
        ).status_code
        == 200
    )


def test_context_and_normal_report_guard(client, db, analyst_headers, finder_setup):
    from app.models.entity_shadow import EntityShadow

    template, job, _, _ = finder_setup
    doc = EntityShadow(
        iri="urn:finder-source",
        class_iri=ROOT,
        module="document",
        properties_json={"job_id": str(job.id)},
    )
    db.add(doc)
    db.commit()
    result = client.get(
        "/api/ast-templates/recognition-context",
        params={"document_iri": doc.iri},
        headers=analyst_headers,
    )
    assert result.status_code == 200, result.text
    assert result.json()["source_job_id"] == str(job.id)
    assert result.json()["selected"]["template_id"] == str(template.id)
    assert result.json()["selected"]["schema_version"] == 2
    denied = client.post(
        "/api/document-analysis/documents/runs",
        params={"document_iri": doc.iri, "template_id": str(template.id)},
        json={"request_key": "blocked"},
        headers=analyst_headers,
    )
    assert denied.status_code == 409, denied.text
    assert db.scalar(select(func.count()).select_from(DocumentAnalysisRun)) == 0
    assert (
        client.get(
            "/api/ast-templates/recognition-context",
            params={"document_iri": doc.iri, "template_id": str(uuid4())},
            headers=analyst_headers,
        ).status_code
        == 422
    )


def test_finder_slot_configuration_mock_preview_and_formal_boundary(
    client,
    db,
    analyst_headers,
    operator_headers,
    finder_setup,
    monkeypatch,
    tmp_path,
):
    from app.models.mock_data import MockTeamMember
    from app.services.reporting.slot_semantics import apply_patch
    from app.services.reporting.template_v2 import TemplateV2

    template, job, *_ = finder_setup
    monkeypatch.setattr("app.services.reporting.template_preparation.check_plan", lambda *args: [])
    monkeypatch.setattr("app.services.reasoning.rule_service.check_plan", lambda *args: [])
    classes = {
        ROOT: {
            "parents": [],
            "properties": [],
            "relationships": [
                {"iri": DEV + "describes", "range": [DRUG + "DrugProduct"], "label": "产品"}
            ],
        },
        DRUG + "DrugProduct": {
            "parents": [],
            "relationships": [],
            "properties": [
                {
                    "iri": DRUG + "appearance",
                    "label": "性状",
                    "datatype": "http://www.w3.org/2001/XMLSchema#string",
                }
            ],
        },
    }
    monkeypatch.setattr(
        "app.services.extraction.extraction_tasks.semantic_schema_from_engine", lambda _: classes
    )
    monkeypatch.setattr(
        "app.services.reporting.report_run_service.ARTIFACT_ROOT", tmp_path / "reports"
    )
    monkeypatch.setattr(settings, "llm_suggest_slots_enabled", False)
    template.schema_json = TemplateV2.model_validate(
        {
            **template.schema_json,
            "ontology_release_ref": "auto:ontology",
            "style_profile_ref": "urn:report:style:standard:1",
            "publication_policy_ref": "urn:report:policy:qa-review:1",
            "source_slots": [{"source_slot_id": "source", "kind": "document", "class_iri": ROOT}],
            "sections": [
                {
                    "section_id": "s",
                    "groups": [
                        {
                            "group_id": "g",
                            "units": [
                                {
                                    "output_id": "u",
                                    "title": "评估组",
                                    "render": {
                                        "kind": "narrative",
                                        "mode": "composed",
                                        "nodes": [],
                                    },
                                }
                            ],
                        }
                    ],
                }
            ],
        }
    ).model_dump(mode="json")
    db.add(
        MockTeamMember(
            team_type="assessment",
            role_code="qa",
            role_class_iri="urn:qa",
            role_label="质量",
            name="隔离测试员",
            department="QA",
        )
    )
    db.commit()
    base = f"/api/ast-templates/{template.id}"
    options = client.get(base + "/semantic-sources", headers=analyst_headers)
    assert options.status_code == 200, options.text
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 0
    payload = {
        "draft_schema": template.schema_json,
        "source_job_id": str(job.id),
        "choice": {
            "output_id": "u",
            "sources": [{"source_key": "mock:assessment_team", "fields": ["name", "role"]}],
            "render": "table",
        },
    }
    assert (
        client.post(base + "/configure-slot", json=payload, headers=operator_headers).status_code
        == 403
    )
    response = client.post(base + "/configure-slot", json=payload, headers=analyst_headers)
    assert response.status_code == 200, response.text
    draft = apply_patch(
        TemplateV2.model_validate(template.schema_json), response.json()
    ).model_dump(mode="json")
    started = client.post(
        route(template, job),
        headers=analyst_headers,
        json={"request_key": "explicit-report-finder"},
    )
    assert started.status_code == 202, started.text
    source_binding = {
        "source": {"job_id": str(job.id), "finder_execution_id": started.json()["execution_id"]}
    }
    result = client.post(
        "/api/report-previews",
        headers=analyst_headers,
        json={
            "template_id": str(template.id),
            "draft_schema": draft,
            "mode": "report",
            "idempotency_key": "mock-preview",
            "source_bindings": source_binding,
        },
    )
    assert result.status_code == 200, result.text
    report = result.json()
    assert report["demonstration"] is True
    history = client.get(f"/api/extraction/jobs/{job.id}/reports", headers=analyst_headers)
    assert history.status_code == 200
    assert (
        next(r for r in history.json() if r["report_run_id"] == report["run_id"])["demonstration"]
        is True
    )
    assert "隔离测试员" in str(report["body_ast"])
    assert report["artifacts"]
    from app.models.extraction import GeneratedReport

    assert (
        db.scalar(
            select(GeneratedReport).where(GeneratedReport.report_run_id == report["run_id"])
        ).narratives["demonstration"]
        is True
    )
    stale = client.post(
        "/api/report-previews",
        headers=analyst_headers,
        json={
            "template_id": str(template.id),
            "draft_schema": draft,
            "mode": "data",
            "idempotency_key": "stale-finder",
            "source_bindings": {
                "source": {"job_id": str(job.id), "finder_execution_id": str(uuid4())}
            },
        },
    )
    assert stale.status_code == 409 and stale.json()["detail"]["code"] == "SOURCE_VERSION_CHANGED"
    template.status = "published"
    db.commit()
    formal = client.post(
        "/api/report-runs",
        headers=analyst_headers,
        json={
            "template_id": str(template.id),
            "purpose": "formal",
            "idempotency_key": "formal-forbidden",
        },
    )
    assert formal.status_code == 409, formal.text
    denied = client.post(
        f"/api/report-runs/{report['run_id']}/content-versions",
        headers=analyst_headers,
        json={
            "expected_body_hash": report["body_hash"],
            "attempt": report["attempt"],
            "idempotency_key": "deny-content",
        },
    )
    assert denied.status_code == 409, denied.text
    suggestion = client.post(
        base + "/suggest-semantics",
        headers=analyst_headers,
        json={"draft_schema": template.schema_json, "output_ids": ["u"]},
    )
    assert suggestion.status_code == 200, suggestion.text
    assert suggestion.json()["completion"] == "incomplete"
    monkeypatch.setattr(
        "app.services.reporting.slot_semantics.suggest",
        lambda *args, **kwargs: {
            "patches": [response.json()],
            "diagnostics": [],
            "completion": "complete",
        },
    )
    compiled_suggestion = client.post(
        base + "/suggest-semantics",
        headers=analyst_headers,
        json={"draft_schema": template.schema_json, "output_ids": ["u"]},
    )
    assert compiled_suggestion.status_code == 200, compiled_suggestion.text
    assert len(compiled_suggestion.json()["patches"]) == 1, compiled_suggestion.text
    assert compiled_suggestion.json()["completion"] == "complete"
    assert db.scalar(select(func.count()).select_from(AnnotationExecution)) == 1
