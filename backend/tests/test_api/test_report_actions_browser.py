"""Browser -> real Finder/report APIs -> actual DOCX, with disposable test data."""

import json
import os
import subprocess
import threading
from hashlib import sha256
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import select

from app.config import settings
from app.models.entity_shadow import EntityShadow
from app.models.extraction import ExtractionJob, GeneratedReport
from app.models.reporting import ReportRun
from app.services.reporting.demo_sources import PROMPT_REF
from app.services.reporting.template_v2 import TemplateV2
from tests.test_api.test_report_recognition_context import FILENAME, TEMPLATE_NAME
from tests.test_api.test_template_finder import DEV, DRUG, ROOT, route
from tests.test_api.test_template_finder import finder_setup as finder_setup  # noqa: F401
from tests.test_reporting.test_docx_layout import assert_black_text, styled_sample
from tests.test_reporting.test_output_contracts import template as example_template


@pytest.mark.skipif(
    os.environ.get("REPORT_ACTIONS_BROWSER") != "1",
    reason="Requires explicit browser opt-in, ESBUILD_MODULE and PLAYWRIGHT_MODULE",
)
def test_risk_report_browser_roundtrip(client, db, finder_setup, tmp_path, monkeypatch):
    template, source, _, _ = finder_setup
    source.source_filename = FILENAME
    template.name, template.version = TEMPLATE_NAME, "v2.6"
    # A different template default must never replace the report's source.
    other = ExtractionJob(source_type="word", source_filename="another-source.docx")
    db.add(other)
    db.flush()
    template.default_source_job_id = other.id
    db.add(EntityShadow(
        iri="urn:" + str(source.id), class_iri=ROOT, module="document",
        properties_json={"source_job_id": str(source.id)},
    ))
    value = json.dumps(example_template())
    for old, new in {
        "urn:Report": ROOT,
        "urn:Equipment": DRUG + "DrugProduct",
        "urn:uses": DEV + "describes",
        "urn:code": DRUG + "appearance",
        "urn:spec": DRUG + "appearance",
        "ontology-1": "auto:ontology",
        "style-1": "urn:report:style:standard:1",
        "policy-1": "urn:report:policy:qa-review:1",
    }.items():
        value = value.replace(old, new)
    schema = json.loads(value)
    schema.update(template_family_id=str(template.id), template_revision_id=str(template.id))
    sample_path = tmp_path / "report-template.docx"
    sample = styled_sample(sample_path)
    template.sample_docx_path = str(sample_path)
    schema["sections"][0]["groups"][0]["units"][0]["origin"] = {
        "document_hash": sha256(sample_path.read_bytes()).hexdigest(),
        "label_anchor": {"table_path": ["table:0"], "row_index": 0,
                         "column_index": 0, "paragraph_index": 0},
    }
    # Reproduce HRS-1597: a required narrative input has no recognized value.
    # Available table rows must still render, and the draft must retain the gap.
    schema["definitions"]["inputs"]["rows"]["projection"]["fields"]["spec"].update(
        value={"kind": "property", "property_iri": DRUG + "routeOfAdministration"},
        required=True,
    )
    schema["sections"][0]["groups"][0]["units"].append({
        "output_id": "overview", "title": "产品与生产计划概述",
        "bindings": [{"binding_ref": "equipment"}],
        "inputs": [{"input_ref": "rows", "alias": "rows", "required": True}],
        "render": {"kind": "narrative", "mode": "assisted", "prompt": {
            "policy_ref": PROMPT_REF, "input_refs": [{"input_id": "rows"}],
            "required_refs": [{"input_id": "rows"}],
        }},
    })
    template.schema_json = TemplateV2.model_validate(schema).model_dump(mode="json")
    db.commit()
    classes = {
        ROOT: {"parents": [], "properties": [], "relationships": [
            {"iri": DEV + "describes", "range": [DRUG + "DrugProduct"]},
        ]},
        DRUG + "DrugProduct": {"parents": [], "relationships": [], "properties": [
            {"iri": DRUG + "appearance", "datatype": "http://www.w3.org/2001/XMLSchema#string"},
            {"iri": DRUG + "routeOfAdministration",
             "datatype": "http://www.w3.org/2001/XMLSchema#string"},
        ]},
    }
    monkeypatch.setattr(
        "app.services.extraction.extraction_tasks.semantic_schema_from_engine", lambda _: classes,
    )
    monkeypatch.setattr("app.services.reporting.template_preparation.check_plan", lambda *args: [])
    monkeypatch.setattr("app.services.reasoning.rule_service.check_plan", lambda *args: [])
    monkeypatch.setattr(
        "app.services.reporting.report_run_service.ARTIFACT_ROOT", tmp_path / "reports"
    )
    monkeypatch.setattr(settings, "llm_suggest_slots_enabled", False)
    monkeypatch.setattr(settings, "local_llm_enabled", True)

    def draft_provider(_system, payload, _policy, _budget):
        assert "（待补充）" in str(payload)
        return {"nodes": [{"kind": "input_ref", "input_id": "rows"}]}

    monkeypatch.setattr("app.services.reporting.narrative_renderer.local_provider", draft_provider)
    initial = client.post(route(template, source), json={"request_key": "existing-graph"},
                          headers={"X-User": "analyst", "X-Role": "senior_analyst"})
    assert initial.status_code == 202, initial.text
    requests = []

    # TestClient runs the real routes without app lifespan; no live credentials,
    # production database, shared artifacts, or external model calls are involved.
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            self.dispatch()

        def do_POST(self):
            self.dispatch()

        def dispatch(self):
            body = self.rfile.read(int(self.headers.get("Content-Length", 0)))
            result = client.request(self.command, self.path, content=body, headers={
                "X-User": self.headers.get("X-User", "analyst"),
                "X-Role": self.headers.get("X-Role", "senior_analyst"),
                "Content-Type": "application/json",
            })
            requests.append((self.command, self.path, result.status_code))
            self.send_response(result.status_code)
            self.send_header("Content-Type", result.headers.get("Content-Type", "application/json"))
            self.end_headers()
            self.wfile.write(result.content)

        def log_message(self, *_args):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    download = tmp_path / "downloaded.docx"
    frontend = Path(__file__).resolve().parents[3] / "frontend"
    try:
        result = subprocess.run(
            ["node", "tests/report-actions-browser.mjs"], cwd=frontend,
            env={**os.environ, "REPORT_ACTIONS_INTEGRATION": json.dumps({
                "origin": f"http://127.0.0.1:{server.server_port}",
                "template": str(template.id), "source": str(source.id), "download": str(download),
                "initialResult": True,
            })}, capture_output=True, text=True, timeout=90,
        )
        assert result.returncode == 0, result.stdout + result.stderr + repr(requests)
    finally:
        server.shutdown()
        thread.join(timeout=5)
        server.server_close()
    db.expire_all()
    reports = list(db.scalars(select(GeneratedReport)))
    assert len(reports) == 1 and reports[0].job_id == source.id
    run = db.get(ReportRun, reports[0].report_run_id)
    assert run.execution_status == "completed" and run.template_id == template.id
    assert run.material_status == "incomplete"
    assert run.source_bundle["sources"]["doc"]["kind"] == "finder_demo"
    assert run.source_bundle["sources"]["doc"]["execution_id"]
    document = Document(download)
    assert_black_text(document)
    assert document.tables[0].cell(1, 0).text == "重复值😀"
    assert any("待补充" in p.text for p in document.paragraphs)
    assert "SAMPLE" not in document.element.body.xml
    assert document.sections[0]._sectPr.xml == sample.sections[0]._sectPr.xml
    assert document.styles["Normal"].element.xml == sample.styles["Normal"].element.xml
    assert document.sections[0].header.tables[0].cell(0, 0).text == "模板页眉"
    assert document.sections[0].footer._element.xpath(".//w:fldSimple/@w:instr") == ["PAGE"]
    assert not any(code >= 400 for _, _, code in requests), requests
    assert sum(method == "POST" and path.endswith("/finder") for method, path, _ in requests) == 1
    assert sum(
        method == "POST" and path == "/api/report-previews" for method, path, _ in requests
    ) == 1
    assert sum("/artifacts/" in path for _, path, _ in requests) == 2
