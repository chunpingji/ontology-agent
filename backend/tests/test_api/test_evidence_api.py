import uuid

import pytest

from app.models.extraction import ExtractionJob


@pytest.fixture
def evidence_api_job(db, monkeypatch):
    from app.api import evidence

    monkeypatch.setattr(
        evidence,
        "semantic_schema_from_engine",
        lambda engine: {
            "urn:test:Drug": {
                "iri": "urn:test:Drug",
                "label": "药物",
                "properties": [
                    {"iri": "urn:test:strength", "datatype": "decimal"},
                ],
                "relationships": [{"iri": "urn:test:uses", "range": ["urn:test:Equipment"]}],
            },
            "urn:test:Equipment": {
                "iri": "urn:test:Equipment",
                "properties": [],
                "relationships": [],
            },
        },
    )
    job = ExtractionJob(id=uuid.uuid4(), source_type="word", status="completed")
    db.add(job)
    db.commit()
    return job


def manual_request(name="A", **changes):
    return {
        "request_key": name,
        "reason": "人工核对原始记录",
        "candidate": {
            "candidate_id": "client-id",
            "kind": "entity",
            "class_iri": "urn:test:Drug",
            "text": name,
            "validation_status": "passed",
            "review_status": "confirmed",
            "commit_status": "succeeded",
            "provenance": [
                {
                    "kind": "manual",
                    "actor": "forged",
                    "review_id": "fake",
                    "value": name,
                    "reason": "fake",
                }
            ],
            **changes,
        },
    }


def test_creation_stamps_identity_and_requires_separate_review(
    client, analyst_headers, operator_headers, evidence_api_job
):
    path = f"/api/extraction/jobs/{evidence_api_job.id}/evidence"
    assert (
        client.post(
            path + "/candidates", json=manual_request(), headers=operator_headers
        ).status_code
        == 403
    )
    response = client.post(path + "/candidates", json=manual_request(), headers=analyst_headers)
    assert response.status_code == 201, response.text
    candidate = response.json()
    assert candidate["validation_status"] == "passed"
    assert candidate["review_status"] == "pending" and candidate["commit_status"] == "not_requested"
    assert candidate["provenance"][0]["actor"] == "analyst"
    assert candidate["provenance"][0]["review_id"] != "fake"
    assert candidate["candidate_id"] != "client-id"
    commit = client.post(
        path + "/commits",
        headers=analyst_headers,
        json={
            "idempotency_key": "early",
            "items": [
                {"candidate_id": candidate["candidate_id"], "revision": 1},
            ],
        },
    )
    assert commit.status_code == 422
    review = f"/api/extraction/evidence/candidates/{candidate['candidate_id']}/review"
    decision = {"expected_revision": 1, "decision": "confirmed", "reason": "已核对"}
    response = client.put(review, headers=analyst_headers, json=decision)
    assert response.status_code == 200, response.text
    assert response.json()["commit_status"] == "not_requested"
    assert client.put(review, headers=analyst_headers, json=decision).status_code == 409
    assert client.get(path + "/snapshot", headers=analyst_headers).json()["published"] is False
    repeated = client.post(path + "/candidates", json=manual_request(), headers=analyst_headers)
    assert repeated.json()["candidate_id"] == candidate["candidate_id"]


def test_client_cannot_forge_source_or_instance_identity(client, analyst_headers, evidence_api_job):
    path = f"/api/extraction/jobs/{evidence_api_job.id}/evidence/candidates"
    response = client.post(
        path,
        headers=analyst_headers,
        json=manual_request(identity={"instance_iri": "urn:test:Drug"}),
    )
    assert response.status_code == 422
    request = manual_request(
        provenance=[
            {
                "kind": "external_record",
                "system": "unregistered",
                "dataset": "private",
                "record_key": "1",
                "record_version": "fake",
                "field_path": "label",
                "value": "A",
                "record_snapshot": {"label": "A"},
            }
        ]
    )
    assert client.post(path, headers=analyst_headers, json=request).status_code == 422


def test_same_manual_key_with_changed_content_conflicts(client, analyst_headers, evidence_api_job):
    path = f"/api/extraction/jobs/{evidence_api_job.id}/evidence/candidates"
    assert client.post(path, headers=analyst_headers, json=manual_request()).status_code == 201
    request = manual_request()
    request["candidate"]["text"] = "different"
    response = client.post(path, headers=analyst_headers, json=request)
    assert response.status_code == 409, response.text


def test_manual_property_requires_endpoint_and_exact_literal(
    client, analyst_headers, evidence_api_job
):
    path = f"/api/extraction/jobs/{evidence_api_job.id}/evidence/candidates"
    entity = client.post(path, headers=analyst_headers, json=manual_request()).json()
    request = {
        "request_key": "value",
        "reason": "人工规格录入",
        "candidate": {
            "kind": "property",
            "subject": {"candidate_id": entity["candidate_id"], "revision": 1},
            "predicate_iri": "urn:test:strength",
            "literal": {
                "kind": "number",
                "raw_value": "250",
                "normalized_value": "999",
                "datatype_iri": "http://www.w3.org/2001/XMLSchema#decimal",
            },
            "provenance": [
                {
                    "kind": "manual",
                    "actor": "fake",
                    "review_id": "fake",
                    "reason": "fake",
                    "value": "250",
                }
            ],
        },
    }
    response = client.post(path, headers=analyst_headers, json=request)
    assert response.status_code == 201, response.text
    assert response.json()["literal"]["normalized_value"] == "250"
    assert response.json()["bindings"][0]["method"] == "manual_decision"
    request["request_key"] = "missing"
    request["candidate"]["subject"]["candidate_id"] = "missing"
    assert client.post(path, headers=analyst_headers, json=request).status_code == 422


def test_legacy_dictionary_cannot_fake_successful_instance_commit(
    client, analyst_headers, db, evidence_api_job, fake_engine
):
    from app.models.extraction import ExtractionCandidate

    candidate = ExtractionCandidate(
        job_id=evidence_api_job.id,
        target_class_iri="urn:test:Drug",
        extracted_properties={"strength": "250"},
        review_status="pending",
    )
    db.add(candidate)
    db.commit()
    result = client.put(
        f"/api/extraction/candidates/{candidate.id}/review",
        headers=analyst_headers,
        json={"status": "confirmed"},
    )
    assert result.status_code == 409 and "legacy_unverified" in result.text
    db.refresh(candidate)
    assert candidate.review_status == "pending" and candidate.committed_iri is None
    assert not fake_engine.projected


def test_word_extraction_persists_shared_ir_and_never_calls_legacy_pipeline(
    client,
    analyst_headers,
    db,
    evidence_api_job,
    monkeypatch,
    tmp_path,
):
    from docx import Document

    from app.api import extraction
    from app.config import settings
    from app.models.evidence import DocumentAnalysisRecord, EvidenceJobState
    from app.services.extraction.word_analysis import analyze_word_core

    doc = Document()
    doc.add_heading("产品 A", 1)
    doc.add_paragraph("规格：250 mg")
    path = tmp_path / "source.docx"
    doc.save(path)
    evidence_api_job.document_path = str(path)
    db.commit()
    monkeypatch.setattr(settings, "local_llm_enabled", False)
    monkeypatch.setattr(
        extraction, "_annotation_cache_path", lambda job: tmp_path / "annotation.json"
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("legacy pipeline must not run")

    monkeypatch.setattr(extraction, "run_extraction_pipeline", forbidden)
    response = client.post(
        f"/api/extraction/jobs/{evidence_api_job.id}/evidence/extract", headers=analyst_headers
    )
    assert response.status_code == 200, response.text
    assert response.json()["run"]["completion"] == "incomplete"
    state = db.get(EvidenceJobState, evidence_api_job.id)
    ir = db.get(DocumentAnalysisRecord, state.analysis_id)
    assert ir.structure_hash == analyze_word_core(path).ir.structure_hash
    assert response.json()["snapshot_id"] is None
    # Coverage now exposes unresolved tasks; reports still require published facts.
    report = client.post(
        f"/api/extraction/jobs/{evidence_api_job.id}/risk-report", headers=analyst_headers
    )
    coverage = client.get(
        f"/api/extraction/jobs/{evidence_api_job.id}/ast-coverage", headers=analyst_headers
    )
    assert report.status_code == 422
    assert "published fact snapshot" in report.text
    assert coverage.status_code == 200, coverage.text
    assert coverage.json()["snapshot_id"] is None
    assert coverage.json()["filled"] == 0

    endpoint = f"/api/extraction/jobs/{evidence_api_job.id}/evidence/extract"
    assert client.post(endpoint, headers=analyst_headers, json={
        "retry_failed": True, "reason": " ",
    }).status_code == 422
    assert client.post(endpoint, headers=analyst_headers, json={
        "retry_failed": True, "reason": "重试", "checkpoint": {},
    }).status_code == 422
    called = []
    original = extraction._compute_annotation

    def inspect(*args, **kwargs):
        called.append(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(extraction, "_compute_annotation", inspect)
    response = client.post(endpoint, headers=analyst_headers, json={
        "retry_failed": True, "reason": "已恢复本地模型，显式重试失败任务", "pause_after": 8,
    })
    assert response.status_code == 200, response.text
    assert called[0]["retry_failed"] is True
    assert called[0]["pause_after"] == 8
    from app.models.reasoning import AuditLog

    entry = db.query(AuditLog).filter_by(action="evidence.extract_retry").one()
    assert entry.actor == "analyst" and entry.details["reason"]
