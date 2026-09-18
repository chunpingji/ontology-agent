"""Sample-to-template persistence and source workflow metadata contracts."""

from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document

from app.models.extraction import AstTemplate, ExtractionJob

CMC = "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"


@pytest.mark.parametrize("mode", [None, "doc_repo_preview", "template_default"])
def test_job_reads_expose_workflow_without_source_config(client, db, analyst_headers, mode):
    job = ExtractionJob(
        source_type="word", source_filename="input.docx", status="pending",
        source_config={"mode": mode, "private_configuration": "must-not-be-exposed"},
    )
    db.add(job)
    db.commit()
    for url in (f"/api/extraction/jobs/{job.id}", "/api/extraction/jobs"):
        response = client.get(url, headers=analyst_headers)
        assert response.status_code == 200, response.text
        payload = response.json()
        if isinstance(payload, list):
            payload = next(item for item in payload if item["id"] == str(job.id))
        assert payload["source_mode"] == mode
        assert "source_config" not in payload
        assert "must-not-be-exposed" not in response.text


@pytest.mark.parametrize("include_content", [True, False])
def test_sample_can_be_saved_as_draft_and_reopened(
    client, db, analyst_headers, tmp_path, monkeypatch, include_content,
):
    from app.api import ast_templates
    from app.services.extraction import extraction_tasks
    from tests.test_reporting.test_output_contracts import ontology

    classes = ontology()
    classes[CMC] = classes.pop("urn:Report")
    monkeypatch.setattr(extraction_tasks, "semantic_schema_from_engine", lambda _: classes)
    monkeypatch.setattr(ast_templates, "_UPLOADS", tmp_path)
    document = Document()
    document.add_heading("输出报告样例", level=1)
    document.add_paragraph("评估结论填写于此。")
    stream = BytesIO()
    document.save(stream)
    raw = stream.getvalue()
    upload = {"file": ("output.docx", raw,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    parsed = client.post("/api/ast-templates/parse-sample", headers=analyst_headers, files=upload)
    assert parsed.status_code == 200, parsed.text
    sample = parsed.json()
    schema = {
        "schema_version": 2, "template_family_id": "draft", "template_revision_id": "draft",
        "revision_no": 1, "doc_no": "EDITED-01", "ontology_release_ref": "auto:ontology",
        "style_profile_ref": "urn:report:style:standard:1",
        "publication_policy_ref": "urn:report:policy:qa-review:1",
        "source_slots": [{"source_slot_id": "source", "kind": "document", "class_iri": CMC}],
        "definitions": {"bindings": {}, "inputs": {}}, "sections": [],
    }
    request = {
        "name": "CMC 输出样例模板", "version": "v1", "iri_pattern": CMC,
        "doc_no": "EDITED-01", "schema_json": schema,
    }
    response = client.post("/api/ast-templates", headers=analyst_headers, json=request)
    assert response.status_code == 201, response.text
    created = response.json()
    template_id = created["id"]
    url = f"/api/ast-templates/{template_id}/sample"
    if not include_content:
        url += "?include_content=false"
    attached = client.post(url, headers=analyst_headers, files=upload)
    if include_content:
        assert attached.status_code == 200, attached.text
        assert attached.json() == sample
    else:
        assert attached.status_code == 204, attached.text
        assert attached.content == b""
    detail = client.get(f"/api/ast-templates/{template_id}", headers=analyst_headers)
    assert detail.status_code == 200, detail.text
    saved = detail.json()
    assert saved["status"] == "draft"
    assert saved["iri_pattern"] == CMC
    assert saved["doc_no"] == saved["schema_json"]["doc_no"] == "EDITED-01"
    assert saved["schema_json"]["source_slots"][0]["class_iri"] == CMC
    assert saved["sample_content_json"] == sample["content_json"]
    assert "输出报告样例" in saved["sample_text"]
    row = db.get(AstTemplate, UUID(template_id))
    assert Path(row.sample_docx_path).read_bytes() == raw
    assert Path(row.sample_docx_path).is_relative_to(tmp_path)
    listing = client.get("/api/ast-templates", headers=analyst_headers)
    assert any(item["id"] == template_id for item in listing.json())
    duplicate = client.post("/api/ast-templates", headers=analyst_headers, json=request)
    assert duplicate.status_code == 409
    assert db.query(AstTemplate).filter(AstTemplate.name == request["name"]).count() == 1
    assert db.query(ExtractionJob).count() == 0


def test_sample_acknowledgement_requires_commit(client, db, analyst_headers, tmp_path, monkeypatch):
    from app.api import ast_templates
    from tests.test_extraction.test_ast_template_basic_info import _docx_bytes, _seed_template

    uploads = tmp_path / "uploads"
    uploads.mkdir()
    monkeypatch.setattr(ast_templates, "_UPLOADS", uploads)
    row = _seed_template(db)
    old_path = uploads / "old.docx"
    old_path.write_bytes(b"old sample")
    row.sample_docx_path = str(old_path)
    db.commit()
    template_id = row.id

    def fail_commit():
        raise RuntimeError("sample commit failed")

    monkeypatch.setattr(db, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="sample commit failed"):
        client.post(
            f"/api/ast-templates/{template_id}/sample?include_content=false",
            headers=analyst_headers,
            files={"file": ("output.docx", _docx_bytes(), "application/octet-stream")},
        )
    db.refresh(row)
    assert row.sample_docx_path == str(old_path)
    assert old_path.read_bytes() == b"old sample"
    assert list(uploads.iterdir()) == [old_path]
