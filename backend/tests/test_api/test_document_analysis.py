"""Document-analysis runs API integration contract."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document

from app.api import document_analysis
from app.config import settings
from app.dependencies import get_ontology_engine
from app.main import app
from app.models.document_analysis import (
    DocumentAnalysisArtifact,
    DocumentAnalysisControlOperation,
    DocumentAnalysisRun,
    DocumentAnalysisTombstone,
    DocumentRecognitionEvent,
    DocumentRunCandidate,
    DocumentVerificationProof,
    DocumentVerificationProofHead,
)
from app.models.extraction import ExtractionJob
from app.services.document_analysis import retention
from app.services.document_analysis.run_store import (
    DocumentAnalysisRunStore,
    RunDeleted,
)
from app.services.extraction.ontology_guided import model_adapter

ROOT_IRI = "https://ontology.pharma-gmp.cn/slpra/drug-development/CMCReport"
SECOND_ROOT_IRI = "https://ontology.pharma-gmp.cn/slpra/drug-development/DrugProduct"
USES_EQUIPMENT = "https://ontology.pharma-gmp.cn/slpra/drug-development/usesEquipment"
EQUIPMENT = "https://ontology.pharma-gmp.cn/slpra/equipment/Equipment"


def _word_bytes(tmp_path: Path, text: str = "本报告明确描述产品甲。") -> bytes:
    document = Document()
    document.add_heading("第一章 产品概要", level=1)
    document.add_paragraph(text)
    document.add_heading("1.1 质量属性", level=2)
    document.add_paragraph("外观为白色片剂。")
    path = tmp_path / "source.docx"
    document.save(path)
    return path.read_bytes()


def _create(
    client,
    headers,
    raw: bytes,
    *,
    request_key: str = "request-1",
    root_class_iri: str = ROOT_IRI,
):
    return client.post(
        "/api/document-analysis/runs",
        headers=headers,
        files={
            "file": (
                "CMC报告.docx",
                raw,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={
            "root_class_iri": root_class_iri,
            "request_key": request_key,
            "metadata_mode": "structure_only",
        },
    )


def test_current_pause_and_resume_do_not_rewrite_checkpoint_coverage(
    client, db, analyst_headers, tmp_path, monkeypatch,
):
    from copy import deepcopy

    from app.models.document_analysis import DocumentRunArtifactHead

    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "artifacts")
    created = _create(client, analyst_headers, _word_bytes(tmp_path), request_key="pause-view")
    run_id = created.json()["recognition_run_id"]
    run = db.get(DocumentAnalysisRun, run_id)
    head = db.get(DocumentRunArtifactHead, (run_id, "graph"))
    artifact = db.get(DocumentAnalysisArtifact, head.artifact_id)
    frozen_payload = deepcopy(artifact.payload)
    graph_url = f"/api/document-analysis/runs/{run_id}/graph"
    original = client.get(graph_url, headers=analyst_headers).json()

    run.execution_status = "paused"
    run.stop_reason = "ranking_paused"
    # Reproduce a graph/checkpoint batch that predates the terminal run status.
    run.progress = {**run.progress, "stop_reason": "attempted_incomplete"}
    db.commit()
    paused = client.get(graph_url, headers=analyst_headers).json()
    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers).json()
    assert paused["coverage"]["stop_reason"] == "ranking_paused"
    assert status["progress"]["stop_reason"] == "ranking_paused"
    assert {**paused["coverage"], "stop_reason": original["coverage"]["stop_reason"]} == (
        original["coverage"]
    )

    run.execution_status = "running"
    run.stop_reason = None
    db.commit()
    resumed = client.get(graph_url, headers=analyst_headers).json()
    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers).json()
    assert resumed["coverage"]["stop_reason"] is None
    assert status["progress"]["stop_reason"] is None
    db.refresh(artifact)
    assert artifact.payload == frozen_payload
    assert resumed["graph_snapshot"] == original["graph_snapshot"]


@pytest.mark.parametrize("sparse_candidates", [False, True])
def test_run_builds_shared_metadata_and_honest_partial_graph(
    client, db, analyst_headers, tmp_path, monkeypatch, sparse_candidates,
):
    storage = tmp_path / "run-artifacts"
    monkeypatch.setattr(settings, "document_analysis_storage_dir", storage)
    monkeypatch.setattr(settings, "document_analysis_evidence_repair_enabled", sparse_candidates)
    monkeypatch.setattr(settings, "document_analysis_adaptive_retrieval_mode", "disabled")

    raw = _word_bytes(tmp_path)
    created = _create(client, analyst_headers, raw)

    assert created.status_code == 202, created.text
    accepted = created.json()
    assert accepted["status"] == "queued"
    assert accepted["stage"] == "accepted"
    assert accepted["artifact_revision"] == 1
    run_id = accepted["recognition_run_id"]

    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers)
    assert status.status_code == 200, status.text
    run = status.json()
    assert run["status"] == "retryable_failure"
    assert run["stage"] == "extracting"
    assert run["artifacts"] == {
        "source": "ready",
        "structure": "ready",
        "metadata": "ready",
        "graph": "partial",
    }
    assert run["identities"]["analysis_id"]
    assert run["identities"]["metadata_snapshot_id"]
    if sparse_candidates:
        # No model service means no admission/execution took place. Search scope
        # must not be fabricated into candidate tasks, or reported as complete.
        assert run["progress"]["candidate_policy"] == "sparse-candidates-v1"
        assert run["progress"]["completion"] == "incomplete"
        assert run["progress"]["records_planned"] == 0
        assert run["progress"]["records_unattempted"] == 0
    else:
        assert "candidate_policy" not in run["progress"]
        assert run["progress"]["records_unattempted"] > 0
    assert run["progress"]["stop_reason"] == "recognition_model_not_configured"
    assert run["expires_at"] is not None

    metadata = client.get(f"/api/document-analysis/runs/{run_id}/metadata", headers=analyst_headers)
    assert metadata.status_code == 200, metadata.text
    metadata_payload = metadata.json()
    assert metadata_payload["availability"] == "ready"
    assert metadata_payload["content"]["type"] == "doc"
    assert metadata_payload["metadata_snapshot"]["generation_source"] == "structure_only"
    assert metadata_payload["analysis"]["analysis_id"] == run["identities"]["analysis_id"]

    graph = client.get(f"/api/document-analysis/runs/{run_id}/graph", headers=analyst_headers)
    assert graph.status_code == 200, graph.text
    graph_payload = graph.json()
    assert graph_payload["availability"] == "partial"
    assert graph_payload["projection"] == "effective_affirmed"
    assert len(graph_payload["entities"]) == 1
    assert graph_payload["entities"][0]["seed_origin"] == "user_selected"
    assert graph_payload["relationships"] == []
    if sparse_candidates:
        assert graph_payload["coverage"]["candidate_policy"] == "sparse-candidates-v1"
        assert graph_payload["coverage"]["records_planned"] == 0
        assert graph_payload["coverage"]["records_examined"] == 0
    else:
        assert "candidate_policy" not in graph_payload["coverage"]
        assert graph_payload["coverage"]["records_unattempted"] > 0
    assert graph_payload["coverage"]["stop_reason"] == "service_failure"

    source = client.get(f"/api/document-analysis/runs/{run_id}/source", headers=analyst_headers)
    assert source.status_code == 200, source.text
    assert source.json()["analysis_id"] == run["identities"]["analysis_id"]
    original = client.get(
        f"/api/document-analysis/runs/{run_id}/source?format=original",
        headers=analyst_headers,
    )
    assert original.status_code == 200
    assert original.content == raw
    assert original.headers["x-document-hash"] == db.get(DocumentAnalysisRun, run_id).document_hash

    assert db.query(ExtractionJob).count() == 0
    assert (
        client.post(
            "/api/document-analysis/word",
            headers=analyst_headers,
            files={"file": ("old.docx", b"old")},
        ).status_code
        == 404
    )


def test_configured_model_persists_effective_proof_and_replays_role_sources(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    """Exercise the configured adapter path without contacting an external model."""

    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    # The response fixture below implements the original adapter's proof protocol.
    monkeypatch.setattr(settings, "document_analysis_evidence_repair_enabled", False)
    monkeypatch.setattr(settings, "document_analysis_adaptive_retrieval_mode", "disabled")

    def deterministic_model(_client, *, user, **_kwargs):
        request = json.loads(user)
        target = next(
            (
                fragment
                for fragment in request["fragments"]
                if "本报告明确使用设备冻干机A" in fragment["text"]
            ),
            None,
        )
        if request["predicate"]["iri"] != USES_EQUIPMENT or target is None:
            return {"proposals": []}
        evidence_id = target["evidence_id"]
        subject_support = (
            []
            if request["subject"]["is_document_root"]
            else [{"evidence_id": evidence_id, "text": "本报告"}]
        )
        if request.get("stage") == "verification":
            return {"verifications": [{
                "candidate_id": candidate["candidate_id"],
                "target_id": candidate["target_id"],
                "type_verdict": "supported",
                "type_support": [{"evidence_id": evidence_id}],
                "role_verdict": "supported",
                "subject_binding_verdict": "supported",
                "predicate_verdict": "supported",
                "applicability_verdict": "supported",
                "counterevidence_verdict": "undetermined",
                "bridge_verdict": "supported",
                "predicate_support": [{"evidence_id": evidence_id, "text": "使用设备"}],
                "subject_support": subject_support,
                "condition_support": [],
                "counterevidence_support": [],
                "reason": "独立原文核验确认本报告使用冻干机A。",
            } for candidate in request["candidates"]]}
        return {
            "proposals": [
                {
                    "kind": "relationship",
                    "object_class_iri": EQUIPMENT,
                    "object_label": "冻干机A",
                    "object_quote": {"evidence_id": evidence_id, "text": "冻干机A"},
                    "predicate_support": [{"evidence_id": evidence_id, "text": "使用设备"}],
                    "subject_support": subject_support,
                    "bridge_kind": "explicit_assertion",
                    "type_verdict": "supported",
                    "role_verdict": "supported",
                    "predicate_verdict": "supported",
                    "applicability_verdict": "supported",
                    "polarity": "affirmed",
                    "reason": "原文明示本报告使用冻干机A。",
                }
            ]
        }

    monkeypatch.setattr(model_adapter, "chat_with_schema", deterministic_model)
    monkeypatch.setattr(
        model_adapter._ConfiguredInputCounter, "count", lambda _self, text: len(text)
    )
    monkeypatch.setattr(model_adapter, "get_local_llm", lambda: object())
    monkeypatch.setattr(settings, "local_llm_model", "deterministic-model")
    monkeypatch.setattr(settings, "local_llm_model_revision", "test-revision")
    raw = _word_bytes(tmp_path, "本报告明确使用设备冻干机A进行生产。")

    created = _create(
        client,
        analyst_headers,
        raw,
        request_key="configured-model-proof",
    )
    assert created.status_code == 202, created.text
    run_id = created.json()["recognition_run_id"]

    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers)
    assert status.status_code == 200, status.text
    assert status.json()["status"] == "finished", status.text

    graph_response = client.get(
        f"/api/document-analysis/runs/{run_id}/graph", headers=analyst_headers
    )
    assert graph_response.status_code == 200, graph_response.text
    graph = graph_response.json()
    assert graph["availability"] == "ready"
    assert len(graph["relationships"]) == 1
    relationship = graph["relationships"][0]
    assert relationship["predicate_iri"] == USES_EQUIPMENT
    assert relationship["structural_valid"] is True
    assert relationship["model_supported"] is True
    assert relationship["policy_eligible"] is True
    assert relationship["proof_ref"] is not None

    proof_key = (run_id, relationship["proof_ref"]["id"], 1)
    proof = db.get(DocumentVerificationProof, proof_key)
    assert proof is not None
    assert proof.payload["predicate_evidence"]["predicate_iri"] == USES_EQUIPMENT
    assert proof.payload["predicate_evidence"]["proof_id"] == relationship["proof_ref"]["id"]
    proof_head = db.get(
        DocumentVerificationProofHead,
        (run_id, relationship["proof_ref"]["id"]),
    )
    assert proof_head is not None
    assert proof_head.proof_revision == 1

    role_refs = relationship["source_selection_refs"]
    assert role_refs["subject"] == []
    root_ref = graph["graph_snapshot"]["root_ref"]
    assert relationship["subject_ref"] == root_ref
    root = next(item for item in graph["entities"] if item["entity_id"] == root_ref["entity_id"])
    assert root["revision"] == root_ref["revision"] == 1
    assert root["class_iri"] == ROOT_IRI
    assert root["seed_origin"] == "user_selected"
    assert root["source_selection_refs"] == []
    assert role_refs["object"]
    assert role_refs["predicate_bridge"]
    for role in ("object", "predicate_bridge"):
        selection_ref = role_refs[role][0]
        replay = client.get(
            f"/api/document-analysis/runs/{run_id}/source",
            headers=analyst_headers,
            params={"selection_ref": selection_ref},
        )
        assert replay.status_code == 200, replay.text
        replayed = replay.json()
        assert replayed["selection"]["selection_ref"] == selection_ref
        assert replayed["selection"]["selection_role"] == role
        assert replayed["anchors"]
        assert replayed["anchors"][0]["document_hash"] == replayed["document_hash"]
        located = client.get(
            f"/api/document-analysis/runs/{run_id}/source-selection",
            headers=analyst_headers, params={"selection_ref": selection_ref},
        )
        assert located.status_code == 200, located.text
        assert located.json() == {
            key: value for key, value in replayed.items() if key not in {"content", "filename"}
        }


def test_create_is_owner_scoped_idempotent_and_rejects_changed_input(
    client, analyst_headers, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    raw = _word_bytes(tmp_path)
    first = _create(client, analyst_headers, raw, request_key="same-key")
    replay = _create(client, analyst_headers, raw, request_key="same-key")
    conflict = _create(
        client,
        analyst_headers,
        _word_bytes(tmp_path, "不同内容。"),
        request_key="same-key",
    )

    assert first.status_code == replay.status_code == 202
    assert replay.json()["recognition_run_id"] == first.json()["recognition_run_id"]
    assert replay.json()["idempotent_replay"] is True
    assert replay.json()["status"] == "retryable_failure"
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "IDEMPOTENCY_CONFLICT"

    hidden = client.get(
        f"/api/document-analysis/runs/{first.json()['recognition_run_id']}",
        headers={"X-User": "another", "X-Role": "senior_analyst"},
    )
    assert hidden.status_code == 404
    assert hidden.json()["error"]["code"] == "RUN_NOT_FOUND"


def test_create_rejects_absolute_root_iri_missing_from_current_ontology(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    unknown_root = "https://ontology.pharma-gmp.cn/slpra/drug-development/UnknownRoot"

    response = _create(
        client,
        analyst_headers,
        _word_bytes(tmp_path),
        request_key="unknown-root",
        root_class_iri=unknown_root,
    )

    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "INVALID_ROOT_CLASS"
    assert db.query(DocumentAnalysisRun).count() == 0


def test_same_upload_and_distinct_valid_root_types_create_distinct_runs(
    client, db, fake_engine, analyst_headers, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *_args, **_kwargs: None)

    def class_detail(class_iri):
        if class_iri != SECOND_ROOT_IRI:
            return None
        return SimpleNamespace(
            label_zh="制剂产品",
            label_en="Drug Product",
            name="DrugProduct",
            parent_iris=[],
            comment="",
        )

    monkeypatch.setattr(fake_engine, "get_class_detail", class_detail)
    raw = _word_bytes(tmp_path)
    first = _create(
        client,
        analyst_headers,
        raw,
        request_key="cmc-root-run",
        root_class_iri=ROOT_IRI,
    )
    second = _create(
        client,
        analyst_headers,
        raw,
        request_key="product-root-run",
        root_class_iri=SECOND_ROOT_IRI,
    )

    assert first.status_code == second.status_code == 202
    first_payload = first.json()
    second_payload = second.json()
    assert first_payload["recognition_run_id"] != second_payload["recognition_run_id"]
    first_run = db.get(DocumentAnalysisRun, first_payload["recognition_run_id"])
    second_run = db.get(DocumentAnalysisRun, second_payload["recognition_run_id"])
    assert first_run.document_hash == second_run.document_hash
    assert first_run.root_class_iri == ROOT_IRI
    assert second_run.root_class_iri == SECOND_ROOT_IRI
    assert first_run.request_key == "cmc-root-run"
    assert second_run.request_key == "product-root-run"


def test_queued_pause_and_resume_use_revision_and_operation_idempotency(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *_args, **_kwargs: None)
    accepted = _create(client, analyst_headers, _word_bytes(tmp_path)).json()
    run_id = accepted["recognition_run_id"]
    body = {
        "expected_revision": accepted["run_revision"],
        "request_key": "pause-once",
        "reason": "检查已提交结果",
    }

    paused = client.post(
        f"/api/document-analysis/runs/{run_id}/pause",
        headers=analyst_headers,
        json=body,
    )
    replay = client.post(
        f"/api/document-analysis/runs/{run_id}/pause",
        headers=analyst_headers,
        json=body,
    )
    stale = client.post(
        f"/api/document-analysis/runs/{run_id}/resume",
        headers=analyst_headers,
        json={**body, "request_key": "resume-stale"},
    )

    assert paused.status_code == replay.status_code == 202
    assert paused.json()["status"] == "paused"
    assert replay.json()["run_revision"] == paused.json()["run_revision"]
    receipt = db.get(DocumentAnalysisControlOperation, (run_id, "pause", "pause-once"))
    assert receipt is not None
    assert receipt.result_revision == paused.json()["run_revision"]
    assert receipt.result_payload == paused.json()
    assert receipt.result_payload_hash is not None
    paused_status = client.get(
        f"/api/document-analysis/runs/{run_id}", headers=analyst_headers
    ).json()
    assert paused_status["expires_at"] is not None
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "RUN_REVISION_CONFLICT"

    resumed = client.post(
        f"/api/document-analysis/runs/{run_id}/resume",
        headers=analyst_headers,
        json={
            "expected_revision": paused.json()["run_revision"],
            "request_key": "resume-once",
            "reason": "继续分析",
        },
    )
    assert resumed.status_code == 202, resumed.text
    assert resumed.json()["status"] == "queued"
    resumed_status = client.get(
        f"/api/document-analysis/runs/{run_id}", headers=analyst_headers
    ).json()
    assert resumed_status["expires_at"] is None

    late_replay = client.post(
        f"/api/document-analysis/runs/{run_id}/pause",
        headers=analyst_headers,
        json=body,
    )
    assert late_replay.status_code == 202, late_replay.text
    assert late_replay.json() == paused.json()

    current_status = client.get(
        f"/api/document-analysis/runs/{run_id}", headers=analyst_headers
    ).json()
    assert current_status["status"] == "queued"
    assert current_status["run_revision"] == resumed_status["run_revision"]
    assert current_status["event_head"] == resumed_status["event_head"]


def test_create_validates_role_type_size_and_root_before_acceptance(
    client, operator_headers, analyst_headers, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    monkeypatch.setattr(settings, "document_analysis_max_upload_bytes", 16)

    role = _create(client, operator_headers, b"not-used", request_key="role")
    unsupported = client.post(
        "/api/document-analysis/runs",
        headers=analyst_headers,
        files={"file": ("notes.txt", b"plain")},
        data={"root_class_iri": ROOT_IRI, "request_key": "type"},
    )
    oversized = client.post(
        "/api/document-analysis/runs",
        headers=analyst_headers,
        files={"file": ("large.docx", b"x" * 32)},
        data={"root_class_iri": ROOT_IRI, "request_key": "large"},
    )
    invalid_root = client.post(
        "/api/document-analysis/runs",
        headers=analyst_headers,
        files={"file": ("source.docx", _word_bytes(tmp_path))},
        data={"root_class_iri": "not-an-iri", "request_key": "root"},
    )

    assert role.status_code == 403
    assert role.json()["error"]["code"] == "ROLE_FORBIDDEN"
    assert unsupported.status_code == 415
    assert unsupported.json()["error"]["code"] == "UNSUPPORTED_SOURCE_TYPE"
    assert oversized.status_code == 413
    assert oversized.json()["error"]["code"] == "SOURCE_TOO_LARGE"
    assert invalid_root.status_code == 422
    assert invalid_root.json()["error"]["code"] == "INVALID_ROOT_CLASS"


def test_framework_validation_and_ontology_dependency_use_contract_errors(
    client, analyst_headers, tmp_path, monkeypatch
):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    missing_file = client.post(
        "/api/document-analysis/runs",
        headers=analyst_headers,
        data={"root_class_iri": ROOT_IRI, "request_key": "missing-file"},
    )
    assert missing_file.status_code == 400
    assert missing_file.json() == {
        "contract_version": "document-analysis-runs-v1",
        "error": {
            "code": "INVALID_REQUEST",
            "message": "请求字段非法",
            "retryable": False,
            "current_revision": None,
        },
    }

    unauthenticated = client.get(f"/api/document-analysis/runs/{uuid4()}")
    assert unauthenticated.status_code == 401
    assert unauthenticated.json()["error"]["code"] == "UNAUTHENTICATED"

    def unloaded_engine():
        raise RuntimeError("Ontology not loaded")

    app.dependency_overrides[get_ontology_engine] = unloaded_engine
    unavailable = _create(
        client,
        analyst_headers,
        _word_bytes(tmp_path),
        request_key="ontology-unavailable",
    )
    assert unavailable.status_code == 503
    payload = unavailable.json()
    assert payload["contract_version"] == "document-analysis-runs-v1"
    assert payload["error"] == {
        "code": "ONTOLOGY_UNAVAILABLE",
        "message": "当前本体不可用，未创建分析运行",
        "retryable": True,
        "current_revision": None,
    }


def test_delete_api_keeps_shared_snapshot_tombstone_and_fences_late_worker(
    client, db, analyst_headers, tmp_path, monkeypatch
):
    storage = tmp_path / "run-artifacts"
    monkeypatch.setattr(settings, "document_analysis_storage_dir", storage)
    monkeypatch.setattr(document_analysis, "dispatch_run", lambda *_args, **_kwargs: None)
    raw = _word_bytes(tmp_path)
    first = _create(client, analyst_headers, raw, request_key="delete-first").json()
    second = _create(client, analyst_headers, raw, request_key="delete-second").json()
    first_id = first["recognition_run_id"]
    second_id = second["recognition_run_id"]

    store = DocumentAnalysisRunStore(db)
    stale_token = store.claim(
        first_id,
        "analyst",
        actor="dispatcher",
        worker_id="late-worker",
    )
    db.commit()
    current = store.get_owned(first_id, "analyst")
    delete_expected_revision = current.revision
    shared_snapshot_id = current.artifact_manifest["ontology_snapshot"]["artifact_id"]

    deleted = client.delete(
        f"/api/document-analysis/runs/{first_id}",
        headers=analyst_headers,
        params={
            "expected_revision": delete_expected_revision,
            "request_key": "delete-operation-删除",
        },
    )
    assert deleted.status_code == 202, deleted.text
    assert deleted.json()["status"] == "deleting"
    first_delete_payload = deleted.json()
    first_delete_bytes = deleted.content

    db.expire_all()
    assert db.get(DocumentAnalysisRun, first_id) is None
    assert db.get(DocumentAnalysisTombstone, first_id) is not None
    assert db.get(DocumentAnalysisArtifact, shared_snapshot_id) is not None
    assert db.get(DocumentAnalysisRun, second_id) is not None
    assert not (storage / first_id / "source" / "original.docx").exists()
    assert (storage / second_id / "source" / "original.docx").exists()
    with pytest.raises(RunDeleted):
        store.assert_fence(first_id, "analyst", stale_token)

    tombstone = client.get(f"/api/document-analysis/runs/{first_id}", headers=analyst_headers)
    assert tombstone.status_code == 410
    assert tombstone.json()["error"]["code"] == "RUN_DELETED"

    marker = db.get(DocumentAnalysisTombstone, first_id)
    assert marker.delete_request_key == "delete-operation-删除"
    assert marker.delete_request_hash is not None
    assert marker.delete_result_payload_hash is not None
    assert marker.delete_result_payload == first_delete_payload
    assert set(marker.delete_result_payload) == {
        "contract_version",
        "recognition_run_id",
        "ranking_budget_enabled",
        "run_revision",
        "event_head",
        "artifact_revision",
        "status",
        "stage",
        "operation",
        "operation_status",
        "available_actions",
    }

    cleanup_calls = []
    monkeypatch.setattr(
        retention,
        "delete_run",
        lambda *_args, **_kwargs: cleanup_calls.append((_args, _kwargs)),
    )
    replay = client.delete(
        f"/api/document-analysis/runs/{first_id}",
        headers=analyst_headers,
        params={
            "expected_revision": delete_expected_revision,
            "request_key": "delete-operation-删除",
        },
    )
    assert replay.status_code == 202, replay.text
    assert replay.json() == first_delete_payload
    assert replay.content == first_delete_bytes
    assert cleanup_calls == []

    changed_content = client.delete(
        f"/api/document-analysis/runs/{first_id}",
        headers=analyst_headers,
        params={
            "expected_revision": delete_expected_revision + 1,
            "request_key": "delete-operation-删除",
        },
    )
    assert changed_content.status_code == 409
    assert changed_content.json()["error"]["code"] == "IDEMPOTENCY_CONFLICT"

    different_operation = client.delete(
        f"/api/document-analysis/runs/{first_id}",
        headers=analyst_headers,
        params={
            "expected_revision": delete_expected_revision,
            "request_key": "unknown-delete-operation",
        },
    )
    assert different_operation.status_code == 410
    assert different_operation.json()["error"]["code"] == "RUN_DELETED"

    non_writer_headers = {"X-User": "analyst", "X-Role": "operator"}
    for params in (
        {
            "expected_revision": delete_expected_revision,
            "request_key": "delete-operation-删除",
        },
        {
            "expected_revision": delete_expected_revision + 1,
            "request_key": "delete-operation-删除",
        },
        {
            "expected_revision": delete_expected_revision,
            "request_key": "unknown-delete-operation",
        },
    ):
        forbidden = client.delete(
            f"/api/document-analysis/runs/{first_id}",
            headers=non_writer_headers,
            params=params,
        )
        assert forbidden.status_code == 403
        assert forbidden.json()["error"]["code"] == "ROLE_FORBIDDEN"

    for params in (
        {
            "expected_revision": delete_expected_revision,
            "request_key": "delete-operation-删除",
        },
        {
            "expected_revision": delete_expected_revision + 1,
            "request_key": "delete-operation-删除",
        },
        {
            "expected_revision": delete_expected_revision,
            "request_key": "unknown-delete-operation",
        },
    ):
        hidden = client.delete(
            f"/api/document-analysis/runs/{first_id}",
            headers={"X-User": "another", "X-Role": "senior_analyst"},
            params=params,
        )
        assert hidden.status_code == 404
        assert hidden.json()["error"]["code"] == "RUN_NOT_FOUND"
    assert cleanup_calls == []


def test_get_tabs_and_etag_are_side_effect_free(client, db, analyst_headers, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_analysis_storage_dir", tmp_path / "run-artifacts")
    created = _create(
        client, analyst_headers, _word_bytes(tmp_path), request_key="read-only-get"
    ).json()
    run_id = created["recognition_run_id"]
    db.expire_all()
    before = db.get(DocumentAnalysisRun, run_id)
    watermark = (before.revision, before.event_head, before.artifact_revision)
    event_count = db.query(DocumentRecognitionEvent).count()
    candidate_count = db.query(DocumentRunCandidate).count()

    status = client.get(f"/api/document-analysis/runs/{run_id}", headers=analyst_headers)
    etag = status.headers["etag"]
    for suffix in ("/metadata", "/graph", "/source"):
        response = client.get(
            f"/api/document-analysis/runs/{run_id}{suffix}", headers=analyst_headers
        )
        assert response.status_code == 200, response.text
    cached = client.get(
        f"/api/document-analysis/runs/{run_id}",
        headers={**analyst_headers, "If-None-Match": etag},
    )
    assert cached.status_code == 304
    assert cached.content == b""

    db.expire_all()
    after = db.get(DocumentAnalysisRun, run_id)
    assert (after.revision, after.event_head, after.artifact_revision) == watermark
    assert db.query(DocumentRecognitionEvent).count() == event_count
    assert db.query(DocumentRunCandidate).count() == candidate_count
