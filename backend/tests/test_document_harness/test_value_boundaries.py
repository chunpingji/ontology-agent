"""Exact value spans and assertion review under an explicitly supplied subject type."""

from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from pydantic import ValidationError
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.observations import property_value
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.protocols import PropertyMapping, stage_schema
from app.services.document_harness.source import build_windows
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def inputs(tmp_path):
    document = Document()
    document.add_heading("例行记录", 1)
    document.add_paragraph("管理标识： AX-07")
    document.add_paragraph("其他对象： BY-09")
    document.add_paragraph("状态 N/A；示例 AX-07 / AX-07")
    path = tmp_path / "source.docx"
    document.save(path)
    ir = analyze_word_core(path).ir
    catalog = catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:boundary:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class .
        :code a owl:DatatypeProperty ; rdfs:domain :Report ;
          rdfs:range xsd:string ; rdfs:label "管理标识" .
        :title a owl:DatatypeProperty ; rdfs:domain :Report ;
          rdfs:range xsd:string ; rdfs:label "标题" .
    ''', format="turtle"), "urn:boundary:Report")
    return ir, catalog


def quote(sources, text):
    source = next(s for s in sources if text in s["text"])
    return {"source_id": source["source_id"], "text": text, "occurrence": 0}


@pytest.mark.parametrize("selection,error", [
    (None, None), ("AX-07", None),
    ("BY-09", "property_value_quote_outside_field"),
    ("invented", "source_quote_mismatch"),
])
def test_value_selection_keeps_original_observation_and_does_not_block_other_fields(
    inputs, selection, error,
):
    ir, catalog = inputs
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        if stage == "discover":
            return {
                "entities": [], "document_field_ids": [], "unowned_fields": [],
                "relation_hints": [], "complete": True,
                "document_source_fields": [
                    {"label": None, "value": quote(payload["sources"], text)}
                    for text in ("管理标识： AX-07", "例行记录")
                ],
            }
        if stage == "assertion_alignment":
            fields = {f["field_id"]: f for f in payload["subject"]["fields"]}
            result = {}
            for key in payload["property_field_ids"]:
                is_code = fields[key]["value"].startswith("管理标识")
                chosen = quote(payload["sources"], "AX-07" if selection == "invented"
                               else selection) if is_code and selection else None
                if selection == "invented" and chosen:
                    chosen["text"] = "invented"
                result[key] = {"mappings": [{
                    "predicate_iri": "urn:boundary:code" if is_code else "urn:boundary:title",
                    "value_component": "span" if chosen else "whole",
                    "value_quote": chosen, "confidence": 0.99,
                }], "reason": "按属性含义选取值，完整观察不变"}
            return {"properties": result, "relations": [], "complete": True}
        assert stage == "evidence_review"
        for candidate in payload["candidates"]:
            assert "class_definition" not in candidate
            assert candidate["subject"]["type_basis"] == "user_selected"
        return {"type_concerns": [], "judgments": {c["id"]: {
            "verdict": "accepted", "confidence": 0.99, "evidence": c["evidence"],
            "reason": "本项语义与主体归属有依据",
        } for c in payload["candidates"]}}

    engine = Engine(ir=ir, catalog=catalog, state={}, invoke=invoke,
                    save=lambda _: None, should_stop=lambda: False)
    engine.run()
    assert calls == ["discover", "assertion_alignment", "evidence_review"]
    properties = list(engine.state["properties"].values())
    assert any(p["value"] == "例行记录" and p["state"] == "accepted" for p in properties)
    original = next(f for f in engine.state["fields"].values()
                    if f["value"] == "管理标识： AX-07")
    assert original["value_evidence"][0]["text"] == "管理标识： AX-07"
    code = [p for p in properties if p["predicate_iri"] == "urn:boundary:code"]
    if error:
        assert not code
        assert any(error in str(o) for o in engine.state["observations"].values())
        return
    assert len(code) == 1 and code[0]["state"] == "accepted"
    assert code[0]["value"] == (selection or original["value"])
    assert code[0]["source_value"] == original["value"]
    ref = code[0]["value_evidence"][0]
    assert ir.unit(ref["source_id"]).text[ref["start"]:ref["end"]] == code[0]["value"]


@pytest.mark.parametrize("verdict", ["accepted", "rejected", "unresolved"])
def test_type_concern_preserves_subject_and_never_rewrites_assertion_verdict(inputs, verdict):
    ir, catalog = inputs
    window = build_windows(ir)[0]
    ref = window.resolve(ir, quote(window.sources, "例行记录"))
    root = {"id": "document", "label": "用户选定主体", "role": "document_root",
            "class_iri": catalog.root_class_iri, "state": "accepted", "evidence": [],
            "field_ids": ["f"], "window_id": None}
    state = {"entities": {"document": root}, "fields": {"f": {
        "id": "f", "label": "", "value": ref["text"], "missing": False,
        "evidence": [ref], "value_evidence": [ref],
    }}, "properties": {"p": {
        "id": "p", "subject_id": "document", "predicate_iri": "urn:boundary:title",
        "value": ref["text"], "source_value": ref["text"], "value_component": "whole",
        "value_evidence": [ref], "field_id": "f", "state": "candidate",
        "evidence": [ref], "window_id": window.id,
    }}, "cursor": {"main": {"window_index": 0, "stage": "evidence_review",
                            "windows_reviewed": 0}}}

    def invoke(stage, payload, schema):
        assert stage == "evidence_review"
        candidate = payload["candidates"][0]
        assert candidate["subject"]["type_basis"] == "user_selected"
        response = {"judgments": {candidate["id"]: {
            "verdict": verdict, "confidence": 0.99, "evidence": candidate["evidence"],
            "reason": "该标题属于章节，不是当前文档名称" if verdict == "rejected"
            else "本项单独核对结果",
        }}, "type_concerns": [{"entity_id": candidate["subject"]["entity_id"],
                               "evidence": candidate["evidence"], "reason": "独立类型疑点"}]}
        jsonschema.validate(response, schema)
        return response

    engine = Engine(ir=ir, catalog=catalog, state=state, invoke=invoke,
                    save=lambda _: None, should_stop=lambda: False)
    engine.evidence_review(window)
    assert engine.state["entities"]["document"] == root
    assert engine.state["properties"]["p"]["state"] == verdict
    concern = next(o for o in engine.state["observations"].values()
                   if o["reason"] == "独立类型疑点")
    assert concern["candidate_subject_ids"] == ["document"] and concern["evidence"]


@pytest.mark.parametrize("value,occurrence,error", [
    ("AX-07", None, "source_quote_ambiguous"),
    ("AX-07", 1, "property_value_quote_outside_field"),
    ("N/A", 0, "missing_source_value"),
])
def test_span_does_not_borrow_another_occurrence_or_turn_a_missing_marker_into_value(
    inputs, value, occurrence, error,
):
    ir, _ = inputs
    window = build_windows(ir)[0]
    source = next(s for s in window.sources if "状态 N/A" in s["text"])
    first = window.resolve(ir, {"source_id": source["source_id"], "text": value, "occurrence": 0})
    field = {"value": value, "missing": False, "value_evidence": [first]}
    mapping = PropertyMapping(predicate_iri="urn:boundary:code", value_component="span",
                              value_quote={"source_id": source["source_id"], "text": value,
                                           "occurrence": occurrence}, confidence=0.99)
    with pytest.raises(ValueError, match=error):
        property_value(field, mapping, window=window, ir=ir, confirmed=True)


def test_span_contract_requires_exact_quote_and_type_concerns_require_real_endpoint_support():
    schema = stage_schema("assertion_alignment", source_ids=["S1"], field_ids=["F1"],
                          property_iris=["urn:boundary:code"])
    mapping = {"predicate_iri": "urn:boundary:code", "value_component": "span",
               "value_quote": {"source_id": "S1", "text": "A"}, "confidence": 0.9}
    answer = {"properties": {"F1": {"mappings": [mapping], "reason": "原文取值"}},
              "relations": [], "complete": True}
    jsonschema.validate(answer, schema)
    for change in ({"value_quote": None}, {"value_component": "whole"}):
        wrong = {**mapping, **change}
        with pytest.raises(ValidationError):
            PropertyMapping.model_validate(wrong)
        invalid = deepcopy(answer)
        invalid["properties"]["F1"]["mappings"] = [wrong]
        with pytest.raises(jsonschema.ValidationError):
            jsonschema.validate(invalid, schema)
    review_schema = stage_schema("evidence_review", source_ids=["S1"], entity_ids=["E0"])
    valid = {"judgments": {}, "type_concerns": [
        {"entity_id": "E0", "evidence": ["S1"], "reason": "原文明确矛盾"},
    ]}
    jsonschema.validate(valid, review_schema)
    for change in ({"entity_id": "E9"}, {"evidence": []}, {"evidence": ["foreign"]}):
        with pytest.raises(jsonschema.ValidationError):
            jsonschema.validate({**valid, "type_concerns": [{**valid["type_concerns"][0],
                                                            **change}]}, review_schema)
