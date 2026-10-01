"""Integration safety for source proof, current work and bounded model dispatch."""

from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from rdflib import Graph
from test_evidence_gate import document, frozen_rule, state_and_claim

from app.services.document_harness.controller import Engine, Paused
from app.services.document_harness.evidence_gate import assertion_dependency_hash
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.source import reference
from app.services.document_harness.work import make_work
from app.services.document_harness.work_execution import drain_work, gate_state
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def evidence_case(tmp_path):
    doc = Document()
    doc.add_paragraph("Alpha contains Beta and Gamma.")
    doc.add_paragraph("Only under condition C; not outside C.")
    doc.add_paragraph("The two components act in parallel.")
    path = tmp_path / "evidence-work.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    catalog = catalog_from_graph(Graph().parse(data="""
        @prefix : <urn:gate:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class . :Thing a owl:Class .
        :describes a owl:ObjectProperty; rdfs:domain :Report; rdfs:range :Thing .
        :contains a owl:ObjectProperty; rdfs:domain :Thing; rdfs:range :Thing .
    """, format="turtle"), "urn:gate:Report")
    source = ir.evidence_units[0]
    entities = {}
    for key, label in [("a", "Alpha"), ("b", "Beta"), ("c", "Gamma")]:
        start = source.text.index(label)
        ref = reference(ir, source.evidence_id, start, start + len(label))
        entities[key] = {
            "id": key, "label": label, "name": ref, "referent": ref,
            "evidence": [ref], "type_evidence": [ref], "class_iri": "urn:gate:Thing",
            "class_label": "Thing", "role": "object", "state": "accepted",
            "field_ids": [], "window_id": None,
        }
    return ir, catalog, {"entities": entities, "fields": {}, "work": {}}


def create_engine(case, invoke, *, state=None, policy=None, budget=None):
    ir, catalog, original = case
    state = deepcopy(state or original)
    state["entities"].setdefault("document", {
        "id": "document", "class_iri": catalog.root_class_iri, "class_label": "Report",
        "label": "Report", "role": "document_root", "state": "accepted", "reason": "Selected",
        "evidence": [], "field_ids": [], "window_id": None,
    })
    engine = Engine(
        ir=ir, catalog=catalog, state=state, invoke=invoke,
        save=lambda changes: None, should_stop=lambda: False,
        policy={"execution_policy": policy or {}}, max_request_bytes=budget,
    )
    engine.state.setdefault("cursor", {"main": {
        "entity_window_id": None, "active_batches": {}, "phase": "graph", "stage": "planning",
    }})
    return engine


def relation(engine, key="r", *, group=False, **updates):
    source = engine.ir.evidence_units[0]
    row = {
        "id": key, "subject_id": "a", "predicate_iri": "urn:gate:contains",
        "alignment_class_iri": "urn:gate:Thing", "label": "contains",
        "evidence": [reference(engine.ir, source.evidence_id, 0, len(source.text))],
        "polarity": "positive", "conditions": [], "state": "candidate", "reason": "proposal",
        "window_id": None,
    }
    if group:
        row.update(object_ids=["b", "c"], participation="all", selection="unspecified",
                   timing="parallel", timing_state="candidate", timing_reason="pending")
    else:
        row["object_id"] = "b"
    return {**row, **updates}


def add_review(engine, row, *, domain="relations"):
    state = {**engine.state, domain: {**engine.state.get(domain, {}), row["id"]: row}}
    work = make_work("evidence_review", {"domain": domain, "assertion_id": row["id"]},
                     state, engine.catalog, engine.execution_policy)
    engine.commit({domain: {row["id"]: row}, "work": {work["id"]: work}})
    return work["id"]


def review_answer(payload, *, verdict="accepted", timing_verdict=None):
    return {"judgments": {row["id"]: {
        "verdict": timing_verdict if timing_verdict and row["kind"] == "relation_timing"
        else verdict,
        "confidence": 0.99, "reason": "Original source checked",
        "evidence": [source["source_id"] for source in payload["sources"]],
    } for row in payload["candidates"]}, "type_concerns": []}


def seed(engine, *, objects=None, polarity="positive", conditions=None):
    source = engine.ir.evidence_units[0]
    return {
        "subject_id": "a", "object_ids": objects or ["b"],
        "predicate_iri": "urn:gate:contains", "region_id": "source-region",
        "clue_refs": [reference(engine.ir, source.evidence_id, 0, len(source.text))],
        "required_context_refs": [], "origin_ids": ["hint"], "priority": "explicit",
        "polarity_hint": polarity, "condition_hints": conditions or [],
        "endpoint_hypothesis": False,
    }


def add_alignment(engine, data):
    work = make_work("relation_alignment", data, engine.state, engine.catalog,
                     engine.execution_policy)
    engine.commit({"work": {work["id"]: work}})
    return work["id"]


def proposal(payload, *, verdict="proposed", polarity="positive", conditions=None,
             missing_context="none"):
    return {"proposals": {row["candidate_id"]: {
        "verdict": verdict, "polarity": polarity, "conditions": conditions or [],
        "evidence": [source["source_id"] for source in payload["sources"]],
        "participation": None, "selection": None, "timing": None,
        "missing_context": missing_context, "reason": "Source proposal", "confidence": 0.99,
    } for row in payload["items"]}}


def finish(engine):
    for _ in range(20):
        if not drain_work(engine, None):
            return
    pytest.fail("work did not finish within bounded dispatches")


@pytest.mark.parametrize("verdict", ["accepted", "rejected", "unresolved"])
def test_exact_semantic_proof_reuses_method_and_verdict_without_model(evidence_case, verdict):
    calls = []
    engine = create_engine(evidence_case, lambda *args: calls.append(args))
    row = relation(engine)
    row["verification"] = {
        "method": "llm", "semantic_verdict": verdict,
        "dependency_hash": assertion_dependency_hash(
            engine.ir, engine.catalog, gate_state(engine), row,
        ),
    }
    key = add_review(engine, row)
    finish(engine)
    assert not calls
    assert engine.state["relations"]["r"]["state"] == verdict
    assert engine.state["relations"]["r"]["verification"]["method"] == "llm"
    assert engine.state["work"][key]["status"] == "done"


@pytest.mark.parametrize("new_evidence", [False, True])
def test_new_counterevidence_invalidates_current_accepted_proof(evidence_case, new_evidence):
    engine = create_engine(evidence_case, lambda stage, payload, _: review_answer(payload))
    key = add_review(engine, relation(engine))
    finish(engine)
    row = engine.state["relations"]["r"]
    assert row["state"] == "accepted"
    unit = engine.ir.evidence_units[1 if new_evidence else 0]
    engine.commit({"hints": {"counter": {
        "id": "counter", "subject_id": "a", "object_id": "b", "label": "contains",
        "polarity": "negative", "conditions": ["Only under C"],
        "evidence": [reference(engine.ir, unit.evidence_id, 0, len(unit.text))],
    }}})
    assert engine.state["relations"]["r"]["state"] == "unresolved"
    assert engine.state["work"][key]["status"] == "ready"
    def recheck(stage, payload, schema):
        assert stage == "evidence_review"
        assert unit.text in [source["text"] for source in payload["sources"]]
        return review_answer(payload, verdict="rejected")

    engine.invoke = recheck
    finish(engine)
    assert engine.state["relations"]["r"]["state"] == "rejected"
    assert engine.state["relations"]["r"]["verification"]["semantic_verdict"] == "rejected"


@pytest.mark.parametrize("direction,subject,target", [
    ("outgoing", "a", "b"), ("incoming", "b", "a"),
])
def test_remote_reference_counterevidence_invalidates_and_enters_review(
    evidence_case, direction, subject, target,
):
    engine = create_engine(evidence_case, lambda stage, payload, _: review_answer(payload))
    key = add_review(engine, relation(engine))
    finish(engine)
    assert engine.state["relations"]["r"]["state"] == "accepted"
    unit = engine.ir.evidence_units[1]
    ref = reference(engine.ir, unit.evidence_id, 0, len(unit.text))
    engine.commit({"reference_cues": {"counter": {
        "id": "counter", "subject_id": subject, "target_ids": [target],
        "direction": direction, "target_expression": ref, "relation_label": "contains",
        "polarity": "negative", "conditions": ["Only under C"], "evidence": [ref],
    }}})
    assert engine.state["relations"]["r"]["state"] == "unresolved"
    assert engine.state["work"][key]["status"] == "ready"

    def recheck(stage, payload, schema):
        assert stage == "evidence_review"
        assert unit.text in [source["text"] for source in payload["sources"]]
        return review_answer(payload, verdict="rejected")

    engine.invoke = recheck
    finish(engine)
    assert engine.state["relations"]["r"]["state"] == "rejected"
    assert engine.state["relations"]["r"]["verification"]["semantic_verdict"] == "rejected"


def test_endpoint_gate_changes_reuse_proof_but_type_evidence_changes_recheck(evidence_case):
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        return review_answer(payload)

    engine = create_engine(evidence_case, invoke)
    key = add_review(engine, relation(engine))
    finish(engine)
    original_proof = deepcopy(engine.state["relations"]["r"]["verification"])
    for state, expected in [("candidate", "unresolved"), ("accepted", "accepted"),
                            ("unresolved", "unresolved")]:
        engine.commit({"entities": {"b": {**engine.state["entities"]["b"], "state": state}}})
        assert engine.state["relations"]["r"]["state"] == expected
        assert engine.state["relations"]["r"]["verification"] == original_proof
        finish(engine)
        assert calls == ["evidence_review"]
    unit = engine.ir.evidence_units[1]
    engine.commit({"entities": {"b": {
        **engine.state["entities"]["b"],
        "type_evidence": [reference(engine.ir, unit.evidence_id, 0, len(unit.text))],
    }}})
    assert engine.state["relations"]["r"]["state"] == "unresolved"
    assert engine.state["work"][key]["status"] == "ready"


@pytest.mark.parametrize("polarity,conditions", [("negative", []), ("positive", ["Only under C"])])
def test_relation_polarity_and_conditions_survive_alignment_and_review(
    evidence_case, polarity, conditions,
):
    def invoke(stage, payload, schema):
        if stage == "relation_alignment":
            return proposal(payload, polarity=polarity, conditions=conditions)
        assert stage == "evidence_review"
        assert all(row["polarity"] == polarity and row["conditions"] == conditions
                   for row in payload["candidates"])
        return review_answer(payload)

    engine = create_engine(evidence_case, invoke)
    add_alignment(engine, seed(engine, polarity=polarity, conditions=conditions))
    finish(engine)
    row = next(iter(engine.state["relations"].values()))
    assert row["state"] == "accepted"
    assert row["polarity"] == polarity and row["conditions"] == conditions
    assert not any(item["polarity"] == "positive" and not item["conditions"]
                   for item in engine.state["relations"].values())


def test_no_relation_is_not_a_negative_fact(evidence_case):
    engine = create_engine(evidence_case, lambda stage, payload, _: proposal(
        payload, verdict="no_relation", polarity="negative",
    ))
    key = add_alignment(engine, seed(engine))
    finish(engine)
    assert not engine.state.get("relations")
    assert engine.state["work"][key]["reason_code"] == "no_relation"


def test_mixed_alignment_keeps_unresolved_member_and_complete_group_separate(evidence_case):
    batches = []

    def invoke(stage, payload, schema):
        if stage == "relation_alignment":
            batches.append([len(row["object_ids"]) for row in payload["items"]])
            answer = proposal(payload, verdict="unresolved", missing_context="participation")
            for row in payload["items"]:
                if len(row["object_ids"]) > 1:
                    answer["proposals"][row["candidate_id"]].update(
                        verdict="proposed", participation="all", selection="unspecified",
                        timing="parallel", missing_context="none",
                    )
            jsonschema.validate(answer, schema)
            invalid = deepcopy(answer)
            for row in payload["items"]:
                if len(row["object_ids"]) == 1:
                    invalid["proposals"][row["candidate_id"]].update(
                        participation="options", selection="exactly_one",
                    )
                    with pytest.raises(jsonschema.ValidationError):
                        jsonschema.validate(invalid, schema)
                    break
            return answer
        assert stage == "evidence_review"
        return review_answer(payload)

    engine = create_engine(evidence_case, invoke)
    single_key = add_alignment(engine, seed(engine))
    group_key = add_alignment(engine, seed(engine, objects=["b", "c"]))
    finish(engine)
    assert sorted(batches[0]) == [1, 2]
    assert not engine.state.get("relations")
    assert engine.state["work"][single_key]["status"] == "waiting"
    assert engine.state["work"][group_key]["status"] == "done"
    groups = list(engine.state["relation_groups"].values())
    assert len(groups) == 1
    assert groups[0]["object_ids"] == ["b", "c"]
    assert groups[0]["participation"] == "all"
    assert groups[0]["state"] == "accepted"


@pytest.mark.parametrize("omit", ["relation_groups", "relation_timing"])
def test_group_missing_one_judgment_never_partially_commits(evidence_case, omit):
    def invoke(stage, payload, schema):
        answer = review_answer(payload)
        omitted = next(row["id"] for row in payload["candidates"] if row["kind"] == omit)
        answer["judgments"].pop(omitted)
        return answer

    engine = create_engine(evidence_case, invoke)
    key = add_review(engine, relation(engine, group=True), domain="relation_groups")
    with pytest.raises(ValueError, match="evidence_review_answer_set_mismatch"):
        drain_work(engine, None)
    row = engine.state["relation_groups"]["r"]
    assert row["state"] != "accepted" and row["timing_state"] != "accepted"
    assert engine.state["work"][key]["status"] == "failed"


def test_group_r_and_t_are_not_split_to_fit_byte_budget(evidence_case, monkeypatch):
    calls = []
    monkeypatch.setattr("app.services.document_harness.model.request_size",
                        lambda stage, payload, schema: 100 * len(payload["candidates"]))
    engine = create_engine(evidence_case, lambda *args: calls.append(args), budget=150)
    add_review(engine, relation(engine, group=True), domain="relation_groups")
    with pytest.raises(ValueError, match="HARNESS_EVIDENCE_CONTEXT_TOO_LARGE"):
        drain_work(engine, None)
    assert not calls


def test_group_participation_and_timing_keep_independent_verdicts(evidence_case):
    calls = []

    def invoke(stage, payload, schema):
        calls.append(deepcopy(payload))
        return review_answer(payload, timing_verdict="unresolved")

    engine = create_engine(evidence_case, invoke)
    add_review(engine, relation(engine, group=True), domain="relation_groups")
    finish(engine)
    row = engine.state["relation_groups"]["r"]
    assert row["state"] == "accepted" and row["timing_state"] == "unresolved"
    assert len(calls) == 1
    assert {item["kind"] for item in calls[0]["candidates"]} == {
        "relation_groups", "relation_timing",
    }


@pytest.mark.parametrize("separate_batches", [False, True])
def test_competing_group_participation_cannot_be_accepted_together(
    evidence_case, separate_batches,
):
    engine = create_engine(evidence_case, lambda stage, payload, _: review_answer(payload))
    add_review(engine, relation(engine, "all", group=True), domain="relation_groups")
    if separate_batches:
        finish(engine)
    add_review(engine, relation(engine, "options", group=True, participation="options",
                               selection="exactly_one", timing="unspecified"),
               domain="relation_groups")
    finish(engine)
    assert all(row["state"] == "unresolved" and row["timing_state"] == "unresolved"
               for row in engine.state["relation_groups"].values())
    assert all(row["verification"]["semantic_verdict"] == "accepted"
               for row in engine.state["relation_groups"].values())
    for status in ("unresolved", "accepted"):
        engine.commit({"entities": {"b": {**engine.state["entities"]["b"], "state": status}}})
        assert all(row["state"] == "unresolved"
                   for row in engine.state["relation_groups"].values())


def add_interpretation(engine, row):
    state = {**engine.state, "relation_groups": {row["id"]: row}}
    work = make_work("group_interpretation", {"group_id": row["id"]}, state,
                     engine.catalog, engine.execution_policy)
    engine.commit({"relation_groups": {row["id"]: row}, "work": {work["id"]: work}})
    return work["id"]


def test_unknown_participation_without_new_context_stays_waiting_without_llm(evidence_case):
    calls = []
    engine = create_engine(evidence_case, lambda *args: calls.append(args))
    group = relation(engine, group=True, participation="unknown", timing="unspecified")
    group["evidence"] = [reference(engine.ir, unit.evidence_id, 0, len(unit.text))
                         for unit in engine.ir.evidence_units]
    key = add_interpretation(engine, group)
    finish(engine)
    assert not calls
    assert engine.state["work"][key]["status"] == "waiting"
    assert engine.state["relation_groups"]["r"]["participation"] == "unknown"
    assert engine.state["relation_groups"]["r"]["state"] != "accepted"


def test_group_interpretation_once_then_independent_review_preserves_all_members(evidence_case):
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        if stage == "group_interpretation":
            assert len(payload["group"]["objects"]) == 2
            return {
                "verdict": "supported", "participation": "options", "selection": "exactly_one",
                "timing": "unspecified", "reason": "Explicit alternate participation",
                "evidence": [row["source_id"] for row in payload["sources"]],
            }
        assert stage == "evidence_review"
        assert payload["candidates"][0]["participation"] == "options"
        assert len(payload["candidates"][0]["objects"]) == 2
        return review_answer(payload)

    engine = create_engine(evidence_case, invoke)
    group = relation(engine, group=True, participation="unknown", timing="unspecified",
                     polarity="negative", conditions=["Only under C"])
    add_interpretation(engine, group)
    finish(engine)
    assert calls == ["group_interpretation", "evidence_review"]
    assert len(engine.state["relation_groups"]) == 1
    row = next(iter(engine.state["relation_groups"].values()))
    work = next(work for work in engine.state["work"].values()
                if work["kind"] == "group_interpretation")
    assert row["id"] == work["input"]["group_id"]
    assert work["output_ids"] == [row["id"]]
    assert row["object_ids"] == ["b", "c"]
    assert row["polarity"] == "negative" and row["conditions"] == ["Only under C"]
    assert row["state"] == "accepted" and row["timing_state"] == "unresolved"
    assert work["expansions_used"] == 1 and work["status"] == "done"
    finish(engine)
    assert len(calls) == 2


def test_group_timing_gate_revokes_and_restores_without_paid_review(evidence_case):
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        return review_answer(payload)

    engine = create_engine(evidence_case, invoke)
    add_review(engine, relation(engine, group=True), domain="relation_groups")
    finish(engine)
    for status, expected in [("unresolved", "unresolved"), ("accepted", "accepted")]:
        engine.commit({"entities": {"b": {**engine.state["entities"]["b"], "state": status}}})
        row = engine.state["relation_groups"]["r"]
        assert row["state"] == expected and row["timing_state"] == expected
        finish(engine)
        assert calls == ["evidence_review"]


def test_frozen_table_rule_finishes_alignment_and_proof_without_llm(tmp_path, evidence_case):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    calls = []
    rule = frozen_rule(ir)
    engine = create_engine((ir, evidence_case[1], state), lambda *args: calls.append(args),
                           policy={"table_relation_rules": [rule]})
    data = {**seed(engine), "clue_refs": claim["evidence"]}
    key = add_alignment(engine, data)
    finish(engine)
    assert not calls
    assert engine.state["work"][key]["status"] == "done"
    assert len(engine.state["work"]) == 1
    row = next(iter(engine.state["relations"].values()))
    assert row["state"] == "accepted" and row["verification"]["method"] == "rule"
    for ref in row["evidence"]:
        assert engine.ir.resolve(engine.ir.anchor(ref["source_id"], ref["start"], ref["end"])) == (
            ref["text"]
        )


def test_expansion_runs_once_and_waiting_does_not_repeat_after_resume(evidence_case):
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        return proposal(payload, verdict="unresolved", missing_context="condition")

    engine = create_engine(evidence_case, invoke)
    key = add_alignment(engine, seed(engine))
    initial_basis = engine.state["work"][key]["expansion_basis_hash"]
    finish(engine)
    work = engine.state["work"][key]
    assert work["status"] == "waiting" and work["expansions_used"] == 1
    assert work["expansion_basis_hash"] == initial_basis
    assert calls == ["relation_alignment", "relation_alignment"]
    resumed = create_engine(evidence_case, invoke, state=engine.state)
    finish(resumed)
    assert len(calls) == 2
    assert not resumed.state.get("relations")


@pytest.mark.parametrize("failure", ["technical", "pause"])
def test_failed_or_paused_review_never_becomes_semantic_rejection(evidence_case, failure):
    def invoke(*args):
        if failure == "pause":
            raise Paused()
        raise RuntimeError("model unavailable")

    engine = create_engine(evidence_case, invoke)
    key = add_review(engine, relation(engine))
    with pytest.raises(Paused if failure == "pause" else RuntimeError):
        drain_work(engine, None)
    row = engine.state["relations"]["r"]
    assert row["state"] in {"candidate", "unresolved"}
    assert not row.get("verification")
    assert engine.state["work"][key]["status"] == ("ready" if failure == "pause" else "failed")
