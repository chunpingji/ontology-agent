"""Relation context retains original heading/row order and condition labels."""

from types import SimpleNamespace

import pytest
from docx import Document

from app.services.document_harness.planning import context_window
from app.services.document_harness.source import Window, reference
from app.services.document_harness.work_execution import work_context
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def source(tmp_path):
    doc = Document()
    doc.add_heading("步骤3：目标产品的合成", 1)
    table = doc.add_table(rows=4, cols=3)
    data = [
        ("投料说明", "本步骤分两个批次投料", "条件"),
        ("第一批次", "于R1反应釜中加入物料A", "加热"),
        ("第二批次", "于R2反应釜中加入物料B", "仅在氮气保护下"),
        ("后续操作", "使用R3设备", "独立条件"),
    ]
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / "relation-context.docx"
    doc.save(path)
    return build_document_ir(path, parse_docx_structure(path))


def ref(ir, text):
    unit = next(unit for unit in ir.evidence_units if unit.text == text)
    return reference(ir, unit.evidence_id, 0, len(unit.text))


def test_context_reorders_dependency_sources_without_changing_quotes_or_coverage(source):
    heading = ref(source, "步骤3：目标产品的合成")
    operation = ref(source, "于R2反应釜中加入物料B")
    window = context_window(source, Window("w", [], [], []), [], {}, [operation, heading])
    assert [item["text"] for item in window.sources] == [heading["text"], operation["text"]]
    assert window.primary_ids == []
    assert window.resolve(source, {"source_id": "S1", "text": heading["text"]}) == heading
    assert window.resolve(source, {"source_id": "S2", "text": operation["text"]}) == operation
    # Different dependency traversal must produce identical aliases and payload.
    ordered = context_window(source, Window("w", [], [], []), [], {}, [heading, operation])
    assert ordered.payload() == window.payload()


@pytest.mark.parametrize("kind", ["relation_alignment", "evidence_review", "group_interpretation"])
def test_relation_context_includes_its_row_label_and_condition_without_other_batches(source, kind):
    heading = ref(source, "步骤3：目标产品的合成")
    operation = ref(source, "于R2反应釜中加入物料B")
    entities = {
        key: {"id": key, "referent": item, "evidence": [item], "field_ids": []}
        for key, item in (("step", heading), ("equipment", operation))
    }
    assertion = {"id": "r", "subject_id": "step", "object_id": "equipment",
                 "object_ids": ["equipment"]}
    engine = SimpleNamespace(ir=source, state={"entities": entities, "fields": {},
                                              "relations": {"r": assertion},
                                              "relation_groups": {"r": assertion}})
    data = ({"subject_id": "step", "object_ids": ["equipment"]}
            if kind == "relation_alignment" else
            {"domain": "relations", "assertion_id": "r"} if kind == "evidence_review" else
            {"group_id": "r"})
    row = {"kind": kind, "input": data,
           "dependencies": {"source_refs": [operation], "field_ids": []}}
    window = work_context(engine, [row], None)
    assert [item["text"] for item in window.sources] == [
        "步骤3：目标产品的合成", "投料说明", "本步骤分两个批次投料", "条件",
        "第二批次", "于R2反应釜中加入物料B", "仅在氮气保护下",
    ]
    assert window.primary_ids == [] and window.fields == []
    assert window.entity_ids == ["equipment", "step"]
