"""Physical fields and quote rights without an ontology or identity inference."""

import pytest
from docx import Document

from app.services.document_harness.source import (
    build_windows,
    make_window,
    quote_for_reference,
    reference,
)
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def document_ir(tmp_path, doc):
    path = tmp_path / "fields.docx"
    doc.save(path)
    return build_document_ir(path, parse_docx_structure(path))


def table_document(rows):
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for index, row in enumerate(rows):
        for column, text in enumerate(row):
            table.cell(index, column).text = text
    return doc, table


def fields(ir):
    return [field for window in build_windows(ir) for field in window.fields]


def assert_references(ir, original_fields):
    for field in original_fields:
        assert "subject_id" not in field
        assert "predicate_iri" not in field
        for ref in field["evidence"]:
            assert ir.unit(ref["source_id"]).text[ref["start"]:ref["end"]] == ref["text"]


@pytest.mark.parametrize("headings", [("字段", "值"), ("Attribute", "Value")])
def test_explicit_field_value_table_keeps_row_labels_and_missing_value(tmp_path, headings):
    doc, _ = table_document([headings, ("名称", "Alpha😀"), ("编号", "N/A")])
    ir = document_ir(tmp_path, doc)
    extracted = fields(ir)
    assert [(item["label"], item["value"]) for item in extracted] == [
        ("名称", "Alpha😀"), ("编号", "N/A"),
    ]
    assert extracted[1]["missing"]
    assert [ref["text"] for ref in extracted[0]["evidence"]] == ["名称", "Alpha😀"]
    assert_references(ir, extracted)


def test_headerless_prose_fields_include_first_row_without_inferring_an_owner(tmp_path):
    descriptions = [
        f"这是第{number}项的原始描述。" + "详细说明该字段的内容；" * 7 for number in (1, 2)
    ]
    doc, _ = table_document([("说明甲", descriptions[0]), ("说明乙", descriptions[1])])
    ir = document_ir(tmp_path, doc)
    assert ir.tables[0]["header_row_count"] == 0
    extracted = fields(ir)
    assert [(item["label"], item["value"]) for item in extracted] == [
        ("说明甲", descriptions[0]), ("说明乙", descriptions[1]),
    ]
    assert_references(ir, extracted)


def test_explicit_colon_reclaims_first_row_from_parser_header_fallback(tmp_path):
    doc, _ = table_document([("  名称：", "Alpha"), ("颜色", "blue")])
    ir = document_ir(tmp_path, doc)
    assert ir.tables[0]["header_row_count"] == 1
    extracted = fields(ir)
    assert [(item["label"], item["value"]) for item in extracted] == [
        ("名称", "Alpha"), ("颜色", "blue"),
    ]
    assert extracted[0]["evidence"][0]["text"] == "名称"
    assert_references(ir, extracted)


def test_two_column_data_table_uses_column_headers_not_first_cell_as_property(tmp_path):
    doc, _ = table_document([("记录号", "测量值"), ("A01", "17"), ("A02", "18")])
    ir = document_ir(tmp_path, doc)
    extracted = fields(ir)
    assert [(item["label"], item["value"]) for item in extracted] == [
        ("记录号", "A01"), ("测量值", "17"), ("记录号", "A02"), ("测量值", "18"),
    ]
    assert_references(ir, extracted)


def test_numbered_headerless_records_remain_unbound_sources(tmp_path):
    doc, _ = table_document([("1", "对象甲"), ("2", "对象乙"), ("3", "对象丙")])
    ir = document_ir(tmp_path, doc)
    assert ir.tables[0]["header_row_count"] == 0
    windows = build_windows(ir)
    assert not fields(ir)
    assert {source["text"] for window in windows for source in window.sources} == {
        "1", "2", "3", "对象甲", "对象乙", "对象丙",
    }


def test_ambiguous_fallback_does_not_invent_first_row_field_binding(tmp_path):
    doc, _ = table_document([("名称", "Alpha"), ("颜色", "blue")])
    ir = document_ir(tmp_path, doc)
    assert ir.tables[0]["header_row_count"] == 1
    assert not any(item["label"] == "名称" and item["value"] == "Alpha" for item in fields(ir))
    # The model still sees both original cells and can propose their field role.
    assert "Alpha" in [source["text"] for source in build_windows(ir)[0].sources]


def test_merged_column_headers_follow_grid_coverage_and_exact_origin(tmp_path):
    doc, table = table_document([
        ("对象", "测量", ""), ("名称", "温度", "长度"), ("Alpha", "20", "3"),
    ])
    table.cell(0, 1).merge(table.cell(0, 2))
    ir = document_ir(tmp_path, doc)
    assert ir.tables[0]["header_row_count"] == 2
    extracted = fields(ir)
    assert [(item["label"], item["value"]) for item in extracted] == [
        ("对象 / 名称", "Alpha"), ("测量 / 温度", "20"), ("测量 / 长度", "3"),
    ]
    assert [ref["text"] for ref in extracted[2]["evidence"]] == ["测量", "长度", "3"]
    assert_references(ir, extracted)


def test_vertical_field_label_merge_reuses_origin_without_copying_source_text(tmp_path):
    doc, table = table_document([("字段", "值"), ("备注", "说明甲"), ("", "说明乙")])
    table.cell(1, 0).merge(table.cell(2, 0))
    ir = document_ir(tmp_path, doc)
    extracted = fields(ir)
    assert [(item["label"], item["value"]) for item in extracted] == [
        ("备注", "说明甲"), ("备注", "说明乙"),
    ]
    assert extracted[0]["evidence"][0] == extracted[1]["evidence"][0]
    assert_references(ir, extracted)


def test_long_section_keeps_naming_context_citable_in_every_window(tmp_path):
    doc = Document()
    doc.add_heading("对象资料", 1)
    doc.add_paragraph("名称：Alpha😀")
    doc.add_paragraph("角色：容器")
    for index in range(9):
        doc.add_paragraph(f"观察{index}：" + "这一段原文描述对象的某项表现。" * 12)
    ir = document_ir(tmp_path, doc)
    windows = build_windows(ir, max_chars=450, max_sources=5)
    assert len(windows) > 1
    name = next(unit for unit in ir.evidence_units if unit.text.startswith("名称"))
    for window in windows:
        source = next(
            source for source in window.sources if source["evidence_id"] == name.evidence_id
        )
        ref = window.resolve(ir, {"source_id": source["source_id"], "text": "Alpha😀"})
        assert ref["text"] == "Alpha😀"
        assert ref["start"] == 3
        assert "fact_eligible" not in source
        assert source["source_id"] in {item["source_id"] for item in window.payload()["sources"]}
    assert {key for window in windows for key in window.primary_ids} == {
        unit.evidence_id for unit in ir.evidence_units if unit.text.strip()
    }


def test_quotes_preserve_unicode_offsets_and_require_occurrence_for_repeated_text(tmp_path):
    doc = Document()
    doc.add_paragraph("前😀甲，再😀甲。")
    ir = document_ir(tmp_path, doc)
    window = build_windows(ir)[0]
    quote = {"source_id": "S1", "text": "😀甲"}
    with pytest.raises(ValueError, match="source_quote_ambiguous"):
        window.resolve(ir, quote)
    second = window.resolve(ir, {**quote, "occurrence": 1})
    assert (second["start"], second["end"], second["text"]) == (5, 7, "😀甲")
    assert quote_for_reference(window, second) == {**quote, "occurrence": 1}
    with pytest.raises(ValueError, match="source_quote_mismatch"):
        window.resolve(ir, {"source_id": "S1", "text": "不存在的引文"})
    with pytest.raises(ValueError, match="empty_source_quote"):
        window.resolve(ir, {"source_id": "S1", "text": " "})
    with pytest.raises(ValueError, match="source_outside_reading_window"):
        window.resolve(ir, {"source_id": "S99", "text": "😀甲"})


def test_fragment_quotes_translate_to_original_source_code_point_offsets(tmp_path):
    doc = Document()
    doc.add_paragraph("前😀甲，再😀甲。")
    ir = document_ir(tmp_path, doc)
    unit = ir.evidence_units[0]
    window = make_window(ir, [(unit.evidence_id, 4, 8)], [unit.evidence_id])
    actual = window.resolve(ir, {"source_id": "S1", "text": "😀甲"})
    assert actual == reference(ir, unit.evidence_id, 5, 7)


@pytest.mark.parametrize("text, excerpt", [("哈哈哈", "哈哈"), ("😀😀😀", "😀😀")])
def test_overlapping_substrings_require_exact_occurrence_and_round_trip(tmp_path, text, excerpt):
    doc = Document()
    doc.add_paragraph(text)
    ir = document_ir(tmp_path, doc)
    unit = ir.evidence_units[0]
    window = build_windows(ir)[0]
    with pytest.raises(ValueError, match="source_quote_ambiguous"):
        window.resolve(ir, {"source_id": "S1", "text": excerpt})
    for occurrence in (0, 1):
        expected = reference(ir, unit.evidence_id, occurrence, occurrence + len(excerpt))
        quote = {"source_id": "S1", "text": excerpt, "occurrence": occurrence}
        assert window.resolve(ir, quote) == expected
        assert quote_for_reference(window, expected) == quote


def test_invalid_occurrence_never_borrows_matching_text_from_another_visible_source(tmp_path):
    doc = Document()
    doc.add_paragraph("同名")
    doc.add_paragraph("同名同名")
    ir = document_ir(tmp_path, doc)
    window = build_windows(ir)[0]
    with pytest.raises(ValueError, match="source_occurrence_invalid"):
        window.resolve(ir, {"source_id": "S1", "text": "同名", "occurrence": 1})
    actual = window.resolve(ir, {"source_id": "S2", "text": "同名", "occurrence": 1})
    assert actual["source_id"] == ir.evidence_units[1].evidence_id
    assert actual["start"] == 2


def test_adjacent_small_sections_share_reading_window_without_assigning_ownership(tmp_path):
    doc = Document()
    doc.add_heading("记录概况", 1)
    doc.add_heading("对象介绍", 2)
    doc.add_paragraph("名称：Alpha")
    doc.add_heading("对象特征", 1)
    doc.add_heading("尺寸", 2)
    doc.add_paragraph("长度：3 cm")
    doc.add_heading("其他对象", 1)
    doc.add_paragraph("名称：Alpha")
    doc.add_paragraph("角色：部件")
    ir = document_ir(tmp_path, doc)
    windows = build_windows(ir)
    assert len(windows) == 1
    window = windows[0]
    assert len({source["section"] for source in window.sources}) == 5
    assert [source["evidence_id"] for source in window.sources] == [
        unit.evidence_id for unit in ir.evidence_units if unit.text.strip()
    ]
    names = [field for field in window.fields if field["label"] == "名称"]
    assert len(names) == 2
    assert names[0]["id"] != names[1]["id"]
    assert not window.entity_ids
    assert_references(ir, window.fields)


@pytest.mark.parametrize("max_chars,max_sources", [(1000, 4), (35, 40)])
def test_adjacent_batches_respect_each_budget_without_dropping_primary_sources(
    tmp_path, max_chars, max_sources,
):
    doc = Document()
    for index in range(4):
        doc.add_heading(f"章节{index}", 1)
        doc.add_paragraph(f"名称：Object{index}")
    ir = document_ir(tmp_path, doc)
    windows = build_windows(ir, max_chars=max_chars, max_sources=max_sources)
    assert len(windows) > 1
    assert all(len(window.sources) <= max_sources for window in windows)
    assert all(sum(len(source["text"]) for source in window.sources) <= max_chars
               for window in windows)
    assert [key for window in windows for key in window.primary_ids] == [
        unit.evidence_id for unit in ir.evidence_units if unit.text.strip()
    ]


def test_large_table_keeps_its_header_chunks_between_small_section_batches(tmp_path):
    doc = Document()
    doc.add_heading("前节", 1)
    doc.add_paragraph("名称：Before")
    doc.add_heading("测量记录", 1)
    table = doc.add_table(rows=13, cols=2)
    table.cell(0, 0).text = "记录号"
    table.cell(0, 1).text = "测量值"
    for index in range(1, 13):
        table.cell(index, 0).text = f"R{index}"
        table.cell(index, 1).text = str(index * 2)
    doc.add_heading("后节", 1)
    doc.add_paragraph("名称：After")
    ir = document_ir(tmp_path, doc)
    windows = build_windows(ir, max_sources=8)
    assert [source["text"] for source in windows[0].sources] == ["前节", "名称：Before"]
    assert [source["text"] for source in windows[-1].sources] == ["后节", "名称：After"]
    for window in windows[1:-1]:
        assert {"记录号", "测量值"} <= {source["text"] for source in window.sources}
        assert all("Before" not in source["text"] and "After" not in source["text"]
                   for source in window.sources)
        assert len(window.sources) <= 8
        for source in window.sources:
            unit = ir.unit(source["evidence_id"])
            assert source["table"] == unit.table_path
            assert source["row"] == unit.row_index
    assert {key for window in windows for key in window.primary_ids} == {
        unit.evidence_id for unit in ir.evidence_units if unit.text.strip()
    }


def test_initial_budget_fragments_never_become_complete_original_fields(tmp_path):
    doc = Document()
    doc.add_paragraph("名称：Alpha")
    original = "详细说明：" + "这是一个还没有结束的原文属性描述。" * 250
    doc.add_paragraph(original)
    ir = document_ir(tmp_path, doc)
    windows = build_windows(ir, max_chars=1900)
    assert len(windows) > 1
    assert not any(field["label"] == "详细说明"
                   for window in windows for field in window.fields)
    fragments = [source for window in windows for source in window.sources
                 if source["evidence_id"] == ir.evidence_units[1].evidence_id]
    assert len(fragments) > 1
    assert fragments[0]["text"].startswith("详细说明：")
    assert all(len(fragment["text"]) < len(original) for fragment in fragments)
    for window in windows:
        for source in window.sources:
            actual = window.resolve(ir, {"source_id": source["source_id"],
                                         "text": source["text"], "occurrence": 0})
            assert actual["text"] == source["text"]
