from copy import deepcopy

import pytest
from docx import Document
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.protocols import stage_schema
from app.services.document_harness.source import build_windows
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def inputs(tmp_path):
    doc = Document()
    doc.add_heading("对象记录", 1)
    for text in ("名称：Alpha", "颜色：red", "备注：N/A", "对象：Beta", "Alpha contains Beta"):
        doc.add_paragraph(text)
    path = tmp_path / "source.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    graph = Graph().parse(
        data="""
        @prefix : <urn:test:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class ; rdfs:label "报告" .
        :Thing a owl:Class ; rdfs:label "物体" ; rdfs:comment "A physical object" .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :Thing .
        :contains a owl:ObjectProperty ; rdfs:domain :Thing ; rdfs:range :Thing .
        :color a owl:DatatypeProperty ; rdfs:label "颜色" ;
          rdfs:domain :Thing ; rdfs:range xsd:string .
    """,
        format="turtle",
    )
    return ir, catalog_from_graph(graph, "urn:test:Report")


def quote(sources, text):
    source = next(s for s in sources if text in s["text"])
    return {"source_id": source["source_id"], "text": text, "occurrence": None}


class Model:
    def __init__(self, *, bad_review=False, foreign_property=False, incomplete=False):
        self.calls = []
        self.bad_review = bad_review
        self.foreign_property = foreign_property
        self.incomplete = incomplete

    def __call__(self, stage, payload, schema):
        self.calls.append((stage, deepcopy(payload), deepcopy(schema)))
        sources = payload["sources"]
        if stage == "referent_candidates":
            return {"new_spans": [], "expressions": [], "partitions": []}
        if stage == "coreference_review":
            return {
                "judgments": {
                    pair["pair_id"]: {
                        "verdict": "unresolved",
                        "basis": "insufficient",
                        "confidence": 0.95,
                        "evidence": [s["source_id"] for s in sources],
                        "proof": [],
                        "reason": "No direct co-reference proof in this context",
                    }
                    for pair in payload["pairs"]
                }
            }
        if stage == "discover":
            fields = {f["label"]: f["field_id"] for f in payload["fields"]}
            return {
                "document_field_ids": [],
                "document_source_fields": [],
                "unowned_fields": [],
                "entities": [
                    {
                        "local_id": "a",
                        "name": quote(sources, "Alpha"),
                        "anchor": quote(sources, "Alpha"),
                        "role": "container",
                        "evidence": [quote(sources, "Alpha")["source_id"]],
                        "field_ids": [fields["名称"], fields["颜色"], fields["备注"]],
                        "source_fields": [],
                    },
                    {
                        "local_id": "b",
                        "name": quote(sources, "Beta"),
                        "anchor": quote(sources, "Beta"),
                        "role": "component",
                        "evidence": [quote(sources, "Beta")["source_id"]],
                        "field_ids": [fields["对象"]],
                        "source_fields": [],
                    },
                ],
                "relation_hints": [
                    {
                        "subject_id": "a",
                        "object_id": "b",
                        "label": "contains",
                        "evidence": [quote(sources, "Alpha contains Beta")["source_id"]],
                        "polarity": "positive",
                        "conditions": [],
                    }
                ],
                "complete": not self.incomplete,
            }
        if stage == "type_alignment":
            return {
                "entities": {
                    e["entity_id"]: {
                        "class_iri": "urn:test:Thing",
                        "confidence": 0.95,
                        "evidence": e["evidence"],
                        "reason": "Physical object in source",
                    }
                    for e in payload["entities"]
                }
            }
        if stage == "property_alignment":
            subject = payload["subject"]
            properties = {
                key: {"mappings": [], "reason": "Current card has no matching property"}
                for key in payload["property_field_ids"]
            }
            for f in subject["fields"]:
                if f["label"] == "颜色" and f["field_id"] in payload["property_field_ids"]:
                    properties[f["field_id"]] = {
                        "mappings": [
                            {
                                "confidence": 0.95,
                                "value_component": "whole",
                                "value_quote": None,
                                "predicate_iri": "urn:test:foreign"
                                if self.foreign_property
                                else "urn:test:color",
                            }
                        ],
                        "reason": "Original field role",
                    }
            return {"properties": properties}
        if stage == "relation_alignment":
            return {
                "proposals": {
                    item["candidate_id"]: {
                        "verdict": "proposed",
                        "polarity": item["polarity_hint"],
                        "conditions": item["condition_hints"],
                        "participation": None,
                        "selection": None,
                        "timing": None,
                        "missing_context": "none",
                        "evidence": item["clue_sources"],
                        "reason": "Source relation",
                        "confidence": 0.95,
                    }
                    for item in payload["items"]
                }
            }
        assert stage in {"entity_review", "evidence_review"}
        judgments = {}
        for c in payload["candidates"]:
            evidence = [
                *c["evidence"],
                *c["subject"]["evidence"],
                *c.get("object", {}).get("evidence", []),
            ]
            if self.bad_review:
                evidence = [quote(sources, "对象记录")["source_id"]]
            judgments[c["id"]] = {
                "verdict": "accepted",
                "confidence": 0.99,
                "evidence": evidence,
                "reason": "Literal source and role checked",
            }
        return {
            "judgments": judgments,
            **({"type_concerns": []} if stage == "evidence_review" else {}),
        }


def execute(inputs, model, *, stop=lambda: False, state=None, windows=None, max_request_bytes=None):
    ir, catalog = inputs
    snapshots = []
    from app.services.document_harness.calls import MemoryCalls

    if state is None or not hasattr(model, "_call_ports"):
        model._call_ports = MemoryCalls(model)
    engine = Engine(
        ir=ir,
        catalog=catalog,
        state=state or {},
        invoke=model, calls=model._call_ports,
        save=lambda changes: snapshots.append(deepcopy(changes)),
        should_stop=stop,
        max_request_bytes=max_request_bytes,
    )
    if windows is not None:
        engine.windows = windows
    engine.run()
    return engine.state, snapshots


def test_full_new_pipeline_preserves_sources_and_exposes_dashed_before_review(inputs):
    model = Model()
    state, snapshots = execute(inputs, model)
    assert set(stage for stage, _, _ in model.calls) == {
        "discover",
        "type_alignment",
        "entity_review",
        "property_alignment",
        "relation_alignment",
        "evidence_review",
    }
    mentions = [e for e in state["entities"].values() if e["role"] != "document_root"]
    assert len(mentions) == 2
    assert all(e["state"] == "accepted" for e in mentions)
    assert any(
        r["predicate_iri"] == "urn:test:contains" and r["state"] == "accepted"
        for r in state["relations"].values()
    )
    assert any(
        r["state"] == "candidate" for s in snapshots for r in s.get("relations", {}).values()
    )
    assert [p["value"] for p in state["properties"].values()] == ["red"]
    assert any(o["reason"] == "missing_source_value" for o in state["observations"].values())
    assert state["cursor"]["main"]["scope_complete"]
    ir, _ = inputs
    for domain in ("entities", "properties", "relations", "observations"):
        for item in state.get(domain, {}).values():
            for ref in item["evidence"]:
                assert ir.unit(ref["source_id"]).text[ref["start"] : ref["end"]] == ref["text"]
    discovery = model.calls[0][1]
    assert any("名称：Alpha" in s["text"] for s in discovery["sources"])
    assert all("fact_eligible" not in s for s in discovery["sources"])
    assert any(f["label"] == "名称" for f in discovery["fields"])
    assert [card["iri"] for card in discovery["schema_guidance"]["classes"]] == ["urn:test:Thing"]
    assert all("evidence" not in card for card in discovery["schema_guidance"]["classes"])
    assert state["windows"][next(iter(state["windows"]))]["guidance_class_iris"] == [
        "urn:test:Thing",
    ]
    for stage, payload, _ in model.calls:
        if stage == "type_alignment":
            assert all(e["entity_id"].startswith("E") for e in payload["entities"])


def test_real_but_unrelated_review_quote_cannot_accept_claim(inputs):
    state, _ = execute(inputs, Model(bad_review=True))
    assert all(
        e["state"] == "unresolved"
        for e in state["entities"].values()
        if e["role"] != "document_root"
    )
    assert all(r["state"] == "unresolved" for r in state["relations"].values())
    assert state["relations"]  # Discovery was not gated on endpoint acceptance.


def test_illegal_current_card_property_is_retained_as_observation(inputs):
    state, _ = execute(inputs, Model(foreign_property=True))
    assert not state.get("properties")
    assert any(
        "property_outside_current_card_or_missing_value: urn:test:foreign" == o["reason"]
        for row in state["observations"].values()
        for o in row.get("alignment_outcomes", {}).values()
    )


def test_shared_field_preserves_separate_subject_card_outcomes_and_actual_mapping(inputs):
    original = Model()

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == "discover":
            color = next(f["field_id"] for f in payload["fields"] if f["label"] == "颜色")
            answer["document_field_ids"] = [color]
            answer["entities"][1]["field_ids"].append(color)
        return answer

    state, _ = execute(inputs, model)
    observation = next(o for o in state["observations"].values() if o["label"] == "颜色")
    outcomes = list(observation["alignment_outcomes"].values())
    assert len(outcomes) == 3
    root = next(o for o in outcomes if o["subject_id"] == "document")
    assert root == {
        "subject_id": "document",
        "class_iri": "urn:test:Report",
        "predicate_iris": [],
        "state": "unmatched",
        "reason": "当前类型卡无合法属性可对齐",
    }
    mapped = [o for o in outcomes if o["state"] == "mapped"]
    assert len(mapped) == 2
    assert len({o["subject_id"] for o in mapped}) == 2
    assert all(o["class_iri"] == "urn:test:Thing" for o in mapped)
    assert all(o["predicate_iris"] == ["urn:test:color"] for o in mapped)
    properties = list(state["properties"].values())
    assert len(properties) == 2
    assert all(
        p["alignment_class_iri"] == "urn:test:Thing"
        and p["state"] == "accepted"
        and p["field_id"] == observation["field_id"]
        for p in properties
    )
    hint = next(o for o in state["observations"].values() if o["kind"] == "relation")
    assert hint["subject_id"] != hint["object_id"]


def test_split_property_menus_do_not_turn_one_unmatched_menu_into_global_rejection(
    inputs,
    monkeypatch,
):
    from app.services.document_harness import controller
    from app.services.document_harness.ontology import PropertyCard

    ir, catalog = inputs
    card = catalog.classes["urn:test:Thing"]
    extra = PropertyCard(iri="urn:test:size", label="尺寸", description="对象尺寸")
    catalog = catalog.model_copy(
        update={
            "classes": {
                **catalog.classes,
                card.iri: card.model_copy(update={"properties": (*card.properties, extra)}),
            }
        }
    )
    original_size = controller.request_size

    def size(stage, payload, schema):
        if stage == "property_alignment" and len(payload["card"]["properties"]) > 1:
            return 1000000
        return original_size(stage, payload, schema)

    monkeypatch.setattr(controller, "request_size", size)
    original = Model()

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == "property_alignment":
            menu = {p["iri"] for p in payload["card"]["properties"]}
            for choice in answer["properties"].values():
                if any(mapping["predicate_iri"] not in menu for mapping in choice["mappings"]):
                    choice.update(mappings=[], reason="尺寸菜单不匹配颜色")
        return answer

    state, _ = execute((ir, catalog), model, max_request_bytes=32768)
    row = next(o for o in state["observations"].values() if o["label"] == "颜色")
    outcomes = {tuple(a["predicate_iris"]): a for a in row["alignment_outcomes"].values()}
    assert outcomes[("urn:test:size",)]["state"] == "unmatched"
    assert outcomes[("urn:test:size",)]["reason"] == "尺寸菜单不匹配颜色"
    assert outcomes[("urn:test:color",)]["state"] == "mapped"
    assert next(iter(state["properties"].values()))["state"] == "accepted"


@pytest.mark.parametrize("change", ["omit", "extra"])
def test_controller_rejects_inexact_property_keys_even_if_model_skips_dynamic_schema(
    inputs,
    change,
):
    base = Model()

    def model(stage, payload, schema):
        response = base(stage, payload, schema)
        if stage == "property_alignment" and payload["property_field_ids"]:
            if change == "omit":
                response["properties"].pop(payload["property_field_ids"][0])
            else:
                response["properties"]["unknown"] = {
                    "mappings": [],
                    "reason": "not a requested field",
                }
        return response

    with pytest.raises(ValueError, match="property_alignment_field_set_mismatch"):
        execute(inputs, model)


def test_explicit_null_property_mapping_retains_each_original_field_and_reason(inputs):
    base = Model()

    def model(stage, payload, schema):
        response = base(stage, payload, schema)
        if stage == "property_alignment":
            fields = {field["field_id"]: field for field in payload["subject"]["fields"]}
            response["properties"] = {
                key: {
                    "mappings": [],
                    "reason": "No legal match for " + fields[key]["label"],
                }
                for key in reversed(payload["property_field_ids"])
            }
        return response

    state, _ = execute(inputs, model)
    assert not state.get("properties")
    observations = {
        row["label"]: row
        for row in state["observations"].values()
        if any(
            outcome["reason"].startswith("No legal match for ")
            for outcome in row.get("alignment_outcomes", {}).values()
        )
    }
    assert set(observations) == {"名称", "颜色", "对象"}
    for field in state["fields"].values():
        if field["label"] in observations:
            row = observations[field["label"]]
            outcomes = list(row["alignment_outcomes"].values())
            assert len(outcomes) == 1
            assert outcomes[0]["reason"] == "No legal match for " + field["label"]
            assert outcomes[0]["class_iri"] == "urn:test:Thing"
            assert outcomes[0]["state"] == "unmatched"
            assert row["evidence"] == field["evidence"]
            assert any(field["value"] in ref["text"] for ref in row["evidence"])


def test_document_fields_are_aligned_without_body_entities(inputs):
    ir, _ = inputs
    graph = Graph().parse(
        data="""
        @prefix : <urn:test:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class .
        :color a owl:DatatypeProperty ; rdfs:domain :Report ; rdfs:label "颜色" ;
          rdfs:range xsd:string .
    """,
        format="turtle",
    )
    catalog = catalog_from_graph(graph, "urn:test:Report")
    calls = []

    def model(stage, payload, schema):
        calls.append(stage)
        if stage == "discover":
            return {
                "entities": [],
                "relation_hints": [],
                "unowned_fields": [],
                "complete": True,
                "document_source_fields": [],
                "document_field_ids": [
                    f["field_id"] for f in payload["fields"] if f["label"] == "颜色"
                ],
            }
        if stage == "property_alignment":
            assert payload["subject"]["role"] == "document_root"
            return {
                "properties": {
                    payload["property_field_ids"][0]: {
                        "mappings": [
                            {
                                "predicate_iri": "urn:test:color",
                                "confidence": 0.99,
                                "value_component": "whole",
                                "value_quote": None,
                            }
                        ],
                        "reason": "原字段",
                    }
                },
            }
        assert stage == "evidence_review"
        return {
            "type_concerns": [],
            "judgments": {
                c["id"]: {
                    "verdict": "accepted",
                    "confidence": 0.99,
                    "evidence": c["evidence"],
                    "reason": "原值和归属明确",
                }
                for c in payload["candidates"]
            },
        }

    state, _ = execute((ir, catalog), model)
    assert calls == ["discover", "property_alignment", "evidence_review"]
    assert [p["value"] for p in state["properties"].values()] == ["red"]
    assert all(p["state"] == "accepted" for p in state["properties"].values())


def test_discovery_oversized_input_splits_before_model_call(inputs, tmp_path):
    from app.services.document_harness.model import request_size
    from app.services.document_harness.protocols import stage_schema

    _, catalog = inputs
    doc = Document()
    doc.add_heading("Reading", 1)
    for i in range(12):
        doc.add_paragraph(f"Field {i}: " + "detail " * 45)
    path = tmp_path / "large.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    calls = []
    # Keep a bounded source allowance as the prompt/schema evolves. The test
    # still requires an oversized parent, multiple calls and a complete split.
    budget = request_size("discover", {}, stage_schema("discover")) + 3500

    def model(stage, payload, schema):
        assert stage == "discover"
        assert request_size(stage, payload, schema) <= budget
        calls.append(payload)
        return {
            "entities": [],
            "relation_hints": [],
            "document_field_ids": [],
            "document_source_fields": [],
            "unowned_fields": [],
            "complete": True,
        }

    engine = Engine(
        ir=ir,
        catalog=catalog,
        state={},
        invoke=model,
        save=lambda _: None,
        should_stop=lambda: False,
        max_request_bytes=budget,
    )
    engine.windows = build_windows(ir, max_chars=4800)
    engine.run()
    assert any(row["children"] for row in engine.state["windows"].values())
    assert len(calls) > 1
    assert all(r["entity_phase"] == "done" for r in engine.state["windows"].values())
    assert engine.state["cursor"]["main"]["scope_complete"]


def test_reading_limit_does_not_claim_full_scope(inputs):
    calls = []

    def incomplete(stage, payload, schema):
        assert stage == "discover"
        calls.append(payload)
        return {
            "entities": [],
            "relation_hints": [],
            "document_field_ids": [],
            "document_source_fields": [],
            "unowned_fields": [],
            "complete": False,
        }

    state, _ = execute(inputs, incomplete)
    assert state["cursor"]["main"]["stage"] == "complete"
    assert not state["cursor"]["main"]["scope_complete"]
    assert 1 < len(calls) <= 31
    assert any(row["children"] for row in state["windows"].values())


def test_prose_field_can_be_proposed_without_a_preparsed_field(inputs):
    original = Model()

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == "discover":
            answer["entities"][0]["source_fields"] = [
                {
                    "label": quote(payload["sources"], "contains"),
                    "value": quote(payload["sources"], "Beta"),
                }
            ]
        return answer

    state, _ = execute(inputs, model)
    field = next(f for f in state["fields"].values() if f["label"] == "contains")
    owner = next(e for e in state["entities"].values() if e["label"] == "Alpha")
    assert field["id"] in owner["field_ids"]
    assert any(
        f["label"] == "contains"
        for stage, payload, _ in original.calls
        if stage == "type_alignment"
        for e in payload["entities"]
        for f in e["fields"]
    )


@pytest.mark.parametrize("overlap,confirmed", [(True, True), (True, False), (False, True)])
def test_untyped_overlapping_referent_gets_independent_menu_without_borrowing_verdict(
    inputs,
    overlap,
    confirmed,
):
    original = Model()
    comparisons = []

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == "discover":
            first, second = answer["entities"]
            first.update(name=None, anchor=quote(payload["sources"], "Alpha contains Beta"))
            if overlap:
                second.update(name=None, anchor={**first["anchor"], "text": "Beta"})
        if stage == "type_alignment":
            if len(payload["entities"]) > 1:
                untyped = next(e for e in payload["entities"] if e["role"] == "component")
                answer["entities"][untyped["entity_id"]].update(
                    class_iri=None,
                    reason="批量回答将类型留给另一个候选",
                )
            else:
                comparisons.append(payload)
                if not confirmed:
                    for value in answer["entities"].values():
                        value.update(class_iri=None, reason="独立核对后仍无合适类型")
        return answer

    state, _ = execute(inputs, model, stop=lambda: len(original.calls) >= (3 if overlap else 2))
    referent = next(e for e in state["entities"].values() if e["role"] == "component")
    assert len(comparisons) == int(overlap)
    assert (referent["class_iri"] == "urn:test:Thing") == (overlap and confirmed)
    assert referent["state"] != "accepted"
    assert len(state["entities"]) == 3  # No merge, even when physical referents overlap.


def test_overlapping_entity_hypotheses_are_reviewed_in_separate_batches(inputs):
    original = Model()

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == "discover":
            first, second = answer["entities"]
            anchor = quote(payload["sources"], "Alpha contains Beta")
            first.update(name=None, anchor=anchor)
            second.update(name=None, anchor={**anchor, "text": "Beta"})
        return answer

    state, _ = execute(inputs, model)
    reviews = [payload for stage, payload, _ in original.calls if stage == "entity_review"]
    assert len(reviews) == 2 and all(len(r["candidates"]) == 1 for r in reviews)
    assert len(state["entities"]) == 3


def test_combined_review_quotes_cover_a_long_source_value(inputs):
    ir, catalog = inputs
    engine = Engine(
        ir=ir,
        catalog=catalog,
        state={
            "entities": {
                "document": {"role": "document_root"},
            },
            "fields": {"f": {"value_evidence": [{"source_id": "u", "start": 0, "end": 1800}]}},
        },
        invoke=None,
        save=None,
        should_stop=lambda: False,
    )
    candidate = {"subject_id": "document", "field_id": "f"}
    refs = [
        {"source_id": "u", "start": 0, "end": 1000},
        {"source_id": "u", "start": 1000, "end": 1800},
    ]
    assert engine.review_sources_cover("properties", candidate, refs)
    refs[1]["start"] = 1001
    assert not engine.review_sources_cover("properties", candidate, refs)


def test_explicit_cross_window_reference_preserves_endpoint_gate_without_name_merging(
    inputs, tmp_path
):
    _, catalog = inputs
    doc = Document()
    doc.add_heading("第一份对象记录", 1)
    doc.add_paragraph("名称：Alpha")
    doc.add_paragraph("颜色：red")
    doc.add_heading("第二份对象记录", 1)
    doc.add_paragraph("名称：Beta")
    doc.add_paragraph("Alpha contains Beta")
    path = tmp_path / "cross-window.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    original = Model()

    def model(stage, payload, schema):
        if stage == "discover":
            field = next(f for f in payload["fields"] if f["label"] == "名称")
            name = field["value"]
            return {
                "document_field_ids": [],
                "document_source_fields": [],
                "unowned_fields": [],
                "entities": [
                    {
                        "local_id": "e",
                        "name": quote(payload["sources"], name),
                        "anchor": quote(payload["sources"], name),
                        "role": "container" if name == "Alpha" else "component",
                        "evidence": [quote(payload["sources"], name)["source_id"]],
                        "field_ids": [f["field_id"] for f in payload["fields"]],
                        "source_fields": [],
                    }
                ],
                "relation_hints": [],
                "reference_cues": [
                    {
                        "local_subject_id": "e",
                        "reference": quote(payload["sources"], "Alpha"),
                        "relation_label": "contains",
                        "direction": "incoming",
                        "kind": "explicit_reference",
                        "polarity": "positive",
                        "conditions": [],
                        "evidence": [quote(payload["sources"], "Alpha contains Beta")["source_id"]],
                    }
                ]
                if name == "Beta"
                else [],
                "complete": True,
            }
        answer = original(stage, payload, schema)
        if stage in {"entity_review", "evidence_review"} and not any(
            "Alpha contains Beta" in s["text"] for s in payload["sources"]
        ):
            for row in payload["candidates"]:
                judgment = answer["judgments"][row["id"]]
                if row["kind"] == "entities" and row["label"] == "Alpha":
                    judgment.update(verdict="unresolved", reason="Object role needs context")
        return answer

    windows = build_windows(ir, max_sources=3)
    assert len(windows) == 2  # Explicit budget boundary, not a section-per-call assumption.
    state, _ = execute((ir, catalog), model, windows=windows)
    assert len([e for e in state["entities"].values() if e["role"] != "document_root"]) == 2
    alpha = next(e for e in state["entities"].values() if e["label"] == "Alpha")
    assert alpha["state"] == "unresolved"
    relation = next(iter(state["relations"].values()))
    assert relation["state"] == "unresolved"
    assert relation["verification"]["semantic_verdict"] == "accepted"
    before = len(original.calls)
    runner = Engine(
        ir=ir,
        catalog=catalog,
        state=state,
        invoke=model,
        save=lambda _: None,
        should_stop=lambda: False,
    )
    runner.commit({"entities": {alpha["id"]: {**alpha, "state": "accepted"}}})
    assert runner.state["relations"][relation["id"]]["state"] == "accepted"
    assert len(original.calls) == before
    for stage, payload, _ in original.calls:
        if stage == "relation_alignment":
            available = {source["source_id"] for source in payload["sources"]}
            assert all(ref in available for obj in payload["entities"] for ref in obj["evidence"])


def test_pause_after_discovery_preserves_candidate_and_resumes_next_stage(inputs):
    model = Model()
    state, _ = execute(inputs, model, stop=lambda: len(model.calls) == 1)
    assert state["cursor"]["main"]["phase"] == "reading"
    assert len(state["cursor"]["main"]["active_batches"]) == 1
    assert len(state["entities"]) == 1  # Paid answer awaits application on continue.
    continued, _ = execute(inputs, model, state=state)
    assert sum(stage == "discover" for stage, _, _ in model.calls) == 1
    assert continued["cursor"]["main"]["stage"] == "complete"


def test_wrong_source_and_ambiguous_quote_are_not_repaired_to_another_unit(inputs):
    ir, _ = inputs
    window = build_windows(ir)[0]
    wrong = next(s for s in window.sources if "颜色" in s["text"])
    with pytest.raises(ValueError, match="mismatch"):
        window.resolve(ir, {"source_id": wrong["source_id"], "text": "Alpha"})
    with pytest.raises(ValueError, match="empty"):
        window.resolve(ir, {"source_id": wrong["source_id"], "text": " "})


def test_review_schema_requires_evidence_for_accepted():
    import jsonschema

    schema = stage_schema("evidence_review", source_ids=["S1"], candidate_ids=["C1"])
    value = {
        "type_concerns": [],
        "judgments": {
            "C1": {
                "verdict": "accepted",
                "confidence": 0.99,
                "reason": "x",
                "evidence": [],
            }
        },
    }
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(value, schema)
    value["judgments"]["C1"]["verdict"] = "unresolved"
    jsonschema.validate(value, schema)


@pytest.mark.parametrize(
    "target_stage,field,error",
    [
        ("type_alignment", "entities", "type_alignment_answer_set_mismatch"),
        ("evidence_review", "judgments", "evidence_review_answer_set_mismatch"),
    ],
)
def test_fixed_answer_set_is_also_checked_after_model_transport(inputs, target_stage, field, error):
    original = Model()

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == target_stage:
            answer[field].pop(next(iter(answer[field])))
        return answer

    with pytest.raises(ValueError, match=error):
        execute(inputs, model)


def test_fixed_answer_objects_do_not_depend_on_answer_key_order(inputs):
    original = Model()

    def model(stage, payload, schema):
        answer = original(stage, payload, schema)
        for field in (
            ("entities", "judgments")
            if stage
            in {
                "type_alignment",
                "evidence_review",
            }
            else ()
        ):
            if field in answer:
                answer[field] = dict(reversed(list(answer[field].items())))
        return answer

    state, _ = execute(inputs, model)
    assert all(row["state"] == "accepted" for row in state["entities"].values())
    assert all(row["state"] == "accepted" for row in state["properties"].values())
    assert all(row["state"] == "accepted" for row in state["relations"].values())
