"""An interleaved repeated identifier needs separate physical mentions on correction."""

from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from test_grouped_alignment import engine_fixture, proposal
from test_lookup import NS
from test_lookup import lookup_fixture as lookup_fixture

from app.services.document_harness.controller import Engine, Paused
from app.services.document_harness.referents import (
    initial_spans,
    register_span,
    run_referent_alignment,
)
from app.services.document_harness.source import identity
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def saved_overlapping_proposal(fixture, tmp_path):
    doc = Document()
    doc.add_paragraph("Use A1/A2 objects; later use A1 again.")
    path = tmp_path / "repeated-identifier.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    engine, window, _ = engine_fixture((ir, *fixture[1:]))
    # A collective draft genuinely covering both mentions still gets a complete
    # alternative, while disjoint drafts are separate local tasks (tested below).
    ref = window.resolve(ir, {"source_id": "S1", "text": window.sources[0]["text"],
                              "occurrence": None})
    engine.state["entities"]["old"].update(referent=ref, name=ref, evidence=[ref])
    entities = engine.entities(window)
    task_id = identity("referent_group", NS + "Object", sorted(
        (e["referent"]["source_id"], e["referent"]["start"], e["referent"]["end"])
        for e in entities
    ))
    context = {"key_candidates": [], "source_status": [], "issues": []}
    spans, identifiers = initial_spans(engine, window, entities, context)
    expressions = []
    for key, text, occurrence in (
        ("whole", "A1/A2", None), ("a", "A1", 0),
        ("b", "A2", None), ("later", "A1", 1),
    ):
        ref = window.resolve(ir, {"source_id": "S1", "text": text, "occurrence": occurrence})
        span_id = register_span(ir, spans, ref)
        identifiers.append(span_id)
        expressions.append({"id": key, "span_id": span_id, "property_iri": NS + "code"})
    proposal = {
        "new_spans": [], "expressions": expressions,
        "partitions": [
            {"id": "single", "members": [
                {"id": "whole", "expression_ids": ["whole"]},
                {"id": "later", "expression_ids": ["later"]},
            ], "reason": "Compound identifier and a later mention"},
            {"id": "members", "members": [
                {"id": "a", "expression_ids": ["a", "later"]},
                {"id": "b", "expression_ids": ["b"]},
            ], "reason": "Incorrectly merge repeated A1 across A2"},
        ],
    }
    engine.commit({"referent_work": {task_id: {
        "key_context": context, "span_catalog": spans,
        "identifier_span_ids": list(dict.fromkeys(identifiers)), "proposal": proposal,
    }}})
    return engine, window, task_id, proposal


def corrected_response(stage, payload, schema):
    if stage == "referent_candidates":
        feedback = payload["proposal_feedback"]
        assert feedback["issue"] == "overlapping_members_in_same_partition"
        proposal = deepcopy(feedback["rejected_proposal"])
        proposal["partitions"][1]["members"] = [
            {"id": key, "expression_ids": [key]} for key in ("a", "b", "later")
        ]
        return proposal
    assert stage == "referent_selection"
    assert len(payload["proposal"]["partitions"][1]["members"]) == 3
    result = {
        "verdict": "supported", "selected_partition_id": "members", "confidence": 0.9,
        "evidence_span_ids": [max(payload["evidence_spans"], key=lambda s: len(s["text"]))[
            "span_id"
        ]], "reason": "Three distinct physical mentions; identity remains unchecked",
    }
    jsonschema.validate(result, schema)
    return result


@pytest.mark.parametrize("pause_before_correction", [False, True])
def test_saved_overlap_can_be_corrected_once_without_merging_repeated_mentions(
    lookup_fixture, tmp_path, pause_before_correction,
):
    engine, window, task_id, rejected = saved_overlapping_proposal(lookup_fixture, tmp_path)
    if pause_before_correction:
        def pause(stage, payload, schema):
            assert stage == "referent_candidates"
            assert payload["proposal_feedback"]["rejected_proposal"] == rejected
            raise Paused()

        engine.invoke = pause
        with pytest.raises(Paused):
            run_referent_alignment(engine, window)
        work = engine.state["referent_work"][task_id]
        assert work["proposal_corrected"] and "proposal" not in work
        assert list(engine.state["entities"]) == ["document", "old"]
        engine = Engine(ir=engine.ir, catalog=engine.catalog, state=deepcopy(engine.state),
                        invoke=corrected_response, save=lambda _: None,
                        should_stop=lambda: False, lookup=engine.lookup,
                        calls=engine._memory_calls)
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        return corrected_response(stage, payload, schema)

    engine.invoke = invoke
    run_referent_alignment(engine, window)
    assert calls == ["referent_candidates", "referent_selection"]
    work = engine.state["referent_work"][task_id]
    assert work["done"] and work["proposal_corrected"]
    bound = [e for e in engine.entities(window) if e.get("identity_binding")]
    assert sorted(e["label"] for e in bound) == ["A1", "A1", "A2"]
    assert len({e["referent"]["start"] for e in bound}) == 3
    assert all(b["identity_status"] == "not_checked" for e in bound
               for b in e["identity_binding"]["identifiers"])
    run_referent_alignment(engine, window)
    assert calls == ["referent_candidates", "referent_selection"]


def test_still_overlapping_correction_stays_failed_and_is_not_retried_on_resume(
    lookup_fixture, tmp_path,
):
    engine, window, task_id, rejected = saved_overlapping_proposal(lookup_fixture, tmp_path)
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        assert stage == "referent_candidates"
        return deepcopy(rejected)

    engine.invoke = invoke
    with pytest.raises(ValueError, match="overlapping_members_in_same_partition"):
        run_referent_alignment(engine, window)
    restored = Engine(ir=engine.ir, catalog=engine.catalog, state=deepcopy(engine.state),
                      invoke=invoke, save=lambda _: None, should_stop=lambda: False,
                      lookup=engine.lookup, calls=engine._memory_calls)
    with pytest.raises(ValueError, match="overlapping_members_in_same_partition"):
        run_referent_alignment(restored, window)
    assert calls == ["referent_candidates"]
    assert not restored.state["referent_work"][task_id].get("done")
    assert list(restored.state["entities"]) == ["document", "old"]
    assert not restored.state["fields"]


def test_seven_disjoint_mentions_in_one_source_have_independent_complete_partitions(
    lookup_fixture, tmp_path,
):
    doc = Document()
    doc.add_paragraph("Use A1/A2 objects" + "; then use A1" * 6)
    path = tmp_path / "many-mentions.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    engine, window, _ = engine_fixture((ir, *lookup_fixture[1:]))
    original = deepcopy(engine.state["entities"]["old"])
    for occurrence in range(1, 7):
        key = f"mention-{occurrence}"
        ref = window.resolve(ir, {"source_id": "S1", "text": "A1", "occurrence": occurrence})
        engine.state["entities"][key] = {
            **original, "id": key, "label": "A1", "referent": ref,
            "name": ref, "evidence": [ref],
        }
        engine.state["window_entities"][window.id]["ids"].append(key)
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        assert len(payload["mentions"]) == 1
        mention = payload["mentions"][0]
        if stage == "referent_candidates":
            for candidate in payload["key_candidates"]:
                for occurrence in candidate["key_occurrences"]:
                    ref = window.resolve(ir, occurrence["quote"])
                    assert mention["start"] <= ref["start"] < ref["end"] <= mention["end"]
            available = {s["text"]: s["span_id"] for s in payload["evidence_spans"]
                         if s["identifier_candidate"]}
            if mention["label"] == "A1/A2":
                if "A1/A2" not in available:
                    result = {"new_spans": [{"source_id": "S1", "text": "A1/A2",
                                             "occurrence": None}],
                              "expressions": [], "partitions": []}
                else:
                    result = proposal(available)
            else:
                assert set(available) == {"A1"}
                result = {"new_spans": [], "expressions": [
                    {"id": "a", "property_iri": NS + "code", "span_id": available["A1"]},
                ], "partitions": [{"id": "single", "members": [
                    {"id": "a", "expression_ids": ["a"]},
                ], "reason": "This physical occurrence alone"}]}
        else:
            assert stage == "referent_selection"
            result = {
                "verdict": "supported",
                "selected_partition_id": "members" if mention["label"] == "A1/A2" else "single",
                "confidence": 0.9,
                "evidence_span_ids": [max(payload["evidence_spans"],
                                          key=lambda s: len(s["text"]))["span_id"]],
                "reason": "The local physical mention is complete",
            }
        jsonschema.validate(result, schema)
        return result

    engine.invoke = invoke
    run_referent_alignment(engine, window)
    work = engine.state["referent_work"]
    assert len(work) == 7 and all(w["done"] for w in work.values())
    bound = [e for e in engine.entities(window) if e.get("identity_binding")]
    assert len(bound) == 8
    assert len({e["referent"]["start"] for e in bound}) == 8
    assert all(not e.get("referent_unresolved") for e in bound)
    count = len(calls)
    run_referent_alignment(engine, window)
    assert len(calls) == count
