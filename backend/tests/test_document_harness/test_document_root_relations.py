"""A document without a textual anchor must still get source-reviewed first hops."""

from copy import deepcopy

import pytest
from docx import Document
from rdflib import Graph
from test_controller import Model, execute, quote

from app.services.document_harness.controller import Engine
from app.services.document_harness.evidence_gate import route_semantic_review
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.planning import (
    admit_relation_work,
    build_source_index,
    collect_relation_seeds,
)
from app.services.document_harness.projection import build_graph_base
from app.services.document_harness.ranking import discovery_guidance, root_target_iris
from app.services.document_harness.source import reference
from app.services.document_harness.work import DEFAULT_POLICY
from app.services.extraction.word_analysis import analyze_word_core


@pytest.fixture
def root_inputs(tmp_path):
    doc = Document()
    doc.add_heading("Alpha", 1)
    doc.add_paragraph("原料药Alpha计划生产1批，预计批量5 kg。")
    doc.add_paragraph("背景：其他报告中的原料药Beta，仅作对照。")
    doc.add_paragraph("本次合成路线包括两次反应。")
    path = tmp_path / "root.docx"
    doc.save(path)
    catalog = catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:root:> . @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class ; rdfs:label "报告" .
        :Drug a owl:Class ; rdfs:label "原料药" .
        :SpecialDrug a owl:Class ; rdfs:subClassOf :Drug .
        :Route a owl:Class ; rdfs:label "合成路线" ; rdfs:comment "完整工艺路线" .
        :Step a owl:Class ; rdfs:label "合成步骤" .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :Drug ;
          rdfs:label "描述" ; rdfs:comment "报告描述的正文对象，不含背景和他文引用。" .
        :hasRoute a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :Route .
        :hasStep a owl:ObjectProperty ; rdfs:domain :Route ; rdfs:range :Step .
    ''', format="turtle"), "urn:root:Report")
    return analyze_word_core(path).ir, catalog


def state_for(ir):
    root = {"id": "document", "label": "Alpha", "role": "document_root",
            "class_iri": "urn:root:Report", "state": "accepted", "evidence": [],
            "field_ids": [], "window_id": None}
    entities = {"document": root}
    for key, name, cls in (("a", "Alpha", "Drug"), ("b", "Beta", "Drug"),
                           ("route", "本次合成路线包括两次反应。", "Route")):
        unit = next(u for u in ir.evidence_units if name in u.text and u.kind != "heading")
        start = unit.text.index(name)
        ref = reference(ir, unit.evidence_id, start, start + len(name))
        entities[key] = {"id": key, "label": name, "role": "正文对象", "referent": ref,
                         "evidence": [ref], "field_ids": [], "class_iri": "urn:root:"+cls,
                         "state": "accepted", "window_id": None}
    return {"entities": entities, "fields": {}, "hints": {}, "work": {}}


def seeds_for(ir, catalog, state):
    return collect_relation_seeds(ir, catalog, state, state, build_source_index(ir, state),
                                  DEFAULT_POLICY)


def test_root_targets_survive_zero_ranking_and_root_is_not_a_discoverable_entity(root_inputs):
    ir, catalog = root_inputs
    seen = []

    def invoke(stage, payload, schema):
        assert stage == "discover"
        seen.append((deepcopy(payload), deepcopy(schema)))
        return {"entities": [], "document_field_ids": [], "document_source_fields": [],
                "unowned_fields": [], "relation_hints": [], "complete": True}

    engine = Engine(ir=ir, catalog=catalog, state={}, invoke=invoke, save=lambda _: None,
                    should_stop=lambda: False,
                    rank=lambda *_: {"snapshot_id": catalog.snapshot_id, "selected_iris": []})
    engine.run()
    payload, schema = seen[0]
    assert set(root_target_iris(catalog)) == {"urn:root:Drug", "urn:root:Route"}
    assert set(engine.state["windows"][engine.windows[0].id]["guidance_class_iris"]) == {
        "urn:root:Drug", "urn:root:Route",
    }
    assert payload["document"]["entity_id"] == "document"
    assert {r["iri"] for r in payload["document"]["relation_guidance"]} == {
        "urn:root:describes", "urn:root:hasRoute",
    }
    assert set(schema["$defs"]["Mention"]["properties"]["candidate_class_iri"]["enum"]) == {
        "urn:root:Drug", "urn:root:Route", None,
    }
    assert all(c["iri"] != catalog.root_class_iri for c in payload["schema_guidance"]["classes"])
    guidance = discovery_guidance(catalog, ["urn:root:Drug", "urn:root:Step"])
    assert len(guidance["classes"]) == 3  # Ranked direct target is not duplicated.


def test_body_targets_create_work_without_a_root_anchor_or_literal_predicate(root_inputs):
    ir, catalog = root_inputs
    state = state_for(ir)
    seeds = seeds_for(ir, catalog, state)
    assert {(s["subject_id"], s["predicate_iri"], tuple(s["object_ids"])) for s in seeds} == {
        ("document", "urn:root:describes", ("a",)),
        ("document", "urn:root:describes", ("b",)),
        ("document", "urn:root:hasRoute", ("route",)),
    }
    work = admit_relation_work(seeds, state, catalog, DEFAULT_POLICY)
    assert all(w["status"] == "ready" for w in work)
    assert all(w["input"]["polarity_hint"] == "uncertain" for w in work)
    assert not state.get("relations")  # Neither range compatibility nor background is a fact.
    assert not state["entities"]["document"].get("referent")
    for seed in seeds:
        assertion = {"id": "proposal", "subject_id": "document",
                     "object_id": seed["object_ids"][0], "predicate_iri": seed["predicate_iri"],
                     "evidence": seed["clue_refs"], "polarity": "positive", "conditions": []}
        assert route_semantic_review(ir, catalog, state, assertion)["action"] == "semantic"


def test_root_fallback_keeps_explicit_negation_conditions_and_paid_work(root_inputs):
    ir, catalog = root_inputs
    state = state_for(ir)
    unit = ir.unit(state["entities"]["a"]["referent"]["source_id"])
    state["hints"]["hint"] = {
        "id": "hint", "subject_id": "document", "object_id": "a", "label": "描述",
        "evidence": [reference(ir, unit.evidence_id, 0, len(unit.text))],
        "polarity": "negative", "conditions": ["仅在对照条件下"],
    }
    work = admit_relation_work(seeds_for(ir, catalog, state), state, catalog, DEFAULT_POLICY)
    root_work = next(w for w in work if w["input"]["object_ids"] == ["a"])
    assert root_work["input"]["polarity_hint"] == "negative"
    assert root_work["input"]["condition_hints"] == ["仅在对照条件下"]
    assert root_work["input"]["priority"] == "explicit"
    root_work.update(status="done", last_call_key="paid-call")
    state["work"][root_work["id"]] = root_work
    repeated = admit_relation_work(seeds_for(ir, catalog, state), state, catalog, DEFAULT_POLICY)
    saved = next(w for w in repeated if w["id"] == root_work["id"])
    assert saved["status"] == "done" and saved["last_call_key"] == "paid-call"


def test_direct_target_summaries_preserve_intersection_constraints():
    catalog = catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:root:> . @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class . :A a owl:Class . :B a owl:Class .
        :Both a owl:Class ; rdfs:subClassOf :A, :B .
        :Child a owl:Class ; rdfs:subClassOf :Both .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ;
          rdfs:range [owl:intersectionOf (:A :B)] .
    ''', format="turtle"), "urn:root:Report")
    assert root_target_iris(catalog) == ["urn:root:Both"]


def test_explicit_root_group_is_not_split_into_structural_singletons(root_inputs):
    ir, catalog = root_inputs
    state = state_for(ir)
    state["hints"]["group"] = {
        "id": "group", "subject_id": "document", "object_ids": ["a", "b"],
        "label": "描述", "evidence": [state["entities"][key]["referent"] for key in ("a", "b")],
        "polarity": "uncertain", "conditions": ["二者择一"],
    }
    describes = [s for s in seeds_for(ir, catalog, state)
                 if s["predicate_iri"] == "urn:root:describes"]
    assert len(describes) == 1 and describes[0]["object_ids"] == ["a", "b"]
    assert describes[0]["condition_hints"] == ["二者择一"]


@pytest.mark.parametrize("case", ["heading", "toc", "rejected", "wrong_type"])
def test_non_body_or_ineligible_targets_do_not_generate_root_work(root_inputs, case):
    ir, catalog = root_inputs
    state = state_for(ir)
    state["entities"] = {k: v for k, v in state["entities"].items() if k in {"document", "a"}}
    entity = state["entities"]["a"]
    if case == "heading":
        unit = ir.evidence_units[0]
        entity["referent"] = reference(ir, unit.evidence_id, 0, len(unit.text))
    elif case == "toc":
        unit = ir.unit(entity["referent"]["source_id"])
        updated = unit.model_copy(update={"navigation_role": "toc_entry"})
        ir = ir.model_copy(update={"evidence_units": [
            updated if u.evidence_id == unit.evidence_id else u for u in ir.evidence_units
        ]})
    elif case == "rejected":
        entity["state"] = "rejected"
    else:
        entity["class_iri"] = "urn:root:Step"  # Reachable only at the second hop.
    assert seeds_for(ir, catalog, state) == []


class RootModel(Model):
    def __call__(self, stage, payload, schema):
        sources = payload["sources"]
        if stage == "discover":
            self.calls.append((stage, deepcopy(payload), deepcopy(schema)))
            return {"document_field_ids": [], "document_source_fields": [],
                    "unowned_fields": [], "complete": True, "relation_hints": [],
                    "entities": [{"local_id": key, "candidate_class_iri": "urn:root:Drug",
                                  "name": quote(sources, name),
                                  "anchor": quote(sources, text), "role": "原料药",
                                  "evidence": [quote(sources, text)["source_id"]],
                                  "field_ids": [], "source_fields": []}
                                 for key, name, text in (
                                     ("a", "Alpha", "原料药Alpha计划生产1批，预计批量5 kg。"),
                                     ("b", "Beta", "背景：其他报告中的原料药Beta，仅作对照。"))]}
        result = super().__call__(stage, payload, schema)
        if stage == "type_alignment":
            for choice in result["entities"].values():
                choice["class_iri"] = "urn:root:Drug"
        if stage == "relation_alignment":
            for item in payload["items"]:
                result["proposals"][item["candidate_id"]]["polarity"] = "positive"
        if stage == "evidence_review":
            for candidate in payload["candidates"]:
                if candidate.get("object", {}).get("label") == "Beta":
                    result["judgments"][candidate["id"]].update(
                        verdict="rejected", reason="他文中的背景对象，不是本报告描述的对象",
                    )
        return result


def test_full_flow_reviews_first_hops_and_rejects_background_without_merging_names(root_inputs):
    model = RootModel()
    state, _ = execute(root_inputs, model, rank=lambda *_: {
        "snapshot_id": root_inputs[1].snapshot_id, "selected_iris": [],
    })
    graph = build_graph_base(state, root_inputs[1].model_dump(mode="json"))
    calls = [p for stage, p, _ in model.calls if stage == "relation_alignment"]
    assert calls and all(item["subject_id"] == "E0" for p in calls for item in p["items"])
    relations = [r for r in graph["relations"] if r["subject_id"] == "document"]
    accepted = [r for r in relations if r["state"] == "accepted"]
    assert len(accepted) == 1
    target = state["entities"][accepted[0]["object_id"]]
    assert target["label"] == state["entities"]["document"]["label"] == "Alpha"
    assert target["id"] != "document" and target["role"] != "document_root"
    assert len([r for r in relations if r["state"] == "rejected"]) == 1
    assert any(stage == "evidence_review" for stage, _, _ in model.calls)
