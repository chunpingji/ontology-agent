"""Co-reference never erases mentions or turns incomplete/contradictory evidence into identity."""

from copy import deepcopy
from itertools import combinations

import jsonschema
import pytest
from docx import Document
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.coreference import (
    candidate_pairs,
    canonical_mentions,
    pair_id,
    project_coreferences,
)
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.protocols import stage_schema
from app.services.document_harness.source import reference
from app.services.document_harness.work import DEFAULT_POLICY, make_work
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def case(tmp_path):
    doc = Document()
    texts = ["清洗设备甲（以下简称清洗单元）编号 EQ-17。", "清洗单元在周一使用。",
             "同名设备用于周二，编号 EQ-18；与清洗设备甲、清洗单元并非同一设备。",
             "另有一台未编号清洗设备。"]
    for i, text in enumerate(texts):
        doc.add_heading(f"第{i + 1}节", 1)
        doc.add_paragraph(text)
    path = tmp_path / "source.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    rdf = Graph().parse(data='''
        @prefix : <urn:test:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class . :Equipment a owl:Class . :Specific a owl:Class;
          rdfs:subClassOf :Equipment . :Batch a owl:Class .
        :describes a owl:ObjectProperty; rdfs:domain :Report; rdfs:range :Equipment, :Batch .
    ''', format="turtle")
    catalog = catalog_from_graph(rdf, "urn:test:Report")
    entities = {}
    for i, text in enumerate(texts):
        unit = next(u for u in ir.evidence_units if u.text == text)
        ref = reference(ir, unit.evidence_id, 0, len(text))
        entities[str(i)] = {"id": str(i), "label": text, "name": None, "role": "设备",
                           "class_iri": "urn:test:Equipment", "class_label": "设备",
                           "state": "accepted", "reason": "已核对", "field_ids": [],
                           "referent": ref, "evidence": [ref], "window_id": str(i)}
    entities["document"] = {"id": "document", "label": "报告", "role": "document_root",
                            "class_iri": "urn:test:Report", "class_label": "报告",
                            "state": "accepted", "reason": "任务类型", "evidence": [],
                            "field_ids": [], "window_id": None}
    state = {"entities": entities}
    add_work(state, catalog, [("0", "1")])
    return ir, catalog, state


def add_work(state, catalog, pairs):
    for left, right in pairs:
        work = make_work("coreference_review", {
            "left_mention_id": left, "right_mention_id": right,
            "clue_refs": [state["entities"][left]["referent"],
                          state["entities"][right]["referent"]],
        }, state, catalog, DEFAULT_POLICY)
        state.setdefault("work", {})[work["id"]] = work


def answer(payload, verdict="same", basis="explicit_alias"):
    sources = payload["sources"]
    by_id = {row["entity_id"]: row for row in payload["entities"]}
    judgments = {}
    for pair in payload["pairs"]:
        binding = None
        if verdict == "same" and basis == "explicit_alias":
            a, b = by_id[pair["left_id"]]["anchor"], by_id[pair["right_id"]]["anchor"]
            # The positive fixture has a literal alias declaration in the first section.
            binding = {
                "left": {**a, "text": "清洗设备甲"},
                "right": {**b, "text": "清洗单元"},
                "declaration": {**a, "text": "清洗设备甲（以下简称清洗单元）"},
            }
        judgments[pair["pair_id"]] = {
            "verdict": verdict, "basis": basis, "confidence": 0.96,
            "evidence": [s["source_id"] for s in sources],
            "proof": [binding["declaration"]] if binding else [],
            "alias_binding": binding, "reason": "根据给定原文核对",
        }
        if verdict == "different":
            source = next(s for s in sources if s["kind"] != "heading")
            judgments[pair["pair_id"]]["proof"] = [{"source_id": source["source_id"],
                                                       "text": source["text"], "occurrence": None}]
    return {"judgments": judgments}


def runner(case, invoke, *, state=None, stop=lambda: False, budget=32768, save=None):
    ir, catalog, initial = case
    state = deepcopy(state or initial)
    engine = Engine(ir=ir, catalog=catalog, state=state, invoke=invoke,
                    save=save or (lambda changes: None), should_stop=stop, max_request_bytes=budget)
    if not engine.state.get("windows"):
        engine.state["windows"] = {
            window.id: {**Engine.window_row(window, [i]), "entity_phase": "done",
                        "reading_state": "complete"}
            for i, window in enumerate(engine.windows)
        }
    engine.state.setdefault("cursor", {"main": {
        "stage": "planning", "entity_window_id": None, "active_batches": {}, "phase": "coreference",
        "scope_complete": False,
    }})
    return engine


def two(case):
    state = deepcopy(case[2])
    for key in ("2", "3"):
        del state["entities"][key]
    return state


def test_explicit_cross_section_alias_work_and_exact_proof(case):
    calls = []
    state = two(case)
    before = deepcopy(state)
    def invoke(stage, payload, schema):
        assert stage == "coreference_review"
        batch = next(iter(engine.state["cursor"]["main"]["active_batches"].values()))
        assert set(batch["source_bindings"]) == {s["source_id"] for s in payload["sources"]}
        assert batch["targets"] == [{
            "domain": "work", "id": row["id"], "dependency_hash": row["dependency_hash"],
        } for row in engine.state["work"].values() if row["status"] == "ready"]
        assert len({source["section"] for source in payload["sources"]}) == 2
        assert any(source["kind"] == "heading" for source in payload["sources"])
        result = answer(payload)
        jsonschema.validate(result, schema)
        calls.append(payload)
        return result
    engine = runner(case, invoke, state=state)
    engine.run()
    assert engine.state["entities"] == before["entities"]
    assert engine.state["cursor"]["main"]["stage"] == "complete"
    decision = next(iter(engine.state["coreferences"].values()))
    assert decision["verdict"] == "same" and decision["proof"]
    work = next(iter(engine.state["work"].values()))
    assert work["status"] == "done" and work["output_ids"] == [decision["id"]]
    assert work["applied_dependency_hash"] == work["dependency_hash"]
    for ref in decision["evidence"] + decision["proof"]:
        assert case[0].unit(ref["source_id"]).text[ref["start"]:ref["end"]] == ref["text"]
    engine.run()
    assert len(calls) == 1  # A completed run cannot cause fresh calls through a read/resume.


@pytest.mark.parametrize("defect", ["missing_endpoint", "outside_source", "fake_proof",
                                   "no_proof", "low_confidence", "weak_basis", "same_alias",
                                   "foreign_alias", "missing_alias"])
def test_wrong_or_insufficient_proof_cannot_merge(case, defect):
    def invoke(stage, payload, schema):
        result = answer(payload)
        judgment = result["judgments"]["P1"]
        if defect == "missing_endpoint":
            judgment["evidence"] = [payload["sources"][0]["source_id"]]
        elif defect == "outside_source":
            judgment["evidence"] = ["foreign"]
        elif defect == "fake_proof":
            judgment["proof"][0]["text"] = "原文不存在的别名定义"
        elif defect == "no_proof":
            judgment["proof"] = []
        elif defect == "low_confidence":
            judgment["confidence"] = 0.5
        elif defect == "same_alias":
            judgment["alias_binding"]["left"]["text"] = "清洗单元"
        elif defect == "foreign_alias":
            judgment["alias_binding"]["right"] = judgment["alias_binding"]["left"]
        elif defect == "missing_alias":
            judgment["alias_binding"] = None
        else:
            judgment["basis"] = "insufficient"
        return result
    engine = runner(case, invoke, state=two(case))
    engine.run()
    decision = next(iter(engine.state["coreferences"].values()))
    assert decision["verdict"] == "unresolved"
    assert "coreference_" in decision["reason"] or "source_" in decision["reason"]


@pytest.mark.parametrize("verdict,basis", [("different", "distinct"),
                                           ("unresolved", "insufficient")])
def test_negative_and_unknown_do_not_merge_or_disappear(case, verdict, basis):
    engine = runner(case, lambda s, p, c: answer(p, verdict, basis), state=two(case))
    engine.run()
    decisions = list(engine.state["coreferences"].values())
    assert decisions[0]["verdict"] == verdict
    aliases = canonical_mentions(engine.state["entities"], decisions, case[1].classes)
    assert aliases["0"] != aliases["1"]


def test_type_and_confirmation_boundaries_do_not_depend_on_names(case):
    state = deepcopy(case[2])
    add_work(state, case[1], [("0", "2"), ("0", "3")])
    state["entities"]["1"]["class_iri"] = "urn:test:Specific"
    state["entities"]["2"]["class_iri"] = "urn:test:Batch"
    state["entities"]["3"]["state"] = "unresolved"
    pairs = list(candidate_pairs(state, case[1].classes))
    assert [(a["id"], b["id"]) for _, a, b in pairs] == [("0", "1")]


def test_compatible_types_and_same_name_without_clue_work_do_not_generate_pairs(case):
    state = deepcopy(case[2])
    state["work"] = {}
    state["entities"]["1"]["label"] = state["entities"]["0"]["label"]
    assert list(candidate_pairs(state, case[1].classes)) == []
    calls = []
    engine = runner(case, lambda *args: calls.append(args), state=state)
    engine.run()
    assert not calls and not engine.state.get("coreferences")


def test_ineligible_work_waits_and_does_not_block_reading_completion(case):
    state = two(case)
    state["entities"]["1"]["state"] = "unresolved"
    calls = []
    engine = runner(case, lambda *args: calls.append(args), state=state)
    engine.run()
    assert not calls
    assert next(iter(engine.state["work"].values()))["status"] == "waiting"
    assert engine.state["cursor"]["main"]["stage"] == "complete"


def decision(left, right, verdict="same"):
    return {"id": pair_id(left, right), "left_mention_id": left, "right_mention_id": right,
            "verdict": verdict, "basis": "explicit_alias", "reason": "依据", "evidence": [],
            "proof": []}


@pytest.mark.parametrize("closing", [None, "different", "unresolved"])
def test_contradictory_or_unproved_triangle_blocks_whole_component(case, closing):
    decisions = [decision("0", "1"), decision("1", "2")]
    if closing:
        decisions.append(decision("0", "2", closing))
    aliases = canonical_mentions(case[2]["entities"], decisions, case[1].classes)
    assert len({aliases[key] for key in ("0", "1", "2")}) == 3


def test_complete_same_group_projection_keeps_mentions_conditions_and_original_owners(case):
    entities = [{key: row[key] for key in ("id", "label", "role", "class_iri", "class_label",
                                           "state", "reason", "evidence")}
                for row in case[2]["entities"].values()]
    result = {"entities": entities,
              "properties": [{"id": "p", "subject_id": "1", "value": "5", "evidence": []}],
              "relations": [{"id": "r", "subject_id": "2", "object_id": "3",
                             "polarity": "negative", "conditions": ["仅周一"], "evidence": []}],
              "observations": [{"candidate_subject_ids": ["0", "1"], "object_id": "2",
                                "alignments": [{"subject_id": "1", "card": "original"}]}]}
    decisions = [decision(a, b) for a, b in combinations(("0", "1", "2"), 2)]
    original = deepcopy(case[2])
    project_coreferences(result, decisions, case[1].classes)
    assert len(result["entities"]) == 3  # merged object, another object, document
    group = next(row for row in result["entities"] if len(row["mentions"]) == 3)
    assert {row["id"] for row in group["mentions"]} == {"0", "1", "2"}
    assert result["properties"][0]["subject_id"] == group["id"]
    assert result["properties"][0]["subject_mention_id"] == "1"
    relation = result["relations"][0]
    assert relation["subject_mention_id"] == "2" and relation["object_mention_id"] == "3"
    assert relation["polarity"] == "negative" and relation["conditions"] == ["仅周一"]
    assert all(row["applied"] for row in result["coreferences"])
    assert case[2] == original


def test_budget_split_pause_resume_saves_pairs_and_never_truncates_sources(case, monkeypatch):
    from app.services.document_harness import coreference
    calls = []
    paused = False
    def size(stage, payload, schema):
        return 100 * len(payload["pairs"])
    monkeypatch.setattr(coreference, "request_size", size)
    # Controller's real budget remains large; only scheduling is forced to split.
    monkeypatch.setattr("app.services.document_harness.controller.request_size", size)
    def invoke(stage, payload, schema):
        nonlocal paused
        calls.extend(tuple(sorted((pair["left_id"], pair["right_id"])))
                     for pair in payload["pairs"])
        paused = len(calls) == 1
        return answer(payload, "unresolved", "insufficient")
    state = deepcopy(case[2])
    add_work(state, case[1], [("0", "2"), ("1", "2")])
    engine = runner(case, invoke, state=state, budget=200, stop=lambda: paused)
    engine.run()
    assert len(engine.state["coreferences"]) == 1
    assert engine.state["cursor"]["main"]["stage"] == "coreference_review"
    continued = runner(case, invoke, state=engine.state, budget=200)
    continued.run()
    assert len(continued.state["coreferences"]) == 3
    assert len(calls) == 3
    assert all(row["status"] == "done" for row in continued.state["work"].values())
    # Unbudgetable single pairs fail explicitly, never become a successful empty result.
    too_small = runner(case, invoke, state=two(case), budget=99)
    with pytest.raises(ValueError, match="HARNESS_EVIDENCE_CONTEXT_TOO_LARGE"):
        too_small.run()


def test_pair_schema_rejects_foreign_ids_and_controller_rejects_missing_answers(case):
    schema = stage_schema("coreference_review", source_ids=["S1"], candidate_ids=["P1"])
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate({"judgments": {}}, schema)
    with pytest.raises(ValueError, match="pair_set_mismatch"):
        runner(case, lambda *args: {"judgments": {}}, state=two(case)).run()


def test_alias_decoding_cannot_quote_both_names_from_the_first_endpoint():
    schema = stage_schema("coreference_review", source_ids=["S1", "S2"],
                          candidate_ids=["P1"], pair_sources={"P1": ("S1", "S2")})
    quote = {"source_id": "S1", "text": "甲", "occurrence": None}
    value = {"judgments": {"P1": {
        "verdict": "same", "basis": "explicit_alias", "confidence": 0.99,
        "evidence": ["S1", "S2"], "proof": [quote], "reason": "原文声明简称",
        "alias_binding": {"left": quote, "right": {**quote, "text": "乙"},
                          "declaration": {**quote, "text": "甲以下简称乙"}},
    }}}
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(value, schema)
    value["judgments"]["P1"]["alias_binding"]["right"]["source_id"] = "S2"
    jsonschema.validate(value, schema)
