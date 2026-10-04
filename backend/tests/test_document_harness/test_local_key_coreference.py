"""Complete source-bound ontology keys recall identity checks without external mappings."""

from copy import deepcopy

import pytest
from docx import Document
from rdflib import Graph
from test_work_evidence import create_engine

from app.services.document_harness.coreference import canonical_mentions, review_coreferences
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.source import reference
from app.services.document_harness.work_execution import plan_coreference_work
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def local_keys(tmp_path):
    doc = Document()
    texts = ["车间甲：设备编号 EQ-17，装置一。", "车间甲：设备编号 EQ-17，装置一用于清洗。",
             "车间乙：设备编号 EQ-17，另一装置。"]
    for i, text in enumerate(texts):
        doc.add_heading(f"第{i + 1}节", 1)
        doc.add_paragraph(text)
    path = tmp_path / "keys.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    catalog = catalog_from_graph(Graph().parse(data="""
        @prefix : <urn:keys:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class .
        :Equipment a owl:Class; owl:hasKey (:code :workshop) .
        :describes a owl:ObjectProperty; rdfs:domain :Report; rdfs:range :Equipment .
        :code a owl:DatatypeProperty; rdfs:domain :Equipment; rdfs:range xsd:string .
        :workshop a owl:DatatypeProperty; rdfs:domain :Equipment; rdfs:range xsd:string .
    """, format="turtle"), "urn:keys:Report")
    entities = {}
    for i, text in enumerate(texts):
        unit = next(u for u in ir.evidence_units if u.text == text)
        ref = reference(ir, unit.evidence_id, 0, len(text))
        identifiers = []
        for prop, value in [("code", "EQ-17"), ("workshop", "车间乙" if i == 2 else "车间甲")]:
            start = text.index(value)
            identifiers.append({
                "property_iri": "urn:keys:" + prop, "value": value,
                "quote": reference(ir, unit.evidence_id, start, start + len(value)),
            })
        entities[str(i)] = {"id": str(i), "label": "装置", "role": "原文设备", "state": "accepted",
                           "class_iri": "urn:keys:Equipment", "class_label": "设备",
                           "referent": ref, "evidence": [ref], "field_ids": [], "window_id": None,
                           "identity_binding": {"group_id": f"g{i}", "identifiers": identifiers}}
    return ir, catalog, {"entities": entities, "fields": {}, "work": {}}


def test_complete_composite_key_recalls_cross_section_pair_then_requires_source_review(local_keys):
    calls = []

    def invoke(stage, payload, schema):
        assert stage == "coreference_review"
        calls.append(payload)
        return {"judgments": {p["pair_id"]: {
            "verdict": "same", "basis": "scoped_identifier", "confidence": 0.95,
            "evidence": [s["source_id"] for s in payload["sources"]],
            "proof": [{"source_id": s["source_id"], "text": s["text"], "occurrence": None}
                      for s in payload["sources"] if "EQ-17" in s["text"]],
            "alias_binding": None, "reason": "两处原文明确给出相同设备编号和车间作用域",
        } for p in payload["pairs"]}}

    engine = create_engine(local_keys, invoke)
    plan_coreference_work(engine)
    assert len(engine.state["work"]) == 1
    row = next(iter(engine.state["work"].values()))
    assert {row["input"]["left_mention_id"], row["input"]["right_mention_id"]} == {"0", "1"}
    assert not engine.state.get("coreferences")  # Matching keys only plan work.
    review_coreferences(engine)
    assert len(calls) == 1
    aliases = canonical_mentions(engine.state["entities"], engine.state["coreferences"].values(),
                                 engine.catalog.classes)
    assert aliases["0"] == aliases["1"] != aliases["2"]


@pytest.mark.parametrize("defect", ["missing_scope", "different_scope", "different_code",
                                  "unbound_value", "ambiguous_component", "unresolved_property",
                                  "identity_flag_only"])
def test_incomplete_or_incompatible_keys_never_recall_a_pair(local_keys, defect):
    ir, catalog, state = deepcopy(local_keys)
    state["entities"].pop("2")
    identifiers = state["entities"]["1"]["identity_binding"]["identifiers"]
    if defect == "missing_scope":
        identifiers.pop()
    elif defect in {"different_scope", "different_code"}:
        item = identifiers[-1 if defect == "different_scope" else 0]
        item["value"] = item["quote"]["text"] = "不同值"
    elif defect == "unbound_value":
        identifiers[0]["value"] = "not-the-source-value"
    elif defect == "ambiguous_component":
        identifiers.append(deepcopy(identifiers[0]))
    elif defect == "unresolved_property":
        card = catalog.classes["urn:keys:Equipment"]
        catalog.classes[card.iri] = card.model_copy(update={"properties": tuple(
            p.model_copy(update={"constraint_status": "unresolved"}) for p in card.properties)})
    else:
        card = catalog.classes["urn:keys:Equipment"]
        catalog.classes[card.iri] = card.model_copy(update={"identity_key_groups": ()})
    engine = create_engine((ir, catalog, state), lambda *_: pytest.fail("unexpected model call"))
    plan_coreference_work(engine)
    assert not engine.state["work"]
    assert not engine.state.get("coreferences")
