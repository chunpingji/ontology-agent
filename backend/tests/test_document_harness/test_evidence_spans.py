"""Program-owned identifier spans, bounded missing-span registration and exact grounding."""

from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from test_grouped_alignment import engine_fixture, proposal
from test_lookup import lookup_fixture as lookup_fixture
from test_lookup import quote

from app.services.document_harness.controller import Engine, Paused
from app.services.document_harness.model import request_size
from app.services.document_harness.referents import run_referent_alignment
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def additions(*texts):
    return {"new_spans": [quote(t) for t in texts], "expressions": [], "partitions": []}


def resume(engine):
    return Engine(ir=engine.ir, catalog=engine.catalog, state=engine.state, invoke=engine.invoke,
                  save=lambda _: None, should_stop=lambda: False, lookup=engine.lookup,
                  max_request_bytes=engine.max_request_bytes)


def with_missing_spans(engine, final=None):
    engine.lookup = None
    original = engine.invoke
    calls = []

    def respond(stage, payload, schema):
        calls.append((stage, deepcopy(payload), deepcopy(schema)))
        if stage == "referent_candidates":
            if payload["new_spans_allowed"]:
                return additions("A1/A2", "A1", "A2")
            if final is not None:
                return deepcopy(final)
        return original(stage, payload, schema)

    engine.invoke = respond
    return calls


def test_existing_spans_need_no_model_quotes_and_values_come_from_source(lookup_fixture):
    engine, window, calls = engine_fixture(lookup_fixture)
    original = engine.invoke

    def respond(stage, payload, schema):
        result = original(stage, payload, schema)
        if stage == "referent_candidates":
            result = proposal({s["text"]: s["span_id"] for s in payload["evidence_spans"]})
            result["expressions"] = result["expressions"][1:]
            result["partitions"] = result["partitions"][1:]
            jsonschema.validate(result, schema)
            assert not result["new_spans"]
            assert all(set(e) == {"id", "property_iri", "span_id"} for e in result["expressions"])
            for change in ("unknown_id", "fabricated_text"):
                invalid = deepcopy(result)
                if change == "unknown_id":
                    invalid["expressions"][0]["span_id"] = "foreign-span"
                else:
                    invalid["expressions"][0]["quote"] = quote("A1 and A2")
                with pytest.raises(jsonschema.ValidationError):
                    jsonschema.validate(invalid, schema)
        else:
            assert set(result) == {"verdict", "selected_partition_id", "confidence",
                                   "evidence_span_ids", "reason"}
        return result

    engine.invoke = respond
    run_referent_alignment(engine, window)
    assert [c[0] for c in calls] == ["referent_candidates", "referent_selection"]
    for entity in engine.entities(window):
        binding = entity["identity_binding"]["identifiers"][0]
        assert binding["value"] == binding["quote"]["text"] == entity["label"]
        assert binding["identity_status"] == "not_checked"
        assert binding["span_id"] in next(iter(engine.state["referent_work"].values()))[
            "span_catalog"
        ]


def test_same_text_in_different_positions_and_documents_has_different_ids(lookup_fixture, tmp_path):
    doc = Document()
    doc.add_paragraph("Use A1/A2 objects; A1 appears again.")
    path = tmp_path / "repeated.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    catalogs = []
    for current in (ir, ir.model_copy(update={"document_hash": "f" * 64})):
        engine, window, calls = engine_fixture((current, *lookup_fixture[1:]))
        # This collective draft includes both occurrences in its physical scope.
        engine.state["entities"]["old"]["referent"] = window.resolve(
            current, quote("Use A1/A2 objects; A1 appears again."),
        )
        run_referent_alignment(engine, window)
        spans = [s for s in calls[0][1]["evidence_spans"] if s["text"] == "A1"]
        assert len(spans) == 2
        assert len({s["span_id"] for s in spans}) == 2
        assert len({s["start"] for s in spans}) == 2
        catalogs.append({s["span_id"] for s in spans})
    assert catalogs[0].isdisjoint(catalogs[1])


def test_missing_mock_keys_are_registered_once_then_selected_by_program_ids(lookup_fixture):
    engine, window, _ = engine_fixture(lookup_fixture)
    calls = with_missing_spans(engine)
    run_referent_alignment(engine, window)
    assert [c[0] for c in calls] == [
        "referent_candidates", "referent_candidates", "referent_selection",
    ]
    assert all(s["text"] not in {"A1", "A2"} for s in calls[0][1]["evidence_spans"])
    assert {s["text"] for s in calls[1][1]["evidence_spans"]} >= {"A1", "A2"}
    assert not calls[1][1]["new_spans_allowed"]
    assert calls[1][2]["properties"]["new_spans"]["maxItems"] == 0
    assert {e["label"] for e in engine.entities(window)} == {"A1", "A2"}
    assert not lookup_fixture[3]
    assert all(not i["source_candidates"] for e in engine.entities(window)
               for i in e["identity_binding"]["identifiers"])


def test_original_field_values_remain_candidates_without_mock(lookup_fixture):
    engine, window, calls = engine_fixture(lookup_fixture)
    engine.lookup = None
    for value in ("A1", "A2"):
        ref = window.resolve(engine.ir, quote(value))
        engine.state["fields"][value] = {
            "id": value, "alias": value, "label": "", "value": value, "missing": False,
            "evidence": [ref], "value_evidence": [ref], "source_aliases": ["S1"], "row": None,
        }
    engine.state["entities"]["old"]["field_ids"] = ["A1", "A2"]
    original = engine.invoke

    def respond(stage, payload, schema):
        result = original(stage, payload, schema)
        if stage == "referent_candidates":
            eligible = {s["text"] for s in payload["evidence_spans"] if s["identifier_candidate"]}
            assert eligible == {"A1", "A2"}
            result = proposal({s["text"]: s["span_id"] for s in payload["evidence_spans"]})
            result["expressions"] = result["expressions"][1:]
            result["partitions"] = result["partitions"][1:]
            jsonschema.validate(result, schema)
        return result

    engine.invoke = respond
    run_referent_alignment(engine, window)
    assert len(calls) == 2
    assert {e["label"] for e in engine.entities(window)} == {"A1", "A2"}
    assert not next(iter(engine.state["referent_work"].values())).get("spans_extended")


@pytest.mark.parametrize("kind,error", [
    ("fabricated", "source_quote_mismatch"),
    ("foreign", "source_outside_reading_window"),
    ("duplicate", "identifier_span_already_available"),
    ("mixed", "identifier_span_request_with_partitions"),
])
def test_invalid_new_span_cannot_register_or_retry_on_resume(lookup_fixture, kind, error):
    engine, window, _ = engine_fixture(lookup_fixture)
    engine.lookup = None
    calls = []

    def respond(stage, payload, schema):
        calls.append(stage)
        result = additions("A1 and A2")
        if kind == "foreign":
            result = additions("A1")
            result["new_spans"][0]["source_id"] = "outside"
        elif kind == "duplicate":
            result = additions("A1", "A1")
        elif kind == "mixed":
            result = {**proposal(), "new_spans": [quote("A1")]}
        return result

    engine.invoke = respond
    with pytest.raises(ValueError, match=error):
        run_referent_alignment(engine, window)
    restored = resume(engine)
    with pytest.raises(ValueError, match=error):
        run_referent_alignment(restored, window)
    assert len(calls) == 1
    assert list(restored.state["entities"]) == ["document", "old"]
    assert not restored.state["fields"]
    assert not next(iter(restored.state["referent_work"].values())).get("spans_extended")


def test_other_source_identifier_gets_one_scoped_correction_without_binding(
    lookup_fixture, tmp_path,
):
    doc = Document()
    doc.add_paragraph("Use A1/A2 objects for the trial.")
    doc.add_paragraph("A1 appears again in a different source.")
    path = tmp_path / "other-source.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    engine, window, _ = engine_fixture((ir, *lookup_fixture[1:]))
    calls = []

    def respond(stage, payload, schema):
        assert stage == "referent_candidates"
        calls.append(deepcopy(payload))
        assert payload["group_source_ids"] == ["S1"]
        assert schema["$defs"]["Quote"]["properties"]["source_id"]["enum"] == ["S1"]
        if len(calls) == 1:
            # Simulate a completed answer saved before the source enum was narrowed.
            return {"new_spans": [{"source_id": "S2", "text": "A1",
                                   "occurrence": None}],
                    "expressions": [], "partitions": []}
        assert payload["proposal_feedback"] == {
            "issue": "identifier_span_outside_group", "allowed_source_ids": ["S1"],
        }
        return {"new_spans": [], "expressions": [], "partitions": []}

    engine.invoke = respond
    run_referent_alignment(engine, window)

    work = next(iter(engine.state["referent_work"].values()))
    assert len(calls) == 2
    assert work["proposal_corrected"] and work["proposal"]["expressions"] == []
    assert not work.get("spans_extended")
    assert not any(e.get("identity_binding") for e in engine.entities(window))
    run_referent_alignment(resume(engine), window)
    assert len(calls) == 2


@pytest.mark.parametrize("kind,error", [
    ("repeat", "identifier_span_extension_exhausted"),
    ("foreign", "unknown_evidence_span"),
    ("coverage", "partition_omits_group_mention"),
])
def test_extended_catalog_still_requires_valid_complete_partitions(lookup_fixture, kind, error):
    engine, window, _ = engine_fixture(lookup_fixture)
    calls = with_missing_spans(engine)
    original = engine.invoke

    def respond(stage, payload, schema):
        answer = original(stage, payload, schema)
        if stage == "referent_candidates" and not payload["new_spans_allowed"]:
            if kind == "repeat":
                return additions("A1")
            if kind == "foreign":
                answer["expressions"][0]["span_id"] = "another-window-span"
            else:
                answer["partitions"][1]["members"].pop()
        return answer

    engine.invoke = respond
    with pytest.raises(ValueError, match=error):
        run_referent_alignment(engine, window)
    with pytest.raises(ValueError, match=error):
        run_referent_alignment(resume(engine), window)
    assert len(calls) == (3 if kind == "coverage" else 2)
    assert "old" in engine.state["entities"]


@pytest.mark.parametrize("pause_at", ["extended", "answered"])
def test_pause_reuses_registered_spans_and_final_answer(lookup_fixture, pause_at):
    engine, window, _ = engine_fixture(lookup_fixture)
    calls = with_missing_spans(engine)
    stopped = False

    def save(changes):
        nonlocal stopped
        for row in changes.get("referent_work", {}).values():
            if row.get("spans_extended") and (
                pause_at == "extended" or "proposal" in row
            ):
                stopped = True

    engine.save = save
    engine.should_stop = lambda: stopped
    with pytest.raises(Paused):
        run_referent_alignment(engine, window)
    restored = resume(engine)
    run_referent_alignment(restored, window)
    run_referent_alignment(restored, window)
    assert [c[0] for c in calls] == [
        "referent_candidates", "referent_candidates", "referent_selection",
    ]
    assert {e["label"] for e in restored.entities(window)} == {"A1", "A2"}


def test_expanded_catalog_cannot_bypass_input_budget(lookup_fixture):
    engine, window, _ = engine_fixture(lookup_fixture)
    original = engine.invoke
    calls = []
    engine.lookup = None

    def respond(stage, payload, schema):
        calls.append(stage)
        engine.max_request_bytes = request_size(stage, payload, schema)
        return additions("A1/A2", "A1", "A2")

    engine.invoke = respond
    with pytest.raises(ValueError, match="HARNESS_EVIDENCE_CONTEXT_TOO_LARGE"):
        run_referent_alignment(engine, window)
    assert calls == ["referent_candidates"]
    assert next(iter(engine.state["referent_work"].values()))["spans_extended"]
    engine.max_request_bytes = 32768
    engine.invoke = original
    run_referent_alignment(engine, window)
    assert {e["label"] for e in engine.entities(window)} == {"A1", "A2"}


def test_selection_cannot_cite_unprovided_span(lookup_fixture):
    engine, window, _ = engine_fixture(lookup_fixture)
    original = engine.invoke

    def respond(stage, payload, schema):
        result = original(stage, payload, schema)
        if stage == "referent_selection":
            result["evidence_span_ids"] = ["not-in-this-catalog"]
        return result

    engine.invoke = respond
    with pytest.raises(ValueError, match="unknown_evidence_span"):
        run_referent_alignment(engine, window)
    assert "old" in engine.state["entities"]


def test_context_span_is_not_automatically_an_identifier(lookup_fixture):
    engine, window, _ = engine_fixture(lookup_fixture)

    def respond(stage, payload, schema):
        answer = proposal({s["text"]: s["span_id"] for s in payload["evidence_spans"]})
        context = next(s for s in payload["evidence_spans"] if s["text"] == "A1/A2")
        assert not context["identifier_candidate"]
        with pytest.raises(jsonschema.ValidationError):
            jsonschema.validate(answer, schema)
        return answer

    engine.invoke = respond
    with pytest.raises(ValueError, match="span_not_identifier_candidate"):
        run_referent_alignment(engine, window)
    assert "old" in engine.state["entities"]
