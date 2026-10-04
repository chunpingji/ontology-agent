"""Public reading behavior: dense rows, exact continuation and shared prompt definitions."""

import json
from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from test_controller import inputs as inputs  # noqa: F401

from app.services.document_harness.calls import MemoryCalls
from app.services.document_harness.continuation import restore_window_plan
from app.services.document_harness.controller import Engine
from app.services.document_harness.model import gateway_schema, prompt_schema
from app.services.document_harness.protocols import Discovery, stage_schema
from app.services.document_harness.source import build_windows
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def ir_for(tmp_path, texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    path = tmp_path / "reading.docx"
    doc.save(path)
    return build_document_ir(path, parse_docx_structure(path))


def answer(payload, *, through=None):
    primary = [r["source_id"] for r in payload["reading_scope"]]
    ids = primary if through is None else primary[:primary.index(through) + 1]
    return {
        "entities": [{"candidate_class_iri": "urn:test:Thing",
                      "local_id": sid, "name": None, "anchor": {"source_id": sid},
                      "role": "object", "evidence": [sid], "field_ids": [], "source_fields": []}
                     for sid in ids],
        "document_field_ids": [], "document_source_fields": [], "unowned_fields": [],
        "relation_hints": [], "reference_cues": [], "complete": through is None,
        "read_through_source_id": through,
    }


def reader(ir, catalog, invoke, *, state=None, calls=None, stop=lambda: False):
    engine = Engine(
        ir=ir, catalog=catalog, state=state or {}, invoke=invoke, calls=calls,
        save=lambda _: None,
        should_stop=lambda: stop() or engine.state.get("cursor", {}).get("main", {}).get(
            "phase",
        ) not in (None, "discovery"),
        rank=lambda *_: {"snapshot_id": catalog.snapshot_id, "selected_iris": ["urn:test:Thing"]},
        policy={"execution_policy": {"reading_concurrency": 1}},
    )
    return engine


@pytest.mark.parametrize("count", [12, 24, 48])
def test_complete_dense_mentions_need_one_reading_call(tmp_path, inputs, count):
    ir = ir_for(tmp_path, [f"Object {index}" for index in range(count)])
    paid = []

    def invoke(stage, payload, schema):
        paid.append(payload)
        result = answer(payload)
        jsonschema.validate(result, schema)
        return result

    engine = reader(ir, inputs[1], invoke)
    engine.run()
    assert len(paid) == 1
    assert len(engine.state["entities"]) == count + 1
    assert engine.state["cursor"]["main"]["reading"]["complete"]
    assert all(not w["children"] for w in engine.state["windows"].values())


def test_partial_prefix_is_not_registered_again_and_resume_preserves_coverage(tmp_path, inputs):
    ir = ir_for(tmp_path, [f"Object {index}" for index in range(6)])
    paid = []

    def invoke(stage, payload, schema):
        paid.append(deepcopy(payload))
        through = payload["reading_scope"][2]["source_id"] if len(paid) == 1 else None
        return answer(payload, through=through)

    calls = MemoryCalls(invoke)
    engine = reader(ir, inputs[1], invoke, calls=calls)
    engine.should_stop = lambda: any(w.get("completed_ranges")
                                   for w in engine.state.get("windows", {}).values())
    engine.run()
    progress = engine.state["cursor"]["main"]["reading"]
    assert progress["complete_characters"] == sum(len(u.text) for u in ir.evidence_units[:3])
    assert not progress["complete"] and len(paid) == 1
    resumed = reader(ir, inputs[1], invoke, state=engine.state, calls=calls)
    resumed.run()
    assert len(paid) == 2
    primary_texts = [[s["text"] for s in p["sources"]
                      if s["source_id"] in {r["source_id"] for r in p["reading_scope"]}]
                     for p in paid]
    assert primary_texts[1] == ["Object 3", "Object 4", "Object 5"]
    assert len(resumed.state["entities"]) == 7
    assert {e["referent"]["text"] for k, e in resumed.state["entities"].items()
            if k != "document"} == {f"Object {i}" for i in range(6)}
    assert resumed.state["cursor"]["main"]["reading"]["complete_characters"] == sum(
        len(u.text) for u in ir.evidence_units
    )
    for row in resumed.state["windows"].values():
        assert restore_window_plan(ir, row["plan"]).payload()


@pytest.mark.parametrize("marker", ["foreign", "S6"])
def test_invalid_prefix_cannot_credit_unread_text(tmp_path, inputs, marker):
    from app.services.document_harness.reading import decode_local_discovery

    ir = ir_for(tmp_path, [f"Object {i}" for i in range(6)])
    window = build_windows(ir)[0]
    result = answer(window.payload())
    result.update(complete=False, read_through_source_id=marker)
    decoded = decode_local_discovery(ir, window, Discovery.model_validate(result), {},
                                     {"id": "document", "field_ids": []},
                                     class_iris=["urn:test:Thing"])
    assert not decoded["complete"] and decoded["completed_ranges"] == []
    assert any(o["kind"] == "scope" for o in decoded["changes"]["observations"].values())


def test_incomplete_without_cursor_still_splits_instead_of_claiming_coverage(tmp_path, inputs):
    ir = ir_for(tmp_path, [f"Object {i}" for i in range(6)])
    paid = []

    def invoke(stage, payload, schema):
        paid.append(payload)
        result = answer(payload)
        if len(paid) == 1:
            result["complete"] = False
        return result

    engine = reader(ir, inputs[1], invoke)
    engine.run()
    assert len(paid) == 3
    assert engine.state["cursor"]["main"]["reading"]["complete"]
    assert len(engine.state["entities"]) == 7


def test_overlapping_paragraph_pieces_do_not_advance_past_a_cross_boundary_mention(tmp_path):
    from app.services.document_harness.continuation import reading_prefix

    ir = ir_for(tmp_path, ["长段落" * 1800])
    window = build_windows(ir, max_chars=4800)[0]
    assert len(window.sources) > 1
    with pytest.raises(ValueError, match="reading_prefix_splits_physical_source"):
        reading_prefix(window, window.sources[0]["source_id"])


def test_table_packing_keeps_rows_together_and_exact_source_coordinates(tmp_path):
    doc = Document()
    table = doc.add_table(rows=61, cols=4)
    for col, text in enumerate(["对象", "编号", "数量", "单位"]):
        table.cell(0, col).text = text
    for row in range(1, 61):
        for col, text in enumerate([f"A{row}", str(row), "1", "g"]):
            table.cell(row, col).text = text
    path = tmp_path / "table.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    windows = build_windows(ir)
    assert len(windows) == 2  # 244 short cells no longer make seven tiny windows.
    for packed in (windows, build_windows(ir, max_sources=11)):
        rows = {}
        for window in packed:
            for source in window.sources:
                assert source["text"] == ir.unit(source["evidence_id"]).text
                if source["row"] and source["evidence_id"] in window.primary_ids:
                    rows.setdefault(source["row"], set()).add(window.id)
        assert len(rows) == 60 and all(len(ids) == 1 for ids in rows.values())


def test_table_remainder_retains_headers_and_row_subject_only_as_context(tmp_path):
    from app.services.document_harness.continuation import (
        reading_prefix,
        remaining_reading_window,
        serialize_window_plan,
    )

    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for row, values in enumerate([("对象", "编号", "数量"),
                                  ("Alpha", "A1", "7"), ("Beta", "B1", "8")]):
        for col, text in enumerate(values):
            table.cell(row, col).text = text
    path = tmp_path / "remainder-table.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    original = build_windows(ir)[0]
    subject = next(s for s in original.sources if s["text"] == "Beta")
    remainder = remaining_reading_window(ir, original, reading_prefix(
        original, subject["source_id"],
    ))
    assert {"对象", "编号", "数量", "Beta", "B1", "8"} <= {
        s["text"] for s in remainder.sources
    }
    assert {ir.unit(r["evidence_id"]).text for r in remainder.primary()} == {"B1", "8"}
    restored = restore_window_plan(ir, serialize_window_plan(remainder))
    assert restored.payload() == remainder.payload()
    assert {f["id"] for f in remainder.fields} <= {f["id"] for f in original.fields}


def test_invalid_original_quote_withholds_only_its_source_from_reported_prefix(tmp_path):
    from app.services.document_harness.reading import decode_local_discovery

    ir = ir_for(tmp_path, ["Object A", "Object B", "Object C"])
    window = build_windows(ir)[0]
    result = answer(window.payload(), through="S2")
    result["entities"][0]["source_fields"] = [
        {"label": None, "value": {"source_id": "S1", "text": "invented"}},
    ]
    decoded = decode_local_discovery(ir, window, Discovery.model_validate(result), {},
                                     {"id": "document", "field_ids": []},
                                     class_iris=["urn:test:Thing"])
    assert not decoded["complete"]
    assert decoded["completed_ranges"] == [{
        "evidence_id": ir.evidence_units[1].evidence_id, "start": 0, "end": len("Object B"),
    }]
    assert not any(f["value"] == "invented"
                   for f in decoded["changes"].get("fields", {}).values())
    assert any(o["reason"] == "source_quote_mismatch"
               for o in decoded["changes"]["observations"].values())


def test_shared_prompt_schema_reduces_bytes_without_changing_constraints():
    schema = stage_schema("discover", source_ids=[f"S{i}" for i in range(100)],
                          field_ids=[f"F{i}" for i in range(100)])
    readable = prompt_schema("discover", schema)
    wire = gateway_schema(schema)
    assert len(json.dumps(readable)) < len(json.dumps(wire)) * 0.8
    assert gateway_schema(readable) == wire
    value = {"entities": [], "document_field_ids": [], "document_source_fields": [],
             "unowned_fields": [], "relation_hints": [], "reference_cues": [],
             "complete": True, "read_through_source_id": None}
    for contract in (readable, wire):
        jsonschema.validate(value, contract)
        with pytest.raises(jsonschema.ValidationError):
            jsonschema.validate({**value, "read_through_source_id": "foreign"}, contract)
