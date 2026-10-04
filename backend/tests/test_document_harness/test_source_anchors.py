"""Entity anchors select physical sources without asking models to repeat whole cells."""

from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from pydantic import ValidationError
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.protocols import Discovery, SourceAnchor, stage_schema
from app.services.document_harness.source import build_windows, make_window
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def discovery(anchor):
    return {
        "entities": [{"candidate_class_iri": "urn:anchor:Object",
                      "local_id": "E1", "role": "object", "anchor": anchor,
                      "evidence": ["S1"], "field_ids": [], "source_fields": []}],
        "document_field_ids": [], "document_source_fields": [], "unowned_fields": [],
        "relation_hints": [], "complete": True,
    }


@pytest.mark.parametrize("invalid", ["omitted", None, {}, {"source_id": ""}])
def test_missing_or_null_anchor_cannot_pass_discovery_contract(invalid):
    answer = discovery(invalid)
    if invalid == "omitted":
        del answer["entities"][0]["anchor"]
    schema = stage_schema("discover", source_ids=["S1"], class_iris=["urn:anchor:Object"])
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(answer, schema)
    with pytest.raises(ValidationError):
        Discovery.model_validate(answer)


def test_source_selection_is_sufficient_but_foreign_ids_are_not_in_schema():
    schema = stage_schema("discover", source_ids=["S1"], class_iris=["urn:anchor:Object"])
    answer = discovery({"source_id": "S1"})
    jsonschema.validate(answer, schema)
    assert Discovery.model_validate(answer).entities[0].anchor.text is None
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(discovery({"source_id": "S2"}), schema)
    with pytest.raises(ValidationError, match="anchor_occurrence_requires_text"):
        SourceAnchor(source_id="S1", occurrence=1)


@pytest.fixture
def table_input(tmp_path):
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for row, values in enumerate([
        ["名称", "编号", "用途"], ["Alpha", "A01", "存放"], ["Alpha", "A02", "混合"],
    ]):
        for col, text in enumerate(values):
            table.cell(row, col).text = text
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(1, 1).add_paragraph("对象甲与对象乙；对象甲")
    path = tmp_path / "anchors.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    catalog = catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:anchor:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class . :Object a owl:Class .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :Object .
    ''', format="turtle"), "urn:anchor:Report")
    return ir, catalog, build_windows(ir)[0]


def test_cells_with_same_text_keep_distinct_physical_coordinates(table_input):
    ir, _, window = table_input
    sources = [s for s in window.sources if s["text"] == "Alpha"]
    refs = [window.resolve_anchor(ir, {"source_id": s["source_id"]}) for s in sources]
    units = [ir.unit(ref["source_id"]) for ref in refs]
    assert [(u.table_path, u.row_index, u.column_index) for u in units] == [
        (["table:0"], 1, 0), (["table:0"], 2, 0),
    ]
    assert units[0].source_cell_id != units[1].source_cell_id
    assert refs[0] != refs[1] and all(ref["text"] == "Alpha" for ref in refs)


def test_merged_cell_and_multiple_paragraphs_keep_parser_origin(table_input):
    ir, _, window = table_input
    sources = [s for s in window.sources if s["text"] in {"存放", "混合"}]
    units = [ir.unit(window.resolve_anchor(ir, {"source_id": s["source_id"]})["source_id"])
             for s in sources]
    assert len(units) == 2
    assert units[0].source_cell_id == units[1].source_cell_id
    assert units[0].evidence_id != units[1].evidence_id
    assert {(u.row_index, u.column_index) for u in units} == {(1, 2)}
    grid = ir.tables[0]["grid"]
    assert grid[1][2] == grid[2][2] == units[0].source_cell_id


def test_partial_cell_requires_exact_unambiguous_text_and_cannot_expand_window(table_input):
    ir, _, window = table_input
    source = next(s for s in window.sources if s["text"] == "对象甲与对象乙；对象甲")
    anchor = {"source_id": source["source_id"], "text": "对象甲"}
    with pytest.raises(ValueError, match="source_quote_ambiguous"):
        window.resolve_anchor(ir, anchor)
    selected = window.resolve_anchor(ir, {**anchor, "occurrence": 1})
    assert (selected["start"], selected["end"], selected["text"]) == (8, 11, "对象甲")
    with pytest.raises(ValueError, match="source_quote_mismatch"):
        window.resolve_anchor(ir, {**anchor, "text": "对象丙"})
    with pytest.raises(ValueError, match="source_outside_reading_window"):
        window.resolve_anchor(ir, {"source_id": "outside"})
    fragment = make_window(ir, [(source["evidence_id"], 4, 7)], [source["evidence_id"]])
    ref = fragment.resolve_anchor(ir, {"source_id": "S1"})
    assert (ref["text"], ref["start"], ref["end"]) == ("对象乙", 4, 7)


@pytest.mark.parametrize("bad_anchor", [False, True])
def test_registration_uses_source_anchor_without_name_and_accounts_for_failures(
    table_input, bad_anchor,
):
    ir, catalog, window = table_input
    calls = []

    def invoke(stage, payload, schema):
        assert stage == "discover"
        sources = [s for s in payload["sources"] if s["text"] == "Alpha"]
        answer = discovery({"source_id": sources[0]["source_id"]})
        first = answer["entities"][0]
        first["evidence"] = [sources[0]["source_id"]]
        first["field_ids"] = [f["field_id"] for f in payload["fields"] if f["row"] == 1]
        second = deepcopy(first)
        second.update(local_id="E2", anchor={"source_id": sources[1]["source_id"]},
                      evidence=[sources[1]["source_id"]], field_ids=[])
        answer["entities"].append(second)
        if bad_anchor:
            first["anchor"]["text"] = "not in the cell"
        calls.append(answer)
        return answer

    engine = Engine(ir=ir, catalog=catalog, state={}, invoke=invoke,
                    save=lambda changes: None,
                    should_stop=lambda: bool(engine.state.get("window_entities")),
                    rank=lambda *_: {"snapshot_id": catalog.snapshot_id,
                                     "selected_iris": ["urn:anchor:Object"]})
    engine.windows = [window]
    engine.run()
    entities = [e for e in engine.state["entities"].values() if e["id"] != "document"]
    assert len(entities) == (1 if bad_anchor else 2)
    assert all(e["label"] == "Alpha" and e["name"] is None for e in entities)
    assert all(e["state"] == "candidate" and e["class_iri"] is None for e in entities)
    assert (engine.state["windows"][window.id]["reading_state"] == "complete") is not bad_anchor
    if bad_anchor:
        failure = next(o for o in engine.state["observations"].values()
                       if o["reason"] == "source_quote_mismatch")
        assert failure["evidence"][0]["text"] == "Alpha"
        assert any(f["value"] == "A01" for f in engine.state["fields"].values())


def test_invalid_anchor_cannot_finish_an_unsplittable_window(table_input):
    ir, catalog, _ = table_input
    unit = next(u for u in ir.evidence_units if u.text == "Alpha")
    window = make_window(ir, [(unit.evidence_id, 0, len(unit.text))], [unit.evidence_id])

    def invoke(stage, *_):
        assert stage == "discover"
        return discovery({"source_id": "S1", "text": "invalid"})

    engine = Engine(ir=ir, catalog=catalog, state={}, invoke=invoke,
                    save=lambda changes: None, should_stop=lambda: False,
                    rank=lambda *_: {"snapshot_id": catalog.snapshot_id,
                                     "selected_iris": ["urn:anchor:Object"]})
    engine.windows = [window]
    engine.run()
    assert engine.state["cursor"]["main"]["stage"] == "complete"
    assert not engine.state["cursor"]["main"]["scope_complete"]
    assert len(engine.windows) == 1
