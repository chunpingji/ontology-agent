"""Bounded inputs retain exact source bindings and do not repeat historical fields."""

from copy import deepcopy

import jsonschema
import pytest
from docx import Document
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.source import identity, make_window, reference
from app.services.document_harness.work import make_work
from app.services.document_harness.work_execution import drain_work, plan_graph_work, review_for
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def records(tmp_path):
    doc = Document()
    for text in (
        "Name: Alpha",
        "Name: Beta",
        "Old: historical-value",
        "Current: current-value",
        "Alpha has weight 10 kg",
        "ReportOld: historic-report",
        "Unused: unrelated-report",
        "Report: current-report",
        "Current reading",
        "Object: B0",
        "Object: B1",
        "Object: B2",
        "Object: B3",
    ):
        doc.add_paragraph(text)
    path = tmp_path / "budgets.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    units = {unit.text: unit for unit in ir.evidence_units}
    whole = make_window(
        ir,
        [(u.evidence_id, 0, len(u.text)) for u in ir.evidence_units],
        [u.evidence_id for u in ir.evidence_units],
    )
    fields = {field["label"]: field for field in whole.fields}
    graph = Graph().parse(
        data="""
        @prefix : <urn:budget:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class . :A a owl:Class . :B a owl:Class .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ;
            rdfs:range [owl:unionOf (:A :B)] .
        :uses a owl:ObjectProperty ; rdfs:domain :A ; rdfs:range :B .
        :contains a owl:ObjectProperty ; rdfs:domain :A ; rdfs:range :B .
        :value a owl:DatatypeProperty ; rdfs:domain [owl:unionOf (:Report :A)] ;
            rdfs:range xsd:string .
    """,
        format="turtle",
    )
    return ir, catalog_from_graph(graph, "urn:budget:Report"), units, fields


def entity(ir, key, unit, *, class_iri=None, fields=()):
    ref = reference(ir, unit.evidence_id, 0, len(unit.text))
    return {
        "id": key,
        "label": unit.text,
        "role": "source-object",
        "class_iri": class_iri,
        "class_label": None,
        "state": "candidate",
        "reason": "source mention",
        "evidence": [ref],
        "referent": ref,
        "field_ids": list(fields),
        "window_id": None,
    }


def engine(
    records, selected, entities, invoke, *, root_fields=(), budget=100000, stage="type_alignment"
):
    ir, catalog, units, fields = records
    selected_units = [units[text] for text in selected]
    window = make_window(
        ir,
        [(u.evidence_id, 0, len(u.text)) for u in selected_units],
        [u.evidence_id for u in selected_units],
    )
    root = {
        "id": "document",
        "label": "Report",
        "role": "document_root",
        "class_iri": "urn:budget:Report",
        "class_label": "Report",
        "state": "accepted",
        "reason": "input type",
        "evidence": [],
        "field_ids": list(root_fields),
        "window_id": None,
    }
    for item in entities:
        item["window_id"] = window.id
    state = {
        "cursor": {
            "main": {
                "skeleton_window_id": window.id,
                "active_batches": {}, "phase": "skeleton", "semantic_step": None, "planned_steps": [],
                "stage": stage,
                "scope_complete": False,
            }
        },
        "entities": {"document": root, **{item["id"]: item for item in entities}},
        "fields": {field["id"]: field for field in fields.values()},
        "window_entities": {window.id: {"ids": [item["id"] for item in entities]}},
    }
    state["windows"] = {window.id: {**Engine.window_row(window, [0]), "skeleton_step": stage}}
    runner = Engine(
        ir=ir,
        catalog=catalog.model_dump(mode="json"),
        state=state,
        invoke=invoke,
        save=lambda changes: None,
        should_stop=lambda: False,
        max_request_bytes=budget,
    )
    runner.windows = [window]
    return runner, window


def test_discovery_cards_use_available_input_budget_before_optional_lookup(records, monkeypatch):
    _, catalog, _, _ = records
    budgets = []

    def rank(_catalog, _payload, budget):
        budgets.append(budget)
        return {"snapshot_id": catalog.snapshot_id, "selected_iris": []}

    def invoke(stage, _payload, _schema):
        assert stage == "discover"
        return {
            "entities": [],
            "document_field_ids": [],
            "document_source_fields": [],
            "unowned_fields": [],
            "relation_hints": [],
            "complete": True,
        }

    monkeypatch.setattr(
        "app.services.document_harness.controller.request_size",
        lambda _stage, _payload, _schema: 100,
    )
    runner, window = engine(records, ["Current reading"], [], invoke, budget=2000, stage="discover")
    runner.rank = rank
    runner.lookup = lambda _operation, _catalog, _argument: {"capabilities": [], "issues": []}

    runner.state["cursor"]["main"]["phase"] = "discovery"
    runner.read_windows()

    assert budgets == [1900]
    assert runner.state["cursor"]["main"]["phase"] == "skeleton"


def test_guidance_stays_together_on_overflow_and_unselected_types_remain_available(
    records,
    monkeypatch,
):
    ir, _, units, _ = records
    candidate = entity(ir, "a", units["Name: Alpha"])
    menus = []

    def size(stage, payload, schema):
        return 100 * (len(payload["classes"]) + len(payload["entities"]))

    def invoke(stage, payload, schema):
        assert stage == "type_alignment"
        menu = [card["iri"] for card in payload["classes"]]
        menus.append(menu)
        return {
            "entities": {
                item["entity_id"]: {
                    "class_iri": "urn:budget:Report" if "urn:budget:Report" in menu else None,
                    "confidence": 0.9,
                    "evidence": item["evidence"],
                    "reason": "None of the guided classes fits; consider the remaining type",
                }
                for item in payload["entities"]
            }
        }

    monkeypatch.setattr("app.services.document_harness.controller.request_size", size)
    runner, window = engine(records, ["Name: Alpha"], [candidate], invoke, budget=350)
    runner.state["windows"][window.id]["guidance_class_iris"] = ["urn:budget:B", "urn:budget:A"]
    runner.type_alignment(window)
    assert menus == [["urn:budget:B", "urn:budget:A"], ["urn:budget:Report"]]
    assert runner.state["entities"]["a"]["class_iri"] == "urn:budget:Report"
    assert runner.state["entities"]["a"]["state"] == "candidate"


def test_type_batches_split_entities_after_cards_and_bind_short_sources_before_rejoining(
    records,
    monkeypatch,
):
    ir, catalog, units, fields = records
    entities = [entity(ir, "a", units["Name: Alpha"]), entity(ir, "b", units["Name: Beta"])]
    calls = []

    def size(stage, payload, schema):
        return 100 * (len(payload["classes"]) + len(payload["entities"]))

    def invoke(stage, payload, schema):
        assert stage == "type_alignment"
        assert size(stage, payload, schema) <= 250
        calls.append(deepcopy(payload))
        classes = {item["iri"] for item in payload["classes"]}
        return {
            "entities": {
                item["entity_id"]: {
                    "class_iri": target
                    if (target := "urn:budget:A" if "Alpha" in item["label"] else "urn:budget:B")
                    in classes
                    else None,
                    "confidence": 0.9,
                    "evidence": item["evidence"],
                    "reason": "matching role",
                }
                for item in payload["entities"]
            }
        }

    monkeypatch.setattr("app.services.document_harness.controller.request_size", size)
    runner, window = engine(records, ["Current reading"], entities, invoke, budget=250)
    runner.type_alignment(window)
    for key, text, expected_type in (
        ("a", "Name: Alpha", "urn:budget:A"),
        ("b", "Name: Beta", "urn:budget:B"),
    ):
        saved = runner.state["entities"][key]
        assert saved["class_iri"] == expected_type
        assert {ref["source_id"] for ref in saved["type_evidence"]} == {units[text].evidence_id}
        assert {
            card["iri"]
            for call in calls
            if any(e["label"] == text for e in call["entities"])
            for card in call["classes"]
        } == set(catalog.reachable_class_iris)
    assert all(len(call["entities"]) == 1 for call in calls)


def test_single_entity_that_still_exceeds_budget_fails_without_truncating_or_calling(
    records, monkeypatch
):
    ir, _, units, _ = records
    monkeypatch.setattr("app.services.document_harness.controller.request_size", lambda *args: 500)
    calls = []
    runner, window = engine(
        records,
        ["Current reading"],
        [entity(ir, "a", units["Name: Alpha"])],
        lambda *args: calls.append(args),
        budget=250,
    )
    with pytest.raises(ValueError, match="harness_single_type_input_too_large"):
        runner.type_alignment(window)
    assert not calls


@pytest.mark.parametrize("comparison_type", ["urn:budget:A", None])
def test_conflicting_type_shards_require_joint_choice_instead_of_highest_self_confidence(
    records,
    monkeypatch,
    comparison_type,
):
    ir, _, units, _ = records
    calls = []
    compared = {"urn:budget:A", "urn:budget:Report"}
    subject = entity(ir, "a", units["Name: Alpha"], class_iri="urn:budget:B")
    subject["confidence"] = 0.99
    subject["type_evidence"] = subject["evidence"]

    def invoke(stage, payload, schema):
        assert stage == "type_alignment"
        cards = {card["iri"] for card in payload["classes"]}
        calls.append(deepcopy(payload))
        if cards == compared:
            assert len(payload["entities"]) == 1
            assert payload["entities"][0]["role"] == "source-object"
            assert "confidence" not in payload["entities"][0]
            assert "reason" not in payload["entities"][0]
            choice, confidence, reason = comparison_type, 0.82, "compared original referent"
        elif "urn:budget:Report" in cards:
            choice, confidence, reason = "urn:budget:Report", 0.95, "wrong shard confidence"
        else:
            choice, confidence, reason = "urn:budget:A", 0.9, "source object role"
        response = {
            "entities": {
                item["entity_id"]: {
                    "class_iri": choice,
                    "confidence": confidence,
                    "reason": reason,
                    "evidence": item["evidence"] if choice else [],
                }
                for item in payload["entities"]
            }
        }
        jsonschema.validate(response, schema)
        return response

    monkeypatch.setattr(
        "app.services.document_harness.controller.request_size",
        lambda stage, payload, schema: 100 * len(payload["classes"]),
    )
    runner, window = engine(records, ["Current reading"], [subject], invoke, budget=250)
    runner.type_alignment(window)
    assert len(calls) == 3
    assert {card["iri"] for card in calls[-1]["classes"]} == compared
    saved = runner.state["entities"]["a"]
    assert saved["class_iri"] == comparison_type
    assert saved["reason"] == "compared original referent"
    if comparison_type:
        assert saved["state"] == "candidate" and saved["confidence"] == 0.82
        assert {ref["source_id"] for ref in saved["type_evidence"]} == {
            units["Name: Alpha"].evidence_id,
        }
    else:
        assert saved["state"] == "candidate"
        assert saved["class_label"] is None and not saved["type_evidence"]
        assert "confidence" not in saved


def test_joint_type_comparison_over_budget_never_resplits_its_candidate_menu(records, monkeypatch):
    ir, _, units, _ = records
    compared = {"urn:budget:A", "urn:budget:Report"}
    calls = []

    def size(stage, payload, schema):
        cards = {card["iri"] for card in payload["classes"]}
        return 500 if len(cards) == 3 or cards == compared else 100

    def invoke(stage, payload, schema):
        cards = {card["iri"] for card in payload["classes"]}
        calls.append(cards)
        assert cards != compared
        return {
            "entities": {
                item["entity_id"]: {
                    "class_iri": (
                        "urn:budget:Report" if "urn:budget:Report" in cards else "urn:budget:A"
                    ),
                    "confidence": 0.95 if "urn:budget:Report" in cards else 0.9,
                    "evidence": item["evidence"],
                    "reason": "proposal",
                }
                for item in payload["entities"]
            }
        }

    monkeypatch.setattr("app.services.document_harness.controller.request_size", size)
    runner, window = engine(
        records,
        ["Current reading"],
        [entity(ir, "a", units["Name: Alpha"])],
        invoke,
        budget=250,
    )
    with pytest.raises(ValueError, match="harness_single_type_comparison_input_too_large"):
        runner.type_alignment(window)
    assert len(calls) == 2
    assert runner.state["entities"]["a"]["class_iri"] is None
    assert runner.state["cursor"]["main"]["phase"] == "skeleton"


def checked_review(payload):
    return {
        "type_concerns": [],
        "judgments": {
            row["id"]: {
                "verdict": "accepted",
                "confidence": 0.99,
                "evidence": row["evidence"],
                "reason": "original field or relation checked",
            }
            for row in payload["candidates"]
        },
    }


def relation_proposal(payload, *, verdict="no_relation", **changes):
    return {
        "proposals": {
            item["candidate_id"]: {
                "verdict": verdict,
                "evidence": item["clue_sources"],
                "polarity": "positive",
                "conditions": [],
                "participation": None,
                "selection": None,
                "timing": None, "ordered_object_ids": None, "order_evidence": [],
                "missing_context": "none",
                "reason": "original relation clue",
                "confidence": 0.99,
                **changes,
            }
            for item in payload["items"]
        }
    }


def add_hints(runner, window, subject_id, objects, predicates):
    refs = [
        ref for key in (subject_id, *objects) for ref in runner.state["entities"][key]["evidence"]
    ]
    runner.commit(
        {
            "hints": {
                f"{predicate}:{object_id}": {
                    "id": f"{predicate}:{object_id}",
                    "subject_id": subject_id,
                    "object_id": object_id,
                    "label": predicate,
                    "polarity": "positive",
                    "conditions": [],
                    "window_id": window.id,
                    "evidence": refs,
                }
                for predicate in predicates
                for object_id in objects
            }
        }
    )


def execute_planned(runner, window):
    plan_graph_work(runner)
    for _ in range(100):
        if not drain_work(runner, window):
            return
    pytest.fail("work did not drain within the bounded fixture")


def test_object_batches_keep_multiple_predicates_and_only_one_current_property_task(
    records,
    monkeypatch,
):
    ir, _, units, fields = records
    subject = entity(
        ir,
        "a",
        units["Name: Alpha"],
        class_iri="urn:budget:A",
        fields=[fields["Old"]["id"], fields["Current"]["id"]],
    )
    objects = [
        entity(ir, f"b{i}", units[f"Object: B{i}"], class_iri="urn:budget:B") for i in range(4)
    ]
    for row in [subject, *objects]:
        row["state"] = "accepted"
    calls = []

    def invoke(stage, payload, schema):
        calls.append((stage, deepcopy(payload)))
        if stage == "property_alignment":
            result = {
                "properties": {
                    key: {
                        "mappings": [
                            {
                                "predicate_iri": "urn:budget:value",
                                "value_component": "whole",
                                "value_quote": None, "unit_quote": None,
                                "confidence": 0.99,
                            }
                        ],
                        "reason": "original field",
                    }
                    for key in payload["property_field_ids"]
                }
            }
        elif stage == "relation_alignment":
            assert len(payload["items"]) == 1
            result = relation_proposal(payload)
        else:
            assert stage == "evidence_review"
            result = checked_review(payload)
        jsonschema.validate(result, schema)
        return result

    monkeypatch.setattr(
        "app.services.document_harness.controller.request_size",
        lambda stage, payload, schema: (
            500 if stage == "relation_alignment" and len(payload["items"]) > 1 else 1
        ),
    )
    runner, window = engine(
        records,
        ["Current: current-value", "Alpha has weight 10 kg"],
        [subject, *objects],
        invoke,
        budget=250,
        stage="planning",
    )
    unit = units["Alpha has weight 10 kg"]
    label = reference(
        ir, unit.evidence_id, unit.text.index("weight"), unit.text.index("weight") + 6
    )
    value = reference(ir, unit.evidence_id, unit.text.index("10 kg"), len(unit.text))
    field_id = identity("field", [label, value])
    runner.state["fields"][field_id] = {
        "id": field_id,
        "label": "weight",
        "value": "10 kg",
        "missing": False,
        "evidence": [label, value],
        "value_evidence": [value],
        "alias": "",
        "source_aliases": [],
    }
    runner.state["entities"]["a"]["field_ids"].append(field_id)
    old = make_work(
        "property_alignment",
        {"subject_id": "a", "field_id": fields["Old"]["id"]},
        runner.state,
        runner.catalog,
        runner.execution_policy,
        status="done",
    )
    runner.commit({"work": {old["id"]: old}})
    add_hints(runner, window, "a", [o["id"] for o in objects], ["uses", "contains"])
    before = list(runner.state["entities"]["a"]["field_ids"])
    execute_planned(runner, window)
    property_calls = [p for stage, p in calls if stage == "property_alignment"]
    relation_calls = [p for stage, p in calls if stage == "relation_alignment"]
    root_calls = [p for p in relation_calls if p["items"][0]["subject_id"] == "E0"]
    relation_calls = [p for p in relation_calls if p["items"][0]["subject_id"] != "E0"]
    assert len(root_calls) == 5  # The report can describe the five body objects.
    assert len(property_calls) == 1 and len(relation_calls) == 8
    assert all(not call["card"]["relations"] for call in property_calls)
    assert {item["predicate_iri"] for call in relation_calls for item in call["items"]} == {
        "urn:budget:uses",
        "urn:budget:contains",
    }
    assert {row["value"] for row in runner.state["properties"].values()} == {
        "current-value",
        "10 kg",
    }
    assert runner.state["entities"]["a"]["field_ids"] == before
    count = len(calls)
    execute_planned(runner, window)
    assert len(calls) == count


def test_root_alignment_excludes_historical_metadata_without_deleting_it(records):
    _, _, _, fields = records
    calls = []

    def invoke(stage, payload, schema):
        assert stage == "property_alignment"
        calls.append(deepcopy(payload))
        return {
            "properties": {
                key: {"mappings": [], "reason": "No matching predicate"}
                for key in payload["property_field_ids"]
            }
        }

    root_fields = [fields[k]["id"] for k in ("ReportOld", "Unused", "Report")]
    runner, window = engine(
        records, ["Report: current-report"], [], invoke, root_fields=root_fields, stage="planning"
    )
    execute_planned(runner, window)
    assert len(calls) == 1
    assert [field["value"] for field in calls[0]["subject"]["fields"]] == ["current-report"]
    assert all(
        "historic-report" not in source["text"] and "unrelated-report" not in source["text"]
        for source in calls[0]["sources"]
    )
    assert runner.state["entities"]["document"]["field_ids"] == root_fields


@pytest.mark.parametrize(
    "participation,selection,timing",
    [
        ("unknown", "unspecified", "sequential"),
        ("options", "unspecified", "parallel"),
        ("all", "exactly_one", "parallel"),
        ("unknown", "exactly_one", "unspecified"),
    ],
)
def test_contradictory_group_proposal_fails_without_repair_or_new_facts(
    records,
    participation,
    selection,
    timing,
):
    ir, _, units, _ = records
    objects = [
        entity(ir, key, units[f"Name: {name}"], class_iri="urn:budget:A")
        for key, name in (("a", "Alpha"), ("b", "Beta"))
    ]
    for row in objects:
        row["state"] = "accepted"
    calls = []

    def invoke(stage, payload, schema):
        assert stage == "relation_alignment"
        calls.append(deepcopy(payload))
        return relation_proposal(
            payload,
            verdict="proposed",
            participation=participation,
            selection=selection,
            timing=timing,
        )

    runner, window = engine(
        records, ["Name: Alpha", "Name: Beta"], objects, invoke, stage="planning"
    )
    runner.commit(
        {
            "hints": {
                "group": {
                    "id": "group",
                    "subject_id": "document",
                    "object_ids": ["a", "b"],
                    "label": "describes",
                    "polarity": "uncertain",
                    "conditions": [],
                    "evidence": [row["referent"] for row in objects],
                    "window_id": window.id,
                }
            }
        }
    )
    execute_planned(runner, window)
    assert len(calls) == 1 and "proposal_feedback" not in calls[0]
    assert not runner.state.get("relations") and not runner.state.get("relation_groups")
    assert any(row["status"] == "failed" for row in runner.state["work"].values())


def test_root_review_keeps_required_field_but_excludes_unrelated_history(records):
    _, _, _, fields = records
    calls = []

    def invoke(stage, payload, schema):
        assert stage == "evidence_review"
        calls.append(deepcopy(payload))
        result = checked_review(payload)
        for judgment in result["judgments"].values():
            judgment.update(verdict="unresolved", evidence=[], reason="source checked; unresolved")
        return result

    root_fields = [fields[k]["id"] for k in ("ReportOld", "Unused", "Report")]
    runner, window = engine(
        records,
        ["Report: current-report"],
        [],
        invoke,
        root_fields=root_fields,
        stage="evidence_review",
    )
    field = fields["ReportOld"]
    assertion = {
        "id": "p",
        "subject_id": "document",
        "predicate_iri": "urn:budget:value",
        "alignment_class_iri": "urn:budget:Report",
        "label": field["label"],
        "value": field["value"],
        "source_value": field["value"],
        "value_component": "whole",
        "value_evidence": field["value_evidence"],
        "field_id": field["id"],
        "state": "candidate",
        "reason": "source field",
        "evidence": field["evidence"],
        "confidence": 0.99,
        "window_id": window.id,
    }
    runner.state["cursor"]["main"].update(phase="semantic", semantic_step="assertions")
    changes = {}
    review_for(runner, changes, "properties", assertion)
    runner.commit(changes)
    assert not drain_work(runner, window)  # Historical evidence belongs to the final drain.
    assert not calls
    while drain_work(runner, None):
        pass
    assert calls
    assert {field["value"] for field in calls[0]["candidates"][0]["subject"]["fields"]} == {
        "historic-report",
    }
    assert all("unrelated-report" not in source["text"] for source in calls[0]["sources"])
    assert runner.state["entities"]["document"]["field_ids"] == root_fields


@pytest.mark.parametrize("field_count", [32, 33, 40])
def test_property_field_batches_cover_all_fields_without_repeating_relation_tasks(
    records,
    tmp_path,
    field_count,
):
    _, catalog, _, _ = records
    doc = Document()
    for text in [
        "Name: Alpha",
        "Object: Beta",
        "Alpha uses Beta and contains Beta.",
        *(f"Field {index}: value-{index}" for index in range(field_count)),
    ]:
        doc.add_paragraph(text)
    path = tmp_path / "many-fields.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    units = {unit.text: unit for unit in ir.evidence_units}
    window = make_window(
        ir,
        [(u.evidence_id, 0, len(u.text)) for u in ir.evidence_units],
        [u.evidence_id for u in ir.evidence_units],
    )
    fields = {field["label"]: field for field in window.fields}
    owned = [fields[f"Field {index}"]["id"] for index in range(field_count)]
    subject = entity(ir, "a", units["Name: Alpha"], class_iri="urn:budget:A", fields=owned)
    obj = entity(ir, "b", units["Object: Beta"], class_iri="urn:budget:B")
    subject["state"] = obj["state"] = "accepted"
    calls = []

    def invoke(stage, payload, schema):
        calls.append((stage, deepcopy(payload)))
        if stage == "property_alignment":
            response = {
                "properties": {
                    key: {
                        "mappings": [
                            {
                                "predicate_iri": "urn:budget:value",
                                "confidence": 0.99,
                                "value_component": "whole",
                                "value_quote": None, "unit_quote": None,
                            }
                        ],
                        "reason": "matching original meaning",
                    }
                    for key in reversed(payload["property_field_ids"])
                }
            }
        elif stage == "relation_alignment":
            response = relation_proposal(payload, verdict="proposed")
        else:
            assert stage == "evidence_review"
            response = checked_review(payload)
        jsonschema.validate(response, schema)
        return response

    runner, window = engine(
        (ir, catalog, units, fields),
        list(units),
        [subject, obj],
        invoke,
        stage="planning",
        budget=500000,
    )
    add_hints(runner, window, "a", ["b"], ["uses", "contains"])
    execute_planned(runner, window)
    property_calls = [payload for stage, payload in calls if stage == "property_alignment"]
    relation_calls = [payload for stage, payload in calls if stage == "relation_alignment"]
    assert [len(call["property_field_ids"]) for call in property_calls] == (
        [32] if field_count == 32 else [32, field_count - 32]
    )
    # Aliases are local to a paid batch; verify stable field IDs from applied rows.
    assert sum(len(call["property_field_ids"]) for call in property_calls) == field_count
    assert {row["field_id"] for row in runner.state["properties"].values()} == set(owned)
    assert {row["value"] for row in runner.state["properties"].values()} == {
        f"value-{index}" for index in range(field_count)
    }
    assert {item["predicate_iri"] for call in relation_calls for item in call["items"]} == {
        "urn:budget:uses",
        "urn:budget:contains",
        "urn:budget:describes",
    }
    assert sum(len(call["items"]) for call in relation_calls) == 4
    assert len(runner.state["relations"]) == 4
    assert len([r for r in runner.state["relations"].values()
                if r["subject_id"] == "document"]) == 2
    assert all(not call["card"]["relations"] for call in property_calls)


def test_subject_without_legal_properties_preserves_observations_without_model_call(records):
    ir, _, units, fields = records
    graph = Graph().parse(
        data="""
        @prefix : <urn:budget:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class . :A a owl:Class . :B a owl:Class .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :A .
        :uses a owl:ObjectProperty ; rdfs:domain :A ; rdfs:range :B .
        """,
        format="turtle",
    )
    catalog = catalog_from_graph(graph, "urn:budget:Report")
    field = fields["Current"]
    subject = entity(ir, "b", units["Name: Beta"], class_iri="urn:budget:B", fields=[field["id"]])
    subject["state"] = "accepted"
    calls = []
    runner, window = engine(
        (ir, catalog, units, fields),
        ["Current: current-value"],
        [subject],
        lambda *args: calls.append(args),
        stage="planning",
    )
    execute_planned(runner, window)
    assert calls == [] and not runner.state.get("properties")
    rows = [
        row for row in runner.state["observations"].values() if row.get("field_id") == field["id"]
    ]
    assert len(rows) == 1 and rows[0]["label"] == field["label"]
    assert rows[0]["evidence"] == field["evidence"]
    attempts = list(rows[0]["alignment_outcomes"].values())
    assert len(attempts) == 1 and attempts[0]["reason"] == "当前类型卡无合法属性可对齐"
    assert attempts[0]["state"] == "unmatched"
