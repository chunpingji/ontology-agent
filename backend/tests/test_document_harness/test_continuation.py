"""Incomplete discovery rereads smaller, source-exact intervals."""

import json
from copy import deepcopy

import pytest
from docx import Document

from app.services.document_harness.continuation import (
    restore_window_plan,
    serialize_window_plan,
    split_reading_window,
)
from app.services.document_harness.source import make_window, quote_for_reference
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


def document_ir(tmp_path, doc):
    path = tmp_path / "continuation.docx"
    doc.save(path)
    return build_document_ir(path, parse_docx_structure(path))


def complete_window(ir):
    return make_window(ir, [(u.evidence_id, 0, len(u.text)) for u in ir.evidence_units if u.text],
                       [u.evidence_id for u in ir.evidence_units if u.text])


def assert_primary_covered(parent, children):
    original = getattr(parent, "primary_ranges", [
        {"evidence_id": source["evidence_id"], "start": source["offset"],
         "end": source["offset"] + len(source["text"])}
        for source in parent.sources if source["evidence_id"] in parent.primary_ids
    ])
    for expected in original:
        spans = sorted((span["start"], span["end"]) for child in children
                       for span in child.primary_ranges
                       if span["evidence_id"] == expected["evidence_id"])
        pending = expected["start"]
        for start, end in spans:
            if end <= pending:
                continue
            assert start <= pending
            pending = end
        assert pending >= expected["end"]


def test_split_keeps_name_and_structural_context_but_reduces_both_child_inputs(tmp_path):
    doc = Document()
    doc.add_heading("对象资料", 1)
    doc.add_paragraph("名称：Alpha")
    for index in range(8):
        doc.add_paragraph(f"观察{index}：" + "这是这一项的原始描述；" * 12)
    ir = document_ir(tmp_path, doc)
    original = complete_window(ir)
    before = deepcopy(original)
    children = split_reading_window(ir, original)
    assert len(children) == 2
    assert original == before
    assert len({child.id for child in children} | {original.id}) == 3
    assert [child.id for child in children] == [
        child.id for child in split_reading_window(ir, original)
    ]
    parent_size = sum(len(source["text"]) for source in original.sources)
    for child in children:
        assert sum(len(source["text"]) for source in child.sources) < parent_size
        assert "名称：Alpha" in [source["text"] for source in child.sources]
        assert "对象资料" in [source["text"] for source in child.sources]
        for field in child.fields:
            for ref in field["evidence"]:
                quote = quote_for_reference(child, ref)
                assert quote is not None
                assert child.resolve(ir, quote) == ref
    assert_primary_covered(original, children)


def test_table_split_keeps_whole_data_rows_and_original_column_headers(tmp_path):
    doc = Document()
    table = doc.add_table(rows=7, cols=2)
    table.cell(0, 0).text = "记录号"
    table.cell(0, 1).text = "测量值"
    for index in range(1, 7):
        table.cell(index, 0).text = f"A0{index}"
        table.cell(index, 1).text = str(10 + index)
    ir = document_ir(tmp_path, doc)
    assert ir.tables[0]["header_row_count"] == 1
    original = complete_window(ir)
    children = split_reading_window(ir, original)
    assert len(children) == 2
    for child in children:
        assert {"记录号", "测量值"} <= {source["text"] for source in child.sources}
        rows = {source["row"] for source in child.sources if source["row"]}
        for row in rows:
            assert len([source for source in child.sources if source["row"] == row]) == 2
        assert {field["id"] for field in child.fields} <= {field["id"] for field in original.fields}
    assert_primary_covered(original, children)


def test_single_long_source_splits_with_overlap_and_offsets_resolve_exact_unicode(tmp_path):
    doc = Document()
    text = "正文甲😀乙" * 180
    doc.add_paragraph(text)
    ir = document_ir(tmp_path, doc)
    original = complete_window(ir)
    children = split_reading_window(ir, original)
    assert len(children) == 2
    first = children[0].primary_ranges[0]
    second = children[1].primary_ranges[0]
    assert first["start"] == 0
    assert second["end"] == len(text)
    assert first["end"] > second["start"]
    for child in children:
        for source in child.sources:
            actual = ir.unit(source["evidence_id"]).text
            assert actual[source["offset"]:source["offset"] + len(source["text"])] == source["text"]
    assert_primary_covered(original, children)


def test_recursion_strictly_reduces_reading_scope_and_stops_at_small_spans(tmp_path):
    doc = Document()
    doc.add_paragraph("长原文" * 300)
    ir = document_ir(tmp_path, doc)
    pending = [(complete_window(ir), 0)]
    leaves = []
    while pending:
        current, depth = pending.pop()
        assert depth < 15
        children = split_reading_window(ir, current)
        if children:
            assert_primary_covered(current, children)
            current_size = sum(len(source["text"]) for source in current.sources)
            assert all(sum(len(source["text"]) for source in child.sources) < current_size
                       for child in children)
            pending.extend((child, depth + 1) for child in children)
        else:
            leaves.append(current)
    assert len(leaves) > 1
    assert all(sum(len(source["text"]) for source in leaf.sources) < 128 for leaf in leaves)


def test_partial_long_literal_is_not_reintroduced_as_a_complete_original_field(tmp_path):
    doc = Document()
    doc.add_paragraph("描述：" + "长属性原值" * 180)
    ir = document_ir(tmp_path, doc)
    original = complete_window(ir)
    assert len(original.fields) == 1
    children = split_reading_window(ir, original)
    assert len(children) == 2
    assert all(child.fields == [] for child in children)
    assert_primary_covered(original, children)


def test_saved_plan_has_no_source_text_and_restores_short_ids_and_references_exactly(tmp_path):
    doc = Document()
    doc.add_paragraph("名称：Alpha")
    for index in range(6):
        doc.add_paragraph(f"字段{index}：" + "字段原值" * 30)
    ir = document_ir(tmp_path, doc)
    original = complete_window(ir)
    for child in split_reading_window(ir, original):
        plan = serialize_window_plan(child)
        assert "Alpha" not in json.dumps(plan, ensure_ascii=False)
        assert "text" not in json.dumps(plan)
        restored = restore_window_plan(ir, json.loads(json.dumps(plan)))
        assert restored.id == child.id
        assert restored.sources == child.sources
        assert restored.fields == child.fields
        assert restored.payload() == child.payload()
        assert restored.primary_ranges == child.primary_ranges
        assert serialize_window_plan(restored) == plan


def test_restore_rejects_changed_source_content_and_unavailable_field_bindings(tmp_path):
    doc = Document()
    doc.add_paragraph("名称：Alpha")
    doc.add_paragraph("颜色：blue")
    ir = document_ir(tmp_path, doc)
    plan = serialize_window_plan(complete_window(ir))
    changed = deepcopy(plan)
    changed["source_hash"] = "mismatch"
    with pytest.raises(ValueError, match="reading_plan_source_changed"):
        restore_window_plan(ir, changed)
    changed = deepcopy(plan)
    changed["field_ids"] = ["unavailable"]
    with pytest.raises(ValueError, match="reading_plan_original_field_not_available"):
        restore_window_plan(ir, changed)
    changed = deepcopy(plan)
    changed["primary_ranges"][0]["end"] = 100000
    with pytest.raises(ValueError, match="reading_plan_source_range_invalid"):
        restore_window_plan(ir, changed)


def test_indivisible_small_record_returns_no_children_instead_of_repeating_parent(tmp_path):
    doc = Document()
    doc.add_paragraph("只有一个小记录")
    ir = document_ir(tmp_path, doc)
    assert split_reading_window(ir, complete_window(ir)) == []
