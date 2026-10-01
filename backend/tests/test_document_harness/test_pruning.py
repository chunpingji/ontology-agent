"""Source-led candidate bounds and pruning do not imply a semantic rejection."""

from collections import Counter
from copy import deepcopy

import pytest
from docx import Document
from rdflib import Graph

from app.services.document_harness import planning
from app.services.document_harness.evidence_gate import table_scope_hash, try_rule_seed
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.planning import (
    admit_relation_work,
    build_source_index,
    collect_relation_seeds,
    resolve_predicates,
)
from app.services.document_harness.source import reference
from app.services.document_harness.work import DEFAULT_POLICY
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def catalog():
    graph = Graph().parse(data="""
        @prefix : <urn:prune:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        :Report a owl:Class . :Thing a owl:Class . :Other a owl:Class .
        :describes a owl:ObjectProperty; rdfs:domain :Report;
          rdfs:range [owl:unionOf (:Thing :Other)] .
        :uses a owl:ObjectProperty; rdfs:domain :Thing; rdfs:range :Thing;
          rdfs:label "uses" .
    """, format="turtle")
    return catalog_from_graph(graph, "urn:prune:Report")


def source_state(tmp_path, count, *, relation_word=""):
    names = [f"N{index:04}" for index in range(count)]
    doc = Document()
    doc.add_paragraph(" ".join([relation_word, *names]))
    path = tmp_path / "pruning.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    entities = {}
    for name in names:
        unit = next(unit for unit in ir.evidence_units if name in unit.text)
        start = unit.text.index(name)
        ref = reference(ir, unit.evidence_id, start, start + len(name))
        entities[name] = {
            "id": name, "label": name, "referent": ref, "evidence": [ref],
            "class_iri": "urn:prune:Thing", "state": "accepted", "field_ids": [],
        }
    return ir, {"entities": entities, "work": {}, "hints": {}, "fields": {}}


def seeds_for(state, count, *, protected=False):
    return [{
        "subject_id": "N0000", "object_ids": [f"N{index:04}"],
        "predicate_iri": "urn:prune:uses", "region_id": "region",
        "clue_refs": [state["entities"]["N0000"]["referent"]],
        "required_context_refs": [], "origin_ids": [f"hint:{index}"],
        "priority": "explicit" if protected else "structural",
        "polarity_hint": "positive", "condition_hints": [], "endpoint_hypothesis": False,
    } for index in range(1, count + 1)]


def test_one_thousand_unconnected_mentions_do_not_generate_pair_work(
    tmp_path, catalog, monkeypatch,
):
    ir, state = source_state(tmp_path, 1000)
    index = build_source_index(ir, state)
    checks = []
    original = planning.legal_relation

    def checked(*args):
        checks.append(1)
        return original(*args)

    monkeypatch.setattr(planning, "legal_relation", checked)
    seeds = collect_relation_seeds(ir, catalog, state, {"entities": state["entities"]},
                                   index, DEFAULT_POLICY)
    assert seeds == []
    assert checks == [], "type-compatible entities alone must not cause pairwise comparison"


def test_weak_materialization_stops_before_a_large_region_becomes_a_pair_matrix(
    tmp_path, catalog,
):
    ir, state = source_state(tmp_path, 1000, relation_word="uses")
    index = build_source_index(ir, state)
    seeds = collect_relation_seeds(ir, catalog, state, {"entities": state["entities"]},
                                   index, DEFAULT_POLICY)
    assert 0 < len(seeds) <= DEFAULT_POLICY["weak_candidates_per_region"]
    buckets = Counter((s["subject_id"], s["predicate_iri"], s["region_id"]) for s in seeds)
    assert max(buckets.values()) <= (
        2 * DEFAULT_POLICY["weak_candidates_per_subject_predicate_region"]
    )
    assert any(row["weak_pool_truncated"] for row in index.scope_changes.values())


def test_explicit_negative_and_conditional_clues_bypass_weak_quota(tmp_path, catalog):
    _ir, state = source_state(tmp_path, 20)
    weak = seeds_for(state, 8)
    explicit = seeds_for(state, 11, protected=True)[8:]
    negative, conditional = seeds_for(state, 13)[11:]
    negative["polarity_hint"] = "negative"
    conditional["condition_hints"] = ["only when required"]
    rows = admit_relation_work([*weak, *explicit, negative, conditional], state,
                              catalog, DEFAULT_POLICY)
    protected = [r for r in rows if r["input"]["priority"] == "explicit"
                 or r["input"]["polarity_hint"] == "negative" or r["input"]["condition_hints"]]
    assert len(protected) == 5
    assert all(row["status"] == "ready" for row in protected)
    assert Counter(row["status"] for row in rows) == {"ready": 9, "pruned": 4}
    assert all(row["reason_code"] == "weak_quota" for row in rows if row["status"] == "pruned")


def test_done_rows_do_not_consume_newly_changed_bucket_admission_quota(tmp_path, catalog):
    _ir, state = source_state(tmp_path, 8)
    initial = seeds_for(state, 6)
    rows = admit_relation_work(initial, state, catalog, DEFAULT_POLICY)
    for row in rows:
        if row["status"] == "ready":
            row.update(status="done", output_ids=[f"fact:{row['id']}"])
        state["work"][row["id"]] = row
    replanned = admit_relation_work(seeds_for(state, 7), state, catalog, DEFAULT_POLICY)
    assert sum(row["status"] == "done" for row in replanned) == 4
    assert sum(row["status"] == "ready" for row in replanned) == 3
    assert all(row["output_ids"] for row in replanned if row["status"] == "done")


def test_unchanged_bucket_never_rotates_pruned_work_or_repeats_ranking(tmp_path, catalog):
    _ir, state = source_state(tmp_path, 8)
    seeds = seeds_for(state, 7)
    rankings = []

    def rank(pairs):
        rankings.append(pairs)
        return list(range(len(pairs)))

    rows = admit_relation_work(seeds, state, catalog, DEFAULT_POLICY, rank)
    pruned = {row["id"] for row in rows if row["status"] == "pruned"}
    for row in rows:
        if row["status"] == "ready":
            row["status"] = "done"
        state["work"][row["id"]] = row
    resumed = admit_relation_work(seeds, state, catalog, DEFAULT_POLICY, rank)
    assert {row["id"] for row in resumed if row["status"] == "pruned"} == pruned
    assert not any(row["status"] == "ready" for row in resumed)
    assert len(rankings) == 1


def test_equal_ranking_scores_preserve_original_object_positions(tmp_path, catalog):
    _ir, state = source_state(tmp_path, 9)
    rows = admit_relation_work(seeds_for(state, 8), state, catalog, DEFAULT_POLICY,
                               lambda pairs: [0.0] * len(pairs))
    assert {row["input"]["object_ids"][0] for row in rows if row["status"] == "ready"} == {
        "N0001", "N0002", "N0003", "N0004",
    }


def test_unknown_type_is_not_confirmed_ontology_incompatibility(tmp_path, catalog):
    _ir, state = source_state(tmp_path, 2)
    subject, obj = state["entities"].values()
    obj.update(class_iri=None, state="unresolved")
    unknown = resolve_predicates(catalog, subject, [obj], "uses")
    assert not unknown["incompatibility_confirmed"]
    obj.update(class_iri="urn:prune:Other", state="accepted")
    illegal = resolve_predicates(catalog, subject, [obj], "uses")
    assert illegal["incompatibility_confirmed"]
    assert illegal["legal_iris"] == []


def test_explicit_predicate_with_untyped_endpoint_keeps_a_waiting_task(tmp_path, catalog):
    ir, state = source_state(tmp_path, 2)
    state["entities"]["N0001"].update(class_iri=None, state="unresolved")
    hint = {"id": "hint", "subject_id": "N0000", "object_id": "N0001", "label": "uses",
            "evidence": [state["entities"]["N0000"]["referent"]]}
    state["hints"][hint["id"]] = hint
    seeds = collect_relation_seeds(ir, catalog, state, {"hints": state["hints"]},
                                   build_source_index(ir, state), DEFAULT_POLICY)
    assert len(seeds) == 1
    rows = admit_relation_work(seeds, state, catalog, DEFAULT_POLICY)
    assert rows[0]["status"] == "waiting"
    assert rows[0]["reason_code"] == "type_or_constraint_unresolved"


def test_unresolved_ontology_constraint_is_waiting_without_semantic_rejection(tmp_path, catalog):
    _ir, state = source_state(tmp_path, 2)
    catalog = deepcopy(catalog)
    card = catalog.classes["urn:prune:Thing"]
    catalog.classes["urn:prune:Thing"] = card.model_copy(update={
        "relations": (card.relations[0].model_copy(update={"constraint_status": "unresolved"}),),
    })
    resolution = resolve_predicates(
        catalog, state["entities"]["N0000"], [state["entities"]["N0001"]], "uses",
    )
    assert resolution["unresolved_iris"] == ["urn:prune:uses"]
    assert not resolution["incompatibility_confirmed"]
    row = admit_relation_work(seeds_for(state, 1, protected=True), state,
                             catalog, DEFAULT_POLICY)[0]
    assert row["status"] == "waiting"
    assert not row["output_ids"]


def test_explicit_predicate_with_untyped_subject_does_not_require_a_type_card(tmp_path, catalog):
    ir, state = source_state(tmp_path, 2)
    state["entities"]["N0000"].update(class_iri=None, state="unresolved")
    state["hints"] = {"hint": {
        "id": "hint", "subject_id": "N0000", "object_id": "N0001", "label": "uses",
        "evidence": [state["entities"]["N0000"]["referent"]],
    }}
    seeds = collect_relation_seeds(ir, catalog, state, {"hints": state["hints"]},
                                   build_source_index(ir, state), DEFAULT_POLICY)
    assert len(seeds) == 1
    assert admit_relation_work(seeds, state, catalog, DEFAULT_POLICY)[0]["status"] == "waiting"


def table_state(tmp_path, headers=("Container", "Component", "Notes")):
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for column, text in enumerate(headers):
        table.cell(0, column).text = text
        table.cell(1, column).text = f"N{column:04}"
    path = tmp_path / "table-pruning.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    state = {"entities": {}, "fields": {}, "hints": {}, "work": {}}
    for unit in ir.evidence_units:
        if unit.row_index == 1:
            ref = reference(ir, unit.evidence_id, 0, len(unit.text))
            state["entities"][unit.text] = {
                "id": unit.text, "label": unit.text, "referent": ref, "evidence": [ref],
                "class_iri": "urn:prune:Thing", "state": "accepted", "field_ids": [],
            }
    rule = {
        "rule_id": "exact_table_relation", "version": "1",
        "document_root_class_iri": "urn:prune:Report", "table_path": ir.tables[0]["table_path"],
        "header_rows": 1, "header_cells": [
            {"row": 0, "column": col, "text": text} for col, text in enumerate(headers)
        ], "subject_column": 0, "object_column": 1,
        "subject_class_iri": "urn:prune:Thing", "object_class_iri": "urn:prune:Thing",
        "predicate_iri": "urn:prune:uses", "polarity": "positive", "conditions": [],
    }
    rule["static_scope_hash"] = table_scope_hash(ir, rule)
    return ir, state, rule


def test_configured_table_mapping_seeds_only_its_physical_value_cells(tmp_path, catalog):
    ir, state, rule = table_state(tmp_path)
    policy = {**DEFAULT_POLICY, "table_relation_rules": [rule]}
    seeds = collect_relation_seeds(ir, catalog, state, {"entities": state["entities"]},
                                   build_source_index(ir, state), policy)
    assert len(seeds) == 1
    assert (seeds[0]["subject_id"], seeds[0]["object_ids"]) == ("N0000", ["N0001"])
    assert seeds[0]["predicate_iri"] == "urn:prune:uses"
    assert try_rule_seed(ir, catalog, state, seeds[0], [rule]) is not None
    absent = collect_relation_seeds(ir, catalog, state, {"entities": state["entities"]},
                                    build_source_index(ir, state), DEFAULT_POLICY)
    assert absent == []


def test_table_mapping_with_competing_cell_mentions_does_not_pick_one(tmp_path, catalog):
    ir, state, rule = table_state(tmp_path)
    state["entities"]["alternative"] = {**state["entities"]["N0001"], "id": "alternative"}
    policy = {**DEFAULT_POLICY, "table_relation_rules": [rule]}
    seeds = collect_relation_seeds(ir, catalog, state, {"entities": state["entities"]},
                                   build_source_index(ir, state), policy)
    assert seeds == []


def test_table_header_requires_saved_field_ownership_and_value_containment(tmp_path, catalog):
    ir, state, _rule = table_state(tmp_path, headers=("Container", "uses", "Notes"))
    header = next(u for u in ir.evidence_units if u.text == "uses")
    value = state["entities"]["N0001"]["referent"]
    header_ref = reference(ir, header.evidence_id, 0, len(header.text))
    field = {"id": "field", "label": "uses", "value": value["text"],
             "value_evidence": [value], "evidence": [header_ref, value]}
    state["fields"]["field"] = field
    delta = {"fields": state["fields"], "entities": state["entities"]}
    assert collect_relation_seeds(ir, catalog, state, delta,
                                  build_source_index(ir, state), DEFAULT_POLICY) == []
    state["entities"]["N0000"]["field_ids"] = ["field"]
    seeds = collect_relation_seeds(ir, catalog, state, delta,
                                   build_source_index(ir, state), DEFAULT_POLICY)
    assert [(s["subject_id"], s["object_ids"]) for s in seeds] == [("N0000", ["N0001"])]
    assert header_ref in seeds[0]["clue_refs"]
    assert all("N0002" not in s["object_ids"] for s in seeds)
    state["entities"]["N0002"]["field_ids"] = ["field"]
    ambiguous = collect_relation_seeds(ir, catalog, state, delta,
                                       build_source_index(ir, state), DEFAULT_POLICY)
    assert {(s["subject_id"], tuple(s["object_ids"])) for s in ambiguous} == {
        ("N0000", ("N0001",)), ("N0002", ("N0001",)),
    }
