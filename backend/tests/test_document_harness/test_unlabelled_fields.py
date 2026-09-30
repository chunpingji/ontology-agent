"""Independent text keeps exact evidence; its property label comes from the card."""

import pytest
from docx import Document
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.ontology import catalog_from_graph
from app.services.extraction.word_analysis import analyze_word_core


@pytest.mark.parametrize("owner,bad_label", [("document", False), ("body", False),
                                            ("document", True)])
def test_unlabelled_value_has_explicit_ownership_and_never_invents_a_label_quote(
    tmp_path, owner, bad_label,
):
    doc = Document()
    doc.add_heading("工作记录", 1)
    doc.add_paragraph("对象：Alpha")
    doc.add_paragraph("仅限组内阅读")
    path = tmp_path / "source.docx"
    doc.save(path)
    ir = analyze_word_core(path).ir
    catalog = catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:unlabelled:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class . :Thing a owl:Class .
        :describes a owl:ObjectProperty ; rdfs:domain :Report ; rdfs:range :Thing .
        :audience a owl:DatatypeProperty ;
          rdfs:domain [ a owl:Class ; owl:unionOf (:Report :Thing) ] ;
          rdfs:range xsd:string ; rdfs:label "可阅对象" ;
          rdfs:comment "说明谁可以阅读该对象" .
    ''', format="turtle"), "urn:unlabelled:Report")
    calls = []

    def invoke(stage, payload, schema):
        calls.append(stage)
        if stage == "discover":
            assert payload["document"]["property_guidance"] == [
                {"label": "可阅对象", "description": "说明谁可以阅读该对象"},
            ]
            source = next(s for s in payload["sources"] if s["text"] == "仅限组内阅读")
            q = {"source_id": source["source_id"], "text": source["text"], "occurrence": 0}
            field = {"label": {**q, "text": "可阅对象"} if bad_label else None, "value": q}
            name_source = next(s for s in payload["sources"] if "Alpha" in s["text"])
            return {"entities": [{
                # Local mention IDs do not share the internal root namespace.
                "local_id": "document", "name": {"source_id": name_source["source_id"],
                "text": "Alpha", "occurrence": 0}, "role": "body_object",
                "anchor": {"source_id": name_source["source_id"], "text": "Alpha"},
                "evidence": [name_source["source_id"]], "field_ids": [], "source_fields": [field],
            }] if owner == "body" else [], "document_field_ids": [],
                "document_source_fields": [field] if owner == "document" else [],
                "unowned_fields": [], "relation_hints": [], "complete": True}
        if stage == "type_alignment":
            return {"entities": {e["entity_id"]: {
                "class_iri": "urn:unlabelled:Thing", "confidence": 0.99,
                "evidence": e["evidence"], "reason": "原文对象",
            } for e in payload["entities"]}}
        if stage == "assertion_alignment":
            fields = {f["field_id"]: f for f in payload["subject"]["fields"]}
            assert all(fields[f]["label"] == "" for f in payload["property_field_ids"])
            return {"properties": {f: {
                "mappings": [{"predicate_iri": "urn:unlabelled:audience",
                              "value_component": "whole", "value_quote": None, "confidence": 0.99}],
                "reason": "由原值和上下文判断属性含义，属性名来自卡片",
            } for f in payload["property_field_ids"]}, "relations": [], "complete": True}
        assert stage in {"entity_review", "evidence_review"}
        return {**({"type_concerns": []} if stage == "evidence_review" else {}),
                "judgments": {c["id"]: {
            "verdict": "accepted", "confidence": 0.99,
            "evidence": [*c["evidence"], *c["subject"]["evidence"]],
            "reason": "引用和主体归属经独立核对",
        } for c in payload["candidates"]}}

    engine = Engine(ir=ir, catalog=catalog, state={}, invoke=invoke, save=lambda changes: None,
                    should_stop=lambda: len(calls) == 1)
    engine.run()
    if bad_label:
        assert not engine.state.get("properties")
        assert not engine.state["entities"]["document"]["field_ids"]
        assert any(o["reason"] == "source_quote_mismatch"
                   for o in engine.state["observations"].values())
        return
    field = next(f for f in engine.state["fields"].values() if f["value"] == "仅限组内阅读")
    assert field["label"] == "" and len(field["evidence"]) == 1
    assert field["evidence"][0]["text"] == "仅限组内阅读"
    assert not engine.state.get("properties")
    root = engine.state["entities"]["document"]
    assert (field["id"] in root["field_ids"]) == (owner == "document")
    engine.should_stop = lambda: False
    engine.run()
    properties = list(engine.state["properties"].values())
    assert len(properties) == 1
    prop = properties[0]
    assert prop["state"] == "accepted" and prop["label"] == "可阅对象"
    assert prop["value"] == "仅限组内阅读"
    assert (prop["subject_id"] == "document") == (owner == "document")
