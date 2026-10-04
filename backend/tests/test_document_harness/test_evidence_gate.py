"""Source proof boundaries: valid text and legal types alone never prove a relation."""

from copy import deepcopy

import pytest
from docx import Document
from pydantic import ValidationError
from rdflib import Graph

from app.services.document_harness.evidence_gate import (
    assertion_dependency_hash,
    normalize_table_rules,
    prove_table_relation,
    route_semantic_review,
    table_scope_hash,
    try_rule_seed,
)
from app.services.document_harness.ontology import catalog_from_graph
from app.services.document_harness.source import reference
from app.services.extraction.document_ir import build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure


@pytest.fixture
def catalog():
    graph = Graph().parse(data="""
        @prefix : <urn:gate:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class . :Thing a owl:Class .
        :describes a owl:ObjectProperty; rdfs:domain :Report; rdfs:range :Thing .
        :contains a owl:ObjectProperty; rdfs:domain :Thing; rdfs:range :Thing .
        :identifier a owl:DatatypeProperty; rdfs:domain :Thing; rdfs:range xsd:string .
    """, format="turtle")
    return catalog_from_graph(graph, "urn:gate:Report")


def document(tmp_path, *, values=("Alpha", "Beta"), headers=("Container", "Component"),
             footnote="", merged=False, nested=False, multiple_paragraphs=False):
    doc = Document()
    doc.add_heading("Declared composition", 1)
    table = doc.add_table(rows=2, cols=len(headers))
    for column, text in enumerate(headers):
        table.cell(0, column).text = text
    for column, text in enumerate(values):
        table.cell(1, column).text = text
    if merged:
        table.cell(0, 0).merge(table.cell(0, 1))
    if nested:
        table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "nested"
    if multiple_paragraphs:
        table.cell(1, 0).add_paragraph("another value")
    if footnote:
        doc.add_paragraph(footnote)
    path = tmp_path / "gate.docx"
    doc.save(path)
    return build_document_ir(path, parse_docx_structure(path))


def entity(ir, name, key):
    unit = next(unit for unit in ir.evidence_units if unit.text.strip() == name)
    start = len(unit.text) - len(unit.text.lstrip())
    ref = reference(ir, unit.evidence_id, start, len(unit.text.rstrip()))
    return {
        "id": key, "label": name, "name": ref, "referent": ref, "evidence": [ref],
        "type_evidence": [ref], "class_iri": "urn:gate:Thing", "role": "object",
        "state": "accepted", "field_ids": [],
    }


def state_and_claim(ir):
    left, right = entity(ir, "Alpha", "a"), entity(ir, "Beta", "b")
    state = {"entities": {"a": left, "b": right}, "execution_policy": {"table_relation_rules": []}}
    claim = {
        "id": "claim", "subject_id": "a", "object_id": "b",
        "predicate_iri": "urn:gate:contains", "alignment_class_iri": "urn:gate:Thing",
        "polarity": "positive", "conditions": [], "state": "candidate",
        "evidence": [left["referent"], right["referent"]],
    }
    return state, claim


def frozen_rule(ir):
    rule = {
        "rule_id": "exact_table_relation", "version": "1",
        "document_root_class_iri": "urn:gate:Report", "table_path": ir.tables[0]["table_path"],
        "header_rows": 1, "header_cells": [
            {"row": 0, "column": 0, "text": "Container"},
            {"row": 0, "column": 1, "text": "Component"},
        ],
        "subject_column": 0, "object_column": 1,
        "subject_class_iri": "urn:gate:Thing", "predicate_iri": "urn:gate:contains",
        "object_class_iri": "urn:gate:Thing", "static_scope_hash": "0" * 64,
        "polarity": "positive", "conditions": [],
    }
    rule["static_scope_hash"] = table_scope_hash(ir, rule)
    return normalize_table_rules([rule])[0]


def test_exact_quotes_and_legal_classes_are_not_relation_proof(tmp_path, catalog):
    """V01/V04: two sibling items and a plausible predicate still need semantics."""
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    assert route_semantic_review(ir, catalog, state, claim)["action"] == "semantic"
    assert prove_table_relation(ir, catalog, state, claim, []) is None
    assert try_rule_seed(ir, catalog, state, claim, []) is None


def test_frozen_table_proves_seed_before_alignment_and_locates_all_support(tmp_path, catalog):
    """V02: this path is pure and returns a proof without any model port."""
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    rule = frozen_rule(ir)
    state["execution_policy"]["table_relation_rules"] = [rule]
    assertion, proof = try_rule_seed(ir, catalog, state, claim, [rule])
    assert assertion["subject_id"] == "a"
    assert assertion["object_id"] == "b"
    assert proof["rule_id"] == "exact_table_relation"
    assert proof["rule_version"] == "1"
    assert proof["dependency_hash"] == assertion_dependency_hash(ir, catalog, state, assertion)
    quoted = [reference(ir, *key)["text"] for key in proof["evidence_refs"]]
    assert {"Alpha", "Beta", "Container", "Component", "Declared composition"} <= set(quoted)
    assert route_semantic_review(ir, catalog, state, assertion)["action"] == "proven"


def test_rule_seed_uses_physical_clue_references_from_work_contract(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    seed = {
        "subject_id": "a", "object_ids": ["b"], "predicate_iri": "urn:gate:contains",
        "clue_refs": [[ref["source_id"], ref["start"], ref["end"]] for ref in claim["evidence"]],
        "required_context_refs": [], "polarity_hint": "positive", "condition_hints": [],
    }
    assert try_rule_seed(ir, catalog, state, seed, [frozen_rule(ir)]) is not None
    seed["condition_hints"] = ["Only for X"]
    assert try_rule_seed(ir, catalog, state, seed, [frozen_rule(ir)]) is None


@pytest.mark.parametrize("change", [
    {"headers": ("Alternatives", "Component")},
    {"footnote": "The container does not contain this component."},
    {"headers": ("Container", "Component", "Condition"), "values": ("Alpha", "Beta", "If X")},
    {"merged": True}, {"nested": True}, {"multiple_paragraphs": True},
])
def test_changed_source_structure_cannot_reuse_table_template(tmp_path, catalog, change):
    """V03: differences outside the declared dynamic cells invalidate proof."""
    original = document(tmp_path)
    rule = frozen_rule(original)
    changed = document(tmp_path, **change)
    state, claim = state_and_claim(changed)
    assert prove_table_relation(changed, catalog, state, claim, [rule]) is None


@pytest.mark.parametrize("change", ["duplicate", "negation", "condition", "unconfirmed", "group"])
def test_ambiguous_semantic_scope_never_gets_table_proof(tmp_path, catalog, change):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    if change == "duplicate":
        state["entities"]["b2"] = {**state["entities"]["b"], "id": "b2"}
    elif change in {"negation", "condition"}:
        state["hints"] = {"h": {
            **claim, "polarity": "negative" if change == "negation" else "positive",
            "conditions": ["Only if X"] if change == "condition" else [],
        }}
    elif change == "unconfirmed":
        state["entities"]["b"]["state"] = "candidate"
    else:
        claim["object_ids"] = ["a", "b"]
    assert prove_table_relation(ir, catalog, state, claim, [frozen_rule(ir)]) is None


def test_static_hash_ignores_only_dynamic_values_not_document_identity(tmp_path):
    original = document(tmp_path)
    rule = frozen_rule(original)
    changed_values = document(tmp_path, values=("Gamma", "Delta"))
    assert original.document_hash != changed_values.document_hash
    assert table_scope_hash(changed_values, rule) == rule["static_scope_hash"]
    changed_scope = document(
        tmp_path, values=("Gamma", "Delta"), footnote="Unless stated otherwise",
    )
    assert table_scope_hash(changed_scope, rule) != rule["static_scope_hash"]


def test_forged_hash_does_not_bypass_merged_grid_rejection(tmp_path, catalog):
    ir = document(tmp_path, merged=True)
    state, claim = state_and_claim(ir)
    rule = frozen_rule(ir)
    assert prove_table_relation(ir, catalog, state, claim, [rule]) is None


def property_claim(ir):
    state, _ = state_and_claim(ir)
    value = state["entities"]["a"]["referent"]
    field = {"id": "field", "label": "identifier", "value": "Alpha", "missing": False,
             "evidence": [value], "value_evidence": [value]}
    state["fields"] = {"field": field}
    state["entities"]["a"].update(field_ids=["field"], identity_binding={
        "identifiers": [{
            "field_id": "field", "property_iri": "urn:gate:identifier", "value": "Alpha",
        }],
    })
    return state, {
        "id": "p", "subject_id": "a", "predicate_iri": "urn:gate:identifier", "field_id": "field",
        "value": "Alpha", "source_value": "Alpha", "source_unit": None,
        "value_component": "whole", "value_evidence": [value], "evidence": [value],
    }


def test_bound_identifier_is_precheck_not_semantic_proof(tmp_path, catalog):
    """V05: matching an identifier does not prove what that identifier means."""
    ir = document(tmp_path)
    state, claim = property_claim(ir)
    result = route_semantic_review(ir, catalog, state, claim)
    assert result["action"] == "semantic"
    assert result["proof"] is None
    changed = {**claim, "value": "Beta"}
    assert route_semantic_review(ir, catalog, state, changed)["reason_code"] == "bound_value_changed"


def test_value_quote_outside_owned_field_is_invalid(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = property_claim(ir)
    claim.update(value_component="span", value="Beta",
                 value_evidence=[state["entities"]["b"]["referent"]])
    assert route_semantic_review(ir, catalog, state, claim)["reason_code"] == "value_outside_field"


@pytest.mark.parametrize("verdict", ["accepted", "rejected", "unresolved"])
def test_exact_reuse_preserves_llm_verdict_and_method(tmp_path, catalog, verdict):
    """V13: cached rejection or uncertainty is never promoted to rule acceptance."""
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    claim["verification"] = {
        "method": "llm", "semantic_verdict": verdict,
        "dependency_hash": assertion_dependency_hash(ir, catalog, state, claim),
        "rule_id": None, "rule_version": None, "rule_hash": None,
    }
    result = route_semantic_review(ir, catalog, state, claim)
    assert result["action"] == "reuse"
    assert result["reused_verification"] == claim["verification"]
    assert result["proof"] is None


def test_confirmation_is_gate_only_but_changed_meaning_invalidates_proof(tmp_path, catalog):
    """V11/V14: changed status can reuse semantics; a changed role cannot."""
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    claim["verification"] = {
        "method": "llm", "semantic_verdict": "accepted",
        "dependency_hash": assertion_dependency_hash(ir, catalog, state, claim),
    }
    for status in ("candidate", "accepted", "unresolved"):
        state["entities"]["b"]["state"] = status
        assert route_semantic_review(ir, catalog, state, claim)["action"] == "reuse"
    state["entities"]["b"]["role"] = "alternative"
    assert route_semantic_review(ir, catalog, state, claim)["action"] == "semantic"


def test_new_negation_evidence_invalidates_old_proof(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    claim["verification"] = {
        "method": "llm", "semantic_verdict": "accepted",
        "dependency_hash": assertion_dependency_hash(ir, catalog, state, claim),
    }
    state["hints"] = {"h": {"subject_id": "a", "object_id": "b", "evidence": claim["evidence"],
                            "polarity": "negative", "conditions": []}}
    assert route_semantic_review(ir, catalog, state, claim)["action"] == "semantic"


def test_unused_new_field_does_not_invalidate_relation_proof(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    before = assertion_dependency_hash(ir, catalog, state, claim)
    state["entities"]["a"]["field_ids"].append("irrelevant")
    state["fields"] = {"irrelevant": {"id": "irrelevant", "value": "unrelated observation"}}
    assert assertion_dependency_hash(ir, catalog, state, claim) == before


def test_reference_recall_results_do_not_change_semantic_proof_fingerprint(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    state["reference_cues"] = {"cue": {
        "subject_id": "a", "target_expression": state["entities"]["b"]["referent"],
        "evidence": claim["evidence"], "direction": "outgoing",
    }}
    before = assertion_dependency_hash(ir, catalog, state, claim)
    state["reference_cues"]["cue"].update(target_ids=["b"], reference_targets_truncated=False)
    assert assertion_dependency_hash(ir, catalog, state, claim) == before
    state["reference_cues"]["cue"]["direction"] = "incoming"
    assert assertion_dependency_hash(ir, catalog, state, claim) != before


def test_invalid_references_precede_reuse(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    claim["evidence"][0] = {**claim["evidence"][0], "text": "Not Alpha"}
    claim["verification"] = {
        "method": "llm", "semantic_verdict": "accepted",
        "dependency_hash": assertion_dependency_hash(ir, catalog, state, claim),
    }
    result = route_semantic_review(ir, catalog, state, claim)
    assert (result["action"], result["reason_code"]) == ("invalid", "invalid_reference")


@pytest.mark.parametrize("members,participation,selection,timing", [
    (["b", "b"], "all", "unspecified", "unspecified"),
    (["a", "b"], "options", "exactly_one", "parallel"),
    (["a", "b"], "all", "exactly_one", "unspecified"),
])
def test_invalid_group_contract_never_reaches_semantics(
    tmp_path, catalog, members, participation, selection, timing,
):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    claim.pop("object_id")
    claim.update(
        object_ids=members, participation=participation, selection=selection, timing=timing,
    )
    assert route_semantic_review(ir, catalog, state, claim)["reason_code"] == "invalid_group_contract"


def test_unknown_participation_waits_without_flattening_members(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    claim.pop("object_id")
    claim.update(object_ids=["a", "b"], participation="unknown",
                 selection="unspecified", timing="unspecified")
    result = route_semantic_review(ir, catalog, state, claim)
    assert result["action"] == "waiting"
    assert claim["object_ids"] == ["a", "b"]


def test_single_edge_cannot_bypass_whole_referent_group(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    state["entities"]["b"]["identity_binding"] = {"group_id": "options", "identifiers": []}
    state["entities"]["b2"] = {**state["entities"]["b"], "id": "b2"}
    assert route_semantic_review(ir, catalog, state, claim)["reason_code"] == "invalid_group_contract"


def test_hypothesis_type_can_be_reviewed_but_missing_type_waits(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    state["entities"]["a"]["state"] = "candidate"
    assert route_semantic_review(ir, catalog, state, claim)["action"] == "semantic"
    state["entities"]["a"]["class_iri"] = None
    assert route_semantic_review(ir, catalog, state, claim)["action"] == "waiting"


def test_known_illegal_predicate_is_invalid_but_unresolved_constraint_waits(tmp_path, catalog):
    ir = document(tmp_path)
    state, claim = state_and_claim(ir)
    illegal = {**claim, "predicate_iri": "urn:gate:unrelated"}
    assert route_semantic_review(ir, catalog, state, illegal)["action"] == "invalid"
    data = catalog.model_dump(mode="json")
    relation = next(row for row in data["classes"]["urn:gate:Thing"]["relations"]
                    if row["iri"] == "urn:gate:contains")
    relation["constraint_status"] = "unresolved"
    assert route_semantic_review(ir, data, state, claim)["reason_code"] == "ontology_unresolved"


@pytest.mark.parametrize("mutation", [
    {"subject_column": True}, {"subject_column": "0"}, {"extra": True},
    {"object_column": 0}, {"polarity": "negative"}, {"conditions": ["if"]},
    {"table_path": [""]}, {"static_scope_hash": "invalid"},
])
def test_invalid_rule_configuration_fails_new_run(tmp_path, mutation):
    rule = {**frozen_rule(document(tmp_path)), **mutation}
    with pytest.raises((ValueError, ValidationError)):
        normalize_table_rules([rule])


def test_rule_configuration_rejects_duplicates_and_never_mutates_inputs(tmp_path):
    rule = frozen_rule(document(tmp_path))
    original = deepcopy(rule)
    normalized = normalize_table_rules([rule])
    normalized[0]["header_cells"][0]["text"] = "Changed"
    assert rule == original
    with pytest.raises(ValueError, match="duplicate_table_rule"):
        normalize_table_rules([rule, rule])
    with pytest.raises(ValueError, match="must_be_a_list"):
        normalize_table_rules("{}")
