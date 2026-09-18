"""Actual Word artifacts must retain template layout without sample business content."""

from base64 import b64decode
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from app.services.reporting.docx_layout import freeze_layout
from app.services.reporting.output_renderer import render_docx
from app.services.reporting.template_v2 import ReportingError

PIXEL = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aL1sAAAAASUVORK5CYII="
)


def without_text_color(xml):
    """Only text color may differ under the user's black-text export rule."""
    root = etree.fromstring(xml.encode() if isinstance(xml, str) else xml)
    for color in list(root.iter(qn("w:color"))):
        color.getparent().remove(color)
    for properties in list(root.iter(qn("w:rPr"))):
        if not len(properties) and not properties.attrib:
            properties.getparent().remove(properties)
    return etree.tostring(root, method="c14n")


def assert_black_text(document):
    runs = []
    for part in document.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            continue
        for color in root.iter(qn("w:color")):
            assert dict(color.attrib) == {qn("w:val"): "000000"}
        for run in root.iter(qn("w:r")):
            runs.append(run)
            assert run.xpath("./w:rPr/w:color/@w:val") == ["000000"]
    assert runs


def styled_sample(path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "SimSun", Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.space_after = Pt(7)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.7), Cm(2.1)
    section.header_distance, section.footer_distance = Cm(1.1), Cm(1.3)
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    header = section.header.add_table(rows=1, cols=2, width=Cm(15))
    header.cell(0, 0).text = "模板页眉"
    header.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 0, 0)
    header.cell(0, 1).paragraphs[0].add_run().add_picture(BytesIO(PIXEL), width=Cm(0.3))
    section.first_page_header.paragraphs[0].text = "首页页眉"
    section.even_page_header.paragraphs[0].text = "偶数页页眉"
    footer = section.footer.paragraphs[0]
    footer.text = "第 "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    footer.runs[0].font.color.rgb = RGBColor(255, 0, 0)
    section.first_page_footer.paragraphs[0].text = "首页页脚"
    section.even_page_footer.paragraphs[0].text = "偶数页页脚"
    doc.add_paragraph("SAMPLE-BUSINESS-DO-NOT-COPY")
    table = doc.add_table(rows=2, cols=1)
    table.style = "Light Shading Accent 1"
    table.cell(0, 0).text = "SAMPLE-HEADER"
    table.cell(1, 0).text = "SAMPLE-VALUE"
    doc.save(path)
    return doc


def test_anchored_layout_preserves_sections_nested_tables_and_direct_formatting(tmp_path):
    path = tmp_path / "template.docx"
    sample = styled_sample(path)
    p = sample.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(13)
    p.runs[0].font.name, p.runs[0].font.size = "KaiTi", Pt(16)
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string("123456")
    outer = sample.tables[0]
    nested = outer.cell(1, 0).add_table(rows=2, cols=2)
    nested.autofit = False
    nested.columns[0].width, nested.columns[1].width = Cm(2), Cm(5)
    nested.style = "Light Shading Accent 1"
    for i, row in enumerate(nested.rows):
        for cell in row.cells:
            cell.text = "SAMPLE-NESTED"
            cell.paragraphs[0].runs[0].font.size = Pt(12 if i == 0 else 9)
            cell.paragraphs[0].runs[0].bold = i == 0
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "ABCDEF" if i == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shading)
    second = sample.add_section(WD_SECTION_START.NEW_PAGE)
    second.page_width, second.page_height = Cm(29.7), Cm(21)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "第二节页眉"
    sample.add_paragraph("SAMPLE-SECOND-SECTION")
    sample.save(path)
    raw = path.read_bytes()
    digest = sha256(raw).hexdigest()

    def origin(**anchor):
        return {"document_hash": digest, "label_anchor": anchor}

    schema = {"sections": [{"section_id": "sec", "groups": [{"group_id": "g", "units": [
        {"output_id": "prose", "origin": origin(paragraph_index=0)},
        {"output_id": "rows", "origin": origin(
            table_path=["table:0", "cell:1:0", "table:0"],
            row_index=0, column_index=0, paragraph_index=0,
        )},
        {"output_id": "last", "origin": origin(paragraph_index=2)},
        {"output_id": "note", "origin": origin(
            table_path=["table:0", "cell:1:0", "table:0"],
            row_index=1, column_index=0, paragraph_index=0,
        )},
    ]}]}]}
    layout = freeze_layout(SimpleNamespace(sample_docx_path=str(path)), schema)
    # Replacing/deleting the template after queueing cannot affect this run.
    path.unlink()
    ast = {"node_id": "doc", "kind": "document", "children": [
        {"node_id": "a", "kind": "group", "text": "新报告标题", "children": [
            {"node_id": "p", "kind": "paragraph", "text": "当前报告正文"},
        ]},
        {"node_id": "b", "kind": "group", "text": "设备表", "children": [
            {"node_id": "t", "kind": "table", "children": [
                {"node_id": f"r{i}", "kind": "row", "children": [
                    {"node_id": f"c{i}{j}", "kind": "cell", "text": f"新值{i}{j}"}
                    for j in range(3)
                ]} for i in range(4)
            ]},
        ]},
        {"node_id": "c", "kind": "group", "text": "第二节标题", "children": [
            {"node_id": "p2", "kind": "paragraph", "text": "第二节本次内容"},
        ]},
        {"node_id": "d", "kind": "group", "text": "备注", "children": [
            {"node_id": "note-p", "kind": "paragraph", "text": "生成的备注不可丢失"},
        ]},
    ]}
    repeated = deepcopy(ast["children"][1])
    repeated["node_id"] = "repeated"
    repeated["children"][0]["children"][1]["children"][0]["text"] = "第二组设备"
    ast["children"].insert(2, repeated)
    output = render_docx(ast, {"font": "Arial", "font_size_pt": 11}, layout=layout,
                         layout_nodes={"a": "prose", "b": "rows", "c": "last", "d": "note",
                                       "repeated": "rows"})
    result = Document(BytesIO(output))
    assert "SAMPLE" not in result.element.body.xml
    assert len(result.sections) == 2
    for original, generated in zip(sample.sections, result.sections):
        assert generated._sectPr.xml == original._sectPr.xml
    title = next(p for p in result.paragraphs if p.text == "新报告标题")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.paragraph_format.space_before == Pt(13)
    assert title.runs[0].font.name == "KaiTi" and title.runs[0].font.size == Pt(16)
    assert title.runs[0].bold and str(title.runs[0].font.color.rgb) == "000000"
    generated_table = result.tables[0].cell(1, 0).tables[0]
    assert len(generated_table.rows) == 4 and len(generated_table.columns) == 3
    assert generated_table.cell(3, 2).text == "新值32"
    assert result.tables[0].cell(1, 0).tables[1].cell(1, 0).text == "第二组设备"
    assert "生成的备注不可丢失" in result.tables[0].cell(1, 0).text
    assert generated_table.style.name == nested.style.name
    assert generated_table.cell(0, 0).paragraphs[0].runs[0].font.size == Pt(12)
    assert generated_table.cell(3, 0).paragraphs[0].runs[0].font.size == Pt(9)
    assert generated_table.cell(0, 0)._tc.xpath("./w:tcPr/w:shd/@w:fill") == ["ABCDEF"]
    assert abs(sum(c.width for c in generated_table.columns) - Cm(7)) < 2000
    assert result.styles["Normal"].element.xml == sample.styles["Normal"].element.xml
    assert_black_text(result)
    with ZipFile(BytesIO(raw)) as before, ZipFile(BytesIO(output)) as after:
        for name in before.namelist():
            if name.startswith(("word/header", "word/footer")):
                assert without_text_color(after.read(name)) == without_text_color(before.read(name))
            elif name.startswith("word/_rels/header"):
                assert after.read(name) == before.read(name)
        assert any(after.read(n) == PIXEL for n in after.namelist() if n.startswith("word/media/"))


def test_flow_template_uses_sample_fonts_headers_and_custom_body_override(tmp_path):
    path = tmp_path / "sample.docx"
    styled_sample(path)
    schema = {"sections": []}
    layout = freeze_layout(SimpleNamespace(sample_docx_path=str(path)), schema)
    ast = {"node_id": "doc", "kind": "document", "children": [
        {"node_id": "p", "kind": "paragraph", "text": "新正文"},
    ]}
    result = Document(BytesIO(render_docx(ast, layout=layout)))
    assert result.paragraphs[0].text == "新正文"
    assert result.styles["Normal"].font.name == "SimSun"
    assert result.sections[0].header.tables[0].cell(0, 0).text == "模板页眉"
    assert result.sections[0].footer._element.xpath(".//w:fldSimple/@w:instr") == ["PAGE"]
    custom = deepcopy(layout)
    custom["custom_style"] = True
    styled = Document(BytesIO(render_docx(
        ast, {"font": "黑体", "font_size_pt": 14}, layout=custom,
    )))
    run = styled.paragraphs[0].runs[0]
    assert run.font.name == "黑体" and run.font.size == Pt(14)
    assert run._r.xpath("./w:rPr/w:rFonts/@w:eastAsia") == ["黑体"]
    assert styled.sections[0].header._element.xml == result.sections[0].header._element.xml


def test_template_theme_and_conditional_table_colors_export_as_black(tmp_path):
    path = tmp_path / "colored.docx"
    sample = styled_sample(path)
    normal = sample.styles["Normal"].element.get_or_add_rPr()
    color = normal.get_or_add_color()
    color.set(qn("w:val"), "FF0000")
    color.set(qn("w:themeColor"), "accent2")
    color.set(qn("w:themeTint"), "99")
    conditional = OxmlElement("w:tblStylePr")
    conditional.set(qn("w:type"), "firstRow")
    properties = OxmlElement("w:rPr")
    properties.append(deepcopy(color))
    conditional.append(properties)
    sample.tables[0].style.element.append(conditional)
    sample.save(path)
    original = path.read_bytes()
    layout = freeze_layout(SimpleNamespace(sample_docx_path=str(path)), {"sections": []})
    ast = {"node_id": "t", "kind": "table", "children": [
        {"node_id": "r", "kind": "row", "children": [
            {"node_id": "c", "kind": "cell", "text": "待补充"},
        ]},
    ]}
    document = Document(BytesIO(render_docx(ast, layout=layout)))
    assert_black_text(document)
    assert document.tables[0].cell(0, 0).text == "待补充"
    assert without_text_color(document.styles.element.xml) == without_text_color(
        sample.styles.element.xml,
    )
    assert path.read_bytes() == original


def test_missing_template_file_does_not_silently_generate_default_format(tmp_path):
    with pytest.raises(ReportingError, match="模板 Word 样例不可用"):
        freeze_layout(SimpleNamespace(sample_docx_path=str(tmp_path / "missing.docx")),
                      {"sections": []})
