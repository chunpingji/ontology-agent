"""以示例 .docx 作为报告输出格式模板（sample_docx_path）——渲染层回归测试。

覆盖此前无测试保护、导致缺陷上线的问题及需求核心：
  - FIX A: Normal 无 rPr 时强设 eastAsia 会 AttributeError 崩溃、被宽 except 吞掉后
    静默回退空白文档（模板格式被悄悄丢弃）；且强设会篡改模板自有中文字体。
  - FIX B: 模板缺 Heading/Table Grid/List Bullet 内建样式时 add_heading / 表格样式
    KeyError 使整篇渲染失败。
  - 需求：页面布局、页眉、正文字体随模板继承；坏文件容错回退空白文档。

Codex round-2 复审新增锁定：
  - #1 正文表格/矩阵不得固定 latin 字体（须随模板），仅保留紧凑字号。
  - #2 多节模板（纵向封面 + 横向正文）的中间节布局不得被 _clear_body 丢弃。
  - #3 styleId 存在但 name 被改时，_ensure_builtin_styles 后 add_heading 仍须可用。
  - #8 模板省略 pgSz/pgMar（None 几何）时，矩阵缩放不得 None-None 崩溃。
  - #10 无模板路径须锁定横向 A4 + 宋体 10.5pt + 默认页眉；列宽容差收紧到 twip 级。
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.services.reporting.docx_renderer import (
    _add_inline_runs,
    _add_markdown_body,
    _clear_body,
    _ensure_builtin_styles,
    _render_assessment_matrix,
    render_risk_report,
)
from app.services.reporting.risk_report_generator import RiskReport, RiskRow


def _adversarial_template(tmp_path) -> str:
    """构造对抗性示例模板：纵向 A4 + 自定义左边距 + 自有页眉 + Normal eastAsia=仿宋，
    并**移除** Heading/Table Grid/List Bullet 内建样式（模拟 LibreOffice/精简导出）。"""
    tpl = Document()
    styles_el = tpl.styles.element
    for st in list(tpl.styles):
        if st.name in ("Heading 1", "Heading 2", "Heading 3", "Table Grid", "List Bullet"):
            styles_el.remove(st.element)
    sec = tpl.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)  # 纵向
    sec.left_margin = Cm(3.17)
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hdr.paragraphs[0].add_run("★模板自有页眉★")
    normal = tpl.styles["Normal"]
    normal.font.name = "仿宋"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋")
    path = tmp_path / "sample.docx"
    tpl.save(str(path))
    return str(path)


def _two_section_template(tmp_path) -> str:
    """纵向封面 + 横向正文的双节模板：两节页面尺寸不同，用于验证中间节不被清空丢弃。"""
    tpl = Document()
    s0 = tpl.sections[0]
    s0.page_width = Cm(21.0)
    s0.page_height = Cm(29.7)  # 第 0 节：纵向封面（属性进入分节段落的 sectPr）
    tpl.add_paragraph("封面")
    s1 = tpl.add_section(WD_SECTION.NEW_PAGE)
    s1.page_width = Cm(29.7)
    s1.page_height = Cm(21.0)  # 第 1 节：横向正文（属性进入 body 末尾 sectPr）
    tpl.add_paragraph("正文")
    path = tmp_path / "twosec.docx"
    tpl.save(str(path))
    return str(path)


def _report() -> RiskReport:
    return RiskReport(
        subject_description="测试对象",
        source_document_name="HRS-1597 CMCReport.docx",
        source_document_type="CMCReport",
        conclusion="测试结论",
        team_members=[{"name": "张三", "title": "工程师"}],
    )


def _report_with_row() -> RiskReport:
    return RiskReport(
        assessment_rows=[
            RiskRow(
                hazid="H1",
                contributing_factors="致因",
                pre_control_level="高",
                post_control_level="低",
                control_measures="措施 measure",
                traceability="T-01",
                status="进行中",
            )
        ]
    )


class TestSampleTemplateRendering:
    def test_vanilla_normal_has_no_rpr(self):
        """锁定 FIX A 的前置条件：全新 Normal 样式的 rPr 为 None（旧代码据此崩溃）。"""
        assert Document().styles["Normal"].font.element.rPr is None

    def test_renders_without_falling_back(self, tmp_path):
        """带对抗性模板渲染不崩溃、不回退——产出可被重新打开的有效 .docx。"""
        path = _adversarial_template(tmp_path)
        out = render_risk_report(_report(), sample_docx_path=path)
        doc = Document(io.BytesIO(out))
        # 正文与表格被正常渲染（FIX B：内建样式已补齐，未因 KeyError 整篇失败）
        assert len(doc.paragraphs) > 5
        assert len(doc.tables) >= 1

    def test_source_cmc_report_reference_survives_sample_template(self, tmp_path):
        path = _adversarial_template(tmp_path)
        out = render_risk_report(_report(), sample_docx_path=path)
        doc = Document(io.BytesIO(out))
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

        assert text.count("分析文档：CMCReport《HRS-1597 CMCReport.docx》") == 1
        assert text.count("数据依据：该文档经读取与校验后形成的关系图谱数据") == 1

    def test_preserves_template_page_and_header(self, tmp_path):
        path = _adversarial_template(tmp_path)
        out = render_risk_report(_report(), sample_docx_path=path)
        doc = Document(io.BytesIO(out))
        sec = doc.sections[0]
        # 页面布局沿用模板（纵向 + 自定义左边距），而非硬编码横向 A4
        assert round(sec.page_width.cm, 1) == 21.0
        assert round(sec.page_height.cm, 1) == 29.7
        assert round(sec.left_margin.cm, 2) == 3.17
        header_text = "".join(p.text for p in sec.header.paragraphs)
        assert "★模板自有页眉★" in header_text

    def test_does_not_clobber_template_cjk_font(self, tmp_path):
        """FIX A：不得把模板 Normal 的 eastAsia 字体覆盖为宋体。"""
        path = _adversarial_template(tmp_path)
        out = render_risk_report(_report(), sample_docx_path=path)
        doc = Document(io.BytesIO(out))
        rfonts = doc.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts()
        assert rfonts.get(qn("w:eastAsia")) == "仿宋"

    def test_multi_section_layout_preserved(self, tmp_path):
        """#2：双节模板（纵向封面 + 横向正文）渲染后两节布局均须保留。旧 _clear_body 只留
        body 末节 sectPr，会把中间节的页面/页眉页脚全部丢失（两节→一节即回归）。"""
        path = _two_section_template(tmp_path)
        out = render_risk_report(_report(), sample_docx_path=path)
        doc = Document(io.BytesIO(out))
        assert len(doc.sections) >= 2, "中间分节不得被 _clear_body 丢弃"
        assert round(doc.sections[0].page_width.cm, 1) == 21.0, "第 0 节仍须纵向"
        assert round(doc.sections[-1].page_width.cm, 1) == 29.7, "末节仍须横向"

    def test_corrupt_template_falls_back_to_blank(self, tmp_path):
        """坏文件容错：无法打开的模板回退空白文档，报告仍成功生成（横向 A4 硬编码）。"""
        bad = tmp_path / "corrupt.docx"
        bad.write_bytes(b"not a real docx")
        out = render_risk_report(_report(), sample_docx_path=str(bad))
        doc = Document(io.BytesIO(out))
        # 回退分支的硬编码横向 A4
        assert round(doc.sections[0].page_width.cm, 1) == 29.7

    def test_blank_path_locks_landscape_a4_songti_header(self):
        """#10：无模板路径须真正锁定硬编码格式，而非仅「能打开」——否则纵向/换字体/错字号
        /删页眉的回归都会被放行。断言横向 A4 + Normal 宋体 10.5pt + 默认动态页眉。"""
        out = render_risk_report(_report())
        doc = Document(io.BytesIO(out))
        sec = doc.sections[0]
        assert round(sec.page_width.cm, 1) == 29.7  # 横向
        assert round(sec.page_height.cm, 1) == 21.0
        normal = doc.styles["Normal"]
        assert normal.font.name == "宋体"
        assert normal.font.size == Pt(10.5)
        header_text = "".join(p.text for p in sec.header.paragraphs)
        assert "QS-A-020F05" in header_text  # 默认 doc_no 进入动态页眉


class TestBodyFontInheritance:
    """#6（核心需求主路径）：正文主体不得硬编码字号/字体，须继承 Normal——
    套用模板时随模板正文字体走。此前 `_add_inline_runs` 强设 宋体 10.5pt，
    使「完整保留字体、字号」在模板正文的**主要渲染路径**上失效。"""

    def test_inline_runs_do_not_pin_font(self):
        """`_add_inline_runs` 产出的 run 不显式设 size/name（=继承 Normal），但保留加粗语义。"""
        doc = Document()
        p = doc.add_paragraph()
        _add_inline_runs(p, "普通文本 **加粗片段** 收尾", prefix="• ")
        assert p.runs, "应至少产出一个 run"
        for r in p.runs:
            assert r.font.size is None, "不得硬编码字号——须继承模板 Normal"
            assert r.font.name is None, "不得硬编码字体——须继承模板 Normal"
        assert any(r.bold for r in p.runs), "**加粗** 语义应保留"

    def test_markdown_body_inherits_template_normal(self, tmp_path):
        """端到端：对抗性模板（Normal=仿宋）经 `_add_markdown_body` 注入正文后，
        正文 run 无显式字体/字号 → 由模板 Normal（仿宋）继承渲染。"""
        path = _adversarial_template(tmp_path)
        doc = Document(path)
        _clear_body(doc)
        _ensure_builtin_styles(doc)
        _add_markdown_body(doc, "第一段正文内容。\n\n- 列表项一\n- 列表项二")
        body_runs = [r for para in doc.paragraphs for r in para.runs if r.text.strip()]
        assert body_runs, "应渲染出正文 run"
        for r in body_runs:
            assert r.font.size is None and r.font.name is None
        # 模板 Normal 的 eastAsia 仍为仿宋（正文将据此渲染）
        rfonts = doc.styles["Normal"].element.get_or_add_rPr().get_or_add_rFonts()
        assert rfonts.get(qn("w:eastAsia")) == "仿宋"


class TestTableFontNotPinned:
    """#1：评估矩阵/表格正文此前对每个 run 强设 latin 宋体，套用非宋体模板时形成中文随模板、
    西文强制宋体的混合格式。修复后不设 font.name（latin 亦随模板），仅保留紧凑字号 Pt(9)。"""

    def test_matrix_runs_keep_size_but_not_latin_font(self):
        doc = Document()
        _render_assessment_matrix(doc, _report_with_row())
        table = doc.tables[-1]
        runs = [
            r
            for row in table.rows
            for cell in row.cells
            for p in cell.paragraphs
            for r in p.runs
            if r.text.strip()
        ]
        assert runs, "表头 + 数据行应产出 run"
        for r in runs:
            assert r.font.name is None, "表格 run 不得固定 latin 字体——须随模板正文字体"
            assert r.font.size == Pt(9), "紧凑表格字号 Pt(9) 为刻意设计，应保留"


class TestAssessmentMatrixWidth:
    """#7：评估矩阵固定列宽（合计 25.5cm）在纵向 A4 模板下会溢出页面/被裁剪。
    修复后按可用正文宽度等比缩放，使表格总宽 ≤ 可用宽度。"""

    def test_matrix_scales_to_portrait_page(self):
        doc = Document()
        sec = doc.sections[0]
        sec.page_width = Cm(21.0)   # 纵向 A4
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(3.0)  # 可用宽 = 15cm < 固定合计 25.5cm → 必缩放
        _render_assessment_matrix(doc, RiskReport())
        table = doc.tables[-1]
        total = sum(col.width for col in table.columns)
        avail = sec.page_width - sec.left_margin - sec.right_margin
        # 确实按可用宽度缩放（未沿用会溢出纵向页面的固定合计 25.5cm）
        assert total < Cm(25.5)
        # 拟合可用正文宽度内。容差收紧到 twip 级：7 列各以整 twip（dxa）存储，回读舍入使合计
        # 较写入 EMU 至多多出约 3-4 个 twip（≈0.006cm）。Cm(0.01)≈5.5 twip 覆盖此量化余差、
        # 又远小于旧 0.05cm（≈28 twip，可掩盖 ~0.5mm 真实溢出）。
        assert total <= avail + Cm(0.01), "表格总列宽须缩放至可用正文宽度内，避免纵向页面溢出"

    def test_matrix_landscape_not_upscaled(self):
        """横向 A4（可用宽 25.7cm > 合计 25.5cm）：不放大，保留设计列宽。"""
        doc = Document()
        sec = doc.sections[0]
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21.0)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)  # 可用 25.7cm
        _render_assessment_matrix(doc, RiskReport())
        table = doc.tables[-1]
        total = sum(col.width for col in table.columns)
        # 未缩放：合计仍约 25.5cm（容 1mm 余差）
        assert abs(total - Cm(25.5)) < Cm(0.1)

    def test_matrix_scales_to_last_section_not_first(self):
        """Codex R2：矩阵追加在 body 末尾 → 归属末节。多节模板（纵向封面 avail≈16cm + 横向
        正文 avail=25.7cm）下，须按末节（正文）宽度缩放，而非首节（封面）——否则会被错缩到
        ~16cm。断言表格用了正文节宽度（不缩放，≈25.5cm），而非封面节宽度。"""
        doc = Document()
        s0 = doc.sections[0]
        s0.page_width = Cm(21.0)
        s0.page_height = Cm(29.7)  # 纵向封面
        s0.left_margin = Cm(2.5)
        s0.right_margin = Cm(2.5)  # 首节可用 ≈16cm（若误用会把矩阵缩到 ~16cm）
        s1 = doc.add_section(WD_SECTION.NEW_PAGE)
        s1.page_width = Cm(29.7)
        s1.page_height = Cm(21.0)  # 横向正文
        s1.left_margin = Cm(2.0)
        s1.right_margin = Cm(2.0)  # 末节可用 25.7cm > 设计 25.5cm → 不缩放
        _render_assessment_matrix(doc, _report_with_row())
        table = doc.tables[-1]
        total = sum(c.width for c in table.columns)
        assert total > Cm(20), "矩阵须按末节（正文）而非首节（封面）宽度缩放"
        assert abs(total - Cm(25.5)) < Cm(0.1), "末节可用宽 25.7>25.5 → 保留设计列宽不缩放"

    def test_matrix_none_geometry_does_not_crash(self):
        """#8：模板省略 <w:pgSz>/<w:pgMar>（合法，依赖 Word 默认值）时 page_width/margins
        读作 None。缩放须 None-guard 跳过，绝不 `None - None` 抛 TypeError。"""
        doc = Document()
        sectPr = doc.sections[0]._sectPr
        for tag in ("w:pgSz", "w:pgMar"):
            el = sectPr.find(qn(tag))
            if el is not None:
                sectPr.remove(el)
        assert doc.sections[0].page_width is None  # 前置：几何确为 None
        _render_assessment_matrix(doc, _report_with_row())  # 不得抛异常
        assert doc.tables, "None 几何下仍应产出表格（跳过缩放，沿用设计列宽）"

    def test_sectionless_template_falls_back_without_crashing(self, tmp_path):
        """Codex R3：模板整篇省略 <w:sectPr>（合法，Word 用默认版式）时 python-docx 能打开
        但 doc.sections 为空——不仅缩放读 sections[-1] 会崩，python-docx 的 add_table 本身
        （_block_width→sections[-1]）也会 IndexError。render_risk_report 须把此类无几何模板识别
        为不可承载并回退到空白横向 A4，报告仍成功生成、含评估矩阵表格。"""
        tpl = Document()
        body = tpl.element.body
        sectPr = body.find(qn("w:sectPr"))
        if sectPr is not None:
            body.remove(sectPr)
        assert len(tpl.sections) == 0  # 前置：确无节
        path = tmp_path / "sectionless.docx"
        tpl.save(str(path))

        out = render_risk_report(_report_with_row(), sample_docx_path=str(path))
        doc = Document(io.BytesIO(out))
        # 回退到空白横向 A4，未崩溃；评估矩阵表格正常产出
        assert round(doc.sections[0].page_width.cm, 1) == 29.7
        assert doc.tables, "无节模板回退后仍须产出评估矩阵表格"

    def test_matrix_sectionless_standalone_does_not_crash(self):
        """Codex R3 纵深防御闭环：即便绕过公共入口、直接把无节文档传给 _render_assessment_matrix，
        也须自补默认 <w:sectPr> 令 add_table 有末节可测量（否则 _block_width→sections[-1] IndexError），
        几何读作 None → 跳过缩放、沿用设计列宽。断言不崩且产出表格。"""
        doc = Document()
        body = doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        if sectPr is not None:
            body.remove(sectPr)
        assert len(doc.sections) == 0  # 前置：确无节
        _render_assessment_matrix(doc, _report_with_row())  # 不得 IndexError
        assert doc.sections, "helper 须自补 <w:sectPr> 令 add_table 可测量"
        assert doc.tables, "无节直调下仍应产出表格（跳过缩放，沿用设计列宽）"


class TestBuiltinStyleIdCollision:
    """#3：模板已含目标 styleId 但样式 name 被改（LibreOffice/精简导出常见），
    仅按 name 判定「缺失」并原样 append 会产生重复 styleId → 破坏 styles part；而若因
    styleId 冲突直接跳过，规范 name 仍缺失 → 随后 add_heading 抛 KeyError。修复后以新
    styleId 插入携带规范 name 的样式：既不重复 styleId，又让按 name 的查找可解析。"""

    def test_no_duplicate_styleid_on_name_collision(self):
        doc = Document()
        h1 = doc.styles["Heading 1"]
        assert h1.style_id == "Heading1"
        h1.name = "我的标题一"  # name 改掉 → 按 name 看「Heading 1」缺失，但 styleId 仍在
        _ensure_builtin_styles(doc)
        ids = [s.style_id for s in doc.styles]
        assert ids.count("Heading1") == 1, "不得因 name 不匹配而追加重复 styleId"

    def test_add_heading_works_after_name_collision(self):
        """#3 核心：name 冲突修复后，按规范 name 解析的 add_heading(level=1/2/3) 须成功。"""
        doc = Document()
        doc.styles["Heading 1"].name = "我的标题一"
        doc.styles["Heading 2"].name = "我的标题二"
        _ensure_builtin_styles(doc)
        # 规范 name 重新可解析，add_heading 不再 KeyError
        assert doc.styles["Heading 1"] is not None
        doc.add_heading("章节一", level=1)
        doc.add_heading("章节二", level=2)
        doc.add_heading("章节三", level=3)
        # styleId 无重复（styles part 未损坏）
        ids = [s.style_id for s in doc.styles]
        assert len(ids) == len(set(ids)), "补齐样式后不得出现重复 styleId"
