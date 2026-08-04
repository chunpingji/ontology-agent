"""PDE/OEB conflict snapshots must survive generation, persistence payloads, and DOCX."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.services.reporting.ast_template import load_default_template
from app.services.reporting.docx_renderer import render_risk_report
from app.services.reporting.risk_report_generator import (
    PENDING_LEVEL,
    PENDING_STATUS,
    RiskReportGenerator,
)


def _conflict() -> dict:
    return {
        "conflict_key": "shared_line_pde",
        "asserted": {
            "pde_mg_day": 1.8,
            "pde_ug_day": 1800.0,
            "band": 2,
        },
        "derived": {
            "band": 5,
            "band_point": 4,
            "pde_ug_day": 50.0,
            "oel_ug_m3": 5.0,
            "provisional": True,
            "input_source": "extracted",
            "provenance": {
                "method": {
                    "id": "ADE-OEB/test",
                    "version": "1.0",
                    "regulation_ref": "EMA test reference",
                },
                "formula": "PDE = NOAEL·BW / (F1·F2·F3·F4·F5)",
                "inputs": {
                    "noael_mg_kg_day": 0.5,
                    "species": "rat",
                    "study_duration_days": 28,
                },
                "factors": {"F1_interspecies": 5.0, "composite_UF": 500.0},
            },
        },
        "delta_bands": 3,
        "pde_ratio": 36.0,
        "summary": "推导 OEB band 5 与原文 PDE 1.8 mg/日（band 2）不一致。",
    }


def _edge() -> dict:
    return {
        "subject_class_iri": "https://example.test/CMCReport",
        "predicate_iri": "https://example.test/hasSharedLineData",
        "object_class_iri": "https://example.test/SharedLineAssessmentData",
        "object_text": "共线评估",
        "object_data_properties": [],
        "source_ref": "§ 共线评估",
        "conflict": _conflict(),
    }


def _all_docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs]
    cells = [
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *cells])


class TestPdeConflictGeneration:
    def test_pending_conflict_adds_pending_row_and_blocks_acceptable_conclusion(self, db):
        report, _ = RiskReportGenerator(db).generate_with_coverage([_edge()])

        assert len(report.pde_conflicts) == 1
        snapshot = report.pde_conflicts[0]
        assert snapshot["asserted"]["pde_mg_day"] == 1.8
        assert snapshot["derived"]["pde_ug_day"] == 50.0
        assert snapshot["decision"]["chosen"] == "pending"
        assert snapshot["effective"] is None

        row = next(r for r in report.assessment_rows if r.hazid == "PDE/OEB 潜能等级冲突")
        assert row.post_control_level == PENDING_LEVEL
        assert row.status == PENDING_STATUS
        assert "总体风险结论为待评估" in report.conclusion
        assert "不得认定全部风险可以接受" in report.conclusion

    @pytest.mark.parametrize(
        ("chosen", "source", "source_label", "pde_mg", "band"),
        [
            ("derived", "derived", "推导", 0.05, 5),
            ("asserted", "asserted", "原文", 1.8, 2),
        ],
    )
    def test_decision_selects_effective_value_without_overwriting_either_side(
        self, db, chosen, source, source_label, pde_mg, band,
    ):
        decision = {
            "chosen": chosen,
            "actor": "qa.user",
            "decided_at": "2026-08-02T10:30:00+00:00",
            "note": "已核对毒理试验原始记录",
            "version": 3,
        }
        report, _ = RiskReportGenerator(db).generate_with_coverage(
            [_edge()], pde_conflict_decision=decision,
        )

        snapshot = report.pde_conflicts[0]
        assert snapshot["asserted"]["pde_mg_day"] == 1.8
        assert snapshot["derived"]["pde_ug_day"] == 50.0
        assert snapshot["effective"] == {
            "source": source,
            "source_label": source_label,
            "pde_mg_day": pde_mg,
            "pde_ug_day": 50.0 if chosen == "derived" else 1800.0,
            "band": band,
        }
        assert snapshot["decision"]["actor"] == "qa.user"
        assert snapshot["decision"]["note"] == "已核对毒理试验原始记录"
        row = next(r for r in report.assessment_rows if r.hazid == "PDE/OEB 潜能等级冲突")
        assert row.status == f"已裁决（采纳{source_label}）"
        assert f"采用{source_label}数据" in report.conclusion

    def test_duplicate_edges_do_not_duplicate_document_conflict(self, db):
        report, _ = RiskReportGenerator(db).generate_with_coverage([_edge(), _edge()])
        assert len(report.pde_conflicts) == 1
        assert sum(
            row.hazid == "PDE/OEB 潜能等级冲突" for row in report.assessment_rows
        ) == 1

    def test_pending_conflict_is_included_in_rule_progress_result(self, db):
        snapshots: list[list[dict[str, str | None]]] = []
        RiskReportGenerator(db).generate_with_coverage(
            [_edge()], assessment_progress_fn=snapshots.append,
        )
        rules = next(step for step in snapshots[-1] if step["key"] == "rules")
        assert rules["status"] == "completed"
        assert rules["result"] == "1 项 · 1 项待评估"


class TestPdeConflictDocx:
    @pytest.mark.parametrize("mode", ["legacy", "template", "sample"])
    def test_pending_conflict_is_rendered_in_every_output_path(self, db, tmp_path, mode):
        report, manifest = RiskReportGenerator(db).generate_with_coverage([_edge()])
        kwargs = {}
        if mode == "template":
            kwargs["template"] = load_default_template()
        elif mode == "sample":
            sample = tmp_path / "sample.docx"
            Document().save(sample)
            kwargs["sample_docx_path"] = str(sample)

        text = _all_docx_text(render_risk_report(report, manifest, **kwargs))

        assert "PDE/OEB 潜能等级冲突及裁决" in text
        assert "PDE 1.8 mg/日（1800.0 µg/日）；OEB band 2" in text
        assert "PDE 0.05 mg/日（50.0 µg/日）；OEB band 5" in text
        assert "待人工裁决" in text
        assert "不得视为全部风险可以接受" in text
        assert "PDE = NOAEL·BW / (F1·F2·F3·F4·F5)" in text

    def test_decided_conflict_renders_effective_value_and_audit_fields(self, db):
        report, manifest = RiskReportGenerator(db).generate_with_coverage(
            [_edge()],
            pde_conflict_decision={
                "chosen": "derived",
                "actor": "qa.user",
                "decided_at": "2026-08-02T10:30:00+00:00",
                "note": "已核对毒理试验原始记录",
                "version": 3,
            },
        )

        text = _all_docx_text(render_risk_report(report, manifest))

        assert "已裁决（采纳推导值）" in text
        assert "本报告有效值" in text
        assert "PDE 0.05 mg/日（50.0 µg/日）；OEB band 5" in text
        assert "qa.user" in text
        assert "2026-08-02T10:30:00+00:00" in text
        assert "已核对毒理试验原始记录" in text
