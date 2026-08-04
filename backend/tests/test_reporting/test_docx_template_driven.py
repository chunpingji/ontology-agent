"""016+: template-driven docx rendering — the report follows the authored AST
template's Section / Group / Slot structure instead of the fixed QS-A-020F05 skeleton.

Verifies:
  - Section titles → level-2 headings, Group titles → level-3, in template order;
  - a semantic slot's LLM text renders in place, and a Markdown table in it becomes a
    real Word table (the 评估小组 member-roster path);
  - deterministic value slots render ``标签：值`` / missing-required warn placeholder
    from the manifest;
  - equipment_table / assessment_table groups render the deterministic tables;
  - passing no template still renders the legacy skeleton (backward-compatible).
"""

from __future__ import annotations

import io

from docx import Document

from app.services.reporting.ast_template import ReportTemplate
from app.services.reporting.coverage_validator import (
    CoverageManifest,
    SlotCoverage,
)
from app.services.reporting.docx_renderer import _render_equipment_tables, render_risk_report
from app.services.reporting.risk_report_generator import (
    EquipmentEntry,
    RiskReport,
    RiskRow,
)

_TEAM_SLOT_ID = "grp_team.members"


def test_aps_equipment_warning_uses_icon_and_red_text_without_equipment_rows():
    doc = Document()
    report = RiskReport(
        equipment_notes=["APS排期冲突：候选设备 PF64216、PF64616 均无可用排期。"]
    )

    _render_equipment_tables(doc, report)

    warning_run = next(
        run
        for paragraph in doc.paragraphs
        for run in paragraph.runs
        if "APS排期冲突" in run.text
    )
    assert warning_run.text.startswith("⚠ APS排期冲突")
    assert warning_run.bold is True
    assert str(warning_run.font.color.rgb) == "C00000"

_TEMPLATE_DICT = {
    "template_id": "tpl-test",
    "doc_no": "QS-A-020F05",
    "revision": "00",
    "sections": [
        {
            "section_id": "sec-team",
            "title": "评估小组 Assessment Team",
            "groups": [
                {
                    "group_id": "grp_team",
                    "title": "评估小组成员 Members",
                    "kind": "fields",
                    "slots": [
                        {
                            "slot_id": _TEAM_SLOT_ID,
                            "label": "评估小组成员",
                            "source": {"kind": "semantic", "coverage_refs": []},
                        }
                    ],
                }
            ],
        },
        {
            "section_id": "sec-body",
            "title": "SECTIONⅠ 主体",
            "groups": [
                {
                    "group_id": "grp_subject",
                    "title": "1. 风险评估对象",
                    "kind": "fields",
                    "slots": [
                        {
                            "slot_id": "subject.name",
                            "label": "产品名称",
                            "source": {
                                "kind": "extraction",
                                "object_class_iri_contains": "DrugProduct",
                                "text": True,
                            },
                            "required": True,
                        },
                        {
                            "slot_id": "subject.pde",
                            "label": "PDE",
                            "source": {
                                "kind": "extraction",
                                "object_class_iri_contains": "DrugProduct",
                                "label": "PDE",
                            },
                            "required": True,
                        },
                    ],
                },
                {
                    "group_id": "grp_equipment",
                    "title": "2. 设备一览表",
                    "kind": "equipment_table",
                    "repeat": {"by": "workshop"},
                    "slots": [
                        {
                            "slot_id": "equipment.id",
                            "label": "设备编号",
                            "source": {
                                "kind": "extraction",
                                "object_class_iri_contains": "Equipment",
                                "text": True,
                            },
                            "required": True,
                        }
                    ],
                },
                {
                    "group_id": "grp_assessment",
                    "title": "3. 风险评估",
                    "kind": "assessment_table",
                    "repeat": {"by": "rule", "rule_group": "risk_assessment"},
                    "slots": [
                        {
                            "slot_id": "assessment.hazid",
                            "label": "HazID",
                            "source": {"kind": "rule", "field": "hazid"},
                            "required": True,
                        }
                    ],
                },
            ],
        },
    ],
}


def _template() -> ReportTemplate:
    return ReportTemplate.model_validate(_TEMPLATE_DICT)


def _report() -> RiskReport:
    return RiskReport(
        equipment_tables={
            "642车间": [EquipmentEntry(1, "RE001", "反应釜", "500L", "316L")]
        },
        assessment_rows=[
            RiskRow("人员-01", "共线操作", "高", "低", "清洁验证", "SOP-001", "可以接受")
        ],
        semantic_slots=[
            {
                "slot_id": _TEAM_SLOT_ID,
                "section_id": "sec-team",
                "label": "评估小组成员",
                "text": (
                    "评估小组成员如下：\n\n"
                    "| 姓名 | 角色 | 部门 | 签名 |\n"
                    "|---|---|---|---|\n"
                    "| 王玉华 | QA | 质量部 |  |\n"
                    "| 陈志强 | EHS评估 | EHS |  |\n"
                ),
            }
        ],
    )


def _manifest() -> CoverageManifest:
    return CoverageManifest(
        template_id="tpl-test",
        slots=[
            SlotCoverage(
                slot_id="subject.name",
                label="产品名称",
                status="filled",
                source_kind="extraction",
                value="HRS-1234",
            ),
            SlotCoverage(
                slot_id="subject.pde",
                label="PDE",
                status="missing_required",
                source_kind="extraction",
            ),
        ],
    )


def _render() -> Document:
    docx_bytes = render_risk_report(_report(), _manifest(), template=_template())
    return Document(io.BytesIO(docx_bytes))


class TestTemplateDrivenStructure:
    def test_section_and_group_headings_in_template_order(self):
        doc = _render()
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        # sections and groups both present …
        assert "评估小组 Assessment Team" in headings
        assert "SECTIONⅠ 主体" in headings
        assert "1. 风险评估对象" in headings
        assert "2. 设备一览表" in headings
        assert "3. 风险评估" in headings
        # … and in template order (team section before body section)
        assert headings.index("评估小组 Assessment Team") < headings.index("SECTIONⅠ 主体")

    def test_no_fixed_skeleton_headings(self):
        """The legacy hardcoded headings must NOT appear when template-driven."""
        doc = _render()
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "SECTION II  风险回顾" not in all_text  # legacy _add_section_two
        assert "语义化插槽 Semantic Slots" not in all_text  # legacy flat appendix


class TestSemanticSlotTable:
    def test_markdown_table_becomes_word_table(self):
        doc = _render()
        # the team roster Markdown table is rendered as a real Word table
        assert doc.tables, "expected at least one Word table"
        all_cells = {
            c.text for t in doc.tables for row in t.rows for c in row.cells
        }
        assert "王玉华" in all_cells
        assert "陈志强" in all_cells
        assert "QA" in all_cells

    def test_semantic_prose_is_llm_annotated(self):
        doc = _render()
        runs = "".join(r.text for p in doc.paragraphs for r in p.runs)
        assert "ⓘ" in runs  # semantic-slot prose carries the LLM marker
        assert "评估小组成员如下" in runs


class TestDeterministicSlots:
    def test_value_slot_renders_label_and_value(self):
        doc = _render()
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "产品名称：HRS-1234" in all_text

    def test_missing_required_slot_flagged(self):
        doc = _render()
        all_text = "\n".join(p.text for p in doc.paragraphs)
        # the required-but-missing PDE slot surfaces its warn placeholder, not silence
        assert "⚠" in all_text

    def test_equipment_and_assessment_tables_present(self):
        doc = _render()
        all_cells = {
            c.text for t in doc.tables for row in t.rows for c in row.cells
        }
        assert "RE001" in all_cells  # equipment_table
        assert any("人员-01" in c for c in all_cells)  # assessment_table


class TestBackwardCompatibility:
    def test_no_template_uses_legacy_skeleton(self):
        docx_bytes = render_risk_report(_report())
        doc = Document(io.BytesIO(docx_bytes))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        # legacy fixed skeleton still emitted when no template is passed
        assert "SECTION I  风险评估" in all_text
