import io
import uuid

from docx import Document

from app.models.evidence import EvidenceCoverage
from app.models.extraction import AstTemplate, ExtractionJob
from app.services.extraction.candidate_store import CandidateStore
from app.services.fact_commit import FactCommitService
from app.services.ontology_instance_writer import EvidenceInstanceWriter
from app.services.reporting.snapshot_report import (
    build_report_inputs,
    freeze_report_inputs,
    render_snapshot_report,
)
from tests.test_extraction.test_fact_commit import assertions


def setup_report(db, real_world, tmp_path, *, positive=False):
    schema = {
        "urn:test:Drug": {
            "properties": [{"iri": "urn:test:strength", "label": "规格"}],
            "relationships": [{"iri": "urn:test:uses", "range": ["urn:test:Equipment"]}],
        },
        "urn:test:Equipment": {"properties": []},
    }
    template = AstTemplate(
        id=uuid.uuid4(),
        name="Snapshot",
        version="v1",
        status="published",
        iri_pattern="urn:test:Drug",
        schema_json={
            "template_id": "test",
            "sections": [
                {
                    "section_id": "s",
                    "title": "基本信息",
                    "groups": [],
                    "coverage": [
                        {
                            "kind": "ontology_relation",
                            "doc_class_iri": "urn:test:Drug",
                            "predicate_iri": "urn:test:uses",
                            "range_class_iri": "urn:test:Equipment",
                        }
                    ],
                }
            ],
        },
    )
    job = ExtractionJob(
        id=uuid.uuid4(),
        source_type="word",
        status="completed",
        source_filename="source.docx",
        source_config={"doc_class_iri": "urn:test:Drug", "template_id": str(template.id)},
    )
    db.add_all([template, job])
    db.commit()
    store = CandidateStore(db)
    values = assertions()
    if positive:
        values[-1] = values[-1].model_copy(update={"assertion_status": "affirmed"})
    candidates = store.persist_validated(job.id, values, actor="fixture")
    for c in candidates:
        store.review(c.candidate_id, c.revision, "confirmed", "fixture verification", "analyst")
    service = FactCommitService(db, EvidenceInstanceWriter(real_world, tmp_path / "worlds"))
    commit = service.request(
        job.id,
        "initial",
        [{"candidate_id": c.candidate_id, "revision": c.revision} for c in candidates],
        "analyst",
    )
    assert service.apply(commit.id).status == "succeeded"
    return job, template, schema


def test_report_inputs_freeze_template_discovery_and_actual_snapshot(db, real_world, tmp_path):
    job, template, schema = setup_report(db, real_world, tmp_path)
    inputs = build_report_inputs(db, job, schema=schema)
    frozen = freeze_report_inputs(db, job.id, inputs)
    db.commit()
    old_id = frozen.id
    template.version = "v2"
    template.schema_json = {**template.schema_json, "revision": "v2"}
    db.commit()
    assert db.get(EvidenceCoverage, old_id).payload["template_version"] == "v1"
    assert build_report_inputs(db, job, schema=schema)["manifest_id"] != old_id
    report, manifest, data = render_snapshot_report(db.get(EvidenceCoverage, old_id).payload)
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "基本信息" in text and "待" in text
    assert report.evidence_snapshot_id == inputs["snapshot_id"]
    assert manifest.snapshot_id == inputs["snapshot_id"]


def test_snapshot_report_never_calls_legacy_enrichment(db, real_world, tmp_path, monkeypatch):
    job, _, schema = setup_report(db, real_world, tmp_path)
    from app.services.reporting.risk_report_generator import RiskReportGenerator

    def forbidden(*args, **kwargs):
        raise AssertionError("legacy enrichment is forbidden")

    monkeypatch.setattr(RiskReportGenerator, "_enriched_edges", forbidden)
    monkeypatch.setattr(RiskReportGenerator, "_build_equipment_tables", forbidden)
    report, _, data = render_snapshot_report(build_report_inputs(db, job, schema=schema))
    assert data.startswith(b"PK")
    assert not report.team_members and not report.approvers
    assert report.conclusion != "可以接受"


def test_explicit_snapshot_slot_uses_exact_path_not_display_label(db, real_world, tmp_path):
    job, template, schema = setup_report(db, real_world, tmp_path, positive=True)
    payload = template.schema_json
    payload["sections"][0]["groups"] = [
        {
            "group_id": "g",
            "title": "设备",
            "kind": "fields",
            "slots": [
                {
                    "slot_id": "equipment",
                    "label": "任意组合显示名称",
                    "source": {
                        "kind": "snapshot",
                        "root_class_iri": "urn:test:Drug",
                        "predicate_path": [{"predicate_iri": "urn:test:uses"}],
                        "range_class_iri": "urn:test:Equipment",
                        "text": True,
                    },
                }
            ],
        }
    ]
    template.schema_json = dict(payload)
    db.commit()
    report, _, _ = render_snapshot_report(build_report_inputs(db, job, schema=schema))
    value = report.semantic_slots[0]
    assert "待补充" not in value["text"]
    assert value["assertion_ids"]


def test_sample_header_footer_signatures_never_leak_into_snapshot_report(db, real_world, tmp_path):
    from zipfile import ZipFile

    job, template, schema = setup_report(db, real_world, tmp_path)
    sample = Document()
    section = sample.sections[0]
    for part in (
        section.header,
        section.footer,
        section.first_page_header,
        section.first_page_footer,
        section.even_page_header,
        section.even_page_footer,
    ):
        part.paragraphs[0].text = "样例签署人-已批准"
    sample_path = tmp_path / "sample.docx"
    sample.save(sample_path)
    original_bytes = sample_path.read_bytes()
    template.sample_docx_path = str(sample_path)
    db.commit()
    _, _, content = render_snapshot_report(build_report_inputs(db, job, schema=schema))
    with ZipFile(io.BytesIO(content)) as archive:
        assert all(
            "样例签署人" not in archive.read(name).decode()
            for name in archive.namelist()
            if name.endswith(".xml") and ("header" in name or "footer" in name)
        )
    assert sample_path.read_bytes() == original_bytes


def test_snapshot_rules_do_not_use_substring_classes_or_round_decimals():
    from app.services.reasoning.interpreter import TRUE, UNKNOWN, Facts, evaluate

    facts = Facts(
        strict_evidence=True,
        drug_classes=["urn:NonSterileDrug"],
        scalars={"dose": "0.123456789012345678901"},
        data_values={"flag": "false"},
    )
    assert evaluate({"op": "class_present", "class": "SterileDrug"}, facts) is UNKNOWN
    assert (
        evaluate({"op": "boolean_has_value", "property": "flag", "value": True}, facts) is UNKNOWN
    )
    assert (
        evaluate(
            {"op": "literal_cmp", "key": "dose", "cmp": "gt", "value": "0.123456789012345678900"},
            facts,
        )
        is TRUE
    )
