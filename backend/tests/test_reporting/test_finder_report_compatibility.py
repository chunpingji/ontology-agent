"""Values must come from selected Finder objects and explicit local Mock records."""

from types import SimpleNamespace

import pytest

from app.services.reporting.demo_sources import (
    PROMPT_REF,
    builtin_contract,
    contract_ref,
    load_mock,
)
from app.services.reporting.output_renderer import render_docx, render_snapshot
from app.services.reporting.report_snapshot import resolve_snapshot
from app.services.reporting.slot_semantics import SlotChoice, apply_patch, build_slot, catalog
from app.services.reporting.template_compiler import compile_template, require_valid
from app.services.reporting.template_v2 import RecordSource, ReportingError, TemplateV2
from tests.test_reporting.test_output_contracts import contracts, ontology, template


def finder_source():
    def row(code, spec):
        return {
            "predicate_iri": "urn:uses",
            "object_class_iri": "urn:Equipment",
            "object_text": "same visible name",
            "object_data_properties": [
                {"iri": "urn:code", "value": code, "source": {"anchors": [{"block_id": code}]}},
                {"iri": "urn:spec", "value": spec, "source": {"anchors": [{"block_id": spec}]}},
            ],
            "sub_relationships": [],
        }

    return {
        "kind": "finder_demo",
        "state": "ready",
        "execution_id": "run-1",
        "root_entity_id": "root",
        "root_class_iri": "urn:Report",
        "job_id": "source",
        "relationships": [row("EQ-A", "spec-A"), row("EQ-B", "spec-B")],
    }


def resolve_finder(value=None, source=None, extra_contracts=None, records=None):
    cs = {**contracts(), **(extra_contracts or {})}
    plan = require_valid(compile_template(value or template(), ontology(), cs))
    return resolve_snapshot(
        plan,
        {
            "template": plan["template"],
            "schema": ontology(),
            "contracts": cs,
            "sources": {"doc": source or finder_source()},
            "records": records or {},
            "demonstration": True,
            "preview_mode": "report",
            "source_bundle_id": "demo",
        },
    )


def test_finder_rows_keep_local_identity_fields_sources_and_docx():
    result = resolve_finder()
    rows = result["inputs"]["rows"]["items"]
    assert result["material_status"] == "ready"
    assert [(r["fields"]["code"]["value"], r["fields"]["spec"]["value"]) for r in rows] == [
        ("EQ-A", "spec-A"),
        ("EQ-B", "spec-B"),
    ]
    assert rows[0]["entity_id"] != rows[1]["entity_id"]
    assert rows[0]["fields"]["code"]["fact_refs"] == []
    provenance = rows[0]["fields"]["code"]["provenance_refs"][0]
    assert provenance["kind"] == "finder_demo" and provenance["verified_fact"] is False
    assert provenance["source"]["anchors"][0]["block_id"] == "EQ-A"
    rendered = render_snapshot(result)
    assert "EQ-A" in str(rendered["body_ast"]) and "演示草稿" in str(rendered["body_ast"])
    from io import BytesIO

    from docx import Document

    document = Document(BytesIO(render_docx(rendered["body_ast"])))
    assert document.tables[0].cell(1, 0).text == "EQ-A"
    assert document.tables[0].cell(2, 0).text == "EQ-B"


def test_missing_conflicting_and_empty_finder_values_are_not_ready():
    source = finder_source()
    source["relationships"][0]["object_data_properties"] = source["relationships"][0][
        "object_data_properties"
    ][1:]
    assert resolve_finder(source=source)["material_status"] == "incomplete"
    source = finder_source()
    source["relationships"][0]["object_data_properties"].append(
        {"iri": "urn:code", "value": "different"}
    )
    assert resolve_finder(source=source)["material_status"] == "conflict"
    source["relationships"] = []
    result = resolve_finder(source=source)
    assert result["material_status"] == "incomplete"
    assert all(not r["satisfied"] for r in result["coverage"] if r["required"])


@pytest.mark.parametrize("state,marker", [("missing", "待补充"), ("conflict", "数据冲突")])
def test_assisted_finder_draft_keeps_gaps_and_generates_docx(state, marker):
    from copy import deepcopy
    from io import BytesIO

    from docx import Document

    from app.services.reporting.output_ast import plain_text

    value, source = template(), finder_source()
    source["relationships"] = source["relationships"][:1]
    properties = source["relationships"][0]["object_data_properties"]
    if state == "missing":
        properties.pop()
    else:
        properties.append({"iri": "urn:spec", "value": "unreviewed-alternative"})
    value["definitions"]["inputs"]["rows"]["projection"]["fields"]["spec"]["required"] = True
    unit = value["sections"][0]["groups"][0]["units"][0]
    unit["render"] = {"kind": "narrative", "mode": "assisted", "prompt": {
        "policy_ref": PROMPT_REF, "input_refs": [{"input_id": "rows"}],
        "required_refs": [{"input_id": "rows"}],
    }}
    snapshot = resolve_finder(value, source, {PROMPT_REF: builtin_contract(PROMPT_REF)})
    original = deepcopy(snapshot)
    calls = []

    def provider(_system, payload, _policy, _budget):
        calls.append(payload)
        missing = payload["inputs"][0]["value"]["items"][0]["fields"]["spec"]
        expected = {"status": state, "value": "（待补充）"}
        if state == "conflict":
            expected["issues"] = ["FACT_CONFLICT"]
        assert missing == expected
        assert "unreviewed-alternative" not in str(payload)
        return {"nodes": [{"kind": "input_ref", "input_id": "rows"}]}

    result = render_snapshot(snapshot, provider=provider)
    assert result["execution_status"] == "completed"
    assert result["material_status"] == ("incomplete" if state == "missing" else "conflict")
    assert len(calls) == 1 and snapshot == original
    assert marker in plain_text(result["body_ast"]) and "EQ-A" in plain_text(result["body_ast"])
    document = Document(BytesIO(render_docx(result["body_ast"])))
    assert any(marker in p.text for p in document.paragraphs)
    assert "演示草稿" in document.paragraphs[0].text
    assert result["output_results"][0]["citations"][0]["provenance_refs"]

    # The same exception in the strict reporting path remains a failure.
    snapshot["source_bundle"]["demonstration"] = False
    strict = render_snapshot(snapshot, provider=lambda *_: pytest.fail("must not call model"))
    assert strict["execution_status"] == "failed"
    assert "INPUT_CONSUMPTION_BLOCKED" in str(strict) or "OBJECT_UNIVERSE_OPEN" in str(strict)

    snapshot["source_bundle"]["demonstration"] = True
    failed = render_snapshot(snapshot, provider=lambda *_: {"nodes": []})
    assert failed["execution_status"] == "failed"
    assert "OUTPUT_REFERENCE_INVALID" in str(failed)


def test_presence_does_not_establish_full_document_coverage_or_arbitrary_subject():
    value = template()
    value["definitions"]["bindings"]["equipment"]["scope"]["require_complete_set"] = True
    result = resolve_finder(value)
    assert result["inputs"]["rows"]["discovery"]["status"] == "open"
    value["definitions"]["bindings"]["equipment"]["scope"]["root"] = {
        "kind": "entity_ref",
        "entity_id": "other-object",
    }
    assert resolve_finder(value)["material_status"] == "incomplete"


def test_single_value_cannot_choose_first_matching_object():
    value = template()
    value["definitions"]["inputs"]["rows"]["projection"] = {
        "kind": "property",
        "property_iri": "urn:code",
    }
    value["sections"][0]["groups"][0]["units"][0]["render"] = {
        "kind": "narrative",
        "nodes": [{"kind": "input_ref", "input_id": "rows"}],
    }
    result = resolve_finder(value)
    assert result["inputs"]["rows"]["state"] == "conflict"
    assert result["inputs"]["rows"]["value"] is None


def empty_slot_template():
    value = template()
    value["definitions"] = {"bindings": {}, "inputs": {}}
    value["sections"][0]["groups"][0]["units"][0].update(
        bindings=[], inputs=[], render={"kind": "narrative", "nodes": []}
    )
    return TemplateV2.model_validate(value)


def test_empty_slot_is_a_configuration_gap_not_success():
    result = resolve_finder(empty_slot_template())
    assert result["material_status"] == "incomplete"
    assert result["coverage"][0]["satisfied"] is False
    assert result["blocking_issues"][0]["code"] == "SLOT_CONFIGURATION_MISSING"
    assert "待配置" in str(render_snapshot(result)["body_ast"])


def test_slot_conversion_creates_real_authorized_ai_input():
    value = empty_slot_template()
    source = finder_source()
    source["relationships"] = source["relationships"][:1]
    options = catalog(value, ontology(), {"doc": source})
    option = next(o for o in options if o["kind"] == "graph" and len(o["predicate_path"]) == 1)
    selected = next(
        f for f in option["fields"] if f["projection"].get("property_iri") == "urn:code"
    )
    unit = value.sections[0].groups[0].units[0]
    choice = SlotChoice(
        output_id=unit.output_id,
        sources=[{"source_key": option["key"], "fields": [selected["key"]]}],
        instructions="说明设备编号",
    )
    converted = apply_patch(value, build_slot(value, unit, choice, options))
    result = resolve_finder(converted, source, {PROMPT_REF: builtin_contract(PROMPT_REF)})
    calls = []

    def provider(system, payload, policy, budget):
        calls.append(payload)
        assert payload["inputs"][0]["value"] == "EQ-A"
        return {
            "nodes": [
                {"kind": "text", "text": "原文记载："},
                {"kind": "input_ref", **payload["required_refs"][0]},
            ]
        }

    rendered = render_snapshot(result, provider=provider)
    assert rendered["execution_status"] == "completed"
    assert len(calls) == 1 and calls[0]["instructions"] == "说明设备编号"
    assert "EQ-A" in str(rendered["body_ast"])
    assert "same visible name" not in str(calls)
    invalid = SlotChoice(
        output_id="table", sources=[{"source_key": "invented", "fields": ["fake"]}]
    )
    with pytest.raises(ReportingError):
        build_slot(value, unit, invalid, options)


def test_mock_refresh_uses_current_values_and_keeps_prior_input(db):
    from app.models.mock_data import MockTeamMember

    db.add(
        MockTeamMember(
            team_type="assessment",
            name="Original",
            role_label="工程师",
            department="技术",
            role_class_iri="urn:Engineer",
            role_code="engineering",
        )
    )
    db.commit()
    config = RecordSource(
        provider="assessment_team",
        contract_ref=contract_ref("assessment_team"),
        filters={"role_code": "engineering"},
    )
    before = load_mock(db, config, 100)
    assert before["kind"] == "mock" and before["provenance"]["reviewed"] is False
    value = empty_slot_template()
    options = catalog(value, ontology())
    choice = SlotChoice(
        output_id="table",
        render="table",
        sources=[{"source_key": "mock:assessment_team", "fields": ["name", "role", "department"]}],
    )
    converted = apply_patch(
        value, build_slot(value, value.sections[0].groups[0].units[0], choice, options)
    )
    slot = next(iter(converted.record_sources))
    refs = {config.contract_ref: builtin_contract(config.contract_ref)}
    first = resolve_finder(converted, extra_contracts=refs, records={slot: before})
    assert first["material_status"] == "ready"
    row = db.query(MockTeamMember).one()
    row.name = "Updated"
    db.commit()
    after = load_mock(db, config, 100)
    assert before["record_id"] != after["record_id"]
    assert "Original" in str(render_snapshot(first)["body_ast"])
    refreshed = resolve_finder(converted, extra_contracts=refs, records={slot: after})
    assert "Updated" in str(render_snapshot(refreshed)["body_ast"])
    config.filters["role_code"] = "does-not-exist"
    assert (
        resolve_finder(converted, extra_contracts=refs, records={slot: load_mock(db, config, 100)})[
            "material_status"
        ]
        == "incomplete"
    )


def test_old_hashes_do_not_gain_empty_compatibility_fields():
    value = TemplateV2.model_validate(template()).model_dump(mode="json")
    assert "record_sources" not in value
    assert "source_field" not in value["definitions"]["inputs"]["rows"]["projection"]


def test_mock_equipment_reads_legacy_columns_without_inventing_missing_properties(db):
    from app.models.mock_data import MockEquipment

    properties = [
        {"iri": None, "label": "主体材质", "value": "不锈钢"},
        {"iri": "urn:unrelated", "label": "安装位置", "value": "不能按标签误取"},
    ]
    db.add(MockEquipment(
        equipment_id="PF64203", iri="urn:mock:PF64203", label="微压柱系统",
        equipment_class_iri="urn:Equipment", workshop_code="642", data_properties=properties,
    ))
    db.flush()
    config = RecordSource(provider="equipment", contract_ref=contract_ref("equipment"))
    record = load_mock(db, config, 100)
    assert record["values"][0]["material"] == "不锈钢"
    assert record["values"][0]["location"] is None
    assert record["values"][0]["specification"] is None
    assert record["provenance"]["reviewed"] is False
    assert properties[0]["iri"] is None


def test_mock_equipment_preserves_conflicting_legacy_and_iri_values(db):
    from app.models.mock_data import MockEquipment

    db.add(MockEquipment(
        equipment_id="PF64603", iri="urn:mock:PF64603", label="微压柱系统",
        equipment_class_iri="urn:Equipment", workshop_code="646", data_properties=[
            {"iri": None, "label": "主体材质", "value": "不锈钢"},
            {"iri": "https://ontology.pharma-gmp.cn/slpra/equipment/constructedOf",
             "label": "主体材质", "value": "玻璃"},
        ],
    ))
    db.flush()
    config = RecordSource(provider="equipment", contract_ref=contract_ref("equipment"))
    record = load_mock(db, config, 100)
    assert record["values"][0]["material"] == ["不锈钢", "玻璃"]


def test_demo_cannot_create_formal_signing_content():
    from app.services.reporting.report_signing import ReportSigning

    service = ReportSigning(None)
    service.request = lambda *args: (None, None, None)
    service.runs = SimpleNamespace(
        get=lambda _: SimpleNamespace(source_bundle={"demonstration": True})
    )
    with pytest.raises(ReportingError, match="演示草稿"):
        service.content("demo", {}, SimpleNamespace(username="analyst"))


def test_legacy_semantic_scope_and_prompt_are_preserved_without_sample_facts():
    from app.services.reporting.template_migration import migrate_template

    old = {
        "sections": [
            {
                "section_id": "s",
                "title": "Old",
                "prompt": "章节行文要求",
                "coverage": [
                    {
                        "kind": "ontology_relation",
                        "doc_class_iri": "urn:Report",
                        "predicate_iri": "urn:uses",
                        "range_class_iri": "urn:Equipment",
                    }
                ],
                "groups": [
                    {
                        "group_id": "g",
                        "slots": [
                            {
                                "slot_id": "u",
                                "label": "设备",
                                "required": True,
                                "source": {"kind": "semantic", "coverage_refs": []},
                            }
                        ],
                    }
                ],
            }
        ]
    }
    result = migrate_template(old, template_id="legacy", version="1")
    slot = result["target_schema"]["sections"][0]["groups"][0]["units"][0]
    assert slot["origin"]["legacy_slot"]["prompt"] == "章节行文要求"
    assert slot["origin"]["legacy_slot"]["coverage"][0]["predicate_iri"] == "urn:uses"
    old["sections"][0]["groups"][0]["slots"][0]["source"]["prompt"] = "Slot 专用要求"
    result = migrate_template(old, template_id="legacy", version="1")
    assert (
        result["target_schema"]["sections"][0]["groups"][0]["units"][0]["origin"]["legacy_slot"][
            "prompt"
        ]
        == "Slot 专用要求"
    )
    old["sections"][0]["groups"][0]["slots"][0]["source"]["coverage_refs"] = ["missing-key"]
    result = migrate_template(old, template_id="legacy", version="1")
    assert any(
        i["code"] == "LEGACY_COVERAGE_AMBIGUOUS"
        for i in result["target_schema"]["migration_issues"]
    )
    assert (
        result["target_schema"]["sections"][0]["groups"][0]["units"][0]["origin"]["legacy_slot"][
            "coverage"
        ]
        == []
    )


def test_mock_invalid_dates_and_static_scope_before_limit(db):
    from app.models.mock_data import MockEquipment
    from app.services.reporting.demo_sources import validate_record_source

    config = RecordSource(
        provider="equipment_schedules",
        contract_ref=contract_ref("equipment_schedules"),
        filters={
            "equipment_id": "A",
            "product_code": "P",
            "start_date": "invalid",
            "end_date": "2026-09-14",
        },
    )
    with pytest.raises(ReportingError) as error:
        validate_record_source(config)
    assert error.value.code == "MOCK_SCHEDULE_SCOPE_INVALID"
    db.add_all(
        [
            MockEquipment(
                equipment_id=code,
                iri="urn:" + code,
                equipment_class_iri="urn:Equipment",
                label=code,
                workshop_code="W",
            )
            for code in ("A", "B", "C")
        ]
    )
    db.commit()
    config = RecordSource(
        provider="equipment", contract_ref=contract_ref("equipment"), filters={"equipment_id": "C"}
    )
    assert load_mock(db, config, 1)["values"][0]["equipment_id"] == "C"


def test_semantic_model_must_declare_table_and_select_valid_fields(monkeypatch):
    from app.services.reporting.slot_semantics import suggest

    value = TemplateV2.model_validate(template())
    unit = value.sections[0].groups[0].units[0]
    unit.bindings, unit.inputs = [], []
    from app.services.reporting.template_v2 import NarrativeRender

    unit.render = NarrativeRender(kind="narrative")
    options = catalog(value, ontology(), {"doc": finder_source()})
    selected = next(o for o in options if o["kind"] == "graph")
    fields = [
        f["key"] for f in selected["fields"] if f["projection"].get("property_iri") == "urn:code"
    ]

    def model(_client, **kwargs):
        assert {"render", "sources"} <= set(kwargs["schema"]["$defs"]["AIChoice"]["required"])
        assert "preview_values" not in kwargs["user"]
        return {
            "slots": [
                {
                    "output_id": unit.output_id,
                    "title": "设备表",
                    "render": "table",
                    "instructions": "",
                    "sources": [{"source_key": selected["key"], "fields": fields}],
                }
            ]
        }

    monkeypatch.setattr("app.services.llm.local_client.chat_with_schema", model)
    row = SimpleNamespace(sample_analysis=None, sample_content_json=None)
    result = suggest(row, value, options, object(), output_ids=[unit.output_id])
    assert result["completion"] == "complete"
    assert result["patches"][0]["unit"]["render"]["kind"] == "table"
    resolved = resolve_finder(value=apply_patch(value, result["patches"][0]))
    assert resolved["material_status"] == "ready"
    bad = SlotChoice(
        output_id=unit.output_id,
        render="narrative",
        sources=[{"source_key": selected["key"], "fields": fields}],
    )
    with pytest.raises(ReportingError) as error:
        build_slot(value, unit, bad, options)
    assert error.value.code == "SLOT_COLLECTION_REQUIRES_TABLE"


def test_finder_equipment_join_uses_row_keys_and_preserves_missing_mock_fields(db):
    from app.models.mock_data import MockEquipment

    iri = "https://ontology.pharma-gmp.cn/slpra/equipment/"
    db.add_all(
        [
            MockEquipment(
                equipment_id=code,
                iri="urn:mock:" + code,
                equipment_class_iri="urn:Equipment",
                label="same label",
                workshop_code=workshop,
                data_properties=[{"iri": iri + "modelSpecification", "value": "mock-" + code}],
            )
            for code, workshop in [("EQ-A", "642"), ("EQ-B", "646")]
        ]
    )
    db.commit()
    config = RecordSource(provider="equipment", contract_ref=contract_ref("equipment"))
    value = template()
    defs = value["definitions"]
    defs["bindings"]["mock"] = {
        "binding_id": "mock",
        "kind": "context",
        "contract_ref": config.contract_ref,
        "scope": {"record_slot": "mock"},
    }
    defs["inputs"]["mock-rows"] = {
        "input_id": "mock-rows",
        "name": "mock-rows",
        "binding_ref": "mock",
        "projection": {
            "kind": "records",
            "fields": {
                key: {"value": {"kind": "field", "field_path": [key]}, "required": True}
                for key in ["equipment_id", "workshop_code", "specification", "material"]
            },
        },
    }
    defs["bindings"]["joined"] = {
        "binding_id": "joined",
        "kind": "derived",
        "provider": "view",
        "contract_ref": "auto:view",
        "operation": {
            "kind": "join",
            "source": {"input_id": "rows"},
            "right": {"input_id": "mock-rows"},
            "keys": ["code"],
            "right_keys": ["equipment_id"],
            "cardinality": "many_to_one",
            "unmatched": "missing",
            "fields": {
                "code": ["left", "code"],
                "source_spec": ["left", "spec"],
                "mock_spec": ["right", "specification"],
                "workshop": ["right", "workshop_code"],
                "material": ["right", "material"],
            },
        },
    }
    defs["inputs"]["joined-rows"] = {
        "input_id": "joined-rows",
        "name": "joined-rows",
        "binding_ref": "joined",
        "projection": {"kind": "identity"},
        "required": True,
    }
    unit = value["sections"][0]["groups"][0]["units"][0]
    unit["bindings"] += [{"binding_ref": "mock"}, {"binding_ref": "joined"}]
    unit["inputs"] = [{"input_ref": "joined-rows", "alias": "joined", "required": True}]
    unit["render"]["rows"]["input_id"] = "joined-rows"
    refs = {config.contract_ref: builtin_contract(config.contract_ref)}
    result = resolve_finder(
        value, extra_contracts=refs, records={"mock": load_mock(db, config, 100)}
    )
    rows = result["inputs"]["joined-rows"]["items"]
    assert [(r["fields"]["code"]["value"], r["fields"]["workshop"]["value"]) for r in rows] == [
        ("EQ-A", "642"),
        ("EQ-B", "646"),
    ]
    assert rows[0]["fields"]["source_spec"]["value"] == "spec-A"
    assert rows[0]["fields"]["mock_spec"]["value"] == "mock-EQ-A"
    assert rows[0]["fields"]["material"]["state"] == "missing"
    assert result["material_status"] == "incomplete"
    assert rows[0]["fields"]["mock_spec"]["provenance_refs"][0]["kind"] == "mock"
    db.query(MockEquipment).filter_by(equipment_id="EQ-B").delete()
    db.commit()
    missing = resolve_finder(
        value, extra_contracts=refs, records={"mock": load_mock(db, config, 100)}
    )
    assert missing["inputs"]["joined-rows"]["items"][1]["fields"]["workshop"]["state"] == "missing"


def test_legacy_required_coverage_and_properties_cannot_be_dropped():
    value = empty_slot_template()
    unit = value.sections[0].groups[0].units[0]
    unit.origin = {
        "legacy_slot": {
            "required": True,
            "coverage": [
                {
                    "kind": "ontology_relation",
                    "doc_class_iri": "urn:Report",
                    "predicate_iri": "urn:uses",
                    "range_class_iri": "urn:Equipment",
                    "required": True,
                    "required_properties": ["urn:code", "urn:spec"],
                }
            ],
        }
    }
    options = catalog(value, ontology(), {"doc": finder_source()})
    source = next(o for o in options if o["kind"] == "graph")
    code = next(
        f["key"] for f in source["fields"] if f["projection"].get("property_iri") == "urn:code"
    )
    choice = SlotChoice(
        output_id=unit.output_id,
        render="table",
        sources=[{"source_key": source["key"], "fields": [code]}],
    )
    with pytest.raises(ReportingError) as error:
        build_slot(value, unit, choice, options)
    assert error.value.code == "LEGACY_COVERAGE_INCOMPLETE"


def test_local_prose_protocol_keeps_references_and_exposes_model_failure(monkeypatch):
    import json

    from app.services.llm.local_client import StructuredModelError
    from app.services.reporting.narrative_renderer import local_provider
    from app.services.reporting.template_v2 import Budget

    payload = {
        "instructions": "编号",
        "allowed_texts": ["原文记载："],
        "inputs": [
            {"ref": {"input_id": "actual-id", "field_path": [], "scope": "input"}, "value": "EQ-A"}
        ],
        "required_refs": [{"input_id": "actual-id"}],
        "claims": {},
    }
    policy = {
        "model": "fixture",
        "temperature": 0,
        "max_input_tokens": 8192,
        "max_output_tokens": 2048,
        "timeout_s": 120,
    }
    monkeypatch.setattr("app.services.llm.local_client.get_local_llm", lambda: object())

    def model(_client, **kwargs):
        assert "$defs" not in kwargs["schema"]
        assert "maxItems" not in kwargs["schema"]["properties"]["node_indices"]
        assert kwargs["raise_on_error"] is True
        request = json.loads(kwargs["user"])
        assert request["required_indices"] == [1]
        vocabulary = request["node_vocabulary"]
        assert vocabulary[1] == {"kind": "input_ref", **payload["inputs"][0]["ref"]}
        return {"node_indices": [0, 1]}

    monkeypatch.setattr("app.services.llm.local_client.chat_with_schema", model)
    nodes = local_provider("system", payload, policy, Budget())["nodes"]
    assert nodes[1]["input_id"] == "actual-id" and nodes[0]["text"] == "原文记载："

    def failed(*args, **kwargs):
        raise StructuredModelError("model_request_failed")

    monkeypatch.setattr("app.services.llm.local_client.chat_with_schema", failed)
    with pytest.raises(ReportingError) as error:
        local_provider("system", payload, policy, Budget())
    assert error.value.code == "MODEL_REQUEST_FAILED"
    monkeypatch.setattr(
        "app.services.llm.local_client.chat_with_schema",
        lambda *args, **kwargs: {"node_indices": [-1]},
    )
    with pytest.raises(ReportingError) as error:
        local_provider("system", payload, policy, Budget())
    assert error.value.code == "OUTPUT_REFERENCE_INVALID"
