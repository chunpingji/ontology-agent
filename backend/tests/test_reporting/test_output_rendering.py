from io import BytesIO

import pytest
from docx import Document

from app.services.reporting.output_ast import plain_text
from app.services.reporting.output_renderer import render_docx, render_snapshot
from app.services.reporting.template_v2 import Budget, NarrativeRender, ReportingError
from tests.test_reporting.test_fact_selector import property_value
from tests.test_reporting.test_output_resolution import bundle, resolve


def test_docx_and_preview_share_values_and_citations():
    from tests.test_reporting.test_output_contracts import template

    data = bundle(property_value("code", "E", "urn:code", "E-01"))
    data["template"] = template()
    snapshot = resolve(data)
    rendered = render_snapshot(snapshot)
    assert rendered["execution_status"] == "completed"
    assert "E-01" in plain_text(rendered["body_ast"])
    doc = Document(BytesIO(render_docx(rendered["body_ast"])))
    assert doc.tables[0].cell(1, 0).text == "E-01"
    citations = rendered["output_results"][0]["citations"]
    assert citations[0]["input_ref"]["field_path"] == ["code"]
    assert citations[0]["fact_refs"]


def test_assisted_cannot_omit_or_add_references():
    from app.services.reporting.input_resolver import typed_value
    from app.services.reporting.narrative_renderer import assisted_nodes

    render = NarrativeRender.model_validate(
        {
            "kind": "narrative",
            "mode": "assisted",
            "prompt": {
                "policy_ref": "prompt",
                "input_refs": [{"input_id": "code"}],
                "required_refs": [{"input_id": "code"}],
            },
        }
    )
    inputs = {"code": typed_value("code", {"kind": "string"}, "E-01")}
    policy = {"model": "fixture", "max_input_tokens": 8192}
    for proposal in (
        {"nodes": [{"kind": "text", "text": "E-01"}]},
        {"nodes": [{"kind": "input_ref", "input_id": "other"}]},
        {"nodes": [{"kind": "signature_region", "signature_region_id": "qa"}]},
    ):
        with pytest.raises(ReportingError) as error:
            assisted_nodes(render, inputs, {}, policy, Budget(), lambda *args: proposal)
        assert error.value.code == "OUTPUT_REFERENCE_INVALID"


def workflow_output(typ, records, render):
    from app.services.reporting.report_snapshot import resolve_snapshot
    from app.services.reporting.template_compiler import compile_template, require_valid
    from tests.test_reporting.test_output_contracts import contracts, ontology, template

    value, registered = template(), contracts()
    value["source_slots"] = []
    value["definitions"]["bindings"]["equipment"] = {
        "binding_id": "equipment",
        "kind": "workflow",
        "contract_ref": "wf",
        "scope": {"record_slot": "roster"},
    }
    value["definitions"]["inputs"]["rows"]["projection"] = {"kind": "identity"}
    value["sections"][0]["groups"][0]["units"][0]["render"] = render
    registered["wf"] = {
        "kind": "workflow",
        "status": "published",
        "definition": {"output_type": typ},
    }
    plan = require_valid(compile_template(value, ontology(), registered))
    data = {
        "template": plan["template"],
        "contracts": registered,
        "sources": {},
        "schema": ontology(),
        "records": {
            "roster": {
                "record_id": "reviewed-fixture",
                "contract_id": "wf",
                "state": "ready",
                "subject_id": "urn:E",
                "applicable_at": None,
                "provenance": {"record_ref": "fixture"},
                "values": records,
            }
        },
    }
    frozen = resolve_snapshot(plan, data)
    return frozen, render_snapshot(frozen)


def test_numeric_table_sort_and_grouping_use_typed_values_in_docx():
    typ = {
        "kind": "list",
        "item_identity": "entity_id",
        "item_type": {
            "kind": "record",
            "fields": {
                "amount": {"type": {"kind": "decimal"}, "semantic_ref": "urn:amount"},
                "group": {"type": {"kind": "string"}, "semantic_ref": "urn:group"},
            },
        },
    }
    _, output = workflow_output(
        typ,
        [
            {"entity_id": "a", "amount": "10", "group": "A"},
            {"entity_id": "b", "amount": "2", "group": "A"},
            {"entity_id": "c", "amount": "1", "group": "B"},
        ],
        {
            "kind": "table",
            "rows": {"input_id": "rows"},
            "columns": [{"column_id": "a", "title": "Amount", "field_ref": "amount"}],
            "order_by": [{"field_ref": "amount"}],
            "group_by": ["group"],
        },
    )
    assert output["execution_status"] == "completed"
    document = Document(BytesIO(render_docx(output["body_ast"])))
    cells = [row.cells[0].text for row in document.tables[0].rows]
    assert cells == ["Amount", "A", "2", "10", "B", "1"]


def test_nested_repeats_keep_each_record_item_scope():
    children = {
        "kind": "list",
        "item_identity": "entity_id",
        "item_type": {
            "kind": "record",
            "fields": {"code": {"type": {"kind": "string"}, "semantic_ref": "urn:code"}},
        },
    }
    typ = {
        "kind": "list",
        "item_identity": "entity_id",
        "item_type": {
            "kind": "record",
            "fields": {"children": {"type": children, "semantic_ref": "urn:children"}},
        },
    }
    frozen, output = workflow_output(
        typ,
        [
            {
                "entity_id": "parent-a",
                "children": [
                    {"entity_id": "child-a", "code": "A1"},
                    {"entity_id": "child-b", "code": "A2"},
                ],
            },
            {"entity_id": "parent-b", "children": [{"entity_id": "child-c", "code": "B1"}]},
        ],
        {
            "kind": "narrative",
            "nodes": [
                {
                    "kind": "repeat",
                    "input_id": "rows",
                    "children": [
                        {
                            "kind": "repeat",
                            "input_id": "rows",
                            "scope": "item",
                            "field_path": ["children"],
                            "children": [
                                {
                                    "kind": "input_ref",
                                    "input_id": "rows",
                                    "scope": "item",
                                    "field_path": ["code"],
                                }
                            ],
                        }
                    ],
                }
            ],
        },
    )
    assert frozen["material_status"] == "ready", frozen["blocking_issues"]
    assert output["execution_status"] == "completed", output
    assert plain_text(output["body_ast"]).endswith("A1A2B1")
    citations = output["output_results"][0]["citations"]
    assert [c["input_ref"]["record_id"] for c in citations] == ["child-a", "child-b", "child-c"]


def test_blocked_value_preserves_fact_citations():
    from tests.test_reporting.test_output_contracts import template

    data = bundle(
        property_value("one", "E", "urn:code", "one"), property_value("two", "E", "urn:code", "two")
    )
    data["template"] = template()
    output = render_snapshot(resolve(data))
    citations = output["output_results"][0]["citations"]
    assert citations[0]["fact_refs"]
    assert citations[0]["provenance_refs"]
