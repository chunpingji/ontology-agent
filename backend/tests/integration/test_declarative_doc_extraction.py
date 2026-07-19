"""017 — E6/E6b doc_pattern executes through the shared Word interpreter."""

from __future__ import annotations

import json
from io import BytesIO

from docx import Document

from tests.fixtures.feature014 import (
    ANALYST,
    APPROVAL_NUMBER,
    DRUG_PRODUCT,
    RISK_LEVEL,
    add_property_binding,
)

CLASSES = "/api/ontology/classes"
JOBS = "/api/extraction/jobs"


def _upload() -> bytes:
    document = Document()
    document.add_heading("HRS-1597 CMC报告", level=1)
    document.add_heading("产品基本信息", level=2)
    document.add_paragraph("批准文号：HRS-1597")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_doc_pattern_partial_success_marks_binding_drift(drug_client):
    profile = {
        "version": 1,
        "sources": [{
            "locator": "section_kv",
            "anchors": {"any_of": ["产品基本信息"]},
        }],
        "identity": {"pattern": "[A-Z]{2,4}-[0-9]{3,5}"},
    }
    mapping = drug_client.post(
        f"{CLASSES}/{DRUG_PRODUCT}/mappings",
        json={
            "mapping_type": "doc_pattern",
            "target": json.dumps(profile, ensure_ascii=False),
            "source_system": "cmc-word",
        },
        headers=ANALYST,
    )
    assert mapping.status_code == 201, mapping.text
    mapping_id = mapping.json()["id"]

    assert add_property_binding(
        drug_client,
        mapping_id,
        property_iri=APPROVAL_NUMBER,
        source_path="批准文号|注册批准文号",
        is_identifier=True,
    ).status_code == 201
    assert add_property_binding(
        drug_client,
        mapping_id,
        property_iri=RISK_LEVEL,
        source_path="风险级别",
    ).status_code == 201

    response = drug_client.post(
        JOBS,
        data={"source_type": "word", "class_mapping_id": mapping_id},
        files={"file": (
            "HRS-1597.docx",
            _upload(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )},
        headers=ANALYST,
    )
    assert response.status_code == 202, response.text
    job_id = response.json()["id"]

    job = drug_client.get(f"{JOBS}/{job_id}").json()
    assert job["status"] in ("reviewing", "done")
    assert job["total_candidates"] == 1

    body = drug_client.get(f"{JOBS}/{job_id}/candidates").json()
    candidates = body["ungrouped"] + [
        candidate
        for group in body["groups"]
        for candidate in group["candidates"]
    ]
    assert candidates[0]["extracted_properties"] == {
        APPROVAL_NUMBER: "HRS-1597"
    }
    source_ref = json.loads(candidates[0]["source_ref"])
    assert source_ref["system"] == "cmc-word"
    assert source_ref["entity"] == "HRS-1597.docx"
    assert source_ref["location"]["kind"] == "section"

    mappings = drug_client.get(
        f"{CLASSES}/{DRUG_PRODUCT}/mappings"
    ).json()
    current = next(item for item in mappings if item["id"] == mapping_id)
    assert current["health"] == "drift"
