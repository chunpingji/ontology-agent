"""Bounded real-model probe of the production document harness, with separate scoring."""

from __future__ import annotations

import argparse
import json
import shutil
from collections import Counter
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

from docx import Document
from rdflib import Graph

from app.services.document_harness.controller import Engine
from app.services.document_harness.coreference import canonical_mentions
from app.services.document_harness.model import call_model, freeze_policy
from app.services.document_harness.ontology import SchemaCatalog, catalog_from_graph
from app.services.document_harness.protocols import INSTRUCTIONS
from app.services.document_harness.source import reference
from app.services.extraction.document_ir import DocumentIR, build_document_ir
from app.services.extraction.docx_structure import parse_docx_structure
from app.services.llm.model_runtime import model_scope


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n")


def synthetic_catalog():
    return catalog_from_graph(Graph().parse(data='''
        @prefix : <urn:coreference-probe:> .
        @prefix owl: <http://www.w3.org/2002/07/owl#> .
        @prefix rdfs: <http://www.w3.org/2000/01/rdf-schema#> .
        @prefix xsd: <http://www.w3.org/2001/XMLSchema#> .
        :Report a owl:Class; rdfs:label "设备记录报告" .
        :Equipment a owl:Class; rdfs:label "生产设备";
          rdfs:comment "工厂中用于生产或清洗的具体物理设备。设备的使用事件不是设备本身。";
          owl:hasKey (:identifier) .
        :identifier a owl:DatatypeProperty; rdfs:label "设备编号";
          rdfs:comment "在指定厂区的设备登记范围内唯一标识设备，需原文确认其归属及范围。";
          rdfs:domain :Equipment; rdfs:range xsd:string .
        :describes a owl:ObjectProperty; rdfs:label "描述设备";
          rdfs:comment "报告记录或描述的设备。"; rdfs:domain :Report; rdfs:range :Equipment .
    ''', format="turtle"), "urn:coreference-probe:Report")


CASES = {
    "alias": (["清洗设备甲，以下简称清洗单元；该简称适用于本报告后续各节。",
               "清洗单元完成本次清洗工作。"], ["same"]),
    "identifier": (["甲厂设备登记：反应釜的设备编号为 EQ-17，该编号在甲厂内唯一。",
                    "甲厂维护记录：设备编号 EQ-17 的反应釜完成维护。"], ["same"]),
    "reference": (["第一节登记的设备是一台200L反应釜。",
                   "本节使用第一节登记的那台200L反应釜进行清洗。"], ["same"]),
    "different_identifiers": (["设备名称：清洗单元，甲厂内唯一设备编号 EQ-17。",
                               "另有一台设备，名称也为清洗单元，甲厂内唯一设备编号 EQ-18。"],
                              ["different"]),
    "name_only": (["清洗单元完成第一节的工作，未提供设备编号或后文对应说明。",
                   "清洗单元完成本节工作，未说明与第一节设备的对应关系。"], ["unresolved"]),
    "ambiguous_reference": (["本次登记一台清洗设备甲。",
                             "另外登记一台清洗设备乙，与甲是不同设备。",
                             "上述清洗设备需要检查，未指明甲或乙。"],
                            ["different", "unresolved", "unresolved"]),
}


def fixture(root, name, texts, catalog):
    doc = Document()
    for i, text in enumerate(texts):
        doc.add_heading(f"第{i + 1}节", 1)
        doc.add_paragraph(text)
    path = root / f"{name}.docx"
    doc.save(path)
    ir = build_document_ir(path, parse_docx_structure(path))
    root_card = catalog.classes[catalog.root_class_iri]
    entities = {"document": {"id": "document", "label": "设备记录报告", "role": "document_root",
                             "class_iri": root_card.iri, "class_label": root_card.label,
                             "state": "accepted", "reason": "任务根", "evidence": [],
                             "field_ids": [], "window_id": None}}
    for i, text in enumerate(texts):
        unit = next(u for u in ir.evidence_units if u.text == text)
        ref = reference(ir, unit.evidence_id, 0, len(text))
        entities[str(i)] = {"id": str(i), "label": text, "name": None, "role": "原文设备提及",
                           "class_iri": "urn:coreference-probe:Equipment",
                           "class_label": "生产设备",
                           "state": "accepted", "reason": "定向共指输入，类型为测试前提",
                           "evidence": [ref], "referent": ref, "field_ids": [], "window_id": str(i)}
    return ir, {"entities": entities}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--source-run", type=UUID)
    parser.add_argument("--cases", nargs="+", choices=[*CASES, "end_to_end"],
                        default=[*CASES, "end_to_end"])
    args = parser.parse_args()
    root = args.output
    root.mkdir(parents=True, exist_ok=False)
    runtime = Path(__file__).resolve().parents[1] / "services" / "document_harness"
    shutil.copytree(runtime, root / "runtime", ignore=shutil.ignore_patterns("__pycache__"))
    shutil.copy2(__file__, root / "probe.py")
    policy = freeze_policy()
    catalog = synthetic_catalog()
    write(root / "manifest.json", {
        "created_at": datetime.now(UTC), "run_id": root.name, "policy": policy,
        "instructions": INSTRUCTIONS, "call_limit": 60,
        "reference_is_recognition_input": False,
        "scope": "controlled source-pair co-reference + one end-to-end synthetic document"
                 "; optional saved real-document mentions, no source-run modifications",
        "reference_status": "developer-authored regression expectations, not expert gold",
    })
    write(root / "synthetic-catalog.json", catalog.model_dump(mode="json"))
    write(root / "reference.json", {name: expected for name, (_, expected) in CASES.items()})
    calls, summary = [], {}
    current_case = ""

    def invoke(stage, payload, schema):
        if len(calls) >= 60:
            raise RuntimeError("evaluation_call_limit")
        number = len(calls) + 1
        path = root / f"{number:03}-{current_case}-{stage}.json"
        record = {"number": number, "case": current_case, "stage": stage,
                  "payload": payload, "schema": schema}
        write(path, record)
        print(f"CALL {number} {current_case} {stage}", flush=True)
        with model_scope(run_id=root.name, task_id=str(number), stage=stage):
            result = call_model(stage, payload, schema, policy)
        write(path, {**record, "result": result})
        calls.append({"number": number, "case": current_case, "stage": stage,
                      "seconds": result["seconds"], "usage": result["usage"],
                      "error": result["error"]})
        write(root / "costs.json", calls)
        if result["error"]:
            raise RuntimeError(result["error"])
        return result["output"]

    def run_case(name, ir, cards, state, *, full=False):
        nonlocal current_case
        current_case = name
        write(root / f"{name}-ir.json", ir.model_dump(mode="json"))
        write(root / f"{name}-input.json", state)
        saved = deepcopy(state)
        def save(changes):
            for domain, values in changes.items():
                saved.setdefault(domain, {}).update(deepcopy(values))
            write(root / f"{name}-state.json", saved)
        engine = Engine(ir=ir, catalog=cards, state=state, invoke=invoke, save=save,
                        should_stop=lambda: False, max_input_tokens=policy["max_input_tokens"])
        if not full:
            engine.state["cursor"] = {"main": {
                "window_index": len(engine.windows), "stage": "coreference_review",
                "scope_complete": False, "windows_total": len(engine.windows),
                "windows_discovered": 0, "windows_reviewed": 0,
            }}
        engine.run()
        decisions = sorted(engine.state.get("coreferences", {}).values(),
                           key=lambda row: (row["left_mention_id"], row["right_mention_id"]))
        for decision in decisions:
            for ref in decision["evidence"] + decision["proof"]:
                assert ir.unit(ref["source_id"]).text[ref["start"]:ref["end"]] == ref["text"]
        aliases = canonical_mentions(engine.state["entities"], decisions, cards.classes)
        summary[name] = {
            "verdicts": [row["verdict"] for row in decisions], "decisions": decisions,
            "mentions": len(aliases) - 1, "document_local_entities": len(set(aliases.values())) - 1,
            "states": dict(Counter(row["state"] for key, row in engine.state["entities"].items()
                                   if key != "document")),
            "pipeline_complete": engine.state["cursor"]["main"]["stage"] == "complete",
            "full_discovery_run": full,
        }
        write(root / "summary.json", summary)
        print("RESULT", name, summary[name]["verdicts"], flush=True)

    for name, (texts, expected) in CASES.items():
        if name not in args.cases:
            continue
        ir, state = fixture(root, name, texts, catalog)
        run_case(name, ir, catalog, state)
        summary[name]["expected"] = expected
        summary[name]["passed"] = summary[name]["verdicts"] == expected
        write(root / "summary.json", summary)

    if "end_to_end" in args.cases:
        ir, _ = fixture(root, "end_to_end", CASES["alias"][0], catalog)
        run_case("end_to_end", ir, catalog, {}, full=True)
        summary["end_to_end"]["passed"] = (
            summary["end_to_end"]["mentions"] >= 2
            and summary["end_to_end"]["document_local_entities"] == 1
            and summary["end_to_end"]["states"].get("accepted", 0) >= 2
        )
    if args.source_run:
        from sqlalchemy import select

        from app.db import SessionLocal
        from app.models.document_analysis import DocumentAnalysisArtifact, DocumentAnalysisRun
        from app.services.document_analysis.run_store import DocumentAnalysisRunStore
        from app.services.document_harness.coreference import candidate_pairs
        from app.services.document_harness.runtime import get_row, read_rows

        with SessionLocal() as db:
            source_run = db.scalar(select(DocumentAnalysisRun).where(
                DocumentAnalysisRun.recognition_run_id == args.source_run,
            ))
            if source_run is None:
                raise ValueError("source_run_missing")
            ir = DocumentIR.model_validate(get_row(db, source_run, "input", "document"))
            state = read_rows(db, source_run, domains={"entities", "fields"})
            ref = DocumentAnalysisRunStore(db).get_artifact(
                args.source_run, source_run.owner_id, "ontology_snapshot",
            )
            cards = SchemaCatalog.model_validate(
                db.get(DocumentAnalysisArtifact, ref.artifact_id).payload,
            )
        write(root / "real-catalog.json", cards.model_dump(mode="json"))
        pairs = [(left, right) for _, left, right in candidate_pairs(state, cards.classes)
                 if left["referent"]["section_id"] != right["referent"]["section_id"]]
        # Freeze a small exact cross-section scope before calling the model, without a gold filter.
        write(root / "real-scope.json", {"source_run": args.source_run,
                                       "pairs": [[a["id"], b["id"]] for a, b in pairs[:3]]})
        for i, (left, right) in enumerate(pairs[:3]):
            selected = {key: state["entities"][key]
                        for key in ("document", left["id"], right["id"])}
            selected_fields = {key: state["fields"][key]
                               for row in selected.values() for key in row["field_ids"]}
            run_case(f"real_{i + 1}", ir, cards,
                     {"entities": selected, "fields": selected_fields})
    write(root / "summary.json", summary)
    print("OUTPUT_DIR", root, flush=True)
    if not all(item.get("passed", True) for item in summary.values()):
        raise SystemExit("regression_expectations_not_met; see summary.json")


if __name__ == "__main__":
    main()
