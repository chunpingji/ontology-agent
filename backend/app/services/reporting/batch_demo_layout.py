"""Sample formatting and explicitly selected sample content for the HRS-5592 demo."""

import hashlib
import json
from copy import deepcopy
from io import BytesIO
from math import ceil
from pathlib import Path
from uuid import UUID

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from fastapi import HTTPException

from app.models.extraction import AstTemplate
from app.services.reporting.batch_demo_operation_forms import load_forms, validate_forms
from app.services.reporting.batch_demo_sample_content import A14_OPERATIONS, A14_TABLES

REFERENCE_ID = UUID("74cd92d6-41ee-42d5-84a2-b4a736a087b9")
RENDERER_VERSION = "sample-layout-v4"
PROTOTYPES = {"cover": (0, 4), "materials": (1, 5), "equipment": (3, 6),
              "instruments": (4, 7), "site": (9, 4), "operations": (15, 5),
              "clearance": (27, 4), "deviation": (28, 5), "control": (29, 4)}


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(",", ":")).encode()).hexdigest()


def metadata(node):
    return next((p for p in node.provenance_refs if p.get("kind") == "template_layout"), {})


def copy_property(target, source, tag):
    current = target.find(qn(tag))
    if current is not None:
        target.remove(current)
    original = source.find(qn(tag)) if source is not None else None
    if original is not None:
        target.insert(0, deepcopy(original))


def fill_paragraph(target, prototype, text, layout):
    """Copy only formatting, never drawings, links, fields, bookmarks or old text."""
    copy_property(target._p, prototype._p, "w:pPr")
    if target._p.pPr is not None:
        for tag in ("w:sectPr", "w:numPr"):
            for child in target._p.pPr.findall(qn(tag)):
                target._p.pPr.remove(child)
    if str(text) == prototype.text:
        for original in prototype.runs:
            run = target.add_run(original.text)
            copy_property(run._r, original._r, "w:rPr")
        return
    run = target.add_run(str(text))
    original = next((r for r in prototype.runs if r.text.strip()), None)
    if original is None and prototype.runs:
        original = prototype.runs[0]
    if original is not None:
        copy_property(run._r, original._r, "w:rPr")
    style = layout.paragraph_style(prototype)
    run.font.name = style["latin_font_family"]
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), style["font_family"])
    run.font.size = Pt(style["font_size_pt"])


class SampleLayout:
    def __init__(self, db):
        row = db.get(AstTemplate, REFERENCE_ID)
        path = Path(row.sample_docx_path) if row and row.sample_docx_path else None
        if path is None or not path.is_file():
            raise HTTPException(409, "指定版式模板的 Word 样例不可用")
        try:
            raw = path.read_bytes()
            self.document = Document(BytesIO(raw))
            self.tables = self.document.tables
            self.paragraphs = self.document.paragraphs
            self._cells = {}
            self._paragraph_styles = {}
            self.forms = load_forms()
            if self.forms["template_id"] != str(REFERENCE_ID):
                raise ValueError("operation form reference changed")
            validate_forms(self)
            grids = self.document.sections[0]._sectPr.xpath("w:docGrid/@w:linePitch")
            self.line_pitch_pt = int(grids[0]) / 20 if grids else 15.6
            # This adapter deliberately supports the inspected sample structure only.
            if (row.schema_json or {}).get("sections"):
                raise ValueError("reference now has AST sections")
            for role, (index, columns) in PROTOTYPES.items():
                table = self.tables[index]
                required_rows = 6 if role == "cover" else 3 if role == "deviation" else 2
                if len(table.columns) != columns or len(table.rows) < required_rows:
                    raise ValueError(role)
                expected = [3] * 6 if role == "cover" else (
                    [5, 5, 3] if role == "deviation" else [columns, columns]
                )
                if any(len(table.rows[i]._tr.tc_lst) != n for i, n in enumerate(expected)):
                    raise ValueError(role + " cells")
            if self.paragraphs[6].text.strip() != "批生产记录":
                raise ValueError("cover")
            if self.tables[15].cell(0, 0).text.strip() != "原材料/操作":
                raise ValueError("operations")
            if len(self.document.sections) < 6:
                raise ValueError("sections")
            self.header = self.document.sections[2].header.tables[0]
            if len(self.header.columns) != 4 or len(self.header.rows[0]._tr.tc_lst) != 4:
                raise ValueError("header")
            for si in range(1, 6):
                header = self.document.sections[si].header
                if len(header.tables) != 1 or len(header.tables[0].columns) != 4:
                    raise ValueError("section header")
                self.validate_plain_content(header._element)
            for ti in A14_TABLES:
                table = self.tables[ti]
                self.validate_plain_content(table._tbl)
                expected = "原材料/操作" if ti in A14_OPERATIONS else (
                    "称量记录" if ti < 15 else "TLC结果处理："
                )
                if not _Cell(table._tbl.tr_lst[0].tc_lst[0], table).text.startswith(expected):
                    raise ValueError("A14 operation record moved")
        except Exception as exc:
            raise HTTPException(409, "参考 Word 样例结构不兼容，请核对演示版式") from exc
        self.manifest = {
            "reference_template_id": str(REFERENCE_ID), "reference_name": row.name,
            "schema_hash": fingerprint(row.schema_json),
            "sample_sha256": hashlib.sha256(raw).hexdigest(),
            "sample_content_hash": fingerprint(row.sample_content_json),
            "renderer_version": RENDERER_VERSION,
            "operation_forms_hash": fingerprint(self.forms),
        }

    def table(self, role):
        if isinstance(role, int):
            return self.tables[role]
        if role.startswith("header:"):
            return self.document.sections[int(role.split(":")[1])].header.tables[0]
        return self.header if role == "header" else self.tables[PROTOTYPES[role][0]]

    @staticmethod
    def validate_plain_content(element):
        # Selected forms have no package relationships or executable fields.
        if element.xpath(".//w:drawing|.//w:pict|.//w:object|.//w:hyperlink|.//w:fldSimple"
                         "|.//w:instrText|.//w:altChunk|.//w:ins|.//w:del"):
            raise ValueError("unsupported content in selected sample form")

    def paragraph_style(self, paragraph, run=None):
        properties = []
        run = run or next((r for r in paragraph.runs if r.text.strip()), None)
        key = (str(paragraph._p.pPr.xml) if paragraph._p.pPr is not None else "",
               str(run._r.rPr.xml) if run is not None and run._r.rPr is not None else "")
        if key in self._paragraph_styles:
            return self._paragraph_styles[key]
        if run is not None and run._r.rPr is not None:
            properties.append(run._r.rPr)
        style = run.style if run is not None else None
        while style is not None:
            if style.element.rPr is not None:
                properties.append(style.element.rPr)
            style = style.base_style
        properties.extend(paragraph._p.xpath("w:pPr/w:rPr"))
        style = paragraph.style
        while style is not None:
            if style.element.rPr is not None:
                properties.append(style.element.rPr)
            style = style.base_style
        properties.extend(self.document.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr"))

        def attribute(tag, key, default=None):
            for pr in properties:
                el = pr.find(qn(tag))
                if el is not None and (value := el.get(qn(key), default)) is not None:
                    return value
            return None

        size = attribute("w:sz", "w:val")
        font = attribute("w:rFonts", "w:eastAsia")
        latin_font = attribute("w:rFonts", "w:ascii") or attribute("w:rFonts", "w:hAnsi")
        align = paragraph.paragraph_format.alignment
        spacing = paragraph.paragraph_format.line_spacing
        size_pt = float(size) / 2 if size else 10.5
        line_height = spacing.pt if hasattr(spacing, "pt") else (
            size_pt * spacing if spacing else
            ceil(size_pt / self.line_pitch_pt) * self.line_pitch_pt
        )
        self._paragraph_styles[key] = {
                "font_size_pt": size_pt, "line_height_pt": line_height,
                "underline": attribute("w:u", "w:val", "single") not in (None, "none"),
                "font_family": font or "宋体", "latin_font_family": latin_font or "Times New Roman",
                "italic": attribute("w:i", "w:val", "1") not in (None, "0", "false"),
                "script": attribute("w:vertAlign", "w:val"),
                "align": {1: "center", 2: "right", 3: "justify"}.get(
                    align, "left"),
                "bold": attribute("w:b", "w:val", "1") not in (None, "0", "false")}
        return self._paragraph_styles[key]

    def cell_paragraphs(self, cell):
        result = []
        for p in cell.paragraphs:
            runs = []
            for run in p.runs:
                if not run.text:
                    continue
                style = self.paragraph_style(p, run)
                if runs and runs[-1]["style"] == style:
                    runs[-1]["text"] += run.text
                else:
                    runs.append({"text": run.text, "style": style})
            result.append({"runs": runs, "style": self.paragraph_style(p),
                           "space_before_pt": p.paragraph_format.space_before.pt
                           if p.paragraph_format.space_before else 0,
                           "space_after_pt": p.paragraph_format.space_after.pt
                           if p.paragraph_format.space_after else 0})
        return result

    def page_metadata(self, index, role):
        s = self.document.sections[index]
        return {"kind": "template_layout", "role": role, "section_index": index,
                "width_pt": s.page_width.pt, "height_pt": s.page_height.pt,
                "top_pt": s.top_margin.pt, "bottom_pt": s.bottom_margin.pt,
                "left_pt": s.left_margin.pt, "right_pt": s.right_margin.pt}

    def table_metadata(self, role):
        table = self.table(role)
        widths = [c.w.twips for c in table._tbl.tblGrid.gridCol_lst]
        return {"kind": "template_layout", "role": role, "widths": widths,
                "width_pt": sum(widths) / 20,
                "align": {1: "center", 2: "right"}.get(table.alignment, "left")}

    def cell_metadata(self, role, row_index, cell_index):
        key = (role, row_index, cell_index)
        if key in self._cells:
            return self._cells[key]
        table = self.table(role)
        cell = _Cell(table.rows[row_index]._tr.tc_lst[cell_index], table)
        borders = {}
        for side in ("top", "right", "bottom", "left"):
            border = cell._tc.tcPr.find("w:tcBorders/w:" + side, namespaces=cell._tc.nsmap)
            if border is None:
                border = table._tbl.tblPr.find("w:tblBorders/w:" + side,
                                               namespaces=table._tbl.nsmap)
            if border is None:
                borders["border_" + side] = "0.5pt solid black"
            elif border.get(qn("w:val")) in ("none", "nil"):
                borders["border_" + side] = "none"
            else:
                width = int(border.get(qn("w:sz"), "4")) / 8
                borders["border_" + side] = f"{width}pt solid black"
        alignment = cell._tc.xpath("w:tcPr/w:vAlign/@w:val")
        self._cells[key] = {"kind": "template_layout", "span": cell._tc.grid_span, **borders,
                            "vertical_align": alignment[0] if alignment else "top",
                            **self.paragraph_style(cell.paragraphs[0])}
        return self._cells[key]

    def render_table(self, target, node):
        meta = metadata(node)
        if meta.get("verbatim"):
            result = deepcopy(self.table(meta["sample_table"])._tbl)
            target._body._body.insert_element_before(result, "w:sectPr")
            return
        prototype = self.table(meta["role"])
        result = OxmlElement("w:tbl")
        result.append(deepcopy(prototype._tbl.tblPr))
        result.append(deepcopy(prototype._tbl.tblGrid))
        for row in node.children:
            row_table = self.tables[metadata(row)["prototype_table"]] if (
                "prototype_table" in metadata(row)
            ) else prototype
            source_row = row_table.rows[metadata(row)["prototype_row"]]._tr
            tr = OxmlElement("w:tr")
            copy_property(tr, source_row, "w:trPr")
            if tr.trPr is not None:
                # Sample heights may be exact; variable source instructions must not clip.
                for height in tr.trPr.findall(qn("w:trHeight")):
                    height.set(qn("w:hRule"), "atLeast")
                for tag in ("w:tblHeader",):
                    for el in tr.trPr.findall(qn(tag)):
                        tr.trPr.remove(el)
            if row.header:
                tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for index, cell_node in enumerate(row.children):
                cell_meta = metadata(cell_node)
                source_cell = source_row.tc_lst[cell_meta.get("prototype_cell", index)]
                tc = OxmlElement("w:tc")
                copy_property(tc, source_cell, "w:tcPr")
                if tc.tcPr is not None:
                    for tag in ("w:vMerge", "w:tcFitText", "w:noWrap"):
                        for el in tc.tcPr.findall(qn(tag)):
                            tc.tcPr.remove(el)
                if meta["role"] == "operations":
                    # Formula rows in the sample span parameter + record. The print contract
                    # now requires a name column and a separate, empty value/unit column.
                    tc.grid_span = 1
                    tc.width = prototype._tbl.tblGrid.gridCol_lst[index].w
                if cell_meta.get("vmerge"):
                    merge = OxmlElement("w:vMerge")
                    merge.set(qn("w:val"), cell_meta["vmerge"])
                    tc.get_or_add_tcPr().append(merge)
                original = _Cell(source_cell, row_table)
                if cell_node.text == original.text:
                    for original_p in original.paragraphs:
                        p = OxmlElement("w:p")
                        tc.append(p)
                        fill_paragraph(Paragraph(p, target), original_p, original_p.text, self)
                else:
                    p = OxmlElement("w:p")
                    tc.append(p)
                    fill_paragraph(Paragraph(p, target), original.paragraphs[0],
                                   cell_node.text, self)
                tr.append(tc)
            result.append(tr)
        if hasattr(target, "_body"):
            target._body._body.insert_element_before(result, "w:sectPr")
        else:
            target._element.insert(0, result)

    def render(self, ast):
        # A fresh OPC package ensures sample media, embeddings and custom XML never leak.
        document = Document()
        document.styles.element.clear()
        for child in self.document.styles.element:
            document.styles.element.append(deepcopy(child))
        document.part.part_related_by(RT.THEME)._blob = self.document.part.part_related_by(
            RT.THEME,
        ).blob
        for index, section_node in enumerate(ast.children):
            section = document.sections[0] if index == 0 else document.add_section(
                WD_SECTION_START.NEW_PAGE,
            )
            source = self.document.sections[metadata(section_node)["section_index"]]
            for prop in ("page_width", "page_height", "top_margin", "bottom_margin",
                         "left_margin", "right_margin", "header_distance", "footer_distance"):
                setattr(section, prop, getattr(source, prop))
            copy_property(section._sectPr, source._sectPr, "w:docGrid")
            section.header.is_linked_to_previous = False
            section.footer.is_linked_to_previous = False
            for part in (section.header, section.footer):
                for p in part.paragraphs:
                    p._p.clear()
            header_index = metadata(section_node).get("header_section")
            if header_index is not None:
                source_header = self.document.sections[header_index].header._element
                # Preserve namespace declarations too (mc:Ignorable references them).
                section.header.part._element = deepcopy(source_header)
            for child in section_node.children:
                if child.kind == "table":
                    if metadata(child)["role"] != "header":
                        self.render_table(document, child)
                else:
                    if metadata(child).get("verbatim"):
                        original = self.paragraphs[metadata(child)["sample_paragraph"]]
                        document._body._body.insert_element_before(
                            deepcopy(original._p), "w:sectPr",
                        )
                        continue
                    p = document.add_paragraph()
                    prototype_index = metadata(child).get("prototype_paragraph", 16)
                    fill_paragraph(p, self.paragraphs[prototype_index], child.text, self)
                    for prop in ("space_before", "space_after"):
                        setattr(p.paragraph_format, prop, Pt(metadata(child)[prop + "_pt"]))
        document.core_properties.title = "HRS-5592 批生产记录（演示草稿）"
        document.core_properties.author = ""
        document.core_properties.last_modified_by = ""
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()
