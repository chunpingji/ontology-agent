"""Explicitly authorized template content for headers and A14 operation records."""

from docx.oxml.ns import qn
from docx.table import _Cell

A14_TABLES = tuple(range(10, 27))
A14_OPERATIONS = (15, 16, 18, 19, 21, 22, 23, 24, 25, 26)


def reference_table(layout, key, node, role):
    table = layout.table(key)
    grids = []
    for tr in table._tbl.tr_lst:
        grid = []
        for tc in tr.tc_lst:
            grid.extend([tc] * tc.grid_span)
        grids.append(grid)
    rows = []
    for ri, tr in enumerate(table._tbl.tr_lst):
        cells, column = [], 0
        for ci, tc in enumerate(tr.tc_lst):
            cell = _Cell(tc, table)
            meta = {**layout.cell_metadata(key, ri, ci),
                    "paragraphs": layout.cell_paragraphs(cell)}
            if tc.vMerge:
                span = 1
                while (tc.vMerge == "restart" and ri + span < len(grids)
                       and grids[ri + span][column].vMerge == "continue"):
                    span += 1
                meta.update(vmerge=tc.vMerge, row_span=span)
            cells.append(node("cell", meta, text=cell.text))
            column += tc.grid_span
        height = table.rows[ri].height
        rows.append(node("row", {"kind": "template_layout", "prototype_row": ri,
                                 "height_pt": height.pt if height else None},
                         children=cells, header=ri == 0))
    result = node("table", {**layout.table_metadata(key), "role": role,
                            "sample_table": key, "verbatim": True}, children=rows)
    result.provenance_refs.append({"kind": "reference_template",
                                   "template_id": layout.manifest["reference_template_id"],
                                   "sample_sha256": layout.manifest["sample_sha256"],
                                   "table": key})
    return result


def operation_records(layout, node):
    result = []
    paragraphs = {p._p: i for i, p in enumerate(layout.paragraphs)}
    for ti in A14_TABLES:
        previous, between = layout.tables[ti]._tbl.getprevious(), []
        while previous is not None and previous.tag == qn("w:p"):
            between.append(previous)
            previous = previous.getprevious()
        for element in reversed(between):
            pi = paragraphs[element]
            paragraph = layout.paragraphs[pi]
            layout.validate_plain_content(element)
            result.append(node("paragraph", {
                "kind": "template_layout", "verbatim": True, "sample_paragraph": pi,
                "role": "heading" if paragraph.text else "spacer",
                **layout.paragraph_style(paragraph),
            }, text=paragraph.text))
        result.append(reference_table(layout, ti, node, "operations" if ti in A14_OPERATIONS else
                                      "weighing" if ti < 15 else "tlc"))
    return result
