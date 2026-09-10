"""Reviewed source-operation bindings to the reference Word's individual form rows."""

import hashlib
import json
from pathlib import Path

from docx.table import _Cell
from fastapi import HTTPException

FORMS_PATH = Path(__file__).with_name("demo") / "operation-forms.json"
UNITS = {"kg", "mg", "g", "L", "%", "ppm", "次", "分钟", "张", "只", "mbar"}


def load_forms():
    return json.loads(FORMS_PATH.read_text())


def grid_cells(table, row):
    """Physical cell indexes at each grid column, without vertical-merge aliasing."""
    result = []
    for index, tc in enumerate(table._tbl.tr_lst[row].tc_lst):
        result.extend([(index, _Cell(tc, table))] * tc.grid_span)
    if len(result) != 5:
        raise ValueError("operation form must have five grid columns")
    return result


def validate_forms(layout):
    for field in layout.forms["fields"].values():
        table = layout.tables[field["table"]]
        if _Cell(table._tbl.tr_lst[0].tc_lst[0], table).text.strip() != "原材料/操作":
            raise ValueError("operation form table moved")
        cells = grid_cells(table, field["row"])
        merged = cells[1][0] == cells[2][0]
        record = "" if merged else cells[2][1].text
        if (cells[1][1].text != field["sample_parameter"]
                or record != field["sample_record"] or merged != field["split_calculation"]):
            raise ValueError(f"operation parameter changed: {field['table']}:{field['row']}")


def recording_text(value):
    # Units label an empty writing area. No instruction quantities become measured values.
    if value.strip() in UNITS:
        return "\u3000\u3000\u3000" + value.strip()
    if value.strip() == "~":
        return "     ~     "
    return value


def operation_table(step, layout, node):
    order = next((p["value"] for p in step["object_data_properties"]
                  if p["label"] == "步骤序号"), None)
    stage = next((s for s in layout.forms["stages"] if s["order"] == order), None)
    if stage is None:
        raise HTTPException(409, "该工艺阶段尚未配置模板参数行")
    bindings = {op["order"]: op for op in stage["operations"]}
    header = layout.table("operations")
    rows = [node("row", {"kind": "template_layout", "prototype_table": 15,
                         "prototype_row": 0}, header=True, children=[
                             node("cell", layout.cell_metadata("operations", 0, c),
                                  text=header.cell(0, c).text) for c in range(5)
                         ])]
    for op in step.get("operations", []):
        binding = bindings.get(op["order"])
        if (binding is None or binding["label"] != op["label"]
                or binding["instruction_sha256"] != hashlib.sha256(
                    op["instruction"].encode(),
                ).hexdigest()):
            raise HTTPException(409, "操作原文已变化，请核对静态模板参数绑定")
        # A source-only introductory instruction remains a row with empty parameter fields.
        keys = binding["fields"] or [None]
        for i, key in enumerate(keys):
            field = layout.forms["fields"][key] if key else None
            ti, ri = (field["table"], field["row"]) if field else (15, 1)
            table = layout.tables[ti]
            physical = grid_cells(table, ri)
            values = [f"{op['order']}、{op['label']}\n{op['instruction']}" if i == 0 else "",
                      field["parameter_name"] if field else "",
                      recording_text(field["record_format"]) if field else "", "", ""]
            cells = []
            for c, value in enumerate(values):
                ci = physical[c][0]
                meta = {**layout.cell_metadata(ti, ri, ci), "span": 1, "prototype_cell": ci}
                if c in (0, 3, 4) and len(keys) > 1:
                    meta.update(vmerge="restart" if i == 0 else "continue",
                                row_span=len(keys) if i == 0 else 0)
                cells.append(node("cell", meta, text=value,
                                  field_id=f"step-{order}-op-{op['order']}-{key}-{c}"))
            height = table.rows[ri].height
            row = node("row", {"kind": "template_layout", "prototype_table": ti,
                               "prototype_row": ri, "operation_order": op["order"],
                               "parameter_key": key, "height_pt": height.pt if height else None},
                       children=cells, record_id=step["entity_id"])
            row.provenance_refs.append({"kind": "source", "source_ref": op["source_ref"]})
            rows.append(row)
    return node("table", layout.table_metadata("operations"), children=rows)
