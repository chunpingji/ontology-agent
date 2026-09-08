"""Render only authorized frozen inputs. No business-class dispatch or source lookup."""

from copy import deepcopy
from decimal import Decimal
from io import BytesIO

from app.services.extraction.evidence_identity import evidence_hash
from app.services.reporting.condition_resolver import evaluate_expression
from app.services.reporting.input_resolver import (
    ResolvedValue,
    business_value,
    project_value,
    select_nodes,
    sort_key,
    typed_value,
)
from app.services.reporting.narrative_renderer import assisted_nodes
from app.services.reporting.output_ast import OutputNode, walk
from app.services.reporting.template_compiler import walk_groups
from app.services.reporting.template_v2 import Format, ReportingError, TemplateV2

RENDERER_VERSION = "output-renderer-v2.1"
STATE_TEXT = {
    "missing": "待补充",
    "confirmed_absent": "已确认不存在",
    "not_applicable": "不适用",
    "pending_review": "待审核",
    "conflict": "数据冲突",
    "invalid": "数据无效",
    "incomplete": "部分数据",
    "unavailable": "来源不可用",
}


def display(value, fmt):
    if value.state != "ready" and value.display_action == "omit":
        return ""
    if value.resolved_type.kind == "entity":
        if fmt.kind == "entity_id":
            return value.entity_id or STATE_TEXT.get(value.state, "待补充")
        label = value.fields.get("display")
        if label:
            return (
                str(label.value)
                if label.state == "ready"
                else STATE_TEXT.get(label.state, "待补充")
            )
        return "未配置显示属性"
    if value.resolved_type.kind == "list":
        texts = [display(item, fmt) for item in value.items]
        suffix = "（部分数据）" if value.discovery and value.discovery.status != "complete" else ""
        return (
            fmt.separator.join(texts) + suffix if texts else STATE_TEXT.get(value.state, "待补充")
        )
    if value.resolved_type.kind == "record":
        return fmt.separator.join(display(child, fmt) for child in value.fields.values())
    if value.state != "ready":
        return STATE_TEXT.get(value.state, "待补充")
    raw, kind = value.value, value.resolved_type.kind
    if kind == "boolean":
        return "是" if raw else "否"
    if kind == "decimal":
        if fmt.decimal_places is not None:
            return format(Decimal(raw), f".{fmt.decimal_places}f")
        return raw
    if kind == "quantity":
        number = (
            format(Decimal(raw["value"]), f".{fmt.decimal_places}f")
            if fmt.decimal_places is not None
            else raw["value"]
        )
        return number + " " + raw["unit"]
    if kind == "range":

        def endpoint(key):
            item = typed_value(value.node_id + "/" + key, value.resolved_type.item_type, raw[key])
            return display(item, fmt)

        return (
            ("[" if raw["lower_inclusive"] else "(")
            + endpoint("lower")
            + "～"
            + endpoint("upper")
            + ("]" if raw["upper_inclusive"] else ")")
        )
    return str(raw)


class UnitRenderer:
    def __init__(self, unit, instance, snapshot, provider):
        self.unit, self.instance, self.snapshot, self.provider = unit, instance, snapshot, provider
        self.prefix = instance["output_instance_id"]
        raw = snapshot["scopes"][instance["execution_scope_id"]]["inputs"]
        self.inputs = {
            use.input_ref: ResolvedValue.model_validate(raw[use.input_ref])
            for use in unit.inputs
            if use.input_ref in raw
        }
        self.catalog = {}
        for value in self.inputs.values():
            for claim in value.derivation.get("claim_catalog", []):
                if (
                    claim.get("eligible")
                    and claim.get("execution_scope_id") == instance["execution_scope_id"]
                ):
                    self.catalog[claim["claim_id"]] = claim
        self.counter, self.diagnostics = 0, []
        self.budget = TemplateV2.model_validate(snapshot["source_bundle"]["template"]).budget

    def node(self, kind, **kwargs):
        self.counter += 1
        if self.counter > self.budget.max_nodes:
            raise ReportingError("OUTPUT_BUDGET_EXCEEDED")
        return OutputNode(node_id=evidence_hash([self.prefix, self.counter]), kind=kind, **kwargs)

    def value(self, input_id, path, fmt, item=None):
        source = item if item is not None else self.inputs.get(input_id)
        if source is None:
            raise ReportingError("OUTPUT_REFERENCE_INVALID")
        selected = project_value(source, path)
        blocked = any(i.blocks_subtree for i in selected.issues)

        def lineage(node):
            yield node
            for child in [*node.items, *node.fields.values()]:
                yield from lineage(child)

        referenced = list(lineage(selected))
        return self.node(
            "value",
            text="数据不可用" if blocked else display(selected, fmt),
            state=selected.state,
            input_ref={
                "input_id": input_id,
                "field_path": path,
                "execution_scope_id": self.instance["execution_scope_id"],
                "record_id": item.entity_id if item else selected.entity_id,
            },
            fact_refs=sorted({ref for node in referenced for ref in node.fact_refs}),
            provenance_refs=[ref for node in referenced for ref in node.provenance_refs],
        )

    def content(self, nodes, items=None):
        items = items or {}
        output = []
        for node in nodes:
            if node.kind == "text":
                output.append(self.node("text", text=node.text))
            elif node.kind == "paragraph":
                output.append(self.node("paragraph", children=self.content(node.children, items)))
            elif node.kind == "input_ref":
                item = items.get(node.input_id) if node.scope == "item" else None
                if node.scope == "item" and item is None:
                    raise ReportingError("INPUT_ITERATION_REQUIRED")
                output.append(self.value(node.input_id, node.field_path, node.format, item))
            elif node.kind == "claim_ref":
                claim = self.catalog.get(node.claim_id)
                if claim is None:
                    raise ReportingError("CLAIM_PRECONDITION_UNPROVEN")
                output.append(
                    self.node(
                        "value",
                        text=claim["text"],
                        claim_ref=node.claim_id,
                        fact_refs=claim["fact_refs"],
                        provenance_refs=claim["provenance_refs"],
                    )
                )
            elif node.kind == "if":
                proof = evaluate_expression(node.condition, self.inputs, items=items)
                selected = {
                    "TRUE": node.children,
                    "FALSE": node.otherwise,
                    "UNKNOWN": node.unknown,
                }[proof["result"]]
                if proof["result"] == "UNKNOWN" and not selected:
                    output.append(self.node("text", text="条件待确认", state="incomplete"))
                else:
                    output.extend(self.content(selected, items))
            elif node.kind == "repeat":
                source = (items if node.scope == "item" else self.inputs).get(node.input_id)
                if source is not None:
                    source = project_value(source, node.field_path)
                if source is None or source.resolved_type.kind != "list":
                    raise ReportingError("OUTPUT_REFERENCE_INVALID")
                if any(i.blocks_subtree for i in source.issues):
                    output.append(self.node("text", text="数据不可用", state=source.state))
                else:
                    for child in source.items:
                        output.extend(self.content(node.children, {**items, node.input_id: child}))
            else:
                output.append(
                    self.node(
                        "signature_region",
                        text="待签署",
                        signature_region_id=node.signature_region_id,
                    )
                )
        return output

    def render(self):
        render = self.unit.render
        metadata = {"mode": "composed", "renderer_version": RENDERER_VERSION}
        status = "completed"
        try:
            if self.unit.when:
                proof = evaluate_expression(self.unit.when, self.inputs)
                if proof["result"] == "FALSE":
                    return self.result([], metadata, "completed", inactive=True)
                if proof["result"] == "UNKNOWN":
                    return self.result(
                        [self.node("paragraph", text="条件待确认", state="incomplete")],
                        metadata,
                        "completed",
                    )
            if render.kind == "table":
                children = [self.table(render)]
            elif render.kind == "form":
                children = []
                for field in render.fields:
                    values = (
                        [self.value(field.value.input_id, field.value.field_path, field.format)]
                        if field.value
                        else []
                    )
                    if field.signature_region_id:
                        values.append(
                            self.node(
                                "signature_region",
                                text="待签署",
                                signature_region_id=field.signature_region_id,
                            )
                        )
                    children.append(
                        self.node(
                            "form_field", text=field.label, field_id=field.field_id, children=values
                        )
                    )
            elif render.kind == "list":
                source = self.inputs.get(render.items.input_id)
                if source is None:
                    raise ReportingError("OUTPUT_REFERENCE_INVALID")
                rows, blocked = select_nodes(source, render.items.field_path)
                if blocked or len(rows) != 1:
                    raise ReportingError("INPUT_CONSUMPTION_BLOCKED")
                children = [
                    self.node(
                        "list",
                        ordered=render.ordered,
                        children=[
                            self.node(
                                "paragraph",
                                record_id=item.entity_id,
                                children=(
                                    self.content(render.nodes, {render.items.input_id: item})
                                    if render.nodes
                                    else [self.value(render.items.input_id, [], Format(), item)]
                                ),
                            )
                            for item in rows[0].items
                        ],
                    )
                ]
            else:
                nodes = render.nodes
                if (
                    render.kind == "narrative"
                    and render.mode == "assisted"
                    and (self.snapshot["source_bundle"].get("preview_mode") == "data")
                ):
                    from app.services.reporting.template_v2 import ContentNode

                    nodes = [
                        ContentNode(
                            kind="input_ref", input_id=ref.input_id, field_path=ref.field_path
                        )
                        for ref in render.prompt.input_refs
                    ]
                    metadata["mode"] = "data_preview"
                elif render.kind == "narrative" and render.mode == "assisted":
                    policy = self.snapshot["source_bundle"]["contracts"][render.prompt.policy_ref][
                        "definition"
                    ]
                    try:
                        nodes, metadata = assisted_nodes(
                            render, self.inputs, self.catalog, policy, self.budget, self.provider
                        )
                    except ReportingError as exc:
                        if render.fallback is None:
                            raise
                        nodes = render.fallback
                        metadata["fallback_reason"] = exc.code
                children = self.content(nodes)
                if children and any(n.kind in {"text", "value"} for n in children):
                    children = [self.node("paragraph", children=children)]
        except ReportingError as exc:
            status = "failed"
            self.diagnostics.append(exc.detail)
            children = [self.node("paragraph", text="此内容尚未生成", state="unavailable")]
        return self.result(children, metadata, status)

    def table(self, render):
        source = self.inputs.get(render.rows.input_id)
        if source is None:
            raise ReportingError("OUTPUT_REFERENCE_INVALID")
        nodes, ancestors = select_nodes(source, render.rows.field_path)
        if ancestors or len(nodes) != 1:
            raise ReportingError("INPUT_CONSUMPTION_BLOCKED")
        rows = nodes[0]
        if any(i.blocks_subtree for i in rows.issues):
            return self.node("paragraph", text="记录范围存在冲突，待核实", state="conflict")
        items = sorted(rows.items, key=lambda row: row.entity_id or "")
        for order in reversed(render.order_by):
            if not order.field_ref:
                if order.direction == "desc":
                    items.reverse()
                continue
            try:
                items.sort(
                    key=lambda row: sort_key(project_value(row, [order.field_ref])),
                    reverse=order.direction == "desc",
                )
            except (ReportingError, TypeError):
                self.diagnostics.append({"code": "TABLE_ORDER_UNRESOLVED"})
        output_rows = [
            self.node(
                "row",
                header=True,
                children=[
                    self.node("cell", text=column.title, field_id=column.column_id, header=True)
                    for column in render.columns
                ],
            )
        ]
        previous_group = None
        if render.group_by:
            try:
                items.sort(
                    key=lambda row: tuple(
                        sort_key(project_value(row, [key])) for key in render.group_by
                    )
                )
            except (ReportingError, TypeError):
                raise ReportingError("TABLE_GROUP_UNRESOLVED")
        for index, row in enumerate(items):
            if render.group_by:
                group_key = tuple(
                    evidence_hash(business_value(row, [key])) for key in render.group_by
                )
                if group_key != previous_group:
                    labels = [
                        self.value(render.rows.input_id, [key], Format(), row)
                        for key in render.group_by
                    ]
                    output_rows.append(
                        self.node(
                            "row",
                            header=True,
                            children=[self.node("cell", header=True, children=labels)]
                            + [self.node("cell", header=True) for _ in render.columns[1:]],
                        )
                    )
                    previous_group = group_key
            cells = []
            for column in render.columns:
                value = (
                    self.node("text", text=str(index + 1))
                    if column.value
                    else self.value(render.rows.input_id, [column.field_ref], column.format, row)
                )
                cells.append(
                    self.node(
                        "cell",
                        field_id=column.field_ref or column.column_id,
                        record_id=row.entity_id,
                        children=[value],
                    )
                )
            output_rows.append(self.node("row", record_id=row.entity_id, children=cells))
        return self.node(
            "table",
            text="部分数据" if rows.discovery and rows.discovery.status != "complete" else "",
            children=output_rows,
        )

    def result(self, children, metadata, status, inactive=False):
        ast = self.node("group", text="" if inactive else self.unit.title, children=children)
        citations = [
            {
                "node_id": node.node_id,
                "input_ref": node.input_ref,
                "fact_refs": node.fact_refs,
                "provenance_refs": node.provenance_refs,
                "claim_ref": node.claim_ref,
            }
            for node in walk(ast)
            if node.input_ref or node.claim_ref
        ]
        return {
            **self.instance,
            "input_snapshot_id": self.snapshot["input_snapshot_id"],
            "output_definition_hash": evidence_hash(self.unit),
            "output_ast": ast.model_dump(mode="json"),
            "citations": citations,
            "consumed_input_refs": [c["input_ref"] for c in citations if c["input_ref"]],
            "generation_metadata": metadata,
            "execution_status": status,
            "validation_issues": self.diagnostics,
            "review_status": "draft",
            "output_hash": evidence_hash(ast),
            "inactive": inactive,
        }


def render_snapshot(snapshot, *, provider=None, completed=None, checkpoint=None):
    from app.services.reporting.calculation_execution import calculation_sections

    checks = snapshot.get("calculation_checks", {})
    calculation_blocked = any(c["blocks_conclusion"] for c in checks.values())
    template = TemplateV2.model_validate(snapshot["source_bundle"]["template"])
    units = {
        unit.output_id: unit
        for section in template.sections
        for group, _ in walk_groups(section.groups)
        for unit in group.units
    }
    results = []
    completed = completed or {}
    for instance in snapshot["output_instances"]:
        result = completed.get(instance["output_instance_id"])
        if calculation_blocked:
            renderer = UnitRenderer(units[instance["output_id"]], instance, snapshot, None)
            result = renderer.result(
                [renderer.node("paragraph", text="PDE 校验或异议处理未完成，报告结论待评估。")],
                {"mode": "calculation_gate"},
                "completed",
            )
        if result is None:
            result = UnitRenderer(
                units[instance["output_id"]], instance, snapshot, provider
            ).render()
        if checkpoint:
            checkpoint(result)
        results.append(deepcopy(result))
    sections = []
    for section in template.sections:

        def assemble(groups):
            output = []
            for group in groups:
                children = [
                    OutputNode.model_validate(r["output_ast"])
                    for r in results
                    if r["group_id"] == group.group_id and not r.get("inactive")
                ]
                children.extend(assemble(group.groups))
                output.append(
                    OutputNode(
                        node_id=group.group_id, kind="group", text=group.title, children=children
                    )
                )
            return output

        sections.append(
            OutputNode(
                node_id=section.section_id,
                kind="section",
                text=section.title,
                children=assemble(section.groups),
            )
        )
    if calculation_blocked:
        sections = [
            OutputNode(
                node_id="calculation-conclusion-gate",
                kind="paragraph",
                state="conflict",
                text="PDE 校验或异议处理尚未完成。本次仅输出校验记录；总体风险结论为待评估。",
            )
        ]
    sections.extend(calculation_sections(checks))
    ast = OutputNode(
        node_id=evidence_hash([snapshot["input_snapshot_id"], "body"]),
        kind="document",
        children=sections,
    )
    return {
        "output_results": results,
        "body_ast": ast.model_dump(mode="json"),
        "body_hash": evidence_hash(ast),
        "calculation_checks": checks,
        "execution_status": "failed"
        if any(r["execution_status"] == "failed" for r in results)
        else "completed",
        "material_status": snapshot["material_status"],
        "review_status": "draft",
    }


def render_docx(ast, style=None):
    from docx import Document
    from docx.shared import Pt

    from app.services.reporting.output_ast import plain_text

    ast = OutputNode.model_validate(ast)
    doc = Document()
    style = style or {}
    doc.styles["Normal"].font.name = style.get("font", "Arial")
    doc.styles["Normal"].font.size = Pt(style.get("font_size_pt", 11))

    def emit(node):
        if node.kind in {"section", "group"}:
            if node.text:
                doc.add_heading(node.text, level=1 if node.kind == "section" else 2)
        elif node.kind == "table":
            if node.text:
                doc.add_paragraph(node.text)
            rows = node.children
            if rows:
                table = doc.add_table(rows=0, cols=max(len(r.children) for r in rows))
                table.style = "Table Grid"
                for row in rows:
                    cells = table.add_row().cells
                    for index, cell in enumerate(row.children):
                        cells[index].text = plain_text(cell)
            return
        elif node.kind in {
            "paragraph",
            "form_field",
            "signature_region",
            "signature",
            "text",
            "value",
        }:
            doc.add_paragraph(plain_text(node))
            return
        elif node.kind == "list":
            for child in node.children:
                doc.add_paragraph(
                    plain_text(child), style="List Number" if node.ordered else "List Bullet"
                )
            return
        for child in node.children:
            emit(child)

    emit(ast)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
