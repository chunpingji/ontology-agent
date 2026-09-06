"""One canonical Word analysis entrypoint for analysis, templates and extraction."""

from dataclasses import dataclass
from pathlib import Path

from app.services.extraction.document_ir import DocumentIR, build_document_ir
from app.services.extraction.docx_structure import DocStructure, parse_docx_structure


@dataclass
class WordAnalysis:
    structure: DocStructure
    ir: DocumentIR


def analyze_word_core(file_path: str | Path, source_filename: str | None = None, *,
                      role: str = "analysis_source", original_path: str | Path | None = None
                      ) -> WordAnalysis:
    structure = parse_docx_structure(file_path, source_filename=source_filename)
    ir = build_document_ir(file_path, structure, role=role, original_path=original_path)
    return WordAnalysis(structure, ir)


def attach_evidence_preview(content: dict, ir: DocumentIR, document) -> dict:
    """Render canonical source-cell topology; OOXML is used only for visual styles."""
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    from app.services.extraction.document_annotator import (
        _apply_cell_grid,
        _inline_nodes,
        _para_runs_and_text,
        _table_grid_px,
    )

    by_block = {u.block_id: u for u in ir.evidence_units if u.table_path is None}
    by_cell = {
        (u.source_cell_id, u.paragraph_index): u for u in ir.evidence_units if u.table_path
    }

    def attrs(unit):
        return {
            "evidenceId": unit.evidence_id, "sourceBlockId": unit.block_id,
            "sourceParagraphIndex": unit.paragraph_index, "sectionNodeId": unit.section_node_id,
            "sourceTablePath": unit.table_path, "sourceRowIndex": unit.row_index,
            "sourceColumnIndex": unit.column_index,
            "physicalPageNumber": unit.physical_page_number,
        }

    def render_table(table: dict, source, existing=None):
        cell_elements = {}
        for ri, row in enumerate(source.rows):
            column = row.grid_cols_before
            for tc in row._tr.tc_lst:
                cell_elements[(ri, column)] = tc
                column += tc.grid_span
        rows = [[] for _ in table["grid"]]
        grid_px = _table_grid_px(source)
        for cell in table["source_cells"]:
            ri, ci = cell["row_index"], cell["column_index"]
            tc = cell_elements[(ri, ci)]
            wrapper = _Cell(tc, source)
            paragraphs = wrapper.paragraphs
            nested_tables = iter(wrapper.tables)
            children = []
            for block in cell["blocks"]:
                if block["kind"] == "table":
                    nested = next(nested_tables)
                    children.append(render_table(block["table"], Table(nested._tbl, wrapper)))
                else:
                    unit = by_cell[(cell["cell_id"], block["paragraph_index"])]
                    paragraph = Paragraph(paragraphs[unit.paragraph_index]._p, wrapper)
                    visual_text, runs = _para_runs_and_text(paragraph, rich=True)
                    # Style mismatches may not rewrite the source evidence.
                    if visual_text != unit.text:
                        runs = []
                    children.append({
                        "type": "paragraph", "attrs": attrs(unit),
                        "content": _inline_nodes(unit.text, [], runs),
                    })
            node = {"type": "tableCell", "content": children or [{"type": "paragraph"}]}
            _apply_cell_grid(node, tc, ri, ci, grid_px)
            node.setdefault("attrs", {}).update(
                colspan=cell["column_span"], rowspan=cell["row_span"],
            )
            rows[ri].append(node)
        return {
            "type": "table", "attrs": {
                **(existing or {}).get("attrs", {}), "sourceTablePath": table["table_path"],
            },
            "content": [{"type": "tableRow", "content": row} for row in rows],
        }

    rendered = []
    for node in content.get("content", []):
        attributes = node.get("attrs", {})
        if node["type"] == "table":
            index = attributes["sourceTableIndex"]
            node = render_table(ir.tables[index], document.tables[index], node)
        else:
            unit = by_block.get(attributes.get("sourceBlockId"))
            if unit:
                node.setdefault("attrs", {}).update(attrs(unit))
        rendered.append(node)
    return {**content, "content": rendered, "analysis": ir.model_dump(mode="json")}
