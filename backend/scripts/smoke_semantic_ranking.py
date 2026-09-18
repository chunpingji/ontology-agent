#!/usr/bin/env python3
"""Replayable CPU/CUDA 12 ranking smoke with a private scheduler DB.

Run this file with the backend environment. All model paths and the fresh output
path must be absolute. Thread limits are inherited from OMP_NUM_THREADS and
MKL_NUM_THREADS. Only supplied ranking models and a new SQLite scheduler are used.
This checks execution and persistence, not gold-standard graph quality.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import sys
import time
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))


def _absolute_path(value: str) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        raise argparse.ArgumentTypeError("an absolute path is required")
    return path


def _bounded_int(minimum: int, maximum: int):
    def convert(value):
        number = int(value)
        if not minimum <= number <= maximum:
            raise argparse.ArgumentTypeError(f"must be between {minimum} and {maximum}")
        return number

    return convert


def _device_arguments(result):
    result.add_argument("--device", default="cpu", help="cpu (default) or an explicit cuda:N.")
    result.add_argument("--dtype", choices=("float32", "float16"), default="float32")
    result.add_argument("--cuda-version", choices=("12.6",), default="12.6")


def _memory_summary(observations):
    return {
        key: max((item.get("worker_metrics", {}).get(key, 0) for item in observations), default=0)
        for key in ("cuda_peak_allocated_bytes", "cuda_peak_reserved_bytes")
    }


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__)
    _device_arguments(result)
    result.add_argument("--output-dir", type=_absolute_path, required=True,
                        help="New output directory; an existing directory is never overwritten.")
    for kind in ("embedding", "reranker"):
        result.add_argument(f"--{kind}-path", type=_absolute_path, required=True)
        result.add_argument(f"--{kind}-manifest-path", type=_absolute_path, required=True)
    result.add_argument("--batch-size", type=_bounded_int(1, 256), default=4)
    result.add_argument("--max-tokens-per-pair", type=_bounded_int(1, 32768), default=2048)
    result.add_argument("--timeout-seconds", type=_bounded_int(1, 3600), default=600)
    result.add_argument("--long-record-repetitions", type=_bounded_int(1, 128), default=4)
    result.add_argument("--length-boundaries", action="store_true",
                        help="Test near-limit success, oversize rejection and worker recovery.")
    result.add_argument("--length-boundary-batch-size", type=_bounded_int(1, 256), default=1,
                        help="Number of near-limit records per boundary request, <= batch-size.")
    return result


def _write_json(path: Path, payload) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
                    encoding="utf-8")


def _require(condition: bool, reason: str) -> None:
    if not condition:
        raise RuntimeError(reason)


def _hardware() -> dict:
    cpu_model = platform.processor()
    cpu_info = Path("/proc/cpuinfo")
    if cpu_info.is_file():
        cpu_model = next((line.split(":", 1)[1].strip()
                          for line in cpu_info.read_text().splitlines()
                          if line.startswith("model name")), cpu_model)
    return {
        "platform": platform.platform(), "machine": platform.machine(),
        "python": platform.python_version(), "cpu_model": cpu_model,
        "logical_cpu_count": os.cpu_count(),
        "cpu_affinity_count": (
            len(os.sched_getaffinity(0)) if hasattr(os, "sched_getaffinity") else None
        ),
        "threads_environment": {name: os.environ.get(name)
                                for name in ("OMP_NUM_THREADS", "MKL_NUM_THREADS")},
    }


def _synthetic_case(output: Path, repetitions: int, run_id: str):
    from docx import Document

    from app.services.extraction.evidence_identity import evidence_hash
    from app.services.extraction.ontology_guided.contracts import (
        EdgeSpec,
        GraphNode,
        SubjectRef,
        VersionedRef,
    )
    from app.services.extraction.ontology_guided.metadata import prepare_metadata
    from app.services.extraction.ontology_guided.records import RecordIndex
    from app.services.extraction.ontology_guided.retrieval import plan_slot
    from app.services.extraction.word_analysis import analyze_word_core

    short = "本合成报告描述产品甲片。"
    long = "合成产品甲片完成溶出试验，记录仅供排序冒烟使用。" * repetitions
    paragraphs = [short, "产品甲片不含合成成分乙。", "在合成条件丙下，产品甲片改变包装。",
                  "会议室桌椅按行政日程进行整理。", long]
    document = Document()
    document.add_heading("合成产品描述与试验记录", 1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    source = output / "synthetic.docx"
    document.save(source)
    _write_json(output / "inputs.json", {
        "origin": "synthetic", "paragraphs": paragraphs,
        "probe_texts": [short, long],
        "probe_pairs": [("报告描述哪种合成产品？", short),
                        ("报告描述哪种合成产品？", paragraphs[3])],
        "long_record_repetitions": repetitions,
    })
    analysis = analyze_word_core(source)
    index = RecordIndex(analysis.ir)
    metadata = prepare_metadata(
        analysis.ir, section_tree=analysis.structure.section_tree.to_dict(),
        summary_version="synthetic-smoke-v1", generation_source="structure_only",
    )
    subject = SubjectRef(entity_id="synthetic-report", revision=1,
                         class_iri="urn:smoke:Report", is_document_root=True)
    node = GraphNode(entity_id=subject.entity_id, revision=1, class_iri=subject.class_iri,
                     class_label="报告", label=source.name, root=True, root_origin="user_specified")
    predicate = EdgeSpec(iri="urn:smoke:describes", label="描述", description="报告描述合成产品。",
                         range_class_iris=["urn:smoke:Product"], declared_by=[subject.class_iri])
    plan = plan_slot(subject, predicate, index, metadata, ontology_hash=evidence_hash(predicate))
    _write_json(output / "document_ir.json", analysis.ir.model_dump(mode="json"))
    _write_json(output / "metadata.json", metadata.model_dump(mode="json"))
    _write_json(output / "retrieval_plan.json", plan.model_dump(mode="json"))
    return {
        "plan": plan, "index": index, "metadata": metadata, "subject_node": node,
        "predicate": predicate,
        "run_fingerprint": evidence_hash([run_id, analysis.ir.document_hash]),
        "root_ref": VersionedRef(id=node.entity_id, revision=1),
        "root_class_iri": node.class_iri, "permission_scope": run_id,
    }


def _scheduler_report(bind, run_id: str) -> list[dict]:
    from sqlalchemy import select
    from sqlalchemy.orm import Session

    from app.models.model_request import LocalModelRequest

    with Session(bind) as db:
        return [{"sequence": row.sequence, "request_id": row.request_id, "run_id": row.run_id,
                 "stage": row.stage, "status": row.status, "metrics": row.metrics}
                for row in db.scalars(select(LocalModelRequest).where(
                    LocalModelRequest.run_id == run_id).order_by(LocalModelRequest.sequence))]



def _length_boundaries(model, args, bind, report: dict, output: Path) -> None:
    from app.services.llm.semantic_ranking import RankingModelUnavailable

    started = time.monotonic()
    limit = args.max_tokens_per_pair
    boundary_batch_size = args.length_boundary_batch_size
    _require(boundary_batch_size <= args.batch_size, "length_boundary_batch_exceeds_configuration")
    _require(limit >= 128, "length_boundary_requires_at_least_128_tokens")
    phrase = "合成产品甲片检验记录。"
    # Count with the real, untruncated model tokenizer. Leave a small allowance
    # for the query and pair special tokens, then verify actual worker counts.
    high = 1
    while model.count_tokens(phrase * high) <= limit + 64:
        high *= 2
        _require(high <= 131072, "length_boundary_token_search_exceeded")
    oversized = phrase * high
    oversized_tokens = model.count_tokens(oversized)
    low, upper = 0, high
    while low + 1 < upper:
        middle = (low + upper) // 2
        if model.count_tokens(phrase * middle) <= limit - 32:
            low = middle
        else:
            upper = middle
    near = phrase * low
    near_tokens = model.count_tokens(near)
    _require(near_tokens >= limit * 0.95, "near_limit_input_too_short")
    _write_json(output / "length_boundary_inputs.json", {
        "near_limit_text": near, "near_limit_tokenizer_max": near_tokens,
        "oversized_text": oversized, "oversized_tokenizer_max": oversized_tokens,
        "max_tokens_per_pair": limit, "pair_query": "产品",
    })
    successes = []
    for operation, inputs in (("embed", [near]), ("score_pairs", [("产品", near)])):
        inputs *= boundary_batch_size
        result = getattr(model, operation)(inputs)
        _require(len(result) == boundary_batch_size, "near_limit_result_missing")
        measured = model.observations[-1]["input_tokens"]
        _require(isinstance(measured, int) and measured <= limit * boundary_batch_size,
                 "near_limit_actual_tokens_exceeded")
        successes.append({"operation": operation, "measured_input_tokens": measured,
                          "batch_size": boundary_batch_size,
                          "worker_metrics": model.observations[-1].get("worker_metrics", {})})
    _require(max(item["measured_input_tokens"] for item in successes)
             >= limit * 0.95 * boundary_batch_size,
             "near_limit_actual_tokens_too_short")
    boundary = {"configured_limit": limit, "near_limit_tokenizer_max": near_tokens,
                "oversized_tokenizer_max": oversized_tokens, "successful_operations": successes,
                "expected_rejections": [], "recovery": []}
    report["length_boundaries"] = boundary
    for operation, inputs in (("embed", [oversized]),
                              ("score_pairs", [("产品", oversized)])):
        before = len(_scheduler_report(bind, report["run_id"]))
        try:
            getattr(model, operation)(inputs)
        except RankingModelUnavailable as exc:
            _require(str(exc) == "ranking_input_too_long", "oversize_failed_for_unexpected_reason")
        else:
            raise RuntimeError("oversize_input_was_not_rejected")
        requests = _scheduler_report(bind, report["run_id"])[before:]
        _require(len(requests) == 1 and requests[0]["status"] == "failed",
                 "oversize_rejection_request_not_recorded")
        _require(getattr(model, "_process", None) is None, "failed_worker_not_cleared")
        boundary["expected_rejections"].append({
            "request_id": requests[0]["request_id"], "operation": operation,
            "reason": "ranking_input_too_long", "worker_cleared": True,
        })
    for operation, inputs in (("embed", ["合成产品"]),
                              ("score_pairs", [("产品", "合成产品")])):
        _require(len(getattr(model, operation)(inputs)) == 1, "worker_recovery_result_missing")
        _require(model.observations[-1]["status"] == "completed", "worker_recovery_failed")
        boundary["recovery"].append({"operation": operation, "status": "completed"})
    boundary["elapsed_seconds"] = time.monotonic() - started
    _write_json(output / "length_boundaries.json", boundary)


def _execute(args, output: Path, report: dict, model_factory=None) -> None:
    from sqlalchemy import create_engine, inspect

    from app.models.model_request import LocalModelPool, LocalModelRequest
    from app.services.extraction.ontology_guided.semantic_reranker import (
        RankingPolicy,
        RankingService,
    )
    from app.services.llm.model_runtime import model_scope
    from app.services.llm.semantic_ranking import LocalSemanticRanking

    config = {f"{kind}_{suffix}": str(getattr(args, f"{kind}_{suffix}"))
              for kind in ("embedding", "reranker") for suffix in ("path", "manifest_path")}
    config.update(mode="rerank", batch_size=args.batch_size,
                  max_tokens_per_pair=args.max_tokens_per_pair,
                  timeout_seconds=args.timeout_seconds, device=args.device,
                  dtype=args.dtype, cuda_version=args.cuda_version)
    report["configuration"] = config
    bind = create_engine(f"sqlite:///{output / 'scheduler.sqlite3'}")
    LocalModelPool.__table__.create(bind)
    LocalModelRequest.__table__.create(bind)
    report["scheduler_tables"] = inspect(bind).get_table_names()
    model, restored_model, service = None, None, None
    factory = model_factory or LocalSemanticRanking
    try:
        arguments = _synthetic_case(output, args.long_record_repetitions, report["run_id"])
        inputs = json.loads((output / "inputs.json").read_text(encoding="utf-8"))
        model = factory(config)
        report["model_identity"] = model.identity
        with model_scope(bind=bind, run_id=report["run_id"], task_id="synthetic-ranking-smoke"):
            phase_started = time.monotonic()
            counts = [model.count_tokens(text) for text in inputs["probe_texts"]]
            _require(all(isinstance(count, int) and count > 0 for count in counts),
                     "tokenizer_did_not_return_positive_counts")
            _require(counts[1] > counts[0], "long_record_tokenization_not_exercised")
            vectors, scores = [], []
            for position in range(0, len(inputs["probe_texts"]), args.batch_size):
                vectors.extend(model.embed(
                    inputs["probe_texts"][position:position + args.batch_size]))
            for position in range(0, len(inputs["probe_pairs"]), args.batch_size):
                scores.extend(model.score_pairs(
                    inputs["probe_pairs"][position:position + args.batch_size]))
            norms = [math.sqrt(sum(value * value for value in vector)) for vector in vectors]
            tolerance = 1e-3 if args.dtype == "float16" else 1e-4
            _require(len(vectors) == 2 and all(math.isclose(norm, 1.0, abs_tol=tolerance)
                                             for norm in norms), "embedding_not_l2_normalized")
            _require(len(scores) == 2 and all(math.isfinite(value) for value in scores),
                     "raw_reranker_scores_missing_or_nonfinite")
            _require(model.identity.get("score_semantics") == "raw_single_logit",
                     "reranker_score_semantics_not_raw_logits")
            report["probe"] = {"token_counts": counts, "embedding_dimensions": len(vectors[0]),
                               "embedding_norms": norms, "raw_reranker_scores": scores,
                               "elapsed_seconds": time.monotonic() - phase_started,
                               "temperature": "cold_load_then_direct_inference"}
            _write_json(output / "probe_outputs.json", {"vectors": vectors, **report["probe"]})
            if args.length_boundaries:
                _length_boundaries(model, args, bind, report, output)
            policy = RankingPolicy(
                mode="semantic", failure_policy="pause", pool_size=len(arguments["index"].records),
                batch_size=args.batch_size, max_tokens_per_pair=args.max_tokens_per_pair,
                ranking_timeout=args.timeout_seconds, technical_retry_limit=0,
            )
            service = RankingService(policy, model)
            service.before_model_hook = lambda state: _write_json(
                output / "ranking_reserved_state.json", state)
            phase_started = time.monotonic()
            epoch = service.prepare_next_epoch(**arguments)
            _require(epoch is not None, "no_ranking_epoch")
            _write_json(output / "ranking_prepared_epoch.json", epoch.model_dump(mode="json"))
            _require(epoch.actual_ranking_mode == "semantic" and not epoch.degraded,
                     f"semantic_ranking_did_not_complete:{epoch.reason}")
            _require(set(epoch.record_ids) == set(arguments["plan"].frozen_record_ids),
                     "whole_record_pool_not_ranked")
            _require(len(epoch.observations) == 2 * len(epoch.record_ids),
                     "both_query_intents_not_fully_ranked")
            _require(all(item["raw_rerank_score"] is not None for item in epoch.observations),
                     "pool_raw_scores_missing")
            _require(all(not view.get("omitted_refs") for view in epoch.retrieval_views),
                     "source_records_were_truncated")
            committed = service.commit_epoch(epoch)
            _require(committed.status == "committed", "epoch_not_committed")
            _write_json(output / "ranking_epoch.json", committed.model_dump(mode="json"))
            state = service.snapshot()
            _write_json(output / "ranking_state.json", state)
            report["ranking"] = {
                "status": committed.status, "actual_ranking_mode": committed.actual_ranking_mode,
                "degraded": committed.degraded, "records": len(committed.record_ids),
                "observations": len(committed.observations), "costs": committed.costs,
                "elapsed_seconds": time.monotonic() - phase_started,
                "temperature": "warm_model_weights",
            }
            _require(state["costs"]["tokens"] > 0 and state["costs"]["model_calls"] > 0,
                     "ranking_reserved_costs_missing")
            before_restore = len(_scheduler_report(bind, report["run_id"]))
            model.close()
            restored_model = factory(config)
            phase_started = time.monotonic()
            restored = RankingService(policy, restored_model, state=json.loads(
                (output / "ranking_state.json").read_text(encoding="utf-8")))
            _require(restored.prepare_next_epoch(**arguments) is None,
                     "restored_full_pool_prepared_again")
            _require(restored.snapshot() == state, "restored_state_changed")
            after_restore = len(_scheduler_report(bind, report["run_id"]))
            _require(after_restore == before_restore and not restored_model.observations,
                     "restore_recalled_model")
            report["restore"] = {"extra_scheduler_requests": after_restore - before_restore,
                                 "state_unchanged": True,
                                 "elapsed_seconds": time.monotonic() - phase_started}
            _write_json(output / "ranking_restored_state.json", restored.snapshot())
    finally:
        if model is not None:
            model.close()
            report["model_observations"] = model.observations
            report["gpu_memory"] = _memory_summary(model.observations)
            _write_json(output / "model_observations.json", model.observations)
        if restored_model is not None:
            restored_model.close()
        requests = _scheduler_report(bind, report["run_id"])
        _write_json(output / "scheduler_requests.json", requests)
        expected_rejection_ids = {
            item["request_id"] for item in report.get("length_boundaries", {}).get(
                "expected_rejections", [])
        }
        report["observed_costs"] = {
            "expected_rejection_requests": len(expected_rejection_ids),
            "unexpected_unknown_request_count": sum(
                row["metrics"].get("input_tokens") is None
                and row["request_id"] not in expected_rejection_ids for row in requests
            ),
            "scheduler_requests": len(requests),
            "completed_requests": sum(row["status"] == "completed" for row in requests),
            "measured_input_tokens": sum(row["metrics"].get("input_tokens") or 0
                                         for row in requests),
            "unknown_request_count": sum(row["metrics"].get("input_tokens") is None
                                         for row in requests),
            "queue_seconds": sum(row["metrics"].get("queue_seconds") or 0 for row in requests),
            "request_seconds": sum(row["metrics"].get("request_seconds") or 0 for row in requests),
        }
        if service is not None:
            _write_json(output / "ranking_final_state.json", service.snapshot())
        bind.dispose()
    _require(requests and all(
        row["status"] == "completed" or (
            row["request_id"] in expected_rejection_ids and row["status"] == "failed"
        ) for row in requests
    ), "scheduler_has_unexpected_failed_or_incomplete_requests")
    _require(report["observed_costs"]["measured_input_tokens"] > 0,
             "actual_input_token_cost_missing")
    _require(report["observed_costs"]["unexpected_unknown_request_count"] == 0,
             "actual_input_token_cost_incomplete")


def run_smoke(args, *, model_factory=None) -> dict:
    output = args.output_dir
    for model_path in (args.embedding_path, args.reranker_path):
        if output.resolve().is_relative_to(model_path.resolve()):
            raise ValueError("output directory must be outside immutable model artifacts")
    # A new run never replaces a failed or successful earlier run's evidence.
    output.mkdir(parents=True, exist_ok=False)
    started = time.monotonic()
    report = {
        "schema_version": "semantic-ranking-smoke-v2", "run_id": uuid4().hex,
        "started_at": datetime.now(UTC).isoformat(), "status": "running", "hardware": _hardware(),
        "scope": {"input_origin": "synthetic", "device": args.device, "dtype": args.dtype,
                  "main_llm_called": False,
                  "business_facts_written": False, "gold_quality_evaluation": False,
                  "scheduler_database": str(output / "scheduler.sqlite3")},
        "command": [sys.executable, str(Path(__file__).resolve()), *sys.argv[1:]],
    }
    overrides = {"DATABASE_URL": f"sqlite:///{output / 'scheduler.sqlite3'}",
                 "HF_HUB_OFFLINE": "1", "TRANSFORMERS_OFFLINE": "1",
                 "TOKENIZERS_PARALLELISM": "false"}
    previous = {key: os.environ.get(key) for key in overrides}
    os.environ.update(overrides)
    try:
        _execute(args, output, report, model_factory=model_factory)
        report["status"] = "passed"
    except Exception as exc:
        report["status"] = "failed"
        report["failure"] = {"type": type(exc).__name__, "reason": str(exc)}
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        try:
            import resource

            report["child_memory"] = {
                "max_rss": resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss,
                "unit": "bytes" if platform.system() == "Darwin" else "KiB",
                "scope": "largest waited-for child process in this parent lifetime",
                "is_sum_of_process_memory": False,
            }
        except ImportError:
            report["child_memory"] = {"status": "unavailable_on_this_platform"}
        report["elapsed_seconds"] = time.monotonic() - started
        report["finished_at"] = datetime.now(UTC).isoformat()
        report["artifact_sha256"] = {
            path.name: hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted(output.iterdir()) if path.is_file() and path.name != "report.json"
        }
        _write_json(output / "report.json", report)
    return report


def main(argv=None) -> int:
    args = parser().parse_args(argv)
    if args.output_dir.exists():
        parser().error("--output-dir must be a new directory")
    report = run_smoke(args)
    print(json.dumps({"status": report["status"], "report": str(args.output_dir / "report.json"),
                      "failure": report.get("failure")}, ensure_ascii=False, allow_nan=False))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
