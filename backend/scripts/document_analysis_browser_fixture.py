#!/usr/bin/env python3
"""Serve two isolated persisted runs through the real application for UI acceptance.

This deliberately seeds controlled model responses, not quality measurements.
The caller must provision an empty loopback PostgreSQL database ending in
``_browser_test``. Normal application lifespan and background workers stay off.
"""

from __future__ import annotations

import argparse
import asyncio
import io
import json
import os
import secrets
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
DEV = "https://ontology.pharma-gmp.cn/slpra/drug-development/"
EQUIPMENT = "https://ontology.pharma-gmp.cn/slpra/equipment/Equipment"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database-url", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--port", type=int, default=58001)
    parser.add_argument("--frontend-origin", default="http://127.0.0.1:53101")
    args = parser.parse_args()
    from sqlalchemy.engine import make_url

    url = make_url(args.database_url)
    if (
        url.get_backend_name() != "postgresql"
        or url.host not in {"127.0.0.1", "localhost"}
        or not (url.database or "").endswith("_browser_test")
    ):
        parser.error("requires a dedicated loopback PostgreSQL *_browser_test database")
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    os.environ.update(
        {
            "DATABASE_URL": args.database_url,
            "ONTOLOGY_DIR": str(BACKEND.parent / "ontology" / "slpra"),
            "OWL_STORE_PATH": str(output / "ontology.sqlite3"),
            "DOCUMENT_ANALYSIS_STORAGE_DIR": str(output / "run-artifacts"),
            "AUTH_REQUIRED": "true",
            "AUTH_SECRET": secrets.token_hex(32),
            "LOCAL_LLM_ENABLED": "false",
            "LLM_CLOUD_ENABLED": "false",
            "SEMANTIC_ALIGNMENT_ENABLED": "false",
            "GLINER_EXTRACTION_ENABLED": "false",
            "LLM_WORD_TREE_SUMMARY_ENABLED": "false",
        }
    )
    import uvicorn
    from docx import Document
    from fastapi import UploadFile
    from fastapi.middleware.cors import CORSMiddleware
    from fastapi.responses import JSONResponse
    from sqlalchemy import func, inspect, select

    import app.models  # noqa: F401
    from app.auth import hash_password
    from app.db import Base, SessionLocal, engine
    from app.main import app
    from app.models.document_analysis import DocumentAnalysisRun
    from app.models.ontology_meta import ROLE_NAMES, AppRole, AppUser
    from app.services.document_analysis import execution
    from app.services.document_analysis.application import DocumentAnalysisApplication
    from app.services.extraction.ontology_guided import model_adapter
    from app.services.extraction.ontology_guided.semantic_reranker import (
        RankingPolicy,
        RankingService,
    )
    from app.services.ontology_engine import ontology_engine

    if inspect(engine).get_table_names():
        raise RuntimeError(
            "browser fixture requires an empty database; existing tables are protected"
        )
    Base.metadata.create_all(engine)
    ontology_engine.load()
    calls = []

    def controlled_response(_client, *, user, **_kwargs):
        request = json.loads(user)
        calls.append(request["stage"])
        target = next(
            (item for item in request["fragments"] if "本报告明确使用设备" in item["text"]), None
        )
        if request["stage"] == "verification":
            return {
                "verifications": [
                    {
                        "candidate_id": candidate["candidate_id"],
                        "target_id": candidate["target_id"],
                        "type_verdict": "supported",
                        "role_verdict": "supported",
                        "predicate_verdict": "supported",
                        "applicability_verdict": "supported",
                        "subject_binding_verdict": "supported",
                        "counterevidence_verdict": "supported",
                        "bridge_verdict": "supported",
                        "predicate_support": [
                            {"evidence_id": target["evidence_id"], "text": "使用设备"}
                        ],
                        "subject_support": (
                            []
                            if request["subject"]["is_document_root"]
                            else [{"evidence_id": target["evidence_id"], "text": "本报告"}]
                        ),
                        "condition_support": [],
                        "counterevidence_support": [],
                        "reason": "隔离浏览器制品：受控独立验证返回原文明示关系。",
                    }
                    for candidate in request["candidates"]
                ]
            }
        if request["predicate"]["iri"] != DEV + "usesEquipment" or target is None:
            return {"proposals": []}
        label = "冻干机甲" if "冻干机甲" in target["text"] else "冻干机乙"
        return {
            "proposals": [
                {
                    "kind": "relationship",
                    "object_class_iri": EQUIPMENT,
                    "object_label": label,
                    "object_quote": {"evidence_id": target["evidence_id"], "text": label},
                    "bridge_kind": "explicit_assertion",
                    "polarity": "affirmed",
                }
            ]
        }

    class ControlledRanker:
        identity = {"model": "browser-controlled-ranking-v1"}

        def count_tokens(self, text):
            return len(text)

        def embed(self, texts):
            return [[1.0, 0.0] for _ in texts]

        def score_pairs(self, pairs):
            return [-float(len(text)) for _query, text in pairs]

    model_adapter.chat_with_schema = controlled_response
    execution.configured_model_adapter = lambda: model_adapter.LocalModelRecognitionAdapter(
        object(), model_identity="browser-controlled-discovery-and-verification-v1"
    )
    username, password = "browser_reader", secrets.token_urlsafe(24)
    manifest = {
        "scope": "real API/browser integration with controlled persisted recognition fixtures",
        "backend_origin": f"http://127.0.0.1:{args.port}",
        "frontend_origin": args.frontend_origin,
        "runs": [],
    }
    with SessionLocal() as db:
        for name in ROLE_NAMES:
            db.add(AppRole(name=name, description=name))
        db.commit()
        role = db.scalar(select(AppRole).where(AppRole.name == "operator"))
        db.add(
            AppUser(
                username=username,
                display_name="Browser acceptance reader",
                role_id=role.id,
                password_hash=hash_password(password),
            )
        )
        db.commit()
        for number, label in enumerate(("冻干机甲", "冻干机乙"), 1):
            policy = RankingPolicy(mode="semantic")
            ranker = ControlledRanker() if number == 1 else None
            execution._configured_ranking = lambda: (
                RankingService(policy, ranker),
                {"policy": policy.model_dump(), "model": ranker.identity if ranker else None},
            )
            document = Document()
            document.add_heading(f"浏览器隔离报告{number}", level=1)
            document.add_paragraph(f"本报告明确使用设备{label}进行生产。")
            data = io.BytesIO()
            document.save(data)
            data.seek(0)
            service = DocumentAnalysisApplication(db, ontology_engine=ontology_engine)
            run, created = asyncio.run(
                service.create_run(
                    owner_id=username,
                    file=UploadFile(file=data, filename=f"隔离报告{number}.docx"),
                    root_class_iri=DEV + "CMCReport",
                    request_key=f"browser-{number}",
                    metadata_mode="structure_only",
                )
            )
            assert created
            token = service.store.claim(
                run.recognition_run_id,
                username,
                actor="browser-fixture",
                worker_id="isolated-fixture",
                lease_seconds=120,
            )
            db.commit()
            execution._execute_claimed(db, service.store, run, token)
            graph = service.graph_response(run, projection="effective_affirmed")
            if not graph["relationships"]:
                raise RuntimeError(
                    f"controlled fixture did not produce a proven relationship: {run.error}"
                )
            manifest["runs"].append(
                {
                    "run_id": str(run.recognition_run_id),
                    "label": label,
                    "filename": run.filename,
                    "ranking": graph["ranking"],
                }
            )

    def snapshot():
        with SessionLocal() as db:
            return {
                "table_counts": {
                    name: db.scalar(select(func.count()).select_from(table))
                    for name, table in sorted(Base.metadata.tables.items())
                },
                "runs": [
                    {
                        "id": str(run.recognition_run_id),
                        "revision": run.revision,
                        "event_head": run.event_head,
                        "artifact_revision": run.artifact_revision,
                        "model_calls": run.progress.get("model_calls"),
                    }
                    for run in db.scalars(
                        select(DocumentAnalysisRun).order_by(DocumentAnalysisRun.recognition_run_id)
                    )
                ],
                "controlled_requests": len(calls),
            }

    manifest["baseline"] = snapshot()
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    credentials = output / "credentials.json"
    credentials.write_text(json.dumps({"username": username, "password": password}))
    credentials.chmod(0o600)
    request_log = output / "api-requests.jsonl"

    @app.middleware("http")
    async def read_only(request, call_next):
        with request_log.open("a") as log:
            log.write(json.dumps({"method": request.method, "path": request.url.path}) + "\n")
        if (
            request.method not in {"GET", "HEAD", "OPTIONS"}
            and request.url.path != "/api/auth/login"
        ):
            return JSONResponse(
                {"detail": "isolated browser fixture is read-only"}, status_code=405
            )
        return await call_next(request)

    @app.get("/__browser_checkpoint")
    def checkpoint():
        return snapshot()

    app.add_middleware(
        CORSMiddleware,
        allow_origins=[args.frontend_origin],
        allow_methods=["GET", "POST", "OPTIONS"],
        allow_headers=["*"],
    )
    print(f"Ready: {output / 'manifest.json'}", flush=True)
    try:
        uvicorn.run(app, host="127.0.0.1", port=args.port, lifespan="off", access_log=False)
    finally:
        credentials.unlink(missing_ok=True)
        ontology_engine.close()


if __name__ == "__main__":
    main()
